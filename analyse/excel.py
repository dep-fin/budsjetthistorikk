"""Eksporter ett lagret uttrekk. Kildetekst skrives alltid som tekst, aldri formler."""

import json
import os
from datetime import date
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.properties import PageSetupProperties

from innlesning import read_report, read_revision, read_revision_for_run
from innlesning.model import Report
from innlesning.sync import RevisionReport

HEADERS = [
    "Kapittel",
    "Kapittelnavn",
    "Post",
    "Formål",
    "Budsjettår",
    "Beløp (kroner)",
    "Beløpets betydning",
    "Retning",
    "Kontrollstatus",
    "Kontrollbehov",
    "Departement",
    "Sesjon",
    "Vedtaksdato",
    "Vedtaks-ID",
    "Vedtaksnummer",
    "Sak-ID",
    "Kilde",
    "Kilderad",
    "Kildeversjon",
    "Råbeløp",
    "Endringsbeløp",
    "Førbeløp",
    "Etterbeløp",
]


def put(cell: Cell | MergedCell, value: str | int | date | None) -> None:
    assert isinstance(cell, Cell)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
    cell.font = Font(name="Arial", size=10, color="000000")
    cell.alignment = Alignment(vertical="center")


def export_excel(database: Path, run_id: str, output: Path) -> None:
    """Les data gjennom kontrakten og skriv XLSX atomisk ved vellykket eksport."""
    if database.resolve() == output.resolve():
        raise ValueError("Excel-filen må ha en annen bane enn databasen")
    report = read_report(database, run_id)
    _export_report(report, output, revision=read_revision_for_run(database, run_id))


def export_revision_excel(database: Path, revision_id: str, output: Path) -> None:
    """Eksporter nøyaktig valgt innholdsrevisjon med separat ferskhetskontekst."""
    if database.resolve() == output.resolve():
        raise ValueError("Excel-filen må ha en annen bane enn databasen")
    revision = read_revision(database, revision_id)
    _export_report(revision.report, output, revision=revision)


def _export_report(report: Report, output: Path, *, revision: RevisionReport | None = None) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Bevilgninger"
    put(sheet.cell(2, 1), "Ett endelig bevilgningsvedtak")
    sheet.cell(2, 1).font = Font(name="Arial", size=14, bold=True, color="181C62")
    put(
        sheet.cell(3, 1),
        f"Strukturering: {report.status}. Bare valgt vedtak; ingen sesjonsgaranti."
        + (f" {revision.freshness}; {revision.observation_status}." if revision else ""),
    )
    put(sheet.cell(4, 1), "Strukturerte utgifter (kroner)")
    put(sheet.cell(4, 6), report.approved_totals["utgift"])
    put(sheet.cell(5, 1), "Strukturerte inntekter (kroner)")
    put(sheet.cell(5, 6), report.approved_totals["inntekt"])
    if report.status != "strukturert":
        put(sheet.cell(4, 7), "Ufullstendig – ingen godkjent sum")
        put(sheet.cell(5, 7), "Ufullstendig – ingen godkjent sum")
    decision = report.decision
    for column, heading in enumerate(HEADERS, start=1):
        cell = sheet.cell(7, column)
        put(cell, heading)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="181C62")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    if decision:
        for row_index, line in enumerate(report.lines, start=8):
            values: list[str | int | date | None] = [
                line.chapter,
                line.chapter_name,
                int(line.post),
                line.purpose,
                decision.budget_year,
                line.amount_kr,
                line.meaning,
                line.direction,
                line.status,
                line.control_reason,
                "uavklart",
                decision.session,
                date.fromisoformat(decision.date),
                decision.id,
                decision.number,
                decision.case_id,
                "Åpne originalvedtak",
                line.locator,
                report.source_sha256,
                line.amount_text,
                None,
                None,
                None,
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_index, column)
                put(cell, value)
            sheet.cell(row_index, 3).number_format = "00"
            sheet.cell(row_index, 6).number_format = "# ##0;[Red]-# ##0;0"
            sheet.cell(row_index, 13).number_format = "dd.mm.yyyy"
            cell = sheet.cell(row_index, HEADERS.index("Kilde") + 1)
            expected_url = (
                "https://www.stortinget.no/no/Saker-og-publikasjoner/Vedtak/Vedtak/Sak/"
                f"?did={decision.id}&p={decision.case_id}"
            )
            if decision.url == expected_url:
                cell.hyperlink = Hyperlink(ref=cell.coordinate, target=expected_url)
                cell.font = Font(name="Arial", size=10, color="4156A6", underline="single")
            else:
                put(cell, "Kildelenke krever kontroll")
            for column in (2, 4, 10):
                sheet.cell(row_index, column).alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[row_index].height = 48
            if line.status != "strukturert":
                sheet.cell(row_index, 9).fill = PatternFill("solid", fgColor="F15D61")
    sheet.freeze_panes = "D8"
    sheet.auto_filter.ref = f"A7:{get_column_letter(len(HEADERS))}{max(7, sheet.max_row)}"
    widths = [
        12,
        46,
        8,
        44,
        12,
        22,
        25,
        12,
        18,
        44,
        15,
        16,
        16,
        16,
        16,
        14,
        24,
        23,
        68,
        22,
        19,
        19,
        19,
    ]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in (4, 5):
        sheet.cell(row, 6).number_format = "# ##0;[Red]-# ##0;0"
    notes = workbook.create_sheet("Kilde og kontroll")
    rows: list[tuple[str, str | int | None]] = [
        ("Kjøring", report.run_id),
        ("Datakontrakt", report.contract_version),
        ("Struktureringsstatus", report.status),
        ("Omfang", "Bare valgt vedtak; verken komplett sesjon, budsjettår eller historikk"),
        (
            "Prøve",
            "Eksplisitt historisk prøve" if report.historical_sample else "Ordinær innlesning",
        ),
        ("Produksjonsgrense", "Samlet leveranse starter i sesjon 2026–2027"),
        ("Register", "Departementstilknytning uavklart; registervedlikehold kommer i B5"),
        ("Kildeslag", str(report.source_metadata.get("fixture_kind", "http_original"))),
        ("Kilde-URL", str(report.source_metadata["url"])),
        ("Hentet UTC", str(report.source_metadata["retrieved_utc"])),
        ("Faktisk innlest SHA-256", report.source_sha256),
        ("Faktisk innlest byteantall", int(str(report.source_metadata["bytes"]))),
        ("Beløpsenhet", "Kroner; multiplikator 1"),
        ("Beløpsbetydning", "Opprinnelig årsbevilgning, nivå; endring/før/etter er ikke relevant"),
        ("Utgifter i kilden (kroner)", report.source_totals.get("utgift")),
        ("Inntekter i kilden (kroner)", report.source_totals.get("inntekt")),
    ]
    if revision is not None:
        rows.extend(
            [
                ("Synkroniseringskontrakt", "2"),
                ("Vedtaksrevisjon", revision.revision_id),
                ("Innholdsfingeravtrykk", revision.fingerprint),
                ("Fingeravtrykksmetode", revision.fingerprint_version),
                ("Struktureringsversjon", revision.parser_version),
                ("Datasett", revision.dataset),
                ("Versjonsstatus", revision.freshness),
                ("Vurdert UTC", revision.assessed_utc),
                ("Nåværende observert revisjon", revision.current_revision_id),
                ("Nyere/siste observasjon", revision.current_observation_id),
                ("Siste observerte råhash", revision.current_source_sha256),
                ("Siste forsøk", revision.latest_attempt_id),
                ("Siste forsøksstatus", revision.latest_attempt_status),
                ("Siste forsøks råhash", revision.latest_attempt_source_sha256),
                ("Observasjonsstatus", revision.observation_status),
                ("Kontrollgrense", "Henvisninger endrer eller godkjenner ingen analysetall"),
            ]
        )
        rows.extend(("Feil i siste forsøk", error) for error in revision.latest_attempt_errors)
        control_sheet = workbook.create_sheet("Kontrollhenvisninger")
        control_rows = [
            [
                "Kontroll-ID",
                "Grunnlagsrevisjon",
                "Observasjon",
                "Kilderad",
                "Valg/referanse",
                "Begrunnelse",
                "Registrert UTC",
                "Status",
            ]
        ]
        control_rows.extend(
            [
                [
                    choice.control_id,
                    choice.revision_id,
                    choice.observation_id,
                    choice.locator or "",
                    choice.reference,
                    choice.reason,
                    choice.registered_utc,
                    choice.status,
                ]
                for choice in revision.controls
            ]
        )
        for row_index, control_values in enumerate(control_rows, 1):
            for column, value in enumerate(control_values, 1):
                put(control_sheet.cell(row_index, column), value)
                control_sheet.column_dimensions[get_column_letter(column)].width = 40
        control_sheet.freeze_panes = "A2"
    else:
        rows.append(("Versjonsgrunnlag", "B1-prøvearkiv; ingen valgt B2-revisjonsregistrering"))
    rows.extend(("Kontrollbehov", reason) for reason in report.controls)
    metadata_text = json.dumps(report.source_metadata, ensure_ascii=False, indent=2)
    for start in range(0, len(metadata_text), 30000):
        rows.append(("Bevart hentemetadata", metadata_text[start : start + 30000]))
    if decision:
        rows.extend([("Vedtaksdato rå", decision.date_raw), ("Kildelenke rå", decision.url_raw)])
        for start in range(0, len(decision.html), 30000):
            rows.append(("Vedtaksfulltekst (HTML)", decision.html[start : start + 30000]))
    for row_index, (label, value) in enumerate(rows, start=1):
        put(notes.cell(row_index, 1), label)
        put(notes.cell(row_index, 2), value)
        notes.cell(row_index, 2).alignment = Alignment(wrap_text=True, vertical="top")
        notes.cell(row_index, 1).font = Font(name="Arial", size=10, bold=True, color="181C62")
        notes.row_dimensions[row_index].height = 45
    notes.column_dimensions["A"].width = 33
    notes.column_dimensions["B"].width = 110
    for page in workbook:
        page.sheet_view.showGridLines = False
        page.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        page.page_setup.orientation = "landscape"
        page.page_setup.paperSize = page.PAPERSIZE_A3
        page.page_setup.fitToWidth = 1
        page.page_setup.fitToHeight = 0
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        os.replace(temporary, output)
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
