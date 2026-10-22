"""Selvstendig kontrolluttrekk, også uten beløp eller bevilgningskjede."""

import os
import re
from dataclasses import asdict, dataclass
from pathlib import Path
from tempfile import NamedTemporaryFile

from lxml import etree, html
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from innlesning.ingestion import text
from innlesning.model import Decision
from innlesning.observations import Extraction, encode, read_extraction
from innlesning.sync import read_revision

from .excel import put

type Value = str | int | None


def write_rows(workbook: Workbook, title: str, rows: list[list[Value]]) -> None:
    """Skriv alle tegn; lange felt fortsetter med entydig cellepeker."""
    sheet = workbook.create_sheet(title)
    continuations: list[list[Value]] = []
    for row_number, values in enumerate(rows, 1):
        for column, value in enumerate(values, 1):
            if isinstance(value, str) and len(value) > 30000:
                continuations.extend(
                    [
                        f"{title}!{get_column_letter(column)}{row_number}",
                        start // 30000 + 1,
                        value[start : start + 30000],
                    ]
                    for start in range(30000, len(value), 30000)
                )
                value = value[:30000]
            put(sheet.cell(row_number, column), value)
            if row_number == 1:
                sheet.cell(row_number, column).font = Font(name="Arial", bold=True, color="FFFFFF")
                sheet.cell(row_number, column).fill = PatternFill("solid", fgColor="181C62")
            sheet.column_dimensions[get_column_letter(column)].width = 28
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    if continuations:
        continuation = workbook.create_sheet(f"{title[:18]} fortsettelse")
        continuation_rows: list[list[Value]] = [["Kilde-celle", "Del", "Tekst"], *continuations]
        for row_number, values in enumerate(continuation_rows, 1):
            for column, value in enumerate(values, 1):
                put(continuation.cell(row_number, column), value)


@dataclass(frozen=True)
class SourceTextSnapshot:
    extractions: tuple[Extraction, ...]
    decisions: tuple[tuple[str, Decision | None], ...]


def add_source_text(
    workbook: Workbook,
    database: Path,
    extraction_ids: tuple[str, ...],
    *,
    snapshot: SourceTextSnapshot | None = None,
) -> None:
    rows: list[list[Value]] = [
        ["Beløpsuttrekk", "Revisjon", "Artefakt SHA-256", "Del", "Vedtaksfulltekst (HTML)"]
    ]
    information: list[list[Value]] = [
        [
            "Beløpsuttrekk",
            "Revisjon",
            "Lokator",
            "Opplysningstype",
            "Kildetekst",
            "Kilde-HTML",
            "Kontrollgrense",
        ]
    ]
    for extraction_id in extraction_ids:
        extraction = (
            next(e for e in snapshot.extractions if e.id == extraction_id)
            if snapshot is not None
            else read_extraction(database, extraction_id)
        )
        if extraction.source_system != "storting":
            information.append(
                [
                    extraction_id,
                    extraction.source_revision_id,
                    None,
                    "uavklart",
                    None,
                    None,
                    "Ingen B2-fulltekst for denne kildetypen",
                ]
            )
            continue
        decision = (
            dict(snapshot.decisions)[extraction.source_revision_id]
            if snapshot is not None
            else read_revision(database, extraction.source_revision_id).report.decision
        )
        if decision is None:
            continue
        rows.extend(
            [
                extraction_id,
                extraction.source_revision_id,
                extraction.artifact.sha256,
                start // 30000 + 1,
                decision.html[start : start + 30000],
            ]
            for start in range(0, len(decision.html), 30000)
        )
        parser = html.HTMLParser(no_network=True, recover=False)
        try:
            document = html.fromstring(decision.html, parser=parser)
        except etree.LxmlError:
            continue
        if any(e.level_name in {"ERROR", "FATAL"} for e in parser.error_log):
            continue
        # Tekstopplysningene er separate fra beløpsobservasjonene. Ingen tall normaliseres her.
        for node in document.xpath(
            ".//p[not(ancestor::table)] | .//table[not(.//td[normalize-space()='Kap.'])]"
        ):
            content = text(node)
            if not content:
                continue
            if re.search(r"fullmakt|forpliktelser|forpliktingar", content, re.I):
                kind = "fullmaktsramme"
            elif re.search(
                r"ber (?:regjeringen|regjeringa|Stortingets)|fremme forslag|fremmes", content, re.I
            ):
                kind = "anmodning"
            elif re.search(r"sats|avgift per|kroner per", content, re.I):
                kind = "sats"
            elif re.search(r"overførte midler|overførte løyvingar", content, re.I):
                kind = "overforte_midler"
            else:
                kind = "uavklart_tekst"
            information.append(
                [
                    extraction_id,
                    extraction.source_revision_id,
                    node.getroottree().getpath(node),
                    kind,
                    content,
                    etree.tostring(node, encoding="unicode", method="html"),
                    "Tekstklassifikasjon v1 til kontroll; aldri additiv bevilgning",
                ]
            )
    write_rows(workbook, "Vedtaksfulltekst", rows)
    write_rows(workbook, "Tekstopplysninger", information)


def export_extraction_excel(database: Path, extraction_id: str, output: Path) -> None:
    if database.resolve() == output.resolve():
        raise ValueError("Excel-filen må ha en annen bane enn databasen")
    extraction = read_extraction(database, extraction_id)
    workbook = Workbook()
    assert workbook.active is not None
    workbook.remove(workbook.active)
    observations: list[list[Value]] = [
        [
            "Observasjon-ID",
            "Ledd-ID",
            "Kapittel",
            "Post",
            "Underpost rå",
            "Underpost",
            "Detaljnivå",
            "Komponent av",
            "Side",
            "Budsjettår",
            "Formål og stikkord",
            "Beløpsbetydning",
            "Kroner",
            "Råbeløp",
            "Kildelokator",
            "Beløpsfelt",
            "Tolkning",
            "Revisjon",
            "Opplysningstype",
            "Strukturstatus",
            "Alle felt (JSON)",
        ]
    ]
    hierarchy: list[list[Value]] = [
        ["Komponent-ID", "Observert posttotal-ID", "Relasjoner med kildebevis"]
    ]
    organizations: list[list[Value]] = [
        [
            "Observasjon-ID",
            "Kapittel",
            "Budsjettår",
            "fdep_nr",
            "gdep_nr",
            "Gul bok-navn",
            "Administrativt departement",
            "Registerversjon/hash",
            "Gyldig fra",
            "Gyldig til",
            "Registerkilde",
            "Alle organisasjonsfelt",
        ]
    ]
    anomalies: list[list[Value]] = [["Grunnlag", "Avvikskode", "Begrunnelse", "Alle felt"]]
    for o in extraction.observations:
        observations.append(
            [
                o.id,
                o.step_id,
                o.chapter,
                o.post,
                o.subpost_raw,
                o.subpost,
                o.detail_level,
                o.component_of,
                o.budget_side,
                o.budget_year,
                o.purpose,
                o.meaning,
                o.amount_kr,
                o.raw_value,
                o.pointer.locator,
                o.pointer.amount_field,
                o.interpretation_id,
                o.source_revision_id,
                o.information_type,
                o.structure_status,
                encode(asdict(o)),
            ]
        )
        if o.component_of or o.relations:
            hierarchy.append([o.id, o.component_of, encode(o.relations)])
        org = o.organization
        organizations.append(
            [
                o.id,
                o.chapter,
                o.budget_year,
                org.fdep_nr,
                org.gdep_nr,
                org.gdep_name,
                org.mapped_department_name,
                org.register_version,
                org.valid_from,
                org.valid_to,
                org.register_source,
                encode(asdict(org)),
            ]
        )
        anomalies.extend([o.id, a.code, a.reason, encode(asdict(a))] for a in o.anomalies)
    anomalies.extend(
        [extraction.id, a.code, a.reason, encode(asdict(a))] for a in extraction.anomalies
    )
    for title, rows in (
        ("Beløpsobservasjoner", observations),
        ("Hierarki", hierarchy),
        ("Organisasjonsgrunnlag", organizations),
        ("Avvik", anomalies),
    ):
        write_rows(workbook, title, rows)
    write_rows(
        workbook,
        "Grunnlag",
        [
            ["Felt", "Verdi"],
            ["Beløpsuttrekk", extraction.id],
            ["Tolkning", encode(asdict(extraction.interpretation))],
            ["Artefakt", encode(asdict(extraction.artifact))],
            ["Antall beløpsobservasjoner", len(extraction.observations)],
            [
                "Dekningsgrense",
                "Bare valgt kilde/tolkning. "
                "Ingen bevilgningssum, faglig godkjenning eller årsdekning.",
            ],
        ],
    )
    add_source_text(workbook, database, (extraction.id,))
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        os.replace(temporary, output)
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
