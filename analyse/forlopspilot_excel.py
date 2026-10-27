"""T3s nye Excelprofil: fem brukerflater, tilgjengelig kontroll, én offlinekapsel."""

import hashlib
from copy import copy
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.hyperlink import Hyperlink

from innlesning.forlopspilot import PARENT, RECIPES, _parent, derived_xml
from innlesning.forlopspilot_model import Report
from innlesning.komitepilot import payload_bytes
from innlesning.observations import encode
from innlesning.omtale import compact

from .extraction_excel import Value, write_rows
from .forlopspilot import (
    contents,
    export_json,
    report_from_json,
    report_to_json,
)
from .komitepilot import legacy_payloads
from .omtale_excel_guard import SheetShape, inspect_excel
from .saksdokumenter import _publish

SHEET_NAMES = frozenset(
    {
        "Forløp",
        "Dokumenter",
        "Omtale",
        "Sammenligninger",
        "Dekning og avvik",
        "Kontroll",
        "Offlinegrunnlag",
    }
)


def _chunks(text: str) -> list[str]:
    return [text[i : i + 30000] for i in range(0, len(text), 30000)] or [""]


def _tables(report: Report, frozen: str) -> dict[str, list[list[Value]]]:
    basis = contents(report.basis)
    docs = {d.id: d for d in basis.documents}
    ctx = {c.observation_id: c for c in basis.contexts}
    forlop: list[list[Value]] = [
        [
            "Konto",
            "Budsjettår",
            "Dokument",
            "Dato",
            "Runde / stadium",
            "Avsender",
            "Nivå kr",
            "Før kr",
            "Endring kr",
            "Etter kr",
            "Status",
            "Felt-ID-er",
            "Ledd-ID",
            "Dokumentversjon",
        ]
    ]
    groups: dict[str, list[object]] = {}
    for o in basis.observations:
        if o.id in report.view.observation_ids:
            groups.setdefault(o.step_id, []).append(o)
    # Tre kildefelt vises samlet per behandlingsledd. ID-ene står i kontrollområdet.
    for step in sorted(
        groups,
        key=lambda sid: min(
            (
                o.decision_date or docs[ctx[o.id].document_id].date,
                {"regjering": 0, "tilraading": 1, "vedtak": 2}.get(ctx[o.id].stage, 3),
                int(o.post or "0"),
            )
            for o in basis.observations
            if o.step_id == sid
        ),
    ):
        observations = [
            o
            for o in basis.observations
            if o.step_id == step and o.id in report.view.observation_ids
        ]
        first = observations[0]
        context = ctx[first.id]
        doc = docs[context.document_id]
        values = {o.meaning: o.amount_kr for o in observations}
        forlop.append(
            [
                f"{first.chapter}/{first.post}",
                first.budget_year,
                doc.title,
                doc.date,
                context.round + " / " + context.stage,
                context.actor,
                values.get("nivaa"),
                values.get("for_belop"),
                values.get("endring"),
                values.get("etter_belop"),
                context.status + "; " + first.professional_status,
                encode(tuple(o.id for o in observations)),
                step,
                doc.id,
            ]
        )
    documents: list[list[Value]] = [
        [
            "Type rad",
            "ID",
            "Dokument",
            "Type/rolle",
            "Dato",
            "Sesjon",
            "Budsjettår",
            "Kapittel",
            "Post",
            "Råkildedekning",
            "Strukturstatus",
            "Faglig status",
            "Beløps-ID-er / kontoer",
            "Kilde-ID-er",
            "Lokator",
            "Kildetekst (kompakt)",
            "Valgt",
        ]
    ]
    for d in basis.documents:
        documents.append(
            [
                "Dokument",
                d.id,
                d.title,
                d.kind,
                d.date,
                d.session,
                None,
                None,
                None,
                d.raw_coverage,
                d.structured_coverage,
                d.professional_status,
                encode(d.accounts),
                encode(d.source_ids),
                None,
                None,
                "ja" if d.id in report.view.document_ids else "nei",
            ]
        )
        for item in basis.inventory:
            if item.document_id == d.id:
                documents.append(
                    [
                        "Post",
                        item.id,
                        d.title,
                        item.role,
                        d.date,
                        d.session,
                        item.budget_year,
                        item.chapter,
                        item.post,
                        "se dokument/kilde",
                        item.structure_status,
                        d.professional_status,
                        encode(item.observation_ids),
                        item.source_id,
                        item.locator,
                        compact(item.text),
                        "ja" if item.id in report.view.inventory_ids else "nei",
                    ]
                )
    statements: list[list[Value]] = [
        [
            "Utsagn-ID",
            "Dokument",
            "Rolle",
            "Avsender",
            "Periode/betydning",
            "Kontoer",
            "Hierarki",
            "Del",
            "Originaltekst (åpne celle for fulltekst)",
            "Originalstruktur",
            "Kilde-ID",
            "Lokator",
            "Status",
            "Valgt",
        ]
    ]
    for s in basis.statements:
        text, markup = _chunks(s.text), _chunks(s.html)
        for n in range(max(len(text), len(markup))):
            statements.append(
                [
                    s.id,
                    docs[s.document_id].title,
                    s.role,
                    s.actor,
                    s.period,
                    encode(s.accounts),
                    encode(s.hierarchy),
                    n + 1,
                    text[n] if n < len(text) else "",
                    markup[n] if n < len(markup) else "",
                    s.source_id,
                    s.locator,
                    s.status,
                    "ja" if s.id in report.view.statement_ids else "nei",
                ]
            )
    observations_by_id = {o.id: o for o in basis.observations}
    comparisons: list[list[Value]] = [
        [
            "Konto",
            "Betydning",
            "Fra dokument",
            "Til dokument",
            "Fra beløp kr",
            "Til beløp kr",
            "Avledet forskjell kr",
            "Status",
            "Grunn",
            "Regel",
            "Venstre observasjons-ID",
            "Høyre observasjons-ID",
            "Sammenlignings-ID",
            "Nullgrunn",
        ]
    ]
    for c in report.view.comparisons:
        left, right = (
            observations_by_id.get(c.left_id or ""),
            observations_by_id.get(c.right_id or ""),
        )
        comparisons.append(
            [
                c.account,
                c.meaning,
                docs[ctx[left.id].document_id].title if left else "manglende stadium",
                docs[ctx[right.id].document_id].title if right else "manglende stadium",
                left.amount_kr if left else None,
                right.amount_kr if right else None,
                c.difference_kr,
                c.status,
                c.reason,
                c.rule,
                c.left_id,
                c.right_id,
                c.id,
                c.null_reason,
            ]
        )
    for calculation in report.calculations:
        comparisons.append(
            [
                calculation.account,
                "beregnet_nivaa",
                "valgt startnivå + vedtatte deltaer",
                "observert siste etterbeløp",
                calculation.amount_kr,
                calculation.observed_kr,
                None
                if calculation.amount_kr is None or calculation.observed_kr is None
                else calculation.amount_kr - calculation.observed_kr,
                calculation.status,
                calculation.reason,
                calculation.rule,
                encode(calculation.input_ids),
                encode(calculation.check_ids),
                calculation.id,
                calculation.reason if calculation.amount_kr is None else None,
            ]
        )
    coverage: list[list[Value]] = [["Globalt forhold før filter", "Verdi/status"]]
    coverage.extend(
        [
            ["Rapport-ID", report.id],
            ["Grunnlags-ID", report.basis.id],
            ["Format", report.version],
            ["Valgt inngang", encode(asdict(report.selection))],
            [
                "Returvalg",
                encode(asdict(report.return_selection)) if report.return_selection else None,
            ],
            ["Valgte tabellfelt", len(report.view.observation_ids)],
            ["Valgte vedtaksfelt", len(report.view.adopted_observation_ids)],
            ["Valgte monetære hendelser", len(report.view.adopted_event_ids)],
            ["Globale tabellfelt", len(basis.observations)],
            [
                "Gamle kopier inne i valgfri uendret T1-kapsel",
                "arvet; ikke normalisert" if report.basis.legacy_json else "ingen T1-kapsel",
            ],
        ]
    )
    coverage.extend([issue, "synlig"] for issue in basis.issues)
    control: list[list[Value]] = [
        [
            "Type",
            "ID",
            "Fra/dokument",
            "Til/payload",
            "Rolle/betydning",
            "Omfang/år",
            "Status",
            "Kilde/hash",
            "Lokator/bevis",
            "Kroner/byte",
            "Metadata",
        ]
    ]
    for source in report.basis.sources:
        control.append(
            [
                "Kilde",
                source.id,
                source.document_key,
                source.payload_id,
                source.media_type,
                None,
                source.status,
                source.expected_sha256,
                source.name,
                None,
                source.metadata_json,
            ]
        )
    for p in report.basis.payloads:
        control.append(
            [
                "Payloadindeks",
                p.id,
                None,
                p.legacy_source_id,
                "legacy-referanse" if p.legacy_source_id else "kanonisk payload",
                None,
                "bevart",
                p.sha256,
                None,
                p.byte_count,
                "Originalbyte finnes én gang i Offlinegrunnlag; --originaler gjenfinner dem",
            ]
        )
    for o in basis.observations:
        control.append(
            [
                "Felt",
                o.id,
                ctx[o.id].document_id,
                o.step_id,
                o.meaning,
                o.budget_year,
                o.structure_status,
                o.pointer.artifact_id,
                o.pointer.locator,
                o.amount_kr,
                encode((o.source_system, o.interpretation_id, o.round, o.proposal_author)),
            ]
        )
    for context in basis.contexts:
        control.append(
            [
                "Kontekst",
                context.id,
                context.document_id,
                context.observation_id,
                context.round + " / " + context.stage,
                context.round,
                context.status,
                None,
                encode(context.evidence_ids),
                None,
                context.actor,
            ]
        )
    for r in basis.relations:
        control.append(
            [
                "Relasjon",
                r.id,
                r.from_id,
                r.to_id,
                r.role,
                encode((r.scope, r.accounts, r.budget_year)),
                r.status,
                None,
                encode(r.evidence_ids),
                None,
                None,
            ]
        )
    for calculation in report.calculations:
        control.append(
            [
                "Beregning",
                calculation.id,
                calculation.account,
                None,
                "beregnet_nivaa",
                calculation.budget_year,
                calculation.status,
                None,
                encode((calculation.input_ids, calculation.check_ids)),
                calculation.amount_kr,
                encode(asdict(calculation)),
            ]
        )
    snapshot: list[list[Value]] = [["Del", "Frosset visningspilot-3 med kanonisk råregister"]]
    snapshot.extend([n, chunk] for n, chunk in enumerate(_chunks(frozen), 1))
    return {
        "Forløp": forlop,
        "Dokumenter": documents,
        "Omtale": statements,
        "Sammenligninger": comparisons,
        "Dekning og avvik": coverage,
        "Kontroll": control,
        "Offlinegrunnlag": snapshot,
    }


def _link(book: Workbook, sheet: str, cell: str, target: str, address: str) -> None:
    book[sheet][cell].hyperlink = Hyperlink(ref=cell, location=f"'{target}'!{address}")
    book[sheet][cell].font = Font(name="Arial", size=11, color="0563C1", underline="single")


def _book(report: Report, raw: str) -> Workbook:
    book = Workbook()
    assert book.active is not None
    book.remove(book.active)
    for name, rows in _tables(report, raw).items():
        write_rows(book, name, rows)
        sheet = book[name]
        sheet.row_dimensions[1].height = 42
        for cell in sheet[1]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for n, row in enumerate(sheet.iter_rows(min_row=2), 2):
            sheet.row_dimensions[n].height = (
                48
                if name in {"Forløp", "Dokumenter", "Sammenligninger", "Dekning og avvik"}
                else 24
            )
            for cell in row:
                cell.alignment = Alignment(vertical="top")
                if (
                    (name == "Forløp" and cell.column in {3, 5, 11})
                    or (name == "Dokumenter" and cell.column in {3, 10, 11, 16})
                    or name == "Dekning og avvik"
                ):
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                if sheet.cell(1, cell.column or 1).value == "Dato" and isinstance(cell.value, str):
                    cell.value = datetime.fromisoformat(cell.value)
                    cell.number_format = 'dd"."mm"."yyyy'
                if type(cell.value) is int:
                    cell.number_format = (
                        "0"
                        if sheet.cell(1, cell.column).value in {"Budsjettår", "Del"}
                        else "#,##0"
                    )
        if name == "Forløp":
            for column, width in {
                "A": 14,
                "B": 14,
                "C": 55,
                "D": 14,
                "E": 28,
                "F": 24,
                "G": 21,
                "H": 21,
                "I": 21,
                "J": 21,
                "K": 42,
                "L": 24,
                "M": 22,
                "N": 22,
            }.items():
                sheet.column_dimensions[column].width = width
        elif name == "Dokumenter":
            for column, width in {
                "A": 14,
                "B": 22,
                "C": 46,
                "D": 26,
                "E": 14,
                "F": 14,
                "G": 14,
                "H": 12,
                "I": 10,
                "J": 35,
                "K": 28,
                "L": 26,
                "M": 28,
                "N": 24,
                "O": 28,
                "P": 90,
                "Q": 10,
            }.items():
                sheet.column_dimensions[column].width = width
        elif name == "Omtale":
            for column, width in {
                "A": 22,
                "B": 40,
                "C": 40,
                "D": 35,
                "E": 45,
                "F": 20,
                "G": 60,
                "H": 8,
                "I": 100,
                "J": 35,
                "K": 22,
                "L": 36,
                "M": 24,
                "N": 10,
            }.items():
                sheet.column_dimensions[column].width = width
        elif name == "Sammenligninger":
            for column, width in {
                "A": 14,
                "B": 22,
                "C": 36,
                "D": 36,
                "E": 21,
                "F": 21,
                "G": 24,
                "H": 16,
                "I": 90,
                "J": 28,
                "K": 24,
                "L": 24,
                "M": 24,
                "N": 60,
            }.items():
                sheet.column_dimensions[column].width = width
        elif name == "Dekning og avvik":
            sheet.column_dimensions["A"].width = 100
            sheet.column_dimensions["B"].width = 100
        else:
            sheet.column_dimensions["A"].width = 22
            sheet.column_dimensions["B"].width = 100 if name == "Offlinegrunnlag" else 24
    doc_rows = {
        str(row[1].value): int(row[0].row or 0)
        for row in book["Dokumenter"].iter_rows(min_row=2)
        if row[0].value == "Dokument"
    }
    text_rows: dict[str, int] = {}
    for row in book["Omtale"].iter_rows(min_row=2):
        text_rows.setdefault(str(row[0].value), int(row[0].row or 0))
    controls: dict[str, int] = {
        str(row[1].value): int(row[0].row or 0) for row in book["Kontroll"].iter_rows(min_row=2)
    }
    data = contents(report.basis)
    docs = {d.id: d for d in data.documents}
    contexts = {c.observation_id: c for c in data.contexts}
    first_account: dict[str, int] = {}
    for row in book["Forløp"].iter_rows(min_row=2):
        document_id = str(row[13].value)
        first_account.setdefault(str(row[0].value), int(row[0].row or 0))
        _link(book, "Forløp", f"C{row[0].row}", "Dokumenter", f"C{doc_rows[document_id]}")
    for row in book["Dokumenter"].iter_rows(min_row=2):
        n = int(row[0].row or 0)
        if row[0].value == "Dokument":
            document_id = str(row[1].value)
            first = next((s for s in data.statements if s.document_id == document_id), None)
            if first:
                _link(book, "Dokumenter", f"C{n}", "Omtale", f"B{text_rows[first.id]}")
        else:
            account = str(row[7].value) + "/" + str(row[8].value)
            if account in first_account:
                _link(book, "Dokumenter", f"H{n}", "Forløp", f"A{first_account[account]}")
    for s in data.statements:
        _link(book, "Omtale", f"B{text_rows[s.id]}", "Dokumenter", f"C{doc_rows[s.document_id]}")
    for row in book["Kontroll"].iter_rows(min_row=2):
        if row[0].value == "Felt":
            document_id = contexts[str(row[1].value)].document_id
            _link(book, "Kontroll", f"C{row[0].row}", "Dokumenter", f"C{doc_rows[document_id]}")
        elif row[0].value == "Relasjon":
            for col in (2, 3):
                target = str(row[col].value)
                if target in docs:
                    _link(
                        book,
                        "Kontroll",
                        f"{'CD'[col - 2]}{row[0].row}",
                        "Dokumenter",
                        f"C{doc_rows[target]}",
                    )
    for row in book["Sammenligninger"].iter_rows(min_row=2):
        if str(row[12].value) in controls:
            _link(
                book,
                "Sammenligninger",
                f"M{row[0].row}",
                "Kontroll",
                f"B{controls[str(row[12].value)]}",
            )
        for col in (10, 11):
            target = str(row[col].value)
            if target in controls:
                _link(
                    book,
                    "Sammenligninger",
                    f"{'KL'[col - 10]}{row[0].row}",
                    "Kontroll",
                    f"B{controls[target]}",
                )
                document_id = contexts[target].document_id
                _link(
                    book,
                    "Sammenligninger",
                    f"{'CD'[col - 10]}{row[0].row}",
                    "Dokumenter",
                    f"C{doc_rows[document_id]}",
                )
    return book


def read_excel(path: Path) -> Report:
    raw, shapes = inspect_excel(
        path,
        sheet_names=SHEET_NAMES,
        snapshot_name="Offlinegrunnlag",
        max_expanded_bytes=128 * 1024 * 1024,
    )
    report = report_from_json(raw)
    expected = _book(report, raw)
    try:
        if shapes != {s.title: SheetShape(s.max_row, s.max_column) for s in expected}:
            raise ValueError("T3-Excel har uventede ark/celleområder")
        saved = load_workbook(path, data_only=False, keep_links=False)
        try:
            if saved.sheetnames != expected.sheetnames or saved.views != expected.views:
                raise ValueError("T3-Excel har endret arkrekkefølge/arbeidsbokvisning")
            for wanted_sheet, actual_sheet in zip(expected, saved, strict=True):
                for dimension in wanted_sheet.column_dimensions.values():
                    dimension.reindex()
                if (
                    actual_sheet.sheet_state != wanted_sheet.sheet_state
                    or actual_sheet.sheet_view != wanted_sheet.sheet_view
                    or actual_sheet.sheet_format != wanted_sheet.sheet_format
                    or actual_sheet.auto_filter != wanted_sheet.auto_filter
                    or actual_sheet.merged_cells != wanted_sheet.merged_cells
                    or set(actual_sheet.row_dimensions) != set(wanted_sheet.row_dimensions)
                    or set(actual_sheet.column_dimensions) != set(wanted_sheet.column_dimensions)
                    or any(
                        dict(actual_sheet.row_dimensions[k]) != dict(d)
                        for k, d in wanted_sheet.row_dimensions.items()
                    )
                    or any(
                        dict(actual_sheet.column_dimensions[k]) != dict(d)
                        for k, d in wanted_sheet.column_dimensions.items()
                    )
                ):
                    raise ValueError("T3-Excel har endret synlighet/oppsett: " + actual_sheet.title)
                for wanted_row, actual_row in zip(wanted_sheet, actual_sheet, strict=True):
                    for wanted, actual in zip(wanted_row, actual_row, strict=True):
                        value = None if wanted.value == "" else wanted.value
                        if (
                            actual.value != value
                            or type(actual.value) is not type(value)
                            or isinstance(value, str)
                            and actual.data_type != "s"
                            or actual.data_type == "f"
                            or actual.number_format != wanted.number_format
                            or copy(actual.font) != copy(wanted.font)
                            or copy(actual.fill) != copy(wanted.fill)
                            or copy(actual.border) != copy(wanted.border)
                            or copy(actual.alignment) != copy(wanted.alignment)
                            or copy(actual.protection) != copy(wanted.protection)
                            or (actual.hyperlink.location if actual.hyperlink else None)
                            != (wanted.hyperlink.location if wanted.hyperlink else None)
                            or (actual.hyperlink.target if actual.hyperlink else None)
                            != (wanted.hyperlink.target if wanted.hyperlink else None)
                        ):
                            raise ValueError(
                                "T3-Excel avviker i celle/stil/lenke: "
                                + actual_sheet.title
                                + "!"
                                + actual.coordinate
                            )
        finally:
            saved.close()
    finally:
        expected.close()
    return report


def export_excel(report: Report, output: Path) -> None:
    if output.exists() or output.suffix.lower() != ".xlsx":
        raise ValueError("Velg nytt T3-Excel-filnavn")
    raw = report_to_json(report)
    book = _book(report, raw)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        if read_excel(temporary) != report:
            raise ValueError("Skrevet T3-Excel avviker")
        _publish(temporary, output)
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def export_pair(report: Report, frozen: Path, excel: Path) -> None:
    if frozen.resolve() == excel.resolve() or frozen.exists() or excel.exists():
        raise ValueError("Velg to forskjellige nye filnavn")
    export_json(report, frozen)
    export_excel(report, excel)
    # Begge skrivere har allerede lukket og gjenlest sin fil mot samme rapport.


def export_originals(report: Report, directory: Path) -> None:
    report_to_json(report)
    directory.mkdir(parents=True, exist_ok=False)
    legacy = legacy_payloads(report.basis.legacy_json)
    originals = {p.sha256: payload_bytes(p, legacy) for p in report.basis.payloads}
    originals.update((hashlib.sha256(raw).hexdigest(), raw) for _, raw in legacy)
    for digest, raw in originals.items():
        (directory / (digest + ".bin")).write_bytes(raw)
    parent = next(
        (s for s in report.basis.sources if s.name == PARENT and s.status == "kildeverifisert"),
        None,
    )
    if parent and parent.payload_id:
        expanded = _parent(originals[parent.payload_id])
        for api_id in RECIPES:
            (directory / ("derived-" + api_id + ".xml")).write_bytes(derived_xml(expanded, api_id))
    (directory / "index.json").write_text(
        encode(
            {
                "sources": [asdict(s) for s in report.basis.sources],
                "payloads": {digest: len(raw) for digest, raw in originals.items()},
                "derived": "avledet_ikke_http",
            }
        ),
        "utf8",
    )
