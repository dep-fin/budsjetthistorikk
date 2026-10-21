"""Frosset lokal pilotpakke. Gjenleser alle filer og Excel mot eksakt SQLite-grunnlag."""

import hashlib
import json
import os
import shutil
from dataclasses import asdict
from datetime import UTC, datetime
from pathlib import Path
from tempfile import mkdtemp
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import Cell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from innlesning.historical_excel import HistoricalSnapshot, _connect, read_historical_workbook
from innlesning.observations import encode

PACKAGE_VERSION = "historisk-pakke-1"
EXCEL_VERSION = "historisk-excel-kontroll-1"
CHUNK = 30000
MAX_SHEET_ROWS = 50000
SHEETS = {
    "observasjoner": "Beløpsobservasjoner",
    "celler": "Celler",
    "kandidater": "Kandidater",
    "kolonneregler": "Kolonnegrunnlag",
    "avvik": "Avvik",
    "sammenligning": "Sammenligning",
    "metadata": "Metadata",
    "grunnlag": "Manifest",
}


def _comparison(database: Path, extraction_id: str, comparison_id: str | None) -> dict[str, Any]:
    if comparison_id is None:
        return {
            "version": "historisk-sammenligning-1",
            "status": "ikke_utfort",
            "rows": [],
            "id": None,
        }
    from .historical_comparison import read_historical_comparison

    result = read_historical_comparison(database, comparison_id)
    if result["extraction_id"] != extraction_id:
        raise ValueError("Sammenligningen tilhører et annet uttrekk")
    return result


def _sections(
    snapshot: HistoricalSnapshot, comparison: dict[str, Any]
) -> dict[str, list[dict[str, Any]]]:
    extraction = snapshot.extraction
    anomalies = [
        {"observation_id": o.id, **asdict(a)} for o in extraction.observations for a in o.anomalies
    ] + [{"observation_id": None, **asdict(a)} for a in extraction.anomalies]
    return {
        "observasjoner": [{"id": o.id, **asdict(o)} for o in extraction.observations],
        "celler": list(snapshot.cells),
        "kandidater": list(snapshot.candidates),
        "kolonneregler": list(snapshot.column_rules),
        "avvik": anomalies,
        "sammenligning": comparison.get("rows", []),
        "metadata": [{"key": key, "value": value} for key, value in snapshot.metadata.items()],
        "grunnlag": [
            {
                "package_version": PACKAGE_VERSION,
                "excel_version": EXCEL_VERSION,
                "extraction_id": extraction.id,
                "interpretation_id": extraction.interpretation.id,
                "source_sha256": extraction.artifact.sha256,
                "source_bytes": extraction.artifact.byte_count,
                "artifact": asdict(extraction.artifact),
                "predecessor_id": snapshot.predecessor_id,
                "import_id": extraction.import_id,
                "imported_utc": extraction.imported_utc,
                "comparison_id": comparison["id"],
                "comparison": comparison,
                "professional_approval": False,
                "year_complete": False,
                "scope": "lokal 2024-pilot; sammenligning godkjenner ingen verdi eller kjede",
                "register_snapshot": None,
                "register_null_reason": "historisk register ikke avklart",
            }
        ],
    }


def _value(value: Any) -> str | int | bool | None:
    if value is None or isinstance(value, str | bool | int):
        return value
    return encode(value)


def _tables(
    sections: dict[str, list[dict[str, Any]]],
) -> dict[str, list[list[str | int | bool | None]]]:
    tables: dict[str, list[list[str | int | bool | None]]] = {}
    for name, records in sections.items():
        headers = list(records[0]) if records else ["status"]
        # Hele typede objektet bevarer bl.a. tom streng/null og nested felt.
        headers.append("typed_json")
        for start in range(0, max(len(records), 1), MAX_SHEET_ROWS):
            suffix = "" if start == 0 else f" {start // MAX_SHEET_ROWS + 1}"
            title = SHEETS[name] + suffix
            tables[title] = [[_value(header) for header in headers]] + [
                [_value(record.get(key)) for key in headers[:-1]] + [encode(record)]
                for record in records[start : start + MAX_SHEET_ROWS]
            ]
    return tables


def _excel_expected(
    sections: dict[str, list[dict[str, Any]]],
) -> dict[str, list[list[str | int | bool | None]]]:
    tables = _tables(sections)
    continuation: list[list[str | int | bool | None]] = [["Ark", "Rad", "Kolonne", "Del", "Tekst"]]
    for title, rows in tables.items():
        for row_num, row in enumerate(rows, 1):
            for col_num, value in enumerate(row, 1):
                if isinstance(value, str) and len(value) > CHUNK:
                    for start in range(CHUNK, len(value), CHUNK):
                        continuation.append(
                            [
                                title,
                                row_num,
                                col_num,
                                start // CHUNK + 1,
                                value[start : start + CHUNK],
                            ]
                        )
                    row[col_num - 1] = value[:CHUNK]
                # Excel serialise tom streng som tom celle. typed_json skiller den fra null.
                if row[col_num - 1] == "":
                    row[col_num - 1] = None
    for start in range(0, max(len(continuation) - 1, 1), MAX_SHEET_ROWS):
        title = "Fortsettelser" + ("" if start == 0 else f" {start // MAX_SHEET_ROWS + 1}")
        tables[title] = [continuation[0]] + continuation[start + 1 : start + MAX_SHEET_ROWS + 1]
    return tables


def _write_excel(path: Path, sections: dict[str, list[dict[str, Any]]]) -> None:
    workbook = Workbook(write_only=True)
    for title, rows in _excel_expected(sections).items():
        sheet = workbook.create_sheet(title)
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for col in range(1, len(rows[0]) + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 28
        for row_num, row in enumerate(rows, 1):
            output: list[Cell] = []
            for value in row:
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"  # Ingen formelinjeksjon fra kilde- eller kontrolltekst.
                if row_num == 1:
                    cell.font = Font(name="Arial", color="FFFFFF", bold=True)
                    cell.fill = PatternFill("solid", fgColor="181C62")
                else:
                    cell.font = Font(name="Arial", size=10)
                output.append(cell)
            sheet.append(output)
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(rows[0]))}{len(rows)}"
    workbook.save(path)
    workbook.close()


def _verify_excel(path: Path, sections: dict[str, list[dict[str, Any]]]) -> None:
    tables = _excel_expected(sections)
    with ZipFile(path) as archive:
        infos = archive.infolist()
        if (
            len(infos) > 256
            or len({item.filename for item in infos}) != len(infos)
            or sum(item.file_size for item in infos) > 1024 * 1024 * 1024
            or any(item.file_size > 256 * 1024 * 1024 for item in infos)
        ):
            raise ValueError("Excel-kontrollfil overskrider ressursgrensene")
        inflated = 0
        for info in infos:
            tail = b""
            actual_size = 0
            with archive.open(info) as stream:
                while chunk := stream.read(65536):
                    actual_size += len(chunk)
                    inflated += len(chunk)
                    if actual_size > 256 * 1024 * 1024 or inflated > 1024 * 1024 * 1024:
                        raise ValueError("Faktisk Excel-innhold overskrider ressursgrensene")
                    if info.filename.endswith((".xml", ".rels")):
                        scan = (tail + chunk).replace(b"\x00", b"").upper()
                        if b"<!DOCTYPE" in scan or b"<!ENTITY" in scan:
                            raise ValueError("DTD/entiteter i kontroll-Excel avvises")
                        tail = chunk[-32:]
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != list(tables):
            raise ValueError("Excel-ark stemmer ikke med lagret grunnlag")
        for title, expected in tables.items():
            sheet = workbook[title]
            count = 0
            for count, actual in enumerate(sheet.iter_rows(), 1):
                if count > len(expected) or len(actual) != len(expected[count - 1]):
                    raise ValueError(f"Excel-dimensjoner avviker: {title}")
                for cell, value in zip(actual, expected[count - 1], strict=True):
                    if (
                        cell.data_type == "f"
                        or cell.value != value
                        or (value is not None and type(cell.value) is not type(value))
                    ):
                        raise ValueError(f"Excel-celle avviker: {title}!{cell.coordinate}")
            if count != len(expected):
                raise ValueError(f"Excel-rader mangler: {title}")
    finally:
        workbook.close()


def _file(path: Path, count: int | None = None) -> dict[str, Any]:
    raw = path.read_bytes()
    return {"sha256": hashlib.sha256(raw).hexdigest(), "bytes": len(raw), "count": count}


def export_historical_package(
    database: Path,
    extraction_id: str | None,
    output: Path,
    *,
    comparison_id: str | None = None,
) -> Path:
    snapshot = read_historical_workbook(database, extraction_id)
    extraction_id = snapshot.extraction.id
    comparison = _comparison(database, extraction_id, comparison_id)
    sections = _sections(snapshot, comparison)
    destination = output.resolve()
    if destination == database.resolve() or destination in database.resolve().parents:
        raise ValueError("Pakken kan ikke overskrive database eller dens mappe")
    if destination.exists():
        verify_historical_package(database, extraction_id, destination, comparison_id=comparison_id)
    else:
        destination.parent.mkdir(parents=True, exist_ok=True)
        staging = Path(mkdtemp(prefix=f".{destination.name}-", dir=destination.parent))
        try:
            for name, records in sections.items():
                with (staging / f"{name}.jsonl").open(
                    "w", encoding="utf-8", newline="\n"
                ) as handle:
                    for record in records:
                        handle.write(encode(record) + "\n")
            _write_excel(staging / "kontroll.xlsx", sections)
            manifest: dict[str, Any] = {
                "version": PACKAGE_VERSION,
                "extraction_id": extraction_id,
                "comparison_id": comparison_id,
                "basis": sections["grunnlag"][0],
                "files": {
                    f"{name}.jsonl": _file(staging / f"{name}.jsonl", len(records))
                    for name, records in sections.items()
                },
            }
            manifest["files"]["kontroll.xlsx"] = _file(staging / "kontroll.xlsx")
            (staging / "manifest.json").write_text(encode(manifest) + "\n", encoding="utf-8")
            verify_historical_package(database, extraction_id, staging, comparison_id=comparison_id)
            # Ren rename av ny katalog: ingen halv pakke og ingen erstatning av gamle filer.
            if destination.exists():
                raise ValueError("Pakkemålet ble opprettet samtidig")
            os.rename(staging, destination)
        finally:
            if staging.exists():
                shutil.rmtree(staging)
    manifest_hash = _file(destination / "manifest.json")["sha256"]
    with _connect(database, writable=True) as connection:
        previous = connection.execute(
            "SELECT manifest_sha256 FROM historical_publications WHERE extraction_id=? AND path=?",
            (extraction_id, str(destination)),
        ).fetchone()
        if previous and previous[0] != manifest_hash:
            raise ValueError("Publisert pakke har endret manifest")
        connection.execute(
            "INSERT OR IGNORE INTO historical_publications "
            "(extraction_id,path,manifest_sha256,published_utc) VALUES (?,?,?,?)",
            (extraction_id, str(destination), manifest_hash, datetime.now(UTC).isoformat()),
        )
    return destination


def verify_historical_package(
    database: Path,
    extraction_id: str | None,
    package: Path,
    *,
    comparison_id: str | None = None,
) -> dict[str, Any]:
    snapshot = read_historical_workbook(database, extraction_id)
    extraction_id = snapshot.extraction.id
    comparison = _comparison(database, extraction_id, comparison_id)
    sections = _sections(snapshot, comparison)
    if package.is_symlink() or not package.is_dir():
        raise ValueError("Pakken må være en ordinær mappe")
    expected_names = {"manifest.json", "kontroll.xlsx"} | {f"{name}.jsonl" for name in sections}
    if {path.name for path in package.iterdir()} != expected_names or any(
        p.is_symlink() or not p.is_file() for p in package.iterdir()
    ):
        raise ValueError("Uventede eller manglende pakkefiler")
    manifest_path = package / "manifest.json"
    if manifest_path.stat().st_size > 32 * 1024 * 1024:
        raise ValueError("For stort manifest")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if not isinstance(manifest, dict) or set(manifest) != {
        "version",
        "extraction_id",
        "comparison_id",
        "basis",
        "files",
    }:
        raise ValueError("Manifestets innhold/serialisering er endret")
    if (
        manifest.get("version") != PACKAGE_VERSION
        or manifest.get("extraction_id") != extraction_id
        or manifest.get("comparison_id") != comparison_id
        or manifest.get("basis") != json.loads(encode(sections["grunnlag"][0]))
    ):
        raise ValueError("Manifest har feil versjon, ID eller grunnlag")
    if not isinstance(manifest.get("files"), dict) or set(manifest["files"]) != expected_names - {
        "manifest.json"
    }:
        raise ValueError("Manifest har feil filinventar")
    for name, records in sections.items():
        path = package / f"{name}.jsonl"
        expected = "".join(encode(record) + "\n" for record in records).encode()
        if path.stat().st_size != len(expected) or path.read_bytes() != expected:
            raise ValueError(f"JSONL avviker fra SQLite: {name}")
        if manifest["files"][path.name] != _file(path, len(records)):
            raise ValueError("Manifestets hash/byteantall/radantall avviker")
    excel = package / "kontroll.xlsx"
    if excel.stat().st_size > 128 * 1024 * 1024 or manifest["files"][excel.name] != _file(excel):
        raise ValueError("Excel-filens hash/byteantall avviker")
    _verify_excel(excel, sections)
    return {
        "verified": True,
        "extraction_id": extraction_id,
        "files": len(expected_names),
        "counts": {name: len(rows) for name, rows in sections.items()},
    }


def main() -> int:
    import argparse
    import sqlite3

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=["export", "verify"])
    parser.add_argument("--database", type=Path, required=True)
    parser.add_argument("--id", required=True)
    parser.add_argument("--comparison")
    parser.add_argument("--package", type=Path, required=True)
    arguments = parser.parse_args()
    try:
        if arguments.command == "export":
            export_historical_package(
                arguments.database,
                arguments.id,
                arguments.package,
                comparison_id=arguments.comparison,
            )
        result = verify_historical_package(
            arguments.database, arguments.id, arguments.package, comparison_id=arguments.comparison
        )
        print(encode(result))
        return 0
    except (ValueError, OSError, sqlite3.Error, BadZipFile) as error:
        print(encode({"status": "eksportfeil", "error": str(error)}))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
