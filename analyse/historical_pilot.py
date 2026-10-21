"""Reproduser hele lokalpiloten fra hashfestede kilder; ingen nett eller råpublisering."""

import argparse
import csv
import hashlib
import json
import os
from collections import Counter
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

from openpyxl import load_workbook

from innlesning import read_sync, sync_saved
from innlesning.historical_excel import import_historical_workbook, read_historical_workbook
from innlesning.observations import (
    AmountObservation,
    Interpretation,
    extract_revision,
    read_extraction,
)

from .historical_comparison import (
    VERSION,
    compare_historical_controls,
    freeze_historical_controls,
    historical_control_scope,
)
from .historical_package import export_historical_package, verify_historical_package

ORIGINAL_SHA = "870d30d55b4f7f6791fe48ea2a15aa3e98d61946393385d0411e46a0169f60f3"
REFERENCE_HASHES = {
    "api_events.csv": "ab9ef8abe7f1b508d54ca17e6f88209d061a6a6e2c8f312d5444894686642e31",
    "blaabok_posts.json": "7f85002b543a57b3948170f102dc581f66d5ca1b5c1372198433ecd05b0b337e",
    "description_inventory.csv": "3f267dddbd26434febd976c9a5b9940d40362d0fbedb567485857342871be56e",
    "decision_inventory.csv": "6e7a1676d372798d37c5877d2ebcc2b7275354e482833ecac4ed83c54c8b820e",
}
GROUPS = {
    "i437/23-24": "40033662",
    "i447/23-24": "40033719",
    "i448/23-24": "40033707",
    "i25/24-25": "40033805",
    "i57/24-25": "40033929",
    "i80/24-25": "40034216",
    "i81/24-25": "40034048",
    "i82/24-25": "40034211",
    "i83/24-25": "40034044",
    "i84/24-25": "40034212",
    "i85/24-25": "40034045",
    "i87/24-25": "40034094",
    "i88/24-25": "40034100",
    "i90/24-25": "40034287",
    "i91/24-25": "40034098",
    "i92/24-25": "40034093",
    "i93/24-25": "40034226",
    "i94/24-25": "40034229",
    "i95/24-25": "40034220",
    "i96/24-25": "40034090",
    "i98/24-25": "40034217",
    "i99/24-25": "40034064",
    "i100/24-25": "40034134",
    "i102/24-25": "40034164",
    "i103/24-25": "40034092",
    "i105/24-25": "40034162",
}


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _same_path(left: Path, right: Path) -> bool:
    return left.resolve() == right.resolve() or (
        left.exists() and right.exists() and left.samefile(right)
    )


def _preflight(repo: Path, source: Path, output: Path, proposal_refs: Path | None) -> None:
    # Enumerate every read input and every fixed output before mkdir/SQLite writes.
    research = repo / "docs/research"
    inputs = [source, *(research / "issue-2/aar" / name for name in REFERENCE_HASHES)]
    if proposal_refs is not None:
        inputs.append(proposal_refs)
    manifest = research / "historisk-excel/pilotkontroll/kontrollgrunnlag.json"
    inputs.append(manifest)
    for item in json.loads(manifest.read_text(encoding="utf-8"))["sources"]:
        fixture = repo / item["fixture"]
        inputs.extend([fixture, fixture.with_suffix(".json")])
    outputs = [output / name for name in ("summary.json", "pilot.sqlite", "bounded-api.sqlite")]
    for database in outputs[1:]:
        outputs.extend(Path(str(database) + suffix) for suffix in ("-journal", "-wal", "-shm"))
    package = output / "kontrollpakke"
    outputs.extend(
        package / name
        for name in (
            "manifest.json",
            "kontroll.xlsx",
            "observasjoner.jsonl",
            "celler.jsonl",
            "kandidater.jsonl",
            "kolonneregler.jsonl",
            "avvik.jsonl",
            "sammenligning.jsonl",
            "metadata.jsonl",
            "grunnlag.jsonl",
        )
    )
    for target in outputs:
        if any(_same_path(target, path) for path in inputs):
            raise ValueError("Kollisjon mellom pilotutdata og kilde/kontrollgrunnlag")
    if any(path.resolve().is_relative_to(package.resolve()) for path in inputs):
        raise ValueError("Kollisjon: pakkeområdet inneholder et inputgrunnlag")
    if any(
        _same_path(left, right)
        for index, left in enumerate(outputs)
        for right in outputs[index + 1 :]
    ):
        raise ValueError("Kollisjon mellom pilotens utdatafiler")
    summary = output / "summary.json"
    if summary.exists():
        try:
            preserved = json.loads(summary.read_text(encoding="utf-8"))
            if (
                preserved.get("version") != "historisk-pilotkontroll-3"
                or preserved.get("source_sha256") != ORIGINAL_SHA
            ):
                raise ValueError("Ukjent sammendragsgrunnlag")
            if not verify_historical_package(
                output / "pilot.sqlite",
                preserved["extraction_id"],
                package,
                comparison_id=preserved["comparison_id"],
            )["verified"]:
                raise ValueError("Ugyldig sammendragspakke")
        except (ValueError, OSError, KeyError, TypeError) as error:
            raise ValueError("Eksisterende sammendrag er fremmed eller ugyldig; bevares") from error


def _publish_summary(path: Path, summary: dict[str, Any]) -> None:
    raw = (json.dumps(summary, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    if path.exists():
        if path.read_bytes() != raw:
            raise ValueError("Eksisterende sammendrag avviker; bevares")
        return
    with NamedTemporaryFile(dir=path.parent, prefix=".summary-", delete=False) as handle:
        temporary = Path(handle.name)
        handle.write(raw)
        handle.flush()
        os.fsync(handle.fileno())
    try:
        # Windows rename is exclusive. POSIX rename replaces, so use link there.
        if os.name == "nt":
            os.rename(temporary, path)
        else:
            os.link(temporary, path)
    except FileExistsError:
        if path.read_bytes() != raw:
            raise ValueError("Eksisterende sammendrag avviker; bevares") from None
    finally:
        temporary.unlink(missing_ok=True)


def _account(post: str) -> list[str]:
    return [str(int(post[:-2])), str(int(post[-2:]))]


def _control(
    identifier: str,
    column: str,
    references: list[dict[str, Any]],
    sources: list[dict[str, Any]],
    *,
    groups: list[str] | None = None,
    component: bool = False,
    bounded: bool = False,
) -> dict[str, Any]:
    selector: dict[str, Any] = {
        "sheet": "Rådata" if column != "D" else "Dokumentiltknyntning2024",
        "column": column,
        "key_fields": ["chapter", "post"]
        if column != "D"
        else ["group", "chapter", "post", "subpost"],
    }
    if column == "D":
        selector["groups"] = groups or sorted(GROUPS)
        selector["subposts"] = ["1", "2", "3", "4", "5", "6"] if component else ["0"]
    if bounded:
        selector["keys"] = [r["key"] for r in references]
    scope = historical_control_scope(column, component=component)
    # Preserve independently supplied reference scopes; never silently repair them.
    references = [
        {"scope": {**scope, "candidate_multiplier": "1"}, "source_budget_year": 2024, **r}
        for r in references
    ]
    return {
        "id": identifier,
        "tolerance": "0",
        "selector": selector,
        "scope": {
            **scope,
            "aggregation": "separat kontrollsum, ingen hierarkigodkjenning eller ekstra vedtak",
            "representation": "eksplisitt v2"
            if "bounded" in identifier
            else "hashfestet forskningsuttrekk fra kilde, ingen full v2-replay",
        },
        "references": references,
        "sources": sources,
    }


def run_pilot(
    repo: Path, source: Path, output: Path, *, proposal_refs: Path | None = None
) -> dict[str, Any]:
    _preflight(repo, source, output, proposal_refs)
    if source.stat().st_size != 257569 or _hash(source) != ORIGINAL_SHA:
        raise ValueError("Pilotoriginal har feil hash/byteantall")
    output.mkdir(parents=True, exist_ok=True)
    database = output / "pilot.sqlite"
    attempt = import_historical_workbook(
        database,
        source,
        budget_year=2024,
        year_evidence="kilderegister ved da720e9dd4ebfff31d88c14488c5f3217d045786",
    )
    if attempt.extraction_id is None:
        raise ValueError(f"Pilotimport feilet: {attempt.error}")
    extraction_id = attempt.extraction_id
    snapshot = read_historical_workbook(database, extraction_id)
    original_workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        original_names = {
            f"D{number}": (row[9], row[7])
            for number, row in enumerate(
                original_workbook["Dokumentiltknyntning2024"].iter_rows(
                    min_row=2, values_only=True
                ),
                2,
            )
        }
    finally:
        original_workbook.close()
    document_observations = [
        o for o in snapshot.extraction.observations if o.pointer.sheet == "Dokumentiltknyntning2024"
    ]
    assert len(document_observations) == len(original_names) == 1371
    assert original_names["D2"] == ("Statsministerens kontor", "Driftsutgifter")
    assert all(
        (o.chapter_name, o.purpose) == original_names[o.pointer.cell or ""]
        for o in document_observations
    )
    reference_dir = repo / "docs/research/issue-2/aar"
    sources: list[dict[str, Any]] = []
    for name, expected in REFERENCE_HASHES.items():
        path = reference_dir / name
        if _hash(path) != expected:
            raise ValueError(f"Kontrollreferansen har endret innhold: {name}")
        sources.append(
            {
                "path": path.relative_to(repo).as_posix(),
                "sha256": expected,
                "bytes": path.stat().st_size,
                "kind": "offentlig forskningsavledning",
                "parent_manifest": "docs/research/issue-2/aar/input_manifest.json",
            }
        )
    blue = json.loads((reference_dir / "blaabok_posts.json").read_text(encoding="utf-8"))
    h_references = [
        {
            "key": _account(row["post"]),
            "amount_kr": str(row["amount"]),
            "pointers": [
                {
                    "artifact_sha256": REFERENCE_HASHES["blaabok_posts.json"],
                    "record": n,
                    "parent_sha256": (
                        "7204477be2b1ed089ae2d2eca1307a96d5929d04ea5a0d874ca32eeae1a53bf8"
                    ),
                    "url": "https://www.stortinget.no/globalassets/pdf/budsjettsiden/blaabok_2024.pdf",
                    "page": row["page"],
                    "top": str(row["top"]),
                }
            ],
        }
        for n, row in enumerate(blue, 1)
    ]
    with (reference_dir / "api_events.csv").open(encoding="utf-8-sig", newline="") as handle:
        events = list(csv.DictReader(handle, delimiter=";"))
    with (reference_dir / "description_inventory.csv").open(
        encoding="utf-8-sig", newline=""
    ) as handle:
        mapping = {
            row["source_group"]: row["api_ids"] for row in csv.DictReader(handle, delimiter=";")
        }
    if any(mapping[group] != api_id for group, api_id in GROUPS.items()):
        raise ValueError("Uavhengig dokumentgruppekobling er endret")
    d_references = [
        {
            "key": [row["source_group"], *_account(row["post"]), "0"],
            "amount_kr": row["amount"],
            "pointers": [
                {
                    "artifact_sha256": REFERENCE_HASHES["api_events.csv"],
                    "record": n,
                    "api_id": row["id"],
                    "url": row["url"],
                    "source_html_row_zero_based": row["row"],
                    "source_row": row["source_row"],
                    "decision_date": row["date"],
                    "session": row["session"],
                }
            ],
        }
        for n, row in enumerate(events, 1)
        if row["kind"] == "change"
        and row["source_group"] in GROUPS
        and row["level_not_delta"] != "True"
    ]
    controls = [
        _control("full_H_blaabok", "H", h_references, sources),
        _control("full_D_derived_ledger", "D", d_references, sources),
    ]
    # Uendret B2/v2 replay er en separat representasjon, med eget datasett prove.
    bounded_manifest = json.loads(
        (repo / "docs/research/historisk-excel/pilotkontroll/kontrollgrunnlag.json").read_text(
            encoding="utf-8"
        )
    )
    replay_database = output / "bounded-api.sqlite"
    observations: list[AmountObservation] = []
    replay_sources = []
    for item in bounded_manifest["sources"]:
        fixture = repo / item["fixture"]
        if _hash(fixture) != item["sha256"] or fixture.stat().st_size != item["bytes"]:
            raise ValueError("Autentisk fixture har endret innhold")
        metadata = json.loads(fixture.with_suffix(".json").read_text(encoding="utf-8"))
        job = sync_saved(
            replay_database,
            fixture.read_bytes(),
            metadata,
            historical_sample=True,
            decision_id=item["id"],
        )
        revision = read_sync(replay_database, job).entries[0].revision_id
        extraction = extract_revision(
            replay_database,
            revision,
            interpretation=Interpretation(adapter_version="storting-belop-v2"),
        )
        if (
            read_extraction(replay_database, extraction.id) != extraction
            or extract_revision(replay_database, revision, interpretation=extraction.interpretation)
            != extraction
        ):
            raise ValueError("V2-replay er ikke idempotent/gjenlesbar")
        observations.extend(extraction.observations)
        replay_sources.append(
            {
                "sha256": item["sha256"],
                "bytes": item["bytes"],
                "parent_sha256": item["parent_sha256"],
                "url": item["url"],
                "api_id": item["id"],
                "interpretation_id": extraction.interpretation.id,
            }
        )
    bounded_h = [
        {
            "key": [o.chapter, o.post],
            "amount_kr": str(o.amount_kr),
            "source_budget_year": o.budget_year,
            "pointers": [
                {
                    "observation_id": o.id,
                    "pointer": o.pointer.locator,
                    "api_id": o.api_id,
                    "source_sha256": o.pointer.artifact_id,
                }
            ],
        }
        for o in observations
        if (o.meaning, o.detail_level) == ("nivaa", "post")
    ]
    bounded_groups = bounded_manifest["source_group_mapping"]
    bounded_d = [
        {
            "key": [bounded_groups[o.api_id], o.chapter, o.post, "0"],
            "amount_kr": str(o.amount_kr),
            "source_budget_year": o.budget_year,
            "pointers": [
                {
                    "observation_id": o.id,
                    "pointer": o.pointer.locator,
                    "api_id": o.api_id,
                    "source_sha256": o.pointer.artifact_id,
                }
            ],
        }
        for o in observations
        if (o.meaning, o.detail_level) == ("endring", "post")
    ]
    controls.extend(
        [
            _control("bounded_H_v2", "H", bounded_h, replay_sources, bounded=True),
            _control(
                "bounded_D_v2",
                "D",
                bounded_d,
                replay_sources,
                groups=sorted(bounded_groups.values()),
            ),
        ]
    )
    components: list[dict[str, Any]] = []
    post_totals: list[dict[str, Any]] = []
    for before_id, after_id, chapter, group in [
        ("40032976", "40033719", "5440", "i447/23-24"),
        ("40032850", "40034064", "2470", "i99/24-25"),
        ("40033719", "40034226", "5440", "i93/24-25"),
    ]:
        before = {
            o.subpost: o
            for o in observations
            if (o.api_id, o.chapter, o.post) == (before_id, chapter, "24")
        }
        after = {
            o.subpost: o
            for o in observations
            if (o.api_id, o.chapter, o.post) == (after_id, chapter, "24")
        }
        for code, observation in after.items():
            previous = before[code]
            assert observation.amount_kr is not None and previous.amount_kr is not None
            record = {
                "key": [group, chapter, "24", code],
                "amount_kr": str(observation.amount_kr - previous.amount_kr),
                "pointers": [
                    {
                        "observation_id": item.id,
                        "pointer": item.pointer.locator,
                        "api_id": item.api_id,
                        "source_sha256": item.pointer.artifact_id,
                    }
                    for item in (previous, observation)
                ],
                "derived": True,
                "source_budget_year": observation.budget_year
                if observation.budget_year == previous.budget_year
                else None,
            }
            (post_totals if code is None else components).append(record)
    controls.append(
        _control(
            "bounded_components_derived",
            "D",
            components,
            replay_sources,
            groups=["i447/23-24", "i99/24-25", "i93/24-25"],
            component=True,
        )
    )
    # Konkrete nivå-/endringsankere med faktisk betydning fra API-feltet.
    for col, api_id, meaning, cell in [
        ("J", "40033719", "etter_belop", "J1360"),
        ("K", "40033719", "endring", "K1360"),
        ("L", "40033805", "endring", "L1360"),
    ]:
        observation = next(
            o
            for o in observations
            if (o.api_id, o.chapter, o.post, o.meaning) == (api_id, "2541", "70", meaning)
        )
        control = _control(
            f"anchor_{col}",
            col,
            [
                {
                    "key": ["2541", "70"],
                    "amount_kr": str(observation.amount_kr),
                    "source_budget_year": observation.budget_year,
                    "pointers": [
                        {
                            "observation_id": observation.id,
                            "pointer": observation.pointer.locator,
                            "source_sha256": observation.pointer.artifact_id,
                        }
                    ],
                }
            ],
            replay_sources,
        )
        control["selector"]["cells"] = [cell]
        control["scope"]["meaning"] = meaning
        controls.append(control)
    if proposal_refs is not None:
        proposals = json.loads(proposal_refs.read_text(encoding="utf-8"))
        for item in proposals:
            control = _control(
                f"anchor_{item['column']}_proposal",
                item["column"],
                item["references"],
                item["sources"],
            )
            control["selector"]["cells"] = item["cells"]
            control["scope"].update(figure_type="regjeringsforslag", meaning=item["meaning"])
            controls.append(control)
    tested_proposals = {
        c["selector"]["column"]
        for c in controls
        if c.get("scope", {}).get("figure_type") == "regjeringsforslag"
    }
    for col in ["G", "I", "M", "N"]:
        if col not in tested_proposals:
            controls.append(
                {
                    "id": f"open_{col}",
                    "tolerance": "0",
                    "status": "ikke_utfort",
                    "reason": (
                        "Ingen frosset uavhengig kilde for kolonnens faktiske stadium/omfang. "
                        "Vedtaksbeløp brukes ikke som forslag. Faglig avklaring hos eier."
                    ),
                }
            )
    specification = {
        "version": VERSION,
        "controls": controls,
        "group_mapping": GROUPS,
        "post_totals_separate": post_totals,
        "scope": (
            "hele H og alle 26 dokumentgrupper; bounded v2 holdes separat; "
            "øvrige kolonner bare konkrete ankere"
        ),
    }
    spec_id = freeze_historical_controls(database, extraction_id, specification)
    report = compare_historical_controls(database, spec_id)
    summary = report["summary"]
    assert len(observations) == 2894
    assert snapshot.metadata["counts"] == {
        "cells": 46788,
        "physical": 44652,
        "absent": 2136,
        "observations": 16339,
    }
    assert summary["full_H_blaabok"] == {
        "match": 847,
        "fractional_kr": 726,
        "missing_reference": 120,
        "missing_excel": 1,
    }
    assert summary["full_D_derived_ledger"] == {
        "match": 1354,
        "missing_reference": 1,
        "missing_excel": 4,
    }
    assert summary["bounded_H_v2"] == {"match": 271, "fractional_kr": 103, "missing_excel": 1}
    assert summary["bounded_D_v2"] == {"match": 854, "missing_reference": 3, "missing_excel": 4}
    assert summary["bounded_components_derived"] == {"match": 12, "missing_excel": 4}
    package = output / "kontrollpakke"
    export_historical_package(database, extraction_id, package, comparison_id=report["id"])
    verified = verify_historical_package(
        database, extraction_id, package, comparison_id=report["id"]
    )
    excel = load_workbook(package / "kontroll.xlsx", read_only=True, data_only=False)
    excel_names_checked = 0
    try:
        sheet = excel["Beløpsobservasjoner"]
        header = [cell.value for cell in next(sheet.iter_rows())]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            pointer_value = row[header.index("pointer")]
            assert isinstance(pointer_value, str)
            pointer = json.loads(pointer_value)
            if pointer["sheet"] == "Dokumentiltknyntning2024":
                assert (
                    row[header.index("chapter_name")],
                    row[header.index("purpose")],
                ) == original_names[pointer["cell"]]
                excel_names_checked += 1
    finally:
        excel.close()
    assert excel_names_checked == 1371
    public_summary = {
        "version": "historisk-pilotkontroll-3",
        "source_sha256": ORIGINAL_SHA,
        "source_bytes": 257569,
        "extraction_id": extraction_id,
        "interpretation_id": snapshot.extraction.interpretation.id,
        "adapter_version": snapshot.extraction.interpretation.adapter_version,
        "inventory_version": snapshot.metadata["inventory_version"],
        "comparison_version": VERSION,
        "control_date": (snapshot.extraction.imported_utc or "")[:10],
        "limits": snapshot.metadata["limits"],
        "document_name_mapping": {"chapter_name": "J", "purpose": "H", "keyword_raw_context": "I"},
        "original_name_fields_checked": len(original_names),
        "published_excel_name_fields_checked": excel_names_checked,
        "comparison_id": report["id"],
        "spec_id": spec_id,
        "counts": snapshot.metadata["counts"],
        "summary": summary,
        "column_counts": dict(
            Counter(
                f"{o.pointer.sheet}:{o.pointer.column}" for o in snapshot.extraction.observations
            )
        ),
        "package_manifest_sha256": _hash(package / "manifest.json"),
        "package_verified": verified,
        "bounded_observations": len(observations),
        "source_references": sources,
        "blocked_chains": report["blocked_chains"],
        "professional_approval": False,
        "year_complete": False,
        "privacy": (
            "Ingen original, fullstendig celleinventar eller kandidater publiseres; "
            "bare aggregert kontrollbevis."
        ),
    }
    if _hash(source) != ORIGINAL_SHA:
        raise ValueError("Originalbyte har endret seg under kjøringen")
    _publish_summary(output / "summary.json", public_summary)
    return public_summary


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--repo", type=Path, default=Path(__file__).resolve().parent.parent)
    parser.add_argument("--proposal-refs", type=Path)
    arguments = parser.parse_args()
    print(
        json.dumps(
            run_pilot(
                arguments.repo,
                arguments.source,
                arguments.output,
                proposal_refs=arguments.proposal_refs,
            ),
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
