"""Kontrolluttrekk for én lagret beregning, med adskilte kildebeløp."""

import json
import os
from dataclasses import asdict
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from innlesning.observations import read_extraction

from .excel import put
from .history import read_history


def export_history_excel(database: Path, calculation_id: str, output: Path) -> None:
    if database.resolve() == output.resolve():
        raise ValueError("Excel-filen må ha en annen bane enn databasen")
    result = read_history(database, calculation_id)
    workbook = Workbook()
    history = workbook.active
    assert history is not None
    history.title = "Historikk"
    observation_sheet = workbook.create_sheet("Beløpsobservasjoner")
    notes = workbook.create_sheet("Grunnlag og kontroll")
    history_rows: list[list[str | int | None]] = [
        [
            "Ledd-ID",
            "Kildeenhet",
            "Revisjon",
            "Dato",
            "Runde",
            "Opprinnelig nivå (kr)",
            "Vedtatt delta (kr)",
            "Observert før (kr)",
            "Observert etter (kr)",
            "Beregnet saldo (kr)",
            "Kjedestatus",
            "Avvik",
            "Inn-observasjoner",
        ]
    ]
    for step in result.steps:
        history_rows.append(
            [
                step.step_id,
                step.source_unit_id,
                step.source_revision_id,
                step.decision_date,
                step.round,
                step.original,
                step.delta,
                step.before,
                step.after,
                step.balance,
                step.status,
                "; ".join(step.anomalies),
                "; ".join(step.observation_ids),
            ]
        )
    observation_rows: list[list[str | int | None]] = [
        [
            "Observasjon-ID",
            "Ledd-ID",
            "Datasett",
            "Kildesystem",
            "Kildeenhet",
            "Revisjon",
            "Tolkning",
            "Artefakt SHA-256",
            "Kilderad",
            "Beløpsfelt",
            "Kapittel",
            "Post",
            "Underpost",
            "Detaljnivå",
            "Budsjettår",
            "Vedtaksdato",
            "Talltype",
            "Runde",
            "Beløpsbetydning",
            "Omfang",
            "Råverdi",
            "Beløpstekst",
            "Kroner normalisert",
            "Råenhet",
            "Multiplikator",
            "Strukturering",
            "Faglig kontroll",
            "Avvik",
            "Dokumentreferanser",
            "Nullgrunner",
            "Alle felt (JSON)",
        ]
    ]
    for id in result.selection.extraction_ids:
        extraction = read_extraction(database, id)
        for observation in extraction.observations:
            if observation.id not in result.input_observation_ids:
                continue
            observation_rows.append(
                [
                    observation.id,
                    observation.step_id,
                    observation.dataset,
                    observation.source_system,
                    observation.source_unit_id,
                    observation.source_revision_id,
                    observation.interpretation_id,
                    observation.pointer.artifact_id,
                    observation.pointer.locator,
                    observation.pointer.amount_field,
                    observation.chapter,
                    observation.post,
                    observation.subpost,
                    observation.detail_level,
                    observation.budget_year,
                    observation.decision_date,
                    observation.figure_type,
                    observation.round,
                    observation.meaning,
                    observation.scope,
                    observation.raw_value,
                    observation.amount_text,
                    observation.amount_kr,
                    observation.unit_raw,
                    observation.multiplier,
                    observation.structure_status,
                    observation.professional_status,
                    json.dumps(
                        [asdict(issue) for issue in observation.anomalies], ensure_ascii=False
                    ),
                    json.dumps(
                        [asdict(reference) for reference in observation.documents],
                        ensure_ascii=False,
                    ),
                    json.dumps(dict(observation.null_reasons), ensure_ascii=False),
                    json.dumps(asdict(observation), ensure_ascii=False),
                ]
            )
    note_rows: list[list[str | int | None]] = [
        ["Felt", "Verdi"],
        ["Beregnings-ID", result.id],
        ["Regelversjon", result.rule_version],
        ["Opprettet UTC", result.created_utc],
        ["Aktuell kontekst vurdert UTC", result.assessed_utc],
        ["Status i beregningen", result.status],
        ["Beregnet saldo i bevart utvalg", result.calculated_balance],
        ["Avstemt saldo med samme grunnlag", result.approved_balance],
        ["Ferskhet", result.freshness],
        ["Dekningsgrense", result.coverage],
        ["Utvalg", json.dumps(asdict(result.selection), ensure_ascii=False)],
        ["Inn-observasjoner", json.dumps(result.input_observation_ids)],
    ]
    for issue in result.anomalies:
        note_rows.append(["Beregningsavvik", json.dumps(asdict(issue), ensure_ascii=False)])
    for label, contexts in (
        ("Bevart kildekontekst", result.contexts),
        ("Nåværende kildekontekst", result.current_contexts),
    ):
        for context in contexts:
            note_rows.append([label, json.dumps(asdict(context), ensure_ascii=False)])
    for id in result.selection.extraction_ids:
        extraction = read_extraction(database, id)
        note_rows.extend(
            [
                ["Beløpsuttrekk", id],
                ["Artefakt", json.dumps(asdict(extraction.artifact), ensure_ascii=False)],
                [
                    "Tolkningsregler",
                    json.dumps(asdict(extraction.interpretation), ensure_ascii=False),
                ],
            ]
        )
        note_rows.extend(
            ["Adapteravvik", json.dumps(asdict(issue), ensure_ascii=False)]
            for issue in extraction.anomalies
        )
    for sheet, rows in (
        (history, history_rows),
        (observation_sheet, observation_rows),
        (notes, note_rows),
    ):
        for row_number, values in enumerate(rows, 1):
            for column, value in enumerate(values, 1):
                # Kildetekst kan være lengre enn én Excel-celle. Fortsett i egne rader i grunnlaget.
                if isinstance(value, str) and len(value) > 32767:
                    note_rows.extend(
                        [
                            f"Fortsettelse {sheet.title}!{get_column_letter(column)}{row_number}",
                            value[start : start + 30000],
                        ]
                        for start in range(30000, len(value), 30000)
                    )
                    value = value[:30000]
                put(sheet.cell(row_number, column), value)
                if row_number == 1:
                    sheet.cell(row_number, column).font = Font(
                        name="Arial", bold=True, color="FFFFFF"
                    )
                    sheet.cell(row_number, column).fill = PatternFill("solid", fgColor="181C62")
                sheet.column_dimensions[get_column_letter(column)].width = (
                    25 if sheet != notes else 50
                )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
    for row in history.iter_rows(min_row=2, min_col=6, max_col=10):
        for cell in row:
            cell.number_format = "# ##0;[Red]-# ##0;0"
    for row in observation_sheet.iter_rows(min_row=2, min_col=23, max_col=23):
        row[0].number_format = "# ##0;[Red]-# ##0;0"
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        os.replace(temporary, output)
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)
