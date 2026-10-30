"""T6s selvforsynte produktfiler med egen presentasjons-/ressursprofil."""

import base64
from copy import copy
from dataclasses import asdict
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.writer.theme import theme_xml

from innlesning.innstilling95 import reconstruct
from innlesning.innstilling95_model import Report
from innlesning.observations import encode
from innlesning.visningspilot_model import PilotReport

from . import visningspilot_excel as old_excel
from .extraction_excel import Value, write_rows
from .innstilling95 import contents, export_json, report_from_json, report_to_json
from .omtale import report_from_json as read_omtale
from .omtale_excel_guard import SheetShape, inspect_excel
from .perioderegister_excel import _package_profile, _style_profile
from .saksdokumenter import _publish

SHEET_NAMES = old_excel.SHEET_NAMES | {"Fullmakter", "Grenbevis", "T6-dekning", "Oppskrifter"}


def _selected(report: Report, accounts: tuple[str, ...], document_ids: tuple[str, ...]) -> str:
    selection = report.selection
    match = (
        selection.document_id in document_ids
        if selection.mode == "document"
        else 2024 in selection.years
        and any(
            account.split("/")[0] in selection.chapters
            or tuple(account.split("/")) in selection.accounts
            for account in accounts
        )
    )
    return "ja" if match else "nei"


def _book(report: Report, frozen: str) -> Workbook:
    data = contents(report.basis)
    projection = PilotReport(
        report.id, data, report.selection, report.return_selection, report.view, report.version
    )
    book = old_excel._book(projection, frozen)
    book["Visningsgrunnlag"]["B1"] = "Frosset innstilling95-1 med originalbyte"
    adopted = [o for o in data.observations if o.figure_type == "stortingsvedtak"]
    for row in book["Utvalg"].iter_rows(min_row=2):
        if row[0].value == "Globale vedtaksfelt":
            row[1].value = len(adopted)
        if row[0].value == "Globale vedtatte hendelser":
            row[1].value = len({o.step_id for o in adopted})
    formal = sum(
        d.key in {"vedtak360", "vedtak361", "vedtak362", "vedtak363", "vedtak364"}
        and d.raw_coverage == "komplett_rekonstruert"
        for d in data.documents
    )
    for label, count in (
        ("Nye separate fullmaktsvedtak", formal),
        ("Bevart separat fullmaktsvedtak211", 1),
        ("Samlede separate fullmaktsvedtak", formal + 1),
    ):
        book["Utvalg"].append([label, count])
    book["Utvalg"].column_dimensions["A"].width = 44
    book["Utvalg"].auto_filter.ref = book["Utvalg"].dimensions
    for row_number in (13, 14, 15):
        book["Utvalg"].row_dimensions[row_number].height = 36
        for cell in book["Utvalg"][row_number]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    book["Dokumenter"]["G1"] = "Lagret strukturdekning (T1 bevart)"
    bindings = {
        b.inventory_id: b.observation_ids
        for b in (*report.basis.t5_basis.bindings, *report.basis.bindings)
    }
    current_ids = {b.inventory_id for b in report.basis.bindings}
    inventory = book["Inventar"]
    inventory["G1"] = "Lagret strukturstatus (T1 bevart)"
    inventory["M1"], inventory["N1"] = "T5/T6-feltbinding", "Gjeldende tabellstatus"
    for column_number in (13, 14):
        inventory.cell(1, column_number).font = copy(inventory["G1"].font)
        inventory.cell(1, column_number).fill = copy(inventory["G1"].fill)
    for row_number, item in enumerate(data.inventory, 2):
        inventory.cell(row_number, 13, encode(bindings.get(item.id, ())))
        inventory.cell(
            row_number,
            14,
            "strukturert_T6"
            if item.id in current_ids
            else "strukturert_T5"
            if bindings.get(item.id)
            else "historisk_T1_bevart",
        )
    inventory.auto_filter.ref = inventory.dimensions
    accounts = book["Kontoer"]
    accounts["G1"] = "Lagret strukturstatus (T1 bevart)"
    accounts["H1"] = "Gjeldende T5/T6-tabellstatus"
    accounts["H1"].font = copy(accounts["G1"].font)
    accounts["H1"].fill = copy(accounts["G1"].fill)
    for account_row in accounts.iter_rows(min_row=2):
        identifier = str(account_row[5].value)
        accounts.cell(
            int(account_row[0].row or 0),
            8,
            "strukturert_T6"
            if identifier in current_ids
            else "strukturert_T5"
            if bindings.get(identifier)
            else "historisk_T1_bevart",
        )
    accounts.auto_filter.ref = accounts.dimensions
    book["Relasjoner"]["D1"] = "Kildebundet rolle (T1 bevart)"
    book["Relasjoner"]["E1"] = "Kildebundet omfang; T5-grener i Grenbevis"
    book["Relasjoner"].column_dimensions["E"].width = 52
    coverage_sheet = book["Dekning"]
    for row in coverage_sheet.iter_rows(min_row=2):
        if str(row[0].value).startswith(
            ("prop16_delvis_strukturert:", "inns95_strukturering_ikke_levert:")
        ):
            row[0].value = "Historisk T1: " + str(row[0].value)
    # Råpayload ligger bare i den selvforsynte kapselen. Synlig originalark peker dit.
    originals: list[list[Value]] = [
        [
            "Kilde-ID",
            "Originalnavn",
            "Dokumentnøkkel",
            "Format",
            "SHA-256",
            "Byte",
            "Forventet SHA-256",
            "Status",
            "Metadata",
            "Payloadreferanse",
        ]
    ]
    old = read_omtale(data.o2b_json)
    for source in (*data.sources, *report.basis.t5_basis.sources, *report.basis.sources):
        originals.append(
            [
                source.id,
                source.name,
                source.document_key,
                source.media_type,
                source.sha256,
                source.byte_count,
                source.expected_sha256,
                source.status,
                source.metadata_json,
                "Visningsgrunnlag: " + source.id,
            ]
        )
    originals.extend(
        [
            s.id,
            s.name,
            s.document_id,
            s.media_type,
            s.sha256,
            s.byte_count,
            s.expected_sha256,
            s.status,
            s.metadata_json,
            "Visningsgrunnlag: " + s.id,
        ]
        for s in old.basis.sources
    )
    for recipe in report.basis.recipes:
        originals.append(
            [
                recipe.id,
                recipe.document_key,
                recipe.document_key,
                "application/xml",
                recipe.expected_sha256,
                recipe.expected_bytes,
                recipe.expected_sha256,
                recipe.status,
                encode(asdict(recipe)),
                "Rekonstrueres: " + recipe.id,
            ]
        )
    original_index = book.sheetnames.index("Originaler")
    book.remove(book["Originaler"])
    write_rows(book, "Originaler", originals)
    book.move_sheet(book["Originaler"], offset=original_index - (len(book.sheetnames) - 1))
    tables: dict[str, list[list[Value]]] = {
        "Fullmakter": [
            [
                "Romertall",
                "Punkt",
                "Kontoer",
                "Felles ramme kr",
                "Rolle",
                "Periode",
                "Status",
                "Kilde-ID",
                "Lokator",
                "Inventar-ID",
                "Råtekst med vilkår",
                "Valgt",
                "Dokument",
            ]
        ],
        "Grenbevis": [
            [
                "Prop.-punkt",
                "Kontoer",
                "Sak",
                "Innst.-punkt",
                "Vedtaks-ID-er",
                "Beløpsstadium",
                "Status",
                "Fra dokument-ID",
                "Til dokument-ID",
                "Kildebevis",
                "Gren-ID",
                "Valgt",
            ]
        ],
        "T6-dekning": [
            [
                "Dokument-ID",
                "Tabellgrupper",
                "Tabellfelt",
                "Kontoer",
                "Gjeldende status",
                "Omfang",
                "Historisk T1-avvik",
                "Overlay-ID",
            ]
        ],
    }
    for a in (*report.basis.t5_basis.authorizations, *report.basis.authorizations):
        tables["Fullmakter"].append(
            [
                a.roman,
                a.point,
                ", ".join(a.accounts),
                a.frame_kr,
                a.role,
                a.period,
                a.status,
                a.source_id,
                a.locator,
                a.id,
                a.text,
                _selected(report, a.accounts, (a.document_id,)),
                next(d.title for d in data.documents if d.id == a.document_id),
            ]
        )
    for b in report.basis.t5_basis.branches:
        tables["Grenbevis"].append(
            [
                b.point,
                ", ".join(b.accounts),
                b.case_id,
                b.committee_point,
                ", ".join(b.decision_ids),
                "Historisk T5: " + b.amount_stage,
                b.status,
                b.from_document_id,
                b.to_document_id,
                encode(b.evidence_ids),
                b.id,
                _selected(report, b.accounts, (b.from_document_id, b.to_document_id)),
            ]
        )
    t5_coverages = report.basis.t5_basis.coverage
    for legacy in t5_coverages:
        document = next(d for d in data.documents if d.id == legacy.document_id)
        tables["T6-dekning"].append(
            [
                legacy.document_id,
                legacy.table_groups,
                legacy.table_fields,
                encode(document.accounts),
                legacy.status,
                legacy.scope,
                legacy.historical_issue,
                legacy.id,
            ]
        )
    for c in report.basis.coverage:
        tables["T6-dekning"].append(
            [
                c.document_id,
                c.table_groups,
                c.table_fields,
                encode(c.accounts),
                c.status,
                c.scope,
                c.historical_issue,
                c.id,
            ]
        )
    tables["Oppskrifter"] = [
        [
            "Oppskrift-ID",
            "Dokument",
            "Forelder-ID",
            "Sidevogn-ID",
            "Objektstart",
            "Objektslutt",
            "Prefix",
            "Suffixstart",
            "SHA-256",
            "Byte",
            "Status",
        ]
    ]
    for r in report.basis.recipes:
        tables["Oppskrifter"].append(
            [
                r.id,
                r.document_key,
                r.parent_source_id,
                r.sidecar_source_id,
                r.object_start,
                r.object_end,
                r.prefix_bytes,
                r.suffix_start,
                r.expected_sha256,
                r.expected_bytes,
                r.status,
            ]
        )
    for name, rows in tables.items():
        write_rows(book, name, rows)
    for name in (*tables, "Originaler"):
        sheet = book[name]
        for row in sheet:
            sheet.row_dimensions[int(row[0].row or 0)].height = 48 if row[0].row == 1 else 36
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if type(cell.value) is int:
                    cell.number_format = "#,##0"
        for column_number in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(column_number)].width = 32
    for name, widths in {
        "Fullmakter": {"A": 12, "B": 10, "C": 36, "D": 23, "E": 32, "F": 54, "K": 110},
        "Grenbevis": {"A": 16, "B": 105, "C": 15, "D": 16, "E": 32, "F": 30, "G": 26},
        "T6-dekning": {"A": 24, "B": 18, "C": 18, "D": 80, "E": 52, "F": 80, "G": 80},
        "Forløp": {"B": 44, "G": 24, "H": 22, "I": 28, "J": 24, "K": 22},
        "Inventar": {"G": 36, "H": 80, "M": 42, "N": 32},
        "Kontoer": {"D": 42, "G": 36, "H": 36},
    }.items():
        for column, width in widths.items():
            book[name].column_dimensions[column].width = width
    book["Grenbevis"].row_dimensions[3].height = 66
    comparisons = book["Sammenligninger"]
    for column, width in {
        "B": 16,
        "C": 22,
        "D": 44,
        "E": 44,
        "F": 24,
        "G": 18,
        "H": 24,
        "I": 80,
        "J": 32,
    }.items():
        comparisons.column_dimensions[column].width = width
    for row in comparisons:
        comparisons.row_dimensions[int(row[0].row or 0)].height = 48 if row[0].row == 1 else 36
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for sheet in (book["Forløp"], book["Inventar"], book["Kontoer"]):
        for row in sheet:
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.row_dimensions[int(row[0].row or 0)].height = (
                48 if row[0].row == 1 or sheet.title == "Inventar" else 36
            )
    # Overlegg og fullmakter har klikkbar dokumentreferanse, i begge retninger.
    document_rows = {d.id: row for row, d in enumerate(data.documents, 2)}
    for number, authorization in enumerate(
        (*report.basis.t5_basis.authorizations, *report.basis.authorizations), 2
    ):
        old_excel._link(
            book,
            "Fullmakter",
            f"M{number}",
            "Dokumenter",
            f"B{document_rows[authorization.document_id]}",
        )
    coverage_documents = tuple(c.document_id for c in t5_coverages) + tuple(
        c.document_id for c in report.basis.coverage
    )
    for number, document_id in enumerate(coverage_documents, 2):
        old_excel._link(
            book, "T6-dekning", f"A{number}", "Dokumenter", f"B{document_rows[document_id]}"
        )
        old_excel._link(
            book, "Dokumenter", f"G{document_rows[document_id]}", "T6-dekning", f"E{number}"
        )
    for number, branch in enumerate(report.basis.t5_basis.branches, 2):
        old_excel._link(
            book,
            "Grenbevis",
            f"I{number}",
            "Dokumenter",
            f"B{document_rows[branch.to_document_id]}",
        )
    return book


def read_excel(path: Path) -> Report:
    raw, shapes = inspect_excel(
        path,
        sheet_names=SHEET_NAMES,
        snapshot_name="Visningsgrunnlag",
        max_expanded_bytes=128 * 1024 * 1024,
    )
    _package_profile(path, len(shapes))
    # Pakkprofilen binder temarelasjon og innholdstype. Bind også hele temaet
    # til skriverens faste payload: cellefontenes theme/scheme er bare pekere.
    wanted_theme = theme_xml.encode("utf8")
    with ZipFile(path) as archive:
        if (
            archive.getinfo("xl/theme/theme1.xml").file_size != len(wanted_theme)
            or archive.read("xl/theme/theme1.xml") != wanted_theme
        ):
            raise ValueError("T6-Excel har endret tema for farger/skrift")
    report = report_from_json(raw)
    expected = _book(report, raw)
    try:
        if shapes != {s.title: SheetShape(s.max_row, s.max_column) for s in expected}:
            raise ValueError("T6-Excel har uventede ark/celleområder")
        _style_profile(path, expected)
        saved = load_workbook(path, data_only=False, keep_links=False)
        try:
            if saved.sheetnames != expected.sheetnames or saved.views != expected.views:
                raise ValueError("T6-Excel har endret arkrekkefølge/arbeidsbokvisning")
            for wanted_sheet, actual_sheet in zip(expected, saved, strict=True):
                for dimension in wanted_sheet.column_dimensions.values():
                    dimension.reindex()
                if (
                    actual_sheet.sheet_state != wanted_sheet.sheet_state
                    or actual_sheet.sheet_view != wanted_sheet.sheet_view
                    or actual_sheet.sheet_format != wanted_sheet.sheet_format
                    or actual_sheet.auto_filter != wanted_sheet.auto_filter
                    or actual_sheet.merged_cells != wanted_sheet.merged_cells
                    or set(actual_sheet.row_dimensions) != set(wanted_sheet.row_dimensions)
                    or set(actual_sheet.column_dimensions) != set(wanted_sheet.column_dimensions)
                    or any(
                        dict(actual_sheet.row_dimensions[k]) != dict(d)
                        for k, d in wanted_sheet.row_dimensions.items()
                    )
                    or any(
                        dict(actual_sheet.column_dimensions[k]) != dict(d)
                        for k, d in wanted_sheet.column_dimensions.items()
                    )
                ):
                    raise ValueError("T6-Excel har endret synlighet/oppsett: " + actual_sheet.title)
                for wanted_row, actual_row in zip(wanted_sheet, actual_sheet, strict=True):
                    for wanted, actual in zip(wanted_row, actual_row, strict=True):
                        value = None if wanted.value == "" else wanted.value
                        if (
                            actual.value != value
                            or type(actual.value) is not type(value)
                            or isinstance(value, str)
                            and actual.data_type != "s"
                            or actual.data_type == "f"
                            or actual.number_format != wanted.number_format
                            or copy(actual.font) != copy(wanted.font)
                            or copy(actual.fill) != copy(wanted.fill)
                            or copy(actual.border) != copy(wanted.border)
                            or copy(actual.alignment) != copy(wanted.alignment)
                            or copy(actual.protection) != copy(wanted.protection)
                            or (actual.hyperlink.location if actual.hyperlink else None)
                            != (wanted.hyperlink.location if wanted.hyperlink else None)
                            or (actual.hyperlink.target if actual.hyperlink else None)
                            != (wanted.hyperlink.target if wanted.hyperlink else None)
                        ):
                            raise ValueError(
                                "T6-Excel avviker i celle/stil/lenke: "
                                + actual_sheet.title
                                + "!"
                                + actual.coordinate
                            )
        finally:
            saved.close()
    finally:
        expected.close()
    return report


def export_excel(report: Report, output: Path) -> None:
    if output.exists() or output.suffix.lower() != ".xlsx":
        raise ValueError("Velg nytt T6-Excel-filnavn")
    raw = report_to_json(report)
    book = _book(report, raw)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        if read_excel(temporary) != report:
            raise ValueError("Skrevet T6-Excel avviker")
        _publish(temporary, output)
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def export_pair(report: Report, frozen: Path, excel: Path) -> None:
    if frozen.resolve() == excel.resolve() or frozen.exists() or excel.exists():
        raise ValueError("Velg to forskjellige nye filnavn")
    export_json(report, frozen)
    export_excel(report, excel)


def export_originals(report: Report, directory: Path) -> None:
    report_to_json(report)
    directory.mkdir(parents=True, exist_ok=False)
    old = read_omtale(report.basis.t5_basis.t1_basis.o2b_json)
    entries = [("t1/" + s.name, s.content_base64) for s in report.basis.t5_basis.t1_basis.sources]
    entries += [("o2b/" + s.name, s.content_base64) for s in old.basis.sources]
    entries += [(s.name, s.content_base64) for s in report.basis.sources]
    entries += [("t5/" + s.name, s.content_base64) for s in report.basis.t5_basis.sources]
    for recipe in report.basis.recipes:
        raw = reconstruct(report.basis.t5_basis, report.basis.sources, recipe)
        if raw is not None:
            identifier = str(40034220 + int(recipe.document_key[6:]) - 359)
            entries.append(("vedtak/" + identifier + ".xml", base64.b64encode(raw).decode("ascii")))
    for name, content in entries:
        if content is None:
            continue
        path = (directory / name).resolve()
        if not path.is_relative_to(directory.resolve()):
            raise ValueError("T6-originalsti utenfor eksportkatalog")
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("xb") as handle:
            handle.write(base64.b64decode(content, validate=True))
    (directory / "index.json").write_text(
        encode(
            {
                "basis_id": report.basis.id,
                "coverage": [asdict(c) for c in report.basis.coverage],
                "sources": [s.id for s in report.basis.sources],
                "recipes": [asdict(r) for r in report.basis.recipes],
            }
        ),
        encoding="utf8",
    )
