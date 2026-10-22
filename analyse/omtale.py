"""Opt-in O2a: selvforsynt omtale-pilot-1 og faktisk gjenlest Excel."""

import argparse
import json
import re
import tempfile
from dataclasses import asdict, dataclass
from pathlib import Path

from lxml import etree, html
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from innlesning.ingestion import NS, field, structure_node, text, validate_source
from innlesning.model import Decision, Report
from innlesning.observations import (
    AmountObservation,
    Artifact,
    DocumentReference,
    Extraction,
    SourcePointer,
    encode,
    identity,
)
from innlesning.omtale import (
    EXTRACTOR,
    OmtaleBasis,
    Section,
    compact,
    identified,
    import_omtale,
    source_bytes,
    validate_basis,
)
from innlesning.sync import FINGERPRINT_VERSION, PARSER_VERSION

from .extraction_excel import Value, verify_continued_text
from .omtale_excel_guard import SheetShape, inspect_excel
from .saksdokumenter import (
    SaksReport,
    _decorate,
    _publish,
    export_built_saks_excel,
    verify_saks_excel,
)
from .saksdokumenter import (
    report_from_json as saks_from_json,
)
from .saksdokumenter import (
    report_to_json as saks_to_json,
)
from .session import SessionRow
from .session_excel import write_rows
from .session_serialization import MAX_REPORT_BYTES, _constant, _decode, _object

FORMAT = "omtale-pilot-1"


@dataclass(frozen=True)
class Relation:
    id: str
    section_id: str
    event_id: str
    decision_revision_id: str
    observation_ids: tuple[str, ...]
    accounts: tuple[str, ...]
    role: str
    actor: str
    evidence_ids: tuple[str, ...]
    status: str
    professional_status: str
    reason: str


@dataclass(frozen=True)
class OmtaleReport:
    id: str
    base_json: str
    basis: OmtaleBasis
    relations: tuple[Relation, ...]
    visible_section_ids: tuple[str, ...]
    issues: tuple[str, ...]
    version: str = FORMAT


def _post1_group(rows: tuple[etree._Element, ...]) -> bool:
    """Konto, endring og fra/til må stå i samme lokale trippel av tabellrader."""
    if len(rows) != 3 or any(row.tag != "tr" or len(row.findall("td")) != 4 for row in rows):
        return False
    chapter, delta, levels = [
        [compact("".join(cell.itertext())) for cell in row.findall("td")] for row in rows
    ]
    return (
        chapter[0] == "1429"
        and chapter[1] == chapter[3] == ""
        and bool(chapter[2])
        and delta[0] == ""
        and delta[1] in {"1", "01"}
        and delta[2] == "Driftsutgifter, økes med"
        and delta[3] == "5 796 000"
        and levels == ["", "", "fra kr 160 243 000 til kr 166 039 000", ""]
    )


def _table_proves_post1(section: Section, basis: OmtaleBasis) -> bool:
    """Forslagsrader er relasjonsbelegg, aldri nye beløpsobservasjoner."""
    if section.key == "PT":
        source = next(s for s in basis.sources if s.id == section.source_id)
        raw = source_bytes(source)
        assert raw is not None
        doc = html.fromstring(raw.decode("utf-8"), parser=html.HTMLParser(no_network=True))
        rows = doc.xpath(section.locator)
        if len(rows) != 1:
            return False
        row = rows[0]
        before, after = row.getprevious(), row.getnext()
        return before is not None and after is not None and _post1_group((before, row, after))
    doc = html.fromstring(section.html)
    matches = 0
    for table in doc.findall(".//table"):
        rows = table.findall(".//tr")
        matches += sum(_post1_group(tuple(rows[n : n + 3])) for n in range(len(rows) - 2))
    return matches == 1


def _original_decision(basis: OmtaleBasis) -> tuple[Decision, Artifact]:
    """Les både konvolutt og vedtaksmetadata fra de bevarte originalbytene."""
    sources = {s.name: s for s in basis.sources}
    original, provenance = sources["b6/40034048.xml"], sources["b6/40034048.json"]
    if original.status != "kildeverifisert" or provenance.status != "kildeverifisert":
        raise ValueError("Vedtaksoriginal eller proveniens er ikke kildeverifisert")
    raw, metadata_raw = source_bytes(original), source_bytes(provenance)
    if raw is None or metadata_raw is None:
        raise ValueError("Vedtaksoriginal eller proveniens mangler")
    metadata = json.loads(metadata_raw, object_pairs_hook=_object, parse_constant=_constant)
    duplicated = json.loads(
        original.metadata_json, object_pairs_hook=_object, parse_constant=_constant
    )
    if not isinstance(metadata, dict) or not isinstance(duplicated, dict):
        raise ValueError("Ugyldig vedtaksproveniens")
    if encode(duplicated) != encode(metadata):
        raise ValueError("XML-kildens metadata avviker fra bevart vedtaksproveniens")
    session = validate_source(raw, metadata)
    root = etree.fromstring(
        raw, etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)
    )
    if root.getroottree().docinfo.doctype or root.tag != NS + "stortingsvedtak_oversikt":
        raise ValueError("Ugyldig vedtaksoriginal")
    for name in ("sesjon_id", "versjon", "stortingsvedtak_liste"):
        if len(root.findall(NS + name)) != 1:
            raise ValueError("Tvetydig vedtakskonvolutt")
    if field(root, "sesjon_id") != session or field(root, "versjon") != "1.6":
        raise ValueError("Vedtaksoriginalens sesjon/format avviker fra proveniensen")
    nodes = root.findall(NS + "stortingsvedtak_liste/" + NS + "stortingsvedtak")
    if len(nodes) != 1:
        raise ValueError("Piloten krever ett entydig originalvedtak")
    node = nodes[0]
    for name in (
        "id",
        "sak_id",
        "stortingsvedtak_nummer",
        "stortingsvedtak_dato_tid",
        "stortingsvedtak_lenke_url",
        "stortingsvedtak_tittel",
        "stortingsvedtak_tekst",
        "stortingsvedtak_type",
    ):
        if len(node.findall(NS + name)) != 1:
            raise ValueError("Tvetydig originalfelt: " + name)
    parsed = Report("", original.sha256 or "", metadata, True)
    structure_node(parsed, node, session)
    decision = parsed.decision
    if decision is None or (decision.id, decision.case_id, decision.number, decision.session) != (
        "40034048",
        "100927",
        "210",
        "2024-2025",
    ):
        raise ValueError("Feil vedtaksidentitet")
    artifact = Artifact(
        parsed.source_sha256,
        f"stortingsvedtak-{session}.xml",
        str(metadata.get("fixture_kind", "http_original")),
        "text/xml",
        len(raw),
        encode(metadata),
    )
    return decision, artifact


def _source_binding(base: SaksReport, basis: OmtaleBasis) -> None:
    """B2-revisjons-ID er en bevart UUID, bundet gjennom samme originale respons."""
    decision, artifact = _original_decision(basis)
    snapshots = [(r, d) for r, d in base.source_text.decisions if d and d.id == decision.id]
    if len(snapshots) != 1 or snapshots[0][1] != decision:
        raise ValueError("Frosne vedtaksmetadata avviker fra original")
    revision = snapshots[0][0]
    extractions = [e for e in base.source_text.extractions if e.source_revision_id == revision]
    unit = f"storting:{decision.session}:{decision.id}"
    if not extractions or any(
        e.artifact != artifact
        or e.source_unit_id != unit
        or e.source_system != "storting"
        or e.dataset != base.selection.base.dataset
        or e.id not in base.selection.base.extraction_ids
        for e in extractions
    ):
        raise ValueError("Vedtaksuttrekkets identitet/proveniens avviker fra original")
    inventory = [r for r in base.base.inventory if r.decision_id == decision.id]
    if len(inventory) != 1:
        raise ValueError("Vedtaksrevisjonen mangler entydig inventarbinding")
    row = inventory[0]
    if (
        (
            row.revision_id,
            row.session,
            row.api_type,
            row.date,
            row.title,
            row.url,
            row.source_sha256,
        )
        != (
            revision,
            decision.session,
            decision.api_type,
            decision.date,
            decision.title,
            decision.url,
            artifact.sha256,
        )
        or row.context.dataset != base.selection.base.dataset
        or row.context.session != decision.session
        or row.context.decision_id != decision.id
        or revision not in row.context.selected_revision_ids
        or not {e.id for e in extractions} <= set(row.extraction_ids)
    ):
        raise ValueError("Vedtaksinventar/revisjon avviker fra original")
    jobs = [j for j in base.base.snapshot.jobs if j.job_id in row.job_ids]
    if not jobs or any(
        j.session != decision.session
        or j.dataset != base.selection.base.dataset
        or j.fingerprint_version != FINGERPRINT_VERSION
        or j.parser_version != PARSER_VERSION
        or not any(e.decision_id == decision.id and e.revision_id == revision for e in j.entries)
        for j in jobs
    ):
        raise ValueError("Vedtaksrevisjonens kildeobservasjon avviker fra original")
    contexts = [c for c in base.contexts if c.decision_revision_id == revision]
    if len(contexts) != 1 or (
        contexts[0].case_id,
        contexts[0].formal_number,
        contexts[0].session,
    ) != (decision.case_id, decision.number, decision.session):
        raise ValueError("Sak/dokumentbinding avviker fra original")


def _chain_issues(base: SaksReport, basis: OmtaleBasis) -> tuple[str, ...]:
    issues: list[str] = []
    sections = {s.key: s for s in basis.sections}
    required = {"P1", "PT", "K1", "KS", "KT", "V", "P1_PDF"}
    issues.extend(
        i for i in basis.issues if not i.startswith("mindretall_kontrollert_fravaerende:")
    )
    issues.extend("belegg_mangler:" + k for k in sorted(required - set(sections)))
    for key in ("PT", "KT", "V"):
        if key in sections and not _table_proves_post1(sections[key], basis):
            issues.append("konto_eller_forslagsbelop_motstrid:" + key)
    for key in ("P1", "K1"):
        section = sections.get(key)
        if section is not None:
            text = compact(section.text)
            hierarchy = tuple(compact(h.text) for h in section.hierarchy)
            if not any(h.startswith("Kap. 1429 ") for h in hierarchy) or not any(
                h.startswith(("Post 01 ", "Post 1 ")) for h in hierarchy
            ):
                issues.append("konto_hierarki_motstrid:" + key)
            if not all(v in text for v in ("post 22", "2,8", "5,8", "3", "lønnsoppgjøret")):
                issues.append("forklaring_omfang_motstrid:" + key)
    kt, voting = sections.get("KT"), sections.get("V")
    if kt and not all(
        v in compact(kt.text) for v in ("en samlet komité", "viser til proposisjonen")
    ):
        issues.append("tilrading_uavklart")
    if voting and not all(
        v in compact(voting.text)
        for v in ("Innst. 81 S", "Prop. 16 S", "1429", "enstemmig vedtatt")
    ):
        issues.append("vedtakelse_uavklart")
    if not basis.representation_links or any(
        link.status != "kildeverifisert" for link in basis.representation_links
    ):
        issues.append("representasjonsforbindelse_uavklart")
    sources = {s.name: s for s in basis.sources}
    decision = sources.get("b6/40034048.xml")
    case = sources.get("b8/sak-100927.xml")
    if decision is None or decision.status != "kildeverifisert":
        issues.append("vedtaksoriginal_mangler_eller_motstrid")
    else:
        try:
            _source_binding(base, basis)
        except (etree.LxmlError, ValueError, KeyError, TypeError) as error:
            issues.append("vedtaksbinding_motstrid:" + str(error))
    if case is None or case.status != "kildeverifisert":
        issues.append("saksoriginal_mangler_eller_motstrid")
    elif not any(
        e.case_id == "100927"
        and e.api_id == "40034048"
        and e.content
        and e.content.source_sha256 == case.sha256
        for e in base.enrichments
    ):
        issues.append("saksbinding_motstrid")
    return tuple(issues)


def _rows_bound_to_extraction(base: SaksReport, basis: OmtaleBasis, rows: list[SessionRow]) -> bool:
    """Et omberegnet snapshot er ikke bevis for en observasjon eller hendelse."""
    if len({r.extraction_id for r in rows}) != 1:
        return False
    matches = [
        e
        for e in base.source_text.extractions
        if e.id == rows[0].extraction_id and e.id in base.selection.base.extraction_ids
    ]
    if len(matches) != 1:
        return False
    extraction = matches[0]
    expected = [
        o
        for o in extraction.observations
        if (o.api_id, o.chapter, o.post, o.budget_year) == ("40034048", "1429", "1", 2024)
    ]
    if (
        len(expected) != 3
        or {o.id for o in expected} != {r.observation.id for r in rows}
        or not _extraction_matches_original(basis, extraction, expected)
    ):
        return False
    by_id = {o.id: o for o in expected}
    selected = [
        r
        for r in base.base.rows
        if r.observation.id in by_id
        or (
            r.observation.api_id,
            r.observation.chapter,
            r.observation.post,
            r.observation.budget_year,
        )
        == ("40034048", "1429", "1", 2024)
    ]
    return all(
        row.observation == by_id.get(row.observation.id)
        and row.extraction_id == extraction.id
        and row.observation.source_revision_id == extraction.source_revision_id
        and row.observation.interpretation_id == extraction.interpretation.id
        and row.observation.source_unit_id == extraction.source_unit_id
        and row.observation.pointer.artifact_id == extraction.artifact.sha256
        for row in [*rows, *selected]
    )


def _extraction_matches_original(
    basis: OmtaleBasis, extraction: Extraction, observations: list[AmountObservation]
) -> bool:
    """Kontroller også lokator/råfelt/hendelse hvis uttrekket selv er omberegnet."""
    try:
        decision, artifact = _original_decision(basis)
    except (etree.LxmlError, ValueError, KeyError, TypeError):
        return False
    if extraction.artifact != artifact:
        return False
    doc = html.fromstring(decision.html, parser=html.HTMLParser(no_network=True))
    tables = doc.findall(".//table")
    if len(tables) != 1:
        return False
    rows = tables[0].findall(".//tr")
    groups = [n for n in range(len(rows) - 2) if _post1_group(tuple(rows[n : n + 3]))]
    if len(groups) != 1:
        return False
    start = groups[0]
    chapter, delta, levels = rows[start : start + 3]
    if [text(cell) for cell in rows[0]] != ["Kap.", "Post", "Formål", "Kroner"]:
        return False
    years = set(re.findall(r"statsbudsjettet for (20\d{2})", text(doc), re.I))
    year_evidence = "Uttrykkelig budsjettår i vedtaksteksten"
    if not years:
        years = set(re.findall(r"statsbudsjettet (20\d{2})", decision.title, re.I))
        year_evidence = "Uttrykkelig budsjettår i vedtakstittelen: " + decision.title
    if years != {"2024"}:
        return False
    unit = f"storting:{decision.session}:{decision.id}"
    delta_locator = f"table[1]/tr[{start + 2}]"
    event = identity(unit, extraction.source_revision_id, delta_locator)
    reference = DocumentReference(
        decision.number,
        "stortingsvedtak_nummer",
        "vedtak",
        "stortingsvedtak",
        decision.number,
        decision.session,
        decision.id,
        decision.url,
        "identifisert",
    )
    source_fields = {
        "source_system": "storting",
        "source_unit_id": unit,
        "chapter_raw": text(chapter.findall("td")[0]),
        "chapter": "1429",
        "post_raw": text(delta.findall("td")[1]),
        "post": "1",
        "subpost_raw": None,
        "subpost": None,
        "chapter_name": text(chapter.findall("td")[2]),
        "purpose": text(delta.findall("td")[2]),
        "detail_level": "post",
        "component_of": None,
        "budget_year": 2024,
        "year_evidence": year_evidence,
        "decision_date": decision.date,
        "decision_session": decision.session,
        "api_id": decision.id,
        "event_id": None,
        "raw_type": "tekst",
        "formula": None,
        "cached_value": None,
        "unit_raw": "Kroner",
        "currency": "NOK",
        "multiplier": "1",
        "precision": "Hele kroner",
        "figure_type": "stortingsvedtak",
        "stage_raw": decision.title,
        "documents": (reference,),
    }
    for observation in observations:
        is_delta = observation.meaning == "endring"
        row = delta if is_delta else levels
        locator = delta_locator + "/td[4]" if is_delta else f"table[1]/tr[{start + 3}]/td[3]"
        raw_value = (
            text(delta.findall("td")[3])
            if is_delta
            else "160 243 000"
            if observation.meaning == "for_belop"
            else "166 039 000"
        )
        pointer = SourcePointer(
            extraction.artifact.sha256,
            locator,
            observation.meaning,
            raw_record=etree.tostring(row, encoding="unicode", method="html"),
        )
        if (
            observation.pointer != pointer
            or observation.step_id != event
            or observation.raw_value != raw_value
            or observation.amount_text != raw_value
            or any(getattr(observation, name) != value for name, value in source_fields.items())
        ):
            return False
    return True


def build_report(base: SaksReport, basis: OmtaleBasis) -> OmtaleReport:
    """Bare uttrykkelig valgt snapshot; ingen joins på bare konto eller beløp."""
    base_json = saks_to_json(base)
    validate_basis(basis)
    chain_issues = _chain_issues(base, basis)
    issues = [*basis.issues, *chain_issues]
    rows = [
        r
        for r in base.base.unfiltered_rows
        if r.observation.api_id == "40034048"
        and r.observation.chapter == "1429"
        and r.observation.post == "1"
        and r.observation.budget_year == 2024
    ]
    values = {r.observation.meaning: r.observation.amount_kr for r in rows}
    events = {r.observation.step_id for r in rows}
    revisions = {r.observation.source_revision_id for r in rows}
    if (
        len(rows) != 3
        or values != {"endring": 5_796_000, "for_belop": 160_243_000, "etter_belop": 166_039_000}
        or len(events) != 1
        or None in events
        or len(revisions) != 1
        or not _rows_bound_to_extraction(base, basis, rows)
    ):
        issues.append("210_post1_grunnlag_mangler_eller_motstrid")
        rows = []
    sections = {s.key: s for s in basis.sections}
    evidence = tuple(s.id for s in basis.sections if s.key in {"P1", "PT", "K1", "KS", "KT", "V"})
    evidence += tuple(s.id for s in basis.sources if s.media_type == "application/xml")
    relations: list[Relation] = []
    if rows:
        event = rows[0].observation.step_id
        for key in ("P1", "PT", "K1", "KS", "KT", "V"):
            section = sections.get(key)
            if section is None:
                continue
            valid = (
                not chain_issues
                and all(s.status == "kildeverifisert" for s in basis.sources)
                and all(s.status == "kildeverifisert" for s in basis.sections)
            )
            relations.append(
                identified(
                    "omtale-relasjon-1",
                    Relation(
                        "",
                        section.id,
                        event,
                        rows[0].observation.source_revision_id,
                        tuple(r.observation.id for r in rows),
                        section.accounts,
                        section.role,
                        section.actor,
                        evidence,
                        "kildeverifisert" if valid else "uavklart",
                        "ikke_faglig_godkjent",
                        "P1 → PT → K1 + KS + KT → V → 210. Avrundet lønn er ikke eksakt komponent."
                        if valid
                        else "Belegg/rolle/kilde er uavklart; se globale Omtaleavvik.",
                    ),
                )
            )
    # Kontofiltre begrenser visningen. Originaler, relasjoner og avvik forblir globale.
    show = any(r.observation.id in {o.observation.id for o in rows} for r in base.base.rows)
    visible = tuple(s.id for s in basis.sections if show and s.key != "PN")
    result = OmtaleReport(
        "", base_json, basis, tuple(relations), visible, tuple(dict.fromkeys(issues))
    )
    return identified(FORMAT, result)


def report_to_json(report: OmtaleReport) -> str:
    _validate(report)
    raw = encode({"format": FORMAT, "report": asdict(report)})
    if len(raw.encode("utf-8")) > MAX_REPORT_BYTES:
        raise ValueError("Omtalerapport overstiger 128 MiB")
    return raw


def _validate(report: OmtaleReport) -> None:
    base = saks_from_json(report.base_json)
    if report.version != FORMAT or report != build_report(base, report.basis):
        raise ValueError("Omtalerapport avviker fra gjenleste kilder/relasjonsbelegg")


def report_from_json(raw: str) -> OmtaleReport:
    if len(raw.encode("utf-8")) > MAX_REPORT_BYTES:
        raise ValueError("Omtalerapport overstiger 128 MiB")
    data = json.loads(raw, object_pairs_hook=_object, parse_constant=_constant)
    if not isinstance(data, dict) or set(data) != {"format", "report"} or data["format"] != FORMAT:
        raise ValueError("Ukjent omtaleformat")
    report: OmtaleReport = _decode(OmtaleReport, data["report"])
    _validate(report)
    return report


def read_json(path: Path) -> OmtaleReport:
    if path.stat().st_size > MAX_REPORT_BYTES:
        raise ValueError("Omtalerapport overstiger 128 MiB")
    return report_from_json(path.read_text("utf-8"))


def _tables(report: OmtaleReport) -> dict[str, list[list[Value]]]:
    sections: list[list[Value]] = [
        [
            "Seksjons-ID",
            "Belegg",
            "Dokument-ID",
            "Kilde-ID",
            "Lokator",
            "Avsender",
            "Rolle",
            "Omfang",
            "Status",
            "Hierarki med lokator",
            "Full tekst",
            "Full HTML",
            "Uttrekksversjon",
            "Valgt visning",
        ]
    ]
    for s in report.basis.sections:
        sections.append(
            [
                s.id,
                s.key,
                s.document_id,
                s.source_id,
                s.locator,
                s.actor,
                s.role,
                encode(s.accounts),
                s.status,
                encode([asdict(h) for h in s.hierarchy]),
                s.text,
                s.html,
                s.extractor,
                "ja" if s.id in report.visible_section_ids else "nei",
            ]
        )
    relations: list[list[Value]] = [
        [
            "Relasjons-ID",
            "Seksjons-ID",
            "Hendelses-ID",
            "Vedtaksrevisjon",
            "Observasjons-ID-er",
            "Omfang",
            "Rolle",
            "Avsender",
            "Belegg-ID-er",
            "Teknisk status",
            "Faglig status",
            "Grense",
        ]
    ]
    relations.extend(
        [
            [
                r.id,
                r.section_id,
                r.event_id,
                r.decision_revision_id,
                encode(r.observation_ids),
                encode(r.accounts),
                r.role,
                r.actor,
                encode(r.evidence_ids),
                r.status,
                r.professional_status,
                r.reason,
            ]
            for r in report.relations
        ]
    )
    sources: list[list[Value]] = [
        [
            "Kilde-ID",
            "Fil",
            "Dokument-ID",
            "Format",
            "SHA-256",
            "Byte",
            "Forventet SHA-256",
            "Status",
            "Transport-/innhentingsmetadata",
        ]
    ]
    sources.extend(
        [
            [
                s.id,
                s.name,
                s.document_id,
                s.media_type,
                s.sha256,
                s.byte_count,
                s.expected_sha256,
                s.status,
                s.metadata_json,
            ]
            for s in report.basis.sources
        ]
    )
    sources.extend(
        [
            [
                link.id,
                "PDF ↔ HTML",
                None,
                "representasjonsrelasjon",
                None,
                None,
                None,
                link.status,
                encode(asdict(link)),
            ]
            for link in report.basis.representation_links
        ]
    )
    # Frosset JSON i ordnede tekststykker gjør arbeidsboken selvforsynt uten eksterne filer.
    frozen = report_to_json(report)
    snapshot: list[list[Value]] = [["Del", "Frosset omtale-pilot-1 (inkludert originalbyte)"]]
    snapshot.extend([[i // 30000 + 1, frozen[i : i + 30000]] for i in range(0, len(frozen), 30000)])
    issue_rows: list[list[Value]] = [["Globalt avvik før kontofilter"]]
    issue_rows.extend([i] for i in report.issues)
    return {
        "Omtale": sections,
        "Relasjoner": relations,
        "Omtalekilder": sources,
        "Omtaleavvik": issue_rows,
        "Omtalegrunnlag": snapshot,
    }


def _read_full(sheet: Worksheet, row: int, column: int) -> Value:
    value = sheet.cell(row, column).value
    if value is None or type(value) is int:
        return value
    if not isinstance(value, str) or sheet.cell(row, column).data_type != "s":
        raise ValueError("Omtaletekst må være lagret som tekst")
    book = sheet.parent
    assert book is not None
    name = f"{sheet.title[:18]} fortsettelse"
    if name in book:
        pointer = sheet.cell(row, column).coordinate
        chunks = [
            (r[1].value, r[2].value)
            for r in book[name].iter_rows(min_row=2)
            if r[0].value == f"{sheet.title}!{pointer}"
        ]
        if [i for i, _ in chunks] != list(range(2, len(chunks) + 2)) or any(
            not isinstance(t, str) for _, t in chunks
        ):
            raise ValueError("Feil omtalevidereføring")
        value += "".join(str(t) for _, t in chunks)
    return value


def read_excel(path: Path) -> OmtaleReport:
    raw, _ = inspect_excel(path)
    report = report_from_json(raw)
    verify_excel(path, report)
    return report


def verify_excel(path: Path, report: OmtaleReport) -> None:
    _validate(report)
    raw, shapes = inspect_excel(path)
    if raw != report_to_json(report):
        raise ValueError("Excel-grunnlaget avviker fra rapporten")
    # Rekonstruer den komplette forventede B9/O2a-visningen fra den frosne rapporten.
    # Bare egenprodusert Excel lastes før inputens ark og dimensjoner er avklart.
    with tempfile.TemporaryDirectory(dir=path.parent) as directory:
        expected_path = Path(directory) / "expected.xlsx"
        export_built_saks_excel(
            Path(directory) / "unavailable.sqlite", saks_from_json(report.base_json), expected_path
        )
        expected_book = load_workbook(expected_path, data_only=False, keep_links=False)
    try:
        for name, rows in _tables(report).items():
            write_rows(expected_book, name, rows)
        expected_shapes = {
            sheet.title: SheetShape(sheet.max_row, sheet.max_column) for sheet in expected_book
        }
        if shapes != expected_shapes:
            raise ValueError("Excel har uventede ark, arkrader eller celleområder")
        _verify_cells(path, report, expected_book)
    finally:
        expected_book.close()


def _verify_cells(path: Path, report: OmtaleReport, expected_book: Workbook) -> None:
    verify_saks_excel(path, saks_from_json(report.base_json))
    saved = load_workbook(path, data_only=False, keep_links=False)
    try:
        for expected_sheet in expected_book:
            actual_sheet = saved[expected_sheet.title]
            for wanted_row, actual_row in zip(expected_sheet, actual_sheet, strict=True):
                for wanted_cell, actual_cell in zip(wanted_row, actual_row, strict=True):
                    value = None if wanted_cell.value == "" else wanted_cell.value
                    if (
                        actual_cell.value != value
                        or type(actual_cell.value) is not type(value)
                        or (isinstance(value, str) and actual_cell.data_type != "s")
                        or (actual_cell.hyperlink.target if actual_cell.hyperlink else None)
                        != (wanted_cell.hyperlink.target if wanted_cell.hyperlink else None)
                        or (actual_cell.hyperlink.location if actual_cell.hyperlink else None)
                        != (wanted_cell.hyperlink.location if wanted_cell.hyperlink else None)
                    ):
                        raise ValueError(
                            "Excel-celle/videreført tekst avviker fra frosset grunnlag: "
                            + actual_sheet.title
                            + "!"
                            + actual_cell.coordinate
                        )
        for name, rows in _tables(report).items():
            sheet = saved[name]
            if sheet.max_row != len(rows) or sheet.max_column != len(rows[0]):
                raise ValueError("Omtalearket har feil dimensjoner: " + name)
            # Kontroller også overflødige/feil pekere og fortsettelsesrader.
            expected = Workbook()
            try:
                write_rows(expected, name, rows)
                verify_continued_text(sheet, expected[name])
            finally:
                expected.close()
            for r, values in enumerate(rows, 1):
                for c, wanted in enumerate(values, 1):
                    actual = _read_full(sheet, r, c)
                    if actual != (None if wanted == "" else wanted):
                        raise ValueError(
                            "Omtalecelle avviker fra original/relasjon: "
                            + name
                            + "!"
                            + sheet.cell(r, c).coordinate
                        )
    finally:
        saved.close()


def export_pair(report: OmtaleReport, json_path: Path, excel_path: Path) -> None:
    """Verifiser begge stagingfiler før publisering. Ingen database åpnes."""
    if (
        json_path.resolve() == excel_path.resolve()
        or json_path.suffix.lower() != ".json"
        or (excel_path.suffix.lower() != ".xlsx")
        or json_path.exists()
        or excel_path.exists()
    ):
        raise ValueError("Omtaleeksport krever forskjellige nye .json- og .xlsx-filer")
    raw = report_to_json(report)
    json_path.parent.mkdir(parents=True, exist_ok=True)
    excel_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(dir=excel_path.parent) as directory:
        staging = Path(directory)
        frozen, excel = staging / "omtale.json", staging / "omtale.xlsx"
        frozen.write_text(raw, "utf-8")
        saved = read_json(frozen)
        export_built_saks_excel(
            staging / "unavailable.sqlite", saks_from_json(saved.base_json), excel
        )
        book = load_workbook(excel, data_only=False, keep_links=False)
        try:
            for name, rows in _tables(saved).items():
                write_rows(book, name, rows)
                _decorate(book[name])
            # Fulltekst får egne brede kolonner og høyde for praktisk lesing.
            book["Omtale"].column_dimensions["K"].width = 100
            book["Omtale"].column_dimensions["L"].width = 100
            for n in range(2, book["Omtale"].max_row + 1):
                book["Omtale"].row_dimensions[n].height = 300
            book.save(excel)
        finally:
            book.close()
        verify_excel(excel, saved)
        _publish(excel, excel_path)
        # JSON kan ligge på et annet filsystem; stage i den valgte målkatalogen.
        with tempfile.NamedTemporaryFile(dir=json_path.parent, suffix=".json", delete=False) as f:
            temporary = Path(f.name)
        try:
            temporary.write_bytes(frozen.read_bytes())
            _publish(temporary, json_path)
        finally:
            temporary.unlink(missing_ok=True)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    inputs = parser.add_mutually_exclusive_group(required=True)
    inputs.add_argument("--base", type=Path, help="Eksakt frosset b6-frozen-2 fra B8")
    inputs.add_argument(
        "--frosset", type=Path, help="Gjenles tidligere omtale-pilot-1 uten kilder/base"
    )
    parser.add_argument("--kilder", type=Path, help="O1-katalog med manifest/originals")
    parser.add_argument("--stotte", type=Path, help="Katalog med b6-/b8-støtteoriginaler")
    parser.add_argument("--uttrekksversjon", default=EXTRACTOR)
    parser.add_argument("--json", type=Path, required=True)
    parser.add_argument("--excel", type=Path, required=True)
    args = parser.parse_args()
    try:
        if args.frosset:
            if args.kilder or args.stotte:
                raise ValueError("Frosset gjenlesing bruker bare den valgte filen")
            report = read_json(args.frosset)
        else:
            if not args.kilder or not args.stotte:
                raise ValueError("--base krever --kilder og --stotte")
            if args.base.stat().st_size > MAX_REPORT_BYTES:
                raise ValueError("B8-grunnlaget overstiger 128 MiB")
            base = saks_from_json(args.base.read_text("utf-8"))
            basis = import_omtale(args.kilder, args.stotte, extractor=args.uttrekksversjon)
            report = build_report(base, basis)
        export_pair(report, args.json, args.excel)
        print(
            encode(
                {
                    "status": "eksportert",
                    "format": FORMAT,
                    "rapport": report.id,
                    "json": str(args.json),
                    "excel": str(args.excel),
                    "avvik": report.issues,
                }
            )
        )
        return 0
    except (OSError, ValueError, TypeError, KeyError) as error:
        print(encode({"status": "eksportfeil", "feil": str(error)}))
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
