"""P1s fem hovedfaner og én streng, selvforsynt offlinekapsel."""

import hashlib
import json
from collections import Counter
from copy import copy
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZipFile

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.styles.stylesheet import write_stylesheet
from openpyxl.utils import get_column_letter
from openpyxl.xml.functions import tostring

from innlesning.observations import encode
from innlesning.perioderegister import shared_sources
from innlesning.perioderegister_model import Report

from . import fleraarspilot_excel as old_excel
from .extraction_excel import Value, write_rows
from .komitepilot import legacy_payloads, legacy_report
from .omtale_excel_guard import SheetShape, _metadata, inspect_excel
from .perioderegister import (
    contents,
    export_json,
    monetary_report,
    report_from_json,
    report_to_json,
)
from .saksdokumenter import _publish

SHEET_NAMES = old_excel.SHEET_NAMES
_chunks = old_excel._chunks


def _tables(report: Report, frozen: str) -> dict[str, list[list[Value]]]:
    old = monetary_report(report.basis, report.selection)
    tables = old_excel._tables(old, "")
    basis = report.basis
    data = contents(basis)
    docs = {d.id: d for d in data.documents}
    proofs = {e.id: e for e in basis.evidence}
    tables["Forløp"][0].extend(["Postkobling", "Originalreferanse", "Lokator", "Registerledd-ID"])
    old_doc_ids = {
        c.document_id for c in data.contexts if c.observation_id in old.view.observation_ids
    }
    observations = {o.id: o for o in data.observations}
    for row in tables["Forløp"][1:]:
        ids = json.loads(str(row[11]))
        field = observations[ids[0]]
        row.extend(["proven_selected", field.pointer.artifact_id, field.pointer.locator, None])
    for e in basis.entries:
        if e.id not in report.view.entry_ids or e.document_id in old_doc_ids:
            continue
        doc = docs[e.document_id]
        first = next(
            (
                proofs[i]
                for i in e.evidence_ids
                if proofs[i].role
                in {"vedtaksoriginal", "hashbundet_tabellokator", "originalreferanse"}
            ),
            proofs[e.evidence_ids[0]] if e.evidence_ids else None,
        )
        tables["Forløp"].append(
            [
                ", ".join(e.accounts),
                e.budget_year,
                doc.title,
                doc.date,
                e.round + " / " + e.role,
                e.actor,
                None,
                None,
                None,
                None,
                e.status + "; " + e.professional_status,
                encode(e.observation_ids),
                e.id,
                doc.id,
                e.post_link,
                first.source_id if first else None,
                first.locator if first else None,
                e.id,
            ]
        )
    # Registerets dokumenterte behandlingsrekkefølge gjelder også de tekstlige
    # leddene. Forslag/fullmakter skal stå sammen med sin runde i Forløp.
    entry_order = {e.id: n for n, e in enumerate(basis.entries)}
    document_order = {e.document_id: n for n, e in reversed(tuple(enumerate(basis.entries)))}
    tables["Forløp"][1:] = sorted(
        tables["Forløp"][1:],
        key=lambda row: (
            int(str(row[1] or 0)),
            entry_order.get(str(row[17]), document_order.get(str(row[13]), len(basis.entries))),
            str(row[0]),
        ),
    )
    for d in basis.documents:
        tables["Dokumenter"].append(
            [
                "Dokument",
                d.id,
                d.title,
                d.kind,
                d.date,
                d.session,
                None,
                None,
                None,
                d.raw_coverage,
                d.structured_coverage,
                d.professional_status,
                encode(d.accounts),
                encode(d.source_ids),
                None,
                None,
                "ja" if d.id in report.view.document_ids else "nei",
            ]
        )
        for item in basis.inventory:
            if item.document_id == d.id:
                tables["Dokumenter"].append(
                    [
                        "Post",
                        item.id,
                        d.title,
                        item.role,
                        d.date,
                        d.session,
                        item.budget_year,
                        item.chapter,
                        item.post,
                        d.raw_coverage,
                        item.structure_status,
                        d.professional_status,
                        encode(item.observation_ids),
                        item.source_id,
                        item.locator,
                        item.text,
                        "ja" if d.id in report.view.document_ids else "nei",
                    ]
                )
    # CaseCandidate inneholder indeksinventaret uten oppdiktet budsjettår.
    for c in basis.case_candidates:
        d = docs[c.document_id]
        proof = proofs[c.evidence_ids[-1]]
        tables["Dokumenter"].append(
            [
                "Indeksinventar",
                c.id,
                d.title,
                c.status,
                None,
                c.index_session,
                c.budget_year,
                None,
                None,
                "komplett_indeksobjekt_fulltekst_mangler",
                c.reason,
                "ikke_faglig_godkjent",
                encode(c.accounts),
                proof.source_id,
                proof.locator,
                c.metadata_json,
                "ja" if c.document_id in report.view.document_ids else "nei",
            ]
        )
    for decision in basis.decision_candidates:
        d = docs[decision.document_id]
        proof = proofs[decision.evidence_ids[-1]]
        tables["Dokumenter"].append(
            [
                "Vedtakskandidat",
                decision.id,
                d.title,
                decision.status,
                d.date,
                decision.session,
                decision.budget_year,
                None,
                None,
                decision.raw_preservation,
                decision.product_structuring,
                decision.professional_approval,
                encode(decision.accounts),
                proof.source_id,
                proof.locator,
                decision.reason,
                "global",
            ]
        )
    for s in basis.statements:
        for n, text in enumerate(_chunks(s.text), 1):
            tables["Omtale"].append(
                [
                    s.id,
                    docs[s.document_id].title,
                    s.role,
                    s.actor,
                    s.period,
                    encode(s.accounts),
                    encode(s.hierarchy),
                    n,
                    text,
                    "",
                    s.source_id,
                    s.locator,
                    s.status,
                    "ja" if s.document_id in report.view.document_ids else "nei",
                ]
            )
    coverage: list[list[Value]] = [["Globalt forhold før filter", "Verdi/status"]]
    coverage.extend(
        [
            ["Rapport-ID", report.id],
            ["Grunnlags-ID", basis.id],
            ["Format", report.version],
            ["Valgt inngang", encode(asdict(report.selection))],
            [
                "Returvalg",
                encode(asdict(report.return_selection)) if report.return_selection else None,
            ],
            ["Årsdekning", basis.year_coverage],
            ["Faglig årsbalanse godkjent", "nei"],
            ["Valgte vedtaksfelt", len(report.view.monetary.adopted_observation_ids)],
            ["Valgte monetære hendelser", len(report.view.monetary.adopted_event_ids)],
            ["Indekskandidatposter", len(basis.case_candidates)],
            ["Vedtakskandidater", len(basis.decision_candidates)],
            ["Globale hull", len(basis.gaps)],
            ["Kildeunivers", "2 183 vedtaksobjekter; 1 364 saksindeksposter; 25 er bare søketreff"],
            ["Indeksstatus", encode(dict(Counter(c.status for c in basis.case_candidates)))],
            ["Gamle immutable kapsler", "Én T1-kapsel i T4; arvet fysisk duplisering er bevart"],
        ]
    )
    coverage.extend([g.key + " " + g.title, g.status + ": " + g.reason] for g in basis.gaps)
    coverage.extend(
        [c.id + " / sak " + c.case_id, c.status + ": " + c.reason] for c in basis.case_candidates
    )
    coverage.extend(
        [
            c.id + " / vedtak " + c.number,
            c.status
            + ": "
            + c.reason
            + "; "
            + encode(
                (
                    c.budget_year,
                    c.accounts,
                    c.post_link,
                    c.raw_preservation,
                    c.product_structuring,
                    c.professional_approval,
                    c.monetary_effect,
                )
            ),
        ]
        for c in basis.decision_candidates
    )
    coverage.extend(
        [
            issue,
            "P1-kildekontroll"
            if issue in basis.issues
            else "historisk T4-kapsel; P1-kilder og status står separat",
        ]
        for issue in data.issues
    )
    tables["Dekning og avvik"] = coverage
    for source in basis.sources:
        tables["Kontroll"].append(
            [
                "Kilde",
                source.id,
                source.document_key,
                source.payload_id,
                source.media_type,
                None,
                source.status,
                source.expected_sha256,
                source.name,
                None,
                source.metadata_json,
            ]
        )
    for evidence in basis.evidence:
        tables["Kontroll"].append(
            [
                "Bevis",
                evidence.id,
                evidence.source_id,
                None,
                evidence.role,
                None,
                "kildebundet",
                evidence.representation_sha256,
                evidence.locator,
                None,
                None,
            ]
        )
    for e in basis.entries:
        tables["Kontroll"].append(
            [
                "Registerledd",
                e.id,
                e.document_id,
                None,
                e.role,
                e.budget_year,
                e.status,
                None,
                encode(e.evidence_ids),
                None,
                encode(asdict(e)),
            ]
        )
    for r in basis.relations:
        tables["Kontroll"].append(
            [
                "Relasjon",
                r.id,
                r.from_id,
                r.to_id,
                r.role,
                encode((r.accounts, r.budget_year)),
                r.status,
                None,
                encode(r.evidence_ids),
                None,
                r.scope,
            ]
        )
    tables["Offlinegrunnlag"] = [["Del", "Frosset perioderegister-1 med kanonisk råregister"]]
    tables["Offlinegrunnlag"].extend([n, chunk] for n, chunk in enumerate(_chunks(frozen), 1))
    return tables


def _book(report: Report, raw: str) -> Workbook:
    book = Workbook()
    assert book.active is not None
    book.remove(book.active)
    widths = {
        "Forløp": (18, 14, 55, 14, 42, 22, 20, 20, 20, 20, 48, 24, 24, 24, 24, 30, 40, 24),
        "Dokumenter": (20, 24, 65, 35, 14, 16, 14, 12, 12, 45, 55, 30, 28, 26, 45, 95, 12),
        "Omtale": (24, 65, 42, 22, 20, 26, 30, 10, 100, 30, 26, 45, 35, 12),
        "Sammenligninger": (16, 24, 55, 55, 22, 22, 24, 24, 95, 28, 26, 26, 26, 65),
        "Dekning og avvik": (75, 120),
        "Kontroll": (22, 26, 26, 26, 42, 28, 35, 28, 70, 22, 100),
        "Offlinegrunnlag": (12, 120),
    }
    for name, rows in _tables(report, raw).items():
        if any(isinstance(value, str) and len(value) > 30000 for row in rows for value in row):
            raise ValueError("P1-visning krever eksplisitt tapsfri tekstdeling")
        write_rows(book, name, rows)
        sheet = book[name]
        for n, width in enumerate(widths[name], 1):
            sheet.column_dimensions[get_column_letter(n)].width = width
        for row in sheet:
            sheet.row_dimensions[int(row[0].row or 0)].height = (
                44
                if row[0].row == 1
                else (72 if name in {"Forløp", "Dokumenter", "Dekning og avvik"} else 30)
            )
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if (
                    sheet.cell(1, int(cell.column or 0)).value == "Dato"
                    and isinstance(cell.value, str)
                    and int(cell.row or 0) > 1
                    and cell.value
                ):
                    cell.value = datetime.fromisoformat(cell.value)
                    cell.number_format = 'dd"."mm"."yyyy'
                if type(cell.value) is int:
                    cell.number_format = (
                        "0"
                        if sheet.cell(1, int(cell.column or 0)).value in {"Budsjettår", "Del"}
                        else "#,##0"
                    )
    doc_rows = {
        str(r[1].value): r[0].row
        for r in book["Dokumenter"].iter_rows(min_row=2)
        if r[0].value == "Dokument"
    }
    account_rows: dict[tuple[str, int], int] = {}
    for row in book["Forløp"].iter_rows(min_row=2):
        doc = str(row[13].value)
        if doc in doc_rows:
            old_excel._link(book, "Forløp", f"C{row[0].row}", "Dokumenter", f"C{doc_rows[doc]}")
        if type(row[1].value) is int and str(row[0].value) in {"1429/1", "1429/22"}:
            account_rows.setdefault((str(row[0].value), row[1].value), int(row[0].row or 0))
    for row in book["Dokumenter"].iter_rows(min_row=2):
        account = (str(row[7].value) + "/" + str(row[8].value), int(str(row[6].value or 0)))
        if account in account_rows:
            old_excel._link(
                book, "Dokumenter", f"H{row[0].row}", "Forløp", f"A{account_rows[account]}"
            )
    text_docs = {s.id: s.document_id for s in contents(report.basis).statements}
    first_text: dict[str, int] = {}
    for row in book["Omtale"].iter_rows(min_row=2):
        text_doc = text_docs.get(str(row[0].value))
        if text_doc in doc_rows:
            first_text.setdefault(text_doc, int(row[0].row or 0))
            old_excel._link(
                book, "Omtale", f"B{row[0].row}", "Dokumenter", f"C{doc_rows[text_doc]}"
            )
    for doc_id, row_number in doc_rows.items():
        if doc_id in first_text:
            old_excel._link(
                book, "Dokumenter", f"C{row_number}", "Omtale", f"B{first_text[doc_id]}"
            )
    for row in book["Kontroll"].iter_rows(min_row=2):
        if str(row[2].value) in doc_rows:
            old_excel._link(
                book, "Kontroll", f"C{row[0].row}", "Dokumenter", f"C{doc_rows[str(row[2].value)]}"
            )
    return book


def _package_profile(path: Path, sheet_count: int) -> None:
    """Bind alle relasjoner og deklarerte deltyper til P1s lokale eksportprofil."""
    office = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
    package = "http://schemas.openxmlformats.org/package/2006/relationships"
    content = "application/vnd.openxmlformats-"
    worksheets = {f"/xl/worksheets/sheet{n}.xml" for n in range(1, sheet_count + 1)}
    relationships = {
        "_rels/.rels": {
            (office + "officeDocument", "xl/workbook.xml"),
            (package + "/metadata/core-properties", "docProps/core.xml"),
            (office + "extended-properties", "docProps/app.xml"),
        },
        "xl/_rels/workbook.xml.rels": {
            (office + "styles", "styles.xml"),
            (office + "theme", "theme/theme1.xml"),
            *((office + "worksheet", name) for name in worksheets),
        },
    }
    overrides = {
        "/xl/workbook.xml": content + "officedocument.spreadsheetml.sheet.main+xml",
        "/xl/styles.xml": content + "officedocument.spreadsheetml.styles+xml",
        "/xl/theme/theme1.xml": content + "officedocument.theme+xml",
        "/docProps/core.xml": content + "package.core-properties+xml",
        "/docProps/app.xml": content + "officedocument.extended-properties+xml",
        **{name: content + "officedocument.spreadsheetml.worksheet+xml" for name in worksheets},
    }
    namespace = "{http://schemas.openxmlformats.org/package/2006/content-types}"
    declared = {
        (namespace + "Override", frozenset({"PartName": name, "ContentType": kind}.items()))
        for name, kind in overrides.items()
    } | {
        (
            namespace + "Default",
            frozenset({"Extension": extension, "ContentType": kind}.items()),
        )
        for extension, kind in (
            ("rels", content + "package.relationships+xml"),
            ("xml", "application/xml"),
        )
    }
    with ZipFile(path) as archive:
        if {name for name in archive.namelist() if name.endswith(".rels")} != set(relationships):
            raise ValueError("P1-Excel har uventede pakkrelasjoner")
        for name, wanted in relationships.items():
            root = _metadata(archive, name)
            identifiers = {element.get("Id", "") for element in root}
            actual = {(element.get("Type"), element.get("Target")) for element in root}
            if (
                root.tag != "{" + package + "}Relationships"
                or root.attrib
                or len(identifiers) != len(root)
                or "" in identifiers
                or len(root) != len(wanted)
                or actual != wanted
                or any(
                    element.tag != "{" + package + "}Relationship"
                    or set(element.attrib) != {"Id", "Type", "Target"}
                    or len(element)
                    for element in root
                )
            ):
                raise ValueError("P1-Excel har ugyldig lokal pakkrelasjon: " + name)
        types = _metadata(archive, "[Content_Types].xml")
        if (
            types.tag != namespace + "Types"
            or types.attrib
            or len(types) != len(declared)
            or {(e.tag, frozenset(e.attrib.items())) for e in types} != declared
            or any(len(e) for e in types)
            or any(name.lstrip("/") not in archive.namelist() for name in overrides)
        ):
            raise ValueError("P1-Excel har feil deklarerte innholdstyper")


def _style_profile(path: Path, expected: Workbook) -> None:
    """Avstem rå xf-flagg som openpyxls celleobjekter ikke bevarer."""
    # Lagring registrerer stiler i arkenes/cellenes rekkefølge. Registrer samme
    # profil uten å serialisere rapportens store rå-/offlineinnhold én gang til.
    for sheet in expected:
        for row in sheet:
            for cell in row:
                _ = cell.style_id
    wanted = etree.fromstring(tostring(write_stylesheet(expected)))
    with ZipFile(path) as archive:
        actual = _metadata(archive, "xl/styles.xml")
        if etree.tostring(actual, method="c14n") != etree.tostring(wanted, method="c14n"):
            raise ValueError("P1-Excel avviker i rå OOXML-stilprofil")


def read_excel(path: Path) -> Report:
    raw, shapes = inspect_excel(
        path,
        sheet_names=SHEET_NAMES,
        snapshot_name="Offlinegrunnlag",
        max_expanded_bytes=128 * 1024 * 1024,
    )
    _package_profile(path, len(shapes))
    report = report_from_json(raw)
    expected = _book(report, raw)
    try:
        if shapes != {s.title: SheetShape(s.max_row, s.max_column) for s in expected}:
            raise ValueError("P1-Excel har uventede ark/celleområder")
        _style_profile(path, expected)
        saved = load_workbook(path, data_only=False, keep_links=False)
        try:
            if saved.sheetnames != expected.sheetnames or saved.views != expected.views:
                raise ValueError("P1-Excel har endret arkrekkefølge/arbeidsbokvisning")
            for wanted_sheet, actual_sheet in zip(expected, saved, strict=True):
                for dimension in wanted_sheet.column_dimensions.values():
                    dimension.reindex()
                if (
                    actual_sheet.sheet_state != wanted_sheet.sheet_state
                    or actual_sheet.sheet_view != wanted_sheet.sheet_view
                    or actual_sheet.sheet_format != wanted_sheet.sheet_format
                    or actual_sheet.auto_filter != wanted_sheet.auto_filter
                    or actual_sheet.merged_cells != wanted_sheet.merged_cells
                    or set(actual_sheet.row_dimensions) != set(wanted_sheet.row_dimensions)
                    or set(actual_sheet.column_dimensions) != set(wanted_sheet.column_dimensions)
                    or any(
                        dict(actual_sheet.row_dimensions[k]) != dict(d)
                        for k, d in wanted_sheet.row_dimensions.items()
                    )
                    or any(
                        dict(actual_sheet.column_dimensions[k]) != dict(d)
                        for k, d in wanted_sheet.column_dimensions.items()
                    )
                ):
                    raise ValueError("P1-Excel har endret synlighet/oppsett: " + actual_sheet.title)
                for wanted_row, actual_row in zip(wanted_sheet, actual_sheet, strict=True):
                    for wanted, actual in zip(wanted_row, actual_row, strict=True):
                        value = None if wanted.value == "" else wanted.value
                        if (
                            actual.value != value
                            or type(actual.value) is not type(value)
                            or isinstance(value, str)
                            and actual.data_type != "s"
                            or actual.data_type == "f"
                            or actual.number_format != wanted.number_format
                            or copy(actual.font) != copy(wanted.font)
                            or copy(actual.fill) != copy(wanted.fill)
                            or copy(actual.border) != copy(wanted.border)
                            or copy(actual.alignment) != copy(wanted.alignment)
                            or copy(actual.protection) != copy(wanted.protection)
                            or (actual.hyperlink.location if actual.hyperlink else None)
                            != (wanted.hyperlink.location if wanted.hyperlink else None)
                            or (actual.hyperlink.target if actual.hyperlink else None)
                            != (wanted.hyperlink.target if wanted.hyperlink else None)
                        ):
                            raise ValueError(
                                "P1-Excel avviker i celle/stil/lenke: "
                                + actual_sheet.title
                                + "!"
                                + actual.coordinate
                            )
        finally:
            saved.close()
    finally:
        expected.close()
    return report


def export_excel(report: Report, output: Path) -> None:
    if output.exists() or output.suffix.lower() != ".xlsx":
        raise ValueError("Velg nytt P1-Excel-filnavn")
    raw = report_to_json(report)
    book = _book(report, raw)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        if read_excel(temporary) != report:
            raise ValueError("Skrevet P1-Excel avviker")
        _publish(temporary, output)
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def export_pair(report: Report, frozen: Path, excel: Path) -> None:
    if frozen.resolve() == excel.resolve() or frozen.exists() or excel.exists():
        raise ValueError("Velg to forskjellige nye filnavn")
    export_json(report, frozen)
    export_excel(report, excel)
    # Begge skrivere har allerede lukket og gjenlest sin fil mot samme rapport.


def export_originals(report: Report, directory: Path) -> None:
    from innlesning.komitepilot import payload_bytes

    report_to_json(report)
    directory.mkdir(parents=True, exist_ok=False)
    old = report.basis.t4_basis
    legacy = legacy_payloads(old.legacy_json)
    previous = legacy_report(old.legacy_json)
    sources = shared_sources(old, legacy, previous.basis if previous else None)
    by_digest = {hashlib.sha256(raw).hexdigest(): raw for _, raw, _ in sources.values()}
    by_digest.update((p.id, payload_bytes(p, legacy)) for p in report.basis.payloads)
    for digest, raw in by_digest.items():
        (directory / (digest + ".bin")).write_bytes(raw)
    (directory / "index.json").write_text(
        encode(
            {
                "sources": [asdict(s) for s in report.basis.sources],
                "reused_sources": {name: sid for name, (sid, _, _) in sources.items()},
                "payloads": {digest: len(raw) for digest, raw in by_digest.items()},
                "evidence": [asdict(e) for e in report.basis.evidence],
            }
        ),
        encoding="utf8",
    )
