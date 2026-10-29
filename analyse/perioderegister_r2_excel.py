"""P2: kilde-/rollebundet Excel med én selvforsynt kapsel og P1s defensive profil."""

import hashlib
import json
from copy import copy
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from innlesning.komitepilot import payload_bytes
from innlesning.observations import encode
from innlesning.perioderegister_model import Report as P1Report
from innlesning.perioderegister_r2 import shared_sources
from innlesning.perioderegister_r2_model import Report
from innlesning.visningspilot_model import PilotSelection

from . import fleraarspilot_excel as old_excel
from . import perioderegister_excel as p1_excel
from .extraction_excel import Value, write_rows
from .komitepilot import legacy_payloads, legacy_report
from .omtale_excel_guard import SheetShape, inspect_excel
from .perioderegister_r2 import contents, export_json, report_from_json, report_to_json
from .saksdokumenter import _publish

SHEET_NAMES = p1_excel.SHEET_NAMES
_chunks = p1_excel._chunks
_package_profile = p1_excel._package_profile
_style_profile = p1_excel._style_profile


def _tables(report: Report, frozen: str) -> dict[str, list[list[Value]]]:
    old_selection = report.selection
    if old_selection.mode == "document" and old_selection.document_id in {
        d.id for d in report.basis.documents
    }:
        old_selection = PilotSelection()
    old_report = P1Report(
        "", report.basis.p1_basis, old_selection, None, report.view.p1, report.calculations
    )
    tables = p1_excel._tables(old_report, "")
    basis = report.basis
    docs = {d.id: d for d in contents(basis).documents}
    proofs = {e.id: e for e in basis.evidence}
    tables["Forløp"][0].extend(
        [
            "R2-funn",
            "Forskningsverdi",
            "Verdiens rolle/enhet",
            "Betinget konto",
            "Overføringsmål",
            "Kildebelegg-ID-er",
            "Representasjon SHA-256",
            "Bevisgrense/råverdi",
        ]
    )
    for row in tables["Forløp"][1:]:
        row.extend([None] * 8)
    for f in basis.findings:
        if f.id not in report.view.finding_ids:
            continue
        doc = docs[f.document_id]
        e = proofs[f.evidence_ids[0]]
        value = None
        meaning = f.role
        if f.key == "FIN-1429-1":
            value = json.loads(f.amounts_json)["source_observed_compensation_estimate"]
            meaning = "foreslått kompensasjonsanslag, kroner"
            e = next(e for e in basis.evidence if e.key == "R28")
        elif f.key == "D230":
            value = json.loads(f.amounts_json)["delta"]
            meaning = "vedtatt endring, annen konto 1560/71, kroner; kildekontekst"
        elif f.key == "O1-OMDISPONERING":
            meaning = "omdisponering /22 til /1; avrunding/lønnsrest gjelder bare /1"
        tables["Forløp"].append(
            [
                ", ".join(f.accounts),
                f.budget_year,
                doc.title,
                doc.date,
                "R2 / " + f.role,
                f.actor,
                None,
                None,
                None,
                None,
                f.status + "; " + f.reason,
                "[]",
                f.id,
                doc.id,
                f.post_link,
                e.source_id,
                e.locator + "/td[4]" if f.key == "FIN-1429-1" else e.locator,
                None,
                f.key,
                value,
                meaning,
                encode(f.conditional_accounts),
                f.transfer_target_year,
                encode(f.evidence_ids),
                e.representation_sha256,
                f.amounts_json,
            ]
        )
    for row in tables["Dokumenter"][1:]:
        if row[0] == "Dokument":
            row[16] = "ja" if row[1] in report.view.document_ids else "nei"
    for d in basis.documents:
        tables["Dokumenter"].append(
            [
                "Dokument",
                d.id,
                d.title,
                d.kind,
                d.date,
                d.session,
                None,
                None,
                None,
                d.raw_coverage,
                d.structured_coverage,
                d.professional_status,
                encode(d.accounts),
                encode(d.source_ids),
                None,
                None,
                "ja" if d.id in report.view.document_ids else "nei",
            ]
        )
    for item in basis.inventory:
        d = docs[item.document_id]
        for n, chunk in enumerate(_chunks(item.text), 1):
            tables["Dokumenter"].append(
                [
                    "R2-inventar",
                    item.id,
                    d.title,
                    item.role,
                    d.date,
                    d.session,
                    item.budget_year,
                    item.chapter,
                    item.post,
                    "R2: ekstra rårepresentasjon/struktur, P1-dokumentpost uendret",
                    item.structure_status + "; tekstdel " + str(n),
                    d.professional_status,
                    encode(item.observation_ids),
                    item.source_id,
                    item.locator,
                    chunk,
                    "ja" if d.id in report.view.document_ids else "nei",
                ]
            )
    for s in basis.statements:
        for n, text in enumerate(_chunks(s.text), 1):
            tables["Omtale"].append(
                [
                    s.id,
                    docs[s.document_id].title,
                    s.role,
                    s.actor,
                    s.period,
                    encode(s.accounts),
                    encode(s.hierarchy),
                    n,
                    text,
                    "",
                    s.source_id,
                    s.locator,
                    s.status,
                    "ja" if s.document_id in report.view.document_ids else "nei",
                ]
            )
    coverage = tables["Dekning og avvik"]
    replacements: dict[str, Value] = {
        "Rapport-ID": report.id,
        "Grunnlags-ID": basis.id,
        "Format": report.version,
        "Valgt inngang": encode(asdict(report.selection)),
        "Returvalg": encode(asdict(report.return_selection)) if report.return_selection else None,
    }
    for row in coverage[1:]:
        if str(row[0]) in replacements:
            row[1] = replacements[str(row[0])]
    coverage.extend(
        [
            ["Historisk P1-basis", basis.p1_basis.id],
            ["R2-tillegg", "Tre avgrensede grener; ingen nye monetære hendelser"],
            [
                "Historisk indeksgrense",
                "67 indeksforekomster / 66 saker / 60 åpne / 25 vedtakskandidater; "
                "2 158 uten ordbokstreff",
            ],
            [
                "Faktisk anvendelse 231/429",
                "uavklart; 2025 er overføringsmål, ingen ny bevilgning/kontinuitet",
            ],
        ]
    )
    coverage.extend(
        [
            s.id + " / separat status for sak " + s.case_id,
            s.status + ": " + s.reason + "; kandidat " + s.case_candidate_id,
        ]
        for s in basis.status_updates
    )
    coverage.extend(
        [c.key + " " + c.target, c.status + ": " + c.result] for c in basis.cross_references
    )
    coverage.extend([issue, "P2-kildekontroll"] for issue in basis.issues)
    for source_item in basis.sources:
        tables["Kontroll"].append(
            [
                "R2-kilde",
                source_item.id,
                source_item.document_key,
                source_item.payload_id,
                source_item.media_type,
                None,
                source_item.status,
                source_item.expected_sha256,
                source_item.name,
                None,
                source_item.metadata_json,
            ]
        )
    for e in basis.evidence:
        for n, text in enumerate(_chunks(e.text), 1):
            tables["Kontroll"].append(
                [
                    "R2-belegg",
                    e.id,
                    e.key,
                    e.source_id,
                    e.role,
                    n,
                    e.kind,
                    e.representation_sha256,
                    e.locator,
                    e.nested_html_field_xpath,
                    encode({"text": text, "rows_json": e.rows_json}),
                ]
            )
    for kind, items in (
        ("R2-funn", basis.findings),
        ("R2-relasjon", basis.relations),
        ("R2-status", basis.status_updates),
        ("R2-kryssreferanse", basis.cross_references),
    ):
        for control_item in items:
            for n, text in enumerate(_chunks(encode(asdict(control_item))), 1):
                tables["Kontroll"].append(
                    [kind, control_item.id, None, None, None, n, None, None, None, None, text]
                )
    tables["Offlinegrunnlag"] = [["Del", "Frosset perioderegister-r2-1 med én uendret P1-basis"]]
    tables["Offlinegrunnlag"].extend([n, chunk] for n, chunk in enumerate(_chunks(frozen), 1))
    return tables


def _book(report: Report, raw: str) -> Workbook:
    book = Workbook()
    assert book.active is not None
    book.remove(book.active)
    widths = {
        "Forløp": (
            18,
            14,
            55,
            14,
            42,
            22,
            20,
            20,
            20,
            20,
            48,
            24,
            24,
            24,
            24,
            30,
            40,
            24,
            24,
            20,
            55,
            25,
            16,
            40,
            40,
            100,
        ),
        "Dokumenter": (20, 24, 65, 35, 14, 16, 14, 12, 12, 45, 55, 30, 28, 26, 45, 95, 12),
        "Omtale": (24, 65, 42, 22, 20, 26, 30, 10, 100, 30, 26, 45, 35, 12),
        "Sammenligninger": (16, 24, 55, 55, 22, 22, 24, 24, 95, 28, 26, 26, 26, 65),
        "Dekning og avvik": (75, 120),
        "Kontroll": (22, 26, 26, 26, 42, 28, 35, 28, 70, 22, 100),
        "Offlinegrunnlag": (12, 120),
    }
    for name, rows in _tables(report, raw).items():
        if any(isinstance(value, str) and len(value) > 30000 for row in rows for value in row):
            raise ValueError("P1-visning krever eksplisitt tapsfri tekstdeling")
        write_rows(book, name, rows)
        sheet = book[name]
        for n, width in enumerate(widths[name], 1):
            sheet.column_dimensions[get_column_letter(n)].width = width
        for row in sheet:
            sheet.row_dimensions[int(row[0].row or 0)].height = (
                44
                if row[0].row == 1
                else (72 if name in {"Forløp", "Dokumenter", "Dekning og avvik"} else 30)
            )
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if (
                    sheet.cell(1, int(cell.column or 0)).value == "Dato"
                    and isinstance(cell.value, str)
                    and int(cell.row or 0) > 1
                    and cell.value
                ):
                    cell.value = datetime.fromisoformat(cell.value)
                    cell.number_format = 'dd"."mm"."yyyy'
                if type(cell.value) is int:
                    cell.number_format = (
                        "0"
                        if sheet.cell(1, int(cell.column or 0)).value
                        in {"Budsjettår", "Overføringsmål", "Del"}
                        else "#,##0"
                    )
    doc_rows = {
        str(r[1].value): r[0].row
        for r in book["Dokumenter"].iter_rows(min_row=2)
        if r[0].value == "Dokument"
    }
    account_rows: dict[tuple[str, int], int] = {}
    for row in book["Forløp"].iter_rows(min_row=2):
        doc = str(row[13].value)
        if doc in doc_rows:
            old_excel._link(book, "Forløp", f"C{row[0].row}", "Dokumenter", f"C{doc_rows[doc]}")
        if type(row[1].value) is int and str(row[0].value) in {"1429/1", "1429/22"}:
            account_rows.setdefault((str(row[0].value), row[1].value), int(row[0].row or 0))
    for row in book["Dokumenter"].iter_rows(min_row=2):
        account = (str(row[7].value) + "/" + str(row[8].value), int(str(row[6].value or 0)))
        if account in account_rows:
            old_excel._link(
                book, "Dokumenter", f"H{row[0].row}", "Forløp", f"A{account_rows[account]}"
            )
    text_docs = {s.id: s.document_id for s in contents(report.basis).statements}
    first_text: dict[str, int] = {}
    for row in book["Omtale"].iter_rows(min_row=2):
        text_doc = text_docs.get(str(row[0].value))
        if text_doc in doc_rows:
            first_text.setdefault(text_doc, int(row[0].row or 0))
            old_excel._link(
                book, "Omtale", f"B{row[0].row}", "Dokumenter", f"C{doc_rows[text_doc]}"
            )
    for doc_id, row_number in doc_rows.items():
        if doc_id in first_text:
            old_excel._link(
                book, "Dokumenter", f"C{row_number}", "Omtale", f"B{first_text[doc_id]}"
            )
    for row in book["Kontroll"].iter_rows(min_row=2):
        if str(row[2].value) in doc_rows:
            old_excel._link(
                book, "Kontroll", f"C{row[0].row}", "Dokumenter", f"C{doc_rows[str(row[2].value)]}"
            )
    return book


def read_excel(path: Path) -> Report:
    raw, shapes = inspect_excel(
        path,
        sheet_names=SHEET_NAMES,
        snapshot_name="Offlinegrunnlag",
        max_expanded_bytes=128 * 1024 * 1024,
    )
    _package_profile(path, len(shapes))
    report = report_from_json(raw)
    expected = _book(report, raw)
    try:
        if shapes != {s.title: SheetShape(s.max_row, s.max_column) for s in expected}:
            raise ValueError("P1-Excel har uventede ark/celleområder")
        _style_profile(path, expected)
        saved = load_workbook(path, data_only=False, keep_links=False)
        try:
            if saved.sheetnames != expected.sheetnames or saved.views != expected.views:
                raise ValueError("P1-Excel har endret arkrekkefølge/arbeidsbokvisning")
            for wanted_sheet, actual_sheet in zip(expected, saved, strict=True):
                for dimension in wanted_sheet.column_dimensions.values():
                    dimension.reindex()
                if (
                    actual_sheet.sheet_state != wanted_sheet.sheet_state
                    or actual_sheet.sheet_view != wanted_sheet.sheet_view
                    or actual_sheet.sheet_format != wanted_sheet.sheet_format
                    or actual_sheet.auto_filter != wanted_sheet.auto_filter
                    or actual_sheet.merged_cells != wanted_sheet.merged_cells
                    or set(actual_sheet.row_dimensions) != set(wanted_sheet.row_dimensions)
                    or set(actual_sheet.column_dimensions) != set(wanted_sheet.column_dimensions)
                    or any(
                        dict(actual_sheet.row_dimensions[k]) != dict(d)
                        for k, d in wanted_sheet.row_dimensions.items()
                    )
                    or any(
                        dict(actual_sheet.column_dimensions[k]) != dict(d)
                        for k, d in wanted_sheet.column_dimensions.items()
                    )
                ):
                    raise ValueError("P1-Excel har endret synlighet/oppsett: " + actual_sheet.title)
                for wanted_row, actual_row in zip(wanted_sheet, actual_sheet, strict=True):
                    for wanted, actual in zip(wanted_row, actual_row, strict=True):
                        value = None if wanted.value == "" else wanted.value
                        if (
                            actual.value != value
                            or type(actual.value) is not type(value)
                            or isinstance(value, str)
                            and actual.data_type != "s"
                            or actual.data_type == "f"
                            or actual.number_format != wanted.number_format
                            or copy(actual.font) != copy(wanted.font)
                            or copy(actual.fill) != copy(wanted.fill)
                            or copy(actual.border) != copy(wanted.border)
                            or copy(actual.alignment) != copy(wanted.alignment)
                            or copy(actual.protection) != copy(wanted.protection)
                            or (actual.hyperlink.location if actual.hyperlink else None)
                            != (wanted.hyperlink.location if wanted.hyperlink else None)
                            or (actual.hyperlink.target if actual.hyperlink else None)
                            != (wanted.hyperlink.target if wanted.hyperlink else None)
                        ):
                            raise ValueError(
                                "P1-Excel avviker i celle/stil/lenke: "
                                + actual_sheet.title
                                + "!"
                                + actual.coordinate
                            )
        finally:
            saved.close()
    finally:
        expected.close()
    return report


def export_excel(report: Report, output: Path) -> None:
    if output.exists() or output.suffix.lower() != ".xlsx":
        raise ValueError("Velg nytt P1-Excel-filnavn")
    raw = report_to_json(report)
    book = _book(report, raw)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        if read_excel(temporary) != report:
            raise ValueError("Skrevet P1-Excel avviker")
        _publish(temporary, output)
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def export_pair(report: Report, frozen: Path, excel: Path) -> None:
    if frozen.resolve() == excel.resolve() or frozen.exists() or excel.exists():
        raise ValueError("Velg to forskjellige nye filnavn")
    export_json(report, frozen)
    export_excel(report, excel)
    # Begge skrivere har allerede lukket og gjenlest sin fil mot samme rapport.


def export_originals(report: Report, directory: Path) -> None:
    report_to_json(report)
    directory.mkdir(parents=True, exist_ok=False)
    old = report.basis.p1_basis.t4_basis
    legacy = legacy_payloads(old.legacy_json)
    previous = legacy_report(old.legacy_json)
    shared = shared_sources(report.basis.p1_basis, legacy, previous.basis if previous else None)
    by_digest = {hashlib.sha256(raw).hexdigest(): raw for _, raw, _ in shared.values()}
    by_digest.update((p.id, payload_bytes(p, legacy)) for p in report.basis.payloads)
    for digest, raw in by_digest.items():
        (directory / (digest + ".bin")).write_bytes(raw)
    (directory / "index.json").write_text(
        encode(
            {
                "sources": [asdict(s) for s in report.basis.sources],
                "reused_sources": {name: sid for name, (sid, _, _) in shared.items()},
                "payloads": {digest: len(raw) for digest, raw in by_digest.items()},
                "evidence": [asdict(e) for e in report.basis.evidence],
            }
        ),
        encoding="utf8",
    )
