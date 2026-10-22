"""B9s leseflate fra frosne kildeobservasjoner; ingen lagring eller ny tolkning."""

import math
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from typing import TYPE_CHECKING

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from innlesning.saksdokumenter import CaseEnrichment, Publication

from .excel import put
from .extraction_excel import verify_continued_text
from .session import SessionRow
from .session_excel import write_rows

if TYPE_CHECKING:
    from .saksdokumenter import SaksReport

HEADERS = (
    "Kapittelnummer",
    "Postnummer",
    "Kapittelnavn",
    "Postnavn",
    "Stikkord",
    "Dato",
    "Vedtaksnummer",
    "Innstillingsnummer",
    "Proposisjon",
    "Beløp før vedtak",
    "Vedtatt endring",
    "Beløp etter vedtak",
    "Kontrollstatus",
)

# Bare eksakte, dokumenterte handlingssuffikser fjernes fra navnet.
ACTION = re.compile(
    r",\s*(?:økes|aukes|auka|reduseres|redusert|reduserast|minka|bevilges)\s+med\s*$",
    re.I,
)
ISSUES = {
    "manglende_start": "Mangler historisk startgrunnlag",
    "mangler_organisasjonsoppslag": "Mangler organisasjonsoppslag",
    "mangler_valgt_kontrollresultat": "Mangler valgt kontrollresultat",
    "ikke_godkjent_analysebelop": "Ikke faglig godkjent analysebeløp",
    "valgt_delinventar": "Valgt delinventar",
    "mangler_belopsuttrekk": "Mangler beløpsuttrekk",
    "saksberikelse_ikke_valgt": "Saksberikelse er ikke valgt",
    "ingen_tolket_saksrespons": "Ingen tolket saksrespons",
    "flere_innstillinger": "Flere innstillinger på saksnivå",
    "flere_proposisjoner": "Flere proposisjoner på saksnivå",
    "uavklart_stikkord": "Uavklart stikkord",
    "gjentatt_belopsrolle": "Flere kildebeløp med samme betydning",
    "flere_startniva": "Flere inngående nivåbeløp",
    "uavklart_rekkefolge": "Uavklart rekkefølge",
    "mangler_delta": "Mangler endringsbeløp",
    "kjedekonflikt": "Konflikt i historisk kjede",
    "lokal_konflikt": "Før, endring og etter stemmer ikke lokalt",
    "etter_konflikt": "Etterbeløpet avviker fra kjeden",
    "nyere_kildeforsok": "Nyere kildeforsøk finnes",
    "historisk_kilderevisjon": "Historisk kilderevisjon",
    "flere_valgte_revisjoner": "Flere valgte revisjoner av samme vedtak",
    "mangler_i_valgt_inventar": "Kjent vedtak mangler i valgt inventar",
    "kjent_id_utenfor_valgt_inventar": "Kjent vedtak utenfor valgt inventar",
}


def _issue_label(issue: str) -> str:
    code, _, detail = issue.partition(":")
    if re.fullmatch(r"[a-f0-9-]{36}", code) and detail:
        return _issue_label(detail)
    if code in ISSUES:
        return ISSUES[code]
    # Også nye koder og deres kildebevis skal være lesbare og synlige.
    return "Avvik: " + issue.replace("_", " ")


def _names(values: list[str | None], label: str, issues: list[str]) -> str:
    names = tuple(dict.fromkeys(v.strip() for v in values if v and v.strip()))
    if len(names) == 1 and all(v and v.strip() for v in values):
        return names[0]
    reason = "motstridende" if len(names) > 1 else "manglende"
    issues.append(f"Uavklart {label.lower()}: {reason} kildenavn")
    return "Uavklart" + (": " + " / ".join(names) if names else ": mangler i kilden")


def _post_name(
    value: str | None, enrichment: CaseEnrichment | None, chapter: str | None, post: str | None
) -> str | None:
    if value is None:
        return None
    name = ACTION.sub("", value).strip()
    # Fjern bare kjente stikkord fra den eksakte bevarte formålsraden.
    candidates = (
        [
            k
            for k in enrichment.keywords
            if (k.raw_text, k.chapter, k.post) == (value, chapter, post)
        ]
        if enrichment
        else []
    )
    if len(candidates) == 1:
        for word in candidates[0].words:
            name = re.sub(r",\s*" + re.escape(word) + r"(?=\s*(?:,|$))", "", name)
    return name.strip(" ,;") or None


def _publications(enrichment: CaseEnrichment | None, kind: str) -> tuple[Publication, ...]:
    if not enrichment or not enrichment.content:
        return ()
    return tuple(p for p in enrichment.content.publications if p.kind == kind)


def _document_label(publications: tuple[Publication, ...]) -> str:
    if not publications:
        return "Uavklart; se Dokumentavvik"
    labels = []
    for p in publications:
        raw = p.fields.get("lenke_tekst")
        label = raw.text if raw and raw.text else "Uavklart dokumentnummer"
        labels.append(
            label + (f" [{p.identity_status}]" if p.identity_status != "identifisert" else "")
        )
    return "\n".join(labels) + ("\nSe alle lenker i Saksdokumenter" if len(labels) > 1 else "")


@dataclass(frozen=True)
class HistoryRow:
    values: tuple[str | int | date | None, ...]
    publications: tuple[tuple[Publication, ...], tuple[Publication, ...]]


def _event_key(row: SessionRow) -> tuple[object, ...]:
    o = row.observation
    return (o.api_id, o.decision_session, o.budget_year, o.chapter, o.post, o.step_id)


def _freshness(decision: str, documents: str) -> str:
    return (
        "Ferskhet i fryst lokal vurdering: vedtak "
        + decision.replace("_", " ")
        + ", saksgrunnlag "
        + documents.replace("_", " ")
    )


def _rows(report: "SaksReport") -> list[HistoryRow]:
    contexts = {c.decision_revision_id: c for c in report.contexts}
    enrichments = {e.id: e for e in report.enrichments}
    groups: dict[tuple[object, ...], list[SessionRow]] = defaultdict(list)
    for row in report.base.rows:
        o = row.observation
        # Ulike tolkninger, revisjoner, omfang og komponenter er aldri én hendelse.
        key = (
            o.source_system,
            o.source_unit_id,
            o.source_revision_id,
            o.interpretation_id,
            o.step_id,
            o.budget_year,
            o.chapter,
            o.post,
            o.subpost,
            o.detail_level,
            o.component_of,
            o.currency,
            o.multiplier,
            o.scope,
            o.figure_type,
            o.information_type,
            o.decision_date,
            o.decision_session,
            o.api_id,
            o.round,
        )
        groups[key].append(row)
    global_issues = tuple(
        dict.fromkeys(_issue_label(i) for i in (*report.base.issues, *report.issues))
    )
    alternatives = Counter(_event_key(group[0]) for group in groups.values())
    result = []
    for group in sorted(
        groups.values(),
        key=lambda g: (
            g[0].observation.decision_date or "9999",
            str(g[0].observation.api_id),
            str(g[0].observation.chapter),
            str(g[0].observation.post),
            g[0].observation.source_revision_id,
            g[0].observation.interpretation_id,
            g[0].observation.step_id,
        ),
    ):
        o = group[0].observation
        context = contexts[o.source_revision_id]
        enrichment = enrichments.get(context.enrichment_id or "")
        local = list(dict.fromkeys(_issue_label(i) for r in group for i in r.issues))
        local.extend(_issue_label(i) for i in context.issues)
        if alternatives[_event_key(group[0])] > 1:
            local.append("Flere adskilte kildegrunnlag for samme konto og ledd; ikke summer radene")
        chapter_name = _names([r.observation.chapter_name for r in group], "Kapittelnavn", local)
        post_name = _names(
            [_post_name(r.observation.purpose, enrichment, o.chapter, o.post) for r in group],
            "Postnavn",
            local,
        )
        keywords = (
            [k for k in enrichment.keywords if (k.chapter, k.post) == (o.chapter, o.post)]
            if enrichment
            else []
        )
        if len(keywords) == 1:
            k = keywords[0]
            keyword = ", ".join(k.words) or "Ingen stikkord funnet"
            matching_source = all(
                (r.observation.purpose, r.observation.pointer.artifact_id)
                == (k.raw_text, k.source_sha256)
                for r in group
            )
            if k.status == "uavklart" or not matching_source:
                reason = (
                    k.reason if matching_source else "stikkordgrunnlaget avviker fra formålsraden"
                )
                keyword = (", ".join(k.words) + "; " if k.words else "") + "Uavklart: " + reason
                local.append("Uavklart stikkord og skille mot postnavn")
        else:
            keyword = "Uavklart: manglende eller flertydig stikkordtolkning"
            local.append("Uavklart stikkord og skille mot postnavn")
        amounts: list[int | None] = []
        for meaning, label in (
            ("for_belop", "førbeløp"),
            ("endring", "endringsbeløp"),
            ("etter_belop", "etterbeløp"),
        ):
            candidates = [r for r in group if r.observation.meaning == meaning]
            amounts.append(candidates[0].observation.amount_kr if len(candidates) == 1 else None)
            if len(candidates) > 1:
                local.append(f"Flere {label}; ingen verdi valgt, se Historikkdetaljer")
            elif not candidates or candidates[0].observation.amount_kr is None:
                local.append(f"Mangler entydig kildeobservert {label}")
        other = [
            r.observation
            for r in group
            if r.observation.meaning not in {"for_belop", "endring", "etter_belop"}
        ]
        for source in other:
            amount = source.amount_kr if source.amount_kr is not None else source.raw_value
            local.append(
                f"Annen beløpsbetydning ({source.meaning}): {amount}; se Historikkdetaljer"
            )
        if o.detail_level != "post" or o.component_of or o.subpost:
            local.append(
                f"Eget detaljeringsnivå: {o.detail_level}, "
                f"underpost {o.subpost or 'ikke oppgitt'}, "
                f"komponent av {o.component_of or 'ikke oppgitt'}; ikke posttotal"
            )
        if (
            o.scope != "enkeltledd"
            or o.currency != "NOK"
            or o.multiplier != "1"
            or o.figure_type != "stortingsvedtak"
            or o.information_type != "bevilgning"
        ):
            local.append(
                f"Særskilt eller uavklart omfang: {o.scope}, {o.currency}, faktor {o.multiplier}, "
                f"{o.figure_type}, {o.information_type}; ikke additivt avklart"
            )
        observed_date: date | str | None = o.decision_date
        try:
            observed_date = date.fromisoformat(o.decision_date or "")
        except ValueError:
            local.append("Mangler entydig vedtaksdato")
        approved = sum(r.approved_amount_kr is not None for r in group)
        status = [
            "Kildebeløp. "
            + (
                "Ikke faglig godkjent"
                if not approved
                else f"Faglig kontroll: {approved} av {len(group)} enkeltbeløp; "
                "se Beløpsobservasjoner"
            )
        ]
        status.extend(dict.fromkeys(local))
        status.append(
            "Kontrollgrunnlagets ferskhet: "
            + ", ".join(dict.fromkeys(r.freshness.replace("_", " ") for r in group))
        )
        status.append(_freshness(context.decision_freshness, context.source_freshness))
        status.append("Saksrelasjon: postvis forslagstilknytning er ikke bevist")
        remainder = [i for i in global_issues if i not in local]
        status.append(
            "Globalt før filtre: " + (". ".join(remainder) if remainder else "se Avvik og Dekning")
        )
        status.append("Ingen uttømmende sesjons- eller årsdekning. Se Historikkdetaljer og Avvik")
        pubs = (_publications(enrichment, "innstilling"), _publications(enrichment, "proposisjon"))
        result.append(
            HistoryRow(
                (
                    o.chapter,
                    o.post,
                    chapter_name,
                    post_name,
                    keyword,
                    observed_date,
                    f"{context.formal_number} ({context.session})",
                    _document_label(pubs[0]),
                    _document_label(pubs[1]),
                    *amounts,
                    ". ".join(status),
                ),
                pubs,
            )
        )
    if not result:
        empty_status = "Ingen beløpsrader i utvalget. Globalt før filtre: " + ". ".join(
            global_issues
        )
        empty_status += ". " + ". ".join(
            dict.fromkeys(
                _freshness(c.decision_freshness, c.source_freshness) for c in report.contexts
            )
        )
        result.append(
            HistoryRow(
                (None,) * 12
                + (
                    empty_status + ". Se Avvik, Dokumentavvik og Dekning. Ingen årsdekning erklært",
                ),
                ((), ()),
            )
        )
    return result


def write_history_view(book: Workbook, report: "SaksReport") -> None:
    """Behold den tidligere detaljvisningen og legg den nye hovedvisningen først."""
    book["Historikk"].title = "Historikkdetaljer"
    if "Historikk fortsettelse" in book:
        continuation = book["Historikk fortsettelse"]
        continuation.title = "Historikkdetaljer fortsettelse"
        for row in continuation.iter_rows(min_row=2):
            put(row[0], str(row[0].value).replace("Historikk!", "Historikkdetaljer!", 1))
    rows = _rows(report)
    # write_rows bevarer lang tekst med pekere i en separat videreføringsfane.
    write_rows(
        book,
        "Historikk",
        [list(HEADERS)]
        + [[v if not isinstance(v, date) else v.isoformat() for v in r.values] for r in rows],
    )
    sheet = book["Historikk"]
    book.move_sheet(sheet, offset=-book.index(sheet))
    book.active = sheet
    sheet.sheet_properties.tabColor = "181C62"
    sheet.freeze_panes = "C2"
    widths = (16, 13, 26, 46, 30, 15, 24, 32, 32, 24, 24, 24, 88)
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.cell(1, column).alignment = Alignment(
            wrap_text=True, vertical="center", horizontal="center"
        )
    sheet.row_dimensions[1].height = 36
    for number, projected in enumerate(rows, 2):
        lines = 1
        for column, value in enumerate(projected.values, 1):
            cell = sheet.cell(number, column)
            if isinstance(value, date):
                put(cell, value)
                cell.number_format = "dd.mm.yyyy"
            if column in (10, 11, 12):
                cell.number_format = '#,##0" kr"'
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top",
                horizontal="right" if isinstance(value, int) else "left",
            )
            if number % 2:
                cell.fill = PatternFill("solid", fgColor="EDEDEE")
            lines = max(
                lines,
                sum(
                    max(1, math.ceil(len(s) / (widths[column - 1] - 3)))
                    for s in str(value or "").split("\n")
                ),
            )
        sheet.row_dimensions[number].height = min(409, max(48, lines * 14 + 12))
        for column, publications in zip((8, 9), projected.publications, strict=True):
            if not publications:
                continue
            cell = sheet.cell(number, column)
            if len(publications) == 1 and publications[0].url:
                cell.hyperlink = Hyperlink(ref=cell.coordinate, target=publications[0].url)
            else:
                first = next(
                    n
                    for n, r in enumerate(book["Saksdokumenter"].iter_rows(min_row=2), 2)
                    if r[3].value == publications[0].id
                )
                cell.hyperlink = Hyperlink(
                    ref=cell.coordinate, location=f"'Saksdokumenter'!E{first}"
                )
            cell.font = Font(name="Arial", size=10, color="4156A6", underline="single")
        sheet.cell(number, 13).hyperlink = Hyperlink(ref=f"M{number}", location="'Avvik'!A1")


def verify_history_view(book: Workbook, report: "SaksReport") -> None:
    sheet = book["Historikk"]
    expected = _rows(report)
    if (
        book.sheetnames[0] != "Historikk"
        or book.active != sheet
        or tuple(c.value for c in sheet[1]) != HEADERS
        or sheet.max_row != len(expected) + 1
    ):
        raise ValueError("Historikk har feil hovedvisning eller kolonner/rader")
    expected_book = Workbook()
    try:
        write_rows(
            expected_book,
            "Historikk",
            [list(HEADERS)]
            + [
                [v if not isinstance(v, date) else v.isoformat() for v in row.values]
                for row in expected
            ],
        )
        full_text = verify_continued_text(sheet, expected_book["Historikk"])
    finally:
        expected_book.close()
    for number, row in enumerate(expected, 2):
        for column, value in enumerate(row.values, 1):
            cell = sheet.cell(number, column)
            actual = full_text.get(cell.coordinate, cell.value)
            if isinstance(actual, datetime):
                actual = actual.date()
            if actual != value or (isinstance(value, str) and cell.data_type != "s"):
                raise ValueError("Historikk avviker fra frosne kildeobservasjoner")
            if isinstance(value, date) and (
                not isinstance(cell.value, date) or cell.number_format != "dd.mm.yyyy"
            ):
                raise ValueError("Historikk mangler datocelle med norsk format")
            if column in (10, 11, 12) and (
                cell.number_format != '#,##0" kr"'
                or (value is not None and type(actual) is not int)
            ):
                raise ValueError("Historikk mangler numeriske kronebeløp")
        for column, publications in zip((8, 9), row.publications, strict=True):
            if not publications:
                continue
            link = sheet.cell(number, column).hyperlink
            if len(publications) == 1 and publications[0].url:
                if link is None or link.target != publications[0].url:
                    raise ValueError("Historikk har feil dokumentlenke")
            else:
                first = next(
                    n
                    for n, r in enumerate(book["Saksdokumenter"].iter_rows(min_row=2), 2)
                    if r[3].value == publications[0].id
                )
                if link is None or link.location != f"'Saksdokumenter'!E{first}" or link.target:
                    raise ValueError("Historikk mangler intern lenke til alle dokumentreferansene")
