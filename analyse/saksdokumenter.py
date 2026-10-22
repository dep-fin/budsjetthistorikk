"""Opt-in B8-sammenstilling og frosset v2. B6/B7s standardformat er uendret."""

import json
import os
import re
import sqlite3
from dataclasses import asdict, dataclass, replace
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet

from innlesning.observations import encode, identity, read_extraction
from innlesning.saksdokumenter import (
    VERSION,
    CaseContext,
    CaseEnrichment,
    Publication,
    read_case_context,
    read_enrichment,
    safe_document_url,
)
from innlesning.sync import read_revision

from .excel import put
from .extraction_excel import SourceTextSnapshot, Value, verify_continued_text
from .saks_history_excel import verify_history_view, write_history_view
from .session import (
    SessionReport,
    SessionSelection,
    _ids,
    build_session_report,
    selection_from_dict,
)
from .session_excel import _history_rows, _write, write_rows
from .session_serialization import MAX_REPORT_BYTES, _constant, _decode, _object
from .session_serialization import report_from_json as base_from_json
from .session_serialization import report_to_json as base_to_json


@dataclass(frozen=True)
class SaksSelection:
    base: SessionSelection
    enrichment_ids: tuple[str, ...] = ()

    def __post_init__(self) -> None:
        if not isinstance(self.base, SessionSelection):
            raise ValueError("Saksutvalg krever et eksplisitt B6-utvalg")
        _ids(self.enrichment_ids, "enrichment_ids")


def saks_selection_from_dict(data: dict[str, Any]) -> SaksSelection:
    if (
        set(data) != {"base", "enrichment_ids"}
        or not isinstance(data["base"], dict)
        or not isinstance(data["enrichment_ids"], list)
    ):
        raise ValueError("Saksutvalg krever base-objekt og enrichment_ids-liste")
    return SaksSelection(selection_from_dict(data["base"]), tuple(data["enrichment_ids"]))


@dataclass(frozen=True)
class DocumentContext:
    decision_revision_id: str
    case_id: str
    formal_number: str
    session: str
    enrichment_id: str | None
    selected_attempt_id: str | None
    case_context: CaseContext
    source_freshness: str
    decision_freshness: str
    issues: tuple[str, ...]
    relation_level: str = "sak; postvis forslagstilknytning er ikke bevist"


@dataclass(frozen=True)
class SaksReport:
    id: str
    selection: SaksSelection
    base: SessionReport
    enrichments: tuple[CaseEnrichment, ...]
    contexts: tuple[DocumentContext, ...]
    source_text: SourceTextSnapshot
    issues: tuple[str, ...]
    version: str = "b6-frozen-2"


def build_saks_report(database: Path, selection: SaksSelection) -> SaksReport:
    """Globalt inventar vurderes før beløpsfiltre; hele lesingen har data_version-vern."""
    with sqlite3.connect(f"{database.resolve().as_uri()}?mode=ro", uri=True) as guard:
        version = guard.execute("PRAGMA data_version").fetchone()[0]
        base = build_session_report(database, selection.base)
        enrichments = tuple(read_enrichment(database, i) for i in selection.enrichment_ids)
        by_revision: dict[str, CaseEnrichment] = {}
        revisions = {r.revision_id for r in base.inventory}
        for selected_enrichment in enrichments:
            if (
                selected_enrichment.dataset != selection.base.dataset
                or selected_enrichment.decision_revision_id not in revisions
            ):
                raise ValueError("Saksberikelsen må høre til valgt inventar og datasett")
            if selected_enrichment.decision_revision_id in by_revision:
                raise ValueError("Velg én saksberikelse per vedtaksrevisjon")
            by_revision[selected_enrichment.decision_revision_id] = selected_enrichment
        decisions = tuple(
            (r.revision_id, read_revision(database, r.revision_id).report.decision)
            for r in base.inventory
        )
        source_text = SourceTextSnapshot(
            tuple(read_extraction(database, i) for i in selection.base.extraction_ids), decisions
        )
        contexts: list[DocumentContext] = []
        issues: list[str] = []
        for row in base.inventory:
            decision = dict(decisions)[row.revision_id]
            if decision is None:
                raise ValueError("Saksinventaret mangler vedtaksidentitet")
            enrichment = by_revision.get(row.revision_id)
            context = read_case_context(
                database,
                dataset=selection.base.dataset,
                session=row.session,
                case_id=decision.case_id,
            )
            local: list[str] = []
            source_freshness = "ikke_valgt"
            if enrichment is None:
                local.append("saksberikelse_ikke_valgt")
            else:
                local.extend(enrichment.issues)
                if enrichment.content is None:
                    source_freshness = "ingen_tolket_saksrespons"
                    local.append("ingen_tolket_saksrespons")
                else:
                    content = enrichment.content
                    local.extend(content.issues)
                    for publication in content.publications:
                        local.extend(i + ":" + publication.locator for i in publication.issues)
                    source_freshness = (
                        "gjeldende"
                        if enrichment.attempt_id == context.current_attempt_id
                        else "historisk_observasjon_samme_innhold"
                        if content.revision_id == context.current_revision_id
                        else "historisk_innhold"
                    )
                    if source_freshness != "gjeldende":
                        local.append("saksgrunnlag_" + source_freshness)
                for keyword in enrichment.keywords:
                    if keyword.status == "uavklart":
                        local.append("uavklart_stikkord:" + keyword.locator)
            if context.latest_status != "observert":
                local.append(
                    "siste_saksforsok_" + context.latest_status + ":" + context.latest_reason
                )
                if enrichment and enrichment.content:
                    source_freshness = "historiske_referanser_etter_" + context.latest_status
            decision_freshness = (
                "gjeldende"
                if row.revision_id == row.context.current_revision_id
                and row.context.observation_status == "observert"
                else "historisk_eller_ikke_vurderbar"
            )
            if decision_freshness != "gjeldende":
                local.append("vedtaksgrunnlag_" + decision_freshness)
            contexts.append(
                DocumentContext(
                    row.revision_id,
                    decision.case_id,
                    decision.number,
                    row.session,
                    enrichment.id if enrichment else None,
                    enrichment.attempt_id if enrichment else None,
                    context,
                    source_freshness,
                    decision_freshness,
                    tuple(local),
                )
            )
            issues.extend(row.revision_id + ":" + i for i in local)
        result = SaksReport(
            "", selection, base, enrichments, tuple(contexts), source_text, tuple(issues)
        )
        result = replace(result, id=identity("b6-frozen-2", asdict(result)))
        if guard.execute("PRAGMA data_version").fetchone()[0] != version:
            raise ValueError("Grunnlaget endret seg under sakssammenstilling; kjør på nytt")
    return result


def report_to_json(report: SaksReport) -> str:
    validate_report(report)
    raw = encode({"format": "b6-frozen-2", "report": asdict(report)})
    if len(raw.encode("utf-8")) > MAX_REPORT_BYTES:
        raise ValueError("Frosset rapport overstiger 128 MiB")
    return raw


def report_from_json(raw: str) -> SaksReport:
    if len(raw.encode("utf-8")) > MAX_REPORT_BYTES:
        raise ValueError("Frosset rapport overstiger 128 MiB")
    data = json.loads(raw, object_pairs_hook=_object, parse_constant=_constant)
    if (
        not isinstance(data, dict)
        or set(data) != {"format", "report"}
        or data["format"] != "b6-frozen-2"
    ):
        raise ValueError("Ukjent B8-rapportkonvolutt/-versjon")
    result: SaksReport = _decode(SaksReport, data["report"])
    validate_report(result)
    return result


def validate_report(report: SaksReport) -> None:
    base_from_json(base_to_json(report.base))
    if (
        report.version != "b6-frozen-2"
        or report.id != identity("b6-frozen-2", asdict(replace(report, id="")))
        or report.base.selection != report.selection.base
        or tuple(e.id for e in report.enrichments) != report.selection.enrichment_ids
        or {c.decision_revision_id for c in report.contexts}
        != {r.revision_id for r in report.base.inventory}
    ):
        raise ValueError("Frosset saksrapport har feil innholdsbinding")
    for e in report.enrichments:
        if e.version != VERSION or e.id != identity(VERSION, asdict(replace(e, id=""))):
            raise ValueError("Frosset saksberikelse har feil identitet")
        if e.content and any(
            p.url is not None and safe_document_url(p.url) != p.url for p in e.content.publications
        ):
            raise ValueError("Frosset saksberikelse har utrygg dokumentlenke")


def _publish(temporary: Path, output: Path) -> None:
    if os.name == "nt":
        os.rename(temporary, output)
    else:
        os.link(temporary, output)


def export_saks_json(
    report: SaksReport, output: Path, *, database: Path, input_path: Path | None = None
) -> None:
    _output(database, output, input_path, ".json")
    raw = report_to_json(report)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with NamedTemporaryFile(
            dir=output.parent,
            suffix=".json",
            delete=False,
            mode="w",
            encoding="utf-8",
            newline="\n",
        ) as handle:
            temporary = Path(handle.name)
            handle.write(raw)
        if report_from_json(temporary.read_text("utf-8")) != report:
            raise ValueError("Frosset JSON avviker fra rapporten")
        _publish(temporary, output)
    finally:
        if temporary:
            temporary.unlink(missing_ok=True)


def _output(database: Path, output: Path, input_path: Path | None, suffix: str) -> None:
    if output.resolve() in {
        database.resolve(),
        Path(str(database.resolve()) + "-wal"),
        Path(str(database.resolve()) + "-shm"),
        Path(str(database.resolve()) + "-journal"),
    }:
        raise ValueError("Eksport kan ikke overskrive database eller SQLite-sidefil")
    if output.exists() or output.suffix.lower() != suffix:
        raise ValueError(f"Sakseksport krever nytt filnavn med {suffix}")
    if input_path and input_path.resolve() in {database.resolve(), output.resolve()}:
        raise ValueError("Database, eksport og utvalgsfil må ha forskjellige baner")


def _label(publication: Publication) -> str:
    value = publication.fields.get("lenke_tekst")
    return value.text or "Uavklart dokumentetikett" if value else "Uavklart dokumentetikett"


def _keyword(enrichment: CaseEnrichment | None, chapter: str, post: str) -> str:
    if enrichment is None:
        return "Uavklart: saksberikelse ikke valgt"
    candidates = [k for k in enrichment.keywords if (k.chapter, k.post) == (chapter, post)]
    if len(candidates) != 1:
        return "Uavklart: manglende eller flertydig stikkordtolkning"
    keyword = candidates[0]
    words = ", ".join(keyword.words)
    if keyword.status == "uavklart":
        return (words + "; " if words else "") + "Uavklart: " + keyword.reason
    return words or "Ingen stikkord funnet"


def _decorate(sheet: Worksheet) -> None:
    sheet.auto_filter.ref = sheet.dimensions
    for column, cell in enumerate(sheet[1], 1):
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="181C62")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.column_dimensions[get_column_letter(column)].width = 28
    sheet.row_dimensions[1].height = 60
    for number, row in enumerate(sheet.iter_rows(min_row=2), 2):
        sheet.row_dimensions[number].height = 60
        for cell in row:
            cell.alignment = Alignment(
                wrap_text=True,
                vertical="top",
                horizontal="left" if isinstance(cell.value, str) else "right",
            )
    for column, cell in enumerate(sheet[1], 1):
        label = str(cell.value)
        if label == "Stikkord":
            sheet.column_dimensions[get_column_letter(column)].width = 34
        if label == "Relasjonsnivå":
            sheet.column_dimensions[get_column_letter(column)].width = 42
        if "kr" in label.lower() and any(
            word in label
            for word in ("beløp", "kr /", "Rå kr", "Effektiv", "Korrigert", "godkjent")
        ):
            for row in sheet.iter_rows(min_row=2, min_col=column, max_col=column):
                row[0].number_format = "#,##0"


def _amount_documents(sheet: Worksheet, report: SaksReport) -> None:
    contexts = {c.decision_revision_id: c for c in report.contexts}
    enrichments = {e.id: e for e in report.enrichments}
    headers = {c.value: n for n, c in enumerate(sheet[1], 1)}
    original_headers = [str(c.value) for c in sheet[1]]
    post_column = headers["Post"]
    sheet.insert_cols(post_column + 1)
    put(sheet.cell(1, post_column + 1), "Stikkord")
    headers = {c.value: n for n, c in enumerate(sheet[1], 1)}
    api_column = headers["Vedtak"]
    put(sheet.cell(1, api_column), "Teknisk API-ID")
    counts = {
        kind: max(
            1,
            max(
                (
                    sum(p.kind == kind for p in e.content.publications)
                    for e in report.enrichments
                    if e.content
                ),
                default=0,
            ),
        )
        for kind in ("innstilling", "proposisjon")
    }
    added = ["Formelt vedtak", "Ferskhet for saksgrunnlag", "Relasjonsnivå"] + [
        f"{label} {n}"
        for kind, label in (("innstilling", "Innstilling"), ("proposisjon", "Proposisjon"))
        for n in range(1, counts[kind] + 1)
    ]
    sheet.insert_cols(api_column + 1, len(added))
    for n, label in enumerate(added, api_column + 1):
        put(sheet.cell(1, n), label)
    headers = {c.value: n for n, c in enumerate(sheet[1], 1)}
    for n in range(2, sheet.max_row + 1):
        context = contexts[str(sheet.cell(n, headers["Revisjon"]).value)]
        enrichment = enrichments.get(context.enrichment_id or "")
        put(
            sheet.cell(n, post_column + 1),
            _keyword(
                enrichment,
                str(sheet.cell(n, headers["Kapittel"]).value),
                str(sheet.cell(n, headers["Post"]).value),
            ),
        )
        put(
            sheet.cell(n, headers["Formelt vedtak"]), f"{context.formal_number} ({context.session})"
        )
        put(sheet.cell(n, headers["Ferskhet for saksgrunnlag"]), context.source_freshness)
        put(sheet.cell(n, headers["Relasjonsnivå"]), context.relation_level)
        for kind, label in (("innstilling", "Innstilling"), ("proposisjon", "Proposisjon")):
            pubs = (
                [p for p in enrichment.content.publications if p.kind == kind]
                if enrichment and enrichment.content
                else []
            )
            if not pubs:
                put(sheet.cell(n, headers[f"{label} 1"]), "Uavklart; se Dokumentavvik")
            for index, publication in enumerate(pubs, 1):
                cell = sheet.cell(n, headers[f"{label} {index}"])
                put(
                    cell,
                    _label(publication)
                    + (
                        " [" + publication.identity_status + "]"
                        if publication.identity_status != "identifisert"
                        else ""
                    ),
                )
                if publication.url:
                    cell.hyperlink = Hyperlink(ref=cell.coordinate, target=publication.url)
                    cell.style = "Hyperlink"
    # Tekstvidereføringene ble skrevet før B8-kolonnene ble satt inn.
    continuation_name = f"{sheet.title[:18]} fortsettelse"
    parent = sheet.parent
    assert parent is not None
    if continuation_name in parent:
        for row in parent[continuation_name].iter_rows(min_row=2):
            address = re.fullmatch(re.escape(sheet.title) + r"!([A-Z]+)([0-9]+)", str(row[0].value))
            if address:
                label = original_headers[column_index_from_string(address[1]) - 1]
                column = headers["Teknisk API-ID" if label == "Vedtak" else label]
                put(row[0], f"{sheet.title}!{get_column_letter(column)}{address[2]}")
    _decorate(sheet)


def _document_sheets(book: Workbook, report: SaksReport) -> None:
    documents: list[list[Value]] = [
        [
            "Berikelse",
            "Vedtaksrevisjon",
            "Saks-ID",
            "Dokument-ID",
            "Råetikett",
            "Type",
            "Nummer",
            "Serie",
            "Sesjon",
            "Lenke (metadata)",
            "Identitet",
            "Lenketilstand",
            "Saksgrunnlag",
            "SHA-256",
            "Lokator",
            "Råfelt med nil/fravær",
            "Avvik",
        ]
    ]
    relations: list[list[Value]] = [
        ["Fra", "Til", "Relasjon", "Nivå", "SHA-256", "Lokator", "Råfelt"]
    ]
    issues: list[list[Value]] = [
        [
            "Vedtaksrevisjon",
            "Saks-ID",
            "Berikelse",
            "Avvik",
            "Siste forsøk",
            "Forsøksstatus",
            "Forsøksgrunn",
            "Forsøkets SHA-256",
        ]
    ]
    keywords: list[list[Value]] = [
        [
            "Berikelse",
            "Kapittel",
            "Post",
            "Stikkord",
            "Tolkningsstatus",
            "Grunn",
            "Råtekst",
            "Rå-HTML",
            "SHA-256",
            "Lokator",
            "Versjon",
        ]
    ]
    contexts = {c.decision_revision_id: c for c in report.contexts}
    for e in report.enrichments:
        context = contexts[e.decision_revision_id]
        relations.append(
            [
                e.decision_revision_id,
                e.case_id,
                "vedtak_i_sak",
                context.relation_level,
                None,
                "stortingsvedtak/sak_id",
                None,
            ]
        )
        for k in e.keywords:
            keywords.append(
                [
                    e.id,
                    k.chapter,
                    k.post,
                    ", ".join(k.words),
                    k.status,
                    k.reason,
                    k.raw_text,
                    k.raw_html,
                    k.source_sha256,
                    k.locator,
                    k.version,
                ]
            )
        if e.content:
            for p in e.content.publications:
                documents.append(
                    [
                        e.id,
                        e.decision_revision_id,
                        e.case_id,
                        p.id,
                        _label(p),
                        p.kind,
                        p.number,
                        p.series,
                        p.session,
                        "Åpne dokument" if p.url else None,
                        p.identity_status,
                        p.link_status,
                        context.source_freshness,
                        p.source_sha256,
                        p.locator,
                        encode(asdict(p)["fields"]),
                        encode(p.issues),
                    ]
                )
                relations.append(
                    [
                        e.case_id,
                        p.id,
                        "publikasjon_referanse",
                        context.relation_level,
                        p.source_sha256,
                        p.locator,
                        encode(asdict(p)["fields"]),
                    ]
                )
            for r in e.content.relations:
                target, kind = r.fields.get("relatert_sak_id"), r.fields.get("relasjon_type")
                relations.append(
                    [
                        e.case_id,
                        target.text if target else None,
                        kind.text if kind else "uavklart",
                        "sak; ingen arvet posttilknytning",
                        r.source_sha256,
                        r.locator,
                        encode(asdict(r)["fields"]),
                    ]
                )
    for c in report.contexts:
        for issue in c.issues:
            issues.append(
                [
                    c.decision_revision_id,
                    c.case_id,
                    c.enrichment_id,
                    issue,
                    c.case_context.latest_attempt_id,
                    c.case_context.latest_status,
                    c.case_context.latest_reason,
                    c.case_context.latest_source_sha256,
                ]
            )
    for name, rows in (
        ("Saksdokumenter", documents),
        ("Dokumentrelasjoner", relations),
        ("Dokumentavvik", issues),
        ("Stikkordgrunnlag", keywords),
    ):
        write_rows(book, name, rows)
        _decorate(book[name])
    urls = {p.id: p.url for e in report.enrichments if e.content for p in e.content.publications}
    for row in book["Saksdokumenter"].iter_rows(min_row=2):
        if row[9].value:
            row[9].hyperlink = Hyperlink(ref=row[9].coordinate, target=urls[str(row[3].value)])
            row[9].style = "Hyperlink"
            row[9].font = Font(name="Arial", size=10, color="0000FF")
            row[9].alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")


def verify_saks_excel(path: Path, report: SaksReport) -> None:
    saved = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        new_layout = "Historikkdetaljer" in saved
        if new_layout:
            verify_history_view(saved, report)
        reserve_book = Workbook()
        try:
            write_rows(reserve_book, "Historikk", _history_rows(report.base))
            continuation = (
                reserve_book["Historikk fortsettelse"]
                if "Historikk fortsettelse" in reserve_book
                else None
            )
            # Original B8 lagret B6-pekerne før dokumentkolonnene ble satt inn.
            legacy_pointers = (
                [str(row[0].value) for row in continuation.iter_rows(min_row=2)]
                if not new_layout and continuation is not None
                else []
            )
            _amount_documents(reserve_book["Historikk"], report)
            stored_pointers = (
                {
                    str(row[0].value): old
                    for row, old in zip(
                        continuation.iter_rows(min_row=2), legacy_pointers, strict=True
                    )
                }
                if legacy_pointers and continuation is not None
                else None
            )
            reserve = saved["Historikkdetaljer" if new_layout else "Historikk"]
            if (reserve.max_row, reserve.max_column) != (
                reserve_book["Historikk"].max_row,
                reserve_book["Historikk"].max_column,
            ):
                raise ValueError("Historikkdetaljer mangler tidligere rader eller kolonner")
            for expected_row, actual_row in zip(reserve_book["Historikk"], reserve, strict=True):
                for wanted, cell in zip(expected_row, actual_row, strict=True):
                    if cell.value != wanted.value or (
                        cell.hyperlink.target if cell.hyperlink else None
                    ) != (wanted.hyperlink.target if wanted.hyperlink else None):
                        raise ValueError("Historikkdetaljer avviker fra tidligere detaljfelt")
            verify_continued_text(
                reserve, reserve_book["Historikk"], stored_pointers=stored_pointers
            )
        finally:
            reserve_book.close()
        sheet = saved["Beløpsobservasjoner"]
        headers = {c.value: n for n, c in enumerate(sheet[1], 1)}
        actual = {str(row[0].value): row for row in sheet.iter_rows(min_row=2)}
        if len(actual) != sheet.max_row - 1 or set(actual) != {
            r.observation.id for r in report.base.rows
        }:
            raise ValueError("Sakseksport har feil/dupliserte observasjoner")
        contexts = {c.decision_revision_id: c for c in report.contexts}
        enrichments = {e.id: e for e in report.enrichments}
        for r in report.base.rows:
            values = actual[r.observation.id]
            context = contexts[r.observation.source_revision_id]
            enrichment = enrichments.get(context.enrichment_id or "")
            for label, expected in (
                ("Råbeløp kr", r.observation.amount_kr),
                ("Gjeldende godkjent beløp kr", r.approved_amount_kr),
                ("Teknisk API-ID", r.observation.api_id),
                ("Revisjon", r.observation.source_revision_id),
                ("Formelt vedtak", f"{context.formal_number} ({context.session})"),
                ("Ferskhet for saksgrunnlag", context.source_freshness),
                (
                    "Stikkord",
                    _keyword(enrichment, r.observation.chapter or "", r.observation.post or ""),
                ),
            ):
                if values[headers[label] - 1].value != expected:
                    raise ValueError("Sakseksport avviker fra uendrede beløpsobservasjoner")
            for kind, label in (("innstilling", "Innstilling"), ("proposisjon", "Proposisjon")):
                pubs = (
                    [p for p in enrichment.content.publications if p.kind == kind]
                    if enrichment and enrichment.content
                    else []
                )
                for number, publication in enumerate(pubs, 1):
                    cell = values[headers[f"{label} {number}"] - 1]
                    expected_label = _label(publication) + (
                        " [" + publication.identity_status + "]"
                        if publication.identity_status != "identifisert"
                        else ""
                    )
                    if (
                        cell.value != expected_label
                        or (cell.hyperlink.target if cell.hyperlink else None) != publication.url
                    ):
                        raise ValueError("Sakseksport avviker fra valgte dokumentetiketter/lenker")
        expected_documents = {
            (e.id, p.id): p for e in report.enrichments if e.content for p in e.content.publications
        }
        seen_documents: set[tuple[str, str]] = set()
        for row in saved["Saksdokumenter"].iter_rows(min_row=2):
            key = (str(row[0].value), str(row[3].value))
            saved_publication = expected_documents.get(key)
            if (
                saved_publication is None
                or key in seen_documents
                or row[4].value != _label(saved_publication)
                or (row[9].hyperlink.target if row[9].hyperlink else None) != saved_publication.url
            ):
                raise ValueError("Sakseksport har feil/dupliserte dokumentreferanser")
            seen_documents.add(key)
        if seen_documents != set(expected_documents):
            raise ValueError("Sakseksport mangler valgte dokumentreferanser")
        if saved["Saksgrunnlag"].cell(2, 2).value != report.id:
            raise ValueError("Sakseksport har feil rapport-ID")
    finally:
        saved.close()


def export_built_saks_excel(
    database: Path, report: SaksReport, output: Path, *, input_path: Path | None = None
) -> SaksReport:
    """Bruk bare frosne data. Databasen leses aldri under filskrivingen."""
    validate_report(report)
    _output(database, output, input_path, ".xlsx")
    book = Workbook()
    assert book.active is not None
    book.remove(book.active)
    temporary: Path | None = None
    try:
        _write(book, database, report.base, source_snapshot=report.source_text)
        for name in ("Beløpsobservasjoner", "Historikk"):
            _amount_documents(book[name], report)
        for name in ("Kildeinventar", "Inventarfulltekst"):
            if name in book:
                sheet = book[name]
                for cell in sheet[1]:
                    if cell.value == "Vedtak":
                        put(cell, "Teknisk API-ID")
                headers = {c.value: n for n, c in enumerate(sheet[1], 1)}
                column = headers["Teknisk API-ID"] + 1
                sheet.insert_cols(column)
                put(sheet.cell(1, column), "Formelt vedtak")
                headers = {c.value: n for n, c in enumerate(sheet[1], 1)}
                contexts = {c.decision_revision_id: c for c in report.contexts}
                for n in range(2, sheet.max_row + 1):
                    context = contexts[str(sheet.cell(n, headers["Revisjon"]).value)]
                    put(sheet.cell(n, column), f"{context.formal_number} ({context.session})")
                _decorate(sheet)
        _document_sheets(book, report)
        write_rows(
            book,
            "Saksgrunnlag",
            [
                ["Felt", "Verdi"],
                ["Rapport-ID", report.id],
                ["Format", report.version],
                ["Saksutvalg", encode(asdict(report.selection))],
                ["Inventarvedtak før beløpsfiltre", len(report.contexts)],
                [
                    "Vedtak med valgt saksberikelse",
                    sum(c.enrichment_id is not None for c in report.contexts),
                ],
                [
                    "Vedtak uten valgt saksberikelse",
                    sum(c.enrichment_id is None for c in report.contexts),
                ],
                [
                    "Vedtak med siste saksforsøk feilet",
                    sum(c.case_context.latest_status == "feilet" for c in report.contexts),
                ],
                ["Globale dokumentavvik før beløpsfiltre", len(report.issues)],
                ["Uttømmende dokumentdekning", "ikke erklært"],
                [
                    "Kilde",
                    "Stortinget. Publikasjonslenker er metadata; dokumentinnhold er ikke hentet.",
                ],
                [
                    "Grense",
                    "Saksrelasjoner beviser ingen postvis forslagstilknytning. "
                    "Ingen ny faglig godkjenning eller årsdekning.",
                ],
            ],
        )
        _decorate(book["Saksgrunnlag"])
        book["Saksgrunnlag"].column_dimensions["A"].width = 52
        book["Saksgrunnlag"].column_dimensions["B"].width = 90
        write_history_view(book, report)
        output.parent.mkdir(parents=True, exist_ok=True)
        with NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        verify_saks_excel(temporary, report)
        _publish(temporary, output)
    finally:
        book.close()
        if temporary:
            temporary.unlink(missing_ok=True)
    return report
