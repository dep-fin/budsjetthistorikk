"""Sikker, atomisk sesjonseksport fra eksplisitt lagret grunnlag."""

import os
from collections import defaultdict
from collections.abc import Sequence
from dataclasses import asdict
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook, load_workbook

from innlesning.observations import encode
from innlesning.sync import read_revision
from innlesning.workfiles import safe_output

from .extraction_excel import Value, add_source_text
from .extraction_excel import write_rows as _write_rows
from .session import SessionReport, SessionRow, SessionSelection, build_session_report


def write_rows(book: Workbook, title: str, rows: Sequence[Sequence[Value]]) -> None:
    if len(rows) > 1_048_576 or any(len(row) > 16_384 for row in rows):
        raise ValueError("Sesjonsuttrekket overstiger Excel-grensen")
    _write_rows(book, title, [list(row) for row in rows])


def export_session_excel(
    database: Path, selection: SessionSelection, output: Path, *, input_path: Path | None = None
) -> SessionReport:
    safe_output(database, output)
    if output.exists():
        raise ValueError(
            "Sesjonseksport krever et nytt filnavn; eksisterende filer overskrives ikke"
        )
    if output.suffix.lower() != ".xlsx":
        raise ValueError("Sesjonseksport krever .xlsx")
    if input_path and (
        input_path.resolve() == output.resolve() or input_path.resolve() == database.resolve()
    ):
        raise ValueError("Database, Excel og utvalgsfil må ha forskjellige baner")
    report = build_session_report(database, selection)
    book = Workbook()
    assert book.active is not None
    book.remove(book.active)
    try:
        _write(book, database, report)
        output.parent.mkdir(parents=True, exist_ok=True)
        temporary: Path | None = None
        try:
            with NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as handle:
                temporary = Path(handle.name)
            book.save(temporary)
            _verify_export(temporary, report)
            # Windows rename og POSIX link nekter atomisk et eksisterende mål.
            if os.name == "nt":
                os.rename(temporary, output)
            else:
                os.link(temporary, output)
        finally:
            if temporary:
                temporary.unlink(missing_ok=True)
    finally:
        book.close()
    return report


def _verify_export(path: Path, report: SessionReport) -> None:
    """Les de faktiske Excel-cellene tilbake før atomisk publisering."""
    saved = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        expected = {r.observation.id: r for r in report.rows}
        seen: set[str] = set()
        for values in saved["Beløpsobservasjoner"].iter_rows(min_row=2, values_only=True):
            identifier = str(values[0])
            row = expected.get(identifier)
            if row is None or identifier in seen:
                raise ValueError("Eksportavstemming: feil/duplisert observasjons-ID")
            seen.add(identifier)
            if (values[1], values[18], values[20], values[21], values[22], values[23]) != (
                row.extraction_id,
                row.observation.amount_kr,
                row.historical_effective_amount_kr,
                row.approved_amount_kr,
                encode(row.controlled_result_ids),
                row.controlled.decision_id if row.controlled else None,
            ):
                raise ValueError("Eksportavstemming: Excel avviker fra lagrede B5-rader og ID-er")
        if seen != set(expected):
            raise ValueError("Eksportavstemming: manglende observasjoner")
    finally:
        saved.close()


def _history_rows(report: SessionReport) -> list[list[Value]]:
    meanings = ("nivaa", "for_belop", "endring", "etter_belop")
    headers: list[Value] = [
        "Budsjettår",
        "Kapittel",
        "Post",
        "Underpost",
        "Nivå",
        "Dato",
        "Vedtak",
        "Ledd",
        "Sesjon",
        "Runde",
    ]
    for label in ("Rå kr", "Korrigert kr", "Effektiv historisk kr", "Gjeldende godkjent kr"):
        headers.extend(f"{meaning} / {label}" for meaning in meanings)
    headers.extend(
        [
            "Observasjons-ID-er",
            "Lokatorer",
            "Kilde SHA-256",
            "Revisjon",
            "Tolkning",
            "Kilde",
            "Avvik",
        ]
    )
    groups: dict[
        tuple[str, int | None, str | None, str | None, str | None, str], list[SessionRow]
    ] = defaultdict(list)
    for row in report.rows:
        o = row.observation
        groups[(o.step_id, o.budget_year, o.chapter, o.post, o.subpost, o.detail_level)].append(row)
    result: list[list[Value]] = [headers]
    sources = {r.revision_id: r.url for r in report.inventory}
    for group in sorted(
        groups.values(),
        key=lambda g: (
            g[0].observation.decision_date or "9999",
            str(g[0].observation.api_id),
            str(g[0].observation.chapter),
            str(g[0].observation.post),
        ),
    ):
        o = group[0].observation
        by_meaning = {m: [r for r in group if r.observation.meaning == m] for m in meanings}
        values: list[Value] = [
            o.budget_year,
            o.chapter,
            o.post,
            o.subpost,
            o.detail_level,
            o.decision_date,
            o.api_id,
            o.step_id,
            o.decision_session,
            o.round,
        ]
        for field in (
            "raw",
            "corrected_amount_kr",
            "historical_effective_amount_kr",
            "approved_amount_kr",
        ):
            for meaning in meanings:
                candidates = by_meaning[meaning]
                value = None
                if len(candidates) == 1:
                    value = (
                        candidates[0].observation.amount_kr
                        if field == "raw"
                        else getattr(candidates[0], field)
                    )
                values.append(value)
        values.extend(
            [
                encode([r.observation.id for r in group]),
                encode([asdict(r.observation.pointer) for r in group]),
                o.pointer.artifact_id,
                o.source_revision_id,
                o.interpretation_id,
                sources.get(o.source_revision_id),
                encode(sorted({i for r in group for i in r.issues})),
            ]
        )
        result.append(values)
    return result


def _write(book: Workbook, database: Path, report: SessionReport) -> None:
    write_rows(
        book,
        "Sesjonsgrunnlag",
        [
            ["Felt", "Verdi"],
            ["Rapport-ID", report.id],
            ["Vurdert UTC", report.assessed_utc],
            ["Utvalg (JSON)", encode(asdict(report.selection))],
            ["Dekningsgrense", report.coverage.boundary],
            ["Avstemming", "utført" if report.reconciliations else "ikke utført"],
            ["Kildeprøver", "utført" if report.source_checks else "ikke utført"],
            ["Uavhengig kontroll #2", "utført" if report.independent_control else "ikke utført"],
        ],
    )
    inventory: list[list[Value]] = [
        [
            "Sesjon",
            "Vedtak",
            "Revisjon",
            "Inventarjobber",
            "Supplerende",
            "API-type",
            "Vedtaksdato",
            "Tittel",
            "Kilde",
            "Rå SHA-256",
            "Uttrekk",
            "Kildeinnhold",
            "Struktur",
            "Observasjoner",
            "Treff",
            "Uavklart treff",
            "Utenfor filtre",
            "Godkjente enkeltbeløp",
            "Aktuell revisjon",
            "Kildeovergang",
            "Siste forsøk",
            "Forsøksstatus",
            "Observasjonsstatus",
            "Avvik",
            "Siste forsøks feil",
            "Siste forsøks rå SHA-256",
            "Siste forsøks omfang",
        ]
    ]
    for r in report.inventory:
        inventory.append(
            [
                r.session,
                r.decision_id,
                r.revision_id,
                encode(r.job_ids),
                str(r.supplemental),
                r.api_type,
                r.date,
                r.title,
                r.url,
                r.source_sha256,
                encode(r.extraction_ids),
                r.content_status,
                r.structure_status,
                r.observation_count,
                r.selected_count,
                r.unresolved_count,
                r.excluded_count,
                r.approved_count,
                r.context.current_revision_id,
                r.context.generation,
                r.context.latest_attempt_id,
                r.context.latest_attempt_status,
                r.context.observation_status,
                encode(r.issues),
                encode(r.context.latest_attempt_errors),
                r.context.latest_attempt_source_sha256,
                r.context.latest_attempt_scope,
            ]
        )
    write_rows(book, "Kildeinventar", inventory)
    write_rows(
        book,
        "Inventarjobber",
        [["Jobb", "Alle felt"]] + [[j.job_id, encode(asdict(j))] for j in report.snapshot.jobs],
    )
    write_rows(
        book,
        "Dekning",
        [
            [
                "Sesjon",
                "Supplerende",
                "Inventar-ID-er",
                "Bevart innhold",
                "Strukturerte vedtak",
                "Beløpsobservasjoner",
                "Godkjente enkeltbeløp",
                "Fullt kildeinnhold",
                "Full struktur",
                "Uttømmende faglig dekning",
                "Grense",
            ]
        ]
        + [
            [
                c.session,
                str(c.supplemental),
                c.inventory_count,
                c.source_count,
                c.structured_count,
                c.observation_count,
                c.approved_count,
                str(c.source_complete),
                str(c.structure_complete),
                str(c.professional_complete),
                c.boundary,
            ]
            for c in (report.coverage, *report.supplemental_coverage)
        ],
    )
    observations: list[list[Value]] = [
        [
            "Observasjon-ID",
            "Uttrekk",
            "Sesjon",
            "Budsjettår",
            "Kapittel",
            "Post",
            "Underpost",
            "Detaljnivå",
            "Komponent av",
            "Side",
            "fdep_nr",
            "gdep_nr",
            "Register (historisk)",
            "Vedtaksdato",
            "Vedtak",
            "Ledd",
            "Beløpsbetydning",
            "Råverdi",
            "Råbeløp kr",
            "Korrigert beløp kr",
            "Effektiv verdi (historisk resultat)",
            "Gjeldende godkjent beløp kr",
            "Kontrollresultater",
            "Avgjørelse",
            "Godkjent bruk",
            "Ferskhet",
            "Filterstatus",
            "Supplerende",
            "Kilde SHA-256",
            "Lokator",
            "Beløpsfelt",
            "Revisjon",
            "Tolkning",
            "Avvik",
            "Alle observasjonsfelt",
            "Alle kontrollfelt",
            "Faglig status",
            "Godkjente felter",
            "Bevis-ID",
            "Kontrollør",
            "Begrunnelse",
        ]
    ]
    for session_row in report.rows:
        o = session_row.observation
        observations.append(
            [
                o.id,
                session_row.extraction_id,
                o.decision_session,
                o.budget_year,
                o.chapter,
                o.post,
                o.subpost,
                o.detail_level,
                o.component_of,
                o.budget_side,
                o.organization.fdep_nr,
                o.organization.gdep_nr,
                o.organization.register_version,
                o.decision_date,
                o.api_id,
                o.step_id,
                o.meaning,
                o.raw_value,
                o.amount_kr,
                session_row.corrected_amount_kr,
                session_row.historical_effective_amount_kr,
                session_row.approved_amount_kr,
                encode(session_row.controlled_result_ids),
                session_row.controlled.decision_id if session_row.controlled else None,
                session_row.controlled.approved_use if session_row.controlled else None,
                session_row.freshness,
                session_row.selection_status,
                str(session_row.supplemental),
                o.pointer.artifact_id,
                o.pointer.locator,
                o.pointer.amount_field,
                o.source_revision_id,
                o.interpretation_id,
                encode(session_row.issues),
                encode(asdict(o)),
                encode(asdict(session_row.controlled)) if session_row.controlled else None,
                session_row.controlled.status if session_row.controlled else "ikke vurdert",
                session_row.controlled.approved_fields if session_row.controlled else None,
                session_row.controlled.evidence_id if session_row.controlled else None,
                session_row.controlled.reviewer if session_row.controlled else None,
                session_row.controlled.reason if session_row.controlled else None,
            ]
        )
    write_rows(book, "Beløpsobservasjoner", observations)
    write_rows(book, "Historikk", _history_rows(report))
    write_rows(
        book,
        "Summer",
        [
            [
                "Budsjettår",
                "Beløpsbetydning",
                "Detaljnivå",
                "Side",
                "Valuta",
                "Multiplikator",
                "Godkjent delsum kr",
                "Komplett sum kr",
                "Inkluderte observasjoner",
                "Utelatte observasjoner",
                "Avvik",
                "Sumomfang",
            ]
        ]
        + [
            [
                s.budget_year,
                s.meaning,
                s.detail_level,
                s.budget_side,
                s.currency,
                s.multiplier,
                s.approved_subtotal_kr,
                s.complete_total_kr,
                encode(s.included_observation_ids),
                encode(s.excluded_observation_ids),
                encode(s.issues),
                "Delsum innen valgte beløpsfiltre; se Sesjonsgrunnlag"
                if "filtrert_delsum" in s.issues
                else "Uttrykkelig valgt kildegrunnlag; ingen årsdekning",
            ]
            for s in report.sums
        ],
    )
    write_rows(
        book,
        "Historikkberegninger",
        [["Beregning", "Teknisk saldo kr", "Status", "Ferskhet", "Alle felt"]]
        + [
            [h.id, h.calculated_balance, h.status, h.freshness, encode(asdict(h))]
            for h in report.histories
        ],
    )
    write_rows(
        book,
        "Kontrollresultater",
        [["Resultat", "Ferskhet", "Godkjent kjedesaldo kr", "Dekningsgrense", "Alle felt"]]
        + [
            [r.id, r.freshness, r.approved_chain_balance, r.coverage, encode(asdict(r))]
            for r in report.controlled_results
        ],
    )
    write_rows(
        book,
        "Avvik",
        [
            [
                "Grunnlag",
                "Avvik",
                "Filterstatus",
                "Budsjettår",
                "Kapittel",
                "Post",
                "Underpost",
                "Nivå",
                "Råverdi",
                "Uttrekk",
                "Lokator",
            ]
        ]
        + [["Global oversikt", i] for i in report.issues]
        + [
            [
                r.observation.id,
                issue,
                r.selection_status,
                r.observation.budget_year,
                r.observation.chapter,
                r.observation.post,
                r.observation.subpost,
                r.observation.detail_level,
                r.observation.raw_value,
                r.extraction_id,
                r.observation.pointer.locator,
            ]
            for r in report.unfiltered_rows
            for issue in r.issues
        ],
    )
    write_rows(
        book,
        "Avstemming",
        [
            [
                "Kontroll",
                "Kapittel",
                "Post",
                "Underpost",
                "Faktisk godkjent kr",
                "Forventet kr",
                "Uttrykkelig fravær",
                "Differanse kr",
                "Status",
                "Begrunnelse",
                "Kildereferanse/hash/lokator",
                "Omfang",
                "Kontrollgrunnlag",
            ]
        ]
        + [
            [
                result.specification.name,
                row.chapter,
                row.post,
                row.subpost,
                row.actual_amount_kr,
                row.expected_amount_kr,
                str(row.reference_absent),
                row.difference_kr,
                row.status,
                row.reason,
                encode(asdict(row.reference)) if row.reference else None,
                encode(
                    {
                        "actual_scope": asdict(result.specification.actual_scope),
                        "reference_scope": asdict(result.specification.reference_scope),
                    }
                ),
                f"Sesjonsgrunnlag!B4/reconciliations/{index}",
            ]
            for index, result in enumerate(report.reconciliations)
            for row in result.rows
        ]
        + ([] if report.reconciliations else [["ikke utført"]]),
    )
    write_rows(
        book,
        "Kildeprøvekontroll",
        [["Observasjon", "Rå kr", "Godkjent effektiv kr", "Status", "Forventning og kildebevis"]]
        + [
            [
                c.expectation.observation_id,
                c.actual_raw_amount_kr,
                c.actual_effective_amount_kr,
                c.status,
                encode(asdict(c.expectation)),
            ]
            for c in report.source_checks
        ]
        + ([] if report.source_checks else [["ikke utført"]]),
    )
    write_rows(
        book,
        "Eksportavstemming",
        [
            ["Kontroll", "Antall", "Status"],
            [
                "Eksakte Excel-observasjons-ID-er og lagrede B5-resultat-/avgjørelses-ID-er",
                len(report.rows),
                "Excel leses tilbake før publisering",
            ],
            [
                "Gjeldende godkjente enkeltverdier mot lagret B5 effektiv verdi",
                sum(r.approved_amount_kr is not None for r in report.rows),
                "Eksakt likhet; historiske verdier holdes separat",
            ],
        ],
    )
    if report.independent_control:
        control = report.independent_control
        write_rows(
            book,
            "Uavhengig kontroll 2024",
            [
                ["Felt", "Verdi"],
                ["Kontrollgrunnlag og beregnet resultat", encode(asdict(control))],
                ["Grense", control.boundary],
            ],
        )
        write_rows(
            book,
            "Omfangsforskjeller 2024",
            [["Konto", "DFØ årets bevilgning kr", "PDF Løyving kr", "Forklaring", "Status"]]
            + [
                [r.account, r.dfo_amount_kr, r.pdf_amount_kr, r.classification, r.status]
                for r in control.residuals
            ],
        )
    add_source_text(book, database, report.selection.extraction_ids)
    fulltext: list[list[Value]] = [
        ["Sesjon", "Vedtak", "Revisjon", "SHA-256", "Del", "Vedtaksfulltekst (HTML)", "Grense"]
    ]
    for r in report.inventory:
        if r.extraction_ids:
            continue
        decision = read_revision(database, r.revision_id).report.decision
        if decision:
            fulltext.extend(
                [
                    r.session,
                    r.decision_id,
                    r.revision_id,
                    r.source_sha256,
                    start // 30000 + 1,
                    decision.html[start : start + 30000],
                    "Kilde uten valgt beløpsuttrekk; tekst og fullmakter til kontroll",
                ]
                for start in range(0, len(decision.html), 30000)
            )
    write_rows(book, "Inventarfulltekst", fulltext)
