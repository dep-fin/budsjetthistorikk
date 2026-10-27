"""T1: selvforsynt Excel med reelle celle-/lenke- og tekststykkerkontroller."""

import base64
from copy import copy
from dataclasses import asdict
from pathlib import Path
from tempfile import NamedTemporaryFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.hyperlink import Hyperlink

from innlesning.observations import encode
from innlesning.omtale import compact
from innlesning.visningspilot_model import PilotReport

from .extraction_excel import Value, write_rows
from .omtale import report_from_json as read_omtale
from .omtale_excel_guard import SheetShape, inspect_excel
from .saksdokumenter import _publish
from .visningspilot import export_json, read_json, report_from_json, report_to_json

SHEET_NAMES = frozenset(
    {
        "Forløp",
        "Dokumenter",
        "Kontoer",
        "Inventar",
        "Fulltekst",
        "Beløpsgrunnlag",
        "Relasjoner",
        "Sammenligninger",
        "Dekning",
        "Originaler",
        "Utvalg",
        "Visningsgrunnlag",
    }
)


def _chunks(value: str) -> list[str]:
    return [value[i : i + 30000] for i in range(0, len(value), 30000)] or [""]


def _tables(report: PilotReport, frozen: str) -> dict[str, list[list[Value]]]:
    basis = report.basis
    docs = {d.id: d for d in basis.documents}
    contexts = {c.observation_id: c for c in basis.contexts}
    forlop: list[list[Value]] = [
        [
            "Observasjons-ID",
            "Dokument",
            "Dokumentdato",
            "Budsjettår",
            "Kapittel",
            "Post",
            "Kildebeløp kr",
            "Betydning",
            "Talltype",
            "Forslagsrolle",
            "Råbeløp",
            "Enhet",
            "Lokator",
            "Kilde SHA-256",
            "Rånavn",
            "Runde i gammel observasjon",
            "Separat behandlingskontekst",
            "Ledd-ID",
            "Teknisk status",
            "Faglig status",
        ]
    ]
    for o in basis.observations:
        if o.id not in report.view.observation_ids:
            continue
        context = contexts[o.id]
        document = docs[context.document_id]
        forlop.append(
            [
                o.id,
                document.title,
                document.date,
                o.budget_year,
                o.chapter,
                o.post,
                o.amount_kr,
                o.meaning,
                o.figure_type,
                o.proposal_role,
                o.raw_value,
                o.unit_raw,
                o.pointer.locator,
                o.pointer.artifact_id,
                o.purpose,
                o.round,
                encode(asdict(context)),
                o.step_id,
                o.structure_status,
                o.professional_status,
            ]
        )
    documents: list[list[Value]] = [
        [
            "Dokumentversjon",
            "Dokument",
            "Type",
            "Sesjon",
            "Dato",
            "Råkildedekning",
            "Strukturdekning",
            "Hele dokumentinventaret",
            "Kilder",
            "Faglig status",
            "Valgt dokument",
        ]
    ]
    documents.extend(
        [
            d.id,
            d.title,
            d.kind,
            d.session,
            d.date,
            d.raw_coverage,
            d.structured_coverage,
            encode(d.accounts),
            encode(d.source_ids),
            d.professional_status,
            "ja" if d.id in report.view.document_ids else "nei",
        ]
        for d in basis.documents
    )
    inventory: list[list[Value]] = [
        [
            "Inventar-ID",
            "Dokument",
            "Budsjettår",
            "Kapittel",
            "Post",
            "Rolle",
            "Strukturstatus",
            "Kildetekst (kompakt)",
            "Lokator",
            "Kilde-ID",
            "Beløps-ID-er",
            "Valgt",
        ]
    ]
    inventory.extend(
        [
            i.id,
            docs[i.document_id].title,
            i.budget_year,
            i.chapter,
            i.post,
            i.role,
            i.structure_status,
            compact(i.text),
            i.locator,
            i.source_id,
            encode(i.observation_ids),
            "ja" if i.id in report.view.inventory_ids else "nei",
        ]
        for i in basis.inventory
    )
    accounts: list[list[Value]] = [
        [
            "Kapittel",
            "Post",
            "Budsjettår",
            "Dokument",
            "Rolle",
            "Inventar-ID",
            "Strukturstatus",
        ]
    ]
    accounts.extend(
        [
            i.chapter,
            i.post,
            i.budget_year,
            docs[i.document_id].title,
            i.role,
            i.id,
            i.structure_status,
        ]
        for i in sorted(
            basis.inventory, key=lambda r: (int(r.chapter), int(r.post), r.document_id, r.role)
        )
    )
    fulltext: list[list[Value]] = [
        [
            "Utsagn-ID",
            "Dokument",
            "Kilde-ID",
            "Lokator",
            "Rolle",
            "Avsender",
            "Periode/betydning",
            "Kontoer",
            "Hierarki",
            "Del",
            "Originaltekst (åpne celle for fulltekst)",
            "Originalstruktur",
            "Teknisk status",
            "Valgt",
        ]
    ]
    for statement in basis.statements:
        text, structure = _chunks(statement.text), _chunks(statement.html)
        for number in range(max(len(text), len(structure))):
            fulltext.append(
                [
                    statement.id,
                    docs[statement.document_id].title,
                    statement.source_id,
                    statement.locator,
                    statement.role,
                    statement.actor,
                    statement.period,
                    encode(statement.accounts),
                    encode(statement.hierarchy),
                    number + 1,
                    text[number] if number < len(text) else "",
                    structure[number] if number < len(structure) else "",
                    statement.status,
                    "ja" if statement.id in report.view.statement_ids else "nei",
                ]
            )
    observations: list[list[Value]] = [
        [
            "Observasjons-ID",
            "Dokument",
            "Kapittel",
            "Post",
            "År",
            "Kroner",
            "Betydning",
            "Talltype",
            "Rolle",
            "Ledd",
            "Originalfelt med alle nullgrunner",
            "Kontekst",
        ]
    ]
    observations.extend(
        [
            o.id,
            docs[contexts[o.id].document_id].title,
            o.chapter,
            o.post,
            o.budget_year,
            o.amount_kr,
            o.meaning,
            o.figure_type,
            o.proposal_role,
            o.step_id,
            encode(asdict(o)),
            encode(asdict(contexts[o.id])),
        ]
        for o in basis.observations
    )
    relations: list[list[Value]] = [
        [
            "Relasjons-ID",
            "Fra",
            "Til",
            "Rolle",
            "Omfang",
            "Kontoer",
            "År",
            "Kildebevis",
            "Teknisk status",
            "Faglig status",
            "Fullt originalgrunnlag",
        ]
    ]
    relations.extend(
        [
            r.id,
            r.from_id,
            r.to_id,
            r.role,
            r.scope,
            encode(r.accounts),
            r.budget_year,
            encode(r.evidence_ids),
            r.status,
            "ikke_faglig_godkjent",
            encode(asdict(r)),
        ]
        for r in basis.relations
    )
    old = read_omtale(basis.o2b_json)
    relations.extend(
        [
            r.id,
            r.section_id,
            r.event_id,
            r.role,
            "dokumentert_posthendelse",
            encode(r.accounts),
            2024,
            encode(r.evidence_ids),
            r.status,
            r.professional_status,
            encode(asdict(r)),
        ]
        for r in old.relations
    )
    comparisons: list[list[Value]] = [
        [
            "Sammenlignings-ID",
            "Konto",
            "Betydning",
            "Venstre observasjon",
            "Høyre observasjon",
            "Avledet forskjell kr",
            "Status",
            "Regel",
            "Grunn",
            "Nullgrunn",
        ]
    ]
    comparisons.extend(
        [
            c.id,
            c.account,
            c.meaning,
            c.left_id,
            c.right_id,
            c.difference_kr,
            c.status,
            c.rule,
            c.reason,
            c.null_reason,
        ]
        for c in report.view.comparisons
    )
    sources: list[list[Value]] = [
        [
            "Kilde-ID",
            "Originalnavn",
            "Dokumentnøkkel",
            "Format",
            "SHA-256",
            "Byte",
            "Forventet SHA-256",
            "Status",
            "Metadata",
            "Del",
            "Originalbyte base64",
        ]
    ]
    originals = [
        (
            s.id,
            s.name,
            s.document_key,
            s.media_type,
            s.sha256,
            s.byte_count,
            s.expected_sha256,
            s.status,
            s.metadata_json,
            s.content_base64,
        )
        for s in basis.sources
    ] + [
        (
            s.id,
            s.name,
            s.document_id,
            s.media_type,
            s.sha256,
            s.byte_count,
            s.expected_sha256,
            s.status,
            s.metadata_json,
            s.content_base64,
        )
        for s in old.basis.sources
    ]
    for identifier, name, key, mime, digest, size, pin, status, metadata, content in originals:
        for number, chunk in enumerate(_chunks(content or ""), 1):
            sources.append(
                [identifier, name, key, mime, digest, size, pin, status, metadata, number, chunk]
            )
    selection: list[list[Value]] = [["Felt", "Verdi"]]
    selection.extend(
        [
            ["Rapport-ID", report.id],
            ["Grunnlags-ID", basis.id],
            ["Versjon", report.version],
            ["Valgt inngang", encode(asdict(report.selection))],
            [
                "Retur til inngang",
                encode(asdict(report.return_selection)) if report.return_selection else None,
            ],
            ["Valgte tabellfelt", len(report.view.observation_ids)],
            ["Valgte vedtaksfelt", len(report.view.adopted_observation_ids)],
            ["Valgte vedtatte hendelser", len(report.view.adopted_event_ids)],
            ["Globale tabellfelt", len(basis.observations)],
            ["Globale vedtaksfelt", 6],
            ["Globale vedtatte hendelser", 2],
        ]
    )
    snapshot: list[list[Value]] = [["Del", "Frosset visningspilot-1 med originalbyte"]]
    snapshot.extend([number, chunk] for number, chunk in enumerate(_chunks(frozen), 1))
    coverage: list[list[Value]] = [["Globalt forhold før filter", "Status"]]
    coverage.extend([issue, "synlig"] for issue in basis.issues)
    return {
        "Forløp": forlop,
        "Dokumenter": documents,
        "Kontoer": accounts,
        "Inventar": inventory,
        "Fulltekst": fulltext,
        "Beløpsgrunnlag": observations,
        "Relasjoner": relations,
        "Sammenligninger": comparisons,
        "Dekning": coverage,
        "Originaler": sources,
        "Utvalg": selection,
        "Visningsgrunnlag": snapshot,
    }


def _link(book: Workbook, sheet: str, cell: str, target_sheet: str, target_cell: str) -> None:
    target = book[sheet][cell]
    target.hyperlink = Hyperlink(ref=cell, location=f"'{target_sheet}'!{target_cell}")
    target.font = Font(name="Arial", color="0563C1", underline="single")


def _book(report: PilotReport, frozen: str) -> Workbook:
    book = Workbook()
    assert book.active is not None
    book.remove(book.active)
    for title, rows in _tables(report, frozen).items():
        write_rows(book, title, rows)
        sheet = book[title]
        sheet.row_dimensions[1].height = 40
        for cell in sheet[1]:
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        # Lange kildefelt beholdes komplett, med tydelige delnummer i kontrollarkene.
        for row_number, row in enumerate(sheet.iter_rows(min_row=2), 2):
            sheet.row_dimensions[row_number].height = 18
            for cell in row:
                cell.alignment = Alignment(vertical="top")
                if type(cell.value) is int:
                    cell.number_format = (
                        "0" if sheet.cell(1, cell.column).value in {"Budsjettår", "År"} else "#,##0"
                    )
        if title == "Forløp":
            for column in ("A", "M", "N", "Q", "R"):
                sheet.column_dimensions[column].width = 22
            sheet.column_dimensions["B"].width = 40
            sheet.column_dimensions["G"].width = 22
            for column, width in {"C": 16, "D": 12, "E": 12, "F": 10}.items():
                sheet.column_dimensions[column].width = width
        elif title == "Dokumenter":
            for column, width in {"B": 52, "C": 18, "D": 16, "E": 16, "F": 46}.items():
                sheet.column_dimensions[column].width = width
        elif title in {"Kontoer", "Inventar", "Beløpsgrunnlag"}:
            year_columns = {"Kontoer": "C", "Inventar": "C", "Beløpsgrunnlag": "E"}
            sheet.column_dimensions[year_columns[title]].width = 12
        elif title == "Dekning":
            sheet.column_dimensions["A"].width = 116
            sheet.column_dimensions["B"].width = 16
        elif title == "Utvalg":
            sheet.column_dimensions["A"].width = 32
            sheet.column_dimensions["B"].width = 100
        elif title == "Fulltekst":
            for column, width in {"E": 38, "F": 28, "G": 52, "J": 8, "K": 100}.items():
                sheet.column_dimensions[column].width = width
    doc_rows = {d.id: number for number, d in enumerate(report.basis.documents, 2)}
    ctx = {c.observation_id: c for c in report.basis.contexts}
    item_rows = {i.id: number for number, i in enumerate(report.basis.inventory, 2)}
    account_rows: dict[tuple[str, str], int] = {}
    for row in book["Kontoer"].iter_rows(min_row=2):
        pair = (str(row[0].value), str(row[1].value))
        assert row[0].row is not None
        account_rows.setdefault(pair, row[0].row)
        identifier = str(row[5].value)
        _link(book, "Kontoer", "F" + str(row[0].row), "Inventar", "A" + str(item_rows[identifier]))
    for number, observation_id in enumerate(report.view.observation_ids, 2):
        document_id = ctx[observation_id].document_id
        _link(book, "Forløp", "B" + str(number), "Dokumenter", "B" + str(doc_rows[document_id]))
    for i, item in enumerate(report.basis.inventory, 2):
        _link(book, "Inventar", "B" + str(i), "Dokumenter", "B" + str(doc_rows[item.document_id]))
        _link(
            book,
            "Inventar",
            "D" + str(i),
            "Kontoer",
            "A" + str(account_rows[(item.chapter, item.post)]),
        )
    for document in report.basis.documents:
        target = next(
            (item_rows[i.id] for i in report.basis.inventory if i.document_id == document.id), None
        )
        if target is not None:
            _link(
                book, "Dokumenter", "H" + str(doc_rows[document.id]), "Inventar", "A" + str(target)
            )
        first_text = next(
            (
                number
                for number, row in enumerate(book["Fulltekst"].iter_rows(min_row=2), 2)
                if row[1].value == document.title
            ),
            None,
        )
        if first_text is not None:
            _link(
                book,
                "Dokumenter",
                "B" + str(doc_rows[document.id]),
                "Fulltekst",
                "B" + str(first_text),
            )
    return book


def read_excel(path: Path) -> PilotReport:
    raw, shapes = inspect_excel(
        path,
        sheet_names=SHEET_NAMES,
        snapshot_name="Visningsgrunnlag",
        max_expanded_bytes=128 * 1024 * 1024,
    )
    report = report_from_json(raw)
    expected = _book(report, raw)
    try:
        if shapes != {s.title: SheetShape(s.max_row, s.max_column) for s in expected}:
            raise ValueError("Visnings-Excel har uventede ark/celleområder")
        saved = load_workbook(path, data_only=False, keep_links=False)
        try:
            if saved.sheetnames != expected.sheetnames:
                raise ValueError("Visnings-Excel har feil arkrekkefølge")
            if saved.views != expected.views:
                raise ValueError("Visnings-Excel har endret arbeidsbokvisning")
            for wanted_sheet, actual_sheet in zip(expected, saved, strict=True):
                for dimension in wanted_sheet.column_dimensions.values():
                    dimension.reindex()
                if (
                    actual_sheet.sheet_state != wanted_sheet.sheet_state
                    or actual_sheet.sheet_view != wanted_sheet.sheet_view
                    or actual_sheet.sheet_format != wanted_sheet.sheet_format
                    or actual_sheet.auto_filter != wanted_sheet.auto_filter
                    or actual_sheet.merged_cells != wanted_sheet.merged_cells
                    or set(actual_sheet.row_dimensions) != set(wanted_sheet.row_dimensions)
                    or set(actual_sheet.column_dimensions) != set(wanted_sheet.column_dimensions)
                    or any(
                        dict(actual_sheet.row_dimensions[k]) != dict(d)
                        for k, d in wanted_sheet.row_dimensions.items()
                    )
                    or any(
                        dict(actual_sheet.column_dimensions[k]) != dict(d)
                        for k, d in wanted_sheet.column_dimensions.items()
                    )
                ):
                    raise ValueError(
                        "Visnings-Excel har endret synlighet/oppsett: " + actual_sheet.title
                    )
                for wanted_row, actual_row in zip(wanted_sheet, actual_sheet, strict=True):
                    for wanted, actual in zip(wanted_row, actual_row, strict=True):
                        value = None if wanted.value == "" else wanted.value
                        if (
                            actual.value != value
                            or type(actual.value) is not type(value)
                            or isinstance(value, str)
                            and actual.data_type != "s"
                            or actual.data_type == "f"
                            or actual.number_format != wanted.number_format
                            or copy(actual.font) != copy(wanted.font)
                            or copy(actual.fill) != copy(wanted.fill)
                            or copy(actual.border) != copy(wanted.border)
                            or copy(actual.alignment) != copy(wanted.alignment)
                            or copy(actual.protection) != copy(wanted.protection)
                            or (actual.hyperlink.location if actual.hyperlink else None)
                            != (wanted.hyperlink.location if wanted.hyperlink else None)
                            or (actual.hyperlink.target if actual.hyperlink else None)
                            != (wanted.hyperlink.target if wanted.hyperlink else None)
                        ):
                            raise ValueError(
                                "Visnings-Excel avviker i celle/visning/tekststykke/lenke: "
                                + actual_sheet.title
                                + "!"
                                + actual.coordinate
                            )
        finally:
            saved.close()
    finally:
        expected.close()
    return report


def export_excel(report: PilotReport, output: Path) -> None:
    if output.exists() or output.suffix.lower() != ".xlsx":
        raise ValueError("Velg nytt Excel-filnavn")
    frozen = report_to_json(report)
    book = _book(report, frozen)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        book.save(temporary)
        if read_excel(temporary) != report:
            raise ValueError("Skrevet Excel avviker")
        _publish(temporary, output)
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def export_pair(report: PilotReport, frozen: Path, excel: Path) -> None:
    if frozen.resolve() == excel.resolve() or frozen.exists() or excel.exists():
        raise ValueError("Velg to forskjellige nye filnavn")
    export_json(report, frozen)
    export_excel(report, excel)
    if read_json(frozen) != read_excel(excel):
        raise ValueError("Skrevne JSON og Excel avviker")


def export_originals(report: PilotReport, directory: Path) -> None:
    """Gjenfinn faktiske originaler offline; ingen ukontrollert eksportsti fra input."""
    report_to_json(report)
    directory.mkdir(parents=True, exist_ok=False)
    old = read_omtale(report.basis.o2b_json)
    entries = [(s.name, s.content_base64) for s in report.basis.sources]
    entries += [("o2b/" + s.name, s.content_base64) for s in old.basis.sources]
    for name, content in entries:
        if content is None:
            continue
        path = (directory / name).resolve()
        if not path.is_relative_to(directory.resolve()):
            raise ValueError("Originalsti utenfor eksportkatalog")
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open("xb") as handle:
            handle.write(base64.b64decode(content, validate=True))
