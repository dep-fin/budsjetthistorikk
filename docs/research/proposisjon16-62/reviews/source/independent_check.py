"""Read-only candidate audit; own expectations from original HTML, not product parser.

Run from candidate root with bundled Python and --root --candidate --base.
Candidate may be any locally available commit with the canonical candidate tree.
Temporary exports stay in .local/source-check; evidence stays beside this script.
"""

import argparse
import base64
import gzip
import hashlib
import json
import os
import re
import subprocess
import sys
from dataclasses import asdict
from pathlib import Path

from lxml import etree, html
from openpyxl import load_workbook
from pypdf import PdfReader


def digest(raw):
    return hashlib.sha256(raw).hexdigest()


def clean(node):
    return " ".join(node.text_content().split())


def number(value):
    return int(re.sub(r"\D", "", value))


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--root", required=True, type=Path)
    p.add_argument("--candidate", required=True)
    p.add_argument("--base", required=True)
    p.add_argument("--output", default=".local/source-check/run1")
    a = p.parse_args()
    root = a.root.resolve()
    os.chdir(root)
    sys.path.insert(0, str(root))
    from analyse import proposisjon16 as t5
    from analyse import visningspilot as t1
    from analyse.proposisjon16_excel import export_originals, export_pair, read_excel
    from innlesning.visningspilot_model import PilotSelection

    def git(*args):
        return subprocess.check_output(["git", *args], cwd=root)

    candidate = git("rev-parse", a.candidate + "^{commit}").decode().strip()
    tree = git("rev-parse", a.candidate + "^{tree}").decode().strip()
    assert tree == "df0eed6d6e8a0b2a07962e641ae05d1194a3304d"
    base = git("rev-parse", a.base + "^{commit}").decode().strip()
    assert base == "2f4ee30c9d8074076fc8042e208c4a4b6345db6b"
    git("diff", "--exit-code", "--quiet", a.candidate)
    out = root / a.output
    out.mkdir(parents=True, exist_ok=True)
    record = {
        "candidate": candidate,
        "tree": tree,
        "base": base,
        "method": (
            "Independent traversal of original HTML table DOM and signed numeric text; "
            "product APIs used only after expected values exist."
        ),
    }
    original_path = "docs/research/omtale-o1/originals/prop16s-ch4.html"
    raw = (root / original_path).read_bytes()
    doc = html.fromstring(raw.decode("utf8"))
    expected = []
    groups = []
    for table_index, table in enumerate(doc.xpath("//table")[:2]):
        chapter = None
        rows = table.xpath("./tbody/tr")
        for i, row in enumerate(rows):
            cells = row.xpath("./td")
            values = [clean(c) for c in cells]
            if not values:
                continue
            if values[0].isdigit():
                chapter = values[0]
            if not values[1].isdigit():
                continue
            account = chapter + "/" + str(int(values[1]))
            after_cells = rows[i + 1].xpath("./td")
            text = clean(after_cells[2])
            before_text, after_text = re.fullmatch(r"fra kr (.+) til kr (.+)", text).groups()
            before, after = number(before_text), number(after_text)
            delta = number(values[3]) * (-1 if "reduseres med" in values[2] else 1)
            assert before + delta == after
            groups.append(
                {
                    "account": account,
                    "side": "utgift" if table_index == 0 else "inntekt",
                    "before": before,
                    "delta": delta,
                    "after": after,
                }
            )
            for meaning, amount, node in [
                ("for_belop", before, after_cells[2]),
                ("endring", delta, cells[3]),
                ("etter_belop", after, after_cells[2]),
            ]:
                expected.append(
                    {
                        "account": account,
                        "meaning": meaning,
                        "amount_kr": amount,
                        "cell": clean(node),
                        "locator": doc.getroottree().getpath(node),
                    }
                )
    assert len(groups) == 37 and len(expected) == 111
    assert sum(g["side"] == "utgift" for g in groups) == 27
    assert len({g["account"].split("/")[0] for g in groups}) == 15
    old_raw = gzip.decompress(
        (root / "tests/fixtures/t2/authentic-t1.json.gz").read_bytes()
    ).decode("utf8")
    old = t1.report_from_json(old_raw)
    basis = t5.import_basis(t1.report_to_json(old), root)
    assert basis.t1_basis == old.basis
    assert json.loads(json.dumps(asdict(basis.t1_basis))) == json.loads(old_raw)["report"]["basis"]
    all_data = t5.contents(basis)
    prop = next(d for d in all_data.documents if d.key == "prop16")
    assert prop.structured_coverage == "delvis_1429"
    observations = {
        (o.chapter + "/" + o.post, o.meaning): o
        for o in all_data.observations
        if o.figure_type == "regjeringsforslag"
    }
    assert len(observations) == 111
    for field in expected:
        o = observations[(field["account"], field["meaning"])]
        assert o.amount_kr == field["amount_kr"] and o.pointer.locator == field["locator"]
        assert o.pointer.artifact_id == digest(raw)
        assert o.budget_year == 2024 and o.currency == "NOK" and o.multiplier == "1"
        field["observation_id"] = o.id
    zero = observations[("4400/2", "etter_belop")]
    assert zero.amount_kr == 0 and not any(k == "amount_kr" for k, v in zero.null_reasons)
    assert len(basis.observations) == 105 and len(all_data.observations) == 123
    assert all(o.figure_type == "regjeringsforslag" for o in basis.observations)
    accounts = {g["account"] for g in groups} | {
        "1420/66",
        "1420/87",
        "1429/60",
        "1420/39",
        "1420/79",
        "1533/45",
        "1474/50",
        "1474/70",
    }
    inventory = [i for i in all_data.inventory if i.document_id == prop.id]
    assert {i.chapter + "/" + i.post for i in inventory} == accounts and len(accounts) == 45
    ch2_raw = (root / "docs/research/omtale-o1/originals/prop16s-ch2.html").read_bytes()
    ch2 = html.fromstring(ch2_raw.decode("utf8"))
    reference_evidence = {}
    for chapter, post in [("1533", "45")]:
        nodes = [
            n
            for n in ch2.xpath("//p")
            if chapter in clean(n) and re.search(r"post(?:ene)?\s+" + post, clean(n))
        ]
        if not nodes:
            nodes = [
                n
                for n in ch2.xpath("//p")
                if re.search(r"post\s+" + post, clean(n))
                and chapter in clean(n.getparent().getparent())
            ]
        assert nodes, (chapter, post)
        reference_evidence[chapter + "/" + post] = [
            {"locator": ch2.getroottree().getpath(n), "text": clean(n)} for n in nodes
        ]
    ref1474 = next(
        i for i in inventory if (i.chapter, i.post, i.role) == ("1474", "1", "endringstabell")
    )
    reference_row = doc.xpath(ref1474.locator)[0]
    assert "postene 50 og 70" in clean(reference_row)
    for post in ("50", "70"):
        reference_evidence["1474/" + post] = [
            {"locator": ref1474.locator, "text": clean(reference_row)}
        ]
    record["groups"] = groups
    record["fields"] = expected
    record["accounts"] = sorted(accounts)
    record["context_evidence"] = reference_evidence
    record["source_hashes"] = {original_path: digest(raw), "chapter2": digest(ch2_raw)}
    auth_specs = [
        ("57", "II", "1", "./p", ["1420/21"], 8000000),
        ("58", "III", "1", ".//tbody/tr[2]", ["1420/61"], 601000000),
        ("58", "III", "2", ".//tbody/tr[3]", ["1420/62"], 266300000),
        ("59", "IV", "1", "./p", ["1420/39", "1420/69", "1420/79"], 16300000),
        ("60", "V", "1", "./ol/li[1]", ["1420/66"], None),
        ("60", "V", "2", "./ol/li[2]", ["1420/87"], None),
        ("60", "V", "3", "./ol/li[3]", ["1429/60"], None),
        ("61", "VI", "1", "./ol/li[1]", ["1471/1", "4471/3"], 25000000),
        ("61", "VI", "2", "./ol/li[2]", [], None),
    ]
    auth_checks = []
    for anchor, roman, point, selector, auth_accounts, frame in auth_specs:
        section = doc.get_element_by_id("kap2-p" + anchor).getparent()
        node = section.xpath(selector)[0]
        aobj = next(x for x in basis.authorizations if (x.roman, x.point) == (roman, point))
        assert aobj.locator == doc.getroottree().getpath(node)
        assert aobj.text == section.text_content()
        assert aobj.html == html.tostring(section, encoding="unicode", with_tail=False)
        assert list(aobj.accounts) == auth_accounts and aobj.frame_kr == frame
        if frame is not None:
            million = re.search(r"(\d+(?:,\d+)?)\s*mill", clean(node))
            assert million and int(float(million[1].replace(",", ".")) * 1000000) == frame
        assert aobj.status == "kildebundet_inventar_ikke_bevilgningsdelta"
        auth_checks.append(
            {
                "roman": roman,
                "point": point,
                "locator": aobj.locator,
                "frame_kr": frame,
                "accounts": auth_accounts,
                "raw_point": clean(node),
            }
        )
    assert len(basis.authorizations) == 9
    record["authorizations"] = auth_checks
    record["parents"] = []
    parents = []
    parent_bytes = []
    for path, sha in [
        (
            "docs/research/visningspilot-41/originals/session-2024-2025-9cfe.xml.gz",
            "9cfeaef59510a54ff8b2a695e0be368e67cfbab574f8d3f466c0c2f0c3c92680",
        ),
        (
            "docs/research/to-visninger-37/originals/aar/session-2024-2025.xml.gz",
            "c5f3b90046d23a0e4cacb24f681a01d7e5fd4468ab68fc3d646dc832bb9a4b42",
        ),
    ]:
        original = gzip.decompress((root / path).read_bytes())
        assert digest(original) == sha and len(original) == 7479943
        parents.append(etree.fromstring(original))
        parent_bytes.append(original)
        record["parents"].append(
            {
                "path": path,
                "compressed_sha256": digest((root / path).read_bytes()),
                "uncompressed_sha256": sha,
                "bytes": len(original),
            }
        )
    record["old_excerpt_bindings"] = []
    for index, path in [(0, "tests/fixtures/b6/40034048"), (1, "tests/fixtures/b8/40034049")]:
        metadata = json.loads((root / (path + ".json")).read_text("utf8"))
        child = (root / (path + ".xml")).read_bytes()
        parent = parent_bytes[index]
        start, length = metadata["object_start_byte_0"], metadata["object_bytes"]
        obj = parent[start : start + length]
        prefix = child.index(obj)
        suffix = len(child) - prefix - length
        assert child == parent[:prefix] + obj + parent[-suffix:]
        assert digest(child) == metadata["sha256"]
        assert digest(parent) == metadata["source_response_sha256"]
        record["old_excerpt_bindings"].append(
            {
                "decision": metadata["decision_id"],
                "prefix": prefix,
                "object_start": start,
                "object_bytes": length,
                "suffix": suffix,
                "excerpt_sha256": digest(child),
                "parent_sha256": digest(parent),
            }
        )
    ns = {"d": "http://data.stortinget.no"}
    branch_specs = [
        ("I", "I", "100927", ["40034048"]),
        ("I", "I", "100889", ["40034220"]),
        ("II", "II", "100889", ["40034221"]),
        ("III", "III", "100889", ["40034222"]),
        ("IV", "IV", "100889", ["40034223"]),
        ("V.1-2", "V", "100889", ["40034224"]),
        ("V.3", "II", "100927", ["40034049"]),
        ("VI", "VI", "100889", ["40034225"]),
    ]
    branch_checks = []
    for point, committee, case, ids in branch_specs:
        branch = next(b for b in basis.branches if (b.point, b.case_id) == (point, case))
        assert branch.committee_point == committee and list(branch.decision_ids) == ids
        assert "1533/45" not in branch.accounts
        evidence = []
        for did in ids:
            parent = parents[0] if did == "40034048" else parents[1]
            node = parent.xpath(".//d:stortingsvedtak[d:id='" + did + "']", namespaces=ns)[0]
            assert node.findtext("d:sak_id", namespaces=ns) == case
            text = node.findtext("d:stortingsvedtak_tekst", namespaces=ns)
            decision_html = html.fromstring(text)
            assert committee.lower() in decision_html.xpath(
                "//@data-strtngt-romertallseksjon-romertall"
            ), (did, committee)
            evidence.append(
                {
                    "decision": did,
                    "case": case,
                    "locator": parent.getroottree().getpath(node),
                    "roman": committee,
                }
            )
        if point == "I":
            assert set(branch.accounts) == {
                g["account"]
                for g in groups
                if (g["account"].startswith("1429/") == (case == "100927"))
            }
        if case == "100889":
            assert branch.amount_stage == "ikke_strukturert_T6"
        branch_checks.append(
            {"point": point, "case": case, "accounts": list(branch.accounts), "evidence": evidence}
        )
    record["branches"] = branch_checks
    representations = []
    for f in ["inns95s.html", "inns95s-20261022-reused.html"]:
        b = (root / "docs/research/to-visninger-37/originals/forgrening" / f).read_bytes()
        d = html.fromstring(b.decode("utf8"))
        assert "beriktiget" in clean(d).lower()
        sections = [clean(d.get_element_by_id("m" + str(i))) for i in range(1, 5)]
        representations.append((f, digest(b), sections))
    assert (
        representations[0][1] != representations[1][1]
        and representations[0][2] == representations[1][2]
    )
    record["beriktiget"] = [
        {"file": f, "sha256": sha, "main_sections": "m1-m4 normalized text equal"}
        for f, sha, _ in representations
    ]
    pdf = PdfReader(root / "docs/research/omtale-o1/originals/prop16s.pdf")
    pdf_pages = [{"page": i + 1, "text": page.extract_text()} for i, page in enumerate(pdf.pages)]
    (out / "prop16-pdf-pages.json").write_text(
        json.dumps(pdf_pages, ensure_ascii=False, indent=2), encoding="utf8"
    )
    table_pages = [
        r["page"]
        for r in pdf_pages
        if "Endringer i statsbudsjettet 2024" in r["text"] or "16,3" in r["text"]
    ]
    assert table_pages
    record["pdf"] = {
        "sha256": digest((root / "docs/research/omtale-o1/originals/prop16s.pdf").read_bytes()),
        "pages": len(pdf.pages),
        "table_or_frame_pages": table_pages,
    }
    cases = {
        "utgift": PilotSelection(years=(2024,), accounts=(("1400", "1"),)),
        "inntekt-null": PilotSelection(years=(2024,), accounts=(("4400", "2"),)),
        "tekst": PilotSelection(years=(2024,), accounts=(("1420", "66"),)),
        "1533": PilotSelection(years=(2024,), accounts=(("1533", "45"),)),
        "1429-1": PilotSelection(years=(2024,), accounts=(("1429", "1"),)),
        "1429-22": PilotSelection(years=(2024,), accounts=(("1429", "22"),)),
        "1429": PilotSelection(years=(2024,), chapters=("1429",)),
        "tomt": PilotSelection(years=(2024,)),
        "2025": PilotSelection(years=(2025,), chapters=("1429",)),
        "uten-aar": PilotSelection(chapters=("1429",)),
        "union": PilotSelection(
            years=(2024,), chapters=("1429",), accounts=(("4400", "2"), ("1420", "22"))
        ),
        "dokument": PilotSelection(mode="document", document_id=prop.id),
    }
    reports = {n: t5.build_report(basis, s) for n, s in cases.items()}
    reports["konto-til-dokument"] = t5.open_document(reports["union"], prop.id)
    reports["retur-konto"] = t5.go_back(reports["konto-til-dokument"])
    reports["dokument-til-konto"] = t5.open_accounts(reports["dokument"], cases["union"])
    reports["retur-dokument"] = t5.go_back(reports["dokument-til-konto"])
    assert (
        reports["retur-konto"] == reports["union"]
        and reports["retur-dokument"] == reports["dokument"]
    )
    record["pairs"] = []
    for name, report in reports.items():
        view = report.view
        if report.selection.mode == "document":
            expect = {o.id for o in observations.values()}
        else:
            s = report.selection
            expect = {
                o.id
                for o in all_data.observations
                if o.budget_year in s.years
                and (o.chapter in s.chapters or (o.chapter, o.post) in s.accounts)
            }
        assert set(view.observation_ids) == expect
        fields = len(expect)
        adopted = sum(
            o.figure_type == "stortingsvedtak" for o in all_data.observations if o.id in expect
        )
        assert len(view.adopted_observation_ids) == adopted
        assert len(view.adopted_event_ids) == adopted // 3
        if name == "1533":
            assert (
                next(d.id for d in all_data.documents if d.key == "inns95") not in view.document_ids
            )
        frozen, excel = out / (name + ".json"), out / (name + ".xlsx")
        if not frozen.exists() and not excel.exists():
            export_pair(report, frozen, excel)
        assert t5.read_json(frozen) == read_excel(excel) == report
        assert report.basis.t1_basis == old.basis
        if name == "dokument":
            book = load_workbook(excel, read_only=False, data_only=False)
            rows = list(book["Forløp"].iter_rows(min_row=2))
            actual_fields = {
                (str(r[4].value) + "/" + str(r[5].value), r[7].value, r[6].value) for r in rows
            }
            assert actual_fields == {(f["account"], f["meaning"], f["amount_kr"]) for f in expected}
            assert len(rows) == 111 and all(r[6].data_type == "n" for r in rows)
            assert all(s.sheet_state == "visible" for s in book)
            assert not any(d.hidden for s in book for d in s.row_dimensions.values())
            assert not any(d.hidden for s in book for d in s.column_dimensions.values())
            assert book["Kontoer"]["G1"].value == "Historisk T1-strukturstatus"
            assert book["Kontoer"]["H1"].value == "Gjeldende T5-tabellstatus"
            record["independent_visible_cells"] = {
                "file": excel.name,
                "sheet": "Forløp",
                "numeric_fields": len(rows),
                "same_as_original_HTML": True,
                "all_sheets_visible": True,
                "all_dimensions_visible": True,
            }
            book.close()
        record["pairs"].append(
            {
                "name": name,
                "fields": fields,
                "adopted_fields": adopted,
                "events": len(view.adopted_event_ids),
                "json_sha256": digest(frozen.read_bytes()),
                "xlsx_sha256": digest(excel.read_bytes()),
            }
        )
    if not (out / "originals").exists():
        export_originals(reports["dokument"], out / "originals")
    index = json.loads((out / "originals/index.json").read_text("utf8"))
    record["originals"] = index
    from analyse.omtale import report_from_json as read_o2

    o2 = read_o2(old.basis.o2b_json)
    files = []
    for prefix, sources in [
        ("t1/", basis.t1_basis.sources),
        ("o2b/", o2.basis.sources),
        ("", basis.sources),
    ]:
        for source in sources:
            if source.content_base64 is None:
                continue
            name = prefix + source.name
            body = base64.b64decode(source.content_base64, validate=True)
            assert digest(body) == source.sha256 and len(body) == source.byte_count
            files.append({"path": name, "sha256": source.sha256, "bytes": source.byte_count})
    record["original_files"] = files
    assert len(files) == 25
    for item in files:
        recovered = (out / "originals" / item["path"]).read_bytes()
        assert digest(recovered) == item["sha256"] and len(recovered) == item["bytes"]
    manifest = json.loads(
        (root / "docs/research/proposisjon16-62/delivery-manifest.json").read_text("utf8")
    )
    check_manifest = []
    for item in manifest["files"]:
        path = item["path"]
        blob = git("show", candidate + ":" + path)
        disk = (root / path).read_bytes()
        if item["representation"] == "UTF-8/LF":
            disk = disk.replace(b"\r\n", b"\n")
        assert digest(blob) == item["sha256"] and len(blob) == item["bytes"], path
        assert digest(disk) == item["sha256"], path
        assert git("rev-parse", candidate + ":" + path).decode().strip() == item["git_blob_sha1"], (
            path
        )
        # Actual Git filter output for Windows and Linux checkout settings.
        if item["representation"] == "UTF-8/LF":
            for mode in ("true", "false"):
                checkout = git(
                    "-c", "core.autocrlf=" + mode, "cat-file", "--filters", candidate + ":" + path
                )
                assert digest(checkout.replace(b"\r\n", b"\n")) == item["sha256"], (path, mode)
        check_manifest.append(path)
    record["manifest_checked"] = check_manifest
    record["original_blob_checks"] = []
    for source in json.loads(
        (root / "docs/research/proposisjon16-62/source-manifest.json").read_text("utf8")
    )["sources"]:
        name = source["name"]
        if name.startswith(("originals/", "b6/", "b8/")):
            continue
        blob = git("show", candidate + ":" + name)
        assert digest(blob) == source["sha256"], name
        record["original_blob_checks"].append(name)
    evidence = root / "docs/research/proposisjon16-62/reviews/source/independent-control.json"
    evidence.write_text(json.dumps(record, ensure_ascii=False, indent=2) + "\n", encoding="utf8")
    print(
        json.dumps(
            {
                "candidate": candidate,
                "tree": tree,
                "fields": len(expected),
                "accounts": len(accounts),
                "pairs": len(reports),
                "originals": len(files),
                "manifest_files": len(check_manifest),
                "report": str(evidence),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
