"""Felles begrenset filgrense: ingen formelkjøring, atomisk eksport, samme radvalidator."""

import csv
import hashlib
import io
import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from lxml import etree
from lxml.etree import XMLSyntaxError
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Protection
from openpyxl.utils import coordinate_to_tuple, get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from .maintenance_store import Diagnostic, reading

MAX_BYTES = 8 * 1024 * 1024
MAX_ROWS = 10000
MAX_CELL = 30000
MAX_UNPACKED_BYTES = 64 * 1024 * 1024
_SHEET_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


class FileLimitError(ValueError):
    """Filen avvises før ubegrenset innlesning eller utpakking."""


def _xlsx_rows(archive: ZipFile, worksheet_path: str, columns: int) -> int:
    """Kontroller fysiske koordinater før openpyxl kan fylle ut tomme celler."""
    parser = etree.XMLPullParser(events=("start", "end"), resolve_entities=False, no_network=True)
    row = previous_row = column = cell_count = row_count = unpacked = cell_length = 0
    with archive.open(worksheet_path) as source:
        while chunk := source.read(64 * 1024):
            unpacked += len(chunk)
            if unpacked > MAX_UNPACKED_BYTES:
                raise ValueError("For stor ZIP-utpakking")
            parser.feed(chunk)
            for event, element in parser.read_events():
                if event == "start":
                    if element.tag == f"{_SHEET_NS}row":
                        row_count += 1
                        row = int(element.get("r", previous_row + 1))
                        if not previous_row < row <= MAX_ROWS + 1 or row_count > MAX_ROWS + 1:
                            raise ValueError("For mange eller ugyldig ordnede rader")
                        previous_row, column = row, 0
                    elif element.tag == f"{_SHEET_NS}c":
                        cell_count += 1
                        coordinate = element.get("r")
                        cell_row, cell_column = (
                            coordinate_to_tuple(coordinate) if coordinate else (row, column + 1)
                        )
                        if (
                            not 1 <= cell_row <= MAX_ROWS + 1
                            or not column < cell_column <= columns
                            or cell_count > (MAX_ROWS + 1) * columns
                        ):
                            raise ValueError("For mange rader, kolonner eller celler")
                        if cell_row != row:
                            raise ValueError("Cellekoordinat stemmer ikke med fysisk rad")
                        column, cell_length = cell_column, 0
                else:
                    if element.tag in {f"{_SHEET_NS}t", f"{_SHEET_NS}v", f"{_SHEET_NS}f"}:
                        cell_length += len(element.text or "")
                        if cell_length > MAX_CELL:
                            raise ValueError("For lang celle")
                    # Ingen XML-rader eller celleinnhold beholdes i denne kontrollen.
                    element.clear()
                    parent = element.getparent()
                    if parent is not None:
                        parent.remove(element)
        parser.close()
    return previous_row


def read_file(
    path: Path, headers: tuple[str, ...], sheet: str
) -> tuple[bytes, list[dict[str, str]], list[Diagnostic]]:
    if path.stat().st_size > MAX_BYTES:
        raise FileLimitError("Arbeidsfil overstiger 8 MiB")
    with path.open("rb") as source:
        raw = source.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise FileLimitError("Arbeidsfil overstiger 8 MiB")
    rows: list[dict[str, str]] = []
    errors: list[Diagnostic] = []
    try:
        values: list[list[str]] = []
        if path.suffix.lower() == ".csv":
            csv.field_size_limit(MAX_CELL)
            for n, row in enumerate(
                csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"), 1
            ):
                if n > MAX_ROWS + 1:
                    raise ValueError("For mange rader")
                values.append(
                    [
                        v[1:]
                        if v.startswith("'")
                        and v[1:2] in {"=", "+", "-", "@", "'", " ", "\t", "\r", "\n"}
                        else v
                        for v in row
                    ]
                )
        elif path.suffix.lower() == ".xlsx":
            with ZipFile(io.BytesIO(raw)) as archive:
                members = archive.infolist()
                if len(members) > 128 or sum(m.file_size for m in members) > MAX_UNPACKED_BYTES:
                    raise ValueError("For stor ZIP-utpakking")
                if any(m.filename.endswith("vbaProject.bin") for m in members):
                    raise ValueError("Makroer er ikke tillatt")
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                if book.sheetnames != [sheet]:
                    raise ValueError(f"Krever bare arket {sheet}")
                page = book[sheet]
                worksheet_path = getattr(page, "_worksheet_path", None)
                if not isinstance(worksheet_path, str):
                    raise ValueError("Krever et vanlig arbeidsark")
                # Samme arkdel som read_only-leseren bruker, også ved en annen ZIP-bane.
                with ZipFile(io.BytesIO(raw)) as archive:
                    actual_rows = _xlsx_rows(archive, worksheet_path, len(headers))
                page.reset_dimensions()
                for n, cells in enumerate(
                    page.iter_rows(min_row=1, min_col=1, max_row=actual_rows, max_col=len(headers)),
                    1,
                ):
                    if n > MAX_ROWS + 1 or len(cells) > len(headers):
                        raise ValueError("For mange rader eller kolonner")
                    row_values = []
                    for column_number, c in enumerate(cells):
                        if c.data_type in {"f", "e"}:
                            errors.append(
                                Diagnostic(
                                    sheet,
                                    n,
                                    headers[column_number],
                                    f"Formel/feilcelle ikke tillatt: {c.coordinate}",
                                )
                            )
                        if c.value is not None and not isinstance(c.value, str | int):
                            errors.append(
                                Diagnostic(
                                    sheet,
                                    n,
                                    headers[column_number],
                                    f"Bare tekst og heltall er tillatt: {c.coordinate}",
                                )
                            )
                        value = str(c.value) if c.value is not None else ""
                        if len(value) > MAX_CELL:
                            raise ValueError("For lang celle")
                        row_values.append(value)
                    values.append(row_values)
            finally:
                book.close()
        else:
            raise ValueError("Krever .xlsx eller UTF-8 .csv")
        if not values or tuple(values[0]) != headers:
            raise ValueError("Kolonnenavn eller rekkefølge stemmer ikke med filformatet")
        for n, row in enumerate(values[1:], 2):
            if not any(row):
                continue
            if len(row) > len(headers) or any(len(v) > MAX_CELL for v in row):
                errors.append(
                    Diagnostic(sheet, n, "rad", "For mange kolonner eller for lang celle")
                )
                continue
            rows.append(
                {
                    **dict(zip(headers, row + [""] * (len(headers) - len(row)), strict=True)),
                    "__row__": str(n),
                }
            )
    except (
        ValueError,
        UnicodeError,
        csv.Error,
        BadZipFile,
        KeyError,
        OSError,
        InvalidFileException,
        ParseError,
        XMLSyntaxError,
    ) as error:
        errors.append(Diagnostic(sheet, 0, "fil", str(error)))
    return raw, rows, errors


def safe_output(database: Path, output: Path) -> None:
    resolved = output.resolve()
    if resolved == database.resolve():
        raise ValueError("Eksport kan ikke overskrive databasen")
    if database.exists():
        with reading(database) as connection:
            if resolved.is_file():
                with resolved.open("rb") as existing:
                    digest = hashlib.file_digest(existing, "sha256").hexdigest()
                if connection.execute("SELECT 1 FROM sources WHERE sha256=?", (digest,)).fetchone():
                    raise ValueError("Eksport kan ikke overskrive en original kildefil")
            for table in ("maintenance_imports", "control_evidence"):
                if connection.execute(
                    "SELECT 1 FROM sqlite_master WHERE name=?", (table,)
                ).fetchone():
                    if any(
                        Path(r[0]).resolve() == resolved
                        for r in connection.execute(f"SELECT filename FROM {table}")
                    ):
                        raise ValueError("Eksport kan ikke overskrive bevart import-/kildefil")


def write_file(
    database: Path,
    output: Path,
    sheet: str,
    headers: tuple[str, ...],
    rows: list[dict[str, str]],
    editable: frozenset[str] = frozenset(),
) -> None:
    safe_output(database, output)
    if output.suffix.lower() not in {".csv", ".xlsx"}:
        raise ValueError("Krever .xlsx eller .csv")
    if len(rows) > MAX_ROWS or any(len(v) > MAX_CELL for row in rows for v in row.values()):
        raise ValueError("Eksport overstiger arbeidsfilgrensen")
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary: Path | None = None
    try:
        with NamedTemporaryFile(dir=output.parent, suffix=output.suffix, delete=False) as handle:
            temporary = Path(handle.name)
        if output.suffix.lower() == ".csv":
            with temporary.open("w", encoding="utf-8-sig", newline="") as target:
                writer = csv.writer(target, delimiter=";")
                writer.writerow(headers)
                for row in rows:
                    writer.writerow(
                        [
                            "'" + v
                            if v[:1] in {"=", "+", "-", "@", "'", " ", "\t", "\r", "\n"}
                            else v
                            for v in (row.get(h, "") for h in headers)
                        ]
                    )
        else:
            book = Workbook()
            page = book.active
            assert page is not None
            page.title = sheet
            for n, row in enumerate([dict(zip(headers, headers, strict=True)), *rows], 1):
                for col, header in enumerate(headers, 1):
                    c = page.cell(n, col, row.get(header, ""))
                    c.data_type = "s"
                    c.protection = Protection(locked=n == 1 or header not in editable)
                    if n == 1:
                        c.font = Font(name="Arial", bold=True, color="FFFFFF")
                        c.fill = PatternFill("solid", fgColor="181C62")
                page.row_dimensions[n].height = 18
            page.freeze_panes = "A2"
            page.auto_filter.ref = page.dimensions
            for column_number in range(1, len(headers) + 1):
                page.column_dimensions[get_column_letter(column_number)].width = 24
            # Låsing er veiledning; importen validerer alle bindinger mot lagret eksport.
            book.save(temporary)
            book.close()
        if temporary.stat().st_size > MAX_BYTES:
            raise ValueError("Eksport overstiger 8 MiB")
        os.replace(temporary, output)
    finally:
        if temporary is not None:
            temporary.unlink(missing_ok=True)
