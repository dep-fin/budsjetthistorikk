"""T4/#48: reproduser autentiske filpar, negative prøver og kompatibilitetsbevis.

python -m ops.fleraarspilot_proeve --output work/t4-kontroll --previous-root <hovedrepo>
    --legacy-b6 <ekte accept-b6.xlsx>
"""

import argparse
import gzip
import hashlib
import json
import platform
import shutil
import subprocess
import sys
import tempfile
import time
from collections import Counter
from dataclasses import asdict
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook

from analyse import forlopspilot as t3
from analyse.fleraarspilot import (
    build_report,
    contents,
    go_back,
    import_basis,
    open_accounts,
    open_document,
    read_json,
)
from analyse.fleraarspilot_excel import export_originals, export_pair, read_excel
from analyse.forlopspilot_excel import read_excel as read_t3_excel
from innlesning import fleraarspilot as source
from innlesning import forlopspilot as old_source
from innlesning.fleraarspilot_model import Basis, Report
from innlesning.observations import encode, identity
from innlesning.visningspilot_model import PilotSelection

from .forlopspilot_proeve import _old_pair, _snapshot
from .forlopspilot_proeve import source_audit as old_source_audit
from .komitepilot_proeve import serialized_counts
from .visningspilot_proeve import _compatibility, _ordinary, evidence, verify_legacy_b6

MUTATIONS = (
    "year",
    "amount",
    "meaning",
    "date",
    "role",
    "raw-name",
    "locator",
    "parent",
    "recipe",
    "reference",
    "hidden-gaps",
    "continuity",
    "selection",
    "calculation",
    "calculation-ref",
    "old-field",
    "inventory-cap",
    "legacy-reference",
    "extra-field",
)


def selection_cases(basis: Basis) -> list[tuple[str, PilotSelection, tuple[int, int, int]]]:
    """Tellerfasit fra før-kode-kontrakten, uten å beregne forventning fra produktet."""
    cases = []
    for years, prefix, counts in (
        ((2024,), "2024", ((13, 7, 3), (10, 4, 2), (23, 11, 5))),
        ((2025,), "2025", ((1, 1, 1), (1, 1, 1), (2, 2, 2))),
        ((2024, 2025), "both-years", ((14, 8, 4), (11, 5, 3), (25, 13, 7))),
    ):
        for posts, name, count in zip(
            (("1",), ("22",), ("1", "22")), ("post1", "post22", "both"), counts, strict=True
        ):
            cases.append(
                (
                    prefix + "-" + name,
                    PilotSelection(years=years, accounts=tuple(("1429", p) for p in posts)),
                    count,
                )
            )
    cases.extend(
        [
            ("chapter", PilotSelection(years=(2024, 2025), chapters=("1429",)), (25, 13, 7)),
            ("empty", PilotSelection(), (0, 0, 0)),
            (
                "empty-multiyear",
                PilotSelection(years=(2024, 2025), accounts=(("999", "1"),)),
                (0, 0, 0),
            ),
            ("wrong-year", PilotSelection(years=(2023,), chapters=("1429",)), (0, 0, 0)),
            ("text-fullmakt", PilotSelection(years=(2024,), accounts=(("1429", "60"),)), (0, 0, 0)),
            (
                "pairs-no-crossproduct",
                PilotSelection(years=(2024, 2025), accounts=(("1429", "1"), ("315", "22"))),
                (14, 8, 4),
            ),
            (
                "union-chapter-pair",
                PilotSelection(years=(2024, 2025), chapters=("1429",), accounts=(("315", "22"),)),
                (25, 13, 7),
            ),
        ]
    )
    docs = {d.key: d.id for d in contents(basis).documents}
    for key, count in (
        ("vedtak201", (2, 2, 2)),
        ("vedtak341", (2, 2, 2)),
        ("vedtak857", (3, 3, 1)),
        ("vedtak210", (6, 6, 2)),
        ("prop16", (6, 0, 0)),
        ("inns81", (6, 0, 0)),
        ("vedtak211", (0, 0, 0)),
        ("inns95", (0, 0, 0)),
    ):
        cases.append(
            ("document-" + key, PilotSelection(mode="document", document_id=docs[key]), count)
        )
    return cases


def mutate(raw: str, mutation: str) -> str:
    """Merkede syntetiske feil; nye objekt-/konvolutthasher er ikke kildebevis."""
    data = json.loads(raw)
    report = data["report"]
    basis = report["basis"]
    observation = basis["observations"][-1]
    field_values = {
        "year": ("budget_year", 2024),
        "amount": ("amount_kr", 33_408_001),
        "meaning": ("meaning", "endring"),
        "date": ("decision_date", "2025-12-16"),
        "role": ("figure_type", "regjeringsforslag"),
        "raw-name": ("purpose", "synthetic identical old account name"),
    }
    if mutation in field_values:
        key, value = field_values[mutation]
        observation[key] = value
    elif mutation == "locator":
        observation["pointer"]["locator"] = "table[1]/tr[102]/td[4]"
    elif mutation in {"parent", "recipe"}:
        s = next(s for s in basis["sources"] if s["name"] == "derived:40034038")
        meta = json.loads(s["metadata_json"])
        meta["parent_sha256" if mutation == "parent" else "object_start_byte_0"] = (
            "0" * 64 if mutation == "parent" else 5662497
        )
        s["metadata_json"] = encode(meta)
        s["id"] = identity("fleraarskilde-1", {**s, "id": ""})
    elif mutation == "reference":
        basis["contexts"][-1]["evidence_ids"] = ["synthetic-dangling"]
        c = basis["contexts"][-1]
        c["id"] = identity("fleraarskontekst-1", {**c, "id": ""})
    elif mutation == "hidden-gaps":
        basis["issues"] = [i for i in basis["issues"] if "2025" not in i]
    elif mutation == "continuity":
        basis["issues"] = [i.replace("ukjent", "godkjent_kontinuitet") for i in basis["issues"]]
    elif mutation == "selection":
        report["selection"]["years"] = [2025]
    elif mutation in {"calculation", "calculation-ref"}:
        report["calculations"][0]["amount_kr" if mutation == "calculation" else "input_ids"] = (
            196_581_000 if mutation == "calculation" else ["synthetic-2025-level"]
        )
    elif mutation == "old-field":
        old = json.loads(basis["legacy_json"])
        old["report"]["basis"]["observations"][0]["amount_kr"] += 1
        old["report"]["basis"]["id"] = identity(
            "visningsgrunnlag-1", {**old["report"]["basis"], "id": ""}
        )
        old["report"]["id"] = identity("visningspilot-1", {**old["report"], "id": ""})
        basis["legacy_json"] = encode(old)
    elif mutation == "inventory-cap":
        basis["inventory"] = [basis["inventory"][0]] * 4097
    elif mutation == "legacy-reference":
        next(p for p in basis["payloads"] if p["sha256"] == source.CATALOG[source.PARENT][1])[
            "legacy_source_id"
        ] = "synthetic-wrong-parent-reference"
    elif mutation == "extra-field":
        report["continuity"] = "synthetic-approved"
    else:
        raise ValueError("Ukjent syntetisk mutasjon")
    for c in report["calculations"]:
        c["id"] = identity("forlopskontroll-1", {**c, "id": ""})
    basis["id"] = identity("fleraarsgrunnlag-1", {**basis, "id": ""})
    report["id"] = identity("visningspilot-4", {**report, "id": ""})
    return encode(data)


def bad_sources(root: Path, condition: str) -> Path:
    for name in (*old_source.CATALOG, *source.CATALOG):
        dest = root / name
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(name, dest)
    path = root / (source.PARENT if condition == "corrupt" else source.META)
    if condition in {"missing", "failed"}:
        path.unlink()
        if condition == "failed":
            path.mkdir()
    elif condition == "corrupt":
        path.write_bytes(b"synthetic invalid gzip")
    elif condition == "contradictory":
        path.write_bytes(path.read_bytes().replace(b"7479943", b"7479942"))
    else:
        raise ValueError("Ukjent syntetisk kildefeil")
    return root


def source_audit() -> dict[str, Any]:
    parent = gzip.decompress(Path(source.PARENT).read_bytes())
    raw = parent[:317] + parent[5662496:5779401] + b"\n" + parent[7479888:]
    assert len(parent) == 7479943 and hashlib.sha256(parent).hexdigest() == source.PARENT_SHA
    assert len(raw) == 117278 and hashlib.sha256(raw).hexdigest() == source.ENVELOPE_SHA
    assert raw == Path("docs/research/to-visninger-37/originals/aar/40034038.xml").read_bytes()
    return {
        "t3": old_source_audit(),
        "source_files": [evidence(Path(n), Path(".")) for n in source.CATALOG],
        "parent_bytes": len(parent),
        "parent_sha256": hashlib.sha256(parent).hexdigest(),
        "derived_bytes": len(raw),
        "derived_sha256": hashlib.sha256(raw).hexdigest(),
        "recipe": "parent[:317]+parent[5662496:5779401]+LF+parent[7479888:]",
        "representation": "avledet_ikke_http",
        "original_derived_byteidentical": True,
    }


def mutate_display(source_file: Path, output: Path, mutation: str) -> None:
    """Syntetisk feil i skrevet presentasjon; offlinegrunnlaget røres ikke."""
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(source_file) as original, ZipFile(output, "w") as changed:
        for entry in original.infolist():
            raw = original.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                root = etree.fromstring(raw)
                if mutation == "amount":
                    root.find(f".//{ns}c[@r='G11']/{ns}v").text = "999"
                elif mutation == "hidden-row":
                    root.find(f".//{ns}row[@r='11']").set("hidden", "1")
                elif mutation == "hidden-column":
                    root.find(f".//{ns}col").set("hidden", "1")
                elif mutation == "hyperlink":
                    root.find(f".//{ns}hyperlink").set("location", "'Omtale'!A1")
                elif mutation == "filter":
                    root.find(ns + "autoFilter").set("ref", "A1:N2")
                raw = etree.tostring(root)
                if mutation == "doctype":
                    raw = b'<!DOCTYPE worksheet [<!ENTITY synthetic "bad">]>' + raw
            if mutation == "numberformat" and entry.filename == "xl/styles.xml":
                raw = raw.replace(b'numFmtId="3"', b'numFmtId="1"')
            if mutation == "hidden-sheet" and entry.filename == "xl/workbook.xml":
                raw = raw.replace(b'state="visible"', b'state="hidden"', 1)
            changed.writestr(entry, raw)


def save_pair(
    output: Path, name: str, report: Report, counts: tuple[int, int, int]
) -> dict[str, Any]:
    frozen, excel = output / (name + ".json"), output / (name + ".xlsx")
    export_pair(report, frozen, excel)
    saved = read_json(frozen)
    assert saved == report  # Excel er allerede lukket/gjenlest av skriveren.
    actual = (
        len(saved.view.observation_ids),
        len(saved.view.adopted_observation_ids),
        len(saved.view.adopted_event_ids),
    )
    assert actual == counts
    book = load_workbook(excel, read_only=True, data_only=False)
    try:
        snapshot = "".join(
            str(r[1]) for r in book["Offlinegrunnlag"].iter_rows(min_row=2, values_only=True)
        )
        assert snapshot == frozen.read_text("utf8")
        shapes = {s.title: {"rows": s.max_row, "columns": s.max_column} for s in book}
        groups = [(r[1], r[0]) for r in book["Forløp"].iter_rows(min_row=2, values_only=True)]
        assert groups == sorted(groups, key=lambda x: (x[0], int(x[1].split("/")[1])))
    finally:
        book.close()
    objects = serialized_counts(snapshot)
    assert objects == serialized_counts(frozen.read_text("utf8"))
    assert objects["new"]["serialized_observation_objects"] == 7
    assert all(n == 1 for n in objects["new"]["observation_occurrences"].values())
    assert all(n == 1 for n in objects["new"]["payload_occurrences"].values())
    with ZipFile(excel) as package:
        snapshot_parts = [n for n in package.namelist() if n == "xl/worksheets/sheet7.xml"]
        assert len(snapshot_parts) == 1
    return {
        "name": name,
        "selection": asdict(saved.selection),
        "report_id": saved.id,
        "basis_id": saved.basis.id,
        "expected": counts,
        "actual": actual,
        "calculations": [asdict(c) for c in saved.calculations],
        "shapes": shapes,
        "serialized_counts": objects,
        "excel_snapshot_count": 1,
        "files": [evidence(frozen, output), evidence(excel, output)],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--previous-root", type=Path, required=True)
    parser.add_argument("--legacy-b6", type=Path, required=True)
    args = parser.parse_args()
    started = time.monotonic()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    temporary = output / "reader-tmp"
    temporary.mkdir()
    tempfile.tempdir = str(temporary)
    raw_t1 = gzip.decompress(Path("tests/fixtures/t2/authentic-t1.json.gz").read_bytes()).decode(
        "utf8"
    )
    basis = import_basis(Path("."), t1_json=raw_t1)
    both_selection = PilotSelection(years=(2024, 2025), chapters=("1429",))
    both = build_report(basis, both_selection)
    audit = source_audit()
    records = [save_pair(output, "both", both, (25, 13, 7))]
    for name, selection, count in selection_cases(basis):
        records.append(save_pair(output, name, build_report(basis, selection), count))
        (output / "pairs-journal.json").write_text(
            json.dumps(records, ensure_ascii=False, indent=2), "utf8"
        )
    docs = {d.key: d.id for d in contents(basis).documents}
    opened = open_document(both, docs["vedtak201"])
    records.append(save_pair(output, "opened-201", opened, (2, 2, 2)))
    assert go_back(read_excel(output / "opened-201.xlsx")) == both
    account = open_accounts(opened, PilotSelection(years=(2025,), accounts=(("1429", "22"),)))
    records.append(save_pair(output, "document-to-account", account, (1, 1, 1)))
    records.append(save_pair(output, "return-document", go_back(account), (2, 2, 2)))
    records.append(save_pair(output, "return-account", go_back(opened), (25, 13, 7)))
    records.append(save_pair(output, "offline", read_json(output / "both.json"), (25, 13, 7)))
    assert (output / "both.json").read_bytes() == (output / "offline.json").read_bytes()
    export_originals(read_excel(output / "offline.xlsx"), output / "originals")
    script = """import socket, sqlite3, sys
def unavailable(*a, **kw): raise AssertionError('network/database forbidden')
socket.create_connection=unavailable
socket.socket.connect=unavailable
socket.socket.connect_ex=unavailable
sqlite3.connect=unavailable
sys.argv=['analyse.fleraarspilot','--frosset',sys.argv[1],'--rot',sys.argv[2],
          '--json',sys.argv[3],'--excel',sys.argv[4]]
from analyse.fleraarspilot import main
main()
"""
    cli = subprocess.run(
        [
            sys.executable,
            "-c",
            script,
            str(output / "both.json"),
            str(output / "no-original-root"),
            str(output / "cli-offline.json"),
            str(output / "cli-offline.xlsx"),
        ],
        capture_output=True,
        text=True,
        encoding="utf8",
    )
    assert cli.returncode == 0, cli.stdout + cli.stderr
    assert read_json(output / "cli-offline.json") == read_excel(output / "cli-offline.xlsx") == both
    failures = []
    for condition in ("missing", "failed", "corrupt", "contradictory"):
        directory = output / ("synthetic-" + condition)
        directory.mkdir()
        failed = import_basis(bad_sources(directory / "sources", condition))
        empty = build_report(failed, PilotSelection(years=(2024, 2025), accounts=(("999", "1"),)))
        export_pair(empty, directory / "empty.json", directory / "empty.xlsx")
        assert read_json(directory / "empty.json") == read_excel(directory / "empty.xlsx") == empty
        ordinary = _ordinary(directory)
        assert ordinary["observations"] == 222 and ordinary["events"] == 74
        failures.append(
            {
                "synthetic": True,
                "condition": condition,
                "issues": failed.issues,
                "ordinary_fresh_database": ordinary,
            }
        )
    negatives = []
    for mutation in MUTATIONS:
        raw = mutate((output / "both.json").read_text("utf8"), mutation)
        frozen = output / ("synthetic-" + mutation + ".json")
        excel = frozen.with_suffix(".xlsx")
        frozen.write_text(raw, "utf8")
        _snapshot(output / "both.xlsx", excel, raw)
        rejections = {}
        for kind, reader, file in (("json", read_json, frozen), ("xlsx", read_excel, excel)):
            try:
                reader(file)
            except ValueError as error:
                rejections[kind] = str(error)
            else:
                raise AssertionError("Syntetisk feil akseptert: " + mutation)
        negatives.append({"synthetic": True, "mutation": mutation, "rejections": rejections})
    presentation_negatives = []
    for mutation in (
        "amount",
        "hidden-row",
        "hidden-column",
        "hyperlink",
        "filter",
        "doctype",
        "numberformat",
        "hidden-sheet",
    ):
        path = output / ("synthetic-display-" + mutation + ".xlsx")
        mutate_display(output / "both.xlsx", path, mutation)
        try:
            read_excel(path)
        except ValueError as error:
            presentation_negatives.append(
                {"synthetic": True, "mutation": mutation, "rejection": str(error)}
            )
        else:
            raise AssertionError("Syntetisk presentasjon akseptert: " + mutation)
    compatibility = _compatibility(args.previous_root, output / "compatibility")
    for kind, folder in (("T1", "t1-41-608895f"), ("T2", "t2-44-1898ef3")):
        for file in sorted((args.previous_root / "outputs" / folder).glob("*.json")):
            if file.with_suffix(".xlsx").is_file():
                compatibility.append(
                    _old_pair(
                        file.parent,
                        output / "compatibility" / kind / file.stem,
                        file.stem,
                        "ekte_tidligere_" + kind,
                    )
                )
    old_root = args.previous_root / "outputs/t3-46-ccba906"
    for file in sorted(old_root.glob("*.json")):
        if not file.with_suffix(".xlsx").is_file():
            continue
        copied = output / "compatibility/T3" / file.stem
        copied.mkdir(parents=True)
        before = [evidence(p, old_root) for p in (file, file.with_suffix(".xlsx"))]
        for path, pin in zip((file, file.with_suffix(".xlsx")), before, strict=True):
            shutil.copyfile(path, copied / path.name)
            assert hashlib.sha256((copied / path.name).read_bytes()).hexdigest() == pin["sha256"]
        old = t3.read_json(copied / file.name)
        assert read_t3_excel(copied / file.with_suffix(".xlsx").name) == old
        assert before == [evidence(p, old_root) for p in (file, file.with_suffix(".xlsx"))]
        compatibility.append(
            {
                "kind": "ekte_tidligere_T3",
                "name": file.stem,
                "report_id": old.id,
                "original_hashes": before,
                "unchanged": True,
            }
        )
    compatibility.append(verify_legacy_b6(args.legacy_b6, output / "compatibility/B6"))
    manifest = {
        "kind": "T4 #48: autentiske prøver og merkede syntetiske negativer",
        "python": platform.python_version(),
        "lxml": etree.LXML_VERSION,
        "libxml2_runtime": etree.LIBXML_VERSION,
        "libxml2_compiled": etree.LIBXML_COMPILED_VERSION,
        "basis_id": basis.id,
        "source_audit": audit,
        "selections": records,
        "public_cli_offline": {
            "returncode": cli.returncode,
            "stdout": cli.stdout,
            "network_database_forbidden": True,
            "original_root_present": False,
        },
        "independence": failures,
        "negative_mutations": negatives,
        "presentation_negatives": presentation_negatives,
        "compatibility": compatibility,
        "compatibility_counts": dict(Counter(str(x["kind"]) for x in compatibility)),
        "seconds": round(time.monotonic() - started, 3),
        "limitations": [
            "Ingen kontinuitetsavgjørelse, helår, årsbalanse eller drift",
            "T1s arvede immutable kopier telles separat",
            "Native Excel, gammel selvstendig frozen-v1-JSON og 554e-fullrespons ikke prøvd",
        ],
        "files": [evidence(p, output) for p in sorted(output.rglob("*")) if p.is_file()],
    }
    (output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", "utf8"
    )
    print(
        encode(
            {
                "output": str(output),
                "basis_id": basis.id,
                "pairs": len(records),
                "compatibility_counts": manifest["compatibility_counts"],
                "seconds": manifest["seconds"],
            }
        )
    )


if __name__ == "__main__":
    main()
