"""Reproduser T3/#46 fra bevarte byte til faktiske JSON-/Excel-par og journal.

python -m ops.forlopspilot_proeve --output work/t3-delivery/ny-kontroll
Ekstra gamle filer: --previous-root <hovedrepo> --legacy-b6 <ekte accept-b6.xlsx>
"""

import argparse
import gzip
import hashlib
import json
import platform
import shutil
import subprocess
import sys
import tempfile
import time
from collections import Counter
from dataclasses import asdict
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook

from analyse import komitepilot as t2
from analyse import visningspilot as t1
from analyse.forlopspilot import (
    build_report,
    contents,
    go_back,
    import_basis,
    open_accounts,
    open_document,
    read_json,
)
from analyse.forlopspilot_excel import export_originals, export_pair, read_excel
from analyse.komitepilot_excel import read_excel as read_t2_excel
from analyse.visningspilot_excel import read_excel as read_t1_excel
from innlesning.forlopspilot import CATALOG, CONTROL, META, PARENT, RECIPES, derived_xml
from innlesning.forlopspilot_model import Report
from innlesning.observations import encode, identity
from innlesning.visningspilot_model import PilotSelection

from .komitepilot_proeve import serialized_counts
from .visningspilot_proeve import _compatibility, _ordinary, evidence, verify_legacy_b6


def source_audit() -> dict[str, Any]:
    """Uavhengig bytekontroll av forskerens avledninger og gammel kontrollrolle."""
    parent_gzip = Path(PARENT).read_bytes()
    parent = gzip.decompress(parent_gzip)
    old = Path(CONTROL).read_bytes()
    assert hashlib.sha256(old).hexdigest() == CATALOG[CONTROL][2]
    assert old[317:-55] == parent[2552391:3817852]
    derived = {}
    for api_id, recipe in RECIPES.items():
        raw = derived_xml(parent, api_id)
        derived[api_id] = {
            "bytes": len(raw),
            "sha256": hashlib.sha256(raw).hexdigest(),
            "object_start": recipe[0],
            "object_end": recipe[1],
            "object_sha256": recipe[2],
            "separator_hex": recipe[4].hex(),
        }
        if api_id == "40032928":
            assert (
                raw == Path("docs/research/to-visninger-37/originals/aar/40032928.xml").read_bytes()
            )
    sidecar = json.loads(Path("tests/fixtures/b3/40033719.json").read_text("utf8"))
    t1_compressed = Path("tests/fixtures/t2/authentic-t1.json.gz").read_bytes()
    old_t1 = gzip.decompress(t1_compressed)
    assert hashlib.sha256(old_t1).hexdigest() == (
        "99e871e0ab780cbb548e15699f3093621d0f25229e3fa06c29cfdb5552803f53"
    )
    return {
        "parent_gzip": evidence(Path(PARENT), Path(".")),
        "parent_bytes": len(parent),
        "parent_sha256": hashlib.sha256(parent).hexdigest(),
        "derived": derived,
        "old_857_fixture": evidence(Path(CONTROL), Path(".")),
        "old_857_sidecar": sidecar,
        "old_object_byteidentical": True,
        "old_554e_parent_present": False,
        "old_554e_boundary": "opprinnelig oppgitt respons mangler; 66aa er separat representasjon",
        "authentic_t1_gzip_sha256": hashlib.sha256(t1_compressed).hexdigest(),
        "authentic_t1_unpacked_sha256": hashlib.sha256(old_t1).hexdigest(),
        "source_files": [evidence(Path(name), Path(".")) for name in CATALOG],
    }


def _save(
    output: Path, name: str, report: Report, expected: tuple[int, int, int]
) -> dict[str, Any]:
    started = time.monotonic()
    frozen, excel = output / (name + ".json"), output / (name + ".xlsx")
    export_pair(report, frozen, excel)
    # Hver skriver har allerede lukket og gjenlest sin fil mot samme rapport.
    # Her kontrolleres dessuten faktisk Excel-snapshot mot de faktiske JSON-bytene.
    book = load_workbook(excel, read_only=True, data_only=False)
    try:
        snapshot = "".join(
            str(row[1]) for row in book["Offlinegrunnlag"].iter_rows(min_row=2, values_only=True)
        )
        assert snapshot == frozen.read_text("utf8")
        actual = (
            len(report.view.observation_ids),
            len(report.view.adopted_observation_ids),
            len(report.view.adopted_event_ids),
        )
        assert actual == expected
        assert book["Forløp"].max_row - 1 == len(
            {
                o.step_id
                for o in contents(report.basis).observations
                if o.id in report.view.observation_ids
            }
        )
        shapes = {s.title: {"rows": s.max_row, "columns": s.max_column} for s in book}
        long_parts = sum(
            1
            for row in book["Omtale"].iter_rows(min_row=2, values_only=True)
            if row[7] and row[7] > 1
        )
    finally:
        book.close()
    counts = serialized_counts(snapshot)
    assert counts == serialized_counts(frozen.read_text("utf8"))
    assert counts["new"]["serialized_observation_objects"] == len(report.basis.observations)
    assert all(v == 1 for v in counts["new"]["payload_occurrences"].values())
    assert all(v == 1 for v in counts["new"]["observation_occurrences"].values())
    return {
        "name": name,
        "report_id": report.id,
        "basis_id": report.basis.id,
        "selection": asdict(report.selection),
        "expected": expected,
        "actual": actual,
        "calculations": [asdict(c) for c in report.calculations],
        "shapes": shapes,
        "continued_fulltext_rows": long_parts,
        "serialized_counts": counts,
        "seconds": round(time.monotonic() - started, 3),
        "files": [evidence(frozen, output), evidence(excel, output)],
    }


def _snapshot(source: Path, output: Path, raw: str) -> None:
    s = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(source) as original, ZipFile(output, "w") as changed:
        for entry in original.infolist():
            content = original.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet7.xml":
                root = etree.fromstring(content)
                rows = root.find(s + "sheetData")
                assert rows is not None
                for row in list(rows)[1:]:
                    rows.remove(row)
                for n, start in enumerate(range(0, len(raw), 30000), 2):
                    row = etree.SubElement(rows, s + "row", r=str(n))
                    a = etree.SubElement(row, s + "c", r=f"A{n}", t="n")
                    etree.SubElement(a, s + "v").text = str(n - 1)
                    b = etree.SubElement(row, s + "c", r=f"B{n}", t="inlineStr")
                    etree.SubElement(etree.SubElement(b, s + "is"), s + "t").text = raw[
                        start : start + 30000
                    ]
                content = etree.tostring(root)
            changed.writestr(entry, content)


def _negatives(output: Path) -> list[dict[str, Any]]:
    directory = output / "synthetic-mutations"
    directory.mkdir()
    negatives = []
    for mutation in (
        "amount",
        "sign",
        "year",
        "role",
        "round",
        "parent",
        "reference",
        "order",
        "hidden-gaps",
        "absence",
        "calculation",
        "calculation-status",
        "calculation-ref",
        "old-field",
    ):
        data = json.loads((output / "both.json").read_text("utf8"))
        b, r = data["report"]["basis"], data["report"]
        if mutation == "amount":
            b["observations"][0]["amount_kr"] += 1
        elif mutation == "sign":
            b["observations"][3]["amount_kr"] = 1_000_000
        elif mutation in {"year", "role", "round"}:
            key, value = {
                "year": ("budget_year", 2025),
                "role": ("figure_type", "regjeringsforslag"),
                "round": ("round", "nysaldering"),
            }[mutation]
            b["observations"][0][key] = value
        elif mutation == "parent":
            source = next(s for s in b["sources"] if s["name"].startswith("derived:"))
            metadata = json.loads(source["metadata_json"])
            metadata["parent_sha256"] = "0" * 64
            source["metadata_json"] = encode(metadata)
        elif mutation == "reference":
            b["relations"][0]["evidence_ids"] = ["synthetic-dangling"]
        elif mutation == "order":
            b["documents"][0]["date"] = "2024-12-20"
        elif mutation == "hidden-gaps":
            b["issues"] = []
        elif mutation == "absence":
            next(s for s in b["statements"] if "fravaer" in s["role"])["text"] = (
                "synthetic whole-year absence"
            )
        elif mutation == "calculation":
            r["calculations"][0]["amount_kr"] += 1
        elif mutation == "calculation-status":
            r["calculations"][0]["status"] = "godkjent årsbalanse"
        elif mutation == "old-field":
            old = json.loads(b["legacy_json"])
            old["report"]["basis"]["observations"][0]["amount_kr"] += 1
            old["report"]["basis"]["id"] = identity(
                "visningsgrunnlag-1", {**old["report"]["basis"], "id": ""}
            )
            old["report"]["id"] = identity("visningspilot-1", {**old["report"], "id": ""})
            b["legacy_json"] = encode(old)
        else:
            r["calculations"][0]["input_ids"] = []
        for c in r["calculations"]:
            c["id"] = ""
            c["id"] = identity("forlopskontroll-1", c)
        b["id"] = ""
        b["id"] = identity("forlopsgrunnlag-1", b)
        r["id"] = ""
        r["id"] = identity("visningspilot-3", r)
        raw = encode(data)
        frozen = directory / (mutation + ".json")
        excel = directory / (mutation + ".xlsx")
        frozen.write_text(raw, "utf8")
        _snapshot(output / "both.xlsx", excel, raw)
        rejection = {}
        for extension, reader, path in (("json", read_json, frozen), ("xlsx", read_excel, excel)):
            try:
                reader(path)
            except ValueError as error:
                rejection[extension] = str(error)
            else:
                raise AssertionError("Syntetisk rehash akseptert: " + mutation)
        negatives.append(
            {"synthetic": True, "mutation": mutation, "rehashed": True, "rejections": rejection}
        )
    for mutation in ("false-value", "numberformat", "hidden-gaps", "false-link"):
        path = directory / (mutation + "-display.xlsx")
        book = load_workbook(output / "both.xlsx")
        if mutation == "false-value":
            book["Forløp"]["G2"] = 999
        elif mutation == "numberformat":
            book["Forløp"]["G2"].number_format = '"0"'
        elif mutation == "hidden-gaps":
            book["Dekning og avvik"].sheet_state = "hidden"
        else:
            book["Forløp"]["C2"].hyperlink = "https://example.invalid/synthetic"
        book.save(path)
        book.close()
        try:
            read_excel(path)
        except ValueError as error:
            negatives.append(
                {"synthetic": True, "mutation": mutation, "rejections": {"xlsx": str(error)}}
            )
        else:
            raise AssertionError("Syntetisk presentasjon akseptert")
    return negatives


def _old_pair(source: Path, copied: Path, stem: str, kind: str) -> dict[str, Any]:
    copied.mkdir(parents=True, exist_ok=False)
    before = {
        suffix: hashlib.sha256((source / (stem + suffix)).read_bytes()).hexdigest()
        for suffix in (".json", ".xlsx")
    }
    for suffix, pin in before.items():
        shutil.copyfile(source / (stem + suffix), copied / (stem + suffix))
        assert hashlib.sha256((copied / (stem + suffix)).read_bytes()).hexdigest() == pin
    if kind == "ekte_tidligere_T1":
        old = t1.read_json(copied / (stem + ".json"))
        assert read_t1_excel(copied / (stem + ".xlsx")) == old
    else:
        old = t2.read_json(copied / (stem + ".json"))
        assert read_t2_excel(copied / (stem + ".xlsx")) == old
    assert before == {
        suffix: hashlib.sha256((source / (stem + suffix)).read_bytes()).hexdigest()
        for suffix in before
    }
    return {
        "kind": kind,
        "source": str(source / stem),
        "report_id": old.id,
        "original_hashes": before,
        "unchanged": True,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--previous-root", type=Path)
    parser.add_argument("--legacy-b6", type=Path)
    args = parser.parse_args()
    started = time.monotonic()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    temporary = output / "reader-tmp"
    temporary.mkdir()
    tempfile.tempdir = str(temporary)
    audit = source_audit()
    (output / "source-audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n", "utf8"
    )
    raw_t1 = gzip.decompress(Path("tests/fixtures/t2/authentic-t1.json.gz").read_bytes()).decode(
        "utf8"
    )
    (output / "authentic-t1.json").write_bytes(raw_t1.encode("utf8"))
    basis = import_basis(Path("."), t1_json=raw_t1)
    both_selection = PilotSelection(years=(2024,), accounts=(("1429", "1"), ("1429", "22")))
    both = build_report(basis, both_selection)
    cases = [
        ("both", both, (23, 11, 5)),
        (
            "post1",
            build_report(basis, PilotSelection(years=(2024,), accounts=(("1429", "1"),))),
            (13, 7, 3),
        ),
        (
            "post22",
            build_report(basis, PilotSelection(years=(2024,), accounts=(("1429", "22"),))),
            (10, 4, 2),
        ),
        (
            "chapter",
            build_report(basis, PilotSelection(years=(2024,), chapters=("1429",))),
            (23, 11, 5),
        ),
        ("empty", build_report(basis, PilotSelection()), (0, 0, 0)),
        (
            "wrong-year",
            build_report(basis, PilotSelection(years=(2025,), chapters=("1429",))),
            (0, 0, 0),
        ),
        (
            "text-fullmakt",
            build_report(basis, PilotSelection(years=(2024,), accounts=(("1429", "60"),))),
            (0, 0, 0),
        ),
    ]
    docs = {d.key: d.id for d in contents(basis).documents}
    cases.append(
        (
            "pairs-no-crossproduct",
            build_report(
                basis, PilotSelection(years=(2024,), accounts=(("1429", "1"), ("315", "22")))
            ),
            (13, 7, 3),
        )
    )
    for key, counts in (
        ("vedtak341", (2, 2, 2)),
        ("vedtak857", (3, 3, 1)),
        ("vedtak210", (6, 6, 2)),
        ("prop16", (6, 0, 0)),
        ("inns81", (6, 0, 0)),
        ("vedtak211", (0, 0, 0)),
        ("inns95", (0, 0, 0)),
    ):
        document = open_document(both, docs[key])
        assert go_back(document) == both
        cases.append(("document-" + key, document, counts))
    document = open_document(cases[1][1], docs["vedtak857"])
    account = open_accounts(document, both_selection)
    cases.extend(
        [
            ("account-to-document", document, (3, 3, 1)),
            ("return-account", go_back(document), (13, 7, 3)),
            ("document-to-account", account, (23, 11, 5)),
            ("return-document", go_back(account), (3, 3, 1)),
        ]
    )
    records = []
    for name, report, expected in cases:
        records.append(_save(output, name, report, expected))
        (output / "pairs-journal.json").write_text(
            json.dumps(records, ensure_ascii=False, indent=2) + "\n", "utf8"
        )
    offline = read_json(output / "both.json")
    records.append(_save(output, "offline", offline, (23, 11, 5)))
    assert (output / "both.json").read_bytes() == (output / "offline.json").read_bytes()
    export_originals(read_excel(output / "offline.xlsx"), output / "originals")
    # Offentlig CLI med uttrykkelig umulig originalrot, uten DB og uten nett.
    command = [
        sys.executable,
        "-m",
        "analyse.forlopspilot",
        "--frosset",
        str(output / "both.json"),
        "--rot",
        str(output / "no-originals"),
        "--json",
        str(output / "cli-offline.json"),
        "--excel",
        str(output / "cli-offline.xlsx"),
    ]
    cli = subprocess.run(command, capture_output=True, text=True, encoding="utf8")
    assert cli.returncode == 0, cli.stdout + cli.stderr
    assert read_json(output / "cli-offline.json") == read_excel(output / "cli-offline.xlsx") == both
    cli_evidence = {
        "command": command,
        "stdout": cli.stdout,
        "returncode": cli.returncode,
        "original_root_present": False,
    }
    failures = []
    for condition in ("missing", "failed", "corrupt", "contradictory"):
        directory = output / ("synthetic-" + condition)
        directory.mkdir()
        source_root = directory / "sources"
        for name in CATALOG:
            dest = source_root / name
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(name, dest)
        source = source_root / (PARENT if condition == "corrupt" else META)
        if condition in {"missing", "failed"}:
            source.unlink()
            if condition == "failed":
                source.mkdir()
        elif condition == "corrupt":
            source.write_bytes(b"synthetic invalid gzip")
        else:
            source.write_text(source.read_text("utf8").replace("6816165", "6816164"), "utf8")
        failed = import_basis(source_root)
        empty = build_report(failed, PilotSelection())
        export_pair(empty, directory / "empty.json", directory / "empty.xlsx")
        assert read_json(directory / "empty.json") == read_excel(directory / "empty.xlsx") == empty
        ordinary = _ordinary(directory)
        assert ordinary["observations"] == 222 and ordinary["events"] == 74
        failures.append(
            {
                "synthetic": True,
                "condition": condition,
                "issues": failed.issues,
                "ordinary_fresh_database": ordinary,
            }
        )
    negatives = _negatives(output)
    compatibility = []
    if args.previous_root:
        compatibility.extend(_compatibility(args.previous_root, output / "compatibility"))
        compatibility.append(
            _old_pair(
                args.previous_root / "outputs/t1-41-608895f",
                output / "compatibility/T1",
                "begge",
                "ekte_tidligere_T1",
            )
        )
        t2_root = args.previous_root / "outputs/t2-44-1898ef3"
        for old in sorted(t2_root.glob("*.json")):
            if old.with_suffix(".xlsx").is_file():
                compatibility.append(
                    _old_pair(
                        t2_root,
                        output / "compatibility/T2" / old.stem,
                        old.stem,
                        "ekte_tidligere_T2",
                    )
                )
    if args.legacy_b6:
        compatibility.append(verify_legacy_b6(args.legacy_b6, output / "compatibility/B6"))
    manifest = {
        "kind": "T3 #46: autentiske positive prøver; særskilt merkede syntetiske negativer",
        "python": platform.python_version(),
        "lxml": etree.LXML_VERSION,
        "libxml2_runtime": etree.LIBXML_VERSION,
        "libxml2_compiled": etree.LIBXML_COMPILED_VERSION,
        "basis_id": basis.id,
        "source_audit": audit,
        "selections": records,
        "public_cli_offline": cli_evidence,
        "independence": failures,
        "negative_mutations": negatives,
        "compatibility": compatibility,
        "compatibility_counts": dict(Counter(str(x["kind"]) for x in compatibility)),
        "seconds": round(time.monotonic() - started, 3),
        "limitations": [
            "Ingen fullårsdekning, godkjent årsbalanse, kontinuitet eller drift; #10 åpen",
            "Gamle T1-kapselkopier er uendret og telt separat",
            "857s opprinnelige 554e-fullrespons mangler; kontrollobjekt er bevart",
            "Selvstendig gammel frozen-v1-rapport-JSON ikke funnet",
            "Ingen native Excel-prøve; Artifact Tool-rendering gjøres separat",
        ],
        "files": [evidence(p, output) for p in sorted(output.rglob("*")) if p.is_file()],
    }
    (output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", "utf8"
    )
    print(
        encode(
            {
                "output": str(output),
                "pairs": len(records),
                "basis_id": basis.id,
                "manifest": evidence(output / "manifest.json", output),
                "compatibility_counts": manifest["compatibility_counts"],
                "seconds": manifest["seconds"],
            }
        )
    )


if __name__ == "__main__":
    main()
