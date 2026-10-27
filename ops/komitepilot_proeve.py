"""Reproduser T2 fra repoets originaler og autentiske T1-testarkiv til faktiske filer.

python -m ops.komitepilot_proeve --output work/t2-leveranse
Ekstra historisk kontroll: --previous-root C:/.../Budsjetthistorikk
    --legacy-b6 C:/.../accept-b6.xlsx
"""

import argparse
import gzip
import hashlib
import json
import platform
import shutil
import tempfile
from collections import Counter
from dataclasses import asdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from analyse.komitepilot import (
    build_report,
    contents,
    go_back,
    import_basis,
    open_accounts,
    open_document,
    read_json,
)
from analyse.komitepilot_excel import export_originals, export_pair, read_excel
from analyse.visningspilot import read_json as read_t1_json
from analyse.visningspilot_excel import read_excel as read_t1_excel
from innlesning.komitepilot import CATALOG, K
from innlesning.komitepilot_model import Report
from innlesning.observations import encode, identity
from innlesning.visningspilot_model import PilotSelection

from .visningspilot_proeve import _compatibility, _ordinary, evidence, verify_legacy_b6


def serialized_counts(raw: str) -> dict[str, Any]:
    """Tell faktisk serialiserte objekter, også inne i urørte gamle JSON-kapsler."""
    observations: dict[str, Counter[str]] = {"new": Counter(), "legacy": Counter()}
    payloads: dict[str, Counter[str]] = {"new": Counter(), "legacy": Counter()}

    def walk(value: Any, legacy: bool = False) -> None:
        category = "legacy" if legacy else "new"
        if isinstance(value, dict):
            if {
                "dataset",
                "source_system",
                "source_revision_id",
                "interpretation_id",
                "pointer",
            } <= set(value):
                pointer = value["pointer"]
                observations[category][
                    identity(
                        value["dataset"],
                        value["source_system"],
                        value["source_revision_id"],
                        pointer["locator"],
                        pointer["amount_field"],
                        value["interpretation_id"],
                        value["format_version"],
                    )
                ] += 1
            if value.get("content_base64") is not None and "sha256" in value:
                payloads[category][value["sha256"]] += 1
            for key, item in value.items():
                if (
                    key in {"legacy_json", "o2b_json", "base_json"}
                    and isinstance(item, str)
                    and item
                ):
                    walk(json.loads(item), legacy or key == "legacy_json")
                else:
                    walk(item, legacy)
        elif isinstance(value, list):
            for item in value:
                walk(item, legacy)

    walk(json.loads(raw))
    return {
        category: {
            "serialized_observation_objects": sum(observations[category].values()),
            "unique_observations": len(observations[category]),
            "observation_occurrences": dict(observations[category]),
            "payload_occurrences": dict(payloads[category]),
            "serialized_raw_payloads": sum(payloads[category].values()),
            "unique_raw_hashes": len(payloads[category]),
        }
        for category in observations
    }


def _save(output: Path, name: str, report: Report) -> dict[str, Any]:
    frozen, excel = output / (name + ".json"), output / (name + ".xlsx")
    export_pair(report, frozen, excel)
    assert read_json(frozen) == read_excel(excel) == report
    book = load_workbook(excel, read_only=True, data_only=False)
    try:
        snapshot = "".join(
            str(row[1]) for row in book["Offlinegrunnlag"].iter_rows(min_row=2, values_only=True)
        )
        assert snapshot == frozen.read_text("utf8")
        shapes = {s.title: {"rows": s.max_row, "columns": s.max_column} for s in book}
        actual = {
            "fields": len(report.view.observation_ids),
            "adopted_fields": len(report.view.adopted_observation_ids),
            "events": len(report.view.adopted_event_ids),
            "visible_treatment_rows": book["Forløp"].max_row - 1,
        }
        assert actual["visible_treatment_rows"] * 3 == actual["fields"]
    finally:
        book.close()
    counts = serialized_counts(snapshot)
    assert counts["new"]["serialized_observation_objects"] == 9
    assert all(v == 1 for v in counts["new"]["payload_occurrences"].values())
    return {
        "name": name,
        "report_id": report.id,
        "basis_id": report.basis.id,
        "selection": asdict(report.selection),
        "actual": actual,
        "files": [evidence(frozen, output), evidence(excel, output)],
        "shapes": shapes,
        "serialized_counts": counts,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--previous-root", type=Path)
    parser.add_argument("--legacy-b6", type=Path)
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    temporary = output / "reader-tmp"
    temporary.mkdir()
    tempfile.tempdir = str(temporary)
    single = import_basis(Path("."))
    t1_raw = gzip.decompress(Path("tests/fixtures/t2/authentic-t1.json.gz").read_bytes())
    assert (
        hashlib.sha256(t1_raw).hexdigest()
        == "99e871e0ab780cbb548e15699f3093621d0f25229e3fa06c29cfdb5552803f53"
    )
    union = import_basis(Path("."), t1_json=t1_raw.decode("utf8"))
    selection = PilotSelection(years=(2024,), accounts=(("315", "78"),))
    pairs = PilotSelection(years=(2024,), accounts=(("315", "78"), ("1429", "1")))
    post = build_report(single, selection)
    united = build_report(union, pairs)
    reports = [
        ("standalone", post, {"fields": 9, "adopted_fields": 3, "events": 1}),
        ("union", united, {"fields": 18, "adopted_fields": 6, "events": 2}),
        (
            "empty",
            build_report(single, PilotSelection()),
            {"fields": 0, "adopted_fields": 0, "events": 0},
        ),
        (
            "chapter",
            build_report(single, PilotSelection(years=(2024,), chapters=("315",))),
            {"fields": 9, "adopted_fields": 3, "events": 1},
        ),
        (
            "years",
            build_report(union, PilotSelection(years=(2024, 2025), accounts=pairs.accounts)),
            {"fields": 18, "adopted_fields": 6, "events": 2},
        ),
    ]
    docs = {d.key: d for d in single.documents}
    for key in docs:
        document = open_document(post, docs[key].id)
        assert go_back(document) == post
        expected = {
            "fields": 3 if key in {"prop25", "inns85", "vedtak208"} else 0,
            "adopted_fields": 3 if key == "vedtak208" else 0,
            "events": 1 if key == "vedtak208" else 0,
        }
        reports.append(("document-" + key, document, expected))
    doc = open_document(post, docs["inns85"].id)
    account = open_accounts(doc, selection)
    assert go_back(account) == build_report(single, doc.selection)
    reports.extend(
        [
            ("account-to-document", doc, {"fields": 3, "adopted_fields": 0, "events": 0}),
            ("return-account", go_back(doc), {"fields": 9, "adopted_fields": 3, "events": 1}),
            ("document-to-account", account, {"fields": 9, "adopted_fields": 3, "events": 1}),
            ("return-document", go_back(account), {"fields": 3, "adopted_fields": 0, "events": 0}),
        ]
    )
    records = []
    for name, report, expected in reports:
        record = _save(output, name, report)
        record["expected"] = expected
        assert all(record["actual"][key] == value for key, value in expected.items())
        records.append(record)
    offline = read_json(output / "standalone.json")
    records.append(_save(output, "offline", offline))
    assert (output / "offline.json").read_bytes() == (output / "standalone.json").read_bytes()
    export_originals(read_excel(output / "offline.xlsx"), output / "originals")
    failures = []
    for condition in ("missing", "failed", "corrupt", "contradictory"):
        directory = output / ("synthetic-" + condition)
        directory.mkdir()
        source_root = directory / "sources"
        for name in CATALOG:
            destination = source_root / name
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(name, destination)
        source = source_root / (K + "inns85s.html")
        if condition in {"missing", "failed"}:
            source.unlink()
            if condition == "failed":
                source.mkdir()
        elif condition == "corrupt":
            source.write_bytes(b"synthetic corrupt HTML")
        else:
            source.write_text(
                source.read_text("utf8").replace("129 000 000", "119 000 000"), "utf8"
            )
        basis = import_basis(source_root)
        gap = build_report(basis, selection)
        assert len(gap.view.observation_ids) == 6
        assert all(c.difference_kr is None and c.status == "hull" for c in gap.view.comparisons)
        export_pair(gap, directory / "gap.json", directory / "gap.xlsx")
        empty = build_report(basis, PilotSelection())
        export_pair(empty, directory / "empty.json", directory / "empty.xlsx")
        assert read_json(directory / "gap.json") == read_excel(directory / "gap.xlsx") == gap
        assert read_json(directory / "empty.json") == read_excel(directory / "empty.xlsx") == empty
        ordinary = _ordinary(directory)
        assert ordinary["observations"] == 222 and ordinary["events"] == 74
        failures.append(
            {
                "synthetic": True,
                "condition": condition,
                "issues": basis.issues,
                "fields": len(gap.view.observation_ids),
                "comparisons": [asdict(c) for c in gap.view.comparisons],
                "ordinary_fresh_database": ordinary,
            }
        )
    negative_dir = output / "synthetic-mutations"
    negative_dir.mkdir()
    negatives = []
    for mutation in ("amount", "actor", "vote", "relation"):
        data = json.loads((output / "standalone.json").read_text("utf8"))
        b = data["report"]["basis"]
        if mutation == "amount":
            b["observations"][4]["amount_kr"] = 119_000_000
        elif mutation == "actor":
            b["contexts"][4]["actor"] = "samlet_komite"
        elif mutation == "vote":
            next(s for s in b["statements"] if s["role"] == "postvis_votering_63_35_194948")[
                "text"
            ] = "synthetic unanimous"
        else:
            b["relations"][1]["evidence_ids"] = ["synthetic-dangling"]
        b["id"] = ""
        b["id"] = identity("komitegrunnlag-1", b)
        data["report"]["id"] = ""
        data["report"]["id"] = identity("visningspilot-2", data["report"])
        target = negative_dir / (mutation + ".json")
        target.write_text(encode(data), "utf8")
        try:
            read_json(target)
        except ValueError as error:
            negatives.append(
                {"synthetic": True, "mutation": mutation, "rehashed": True, "rejected": str(error)}
            )
        else:
            raise AssertionError("Synthetic mutation accepted")
    for mutation in ("numberformat", "hidden-issues", "false-link"):
        target = negative_dir / (mutation + ".xlsx")
        shutil.copyfile(output / "standalone.xlsx", target)
        book = load_workbook(target)
        if mutation == "numberformat":
            book["Forløp"]["H2"].number_format = '"0"'
        elif mutation == "hidden-issues":
            book["Dekning og avvik"].sheet_state = "hidden"
        else:
            book["Forløp"]["C2"].hyperlink = "https://example.invalid/synthetic"
        book.save(target)
        book.close()
        try:
            read_excel(target)
        except ValueError as error:
            negatives.append({"synthetic": True, "mutation": mutation, "rejected": str(error)})
        else:
            raise AssertionError("Synthetic presentation accepted")
    compatibility = []
    if args.previous_root:
        compatibility = _compatibility(args.previous_root, output / "compatibility")
        source = args.previous_root / "outputs/t1-41-608895f"
        copied = output / "compatibility/old-t1"
        copied.mkdir()
        before = {
            suffix: hashlib.sha256((source / ("begge" + suffix)).read_bytes()).hexdigest()
            for suffix in (".json", ".xlsx")
        }
        for suffix in before:
            shutil.copyfile(source / ("begge" + suffix), copied / ("begge" + suffix))
        old = read_t1_json(copied / "begge.json")
        assert read_t1_excel(copied / "begge.xlsx") == old
        assert old.basis.observations == tuple(
            o for o in contents(union).observations if o.source_system != "komitepilot"
        )
        assert before == {
            suffix: hashlib.sha256((source / ("begge" + suffix)).read_bytes()).hexdigest()
            for suffix in before
        }
        compatibility.append(
            {
                "kind": "ekte_tidligere_T1",
                "report_id": old.id,
                "original_hashes": before,
                "unchanged": True,
            }
        )
    if args.legacy_b6:
        compatibility.append(verify_legacy_b6(args.legacy_b6, output / "compatibility/old-B6"))
    manifest = {
        "kind": "T2 #44: autentiske positive prøver og separat merkede syntetiske feil",
        "python": platform.python_version(),
        "basis_id": single.id,
        "union_basis_id": union.id,
        "source_ids": [s.id for s in single.sources],
        "observation_ids": [o.id for o in single.observations],
        "selections": records,
        "independence": failures,
        "negative_mutations": negatives,
        "compatibility": compatibility,
        "compatibility_counts": dict(Counter(str(x["kind"]) for x in compatibility)),
        "limitations": [
            "K1 positiv alternativport åpen",
            "Ingen fullårsdekning, faglig godkjenning, kontinuitet eller drift; #10 åpen",
            "Arvede T1-kapselkopier er bevart og telt separat",
            (
                "Selvstendig gammel frozen-v1-JSON ikke funnet; "
                "innebygde kjerner og ekte eldre B6-Excel prøves separat"
            ),
            (
                "Ingen native Excel-prøve; lagrede celler/lenker/stil/byte "
                "og rendering av lagret XLSX kontrolleres"
            ),
        ],
        "files": [evidence(p, output) for p in sorted(output.rglob("*")) if p.is_file()],
    }
    (output / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", "utf8"
    )
    print(
        encode(
            {
                "output": str(output),
                "pairs": len(records),
                "manifest": evidence(output / "manifest.json", output),
                "basis_id": single.id,
                "compatibility_counts": manifest["compatibility_counts"],
            }
        )
    )


if __name__ == "__main__":
    main()
