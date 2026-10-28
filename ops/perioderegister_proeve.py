"""Reproduser P1s autentiske filer, semantiske negativer og uavhengig ordinær innlesning.

python -m ops.perioderegister_proeve --output work/p1-kontroll --previous-root <hovedrepo>
"""

import argparse
import base64
import gzip
import hashlib
import json
import os
import platform
import shutil
import subprocess
import sys
import tempfile
import time
from collections import Counter
from pathlib import Path
from typing import Any

from lxml import etree, html
from openpyxl import load_workbook

from analyse.omtale_excel_guard import inspect_excel
from analyse.perioderegister import (
    build_report,
    contents,
    go_back,
    import_basis,
    open_accounts,
    open_document,
    read_json,
    report_to_json,
)
from analyse.perioderegister_excel import (
    SHEET_NAMES,
    _book,
    export_pair,
    read_excel,
)
from analyse.session_serialization import _decode
from innlesning.observations import encode, identity
from innlesning.omtale import compact
from innlesning.perioderegister import representation
from innlesning.perioderegister_catalog import CATALOG, ROOT
from innlesning.perioderegister_model import Basis, Report
from innlesning.visningspilot_model import PilotSelection

from .fleraarspilot_proeve import mutate_display as _old_display
from .visningspilot_proeve import _compatibility, _ordinary, evidence, verify_legacy_b6

MUTATIONS = (
    "source-bytes",
    "source-hash",
    "parent",
    "locator",
    "year",
    "post",
    "role",
    "ordinal",
    "case-identity",
    "candidate-status",
    "decision-year",
    "decision-reason",
    "fullmakt-amount",
    "alternative-adopted",
    "prop36-event",
    "positive-coverage",
    "annual-approval",
    "hidden-gaps",
    "hidden-global-view",
    "text-loss",
    "duplicate-payload",
    "legacy-raw-copy",
    "entry-cap",
)


def _rehash(data: dict[str, Any]) -> None:
    report, basis = data["report"], data["report"]["basis"]
    namespaces = {
        "sources": "periodekilde-1",
        "evidence": "periodebevis-1",
        "documents": "periodedokument-1",
        "inventory": "periodeinventar-1",
        "statements": "periodeutsagn-1",
        "relations": "perioderelasjon-1",
        "entries": "periodeledd-1",
        "case_candidates": "periodesakskandidat-1",
        "decision_candidates": "periodevedtakskandidat-1",
        "gaps": "periodehull-1",
    }
    replacements: dict[str, str] = {}

    def update(obj: Any) -> Any:
        if isinstance(obj, str):
            return replacements.get(obj, obj)
        if isinstance(obj, list):
            return [update(v) for v in obj]
        if isinstance(obj, dict):
            return {k: v if k == "t4_basis" else update(v) for k, v in obj.items()}
        return obj

    # Flere runder binder også dokument-/bevisreferanser gjennom hele DAG-en.
    for collection, namespace in namespaces.items():
        basis[collection] = update(basis[collection])
        for obj in basis[collection]:
            old = obj["id"]
            obj["id"] = ""
            obj["id"] = identity(namespace, obj)
            replacements[old] = obj["id"]
    # Ikke endre gamle objekter eller deres navnerom i et nytt format.
    changed = update(report)
    report.clear()
    report.update(changed)
    report["basis"]["id"] = ""
    report["basis"]["id"] = identity("periodegrunnlag-1", report["basis"])
    report["id"] = ""
    report["id"] = identity("perioderegister-1", report)


def mutate(raw: str, name: str) -> str:
    data = json.loads(raw)
    report, basis = data["report"], data["report"]["basis"]
    fullmakt = next(e for e in basis["entries"] if e["key"] == "original-346")
    if name == "source-bytes":
        payload = basis["payloads"][0]
        old_digest = payload["id"]
        content = base64.b64decode(payload["content_base64"])
        changed = content[:-1] + bytes([content[-1] ^ 1])
        digest = hashlib.sha256(changed).hexdigest()
        payload.update(id=digest, sha256=digest, content_base64=base64.b64encode(changed).decode())
        for source in basis["sources"]:
            if source["payload_id"] == old_digest:
                source["payload_id"] = digest
                source["expected_sha256"] = digest
                source["status"] = "kildeverifisert"
    elif name == "source-hash":
        basis["evidence"][0]["representation_sha256"] = "0" * 64
    elif name == "parent":
        s = basis["sources"][0]
        meta = json.loads(s["metadata_json"])
        meta["representation_sha256"] = "0" * 64
        s["metadata_json"] = encode(meta)
    elif name == "locator":
        basis["evidence"][0]["locator"] = "synthetic:/wrong-original/1"
    elif name == "year":
        fullmakt["budget_year"] = 2023
    elif name == "post":
        fullmakt["accounts"] = ["1429/22"]
    elif name == "role":
        fullmakt["role"] = "adopted_levels"
    elif name == "ordinal":
        basis["case_candidates"][0]["ordinal_1based"] += 1
    elif name == "case-identity":
        basis["case_candidates"][0]["case_id"] = "999999"
    elif name == "candidate-status":
        basis["case_candidates"][0]["status"] = "closed"
    elif name == "decision-year":
        next(c for c in basis["decision_candidates"] if c["decision_api_id"] == "40034038")[
            "budget_year"
        ] = 2024
    elif name == "decision-reason":
        basis["decision_candidates"][0]["reason"] = "synthetic complete annual exclusion"
    elif name == "fullmakt-amount":
        fullmakt["observation_ids"] = list(
            report["view"]["monetary"]["adopted_observation_ids"][:1]
        )
        fullmakt["event_ids"] = list(report["view"]["monetary"]["adopted_event_ids"][:1])
    elif name == "alternative-adopted":
        next(e for e in basis["entries"] if e["key"] == "rnb-alternative-1")["role"] = (
            "adopted_change"
        )
    elif name == "prop36-event":
        next(e for e in basis["entries"] if e["key"] == "annual-control")["event_ids"] = [
            "synthetic-prop36-event"
        ]
    elif name == "positive-coverage":
        basis["year_coverage"] = "complete"
    elif name == "annual-approval":
        basis["professional_annual_balance_approval"] = True
    elif name == "hidden-gaps":
        basis["gaps"] = []
        report["view"]["gap_ids"] = []
    elif name == "hidden-global-view":
        report["view"]["case_candidate_ids"] = []
    elif name == "text-loss":
        basis["statements"][0]["text"] = "synthetic truncated text"
    elif name == "duplicate-payload":
        basis["payloads"].append(basis["payloads"][0])
    elif name == "legacy-raw-copy":
        from analyse.komitepilot import legacy_payloads

        payload = next(p for p in basis["payloads"] if p["legacy_source_id"])
        old_raw = dict(legacy_payloads(basis["t4_basis"]["legacy_json"]))[
            payload["legacy_source_id"]
        ]
        payload["content_base64"] = base64.b64encode(old_raw).decode()
        payload["legacy_source_id"] = None
    elif name == "entry-cap":
        basis["entries"] = [fullmakt] * 257
    else:
        raise ValueError(name)
    _rehash(data)
    return encode(data)


def selection_cases(basis: Basis) -> list[tuple[str, PilotSelection, tuple[int, int]]]:
    cases = [
        ("post1", PilotSelection(years=(2024,), accounts=(("1429", "1"),)), (7, 3)),
        ("post22", PilotSelection(years=(2024,), accounts=(("1429", "22"),)), (4, 2)),
        ("both", PilotSelection(years=(2024,), accounts=(("1429", "1"), ("1429", "22"))), (11, 5)),
        ("chapter", PilotSelection(years=(2024,), chapters=("1429",)), (11, 5)),
        ("wrong-year", PilotSelection(years=(2025,), chapters=("1429",)), (0, 0)),
        ("empty", PilotSelection(), (0, 0)),
        ("pairs", PilotSelection(years=(2024,), accounts=(("1429", "1"), ("315", "22"))), (7, 3)),
    ]
    docs = {d.key: d.id for d in contents(basis).documents}
    for key, count in (
        ("vedtak341", (2, 2)),
        ("vedtak857", (3, 1)),
        ("vedtak210", (6, 2)),
        ("vedtak346", (0, 0)),
        ("vedtak201", (0, 0)),
    ):
        cases.append((key, PilotSelection(mode="document", document_id=docs[key]), count))
    cases.append(
        (
            "unknown-index",
            PilotSelection(mode="document", document_id=basis.case_candidates[0].document_id),
            (0, 0),
        )
    )
    return cases


def save_pair(output: Path, name: str, report: Report, counts: tuple[int, int]) -> dict[str, Any]:
    frozen, excel = output / (name + ".json"), output / (name + ".xlsx")
    export_pair(report, frozen, excel)
    saved = read_json(frozen)
    assert saved == read_excel(excel) == report
    assert (
        len(saved.view.monetary.adopted_observation_ids),
        len(saved.view.monetary.adopted_event_ids),
    ) == counts
    assert (
        len(saved.view.case_candidate_ids),
        len(saved.view.decision_candidate_ids),
        len(saved.view.gap_ids),
    ) == (67, 25, 7)
    book = load_workbook(excel, read_only=True)
    try:
        typed = [
            {
                "row": row[0].row,
                "account": str(row[0].value),
                "year": row[1].value,
                "values": [c.value for c in row[6:10]],
                "types": [c.data_type for c in row[6:10]],
                "formats": [c.number_format for c in row[6:10]],
            }
            for row in book["Forløp"].iter_rows(min_row=2)
        ]
    finally:
        book.close()
    return {
        "name": name,
        "report_id": report.id,
        "counts": counts,
        "global": (67, 25, 7),
        "typed_cells": typed,
        "files": [evidence(p, output) for p in (frozen, excel)],
    }


def source_audit() -> dict[str, Any]:
    originals = {}
    source_sizes = []
    for name, (pin, count, _, _, _) in CATALOG.items():
        raw = Path(name).read_bytes()
        assert len(raw) == count and hashlib.sha256(raw).hexdigest() == pin, name
        originals[name] = representation(raw, name)
        source_sizes.append(
            {"name": name, "archive_bytes": len(raw), "representation_bytes": len(originals[name])}
        )
    candidates = json.loads(originals[ROOT + "case-candidates.json"])
    for row in candidates:
        root = etree.fromstring(originals[ROOT + "originals/saker-" + row["session"] + ".xml.gz"])
        nodes = root.xpath(f"/*/*[3]/*[{row['ordinal_1based']}]")
        assert len(nodes) == 1
        fields = {etree.QName(n).localname: n.text for n in nodes[0]}
        assert all(fields[k] == value for k, value in row["metadata"].items())
    table_rows = 0
    for row in json.loads(originals[ROOT + "table-locators.json"]):
        raw = originals[ROOT + "originals/" + row["source"] + ".gz"]
        assert hashlib.sha256(raw).hexdigest() == row["body_sha256"]
        doc = html.document_fromstring(raw, parser=html.HTMLParser(encoding="utf8"))
        assert len(doc.xpath(row["table_xpath"])) == 1
        for item in row["rows"]:
            nodes = doc.xpath(item["xpath"])
            assert len(nodes) == 1, item["xpath"]
            assert [compact(n.text_content()) for n in nodes[0]] == item["cells"]
            table_rows += 1
    for item in json.loads(originals[ROOT + "decision-evidence.json"]):
        root = etree.fromstring(originals[item["source_path"]])
        ns = {"s": "http://data.stortinget.no"}
        node = root.xpath(item["xml_locator"], namespaces=ns)
        assert len(node) == 1
        assert hashlib.sha256((node[0].text or "").encode()).hexdigest() == item["text_sha256_utf8"]
    voting = json.loads(originals[ROOT + "voting-evidence.json"])
    for section in ("votes", "decision_links"):
        for row in voting[section]:
            root = etree.fromstring(originals[ROOT + "originals/" + row["source"] + ".gz"])
            nodes = root.xpath(row["xpath"])
            assert len(nodes) == 1
            actual = {etree.QName(n).localname: n.text for n in nodes[0] if not len(n)}
            assert actual == row["fields"]
    return {
        "pinned_sources": len(CATALOG),
        "source_sizes": source_sizes,
        "case_occurrences": 67,
        "table_rows": table_rows,
        "decision_objects": 25,
        "errors": 0,
    }


def mutate_display(source: Path, output: Path, name: str) -> None:
    if name != "amount":
        _old_display(source, output, name)
        return
    book = load_workbook(source)
    book["Forløp"]["G2"] = 999
    book.save(output)
    book.close()


def negative_pairs(output: Path, report: Report) -> list[dict[str, Any]]:
    raw = report_to_json(report)
    directory = output / "negative"
    directory.mkdir()
    records = []
    for name in MUTATIONS:
        poisoned = mutate(raw, name)
        frozen, excel = directory / (name + ".json"), directory / (name + ".xlsx")
        frozen.write_text(poisoned, encoding="utf8")
        # Bygg konsistente dimensjoner/filter/radhøyder rundt den forgiftede kapselen.
        typed: Report = _decode(Report, json.loads(poisoned)["report"])
        book = _book(typed, poisoned)
        book.save(excel)
        book.close()
        capsule, _ = inspect_excel(
            excel,
            sheet_names=SHEET_NAMES,
            snapshot_name="Offlinegrunnlag",
            max_expanded_bytes=128 * 1024 * 1024,
        )
        assert capsule == poisoned, name
        errors = {}
        for kind, reader, path in (("json", read_json, frozen), ("xlsx", read_excel, excel)):
            try:
                reader(path)
            except ValueError as error:
                errors[kind] = str(error)
            else:
                raise AssertionError("Syntetisk angrep akseptert: " + name)
        records.append(
            {
                "mutation": name,
                "synthetic": True,
                "rehashed": True,
                "container_passed_exact_capsule": True,
                "rejections": errors,
            }
        )
    return records


def old_pair(file: Path, copied: Path, kind: str) -> dict[str, Any]:
    from analyse import fleraarspilot as t4
    from analyse import fleraarspilot_excel as t4_excel
    from analyse import forlopspilot as t3
    from analyse import forlopspilot_excel as t3_excel
    from analyse import komitepilot as t2
    from analyse import komitepilot_excel as t2_excel
    from analyse import visningspilot as t1
    from analyse import visningspilot_excel as t1_excel

    readers = {
        "T1": (t1.read_json, t1_excel.read_excel),
        "T2": (t2.read_json, t2_excel.read_excel),
        "T3": (t3.read_json, t3_excel.read_excel),
        "T4": (t4.read_json, t4_excel.read_excel),
    }
    files = (file, file.with_suffix(".xlsx"))
    before = [hashlib.sha256(p.read_bytes()).hexdigest() for p in files]
    copied.mkdir(parents=True)
    for p, pin in zip(files, before, strict=True):
        shutil.copyfile(p, copied / p.name)
        assert hashlib.sha256((copied / p.name).read_bytes()).hexdigest() == pin
    first, second = readers[kind]
    old = first(copied / file.name)
    assert second(copied / file.with_suffix(".xlsx").name) == old
    assert before == [hashlib.sha256(p.read_bytes()).hexdigest() for p in files]
    return {
        "kind": "ekte_tidligere_" + kind,
        "source": str(file),
        "report_id": old.id,
        "original_hashes": before,
        "unchanged": True,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--previous-root", type=Path)
    parser.add_argument("--legacy-b6", type=Path)
    args = parser.parse_args()
    started = time.monotonic()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    (output / "tmp").mkdir()
    tempfile.tempdir = str(output / "tmp")
    os.environ["TMP"] = os.environ["TEMP"] = str(output / "tmp")
    old = gzip.decompress(Path("tests/fixtures/t2/authentic-t1.json.gz").read_bytes()).decode()
    basis = import_basis(Path("."), t1_json=old)
    audit = source_audit()
    records = []
    for name, selection, count in selection_cases(basis):
        report = build_report(basis, selection)
        records.append(save_pair(output, name, report, count))
    both = read_json(output / "both.json")
    fullmakt = next(e for e in basis.entries if e.key == "original-346")
    opened = open_document(both, fullmakt.document_id)
    records.append(save_pair(output, "opened-fullmakt", opened, (0, 0)))
    assert go_back(read_excel(output / "opened-fullmakt.xlsx")) == both
    reverse = open_accounts(opened, PilotSelection(years=(2024,), accounts=(("1429", "22"),)))
    records.append(save_pair(output, "reverse-post22", reverse, (4, 2)))
    assert go_back(reverse) == build_report(basis, opened.selection)
    records.append(save_pair(output, "repeated", read_excel(output / "both.xlsx"), (11, 5)))
    assert (output / "repeated.json").read_bytes() == (output / "both.json").read_bytes()
    script = """import socket, sqlite3, sys
def unavailable(*a, **kw): raise AssertionError('network/database forbidden')
socket.create_connection=unavailable
socket.socket.connect=unavailable
socket.socket.connect_ex=unavailable
sqlite3.connect=unavailable
sys.argv=['analyse.perioderegister','--frosset',sys.argv[1],'--rot',sys.argv[2],
          '--json',sys.argv[3],'--excel',sys.argv[4],'--originaler',sys.argv[5]]
from analyse.perioderegister import main
main()
"""
    offline_cwd = output / "empty-working-directory"
    offline_cwd.mkdir()
    offline_env = dict(os.environ)
    offline_env["PYTHONPATH"] = os.pathsep.join(
        (str(Path.cwd()), str(Path(".local/python").resolve()))
    )
    cli = subprocess.run(
        [
            sys.executable,
            "-c",
            script,
            str(output / "both.json"),
            str(output / "missing-root"),
            str(output / "cli-offline.json"),
            str(output / "cli-offline.xlsx"),
            str(output / "originals"),
        ],
        capture_output=True,
        text=True,
        encoding="utf8",
        cwd=offline_cwd,
        env=offline_env,
    )
    assert cli.returncode == 0, cli.stdout + cli.stderr
    assert read_json(output / "cli-offline.json") == read_excel(output / "cli-offline.xlsx") == both
    negatives = negative_pairs(output, both)
    display = []
    for name in (
        "amount",
        "hidden-row",
        "hidden-column",
        "hyperlink",
        "filter",
        "doctype",
        "numberformat",
        "hidden-sheet",
    ):
        path = output / ("display-" + name + ".xlsx")
        mutate_display(output / "both.xlsx", path, name)
        try:
            read_excel(path)
        except ValueError as error:
            display.append({"name": name, "rejection": str(error)})
        else:
            raise AssertionError(name)
    independence = []
    for condition in ("missing", "failed", "corrupt", "contradictory"):
        directory = output / ("independence-" + condition)
        directory.mkdir()
        # Fire forskjellige P1-kilder/inputfeil, fire separate ordinære databaser.
        if condition == "missing":
            failed = import_basis(directory / "no-root")
            empty = build_report(failed, PilotSelection())
            export_pair(empty, directory / "empty.json", directory / "empty.xlsx")
        elif condition == "failed":
            try:
                import_basis(directory / "no-root", t1_json="failed enrichment")
            except ValueError:
                pass
            else:
                raise AssertionError(condition)
        elif condition == "corrupt":
            (directory / "corrupt.json").write_text("corrupt P1", encoding="utf8")
            try:
                read_json(directory / "corrupt.json")
            except ValueError:
                pass
            else:
                raise AssertionError(condition)
        else:
            (directory / "contradictory.json").write_text(
                mutate(report_to_json(both), "positive-coverage"), encoding="utf8"
            )
            try:
                read_json(directory / "contradictory.json")
            except ValueError:
                pass
            else:
                raise AssertionError(condition)
        import sqlite3

        from innlesning import read_extraction
        from innlesning.registers import read_register, seed_register
        from innlesning.source_attempts import begin_source_attempt

        database = directory / "ordinary.sqlite"
        begin_source_attempt(database, "prove", "2024-2025")
        register_before = seed_register(database)
        ordinary = _ordinary(directory)
        assert ordinary["observations"] == 222 and ordinary["events"] == 74
        assert read_register(database, register_before.id) == register_before
        extracted = read_extraction(database, str(ordinary["extraction_id"]))
        assert {o.round for o in extracted.observations} == {"uavklart"}
        with sqlite3.connect(database) as connection:
            assert connection.execute("PRAGMA user_version").fetchone()[0] == 5
        ordinary["schema"] = 5
        ordinary["b5_register_unchanged"] = register_before.id
        ordinary["old_round_values"] = ["uavklart"]
        independence.append({"condition": condition, "ordinary": ordinary})
        if condition != "missing":
            empty = build_report(basis, PilotSelection())
            export_pair(empty, directory / "empty.json", directory / "empty.xlsx")
            assert read_json(directory / "empty.json") == read_excel(directory / "empty.xlsx")
    compatibility = []
    if args.previous_root:
        compatibility.extend(_compatibility(args.previous_root, output / "compatibility"))
        for kind, folder in (
            ("T1", "t1-41-608895f"),
            ("T2", "t2-44-1898ef3"),
            ("T3", "t3-46-ccba906"),
            ("T4", "t4-48-3f856d8"),
        ):
            directory = args.previous_root / "outputs" / folder
            file = next(
                (
                    p
                    for p in sorted(directory.glob("*.json"))
                    if p.with_suffix(".xlsx").is_file() and not p.name.startswith("synthetic")
                ),
                None,
            )
            if file:
                compatibility.append(
                    old_pair(file, output / "compatibility" / kind / file.stem, kind)
                )
        if args.legacy_b6:
            compatibility.append(verify_legacy_b6(args.legacy_b6, output / "compatibility/B6"))
    result = {
        "python": platform.python_version(),
        "lxml": etree.LXML_VERSION,
        "libxml2_runtime": etree.LIBXML_VERSION,
        "basis_id": basis.id,
        "source_audit": audit,
        "selections": records,
        "negative_pairs": negatives,
        "display_negatives": display,
        "independence": independence,
        "compatibility": compatibility,
        "offline_cli": {
            "returncode": cli.returncode,
            "stdout": cli.stdout,
            "network_sqlite_forbidden": True,
            "original_root_present": False,
        },
        "resources": {
            "report_bytes": len(report_to_json(both).encode()),
            "payloads": len(basis.payloads),
            "sources": len(basis.sources),
            "evidence": len(basis.evidence),
            "entries": len(basis.entries),
            "inventory_per_document": dict(Counter(i.document_id for i in basis.inventory)),
            "limits": {
                "report": 64 * 1024 * 1024,
                "payload": 16 * 1024 * 1024,
                "decompressed": 32 * 1024 * 1024,
                "sources": 512,
                "entries": 256,
                "cases": 67,
                "decisions": 25,
                "gaps": 7,
                "evidence": 4096,
            },
        },
        "seconds": round(time.monotonic() - started, 3),
        "limitations": [
            "Ingen native Excel, full årsdekning, faglig årsbalanse, kontinuitet eller drift",
            "Gamle immutable kapslers fysiske duplisering beholdes",
        ],
        "files": [evidence(p, output) for p in sorted(output.rglob("*")) if p.is_file()],
    }
    (output / "verification.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf8"
    )
    print(
        encode({"pairs": len(records), "negatives": len(negatives), "seconds": result["seconds"]})
    )


if __name__ == "__main__":
    main()
