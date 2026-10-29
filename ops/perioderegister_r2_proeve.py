"""Reproduser P2s faktiske filer, semantiske port, kildegjenbruk og 368-uavhengighet."""

import argparse
import base64
import gzip
import hashlib
import json
import os
import platform
import shutil
import sqlite3
import subprocess
import sys
import time
from collections import Counter
from pathlib import Path
from typing import Any

from lxml import etree
from openpyxl import load_workbook

from analyse import perioderegister as p1
from analyse import perioderegister_excel as p1_excel
from analyse.omtale_excel_guard import inspect_excel
from analyse.perioderegister_r2 import (
    build_report,
    contents,
    go_back,
    import_basis,
    open_accounts,
    open_document,
    read_json,
    report_to_json,
)
from analyse.perioderegister_r2_excel import SHEET_NAMES, _book, export_pair, read_excel
from analyse.session_serialization import _decode
from innlesning.komitepilot import payload_bytes
from innlesning.observations import encode, identity
from innlesning.perioderegister_r2 import representation, shared_sources
from innlesning.perioderegister_r2_catalog import CATALOG
from innlesning.perioderegister_r2_model import Report
from innlesning.visningspilot_model import PilotSelection

from .perioderegister_proeve import _rehash as rehash_p1
from .perioderegister_proeve import mutate_display, old_pair
from .visningspilot_proeve import _compatibility, _ordinary, evidence, verify_legacy_b6

MUTATIONS = (
    "estimate-adopted",
    "o1-exact",
    "authorization-used",
    "authorization-2025",
    "post22",
    "wrong-source",
    "representation",
    "locator",
    "nested-locator",
    "role",
    "actor",
    "year",
    "proposal",
    "vote",
    "decision",
    "hidden-gaps",
    "hidden-candidates",
    "hidden-status",
    "hidden-crossrefs",
    "positive-coverage",
    "annual-approval",
    "old-identity",
    "source-bytes",
    "duplicate-payload",
    "old-payload-copy",
    "text-loss",
)


def _rehash(data: dict[str, Any]) -> None:
    report, basis = data["report"], data["report"]["basis"]
    replacements: dict[str, str] = {}

    def update(value: Any) -> Any:
        if isinstance(value, str):
            return replacements.get(value, value)
        if isinstance(value, list):
            return [update(v) for v in value]
        if isinstance(value, dict):
            return {k: v if k == "p1_basis" else update(v) for k, v in value.items()}
        return value

    namespaces = {
        "sources": "periodekilde-r2-1",
        "evidence": "periodebevis-r2-1",
        "documents": "periodedokument-r2-1",
        "inventory": "periodeinventar-r2-1",
        "statements": "periodeutsagn-r2-1",
        "relations": "perioderelasjon-r2-1",
        "findings": "periodefunn-r2-1",
        "status_updates": "periodestatus-r2-1",
        "cross_references": "periodekryssreferanse-r2-1",
    }
    for collection, namespace in namespaces.items():
        basis[collection] = update(basis[collection])
        for obj in basis[collection]:
            old = obj["id"]
            obj["id"] = ""
            obj["id"] = identity(namespace, obj)
            replacements[old] = obj["id"]
    changed = update(report)
    report.clear()
    report.update(changed)
    report["basis"]["id"] = ""
    report["basis"]["id"] = identity("periodegrunnlag-r2-1", report["basis"])
    report["id"] = ""
    report["id"] = identity("perioderegister-r2-1", report)


def mutate(raw: str, name: str) -> str:
    data = json.loads(raw)
    report, basis = data["report"], data["report"]["basis"]
    findings = {f["key"]: f for f in basis["findings"]}
    estimate, authorization = findings["FIN-1429-1"], findings["FULLMAKT429"]
    if name == "estimate-adopted":
        estimate["role"] = "endelig_stortingsvedtak"
        values = json.loads(estimate["amounts_json"])
        values["adopted_new_event"] = 2796000
        estimate["amounts_json"] = encode(values)
    elif name == "o1-exact":
        findings["O1-OMDISPONERING"]["role"] = "eksakt_primært_lønnsbeløp"
    elif name == "authorization-used":
        authorization["role"] = "faktisk_overføring"
        authorization["amounts_json"] = encode({"transferred": 2796000})
    elif name == "authorization-2025":
        authorization["role"] = "ny_bevilgning"
        authorization["budget_year"] = 2025
        authorization["accounts"] = ["1429/1"]
    elif name == "post22":
        estimate["accounts"] = ["1429/22"]
        authorization["conditional_accounts"] = ["1429/22"]
    elif name == "wrong-source":
        basis["evidence"][0]["source_id"] = basis["evidence"][1]["source_id"]
    elif name == "representation":
        basis["evidence"][0]["representation_sha256"] = "0" * 64
    elif name == "locator":
        next(e for e in basis["evidence"] if e["key"] == "R28")["locator"] += "/td[3]"
    elif name == "nested-locator":
        next(e for e in basis["evidence"] if e["key"] == "D40034098")["nested_html_field_xpath"] = (
            "./s:id"
        )
    elif name == "role":
        findings["D231"]["role"] = "fordelingsfullmakt_1429"
    elif name == "actor":
        findings["KOMITE"]["actor"] = "mindretallet"
    elif name == "year":
        estimate["budget_year"] = 2025
    elif name == "proposal":
        authorization["proposal_ids"] = ["40077003"]
    elif name == "vote":
        authorization["voting_id"] = "23887"
    elif name == "decision":
        findings["D230"]["decision_api_id"] = "40034099"
    elif name in {"hidden-gaps", "old-identity"}:
        old = {"report": {"basis": basis["p1_basis"], "view": report["view"]["p1"]}}
        if name == "hidden-gaps":
            old["report"]["basis"]["gaps"] = []
            old["report"]["view"]["gap_ids"] = []
        else:
            old["report"]["basis"]["case_candidates"][0]["case_id"] = "999999"
        rehash_p1(old)
        basis["p1_basis"] = old["report"]["basis"]
        report["view"]["p1"] = old["report"]["view"]
    elif name == "hidden-candidates":
        report["view"]["p1"]["case_candidate_ids"] = []
        report["view"]["p1"]["decision_candidate_ids"] = []
    elif name == "hidden-status":
        basis["status_updates"] = []
        report["view"]["status_update_ids"] = []
    elif name == "hidden-crossrefs":
        basis["cross_references"] = []
        report["view"]["cross_reference_ids"] = []
    elif name == "positive-coverage":
        basis["year_coverage"] = "complete"
    elif name == "annual-approval":
        basis["professional_annual_balance_approval"] = True
    elif name == "source-bytes":
        payload = basis["payloads"][0]
        old_digest = payload["id"]
        value = base64.b64decode(payload["content_base64"]) + b" "
        digest = hashlib.sha256(value).hexdigest()
        payload.update(
            id=digest,
            sha256=digest,
            byte_count=len(value),
            content_base64=base64.b64encode(value).decode(),
        )
        for source in basis["sources"]:
            if source["payload_id"] == old_digest:
                source.update(payload_id=digest, expected_sha256=digest)
    elif name == "duplicate-payload":
        basis["payloads"].append(basis["payloads"][0])
    elif name == "old-payload-copy":
        basis["payloads"].append(basis["p1_basis"]["payloads"][0])
    elif name == "text-loss":
        basis["evidence"][0]["text"] = "synthetic truncated evidence"
    else:
        raise ValueError(name)
    _rehash(data)
    return encode(data)


def selection_cases(basis: Any) -> list[tuple[str, PilotSelection, tuple[int, int]]]:
    cases = [
        ("post1", PilotSelection(years=(2024,), accounts=(("1429", "1"),)), (7, 3)),
        ("post22", PilotSelection(years=(2024,), accounts=(("1429", "22"),)), (4, 2)),
        ("both", PilotSelection(years=(2024,), chapters=("1429",)), (11, 5)),
        ("wrong-year", PilotSelection(years=(2025,), chapters=("1429",)), (0, 0)),
        ("empty", PilotSelection(), (0, 0)),
    ]
    docs = {d.key: d.id for d in contents(basis).documents}
    for key, count in (
        ("prop34", (0, 0)),
        ("inns91", (0, 0)),
        ("vedtak230", (0, 0)),
        ("vedtak231", (0, 0)),
        ("vedtak429", (0, 0)),
        ("prop36", (0, 0)),
        ("inns90", (0, 0)),
        ("forslag23887", (0, 0)),
        ("votering23887", (0, 0)),
        ("forslag23945", (0, 0)),
        ("votering23945", (0, 0)),
        ("vedtak341", (2, 2)),
        ("vedtak857", (3, 1)),
        ("vedtak210", (6, 2)),
        ("vedtak346", (0, 0)),
    ):
        cases.append((key, PilotSelection(mode="document", document_id=docs[key]), count))
    index = next(c for c in basis.p1_basis.case_candidates if c.case_id == "101000")
    cases.append(
        ("index101000", PilotSelection(mode="document", document_id=index.document_id), (0, 0))
    )
    return cases


def save_pair(output: Path, name: str, report: Report, counts: tuple[int, int]) -> dict[str, Any]:
    frozen, excel = output / (name + ".json"), output / (name + ".xlsx")
    export_pair(report, frozen, excel)
    saved = read_json(frozen)
    assert saved == read_excel(excel) == report
    old = saved.view.p1
    assert (
        len(old.monetary.adopted_observation_ids),
        len(old.monetary.adopted_event_ids),
    ) == counts
    assert old.case_candidate_ids == tuple(c.id for c in report.basis.p1_basis.case_candidates)
    assert old.decision_candidate_ids == tuple(
        c.id for c in report.basis.p1_basis.decision_candidates
    )
    assert old.gap_ids == tuple(g.id for g in report.basis.p1_basis.gaps)
    assert saved.view.status_update_ids == tuple(s.id for s in report.basis.status_updates)
    assert saved.view.cross_reference_ids == tuple(c.id for c in report.basis.cross_references)
    book = load_workbook(excel, read_only=True)
    try:
        coverage = list(book["Dekning og avvik"].values)
        for gap in report.basis.p1_basis.gaps:
            assert [gap.key + " " + gap.title, gap.status + ": " + gap.reason] in [
                list(r) for r in coverage
            ]
        for candidate in report.basis.p1_basis.case_candidates:
            assert any(
                candidate.id in str(r[0]) and candidate.reason in str(r[1]) for r in coverage
            )
        assert any("hovedsak_undersokt_faktisk_postvirkning_uavklart" in str(r) for r in coverage)
        typed = [
            {
                "row": row[0].row,
                "account": row[0].value,
                "year": row[1].value,
                "adopted": [c.value for c in row[6:10]],
                "research": row[19].value,
                "research_type": row[19].data_type,
                "research_format": row[19].number_format,
                "role": row[20].value,
                "conditional": row[21].value,
                "target": row[22].value,
            }
            for row in book["Forløp"].iter_rows(min_row=2)
        ]
    finally:
        book.close()
    return {
        "name": name,
        "report_id": report.id,
        "counts": counts,
        "global": [67, 25, 7],
        "typed_cells": typed,
        "files": [evidence(p, output) for p in (frozen, excel)],
    }


def negative_pairs(output: Path, report: Report) -> list[dict[str, Any]]:
    directory = output / "negative"
    directory.mkdir()
    raw = report_to_json(report)
    records = []
    for name in MUTATIONS:
        poisoned = mutate(raw, name)
        frozen, excel = directory / (name + ".json"), directory / (name + ".xlsx")
        frozen.write_text(poisoned, encoding="utf8")
        typed: Report = _decode(Report, json.loads(poisoned)["report"])
        book = _book(typed, poisoned)
        book.save(excel)
        book.close()
        capsule, _ = inspect_excel(
            excel,
            sheet_names=SHEET_NAMES,
            snapshot_name="Offlinegrunnlag",
            max_expanded_bytes=128 * 1024 * 1024,
        )
        assert capsule == poisoned
        errors = {}
        for kind, reader, path in (("json", read_json, frozen), ("xlsx", read_excel, excel)):
            try:
                reader(path)
            except ValueError as error:
                errors[kind] = str(error)
            else:
                raise AssertionError("Syntetisk angrep akseptert: " + name)
        records.append(
            {
                "mutation": name,
                "synthetic": True,
                "rehashed": True,
                "container_passed_exact_capsule": True,
                "rejections": errors,
                "files": [evidence(p, output) for p in (frozen, excel)],
            }
        )
    return records


def source_audit(report: Report) -> dict[str, Any]:
    from analyse.komitepilot import legacy_payloads, legacy_report

    old = report.basis.p1_basis.t4_basis
    legacy = legacy_payloads(old.legacy_json)
    previous = legacy_report(old.legacy_json)
    shared = shared_sources(report.basis.p1_basis, legacy, previous.basis if previous else None)
    sizes = []
    for name, (pin, count, _, _, _) in CATALOG.items():
        raw = Path(name).read_bytes()
        assert len(raw) == count and hashlib.sha256(raw).hexdigest() == pin
        body = representation(raw, name)
        blob = subprocess.check_output(["git", "show", "HEAD:" + name])
        assert blob == raw
        sizes.append(
            {
                "name": name,
                "bytes": len(raw),
                "representation_bytes": len(body),
                "sha256": pin,
                "git_blob_equal": True,
                "old_source_reused": name in shared,
            }
        )
    old_digests = {hashlib.sha256(raw).hexdigest() for _, raw, _ in shared.values()}
    assert not old_digests.intersection(p.id for p in report.basis.payloads)
    return {
        "pinned_files": len(sizes),
        "files": sizes,
        "new_payloads": len(report.basis.payloads),
        "physical_new_payloads": sum(p.content_base64 is not None for p in report.basis.payloads),
        "duplicate_new_raw_payloads": 0,
        "old_immutable_duplication": "bevart separat i P1/T4/T1-kapslene",
        "max_new_payload_bytes": max(p.byte_count for p in report.basis.payloads),
        "max_representation_bytes": max(r["representation_bytes"] for r in sizes),
    }


def offline(output: Path, report: Report) -> dict[str, Any]:
    frozen = output / "both.json"
    working = output / "empty-working"
    working.mkdir()
    code = r"""
import hashlib,json,runpy,socket,sqlite3,sys
from pathlib import Path
def forbidden(*args, **kwargs):
    raise AssertionError("network/sqlite forbidden")
socket.create_connection=forbidden
socket.socket.connect=forbidden
socket.socket.connect_ex=forbidden
sqlite3.connect=forbidden
sys.argv=["analyse.perioderegister_r2", "--frosset", sys.argv[1], "--rot", sys.argv[2],
          "--json",sys.argv[3],"--excel",sys.argv[4],"--originaler",sys.argv[5]]
runpy.run_module("analyse.perioderegister_r2",run_name="__main__")
from analyse.perioderegister_r2 import read_json
from analyse.perioderegister_r2_excel import read_excel
saved=read_json(Path(sys.argv[sys.argv.index("--json")+1]))
assert read_excel(Path(sys.argv[sys.argv.index("--excel")+1]))==saved
"""
    env = dict(os.environ)
    env["PYTHONPATH"] = os.pathsep.join((str(Path.cwd()), os.environ.get("PYTHONPATH", "")))
    args = [
        sys.executable,
        "-c",
        code,
        str(frozen),
        str(working / "no-original-root"),
        str(output / "offline.json"),
        str(output / "offline.xlsx"),
        str(output / "offline-originals"),
    ]
    result = subprocess.run(args, cwd=working, env=env, capture_output=True, text=True, check=False)
    assert result.returncode == 0, result.stdout + result.stderr
    assert read_json(output / "offline.json") == read_excel(output / "offline.xlsx") == report
    index = json.loads((output / "offline-originals/index.json").read_text("utf8"))
    for digest, count in index["payloads"].items():
        raw = (output / "offline-originals" / (digest + ".bin")).read_bytes()
        assert len(raw) == count and hashlib.sha256(raw).hexdigest() == digest
    for payload in report.basis.payloads:
        assert (output / "offline-originals" / (payload.id + ".bin")).read_bytes() == payload_bytes(
            payload
        )
    assert index["evidence"] == json.loads(encode([e.__dict__ for e in report.basis.evidence]))
    return {
        "returncode": result.returncode,
        "stdout": result.stdout,
        "network_sqlite_forbidden": True,
        "missing_original_root": True,
        "exported_raw_digests": len(index["payloads"]),
        "json_excel_equal": True,
    }


def independent_368(output: Path, report: Report) -> list[dict[str, Any]]:
    from innlesning import read_extraction
    from innlesning.registers import read_register, seed_register
    from innlesning.source_attempts import begin_source_attempt

    records = []
    for condition in ("missing", "failed", "corrupt", "contradictory"):
        directory = output / ("independence-" + condition)
        directory.mkdir()
        if condition == "missing":
            missing = import_basis(directory / "no-root")
            assert missing.issues
            empty = build_report(missing, PilotSelection())
            export_pair(empty, directory / "missing.json", directory / "missing.xlsx")
            assert read_json(directory / "missing.json") == read_excel(directory / "missing.xlsx")
        else:
            try:
                if condition == "failed":
                    import_basis(directory / "no-root", t1_json="failed enrichment")
                else:
                    path = directory / "bad.json"
                    path.write_text(
                        "corrupt P2"
                        if condition == "corrupt"
                        else mutate(report_to_json(report), "estimate-adopted"),
                        encoding="utf8",
                    )
                    read_json(path)
            except ValueError:
                pass
            else:
                raise AssertionError(condition)
        database = directory / "ordinary.sqlite"
        begin_source_attempt(database, "prove", "2024-2025")
        register = seed_register(database)
        ordinary = _ordinary(directory)
        assert ordinary["observations"] == 222 and ordinary["events"] == 74
        assert read_register(database, register.id) == register
        extracted = read_extraction(database, str(ordinary["extraction_id"]))
        assert {o.round for o in extracted.observations} == {"uavklart"}
        with sqlite3.connect(database) as connection:
            assert connection.execute("PRAGMA user_version").fetchone()[0] == 5
        trio = {
            o.meaning: o.amount_kr
            for o in extracted.observations
            if o.chapter == "700" and o.post == "1"
        }
        assert set(trio.values()) == {304351000, 6324000, 310675000}
        records.append(
            {
                "condition": condition,
                "ordinary": ordinary,
                "schema": 5,
                "b5_register_unchanged": register.id,
                "old_round_values": ["uavklart"],
                "700/1": trio,
            }
        )
    return records


def compatibility(previous: Path, output: Path, legacy_b6: Path) -> list[dict[str, Any]]:
    records = _compatibility(previous, output / "compatibility")
    for kind, folder, name in (
        ("T1", "t1-41-608895f", "begge.json"),
        ("T4", "t4-48-3f856d8", "2024-both.json"),
    ):
        records.append(
            old_pair(previous / "outputs" / folder / name, output / "compatibility" / kind, kind)
        )
    original = previous / "outputs/p1-52-973db1f/p1-1429-2024.json"
    copied = output / "compatibility/P1"
    copied.mkdir()
    hashes = []
    for path in (original, original.with_suffix(".xlsx")):
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        hashes.append(digest)
        shutil.copyfile(path, copied / path.name)
        assert hashlib.sha256((copied / path.name).read_bytes()).hexdigest() == digest
    old = p1.read_json(copied / original.name)
    assert p1_excel.read_excel(copied / original.with_suffix(".xlsx").name) == old
    assert hashes == [
        hashlib.sha256(p.read_bytes()).hexdigest()
        for p in (original, original.with_suffix(".xlsx"))
    ]
    records.append(
        {"kind": "ekte_tidligere_P1", "report_id": old.id, "hashes": hashes, "unchanged": True}
    )
    records.append(verify_legacy_b6(legacy_b6, output / "compatibility/B6"))
    return records


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--previous-root", type=Path, required=True)
    parser.add_argument("--legacy-b6", type=Path, required=True)
    parser.add_argument("--resume-after", choices=("pairs", "negatives", "offline", "independence"))
    args = parser.parse_args()
    started = time.monotonic()
    output = args.output.resolve()
    if args.resume_after:
        assert (output / "both.json").is_file()
    else:
        output.mkdir(parents=True, exist_ok=False)
    checkpoint = output / "progress.json"
    result = json.loads(checkpoint.read_text("utf8")) if args.resume_after else {}

    def save() -> None:
        checkpoint.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf8")

    old = gzip.decompress(Path("tests/fixtures/t2/authentic-t1.json.gz").read_bytes()).decode()
    basis = (
        read_json(output / "both.json").basis
        if args.resume_after
        else import_basis(Path("."), t1_json=old)
    )
    if not args.resume_after:
        records = []
        for name, selection, counts in selection_cases(basis):
            report = build_report(basis, selection)
            records.append(save_pair(output, name, report, counts))
            result["selections"] = records
            save()
        both = read_json(output / "both.json")
        # Begge retninger faktisk skrevet, lukket og gjenlest, med samme startutvalg.
        for key in ("prop34", "inns91", "vedtak230", "vedtak231", "prop36", "vedtak429"):
            doc = next(d for d in contents(basis).documents if d.key == key)
            document = open_document(both, doc.id)
            assert go_back(document) == both
            records.append(save_pair(output, "navigate-" + key, document, (0, 0)))
            records.append(
                save_pair(
                    output,
                    "return-" + key,
                    go_back(read_json(output / ("navigate-" + key + ".json"))),
                    (11, 5),
                )
            )
            account = open_accounts(document, both.selection)
            assert go_back(account) == build_report(basis, document.selection)
            records.append(save_pair(output, "reverse-" + key, account, (11, 5)))
            records.append(save_pair(output, "reverse-return-" + key, go_back(account), (0, 0)))
            save()
        records.append(save_pair(output, "repeat-both", both, (11, 5)))
        save()
    both = read_json(output / "both.json")
    if args.resume_after not in {"negatives", "offline", "independence"}:
        result["negative_pairs"] = negative_pairs(output, both)
        save()
    if args.resume_after not in {"offline", "independence"}:
        result["offline"] = offline(output, both)
        save()
    if args.resume_after != "independence":
        result["independence"] = independent_368(output, both)
        save()
    result["compatibility"] = compatibility(
        args.previous_root.resolve(), output, args.legacy_b6.resolve()
    )
    result["source_audit"] = source_audit(both)
    display = []
    for name in (
        "amount",
        "hidden-row",
        "hidden-column",
        "hyperlink",
        "filter",
        "doctype",
        "numberformat",
        "hidden-sheet",
    ):
        path = output / ("display-" + name + ".xlsx")
        mutate_display(output / "both.xlsx", path, name)
        try:
            read_excel(path)
        except ValueError as error:
            display.append({"name": name, "rejection": str(error)})
        else:
            raise AssertionError(name)
    result["display_negatives"] = display
    result.update(
        {
            "python": platform.python_version(),
            "lxml": etree.LXML_VERSION,
            "libxml2_runtime": etree.LIBXML_VERSION,
            "basis_id": basis.id,
            "report_bytes": len(report_to_json(both).encode()),
            "new_sources": len(basis.sources),
            "inventory_per_document": dict(Counter(i.document_id for i in basis.inventory)),
            "seconds_this_invocation": round(time.monotonic() - started, 3),
            "limitations": [
                "Ingen native Excel, årsaksept, kontinuitet eller drift",
                "Gamle immutable kapslers fysiske duplisering beholdes",
            ],
            "files": [evidence(p, output) for p in sorted(output.rglob("*")) if p.is_file()],
        }
    )
    (output / "verification.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf8"
    )
    print(
        encode(
            {
                "pairs": len(result["selections"]),
                "negative_pairs": len(result["negative_pairs"]),
                "seconds": result["seconds_this_invocation"],
            }
        )
    )


if __name__ == "__main__":
    main()
