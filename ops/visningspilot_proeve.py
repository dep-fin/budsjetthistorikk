"""T1-aksept fra ekte O2b-leveranse til nye filer. Gamle filer leses via kopier.

python -m ops.visningspilot_proeve --previous-root C:/.../Budsjetthistorikk
    --output work/t1-leveranse
"""

import argparse
import hashlib
import json
import shutil
import subprocess
import sys
from dataclasses import asdict
from pathlib import Path

from openpyxl import load_workbook

from analyse.omtale import read_excel as read_old_excel
from analyse.omtale import read_json as read_old_json
from analyse.omtale import report_to_json as old_json
from analyse.saksdokumenter import report_from_json as read_saks
from analyse.saksdokumenter import verify_saks_excel
from analyse.session_serialization import report_from_json as read_b6
from analyse.visningspilot import (
    build_report,
    go_back,
    import_basis,
    open_accounts,
    open_document,
    read_json,
)
from analyse.visningspilot_excel import export_originals, export_pair, read_excel
from innlesning import read_extraction, read_source
from innlesning.observations import encode
from innlesning.visningspilot import SOURCE_CATALOG
from innlesning.visningspilot_model import PilotReport, PilotSelection


def evidence(path: Path, root: Path) -> dict[str, str | int]:
    return {
        "file": path.relative_to(root).as_posix(),
        "bytes": path.stat().st_size,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }


def _run(output: Path, name: str, arguments: list[str]) -> dict[str, object]:
    command = [sys.executable, "-m", *arguments]
    result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8")
    record = {
        "command": command,
        "returncode": result.returncode,
        "stdout": result.stdout,
        "stderr": result.stderr,
    }
    (output / (name + ".json")).write_text(json.dumps(record, ensure_ascii=False, indent=2), "utf8")
    if result.returncode not in (0, 2):
        raise ValueError("Ordinær kommandoprøve feilet: " + result.stdout + result.stderr)
    return json.loads(result.stdout)


def _ordinary(output: Path) -> dict[str, object]:
    database = output / "ordinary.sqlite"
    summary = _run(
        output,
        "01-synkroniser",
        [
            "innlesning",
            "synkroniser",
            "--database",
            str(database),
            "--sesjon",
            "2024-2025",
            "--historisk-proeve",
            "--vedtak-id",
            "40034229",
            "--response",
            "tests/fixtures/b6/40034229.xml",
            "--metadata",
            "tests/fixtures/b6/40034229.json",
        ],
    )
    assert summary["antall"]["ny"] == 1 and not summary["feil"]
    revision = summary["vedtak"][0]["revision_id"]
    interpreted = _run(
        output,
        "02-tolk",
        [
            "innlesning",
            "tolk",
            "--database",
            str(database),
            "--revisjon",
            revision,
            "--tolkning",
            "storting-belop-v2",
        ],
    )
    extraction = read_extraction(database, str(interpreted["belopsuttrekk"]))
    assert len({o.step_id for o in extraction.observations}) == 74
    assert (
        read_source(database, extraction.artifact.sha256)
        == Path("tests/fixtures/b6/40034229.xml").read_bytes()
    )
    ordinary = output / "ordinary.xlsx"
    _run(
        output,
        "03-analyse",
        [
            "analyse",
            "--database",
            str(database),
            "--belopsuttrekk",
            extraction.id,
            "--excel",
            str(ordinary),
        ],
    )
    book = load_workbook(ordinary, read_only=True, data_only=False)
    try:
        rows = list(book["Beløpsobservasjoner"].iter_rows(min_row=2, values_only=True))
        assert {r[0] for r in rows} == {o.id for o in extraction.observations}
        headers = {c.value: n for n, c in enumerate(book["Beløpsobservasjoner"][1])}
        trio = {
            (r[headers["Beløpsbetydning"]], r[headers["Kroner"]])
            for r in rows
            if r[headers["Kapittel"]] == "700" and r[headers["Post"]] == "1"
        }
        assert trio == {
            ("for_belop", 304_351_000),
            ("endring", 6_324_000),
            ("etter_belop", 310_675_000),
        }
    finally:
        book.close()
    return {
        "revision": revision,
        "extraction_id": extraction.id,
        "observations": len(extraction.observations),
        "events": 74,
        "raw_sha256": extraction.artifact.sha256,
        "actual_saved_700_1_trio": sorted(trio),
        "raw_preserved": True,
    }


def _compatibility(previous: Path, output: Path) -> list[dict[str, object]]:
    result = []
    for directory, names in (
        (
            "o2b-210-begge-poster-7e70506",
            ("begge", "post1", "post22", "tomt", "offline", "syntetisk-forslag"),
        ),
        ("o2a-210-post1-52f0545", ("omtale-210-post1", "offline-gjenlest")),
    ):
        for name in names:
            source = previous / "outputs" / directory
            copied = output / directory / name
            copied.mkdir(parents=True)
            before = {
                suffix: hashlib.sha256((source / (name + suffix)).read_bytes()).hexdigest()
                for suffix in (".json", ".xlsx")
            }
            for suffix in (".json", ".xlsx"):
                shutil.copyfile(source / (name + suffix), copied / (name + suffix))
                assert (
                    hashlib.sha256((copied / (name + suffix)).read_bytes()).hexdigest()
                    == (before[suffix])
                )
            frozen, excel = copied / (name + ".json"), copied / (name + ".xlsx")
            report = read_old_json(frozen)
            assert read_old_excel(excel) == report
            assert old_json(report) == frozen.read_text("utf8")
            after = {
                suffix: hashlib.sha256((source / (name + suffix)).read_bytes()).hexdigest()
                for suffix in before
            }
            assert before == after
            result.append(
                {
                    "kind": "ekte_tidligere_O2",
                    "name": directory + "/" + name,
                    "report_id": report.id,
                    "original_hashes": before,
                    "unchanged": True,
                }
            )
    for directory, columns in (
        ("vedtak-210-2026-10-22", 39),
        ("vedtak-210-13-kolonner-2026-10-22", 13),
    ):
        source = previous / "outputs" / directory
        target = output / directory
        target.mkdir()
        before = {
            suffix: hashlib.sha256((source / ("vedtak-210" + suffix)).read_bytes()).hexdigest()
            for suffix in (".json", ".xlsx")
        }
        for suffix in before:
            shutil.copyfile(source / ("vedtak-210" + suffix), target / ("vedtak-210" + suffix))
            assert (
                hashlib.sha256((target / ("vedtak-210" + suffix)).read_bytes()).hexdigest()
                == (before[suffix])
            )
        report = read_saks((target / "vedtak-210.json").read_text("utf8"))
        verify_saks_excel(target / "vedtak-210.xlsx", report)
        # B6-kjernen i den ekte gamle B8-konvolutten, ingen regenerert gammel leveranse.
        old_b6 = read_b6(encode({"format": "b6-frozen-1", "report": asdict(report.base)}))
        assert old_b6 == report.base
        book = load_workbook(target / "vedtak-210.xlsx", read_only=True)
        try:
            assert book["Historikk"].max_column == columns
            assert len(book.sheetnames) == (21 if columns == 39 else 22)
        finally:
            book.close()
        assert before == {
            suffix: hashlib.sha256((source / ("vedtak-210" + suffix)).read_bytes()).hexdigest()
            for suffix in before
        }
        result.append(
            {
                "kind": "ekte_tidligere_B8_B9_med_B6_kjerne",
                "name": directory,
                "report_id": report.id,
                "original_hashes": before,
                "unchanged": True,
                "standalone_original_B6_delivery": (
                    "ikke funnet; innebygd original B6-kjerne kontrollert"
                ),
            }
        )
    return result


def verify_legacy_b6(source: Path, output: Path) -> dict[str, object]:
    """Les ekte eldre B6-bok med bevarte observasjonsmodeller, uten historisk DB."""
    from analyse.session_serialization import _decode
    from innlesning.observations import AmountObservation

    expected_hash = "8487909b5e7b2b7498c9ab78de9fb88966be926202df89388ef3f31871ac1f39"
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    assert before == expected_hash and source.stat().st_size == 4_648_442
    output.mkdir(parents=True, exist_ok=False)
    copied = output / "accept-b6.xlsx"
    shutil.copyfile(source, copied)
    assert hashlib.sha256(copied.read_bytes()).hexdigest() == before
    book = load_workbook(copied, read_only=True, data_only=False, keep_links=False)
    try:
        assert len(book.sheetnames) == 21
        assert book["Sesjonsgrunnlag"]["B2"].value == (
            "3ff5698de68eaad7f200d08fb73d83c150adea56ae49ccc2b1bd2efa305a9cfc"
        )
        assert book["Kildeinventar"].max_row == 2184
        assert book["Historikk"].max_column == 33
        sheet = book["Beløpsobservasjoner"]
        assert sheet.max_row == 5493 and sheet.max_column == 36
        identifiers = []
        extraction_ids = set()
        for row in sheet.iter_rows(min_row=2, values_only=True):
            observation: AmountObservation = _decode(AmountObservation, json.loads(str(row[34])))
            assert observation.id == row[0]
            assert observation.amount_kr == row[18]
            assert observation.meaning == row[16]
            assert observation.pointer.artifact_id == row[28]
            assert observation.pointer.locator == row[29]
            assert observation.figure_type == "stortingsvedtak"
            assert row[21] is None and row[23] is None  # ingen B5-godkjenning
            identifiers.append(observation.id)
            extraction_ids.add(row[1])
        assert len(set(identifiers)) == 5492 and len(extraction_ids) == 47
        selection = json.loads(str(book["Sesjonsgrunnlag"]["B4"].value))
        assert len(selection["extraction_ids"]) == 50
        assert extraction_ids <= set(selection["extraction_ids"])
        assert not selection["controlled_result_ids"]
    finally:
        book.close()
    assert hashlib.sha256(source.read_bytes()).hexdigest() == before
    assert hashlib.sha256(copied.read_bytes()).hexdigest() == before
    return {
        "kind": "ekte_eldre_B6_arbeidsbok",
        "source_path": str(source),
        "sha256": before,
        "bytes": 4_648_442,
        "report_id": "3ff5698de68eaad7f200d08fb73d83c150adea56ae49ccc2b1bd2efa305a9cfc",
        "observations": 5492,
        "selected_extractions": 50,
        "extractions_with_amount_rows": 47,
        "observation_ids_sha256": hashlib.sha256(encode(identifiers).encode()).hexdigest(),
        "actual_cells_and_old_observation_models_checked": True,
        "unchanged": True,
        "boundary": (
            "Ekte eldre B6-Excel, ikke omgenerert. Selvstendig gammel frozen-v1-rapport-JSON "
            "er ikke funnet. Utvalgs-/bevis-JSON er ikke SessionReport. "
            "Eksakt historisk writercommit er ikke fastslått; ingen historisk DB åpnet."
        ),
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--previous-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--legacy-b6", type=Path)
    args = parser.parse_args()
    output: Path = args.output
    output.mkdir(parents=True, exist_ok=False)
    old = args.previous_root / "outputs/o2b-210-begge-poster-7e70506/begge.json"
    old_raw = old.read_text("utf8")
    basis = import_basis(old_raw, Path("."))
    docs = {d.key: d for d in basis.documents}
    reports: list[tuple[str, PilotReport]] = []
    for name, selection in (
        ("post1", PilotSelection(years=(2024,), accounts=(("1429", "1"),))),
        ("post22", PilotSelection(years=(2024,), accounts=(("1429", "22"),))),
        ("begge", PilotSelection(years=(2024,), accounts=(("1429", "1"), ("1429", "22")))),
        ("post60", PilotSelection(years=(2024,), accounts=(("1429", "60"),))),
        ("tomt", PilotSelection(years=(2024,))),
        ("kapittel", PilotSelection(years=(2024,), chapters=("1429",))),
        (
            "par-uten-kryssprodukt",
            PilotSelection(years=(2024,), accounts=(("1429", "1"), ("1420", "22"))),
        ),
        (
            "kapittel-og-par",
            PilotSelection(years=(2024,), chapters=("1420",), accounts=(("1429", "1"),)),
        ),
        ("2025-kontekst", PilotSelection(years=(2025,), accounts=(("1429", "60"),))),
    ):
        reports.append((name, build_report(basis, selection)))
    for key in docs:
        reports.append(
            (
                "dokument-" + key,
                build_report(basis, PilotSelection(mode="document", document_id=docs[key].id)),
            )
        )
    start = reports[0][1]
    navigation = open_document(start, docs["inns81"].id)
    assert go_back(navigation) == start
    reports.extend((("konto-til-dokument", navigation), ("tilbake-til-konto", go_back(navigation))))
    start_document = next(r for name, r in reports if name == "dokument-vedtak211")
    navigation = open_accounts(
        start_document, PilotSelection(years=(2024,), accounts=(("1429", "60"),))
    )
    assert go_back(navigation) == start_document
    reports.extend(
        (("dokument-til-konto", navigation), ("tilbake-til-dokument", go_back(navigation)))
    )
    results = []
    for name, report in reports:
        frozen, excel = output / (name + ".json"), output / (name + ".xlsx")
        export_pair(report, frozen, excel)
        assert read_json(frozen) == read_excel(excel) == report
        results.append(
            {
                "selection": name,
                "report_id": report.id,
                "fields": len(report.view.observation_ids),
                "adopted_fields": len(report.view.adopted_observation_ids),
                "adopted_events": len(report.view.adopted_event_ids),
                "comparisons": len(report.view.comparisons),
                "files": [evidence(frozen, output), evidence(excel, output)],
            }
        )
    offline = read_json(output / "begge.json")
    export_pair(offline, output / "offline.json", output / "offline.xlsx")
    assert (output / "offline.json").read_bytes() == (output / "begge.json").read_bytes()
    export_originals(offline, output / "gjenfunne-originaler")
    failures = []
    for condition in ("missing", "failed", "corrupt", "contradictory"):
        directory = output / ("syntetisk-" + condition)
        directory.mkdir()
        source_root = directory / "sources"
        for name in SOURCE_CATALOG:
            target = source_root / name
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(name, target)
        source = source_root / "tests/fixtures/b8/40034049.xml"
        if condition in {"missing", "failed"}:
            source.unlink()
            if condition == "failed":
                source.mkdir()
        elif condition == "corrupt":
            source.write_bytes(b"synthetic corrupt XML")
        else:
            source.write_text(
                Path("tests/fixtures/b8/40034049.xml").read_text("utf8").replace("1429", "1420"),
                "utf8",
            )
        bad = build_report(import_basis(old_raw, source_root), PilotSelection(years=(2024,)))
        export_pair(bad, directory / "tomt.json", directory / "tomt.xlsx")
        failures.append(
            {
                "condition": condition,
                "pilot_issues": bad.basis.issues,
                "ordinary": _ordinary(directory),
            }
        )
    # Faktisk bortfalt numerisk nyadapter: gammel O2b-tekst beholdes, differansen er hull.
    missing_stage = output / "syntetisk-manglende-tilraading"
    missing_stage.mkdir()
    source_root = missing_stage / "sources"
    for name in SOURCE_CATALOG:
        target = source_root / name
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(name, target)
    (source_root / "docs/research/omtale-o1/originals/inns81s.html").unlink()
    gap = build_report(
        import_basis(old_raw, source_root),
        PilotSelection(years=(2024,), accounts=(("1429", "1"),)),
    )
    export_pair(gap, missing_stage / "hull.json", missing_stage / "hull.xlsx")
    assert len(gap.view.observation_ids) == 6
    assert all(c.status == "hull" and c.difference_kr is None for c in gap.view.comparisons)
    compatibility = _compatibility(args.previous_root, output / "compatibility")
    if args.legacy_b6:
        compatibility.append(verify_legacy_b6(args.legacy_b6, output / "compatibility/ekte-B6"))
    manifest = {
        "kind": (
            "T1: autentisk pilot, særskilt merkede syntetiske feil, "
            "ingen faglig/års-/driftsgodkjenning"
        ),
        "basis_id": basis.id,
        "rule": basis.rule,
        "old_o2b_report_id": read_old_json(old).id,
        "selections": results,
        "independence": failures,
        "compatibility": compatibility,
        "missing_numeric_stage": {
            "report_id": gap.id,
            "fields": len(gap.view.observation_ids),
            "comparisons": [asdict(c) for c in gap.view.comparisons],
            "files": [
                evidence(missing_stage / "hull.json", output),
                evidence(missing_stage / "hull.xlsx", output),
            ],
        },
        "limitations": [
            "Selvstendig gammel frozen-v1-rapport-JSON er ikke funnet; "
            "ekte gammel innebygd B6-kjerne kontrollert. Eldre Excel prøves med --legacy-b6.",
            "Native Excel ikke prøvd; lagrede celler/lenker/fulltekst "
            "og rendering kontrolleres separat",
            "T2-T6 og K1-K3 ikke levert; #10 åpen",
        ],
        "files": [evidence(p, output) for p in sorted(output.rglob("*")) if p.is_file()],
    }
    (output / "kontroll.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), "utf8"
    )
    print(
        encode(
            {
                "output": str(output),
                "basis_id": basis.id,
                "pairs": len(reports) + 1,
                "genuine_old_pairs": sum(
                    item["kind"] != "ekte_eldre_B6_arbeidsbok" for item in compatibility
                ),
                "genuine_old_B6_workbooks": sum(
                    item["kind"] == "ekte_eldre_B6_arbeidsbok" for item in compatibility
                ),
                "independence_cases": len(failures),
            }
        )
    )


if __name__ == "__main__":
    main()
