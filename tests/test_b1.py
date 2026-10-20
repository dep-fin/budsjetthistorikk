import hashlib
import json
import re
from pathlib import Path

import pytest

from innlesning import import_saved, read_report, read_source

FIXTURES = Path(__file__).parent / "fixtures"


def synthetic_metadata(raw: bytes) -> dict[str, object]:
    metadata = json.loads((FIXTURES / "40033944.metadata.json").read_text(encoding="utf-8"))
    metadata.update(
        fixture_kind="synthetic_test",
        synthetic_parent_sha256=metadata["sha256"],
        sha256=hashlib.sha256(raw).hexdigest(),
        bytes=len(raw),
    )
    return metadata


def test_original_allocation_is_preserved_and_retrievable(tmp_path: Path) -> None:
    """Fasit transkribert fra kildekontrakten, ikke beregnet med parseren."""
    raw = (FIXTURES / "40033944.derived.xml").read_bytes()
    metadata = json.loads((FIXTURES / "40033944.metadata.json").read_text(encoding="utf-8"))
    database = tmp_path / "budget.sqlite"
    run_id = import_saved(database, raw, metadata, "40033944", historical_sample=True)
    report = read_report(database, run_id)
    assert report.status == "strukturert"
    assert report.decision is not None
    assert (
        report.decision.id,
        report.decision.number,
        report.decision.session,
        report.decision.date,
        report.decision.case_id,
        report.decision.api_type,
        report.decision.budget_year,
    ) == ("40033944", "113", "2024-2025", "2024-12-13", "100215", "BUDSJETT", 2025)
    assert [(line.chapter, line.post, line.amount_kr, line.locator) for line in report.lines] == [
        ("42", "1", 10_478_000, "table[1]/tr[4]"),
        ("42", "75", 250_000, "table[1]/tr[5]"),
        ("43", "1", 118_311_000, "table[1]/tr[7]"),
        ("44", "1", 49_986_000, "table[1]/tr[9]"),
        ("45", "1", 31_322_000, "table[1]/tr[11]"),
        ("51", "1", 669_959_000, "table[1]/tr[13]"),
        ("51", "75", 29_791_000, "table[1]/tr[14]"),
        ("3051", "2", 300_000, "table[1]/tr[18]"),
    ]
    assert [line.direction for line in report.lines] == ["utgift"] * 7 + ["inntekt"]
    assert all(line.meaning == "opprinnelig bevilgning" for line in report.lines)
    assert all(line.status == "strukturert" for line in report.lines)
    assert report.source_totals == {"utgift": 910_097_000, "inntekt": 300_000}
    assert report.approved_totals == {"utgift": 910_097_000, "inntekt": 300_000}
    assert report.source_sha256 == metadata["sha256"]
    assert read_source(database, report.source_sha256) == raw
    assert report.source_metadata["fixture_kind"] == "derived_excerpt"
    assert "På statsbudsjettet for 2025 bevilges under:" in report.decision.html
    assert report.decision.url.endswith("?did=40033944&p=100215")
    assert [line.chapter_name for line in report.lines] == [
        "Stortingets ombudsnemnd for Forsvaret",
        "Stortingets ombudsnemnd for Forsvaret",
        "Stortingets ombud for kontroll med forvaltningen",
        "Stortingets kontrollutvalg for etterretnings-, overvåkings- og sikkerhetstjeneste",
        "Norges institusjon for menneskerettigheter",
        "Riksrevisjonen",
        "Riksrevisjonen",
        "Riksrevisjonen",
    ]
    assert [line.purpose for line in report.lines] == [
        "Driftsutgifter",
        "Internasjonale organisasjoner og nettverk",
        "Driftsutgifter",
        "Driftsutgifter",
        "Driftsutgifter",
        "Driftsutgifter",
        "Internasjonale organisasjoner og nettverk",
        "Refusjon utland",
    ]


def test_excel_matches_the_saved_original_and_has_clickable_source(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    from analyse import export_excel

    raw = (FIXTURES / "40033944.derived.xml").read_bytes()
    metadata = json.loads((FIXTURES / "40033944.metadata.json").read_text(encoding="utf-8"))
    database = tmp_path / "budget.sqlite"
    run_id = import_saved(database, raw, metadata, "40033944", historical_sample=True)
    report = read_report(database, run_id)
    output = tmp_path / "budget.xlsx"
    export_excel(database, run_id, output)
    workbook = load_workbook(output)
    sheet = workbook["Bevilgninger"]
    header = next(row for row in sheet if row[0].value == "Kapittel")
    columns = {cell.value: cell.column for cell in header}
    records = list(sheet.iter_rows(min_row=header[0].row + 1))
    assert len(records) == 8
    assert [row[columns["Beløp (kroner)"] - 1].value for row in records] == [
        line.amount_kr for line in report.lines
    ]
    assert all(
        row[columns["Beløpets betydning"] - 1].value == "opprinnelig bevilgning" for row in records
    )
    assert all(
        row[columns["Kilde"] - 1].hyperlink.target
        == "https://www.stortinget.no/no/Saker-og-publikasjoner/Vedtak/Vedtak/Sak/"
        "?did=40033944&p=100215"
        for row in records
    )
    assert all(row[columns["Endringsbeløp"] - 1].value is None for row in records)
    assert all(row[columns["Førbeløp"] - 1].value is None for row in records)
    assert all(row[columns["Etterbeløp"] - 1].value is None for row in records)
    notes = [cell.value for row in workbook["Kilde og kontroll"] for cell in row]
    assert "derived_excerpt" in notes
    assert metadata["sha256"] in notes
    assert any("historisk prøve" in str(value) for value in notes)
    assert any("Departementstilknytning uavklart" in str(value) for value in notes)
    workbook.close()


@pytest.mark.parametrize("replacement", [b"", b"I 615 833", b"10 478 000 / 11 000 000"])
def test_uncertain_amount_is_retained_without_zero_or_approved_total(
    tmp_path: Path, replacement: bytes
) -> None:
    from openpyxl import load_workbook

    from analyse import export_excel

    raw = (FIXTURES / "40033944.derived.xml").read_bytes().replace(b"10 478 000", replacement)
    database = tmp_path / "budget.sqlite"
    run_id = import_saved(
        database, raw, synthetic_metadata(raw), "40033944", historical_sample=True
    )
    report = read_report(database, run_id)
    assert report.status == "til kontroll"
    assert len(report.lines) == 8
    assert report.lines[0].amount_kr is None
    assert report.lines[0].status == "til kontroll"
    assert report.lines[0].amount_text == replacement.decode()
    assert report.lines[0].control_reason
    assert report.approved_totals == {"utgift": None, "inntekt": None}
    assert read_source(database, report.source_sha256) == raw
    export_excel(database, run_id, tmp_path / "review.xlsx")
    workbook = load_workbook(tmp_path / "review.xlsx")
    sheet = workbook["Bevilgninger"]
    header = next(row for row in sheet if row[0].value == "Kapittel")
    columns = {cell.value: cell.column for cell in header}
    first = list(sheet.iter_rows(min_row=header[0].row + 1))[0]
    assert first[columns["Beløp (kroner)"] - 1].value is None
    assert first[columns["Kontrollstatus"] - 1].value == "til kontroll"
    assert any("ingen godkjent sum" in str(cell.value) for row in sheet for cell in row)
    workbook.close()


def test_unclassified_content_blocks_complete_sample_status(tmp_path: Path) -> None:
    raw = (FIXTURES / "40033944.derived.xml").read_bytes()
    raw = raw.replace(
        b"bevilges under:", b"bevilges under:&lt;/p&gt;&lt;p&gt;Ekstra bevilgning: 8 kroner"
    )
    database = tmp_path / "budget.sqlite"
    run_id = import_saved(
        database, raw, synthetic_metadata(raw), "40033944", historical_sample=True
    )
    report = read_report(database, run_id)
    assert report.status == "til kontroll"
    assert report.approved_totals == {"utgift": None, "inntekt": None}
    assert any("tabellen" in reason for reason in report.controls)


def test_ambiguous_column_spans_do_not_look_like_a_complete_table(tmp_path: Path) -> None:
    raw = (FIXTURES / "40033944.derived.xml").read_bytes()
    raw = raw.replace(b"&lt;td&gt;", b"&lt;td colspan='2'&gt;", 1)
    database = tmp_path / "budget.sqlite"
    run_id = import_saved(
        database, raw, synthetic_metadata(raw), "40033944", historical_sample=True
    )
    report = read_report(database, run_id)
    assert report.status == "til kontroll"
    assert report.approved_totals["utgift"] is None


def test_text_after_table_is_also_retained_as_control_need(tmp_path: Path) -> None:
    raw = (FIXTURES / "40033944.derived.xml").read_bytes()
    raw = raw.replace(b"&lt;/table&gt;", b"&lt;/table&gt;Ekstra bevilgning: 8 kroner")
    database = tmp_path / "budget.sqlite"
    run_id = import_saved(
        database, raw, synthetic_metadata(raw), "40033944", historical_sample=True
    )
    report = read_report(database, run_id)
    assert report.status == "til kontroll"
    assert report.approved_totals["utgift"] is None


@pytest.mark.parametrize("variant", ["invalid_xml", "missing_id", "nil_text", "wrong_session"])
def test_missing_or_invalid_decision_is_a_preserved_error(tmp_path: Path, variant: str) -> None:
    raw = (FIXTURES / "40033944.derived.xml").read_bytes()
    if variant == "invalid_xml":
        raw = b"<invalid"
    elif variant == "missing_id":
        raw = raw.replace(b"<id>40033944</id>", b"<id>99999999</id>")
    elif variant == "nil_text":
        raw = re.sub(
            rb"<stortingsvedtak_tekst>.*?</stortingsvedtak_tekst>",
            b'<stortingsvedtak_tekst i:nil="true" />',
            raw,
            flags=re.DOTALL,
        )
    else:
        raw = raw.replace(b"<sesjon_id>2024-2025</sesjon_id>", b"<sesjon_id>2023-2024</sesjon_id>")
    database = tmp_path / "budget.sqlite"
    run_id = import_saved(
        database, raw, synthetic_metadata(raw), "40033944", historical_sample=True
    )
    report = read_report(database, run_id)
    assert report.status == "innlesningsfeil"
    assert report.lines == []
    assert report.controls
    assert report.approved_totals == {"utgift": None, "inntekt": None}
    assert read_source(database, report.source_sha256) == raw


def test_xml_external_entity_is_preserved_as_data_and_rejected(tmp_path: Path) -> None:
    secret = tmp_path / "private.txt"
    secret.write_text("DO_NOT_LOAD_THIS_CONTENT", encoding="utf-8")
    raw = (FIXTURES / "40033944.derived.xml").read_bytes()
    declaration = (
        '<!DOCTYPE stortingsvedtak_oversikt [<!ENTITY leak SYSTEM "' + secret.as_uri() + '">]>\n'
    ).encode()
    raw = raw.replace(b"<stortingsvedtak_oversikt ", declaration + b"<stortingsvedtak_oversikt ", 1)
    raw = raw.replace(b"Driftsutgifter", b"&leak;", 1)
    database = tmp_path / "budget.sqlite"
    run_id = import_saved(
        database, raw, synthetic_metadata(raw), "40033944", historical_sample=True
    )
    report = read_report(database, run_id)
    assert report.status == "innlesningsfeil"
    assert any("DOCTYPE" in reason for reason in report.controls)
    assert "DO_NOT_LOAD_THIS_CONTENT" not in str(report)
    assert read_source(database, report.source_sha256) == raw


def test_source_formula_remains_text_in_excel(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    from analyse import export_excel

    raw = (FIXTURES / "40033944.derived.xml").read_bytes()
    raw = raw.replace(
        b"Driftsutgifter", b"=HYPERLINK(&quot;https://evil.invalid&quot;,&quot;X&quot;)"
    )
    database = tmp_path / "budget.sqlite"
    run_id = import_saved(
        database, raw, synthetic_metadata(raw), "40033944", historical_sample=True
    )
    output = tmp_path / "review.xlsx"
    export_excel(database, run_id, output)
    workbook = load_workbook(output)
    cells = [cell for sheet in workbook for row in sheet for cell in row]
    assert not any(cell.data_type == "f" for cell in cells)
    assert any(str(cell.value).startswith("=HYPERLINK(") for cell in cells)
    assert all(
        cell.hyperlink.target.startswith("https://www.stortinget.no/")
        for cell in cells
        if cell.hyperlink
    )
    workbook.close()
