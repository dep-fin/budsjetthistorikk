"""B2 through persisted public interfaces; all mutations are synthetic test data."""

import hashlib
import json
import sqlite3
from pathlib import Path

import pytest
from lxml import etree

from innlesning import read_revision, read_source, read_sync, sync_saved

FIXTURES = Path(__file__).parent / "fixtures"
RAW = (FIXTURES / "40033944.derived.xml").read_bytes()


def metadata(raw: bytes, *, tick: int = 0) -> dict[str, object]:
    result = json.loads((FIXTURES / "40033944.metadata.json").read_text(encoding="utf-8"))
    result.update(
        fixture_kind="synthetic_test",
        session_coverage="complete",
        retrieved_utc=f"2026-10-20T14:{tick:02d}:00+00:00",
        sha256=hashlib.sha256(raw).hexdigest(),
        bytes=len(raw),
    )
    return result


def synchronize(database: Path, raw: bytes = RAW, *, tick: int = 0) -> str:
    return sync_saved(database, raw, metadata(raw, tick=tick), historical_sample=True)


def test_repeated_content_and_transport_time_reuse_revision_and_eight_lines(tmp_path: Path) -> None:
    database = tmp_path / "b2.sqlite"
    first = read_sync(database, synchronize(database))
    raw = RAW.replace(b"13:34:46.1397306", b"13:35:46.1397306")
    second = read_sync(database, synchronize(database, raw, tick=1))
    assert first.counts == {"ny": 1, "uendret": 0, "rettelse": 0, "mangler": 0}
    assert second.counts == {"ny": 0, "uendret": 1, "rettelse": 0, "mangler": 0}
    assert first.entries[0].revision_id == second.entries[0].revision_id
    assert first.source_sha256 != second.source_sha256
    revision = read_revision(database, first.entries[0].revision_id)
    assert (
        revision.report.run_id
        == read_revision(database, second.entries[0].revision_id).report.run_id
    )
    assert [(line.chapter, line.post, line.amount_kr) for line in revision.report.lines] == [
        ("42", "1", 10_478_000),
        ("42", "75", 250_000),
        ("43", "1", 118_311_000),
        ("44", "1", 49_986_000),
        ("45", "1", 31_322_000),
        ("51", "1", 669_959_000),
        ("51", "75", 29_791_000),
        ("3051", "2", 300_000),
    ]
    assert read_source(database, first.source_sha256) == RAW
    assert read_source(database, second.source_sha256) == raw
    repeated = read_sync(database, synchronize(database, raw, tick=1))
    assert repeated.counts == {"ny": 0, "uendret": 1, "rettelse": 0, "mangler": 0}
    assert repeated.entries[0].revision_id == first.entries[0].revision_id


def test_amount_correction_keeps_historical_export_and_expires_control_on_return(
    tmp_path: Path,
) -> None:
    from openpyxl import load_workbook

    from analyse import export_revision_excel
    from innlesning import record_control_reference

    database = tmp_path / "b2.sqlite"
    a_job = read_sync(database, synchronize(database))
    a_id = a_job.entries[0].revision_id
    control_id = record_control_reference(
        database,
        a_id,
        '=HYPERLINK("https://evil.invalid","Ekstern kontroll #42")',
        "Sjekket mot original",
        locator="table[1]/tr[4]",
    )
    corrected = RAW.replace(b"10 478 000", b"10 479 000").replace(b"910 097 000", b"910 098 000")
    b_job = read_sync(database, synchronize(database, corrected, tick=1))
    b_id = b_job.entries[0].revision_id
    assert b_job.counts == {"ny": 0, "uendret": 0, "rettelse": 1, "mangler": 0}
    b = read_revision(database, b_id)
    assert b.report.approved_totals == {"utgift": 910_098_000, "inntekt": 300_000}
    assert len(b.report.lines) == 8
    assert b.controls[0].control_id == control_id
    assert b.controls[0].status == "utdatert"
    assert read_revision(database, a_id).freshness == "historisk innhold"
    output = tmp_path / "historisk-a.xlsx"
    export_revision_excel(database, a_id, output)
    workbook = load_workbook(output)
    assert workbook["Bevilgninger"]["F8"].value == 10_478_000
    notes = [cell.value for row in workbook["Kilde og kontroll"] for cell in row]
    assert a_id in notes and b_id in notes and a_job.source_sha256 in notes
    assert "historisk innhold" in notes
    assert any(cell.value == "utdatert" for row in workbook["Kontrollhenvisninger"] for cell in row)
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()
    returned = read_sync(database, synchronize(database, tick=2))
    assert returned.entries[0].revision_id == a_id
    a = read_revision(database, a_id)
    assert a.freshness == "gjeldende observert innhold"
    assert a.controls[0].status == "utdatert"
    assert a.report.approved_totals == {"utgift": 910_097_000, "inntekt": 300_000}


@pytest.mark.parametrize("field", ["reference", "reason", "locator"])
def test_control_text_limit_rejects_before_registration_and_excel_keeps_limit(
    tmp_path: Path, field: str
) -> None:
    from openpyxl import load_workbook

    from analyse import export_revision_excel
    from innlesning import record_control_reference

    database = tmp_path / "b2.sqlite"
    first = read_sync(database, synchronize(database))
    revision_id = first.entries[0].revision_id
    values = {"reference": "Kontroll", "reason": "Kontrollert", "locator": "table[1]/tr[4]"}
    values[field] = "X" * 32768
    with pytest.raises(ValueError, match="32767"):
        record_control_reference(database, revision_id, **values)
    assert read_revision(database, revision_id).controls == []
    reference = "R" * 32760 + "ENDMARK"
    reason = "B" * 32760 + "BEGRUNN"
    control_id = record_control_reference(
        database, revision_id, reference, reason, locator="table[1]/tr[4]"
    )
    controls = read_revision(database, revision_id).controls
    assert len(controls) == 1
    assert controls[0].control_id == control_id
    assert controls[0].reference == reference
    assert controls[0].reason == reason
    output = tmp_path / "kontrollgrense.xlsx"
    export_revision_excel(database, revision_id, output)
    workbook = load_workbook(output)
    assert workbook["Kontrollhenvisninger"]["E2"].value == reference
    assert workbook["Kontrollhenvisninger"]["F2"].value == reason
    assert workbook["Kontrollhenvisninger"]["D2"].value == "table[1]/tr[4]"
    workbook.close()


def variant(*, extra: bool = False, empty: bool = False) -> bytes:
    root = etree.fromstring(RAW)
    ns = "{http://data.stortinget.no}"
    listing = root.find(ns + "stortingsvedtak_liste")
    assert listing is not None
    if extra:
        node = etree.fromstring(etree.tostring(listing[0]))
        for element in node.iter():
            if element.text:
                element.text = element.text.replace("40033944", "90000001")
        node.find(ns + "stortingsvedtak_dato_tid").text = "2001-01-01T00:00:00"
        node.find(ns + "stortingsvedtak_type/" + ns + "id").text = "ANNET"
        node.find(ns + "stortingsvedtak_tekst").text = "Ukjent vedtak uten standardinnledning"
        listing.append(node)
    if empty:
        for node in list(listing):
            listing.remove(node)
    return etree.tostring(root)


def test_full_inventory_keeps_old_dated_other_type_and_missing_return(tmp_path: Path) -> None:
    database = tmp_path / "b2.sqlite"
    first = read_sync(database, synchronize(database))
    combined = variant(extra=True)
    added = read_sync(database, synchronize(database, combined, tick=1))
    assert added.counts == {"ny": 1, "uendret": 1, "rettelse": 0, "mangler": 0}
    unknown = read_revision(database, added.entries[1].revision_id)
    assert unknown.report.decision.api_type == "ANNET"
    assert unknown.report.decision.date == "2001-01-01"
    assert unknown.report.decision.html == "Ukjent vedtak uten standardinnledning"
    assert unknown.report.status == "til kontroll"
    assert unknown.report.approved_totals == {"utgift": None, "inntekt": None}
    absent = read_sync(database, synchronize(database, variant(empty=True), tick=2))
    assert absent.counts == {"ny": 0, "uendret": 0, "rettelse": 0, "mangler": 2}
    assert absent.control_count == 2
    original = read_revision(database, first.entries[0].revision_id)
    assert original.observation_status == "mangler i siste fulle svar"
    assert original.report.lines[0].amount_kr == 10_478_000
    back = read_sync(database, synchronize(database, combined, tick=3))
    assert back.counts == {"ny": 0, "uendret": 2, "rettelse": 0, "mangler": 0}
    assert back.entries[0].revision_id == first.entries[0].revision_id
    assert back.entries[1].revision_id == unknown.revision_id


def test_staged_and_interrupted_publish_resume_without_partial_state(tmp_path: Path) -> None:
    from innlesning import resume_sync, stage_sync

    database = tmp_path / "b2.sqlite"
    original = read_sync(database, synchronize(database))
    changed = variant(extra=True).replace(b"10 478 000", b"I 615 833")
    job_id = stage_sync(database, changed, metadata(changed, tick=1), historical_sample=True)
    assert read_sync(database, job_id).status == "venter"
    with pytest.raises(ValueError, match=job_id):
        synchronize(database, tick=2)
    # Real SQLite fault at publication completion; assertions stay at public interfaces.
    with sqlite3.connect(database) as connection:
        connection.execute(
            "CREATE TRIGGER interruption BEFORE UPDATE ON sync_jobs "
            "WHEN NEW.status='fullfort' BEGIN SELECT RAISE(ABORT, 'power loss'); END"
        )
    with pytest.raises(sqlite3.IntegrityError, match="power loss"):
        resume_sync(database, job_id)
    assert read_sync(database, job_id).status == "venter"
    assert read_sync(database, job_id).entries == []
    before = read_revision(database, original.entries[0].revision_id)
    assert before.current_revision_id == original.entries[0].revision_id
    assert before.report.lines[0].amount_kr == 10_478_000
    with sqlite3.connect(database) as connection:
        connection.execute("DROP TRIGGER interruption")
    assert resume_sync(database, job_id) == job_id
    completed = read_sync(database, job_id)
    assert completed.counts == {"ny": 1, "uendret": 0, "rettelse": 1, "mangler": 0}
    result = read_revision(database, completed.entries[0].revision_id)
    assert result.report.status == "til kontroll"
    assert result.report.lines[0].amount_kr is None
    assert result.report.approved_totals == {"utgift": None, "inntekt": None}
    assert resume_sync(database, job_id) == job_id
    assert read_sync(database, job_id) == completed
    assert read_source(database, completed.source_sha256) == changed


@pytest.mark.parametrize(
    "change",
    [
        (b"<sak_lenke_url>", b"<sak_lenke_url updated='true'>"),
        (b"<stortingsvedtak_nummer>113", b"<stortingsvedtak_nummer>114"),
        (b"2024-12-13T00:00:00", b"2024-12-12T00:00:00"),
        (b"<stortingsvedtak_tittel>", b"<stortingsvedtak_tittel>Ny tittel "),
        (b"<id>BUDSJETT</id>", b"<id>ANNET</id>"),
        (b"</stortingsvedtak>", b"<ukjent i:nil='true'/></stortingsvedtak>"),
        (b"</stortingsvedtak>", b"<ukjent> </ukjent></stortingsvedtak>"),
        (b"Driftsutgifter", b"Driftsutgifter "),
        (b"for 2025 bevilges", b"for 2026 bevilges"),
    ],
)
def test_all_object_values_attributes_and_unknown_fields_are_revision_content(
    tmp_path: Path, change: tuple[bytes, bytes]
) -> None:
    database = tmp_path / "b2.sqlite"
    a = read_sync(database, synchronize(database))
    raw = RAW.replace(*change)
    assert raw != RAW
    b = read_sync(database, synchronize(database, raw, tick=1))
    assert b.status == "fullfort"
    assert b.counts["rettelse"] == 1
    assert b.entries[0].revision_id != a.entries[0].revision_id
    assert (
        read_revision(database, b.entries[0].revision_id).fingerprint
        != read_revision(database, a.entries[0].revision_id).fingerprint
    )


@pytest.mark.parametrize(
    "raw",
    [
        b"<truncated",
        RAW.replace(b"stortingsvedtak_liste", b"missing_list"),
        RAW.replace(b"<id>40033944</id>", b"<id>40033944</id><id>40033944</id>"),
        RAW.replace(b"<id>40033944</id>", b"<id>99999999</id>"),
        RAW.replace(b"<sesjon_id>2024-2025", b"<sesjon_id>2025-2026"),
        RAW.replace(b"<versjon>1.6", b"<versjon>9.9"),
        RAW.replace(b"<stortingsvedtak_tekst>", b"<stortingsvedtak_tekst i:nil='true'>"),
    ],
    ids=[
        "truncated",
        "missing-list",
        "duplicate-id-field",
        "identity-conflict",
        "wrong-session",
        "unknown-format",
        "nil-with-content",
    ],
)
def test_invalid_full_snapshot_preserves_error_source_and_last_good_state(
    tmp_path: Path, raw: bytes
) -> None:
    database = tmp_path / "b2.sqlite"
    a = read_sync(database, synchronize(database))
    failed = read_sync(database, synchronize(database, raw, tick=1))
    assert failed.status == "feilet"
    assert failed.errors
    assert failed.entries == []
    assert read_source(database, failed.source_sha256) == raw
    revision = read_revision(database, a.entries[0].revision_id)
    assert revision.current_revision_id == a.entries[0].revision_id
    assert revision.observation_status == "siste forsøk ikke fullført"


@pytest.mark.parametrize("timestamp", ["0001-01-01T00:00:00+14:00", "9999-12-31T23:59:59-14:00"])
@pytest.mark.parametrize("location", ["source", "metadata"])
@pytest.mark.parametrize("has_previous", [False, True])
def test_unrepresentable_utc_time_is_a_saved_failure_and_releases_pending_job(
    tmp_path: Path, timestamp: str, location: str, has_previous: bool
) -> None:
    from innlesning import resume_sync, stage_sync

    database = tmp_path / "b2.sqlite"
    previous = read_sync(database, synchronize(database)) if has_previous else None
    root = etree.fromstring(RAW)
    if location == "source":
        root.find("{http://data.stortinget.no}respons_dato_tid").text = timestamp
    raw = etree.tostring(root)
    source_metadata = metadata(raw, tick=1)
    if location == "metadata":
        source_metadata["retrieved_utc"] = timestamp
    job_id = stage_sync(database, raw, source_metadata, historical_sample=True)
    assert read_sync(database, job_id).status == "venter"
    assert resume_sync(database, job_id) == job_id
    failed = read_sync(database, job_id)
    assert failed.status == "feilet"
    assert failed.completed_utc is not None
    assert failed.errors and "UTC" in failed.errors[0]
    assert failed.entries == []
    assert failed.source_metadata == source_metadata
    assert read_source(database, failed.source_sha256) == raw
    assert resume_sync(database, job_id) == job_id
    assert read_sync(database, job_id) == failed
    if previous is not None:
        revision_id = previous.entries[0].revision_id
        revision = read_revision(database, revision_id)
        assert revision.current_revision_id == revision_id
        assert revision.observation_status == "siste forsøk ikke fullført"
    valid = read_sync(database, synchronize(database, tick=2))
    assert valid.status == "fullfort"
    assert valid.counts["uendret" if has_previous else "ny"] == 1


def test_unknown_database_is_not_silently_adopted(tmp_path: Path) -> None:
    database = tmp_path / "unknown.sqlite"
    with sqlite3.connect(database) as connection:
        connection.execute("CREATE TABLE unrelated (id TEXT)")
    with pytest.raises(ValueError, match="skjema"):
        synchronize(database)


def test_legacy_database_read_write_and_export_survive_additive_upgrade(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    from analyse import export_excel
    from innlesning import import_saved, read_report

    database = tmp_path / "legacy.sqlite"
    source_metadata = json.loads((FIXTURES / "40033944.metadata.json").read_text(encoding="utf-8"))
    legacy = import_saved(database, RAW, source_metadata, "40033944", historical_sample=True)
    before = read_report(database, legacy)
    export_excel(database, legacy, tmp_path / "before.xlsx")
    job = read_sync(database, synchronize(database))
    assert job.counts["ny"] == 1  # B1-prøvearkivet adopteres ikke automatisk.
    assert read_report(database, legacy) == before
    assert read_source(database, before.source_sha256) == RAW
    second_legacy = import_saved(database, RAW, source_metadata, "40033944", historical_sample=True)
    assert second_legacy != legacy
    assert read_report(database, second_legacy).lines == before.lines
    export_excel(database, legacy, tmp_path / "after.xlsx")
    b2 = read_revision(database, job.entries[0].revision_id)
    output = tmp_path / "via-existing-kjoring.xlsx"
    export_excel(database, b2.report.run_id, output)
    workbook = load_workbook(output)
    notes = [cell.value for row in workbook["Kilde og kontroll"] for cell in row]
    assert b2.revision_id in notes
    workbook.close()


def test_old_saved_snapshot_and_old_completed_resume_cannot_replace_newer_state(
    tmp_path: Path,
) -> None:
    from innlesning import resume_sync

    database = tmp_path / "b2.sqlite"
    a = read_sync(database, synchronize(database))
    raw = RAW.replace(b"13:34:46.1397306", b"13:35:46.1397306").replace(b"10 478 000", b"I 615 833")
    b = read_sync(database, synchronize(database, raw, tick=1))
    stale = read_sync(database, synchronize(database, tick=2))
    assert stale.status == "feilet"
    assert "Eldre" in stale.errors[0]
    resume_sync(database, a.job_id)
    current = read_revision(database, a.entries[0].revision_id)
    assert current.current_revision_id == b.entries[0].revision_id


def test_excerpt_scope_and_samples_do_not_pollute_production_inventory(tmp_path: Path) -> None:
    database = tmp_path / "b2.sqlite"
    raw = RAW.replace(b"2024-2025", b"2026-2027")
    source_metadata = metadata(raw)
    source_metadata.update(
        url="https://data.stortinget.no/eksport/stortingsvedtak?sesjonid=2026-2027",
        fixture_kind="http_original",
    )
    production = read_sync(database, sync_saved(database, raw, source_metadata))
    assert production.dataset == "produksjon"
    changed = raw.replace(b"10 478 000", b"I 615 833")
    sample_metadata = metadata(changed, tick=1)
    sample_metadata.update(url=source_metadata["url"], fixture_kind="synthetic_test")
    sample = read_sync(database, sync_saved(database, changed, sample_metadata))
    assert sample.dataset == "prove"
    assert (
        read_revision(database, production.entries[0].revision_id).freshness
        == "gjeldende observert innhold"
    )
    source_metadata.update(
        sha256=hashlib.sha256(changed).hexdigest(),
        bytes=len(changed),
        fixture_kind="derived_excerpt",
        session_coverage="partial",
    )
    with pytest.raises(ValueError, match="utdrag"):
        sync_saved(database, changed, source_metadata)
    selected = read_sync(
        database, sync_saved(database, changed, source_metadata, decision_id="40033944")
    )
    assert selected.scope == "valgt"
    assert selected.counts["mangler"] == 0


def test_future_database_schema_is_rejected_for_read_and_write(tmp_path: Path) -> None:
    from innlesning import read_report

    database = tmp_path / "b2.sqlite"
    job = read_sync(database, synchronize(database))
    revision = read_revision(database, job.entries[0].revision_id)
    with sqlite3.connect(database) as connection:
        connection.execute("PRAGMA user_version = 999")
    with pytest.raises(ValueError, match="999"):
        synchronize(database)
    with pytest.raises(ValueError, match="999"):
        read_report(database, revision.report.run_id)


def test_control_survives_time_change_but_uncertain_latest_attempt_is_visible(
    tmp_path: Path,
) -> None:
    from innlesning import record_control_reference

    database = tmp_path / "b2.sqlite"
    first = read_sync(database, synchronize(database))
    revision_id = first.entries[0].revision_id
    record_control_reference(database, revision_id, "Kontroll", "Kilde kontrollert")
    synchronize(database, RAW.replace(b"13:34:46.1397306", b"13:35:46.1397306"), tick=1)
    assert read_revision(database, revision_id).controls[0].status == "gjeldende henvisning"
    synchronize(database, b"<broken", tick=2)
    revision = read_revision(database, revision_id)
    assert revision.controls[0].status == "ikke vurderbar"
    assert revision.latest_attempt_status == "feilet"


def test_identical_amounts_under_a_different_id_are_two_decisions(tmp_path: Path) -> None:
    database = tmp_path / "b2.sqlite"
    first = read_sync(database, synchronize(database))
    other = RAW.replace(b"40033944", b"90000001").replace(b"2024-12-13", b"2001-01-01")
    root = etree.fromstring(RAW)
    listing = root.find("{http://data.stortinget.no}stortingsvedtak_liste")
    listing.append(
        etree.fromstring(other).find("{http://data.stortinget.no}stortingsvedtak_liste")[0]
    )
    job = read_sync(database, synchronize(database, etree.tostring(root), tick=1))
    assert job.counts == {"ny": 1, "uendret": 1, "rettelse": 0, "mangler": 0}
    new = read_revision(database, job.entries[1].revision_id)
    assert new.report.decision.id == "90000001"
    assert new.report.decision.date == "2001-01-01"
    assert len(new.report.lines) == 8
    assert new.report.lines[0].amount_kr == 10_478_000
    assert new.report.run_id != read_revision(database, first.entries[0].revision_id).report.run_id


def test_formatting_and_prefix_changes_preserve_revision_but_unknown_nil_is_content(
    tmp_path: Path,
) -> None:
    database = tmp_path / "b2.sqlite"
    first = read_sync(database, synchronize(database))
    root = etree.fromstring(RAW)
    etree.indent(root, space="        ")
    raw = etree.tostring(root).replace(b"xmlns:i=", b"xmlns:newprefix=")
    job = read_sync(database, synchronize(database, raw, tick=1))
    assert job.counts["uendret"] == 1
    assert job.entries[0].revision_id == first.entries[0].revision_id
    unknown = RAW.replace(b"</stortingsvedtak>", b"<ukjent i:nil='true'/></stortingsvedtak>")
    a = read_sync(database, synchronize(database, unknown, tick=2))
    b = read_sync(
        database, synchronize(database, unknown.replace(b"i:nil='true'", b"i:nil='false'"), tick=3)
    )
    assert a.counts["rettelse"] == b.counts["rettelse"] == 1
    assert a.entries[0].revision_id != b.entries[0].revision_id


@pytest.mark.parametrize(
    ("root_space", "list_space", "object_space", "field_space", "classification"),
    [
        ("preserve", None, None, None, "rettelse"),
        (None, "preserve", None, None, "rettelse"),
        ("preserve", "default", None, None, "uendret"),
        ("default", "preserve", None, None, "rettelse"),
        (None, "preserve", "default", None, "uendret"),
        (None, "preserve", None, "default", "uendret"),
    ],
)
def test_inherited_xml_space_and_local_default_control_separator_revisions(
    tmp_path: Path,
    root_space: str | None,
    list_space: str | None,
    object_space: str | None,
    field_space: str | None,
    classification: str,
) -> None:
    from innlesning import record_control_reference

    database = tmp_path / "b2.sqlite"
    root = etree.fromstring(RAW)
    ns = "{http://data.stortinget.no}"
    listing = root.find(ns + "stortingsvedtak_liste")
    node = listing[0]
    unknown = etree.SubElement(node, ns + "ukjent")
    etree.SubElement(unknown, ns + "barn").tail = " "
    etree.SubElement(unknown, ns + "barn")
    for element, space in (
        (root, root_space),
        (listing, list_space),
        (node, object_space),
        (unknown, field_space),
    ):
        if space is not None:
            element.set("{http://www.w3.org/XML/1998/namespace}space", space)
    first = read_sync(database, synchronize(database, etree.tostring(root)))
    revision_id = first.entries[0].revision_id
    record_control_reference(database, revision_id, "Kontroll", "Før endret separator")
    unknown[0].tail = "  "
    second = read_sync(database, synchronize(database, etree.tostring(root), tick=1))
    assert second.status == "fullfort"
    assert second.counts[classification] == 1
    assert (second.entries[0].revision_id == revision_id) == (classification == "uendret")
    assert read_revision(database, revision_id).controls[0].status == (
        "gjeldende henvisning" if classification == "uendret" else "utdatert"
    )


@pytest.mark.parametrize("ignored", [b"<!--note-->", b"<?note ignored?>"])
@pytest.mark.parametrize(
    ("before", "after", "classification"),
    [
        (b"prefix{ignored}A", b"prefix{ignored}B", "rettelse"),
        (b"<barn/>A{ignored}B", b"<barn/>A{ignored}C", "rettelse"),
        (b" <barn/>{ignored}A", b"  <barn/>{ignored}A", "rettelse"),
        (b"prefixAB", b"prefixA{ignored}B", "uendret"),
        (b"<barn/>AB", b"<barn/>A{ignored}B", "uendret"),
        (b" <barn/> ", b" <barn/>{ignored} ", "uendret"),
    ],
)
def test_ignored_xml_nodes_preserve_character_data_and_mixed_text(
    tmp_path: Path, ignored: bytes, before: bytes, after: bytes, classification: str
) -> None:
    from innlesning import record_control_reference

    database = tmp_path / "b2.sqlite"

    def raw(content: bytes) -> bytes:
        return RAW.replace(
            b"</stortingsvedtak>",
            b"<ukjent>" + content.replace(b"{ignored}", ignored) + b"</ukjent></stortingsvedtak>",
        )

    first = read_sync(database, synchronize(database, raw(before)))
    revision_id = first.entries[0].revision_id
    record_control_reference(database, revision_id, "Kontroll", "Før endret tegninnhold")
    second = read_sync(database, synchronize(database, raw(after), tick=1))
    assert second.status == "fullfort"
    assert second.counts[classification] == 1
    assert (second.entries[0].revision_id == revision_id) == (classification == "uendret")
    assert read_revision(database, revision_id).controls[0].status == (
        "gjeldende henvisning" if classification == "uendret" else "utdatert"
    )


def test_duplicates_and_corrupted_empty_list_are_failures_not_absence(tmp_path: Path) -> None:
    database = tmp_path / "b2.sqlite"
    synchronize(database)
    root = etree.fromstring(RAW)
    listing = root.find("{http://data.stortinget.no}stortingsvedtak_liste")
    listing.append(etree.fromstring(etree.tostring(listing[0])))
    duplicate = read_sync(database, synchronize(database, etree.tostring(root), tick=1))
    assert duplicate.status == "feilet"
    assert duplicate.counts["mangler"] == 0
    corrupted = variant(empty=True).replace(
        b"<stortingsvedtak_liste>", b"<stortingsvedtak_liste>truncated text"
    )
    invalid_empty = read_sync(database, synchronize(database, corrupted, tick=2))
    assert invalid_empty.status == "feilet"
    assert invalid_empty.counts["mangler"] == 0


def test_selected_excerpt_cannot_mark_other_known_id_missing(tmp_path: Path) -> None:
    database = tmp_path / "b2.sqlite"
    full = read_sync(database, synchronize(database, variant(extra=True)))
    excerpt_metadata = metadata(RAW, tick=1)
    excerpt_metadata.update(fixture_kind="derived_excerpt", session_coverage="partial")
    selected = read_sync(
        database,
        sync_saved(database, RAW, excerpt_metadata, decision_id="40033944", historical_sample=True),
    )
    assert selected.scope == "valgt"
    assert selected.counts == {"ny": 0, "uendret": 1, "rettelse": 0, "mangler": 0}
    assert (
        read_revision(database, full.entries[1].revision_id).current_revision_id
        == full.entries[1].revision_id
    )


@pytest.mark.parametrize(
    "bad",
    [
        {"sha256": "0" * 64},
        {"bytes": 1},
        {"retrieved_utc": "2026-10-20T00:00:00"},
        {"status": 503},
        {"url": "https://example.com/private"},
    ],
)
def test_invalid_provenance_is_rejected_before_staging(
    tmp_path: Path, bad: dict[str, object]
) -> None:
    database = tmp_path / "b2.sqlite"
    source_metadata = metadata(RAW)
    source_metadata.update(bad)
    with pytest.raises(ValueError):
        sync_saved(database, RAW, source_metadata, historical_sample=True)
    assert not database.exists()


def test_metadata_correction_and_session_scope_keep_their_own_identity(tmp_path: Path) -> None:
    database = tmp_path / "b2.sqlite"
    first = read_sync(database, synchronize(database))
    raw = RAW.replace(b"100215", b"100216").replace(b"Budsjettvedtak", b"Oppdatert typenavn")
    second = read_sync(database, synchronize(database, raw, tick=1))
    assert second.counts["rettelse"] == 1
    assert (
        read_revision(database, second.entries[0].revision_id).report.decision.case_id == "100216"
    )
    assert read_revision(database, first.entries[0].revision_id).report.decision.case_id == "100215"
    another_session = RAW.replace(b"2024-2025", b"2026-2027")
    other_metadata = metadata(another_session)
    other_metadata["url"] = "https://data.stortinget.no/eksport/stortingsvedtak?sesjonid=2026-2027"
    other = read_sync(database, sync_saved(database, another_session, other_metadata))
    assert other.counts["ny"] == 1
    assert other.entries[0].revision_id != first.entries[0].revision_id
    assert (
        read_revision(database, first.entries[0].revision_id).current_revision_id
        == second.entries[0].revision_id
    )
