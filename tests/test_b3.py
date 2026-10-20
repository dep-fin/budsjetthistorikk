"""Kilde → B2 → B3 → gjenåpning → Excel. Mutasjoner merkes syntetiske."""

import hashlib
import json
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from analyse.history import HistorySelection, calculate_history, read_history
from analyse.history_excel import export_history_excel
from innlesning import read_sync, sync_saved
from innlesning.observations import (
    AmountObservation,
    Artifact,
    DocumentReference,
    Extraction,
    Interpretation,
    Organization,
    SourcePointer,
    extract_revision,
    identity,
    read_extraction,
    save_extraction,
)

FIXTURES = Path(__file__).parent / "fixtures" / "b3"


def ingest(database: Path, identifier: str) -> str:
    raw = (FIXTURES / f"{identifier}.xml").read_bytes()
    metadata = json.loads((FIXTURES / f"{identifier}.json").read_text(encoding="utf-8"))
    job = sync_saved(database, raw, metadata, historical_sample=True, decision_id=identifier)
    revision = read_sync(database, job).entries[0].revision_id
    return extract_revision(database, revision).id


def test_authentic_reduction_and_increases_survive_storage_and_excel(tmp_path: Path) -> None:
    database = tmp_path / "b3.sqlite"
    ids = tuple(ingest(database, id) for id in ("40032850", "40033719", "40033805", "40034064"))
    original = read_extraction(database, ids[0])
    selection = HistorySelection("prove", ids, 2024, "2541", "70")
    result = calculate_history(database, selection)
    assert result.approved_balance == 12_890_400_000
    assert [step.delta for step in result.steps] == [None, -608_000_000, 400_000, 680_000_000]
    assert [step.balance for step in result.steps] == [
        12_818_000_000,
        12_210_000_000,
        12_210_400_000,
        12_890_400_000,
    ]
    assert len(result.input_observation_ids) == 10
    assert read_history(database, result.id).approved_balance == 12_890_400_000
    assert calculate_history(database, selection).id == result.id
    assert extract_revision(database, original.source_revision_id).id == original.id
    assert read_extraction(database, original.id) == original
    output = tmp_path / "historikk.xlsx"
    export_history_excel(database, result.id, output)
    workbook = load_workbook(output)
    assert {"Historikk", "Beløpsobservasjoner", "Grunnlag og kontroll"} <= set(workbook.sheetnames)
    values = [cell.value for page in workbook for row in page for cell in row]
    assert 12_890_400_000 in values
    assert "table[1]/tr[92]" in values
    assert "table[1]/tr[1071]" in values
    assert "rnb" in values
    assert original.artifact.sha256 in values
    assert not any(cell.data_type == "f" for page in workbook for row in page for cell in row)
    workbook.close()


@pytest.mark.parametrize(
    "identifier,chapter,post,delta",
    [
        ("40033805", "2541", "70", 400_000),
        ("40034443", "900", "32", 454_000_000),
        ("40034443", "1760", "45", -227_000_000),
    ],
)
def test_authentic_change_and_new_post_do_not_invent_zero_start(
    tmp_path: Path,
    identifier: str,
    chapter: str,
    post: str,
    delta: int,
) -> None:
    database = tmp_path / "missing.sqlite"
    id = ingest(database, identifier)
    result = calculate_history(
        database,
        HistorySelection("prove", (id,), 2024 if identifier == "40033805" else 2025, chapter, post),
    )
    assert result.approved_balance is None
    assert result.steps[0].delta == delta
    assert result.steps[0].balance is None
    assert "manglende_start" in {issue.code for issue in result.anomalies}
    export_history_excel(database, result.id, tmp_path / "missing.xlsx")
    workbook = load_workbook(tmp_path / "missing.xlsx")
    assert workbook["Historikk"]["J2"].value is None
    assert workbook["Historikk"]["G2"].value == delta
    workbook.close()


@pytest.mark.parametrize(
    "chapter,post,final", [("1860", "1", 360_312_000), ("4860", "1", 83_400_000)]
)
def test_authentic_repeated_transitions_keep_both_ids_and_block_chain(
    tmp_path: Path,
    chapter: str,
    post: str,
    final: int,
) -> None:
    database = tmp_path / "havtil.sqlite"
    ids = tuple(ingest(database, id) for id in ("40032850", "40033719", "40034064", "40034226"))
    result = calculate_history(database, HistorySelection("prove", ids, 2024, chapter, post))
    assert result.approved_balance is None
    assert "kjedekonflikt" in {issue.code for issue in result.anomalies}
    assert len({step.source_revision_id for step in result.steps[-2:]}) == 2
    assert [step.after for step in result.steps[-2:]] == [final, final]
    assert "lokal_konflikt" not in {issue.code for issue in result.anomalies}
    export_history_excel(database, result.id, tmp_path / "havtil.xlsx")


def test_authentic_financing_conflict_keeps_zero_and_final_without_synthetic_delta(
    tmp_path: Path,
) -> None:
    database = tmp_path / "finansiering.sqlite"
    ids = tuple(ingest(database, id) for id in ("40033025", "40033719", "40034287"))
    result = calculate_history(database, HistorySelection("prove", ids, 2024, "5999", "90"))
    assert result.approved_balance is None
    assert [step.delta for step in result.steps] == [None, -101_874_353_000, 18_497_356_000]
    assert [step.after for step in result.steps] == [None, 0, 120_371_709_000]
    assert "kjedekonflikt" in {issue.code for issue in result.anomalies}


def test_authentic_raw_I_is_not_repaired_or_zero_filled(tmp_path: Path) -> None:
    database = tmp_path / "599.sqlite"
    id = ingest(database, "40036289")
    extraction = read_extraction(database, id)
    after = next(
        o for o in extraction.observations if o.post == "70" and o.meaning == "etter_belop"
    )
    assert after.raw_value == "I 615 833"
    assert after.amount_kr is None
    assert after.structure_status == "til_kontroll"
    result = calculate_history(database, HistorySelection("prove", (id,), 2026, "5538", "70"))
    assert result.approved_balance is None
    export_history_excel(database, result.id, tmp_path / "599.xlsx")
    workbook = load_workbook(tmp_path / "599.xlsx")
    assert "I 615 833" in [cell.value for row in workbook["Beløpsobservasjoner"] for cell in row]
    workbook.close()


def mutate(database: Path, identifier: str, old: bytes, new: bytes, tick: int = 20) -> str:
    raw = (FIXTURES / f"{identifier}.xml").read_bytes().replace(old, new)
    metadata = json.loads((FIXTURES / f"{identifier}.json").read_text(encoding="utf-8"))
    metadata.update(
        fixture_kind="synthetic_test",
        sha256=hashlib.sha256(raw).hexdigest(),
        bytes=len(raw),
        retrieved_utc=f"2026-10-20T13:{tick:02d}:00+00:00",
    )
    job = sync_saved(database, raw, metadata, historical_sample=True, decision_id=identifier)
    revision = read_sync(database, job).entries[0].revision_id
    return extract_revision(database, revision).id


def test_synthetic_local_conflict_is_preserved_through_storage_and_excel(tmp_path: Path) -> None:
    database = tmp_path / "local.sqlite"
    original = ingest(database, "40032850")
    rnb = ingest(database, "40033719")
    changed = mutate(database, "40033805", b"12 210 400 000", b"12 210 400 001")
    result = calculate_history(
        database, HistorySelection("prove", (original, rnb, changed), 2024, "2541", "70")
    )
    assert result.approved_balance is None
    assert "lokal_konflikt" in {issue.code for issue in result.anomalies}
    assert result.steps[-1].after == 12_210_400_001
    export_history_excel(database, result.id, tmp_path / "local.xlsx")


def generic(
    database: Path,
    *,
    date: str | None = None,
    meaning: str = "nivaa",
    round: str = "opprinnelig",
    value: int | None = 10_032_000,
    year: int = 2013,
    figure: str = "stortingsvedtak",
    scope: str = "enkeltledd",
    info: str = "bevilgning",
    marker: str = "H2",
    structure: str = "strukturert",
    subpost: str | None = "1",
    detail: str = "underpost",
) -> Extraction:
    """Lokal syntetisk kontraktprøve, ingen historisk Excel-import eller enhetsgodkjenning."""
    raw = b"synthetic generic artifact; not an actual workbook"
    digest = hashlib.sha256(raw).hexdigest()
    interpretation = Interpretation(
        adapter_version="synthetic-excel-contract-v1", reason="Representasjonsprøve"
    )
    artifact = Artifact(
        digest,
        "synthetic.xlsx",
        "workbook",
        "application/octet-stream",
        len(raw),
        '{"fixture_kind":"synthetic_test"}',
    )
    observation = AmountObservation(
        dataset="historisk",
        source_system="excel",
        source_unit_id=f"excel:{digest}:{marker}",
        source_revision_id=identity(digest, marker),
        interpretation_id=interpretation.id,
        pointer=SourcePointer(
            digest,
            f"Rådata!{marker}",
            marker,
            "Rådata",
            2,
            "Saldert (f120)",
            marker,
            '=HYPERLINK("https://example.invalid")',
        ),
        step_id=identity("synthetic-step", marker),
        raw_value="10.032",
        amount_text="10.032",
        amount_kr=value,
        chapter_raw="0001",
        post_raw="01",
        subpost_raw=subpost,
        chapter="1",
        post="1",
        subpost=subpost,
        detail_level=detail,
        budget_year=year,
        year_evidence="Syntetisk inventar; ingen generell historisk tolkning",
        decision_date=date,
        unit_raw="mill. kroner",
        currency="NOK",
        multiplier="1000000",
        conversion_rule="Syntetisk desimalverdi 10.032 * 1000000",
        raw_type="n",
        precision="Bevart råverdi",
        normalization_evidence="Kun kontraktprøve",
        stage_raw="Saldert (f120)",
        figure_type=figure,
        round=round,
        meaning=meaning,
        scope=scope,
        information_type=info,
        structure_status=structure,
        organization=Organization(
            dep="15",
            fdep_nr="15",
            gdep_nr="92",
            gdep_name="Regjeringen",
            omr_nr="1",
            kat_nr="10",
            register_source="syntetisk register",
            register_version="test-1",
            valid_from="2013",
            valid_to="2013",
            validity_precision="ar",
        ),
        documents=(
            DocumentReference("p141/16-17", "E", "forslagsgrunnlag"),
            DocumentReference("i445/16-17", "F", "innstilling"),
            DocumentReference("v963/16-17", "G", "vedtak"),
        ),
        null_reasons=(("decision_date", "ikke oppgitt"), ("api_id", "ikke relevant")),
    )
    extraction = Extraction(
        "historisk",
        "excel",
        observation.source_unit_id,
        observation.source_revision_id,
        interpretation,
        artifact,
        (observation,),
    )
    save_extraction(database, extraction, raw)
    return extraction


def test_generic_representation_preserves_excel_metadata_without_api_id_or_date(
    tmp_path: Path,
) -> None:
    database = tmp_path / "generic.sqlite"
    extraction = generic(database)
    reopened = read_extraction(database, extraction.id)
    assert reopened == extraction
    observation = reopened.observations[0]
    assert observation.api_id is None and observation.decision_date is None
    assert dict(observation.null_reasons)["decision_date"] == "ikke oppgitt"
    assert observation.chapter_raw == "0001" and observation.chapter == "1"
    assert observation.organization.gdep_nr == "92"
    selection = HistorySelection("historisk", (extraction.id,), 2013, "1", "1", "1", "underpost")
    missing_order = calculate_history(database, selection)
    assert missing_order.approved_balance is None
    assert "uavklart_rekkefolge" in {issue.code for issue in missing_order.anomalies}
    selected = replace(
        selection,
        step_order=(observation.step_id,),
        order_evidence="Syntetisk kontrollgrunnlag for ett ledd",
    )
    result = calculate_history(database, selected)
    assert result.approved_balance == 10_032_000
    output = tmp_path / "generic.xlsx"
    export_history_excel(database, result.id, output)
    workbook = load_workbook(output)
    row = list(workbook["Beløpsobservasjoner"].values)[1]
    assert row[20] == "10.032" and row[22] == 10_032_000
    assert "p141/16-17" in row[28] and "i445/16-17" in row[28] and "v963/16-17" in row[28]
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()


@pytest.mark.parametrize(
    "figure,scope,info",
    [
        ("regjeringsforslag", "enkeltledd", "bevilgning"),
        ("komiteforslag", "enkeltledd", "bevilgning"),
        ("regnskap", "enkeltledd", "bevilgning"),
        ("stortingsvedtak", "aggregat", "bevilgning"),
        ("stortingsvedtak", "enkeltledd", "fullmaktsramme"),
    ],
)
def test_other_figure_types_aggregates_and_authorities_never_add(
    tmp_path: Path,
    figure: str,
    scope: str,
    info: str,
) -> None:
    database = tmp_path / "excluded.sqlite"
    original = generic(database, date="2013-01-01")
    excluded = generic(
        database,
        marker="I2",
        date="2013-02-01",
        meaning="endring",
        round="rnb",
        value=72_000_000,
        figure=figure,
        scope=scope,
        info=info,
    )
    result = calculate_history(
        database,
        HistorySelection("historisk", (original.id, excluded.id), 2013, "1", "1", "1", "underpost"),
    )
    assert result.approved_balance == 10_032_000
    assert len(result.steps) == 1
    assert len(result.input_observation_ids) == 2  # Ekskludert kildeopplysning er synlig i Excel.


def test_same_date_requires_source_based_order_and_unknown_step_blocks(tmp_path: Path) -> None:
    database = tmp_path / "order.sqlite"
    original = generic(database, date="2013-01-01")
    change = generic(
        database, marker="I2", date="2013-01-01", meaning="endring", round="ovrig", value=-32_000
    )
    selection = HistorySelection(
        "historisk", (change.id, original.id), 2013, "1", "1", "1", "underpost"
    )
    assert calculate_history(database, selection).approved_balance is None
    selected = replace(
        selection,
        step_order=(original.observations[0].step_id, change.observations[0].step_id),
        order_evidence="Syntetisk kilde sier original før endring",
    )
    assert calculate_history(database, selected).approved_balance == 10_000_000
    unknown = generic(
        database, marker="J2", date="2013-02-01", meaning="uavklart", round="uavklart", value=0
    )
    result = calculate_history(
        database, replace(selection, extraction_ids=(original.id, unknown.id))
    )
    assert result.approved_balance is None


def test_corrections_keep_revisions_separate_and_never_reactivate_old_calculation(
    tmp_path: Path,
) -> None:
    database = tmp_path / "correction.sqlite"
    a = ingest(database, "40033805")
    original, rnb = ingest(database, "40032850"), ingest(database, "40033719")
    selection = HistorySelection("prove", (original, rnb, a), 2024, "2541", "70")
    result_a = calculate_history(database, selection)
    assert result_a.approved_balance == 12_210_400_000
    b = mutate(database, "40033805", b"400 000", b"500 000")
    assert b != a
    assert read_history(database, result_a.id).approved_balance is None
    mixed = calculate_history(database, replace(selection, extraction_ids=(original, rnb, a, b)))
    assert mixed.approved_balance is None
    assert "motstridende_grunnlag" in {issue.code for issue in mixed.anomalies}
    # Tilbakekomst er syntetisk A → B → A med senere provenienstid.
    returned = mutate(database, "40033805", b"400 000", b"400 000", tick=21)
    assert returned == a
    assert read_history(database, result_a.id).freshness == "utdatert kilde-/kontrollgrunnlag"
    assert read_history(database, result_a.id).approved_balance is None
    new_a = calculate_history(database, selection)
    assert new_a.id != result_a.id and new_a.approved_balance == 12_210_400_000


def test_interpretation_versions_and_datasets_years_are_explicit(tmp_path: Path) -> None:
    database = tmp_path / "selection.sqlite"
    id = ingest(database, "40033805")
    old = read_extraction(database, id)
    interpretation = Interpretation(
        round="ovrig", round_evidence="Sakstittel om permitteringsregelverket; kildefunn #2"
    )
    new = extract_revision(database, old.source_revision_id, interpretation=interpretation)
    assert new.id != old.id
    assert new.observations[0].id != old.observations[0].id
    assert new.observations[0].amount_kr == old.observations[0].amount_kr
    assert old.observations[0].round == "uavklart" and new.observations[0].round == "ovrig"
    with pytest.raises(ValueError, match="Datasett"):
        calculate_history(database, HistorySelection("produksjon", (old.id,), 2024, "2541", "70"))
    assert (
        calculate_history(database, HistorySelection("prove", (old.id,), 2025, "2541", "70")).steps
        == ()
    )
    mixed = calculate_history(
        database, HistorySelection("prove", (old.id, new.id), 2024, "2541", "70")
    )
    assert "motstridende_grunnlag" in {issue.code for issue in mixed.anomalies}


def test_real_initiative_and_authority_do_not_create_allocation_steps(tmp_path: Path) -> None:
    database = tmp_path / "initiative.sqlite"
    ids = tuple(ingest(database, id) for id in ("40034385", "40035422", "40035429"))
    result = calculate_history(database, HistorySelection("prove", ids, 2025, "41", "45"))
    assert len(result.steps) == 1
    assert result.steps[0].delta == 9_700_000
    assert result.approved_balance is None
    assert read_extraction(database, ids[0]).observations == ()
    assert read_extraction(database, ids[2]).observations == ()


def test_after_without_before_is_checked_against_calculated_balance(tmp_path: Path) -> None:
    database = tmp_path / "after.sqlite"
    original = generic(database, date="2013-01-01")
    change = generic(
        database, marker="I2", date="2013-02-01", meaning="endring", round="ovrig", value=-32_000
    )
    delta = change.observations[0]
    after = replace(
        delta,
        pointer=replace(delta.pointer, amount_field="etter_belop"),
        raw_value="10000001",
        amount_text="10000001",
        amount_kr=10_000_001,
        meaning="etter_belop",
    )
    modified = replace(
        change,
        observations=(delta, after),
        interpretation=replace(change.interpretation, reason="Syntetisk til-felt uten fra"),
    )
    modified = replace(
        modified,
        observations=tuple(
            replace(o, interpretation_id=modified.interpretation.id) for o in modified.observations
        ),
    )
    save_extraction(database, modified, b"synthetic generic artifact; not an actual workbook")
    result = calculate_history(
        database,
        HistorySelection("historisk", (original.id, modified.id), 2013, "1", "1", "1", "underpost"),
    )
    assert result.approved_balance is None
    assert "etter_konflikt" in {issue.code for issue in result.anomalies}


def test_unknown_control_row_cannot_be_silently_dropped(tmp_path: Path) -> None:
    database = tmp_path / "unknown.sqlite"
    original, rnb = ingest(database, "40032850"), ingest(database, "40033719")
    changed = mutate(database, "40033805", b"til kr 12 210 400 000", b"til kroner 12 210 400 000")
    result = calculate_history(
        database, HistorySelection("prove", (original, rnb, changed), 2024, "2541", "70")
    )
    assert result.approved_balance is None
    assert "ukjent_ledd" in {issue.code for issue in result.anomalies}


def test_b1_contract_mapping_and_additive_upgrade_preserve_pending_b2_job(tmp_path: Path) -> None:
    import sqlite3
    from dataclasses import asdict

    from innlesning import import_saved, read_report, read_revision, resume_sync, stage_sync
    from innlesning.storage import read_source

    database = tmp_path / "compatibility.sqlite"
    raw = (FIXTURES.parent / "40033944.derived.xml").read_bytes()
    metadata = json.loads((FIXTURES.parent / "40033944.metadata.json").read_text(encoding="utf-8"))
    b1 = import_saved(database, raw, metadata, "40033944", historical_sample=True)
    b1_before = asdict(read_report(database, b1))
    job = sync_saved(database, raw, metadata, historical_sample=True, decision_id="40033944")
    revision = read_sync(database, job).entries[0].revision_id
    b2_before = read_revision(database, revision)
    pending = stage_sync(database, raw, metadata, historical_sample=True, decision_id="40033944")
    with sqlite3.connect(database) as connection:
        assert connection.execute("PRAGMA user_version").fetchone()[0] == 2
    extraction = extract_revision(database, revision)
    target = next(o for o in extraction.observations if o.chapter == "42" and o.post == "1")
    assert (
        target.amount_kr,
        target.budget_year,
        target.figure_type,
        target.round,
        target.meaning,
        target.scope,
    ) == (10_478_000, 2025, "stortingsvedtak", "opprinnelig", "nivaa", "enkeltledd")
    assert target.pointer.locator == b2_before.report.lines[0].locator
    assert asdict(read_report(database, b1)) == b1_before
    assert read_source(database, extraction.artifact.sha256) == raw
    assert read_revision(database, revision).fingerprint == b2_before.fingerprint
    assert read_revision(database, revision).parser_version == "b1-original-v1"
    assert read_sync(database, pending).status == "venter"
    resume_sync(database, pending)
    assert read_sync(database, pending).entries[0].revision_id == revision
    with sqlite3.connect(database) as connection:
        assert connection.execute("PRAGMA user_version").fetchone()[0] == 3
    import_saved(database, raw, metadata, "40033944", historical_sample=True)
    assert read_extraction(database, extraction.id) == extraction


def test_failed_latest_attempt_suspends_live_balance_without_mutating_saved_result(
    tmp_path: Path,
) -> None:
    from innlesning import resume_sync, stage_sync

    database = tmp_path / "pending.sqlite"
    original = ingest(database, "40032850")
    result = calculate_history(database, HistorySelection("prove", (original,), 2024, "2541", "70"))
    assert result.approved_balance == 12_818_000_000
    raw = (FIXTURES / "40032850.xml").read_bytes()[:-20]
    metadata = json.loads((FIXTURES / "40032850.json").read_text(encoding="utf-8"))
    metadata.update(
        fixture_kind="synthetic_test",
        sha256=hashlib.sha256(raw).hexdigest(),
        bytes=len(raw),
        retrieved_utc="2026-10-20T14:00:00+00:00",
    )
    pending = stage_sync(database, raw, metadata, historical_sample=True, decision_id="40032850")
    assert read_history(database, result.id).approved_balance is None
    resume_sync(database, pending)
    reopened = read_history(database, result.id)
    assert reopened.approved_balance is None
    assert reopened.calculated_balance == 12_818_000_000
    assert reopened.freshness == "ikke vurderbart kildegrunnlag"


def test_old_database_read_does_not_migrate_and_generic_write_does_not_adopt_b1(
    tmp_path: Path,
) -> None:
    import sqlite3

    from innlesning import import_saved, read_report

    database = tmp_path / "legacy.sqlite"
    raw = (FIXTURES.parent / "40033944.derived.xml").read_bytes()
    metadata = json.loads((FIXTURES.parent / "40033944.metadata.json").read_text(encoding="utf-8"))
    b1 = import_saved(database, raw, metadata, "40033944", historical_sample=True)
    with pytest.raises(ValueError, match="Ukjent beløpsuttrekk"):
        read_extraction(database, "unknown")
    assert read_report(database, b1).lines[0].amount_kr == 10_478_000
    with sqlite3.connect(database) as connection:
        assert connection.execute("PRAGMA user_version").fetchone()[0] == 0
    generic(database)
    with sqlite3.connect(database) as connection:
        assert connection.execute("PRAGMA user_version").fetchone()[0] == 3
        assert (
            connection.execute(
                "SELECT name FROM sqlite_master WHERE name='decision_heads'"
            ).fetchone()
            is None
        )


def test_database_path_guard_and_atomic_excel_failure_preserve_existing_output(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from openpyxl import Workbook

    database = tmp_path / "atomic.sqlite"
    extraction = generic(database, date="2013-01-01")
    result = calculate_history(
        database, HistorySelection("historisk", (extraction.id,), 2013, "1", "1", "1", "underpost")
    )
    with pytest.raises(ValueError, match="annen bane"):
        export_history_excel(database, result.id, database)
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"previous deliverable")

    def fail(*args: object, **kwargs: object) -> None:
        raise OSError("Simulert skrivefeil")

    monkeypatch.setattr(Workbook, "save", fail)
    with pytest.raises(OSError, match="Simulert"):
        export_history_excel(database, result.id, output)
    assert output.read_bytes() == b"previous deliverable"
    assert list(tmp_path.glob("*.xlsx")) == [output]


def test_unknown_information_type_and_out_of_range_balance_block_approval(tmp_path: Path) -> None:
    database = tmp_path / "unknown-type.sqlite"
    original = generic(database, date="2013-01-01", value=9_000_000_000_000_000)
    unknown = generic(
        database,
        marker="I2",
        date="2013-02-01",
        meaning="endring",
        round="ovrig",
        value=0,
        info="uavklart",
    )
    selection = HistorySelection(
        "historisk", (original.id, unknown.id), 2013, "1", "1", "1", "underpost"
    )
    assert calculate_history(database, selection).approved_balance is None
    change = generic(
        database,
        marker="J2",
        date="2013-03-01",
        meaning="endring",
        round="ovrig",
        value=9_000_000_000_000_000,
    )
    result = calculate_history(
        database, replace(selection, extraction_ids=(original.id, change.id))
    )
    assert result.approved_balance is None
    assert "saldo_utenfor_presisjon" in {issue.code for issue in result.anomalies}


@pytest.mark.parametrize("uncertainty", ["chapter", "post", "detail_level", "subpost"])
def test_unclassified_selected_source_is_not_dropped_by_account_filter(
    tmp_path: Path,
    uncertainty: str,
) -> None:
    database = tmp_path / "unknown-account.sqlite"
    original = generic(database, date="2013-01-01")
    uncertain = generic(
        database, marker="I2", date="2013-02-01", meaning="endring", round="ovrig", value=1
    )
    modified = replace(
        uncertain,
        interpretation=replace(uncertain.interpretation, reason="Syntetisk uavklart konto"),
    )
    observation = replace(
        uncertain.observations[0],
        interpretation_id=modified.interpretation.id,
        **{uncertainty: "uavklart" if uncertainty == "detail_level" else None},
    )
    modified = replace(modified, observations=(observation,))
    save_extraction(database, modified, b"synthetic generic artifact; not an actual workbook")
    result = calculate_history(
        database,
        HistorySelection("historisk", (original.id, modified.id), 2013, "1", "1", "1", "underpost"),
    )
    assert result.approved_balance is None
    assert observation.id in result.input_observation_ids
    assert "uavklart_kontoutvalg" in {issue.code for issue in result.anomalies}
