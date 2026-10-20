"""B4 ved godkjent grense kilde → B2 → tolkning → lagring → Excel."""

import json
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_b3 import synthetic_html

from analyse import HistorySelection, calculate_history, export_extraction_excel
from innlesning import extract_revision, read_extraction, read_sync, sync_saved
from innlesning.observations import Extraction, Interpretation, save_extraction

FIXTURES = Path(__file__).parent / "fixtures" / "b3"


def ingest(database: Path, identifier: str) -> Extraction:
    directory = FIXTURES if (FIXTURES / f"{identifier}.xml").exists() else FIXTURES.parent / "b4"
    raw = (directory / f"{identifier}.xml").read_bytes()
    metadata = json.loads((directory / f"{identifier}.json").read_text(encoding="utf-8"))
    job = sync_saved(database, raw, metadata, historical_sample=True, decision_id=identifier)
    revision = read_sync(database, job).entries[0].revision_id
    return extract_revision(
        database, revision, interpretation=Interpretation(adapter_version="storting-belop-v2")
    )


def test_authentic_negative_net_preserves_total_components_and_cells(tmp_path: Path) -> None:
    database = tmp_path / "net.sqlite"
    extraction = ingest(database, "40032850")
    rows = [o for o in extraction.observations if (o.chapter, o.post) == ("2470", "24")]
    total = next(o for o in rows if o.detail_level == "post")
    components = [o for o in rows if o.detail_level == "underpost"]
    assert total.amount_kr == -15_000_000
    assert total.pointer.locator == "table[1]/tr[87]/td[5]"
    assert [(o.subpost, o.amount_kr) for o in components] == [
        ("1", -730_000_000),
        ("2", 594_000_000),
        ("3", 90_000_000),
        ("4", 1_000_000),
        ("5", 55_000_000),
        ("6", -25_000_000),
    ]
    assert all(o.component_of == total.id and o.budget_side == "utgift" for o in components)
    assert all(o.pointer.locator.endswith("/td[4]") for o in components)
    assert components[-1].pointer.locator != total.pointer.locator
    assert read_extraction(database, extraction.id) == extraction
    result = calculate_history(
        database, HistorySelection("prove", (extraction.id,), 2024, "2470", "24")
    )
    assert result.approved_balance == -15_000_000
    assert result.input_observation_ids == (total.id,)
    component = calculate_history(
        database, HistorySelection("prove", (extraction.id,), 2024, "2470", "24", "1", "underpost")
    )
    assert component.approved_balance == -730_000_000
    output = tmp_path / "net.xlsx"
    export_extraction_excel(database, extraction.id, output)
    workbook = load_workbook(output)
    values = [cell.value for page in workbook for row in page for cell in row]
    assert total.id in values and components[0].id in values
    assert "utgift" in values
    assert not any(cell.data_type == "f" for page in workbook for row in page for cell in row)
    workbook.close()


@pytest.mark.parametrize("fault", ["missing", "self", "account", "year", "step", "relation"])
def test_storage_rejects_broken_component_parent(tmp_path: Path, fault: str) -> None:
    from innlesning import read_source

    database = tmp_path / "hierarchy.sqlite"
    extraction = ingest(database, "40032850")
    child = next(o for o in extraction.observations if o.component_of)
    changes = {
        "missing": {"component_of": "missing"},
        "self": {"component_of": child.id},
        "account": {"post": "25"},
        "year": {"budget_year": 2025},
        "step": {"step_id": "wrong"},
        "relation": {"relations": ()},
    }
    broken = replace(child, **changes[fault])
    modified = replace(
        extraction,
        observations=tuple(broken if o.id == child.id else o for o in extraction.observations),
    )
    # En annen tolkning er ikke nødvendig: strukturfeil avvises før immutable-kollisjonskontroll.
    with pytest.raises(ValueError, match="hierarki"):
        save_extraction(database, modified, read_source(database, extraction.artifact.sha256))


@pytest.mark.parametrize(
    "identifier,chapter,total,row",
    [
        ("40032763", "2445", -334_642_000, 79),
        ("40032791", "2460", 102_500_000, 101),
        ("40032976", "5440", 295_900_000_000, 62),
    ],
)
def test_authentic_original_net_variants(
    tmp_path: Path, identifier: str, chapter: str, total: int, row: int
) -> None:
    database = tmp_path / "original.sqlite"
    extraction = ingest(database, identifier)
    observations = [o for o in extraction.observations if (o.chapter, o.post) == (chapter, "24")]
    parent = next(o for o in observations if o.detail_level == "post")
    assert parent.amount_kr == total
    assert parent.pointer.locator == f"table[1]/tr[{row}]/td[5]"
    assert all(o.component_of == parent.id for o in observations if o.detail_level == "underpost")
    assert all(o.meaning == "nivaa" for o in observations)
    assert not any(
        a.code == "komponentsum_avvik" and a.chapter == chapter for a in extraction.anomalies
    )
    result = calculate_history(
        database, HistorySelection("prove", (extraction.id,), 2024, chapter, "24")
    )
    assert result.approved_balance == total
    export_extraction_excel(database, extraction.id, tmp_path / "original.xlsx")


@pytest.mark.parametrize(
    "identifier,chapter,total,row,year",
    [
        ("40033719", "5440", 191_800_000_000, 1373, 2024),
        ("40034064", "2470", -15_000_000, 78, 2024),
        ("40034226", "5440", 233_300_000_000, 52, 2024),
        ("40035422", "2445", -369_642_000, 984, 2025),
    ],
)
def test_later_net_statement_is_observed_after_without_invented_delta(
    tmp_path: Path, identifier: str, chapter: str, total: int, row: int, year: int
) -> None:
    database = tmp_path / "after.sqlite"
    extraction = ingest(database, identifier)
    observations = [o for o in extraction.observations if (o.chapter, o.post) == (chapter, "24")]
    parent = next(o for o in observations if o.detail_level == "post")
    assert parent.amount_kr == total and parent.meaning == "etter_belop"
    assert parent.pointer.locator == f"table[1]/tr[{row}]/td[4]"
    assert all(o.meaning == "etter_belop" for o in observations)
    result = calculate_history(
        database, HistorySelection("prove", (extraction.id,), year, chapter, "24")
    )
    assert result.approved_balance is None and result.steps[0].delta is None
    assert result.steps[0].after == total
    assert "mangler_delta" in {a.code for a in result.anomalies}
    export_extraction_excel(database, extraction.id, tmp_path / "after.xlsx")


def synthetic(database: Path, body: str, year: int = 2024) -> Extraction:
    old = synthetic_html(
        database, "40032850", f"<div>På statsbudsjettet for {year} bevilges under:{body}</div>"
    )
    return extract_revision(
        database,
        old.source_revision_id,
        interpretation=Interpretation(adapter_version="storting-belop-v2"),
    )


def table(chapter: str, rows: str, post: str = "24") -> str:
    return (
        "<table><tr><td>Kap.</td><td>Post</td><td>Formål</td><td>Kroner</td><td>Kroner</td></tr>"
        f"<tr><td>{chapter}</td><td></td><td>Kildens kapittelnavn</td><td></td><td></td></tr>"
        f"<tr><td></td><td>{post}</td><td>Netto driftsresultat, kan overføres</td>"
        "<td></td><td></td></tr>" + rows + "</table>"
    )


@pytest.mark.parametrize("code", ["1 ", "1. ", "1.\n", "01 "])
def test_component_number_variants_preserve_source_and_normal_code(
    tmp_path: Path, code: str
) -> None:
    database = tmp_path / "component-code.sqlite"
    extraction = synthetic(
        database,
        table(
            "2445",
            f"<tr><td></td><td></td><td>{code}Driftsinntekter, kan overføres</td>"
            "<td>-7</td><td>-7</td></tr>",
        ),
    )
    child, parent = extraction.observations
    assert child.subpost_raw == code.split(".")[0].strip()
    assert child.subpost == "1" and child.component_of == parent.id
    assert "Driftsinntekter, kan overføres" in child.purpose
    assert child.budget_side == parent.budget_side == "utgift"
    assert read_extraction(database, extraction.id) == extraction
    export_extraction_excel(database, extraction.id, tmp_path / "component-code.xlsx")


@pytest.mark.parametrize(
    "rows,code",
    [
        ("<tr><td></td><td></td><td>1 Første</td><td>2</td><td></td></tr>", "mangler_posttotal"),
        (
            "<tr><td></td><td></td><td>? Ukjent</td><td>2</td><td></td></tr>"
            "<tr><td></td><td></td><td>Sum:</td><td>2</td><td></td></tr>",
            "ukjent_komponent",
        ),
        (
            "<tr><td></td><td></td><td>1 Første</td><td>2</td><td></td></tr>"
            "<tr><td></td><td></td><td>1 Annen</td><td>3</td><td>5</td></tr>",
            "duplisert_komponent",
        ),
        (
            "<tr><td></td><td></td><td>1 Første</td><td>2</td><td></td></tr>"
            "<tr><td></td><td></td><td>3 Tredje</td><td>3</td><td>5</td></tr>",
            "manglende_komponent",
        ),
        ("<tr><td></td><td></td><td>1 Første</td><td>2</td><td>7</td></tr>", "komponentsum_avvik"),
    ],
)
def test_uncertain_component_blocks_chain_but_preserves_source_and_excel(
    tmp_path: Path, rows: str, code: str
) -> None:
    database = tmp_path / "bad-components.sqlite"
    extraction = synthetic(database, table("2445", rows))
    assert code in {a.code for a in extraction.anomalies}
    result = calculate_history(
        database, HistorySelection("prove", (extraction.id,), 2024, "2445", "24")
    )
    assert result.approved_balance is None
    assert code in {a.code for a in result.anomalies}
    export_extraction_excel(database, extraction.id, tmp_path / "bad-components.xlsx")


@pytest.mark.parametrize(
    "chapter,side",
    [
        ("0", "uavklart"),
        ("1", "utgift"),
        ("2999", "utgift"),
        ("3000", "inntekt"),
        ("5999", "inntekt"),
        ("6000", "kapitalregnskap"),
        ("9999", "kapitalregnskap"),
        ("10000", "uavklart"),
    ],
)
def test_synthetic_capital_boundaries_and_90_post(tmp_path: Path, chapter: str, side: str) -> None:
    database = tmp_path / "capital.sqlite"
    body = (
        "<table><tr><td>Kap.</td><td>Post</td><td>Formål</td><td>Kroner</td></tr>"
        f"<tr><td>{chapter}</td><td></td><td>Kapittel</td><td></td></tr>"
        "<tr><td></td><td>90</td><td>Driftsinntekter</td><td>-10</td></tr></table>"
    )
    extraction = synthetic(database, body)
    observation = extraction.observations[0]
    assert observation.amount_kr == -10 and observation.budget_side == side
    result = calculate_history(
        database, HistorySelection("prove", (extraction.id,), 2024, chapter, "90")
    )
    assert result.approved_balance == (-10 if side in {"utgift", "inntekt"} else None)
    export_extraction_excel(database, extraction.id, tmp_path / "capital.xlsx")


@pytest.mark.parametrize(
    "chapter,fdep,gdep",
    [
        ("1", "15", "91"),
        ("41", "16", "93"),
        ("2309", "23", "23"),
        ("2410", "2", "24"),
        ("2445", "15", "26"),
        ("2650", "6", "27"),
        ("5309", "23", "60"),
        ("5501", "16", "30"),
        ("5700", "16", "57"),
        ("5531", None, None),
        ("2531", None, None),
        ("999", None, None),
    ],
)
@pytest.mark.parametrize("year", [2025, 2027, 2028])
def test_explicit_organization_snapshot_has_exact_validity_and_separate_codes(
    tmp_path: Path, chapter: str, fdep: str | None, gdep: str | None, year: int
) -> None:
    database = tmp_path / "register.sqlite"
    extraction = synthetic(
        database,
        table(
            chapter, "<tr><td></td><td></td><td>1 Driftsinntekter</td><td>-7</td><td>-7</td></tr>"
        ),
        year,
    )
    org = extraction.observations[0].organization
    assert (org.fdep_nr, org.gdep_nr) == ((fdep, gdep) if year == 2027 else (None, None))
    if year == 2027 and fdep:
        assert org.valid_from == org.valid_to == "2027"
        assert extraction.interpretation.register_snapshot in org.register_version
        if fdep == "23":
            assert org.mapped_department_code is None and "Særkategori" in org.null_reason
        assert (
            "952757a9964f92a7e2cb1972fcfb6b82b426d4973de641a29fef5ea66d6c1adb"
            in org.register_source
        )
    else:
        assert "mangler_organisasjonsoppslag" in {a.code for a in extraction.anomalies}
    result = calculate_history(
        database, HistorySelection("prove", (extraction.id,), year, chapter, "24")
    )
    assert result.approved_balance == -7  # Regnekontroll godkjenner ingen organisasjon.
    assert read_extraction(database, extraction.id) == extraction
    export_extraction_excel(database, extraction.id, tmp_path / "register.xlsx")


@pytest.mark.parametrize(
    "identifier,kind,marker",
    [
        ("40035429", "fullmaktsramme", "72 mill."),
        ("40034385", "anmodning", "Akersgata"),
    ],
)
def test_zero_amount_extraction_exports_full_authority_or_request_without_history(
    tmp_path: Path, identifier: str, kind: str, marker: str
) -> None:
    from innlesning import read_revision

    database = tmp_path / "text.sqlite"
    extraction = ingest(database, identifier)
    assert extraction.observations == ()
    output = tmp_path / "text.xlsx"
    export_extraction_excel(database, extraction.id, output)
    book = load_workbook(output)
    fulltext = "".join(
        row[4] for row in book["Vedtaksfulltekst"].iter_rows(min_row=2, values_only=True)
    )
    assert fulltext == read_revision(database, extraction.source_revision_id).report.decision.html
    assert marker in fulltext
    information = list(book["Tekstopplysninger"].iter_rows(min_row=2, values_only=True))
    assert any(row[3] == kind for row in information)
    assert book["Beløpsobservasjoner"].max_row == 1
    assert not any(cell.data_type == "f" for page in book for row in page for cell in row)
    book.close()


def test_long_fulltext_is_exact_in_both_exports_and_formula_safe(tmp_path: Path) -> None:
    from analyse import export_history_excel
    from innlesning import read_revision

    database = tmp_path / "long.sqlite"
    content = '=HYPERLINK("http://example.invalid")' + "æøå" * 12000
    extraction = synthetic(
        database,
        "<p>"
        + content
        + "</p>"
        + table(
            "2445", "<tr><td></td><td></td><td>1 Driftsinntekter</td><td>-7</td><td>-7</td></tr>"
        ),
    )
    result = calculate_history(
        database, HistorySelection("prove", (extraction.id,), 2024, "2445", "24")
    )
    expected = read_revision(database, extraction.source_revision_id).report.decision.html
    for name in ("whole", "history"):
        output = tmp_path / f"{name}.xlsx"
        if name == "whole":
            export_extraction_excel(database, extraction.id, output)
        else:
            export_history_excel(database, result.id, output)
        book = load_workbook(output)
        fulltext = "".join(
            row[4] for row in book["Vedtaksfulltekst"].iter_rows(min_row=2, values_only=True)
        )
        assert fulltext == expected
        assert not any(cell.data_type == "f" for page in book for row in page for cell in row)
        first = book["Tekstopplysninger"]["E2"].value
        fragments = [
            row[2]
            for row in book["Tekstopplysninger fortsettelse"].iter_rows(min_row=2, values_only=True)
            if row[0] == "Tekstopplysninger!E2"
        ]
        assert first + "".join(fragments) == content
        book.close()
    original = database.read_bytes()
    with pytest.raises(ValueError, match="bane"):
        export_extraction_excel(database, extraction.id, database)
    assert database.read_bytes() == original


def test_version_one_identity_and_result_are_frozen_beside_v2(tmp_path: Path) -> None:
    from dataclasses import asdict

    from innlesning import read_source
    from innlesning.observations import encode

    database = tmp_path / "versions.sqlite"
    new = ingest(database, "40032850")
    old = extract_revision(database, new.source_revision_id)
    assert (
        old.interpretation.id == "09dab815eeaa206b171ff8fac0f3e51ce6d5d1a4e670853c8a9e82fc43f409a7"
    )
    old_content = encode(asdict(old))
    old_ids = tuple(o.id for o in old.observations)
    assert all(o.detail_level == "post" and o.component_of is None for o in old.observations)
    assert old.id != new.id
    assert extract_revision(database, new.source_revision_id).id == old.id
    assert ingest(database, "40032850") == new
    assert read_extraction(database, old.id) == old
    assert tuple(o.id for o in read_extraction(database, old.id).observations) == old_ids
    assert encode(asdict(read_extraction(database, old.id))) == old_content
    save_extraction(database, old, read_source(database, old.artifact.sha256))
    result = calculate_history(
        database, HistorySelection("prove", (old.id, new.id), 2024, "2470", "24")
    )
    assert result.approved_balance is None
    assert "motstridende_grunnlag" in {a.code for a in result.anomalies}
    assert (
        replace(new.interpretation, register_snapshot="another snapshot").id
        != new.interpretation.id
    )
    assert read_extraction(database, new.id) == new


@pytest.mark.parametrize(
    "identifier,year,chapter,post,value,keyword",
    [
        ("40033025", 2024, "2410", "90", 39_751_273_000, "overslagsbevilgning"),
        ("40033025", 2024, "5310", "90", 13_970_169_000, ""),
        ("40033025", 2024, "5310", "93", 8_287_563_000, ""),
        ("40033025", 2024, "1651", "98", 68_000_000_000, ""),
        ("40033025", 2024, "162", "95", 1_239_864_000, ""),
        ("40033025", 2024, "162", "96", 750_000_000, ""),
        ("40033025", 2024, "5325", "92", 30_500_000, ""),
        ("40034287", 2024, "2800", "96", 42_800_640_000, ""),
        ("40034443", 2025, "900", "32", 454_000_000, "kan overføres"),
        ("40034443", 2025, "1760", "45", -227_000_000, "kan overføres"),
        ("40035422", 2025, "41", "45", 9_700_000, "kan overføres"),
    ],
)
def test_authentic_90_new_and_annet_amounts_survive_export(
    tmp_path: Path, identifier: str, year: int, chapter: str, post: str, value: int, keyword: str
) -> None:
    database = tmp_path / "special.sqlite"
    extraction = ingest(database, identifier)
    row = next(
        o
        for o in extraction.observations
        if o.chapter == chapter and o.post == post and o.meaning in {"nivaa", "endring"}
    )
    assert row.amount_kr == value and keyword in row.purpose
    assert row.information_type == "bevilgning"
    result = calculate_history(
        database, HistorySelection("prove", (extraction.id,), year, chapter, post)
    )
    if row.meaning == "endring":
        assert result.approved_balance is None
        assert "manglende_start" in {a.code for a in result.anomalies}
    export_extraction_excel(database, extraction.id, tmp_path / "special.xlsx")


@pytest.mark.parametrize("chapter,post", [("1860", "1"), ("4860", "1"), ("5999", "90")])
def test_old_authentic_conflicts_remain_blocked_in_v2(
    tmp_path: Path, chapter: str, post: str
) -> None:
    database = tmp_path / "conflict.sqlite"
    identifiers = (
        ("40033025", "40033719", "40034287")
        if chapter == "5999"
        else ("40032850", "40033719", "40034064", "40034226")
    )
    extractions = tuple(ingest(database, id) for id in identifiers)
    result = calculate_history(
        database, HistorySelection("prove", tuple(e.id for e in extractions), 2024, chapter, post)
    )
    assert result.approved_balance is None
    assert "kjedekonflikt" in {a.code for a in result.anomalies}
    assert result.rule_version == "vedtatt-kjede-v2"


def test_authentic_raw_I_remains_unknown_in_v2_and_export(tmp_path: Path) -> None:
    database = tmp_path / "raw-I.sqlite"
    extraction = ingest(database, "40036289")
    row = next(o for o in extraction.observations if o.post == "70" and o.meaning == "etter_belop")
    assert row.raw_value == "I 615 833" and row.amount_kr is None
    assert row.structure_status == "til_kontroll"
    export_extraction_excel(database, extraction.id, tmp_path / "raw-I.xlsx")


@pytest.mark.parametrize(
    "markup,expected",
    [
        ("<td>Kap?</td><td>Post</td><td>Formål</td><td>Kroner</td>", "ukjent_tabell"),
        ("<td>Kap.</td><td>Post</td><td>Formål</td><td>Millioner</td>", "ukjent_tabell"),
    ],
)
def test_unknown_table_is_visible_in_whole_extraction_export(
    tmp_path: Path, markup: str, expected: str
) -> None:
    database = tmp_path / "table.sqlite"
    extraction = synthetic(database, f"<table><tr>{markup}</tr></table>")
    assert extraction.observations == ()
    assert expected in {a.code for a in extraction.anomalies}
    output = tmp_path / "table.xlsx"
    export_extraction_excel(database, extraction.id, output)
    book = load_workbook(output)
    assert expected in [row[1] for row in book["Avvik"].iter_rows(min_row=2, values_only=True)]
    book.close()


@pytest.mark.parametrize(
    "badrow",
    [
        "<td></td><td colspan='2'>900</td><td>Kapittel</td>",
        "<td></td><td rowspan='2'>24</td><td>Netto</td><td>1</td>",
        "<td></td><td>24</td><td>Netto</td><td>1</td><td></td>",
    ],
)
def test_unknown_layout_clears_v2_account_and_blocks_any_affected_chain(
    tmp_path: Path, badrow: str
) -> None:
    database = tmp_path / "layout.sqlite"
    body = (
        "<table><tr><td>Kap.</td><td>Post</td><td>Formål</td><td>Kroner</td></tr>"
        "<tr><td>800</td><td></td><td>Forrige kapittel</td><td></td></tr>"
        + f"<tr>{badrow}</tr>"
        + "<tr><td></td><td>32</td><td>Beløp</td><td>2</td></tr></table>"
    )
    extraction = synthetic(database, body)
    assert extraction.observations[0].chapter is None
    assert any(a.code == "ukjent_rad" and a.chapter is None for a in extraction.anomalies)
    result = calculate_history(
        database, HistorySelection("prove", (extraction.id,), 2024, "900", "32")
    )
    assert result.approved_balance is None
    export_extraction_excel(database, extraction.id, tmp_path / "layout.xlsx")


def test_subpost_zero_is_selected_separately_from_post_total(tmp_path: Path) -> None:
    database = tmp_path / "zero.sqlite"
    extraction = synthetic(
        database,
        table(
            "2445", "<tr><td></td><td></td><td>0 Kildens nullkode</td><td>-7</td><td>-7</td></tr>"
        ),
    )
    child, parent = extraction.observations
    assert child.subpost == "0" and child.detail_level == "underpost"
    assert parent.subpost is None and parent.detail_level == "post"
    result = calculate_history(
        database, HistorySelection("prove", (extraction.id,), 2024, "2445", "24", "0", "underpost")
    )
    assert result.approved_balance == -7
    assert result.input_observation_ids == (child.id,)
    assert parent.id not in result.input_observation_ids
    with pytest.raises(ValueError, match="Underpostutvalg"):
        HistorySelection("prove", (extraction.id,), 2024, "2445", "24", None, "underpost")


def test_v2_can_be_created_while_b2_job_is_pending_and_resume_works(tmp_path: Path) -> None:
    from innlesning import read_revision, resume_sync, stage_sync

    database = tmp_path / "pending.sqlite"
    extraction = ingest(database, "40032850")
    raw = (FIXTURES / "40032850.xml").read_bytes()
    metadata = json.loads((FIXTURES / "40032850.json").read_text(encoding="utf-8"))
    pending = stage_sync(database, raw, metadata, historical_sample=True, decision_id="40032850")
    assert read_sync(database, pending).status == "venter"
    old_report = read_revision(database, extraction.source_revision_id).report
    newer = extract_revision(
        database,
        extraction.source_revision_id,
        interpretation=replace(
            extraction.interpretation, reason="Syntetisk ny tolkning mens B2 venter"
        ),
    )
    assert newer.id != extraction.id
    resume_sync(database, pending)
    assert read_sync(database, pending).status == "fullfort"
    assert read_revision(database, extraction.source_revision_id).report == old_report
    assert read_extraction(database, extraction.id) == extraction


def test_cli_explicit_v2_and_whole_extraction_export(tmp_path: Path) -> None:
    from test_b3_cli import cli

    database = tmp_path / "cli.sqlite"
    extraction = ingest(database, "40035429")
    interpreted = cli(
        "innlesning",
        "tolk",
        "--database",
        str(database),
        "--revisjon",
        extraction.source_revision_id,
        "--tolkning",
        "storting-belop-v2",
    )
    assert interpreted.returncode == 0, interpreted.stdout + interpreted.stderr
    assert json.loads(interpreted.stdout)["belopsuttrekk"] == extraction.id
    output = tmp_path / "cli.xlsx"
    exported = cli(
        "analyse",
        "--database",
        str(database),
        "--belopsuttrekk",
        extraction.id,
        "--excel",
        str(output),
    )
    assert exported.returncode == 0 and output.is_file(), exported.stdout + exported.stderr


def test_generic_storage_requires_component_sum_control_to_be_visible(tmp_path: Path) -> None:
    from innlesning import read_source

    database = tmp_path / "hidden-sum.sqlite"
    extraction = synthetic(
        database,
        table("2445", "<tr><td></td><td></td><td>1 Komponent</td><td>2</td><td>7</td></tr>"),
    )
    assert "komponentsum_avvik" in {a.code for a in extraction.anomalies}
    with pytest.raises(ValueError, match="hierarki"):
        save_extraction(
            database,
            replace(extraction, anomalies=()),
            read_source(database, extraction.artifact.sha256),
        )
