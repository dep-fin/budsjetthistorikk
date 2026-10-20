"""Syntetiske reviewkontroller ved kilde → lagring → Excel-grensen."""

import json
import sqlite3
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_b3 import generic
from test_b4 import ingest, synthetic, table

from analyse import (
    HistorySelection,
    calculate_history,
    export_extraction_excel,
    export_history_excel,
)
from analyse.history import read_history
from innlesning import read_extraction, read_source
from innlesning.observations import Anomaly, encode, save_extraction


@pytest.mark.parametrize(
    "fault",
    ["relation_only", "empty_parent", "extra_target", "duplicate", "no_evidence"],
)
def test_storage_requires_one_consistent_documented_component_relation(
    tmp_path: Path, fault: str
) -> None:
    database = tmp_path / "source.sqlite"
    extraction = ingest(database, "40032850")
    child = next(o for o in extraction.observations if o.component_of)
    parent = next(o for o in extraction.observations if o.id == child.component_of)
    if fault in {"relation_only", "empty_parent"}:
        child = (
            replace(child, component_of=None)
            if fault == "relation_only"
            else replace(child, component_of="", relations=())
        )
        observations = (child,)  # En løs komponent_av uten forelder må også avvises.
    else:
        relations = {
            "extra_target": child.relations + (("komponent_av", "missing", "syntetisk"),),
            "duplicate": child.relations * 2,
            "no_evidence": (("komponent_av", parent.id, ""),),
        }
        child = replace(child, relations=relations[fault])
        observations = tuple(child if o.id == child.id else o for o in extraction.observations)
    with pytest.raises(ValueError, match="hierarki"):
        save_extraction(
            tmp_path / "rejected.sqlite",
            replace(extraction, observations=observations, anomalies=()),
            read_source(database, extraction.artifact.sha256),
        )


@pytest.mark.parametrize("subpost", ["0", "10", "25"])
def test_generic_component_code_is_preserved_without_storting_numbering(
    tmp_path: Path, subpost: str
) -> None:
    template_db = tmp_path / "template.sqlite"
    template = generic(template_db, date="2013-01-01", subpost=None, detail="post")
    parent = template.observations[0]
    child = replace(
        parent,
        pointer=replace(parent.pointer, locator="Rådata!H3", row=3, cell="H3"),
        detail_level="underpost",
        subpost_raw=subpost,
        subpost=subpost,
        component_of=parent.id,
        relations=(("komponent_av", parent.id, "Syntetisk dokumentert Excel-hierarki"),),
    )
    extraction = replace(template, observations=(parent, child))
    database = tmp_path / "generic-component.sqlite"
    save_extraction(database, extraction, read_source(template_db, template.artifact.sha256))
    assert read_extraction(database, extraction.id) == extraction
    assert extraction.anomalies == ()
    for code, detail in ((None, "post"), (subpost, "underpost")):
        result = calculate_history(
            database,
            HistorySelection("historisk", (extraction.id,), 2013, "1", "1", code, detail),
        )
        assert result.approved_balance == 10_032_000
        assert result.rule_version == "vedtatt-kjede-v1"
    output = tmp_path / "generic-component.xlsx"
    export_extraction_excel(database, extraction.id, output)
    book = load_workbook(output)
    try:
        values = [cell.value for sheet in book for row in sheet for cell in row]
        assert parent.id in values and child.id in values and subpost in values
        assert "manglende_komponent" not in values
    finally:
        book.close()


@pytest.mark.parametrize(
    "code", ["1.25", "1.2.3", "1..", "1,25", "1/2", "1-2", "1a", "1.Komponent", "١"]
)
def test_synthetic_ambiguous_component_codes_remain_raw_and_block_history(
    tmp_path: Path, code: str
) -> None:
    database = tmp_path / "ambiguous-code.sqlite"
    extraction = synthetic(
        database,
        table(
            "2445",
            f"<tr><td></td><td></td><td>{code} Komponent</td><td>7</td><td>7</td></tr>",
        ),
        year=2027,
    )
    child = extraction.observations[0]
    assert child.subpost is None and child.subpost_raw == code
    assert child.purpose == f"{code} Komponent"
    assert child.structure_status == "til_kontroll"
    assert "ukjent_komponent" in {a.code for a in extraction.anomalies}
    result = calculate_history(
        database, HistorySelection("prove", (extraction.id,), 2027, "2445", "24")
    )
    assert result.approved_balance is None
    assert read_extraction(database, extraction.id) == extraction
    output = tmp_path / "ambiguous-code.xlsx"
    export_extraction_excel(database, extraction.id, output)
    book = load_workbook(output)
    try:
        values = [cell.value for sheet in book for row in sheet for cell in row]
        assert f"{code} Komponent" in values and "ukjent_komponent" in values
    finally:
        book.close()


@pytest.mark.parametrize("heading", ["Kap.", "Kroner"])
@pytest.mark.parametrize("span", ["colspan", "rowspan"])
def test_synthetic_spanned_header_blocks_amounts_and_exports_control(
    tmp_path: Path, heading: str, span: str
) -> None:
    database = tmp_path / "header.sqlite"
    body = table(
        "2445", "<tr><td></td><td></td><td>1 Komponent</td><td>-7</td><td>-7</td></tr>"
    ).replace(f"<td>{heading}</td>", f'<td {span}="2">{heading}</td>', 1)
    extraction = synthetic(database, body, year=2027)
    assert extraction.observations == ()
    assert "ukjent_tabell" in {a.code for a in extraction.anomalies}
    assert read_extraction(database, extraction.id) == extraction
    result = calculate_history(
        database, HistorySelection("prove", (extraction.id,), 2027, "2445", "24")
    )
    assert result.approved_balance is None and result.status == "ufullstendig"
    output = tmp_path / "header.xlsx"
    export_extraction_excel(database, extraction.id, output)
    book = load_workbook(output)
    try:
        assert book["Beløpsobservasjoner"].max_row == 1
        assert "ukjent_tabell" in [
            row[1] for row in book["Avvik"].iter_rows(min_row=2, values_only=True)
        ]
        assert span in "".join(
            row[4] for row in book["Vedtaksfulltekst"].iter_rows(min_row=2, values_only=True)
        )
    finally:
        book.close()


@pytest.mark.parametrize("archived", [False, True])
def test_generic_organization_control_preserves_fresh_and_archived_v1_history(
    tmp_path: Path, archived: bool
) -> None:
    template_db = tmp_path / "template.sqlite"
    template = generic(template_db, date="2013-01-01", subpost=None, detail="post")
    extraction = replace(
        template,
        anomalies=(
            Anomaly(
                "mangler_organisasjonsoppslag",
                "Unresolved mapping",
                chapter="1",
                post="1",
                budget_year=2013,
            ),
        ),
    )
    database = tmp_path / "v1-history.sqlite"
    save_extraction(database, extraction, read_source(template_db, extraction.artifact.sha256))
    # Syntetisk arkiv produsert av den leverte B3-basen ac8d4e4; ikke dagens implementasjon.
    expected = json.loads(
        (Path(__file__).parent / "fixtures" / "synthetic-history-v1.json").read_text(
            encoding="utf-8"
        )
    )
    assert expected["id"] == "ad24cc3a988e38f2612e692f1cb2ff6fd17a3e6822f5786af86b9a50a195fc09"
    if archived:
        # Databaseoppsett for et allerede lagret resultat; kontrollen leser kun offentlige API-er.
        with sqlite3.connect(database) as connection:
            connection.execute(
                "INSERT INTO history_calculations VALUES (?,?,?)",
                (expected["id"], encode(expected), expected["created_utc"]),
            )
        assert read_history(database, expected["id"]).approved_balance is None
    selection = HistorySelection("historisk", (extraction.id,), 2013, "1", "1")
    result = calculate_history(database, selection)
    assert result.id == expected["id"] and result.rule_version == "vedtatt-kjede-v1"
    assert result.status == "ufullstendig" and result.approved_balance is None
    assert result.calculated_balance is None and result.steps[0].balance is None
    assert [a.code for a in result.anomalies] == ["mangler_organisasjonsoppslag"]
    if archived:
        assert result.created_utc == expected["created_utc"]
    assert calculate_history(database, selection).created_utc == result.created_utc
    output = tmp_path / "v1-history.xlsx"
    export_history_excel(database, result.id, output)
    book = load_workbook(output)
    try:
        assert book["Historikk"]["J2"].value is None
        assert any(
            "mangler_organisasjonsoppslag" in str(cell.value)
            for sheet in book
            for row in sheet
            for cell in row
        )
    finally:
        book.close()


def test_many_synthetic_blocks_keep_only_their_own_observed_parent(tmp_path: Path) -> None:
    database = tmp_path / "many-blocks.sqlite"
    blocks = []
    for index in range(160):
        header = "<tr><td></td><td>24</td><td>Netto driftsresultat</td><td></td><td></td></tr>"
        components = (
            f"<tr><td></td><td></td><td>1 Første</td><td>{index}</td><td></td></tr>"
            "<tr><td></td><td></td><td>2 Andre</td><td>7</td>"
        )
        if index % 2:
            ending = (
                "<td></td></tr><tr><td></td><td></td><td>Sum:</td>"
                f"<td>{index + 7}</td><td></td></tr>"
            )
        else:
            ending = f"<td>{index + 7}</td></tr>"
        blocks.append(header + components + ending)
    body = (
        "<table><tr><td>Kap.</td><td>Post</td><td>Formål</td><td>Kroner</td><td>Kroner</td></tr>"
        "<tr><td>2445</td><td></td><td>Syntetisk kapittel</td><td></td><td></td></tr>"
        + "".join(blocks)
        + "</table>"
    )
    extraction = synthetic(database, body, year=2027)
    assert extraction.anomalies == ()
    assert len(extraction.observations) == 480
    for index in range(160):
        first, second, parent = extraction.observations[index * 3 : index * 3 + 3]
        assert (first.amount_kr, second.amount_kr, parent.amount_kr) == (index, 7, index + 7)
        assert first.component_of == second.component_of == parent.id
        assert first.step_id == second.step_id == parent.step_id
        assert first.relations[0][1] == second.relations[0][1] == parent.id
    assert len({o.step_id for o in extraction.observations}) == 160
    assert read_extraction(database, extraction.id) == extraction
    output = tmp_path / "many-blocks.xlsx"
    export_extraction_excel(database, extraction.id, output)
    book = load_workbook(output)
    try:
        assert book["Beløpsobservasjoner"].max_row == 481
        assert book["Hierarki"].max_row == 321
    finally:
        book.close()


def test_many_invalid_blocks_preserve_every_control_and_source_cell(tmp_path: Path) -> None:
    database = tmp_path / "many-invalid-blocks.sqlite"
    header = "<tr><td></td><td>24</td><td>Netto driftsresultat</td><td></td><td></td></tr>"
    component = "<tr><td></td><td></td><td>1 Komponent</td><td>7</td><td>8</td></tr>"
    body = (
        "<table><tr><td>Kap.</td><td>Post</td><td>Formål</td><td>Kroner</td><td>Kroner</td></tr>"
        "<tr><td>2445</td><td></td><td>Kapittel</td><td></td><td></td></tr>"
        + (header + component) * 320
        + "</table>"
    )
    extraction = synthetic(database, body, year=2027)
    expected_controls = []
    for index in range(320):
        child, parent = extraction.observations[index * 2 : index * 2 + 2]
        assert (child.amount_kr, parent.amount_kr) == (7, 8)
        assert child.component_of == parent.id
        assert child.pointer.locator == f"table[1]/tr[{index * 2 + 4}]/td[4]"
        assert parent.pointer.locator == f"table[1]/tr[{index * 2 + 4}]/td[5]"
        expected_controls.append(
            Anomaly(
                "komponentsum_avvik",
                "Komponenter og observert total stemmer ikke eksakt",
                evidence=parent.id,
                chapter="2445",
                post="24",
                budget_year=2027,
            )
        )
    assert len(extraction.observations) == 640
    assert extraction.anomalies == tuple(expected_controls)
    assert read_extraction(database, extraction.id) == extraction
    raw = read_source(database, extraction.artifact.sha256)
    repeated_controls = replace(
        extraction,
        anomalies=(extraction.anomalies[-1], *extraction.anomalies, extraction.anomalies[0]),
    )
    copied_database = tmp_path / "repeated-controls.sqlite"
    save_extraction(copied_database, repeated_controls, raw)
    assert read_extraction(copied_database, extraction.id) == repeated_controls
    with pytest.raises(ValueError, match="hierarki"):
        save_extraction(
            tmp_path / "hidden-control.sqlite",
            replace(extraction, anomalies=extraction.anomalies[:-1]),
            raw,
        )
    result = calculate_history(
        database, HistorySelection("prove", (extraction.id,), 2027, "2445", "24")
    )
    assert result.status == "ufullstendig" and result.approved_balance is None
    assert read_history(database, result.id).approved_balance is None
    output = tmp_path / "many-invalid-blocks.xlsx"
    export_extraction_excel(database, extraction.id, output)
    book = load_workbook(output)
    try:
        observations = list(book["Beløpsobservasjoner"].iter_rows(min_row=2, values_only=True))
        assert [(row[0], row[12], row[14]) for row in observations] == [
            (o.id, o.amount_kr, o.pointer.locator) for o in extraction.observations
        ]
        controls = [
            json.loads(row[3]) for row in book["Avvik"].iter_rows(min_row=2, values_only=True)
        ]
        assert [control["code"] for control in controls] == ["komponentsum_avvik"] * 320
        assert [control["evidence"] for control in controls] == [
            control.evidence for control in expected_controls
        ]
        assert book["Hierarki"].max_row == 321
    finally:
        book.close()


def test_mixed_v2_selection_keeps_generic_organization_control(tmp_path: Path) -> None:
    database = tmp_path / "mixed.sqlite"
    storting = synthetic(
        database,
        table("2445", "<tr><td></td><td></td><td>1 Komponent</td><td>-7</td><td>-7</td></tr>"),
    )
    assert "mangler_organisasjonsoppslag" in {a.code for a in storting.anomalies}
    template_db = tmp_path / "template.sqlite"
    template = generic(template_db)
    generic_control = replace(
        template,
        dataset="prove",
        observations=(),
        anomalies=(
            Anomaly(
                "mangler_organisasjonsoppslag",
                "Syntetisk generisk kontrollgrunnlag",
                chapter="2445",
                post="24",
                budget_year=2024,
            ),
        ),
    )
    save_extraction(database, generic_control, read_source(template_db, template.artifact.sha256))
    result = calculate_history(
        database,
        HistorySelection("prove", (storting.id, generic_control.id), 2024, "2445", "24"),
    )
    assert result.rule_version == "vedtatt-kjede-v2" and result.approved_balance is None
    assert [a.reason for a in result.anomalies if a.code == "mangler_organisasjonsoppslag"] == [
        "Syntetisk generisk kontrollgrunnlag"
    ]
