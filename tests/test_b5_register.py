"""B5: offentlig filgrense → lagring → ny tolkning; kildeprøver beholdes."""

import csv
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_b4 import ingest

from innlesning import extract_revision, read_extraction, read_source
from innlesning.observations import Interpretation
from innlesning.registers import (
    ChapterEntry,
    DepartmentEntry,
    PostEntry,
    export_register,
    import_register,
    read_register,
    seed_register,
)


def edit_rows(path: Path, edit: object) -> None:
    if path.suffix == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as source:
            reader = csv.DictReader(source, delimiter=";")
            headers = reader.fieldnames
            rows = list(reader)
        edit(rows)  # type: ignore[operator]
        with path.open("w", encoding="utf-8-sig", newline="") as target:
            writer = csv.DictWriter(target, fieldnames=headers, delimiter=";")
            writer.writeheader()
            writer.writerows(rows)
    else:
        book = load_workbook(path)
        sheet = book.active
        headers = [c.value for c in sheet[1]]
        rows = [
            dict(
                zip(
                    headers,
                    [str(c.value) if c.value is not None else "" for c in row],
                    strict=False,
                )
            )
            for row in sheet.iter_rows(min_row=2)
        ]
        edit(rows)  # type: ignore[operator]
        sheet.delete_rows(2, sheet.max_row)
        for row in rows:
            sheet.append([row.get(h, "") for h in headers])
        book.save(path)
        book.close()


@pytest.mark.parametrize("suffix", [".csv", ".xlsx"])
def test_register_roundtrip_reprocesses_90032_without_rewriting_old(
    tmp_path: Path, suffix: str
) -> None:
    db = tmp_path / "b5.sqlite"
    old = ingest(db, "40034443")
    original = read_source(db, old.artifact.sha256)
    seed = seed_register(db)
    assert len(seed.departments) == 17  # 16 ordinære + særkategori Ymse
    assert seed.lookup("5309", 2027).gdep_nr == "60"
    assert seed.lookup("2309", 2027).mapped_department_code is None
    assert seed.lookup("900", 2025).fdep_nr is None
    path = tmp_path / ("register" + suffix)
    export_register(db, seed.id, path)

    def extend(rows: list[dict[str, str]]) -> None:
        rows.extend(
            [
                {
                    "radtype": "departement",
                    **{
                        k: str(v)
                        for k, v in DepartmentEntry(
                            "9",
                            "Nærings- og fiskeridepartementet",
                            "NFD",
                            "NFD",
                            "ordinær",
                            2025,
                            2025,
                            "vedtak 40034443",
                            "tabell 1",
                            "Kildens navn",
                        )
                        .to_row()
                        .items()
                    },
                },
                {
                    "radtype": "kapittel",
                    **{
                        k: str(v)
                        for k, v in ChapterEntry(
                            "900",
                            "9",
                            "",
                            "uavklart",
                            2025,
                            2025,
                            "vedtak 40034443",
                            "table[1]/tr[3]",
                            "Administrativ avklaring",
                        )
                        .to_row()
                        .items()
                    },
                },
                {
                    "radtype": "post",
                    **{
                        k: str(v)
                        for k, v in PostEntry(
                            "900",
                            "32",
                            2025,
                            2025,
                            "vedtak 40034443",
                            "table[1]/tr[4]",
                            "Ny post, ingen inngående null",
                        )
                        .to_row()
                        .items()
                    },
                },
            ]
        )

    edit_rows(path, extend)
    result = import_register(db, path)
    assert not result.errors and result.snapshot_id != seed.id
    snapshot = read_register(db, result.snapshot_id)
    assert snapshot.catalogued("900", "32", 2025)
    assert snapshot.lookup("900", 2025).mapped_department_code == "9"
    new = extract_revision(
        db,
        old.source_revision_id,
        interpretation=Interpretation(
            adapter_version="storting-belop-v2", register_snapshot=snapshot.id
        ),
    )
    amount = next(o for o in new.observations if (o.chapter, o.post) == ("900", "32"))
    assert amount.amount_kr == 454_000_000 and amount.organization.fdep_nr == "9"
    assert amount.id not in {o.id for o in old.observations}
    assert read_extraction(db, old.id) == old and read_source(db, old.artifact.sha256) == original
    assert import_register(db, path) == result
    assert replace(snapshot, departments=tuple(reversed(snapshot.departments))).id == snapshot.id
    from analyse import (
        HistorySelection,
        export_controlled_excel,
        read_history,
        reprocess_controlled,
    )
    from innlesning import (
        activate_basis,
        define_control_basis,
        export_controls,
        import_controls,
        store_evidence,
    )

    basis = define_control_basis(db, "Kildebegrunnet katalogkontroll av ny 900/32")
    activate_basis(db, (new.id,), snapshot.id, basis)
    proof = store_evidence(
        db,
        Path(__file__).parent / "fixtures/b3/40034443.xml",
        version="40034443",
        locator="table[1]/tr[4]/td[4]",
        role="kontrollkilde",
        url="https://www.stortinget.no/no/Saker-og-publikasjoner/Vedtak/Vedtak/Sak/?did=40034443&p=102001",
        parent_sha256="9cfeaef59510a54ff8b2a695e0be368e67cfbab574f8d3f466c0c2f0c3c92680",
    )
    control_file = tmp_path / ("catalogue-control" + suffix)
    export_controls(
        db,
        (new.id,),
        control_file,
        register_snapshot=snapshot.id,
        control_basis=basis,
        observation_ids=(amount.id,),
    )
    edit_rows(
        control_file,
        lambda rows: rows[0].update(
            status="godkjent",
            godkjente_felter="amount_kr,organization,catalogue",
            godkjent_bruk="kjedegrunnlag",
            bevis_id=proof.id,
            kontrollor="Testkontrollør",
            begrunnelse="Kildebegrunnet ny katalogpost",
        ),
    )
    assert not import_controls(db, control_file).errors
    processed = reprocess_controlled(
        db,
        (new.id,),
        snapshot.id,
        basis,
        selection=HistorySelection("prove", (new.id,), 2025, "900", "32"),
    )
    assert processed.rows[0].effective_amount_kr == 454_000_000
    assert processed.approved_chain_balance is None
    assert "manglende_start" in {
        a.code for a in read_history(db, processed.technical_calculation_id).anomalies
    }
    export_controlled_excel(db, processed.id, tmp_path / "new-post.xlsx")
    assert read_extraction(db, old.id) == old


@pytest.mark.parametrize("fault", ["overlap", "year", "gap", "source", "ymse"])
def test_invalid_register_is_durable_and_publishes_nothing(tmp_path: Path, fault: str) -> None:
    from innlesning.maintenance_store import read_import

    db = tmp_path / "invalid.sqlite"
    seed = seed_register(db)
    path = tmp_path / "edit.csv"
    export_register(db, seed.id, path)

    def mutate(rows: list[dict[str, str]]) -> None:
        if fault == "overlap":
            rows.append(dict(rows[1]))
        elif fault == "year":
            rows[1]["fra_ar"] = "2028"
        elif fault == "gap":
            next(r for r in rows if r.get("nummer") == "15")["til_ar"] = "2026"
        elif fault == "source":
            rows[1]["kilde"] = ""
        else:
            next(r for r in rows if r.get("nummer") == "23")["kategori"] = "ordinær"

    edit_rows(path, mutate)
    failed = import_register(db, path)
    assert failed.errors and failed.snapshot_id is None and not failed.new_ids
    assert all(e.sheet == "Register" and e.field and e.row > 1 for e in failed.errors)
    assert read_import(db, failed.id) == failed and read_register(db, seed.id) == seed


def test_free_numbers_can_be_assigned_and_abbreviation_reused_in_other_period(
    tmp_path: Path,
) -> None:
    db = tmp_path / "free.sqlite"
    seed = seed_register(db)
    path = tmp_path / "free.csv"
    export_register(db, seed.id, path)

    def add(rows: list[dict[str, str]]) -> None:
        for code, start in (("10", 2025), ("12", 2026)):
            rows.append(
                {
                    "radtype": "departement",
                    **{
                        k: str(v)
                        for k, v in DepartmentEntry(
                            code,
                            "Kildebegrunnet navn",
                            "UD",
                            "Kildenavn",
                            "ordinær",
                            start,
                            start,
                            "Kilde",
                            "Ref",
                            "Avklaring",
                        )
                        .to_row()
                        .items()
                    },
                }
            )

    edit_rows(path, add)
    result = import_register(db, path)
    assert not result.errors
    assert {d.number for d in read_register(db, result.snapshot_id).departments} >= {"10", "12"}
