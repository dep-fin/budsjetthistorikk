"""Arbeidsfilens faktiske celler valideres selv med feil dimensjonsmetadata."""

import re
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from test_b4 import synthetic
from test_b5_register import edit_rows

from innlesning import extract_revision, read_import, read_import_source
from innlesning.controls import (
    activate_basis,
    define_control_basis,
    export_controls,
    import_controls,
    latest_decision,
)
from innlesning.observations import Interpretation
from innlesning.registers import (
    export_register,
    import_register,
    read_current_register,
    read_register,
    seed_register,
)

DIMENSIONS = ("correct", "under_rows", "under_columns", "over_rows", "over_columns", "missing")


def working_file(database: Path, path: Path, kind: str):
    seed = seed_register(database)
    if kind == "register":
        export_register(database, seed.id, path)
        edit_rows(path, lambda rows: rows.__setitem__(slice(None), rows[:3]))
        return seed, ()
    original = synthetic(
        database,
        "<table><tr><td>Kap.</td><td>Post</td><td>Formål</td><td>Kroner</td></tr>"
        "<tr><td>41</td><td></td><td>Stortinget</td><td></td></tr>"
        "<tr><td></td><td>1</td><td>Drift</td><td>10</td></tr>"
        "<tr><td></td><td>2</td><td>Utstyr</td><td>20</td></tr></table>",
        year=2027,
    )
    extraction = extract_revision(
        database,
        original.source_revision_id,
        interpretation=Interpretation(
            adapter_version="storting-belop-v2", register_snapshot=seed.id
        ),
    )
    basis = define_control_basis(database, "Syntetisk kontroll av filgrensen")
    activate_basis(database, (extraction.id,), seed.id, basis)
    export_controls(
        database,
        (extraction.id,),
        path,
        register_snapshot=seed.id,
        control_basis=basis,
    )

    def reject(rows):
        for row in rows:
            row.update(status="avvist", kontrollor="Test", begrunnelse="Kun filprøve")

    edit_rows(path, reject)
    return seed, tuple(o.id for o in extraction.observations)


def dimension_metadata(path: Path, reference: str | None) -> None:
    with ZipFile(path) as source:
        members = [(member, source.read(member)) for member in source.infolist()]
    with ZipFile(path, "w") as target:
        for member, content in members:
            if member.filename == "xl/worksheets/sheet1.xml":
                replacement = (
                    b"" if reference is None else f'<dimension ref="{reference}"/>'.encode()
                )
                content, count = re.subn(rb"<dimension\b[^>]*/>", replacement, content)
                assert count == 1
            target.writestr(member, content)


def misstate_dimensions(path: Path, dimension: str) -> None:
    if dimension == "correct":
        return
    book = load_workbook(path)
    page = book.active
    column, row = get_column_letter(page.max_column), page.max_row
    book.close()
    dimension_metadata(
        path,
        {
            "under_rows": f"A1:{column}2",
            "under_columns": f"A1:B{row}",
            "over_rows": f"A1:{column}10002",
            "over_columns": f"A1:XFD{row}",
            "missing": None,
        }[dimension],
    )


@pytest.mark.parametrize("kind", ["register", "kontroll"])
@pytest.mark.parametrize("dimension", DIMENSIONS)
def test_later_invalid_rows_are_preserved_and_never_partially_published(
    tmp_path: Path, kind: str, dimension: str
) -> None:
    database = tmp_path / "rows.sqlite"
    path = tmp_path / "rows.xlsx"
    seed, observations = working_file(database, path, kind)
    edit_rows(path, lambda rows: rows[-1].update(begrunnelse=""))
    misstate_dimensions(path, dimension)
    raw = path.read_bytes()
    imported = (import_register if kind == "register" else import_controls)(database, path)
    assert imported.errors and not imported.new_ids and imported.snapshot_id is None
    assert any(error.row >= 3 and error.field == "begrunnelse" for error in imported.errors)
    assert read_import(database, imported.id) == imported
    assert read_import_source(database, imported.id) == raw
    assert read_current_register(database).id == seed.id
    assert all(latest_decision(database, observation) is None for observation in observations)


@pytest.mark.parametrize("kind", ["register", "kontroll"])
@pytest.mark.parametrize("dimension", DIMENSIONS)
def test_valid_workfiles_use_all_cells_independently_of_dimensions(
    tmp_path: Path, kind: str, dimension: str
) -> None:
    database = tmp_path / "valid.sqlite"
    path = tmp_path / "valid.xlsx"
    _, observations = working_file(database, path, kind)
    misstate_dimensions(path, dimension)
    raw = path.read_bytes()
    imported = (import_register if kind == "register" else import_controls)(database, path)
    assert not imported.errors
    assert read_import_source(database, imported.id) == raw
    if kind == "register":
        snapshot = read_register(database, imported.snapshot_id)
        assert len(snapshot.departments) == 2 and not snapshot.chapters and not snapshot.posts
        assert read_current_register(database).id == snapshot.id
    else:
        assert len(imported.new_ids) == 2
        assert all(
            latest_decision(database, observation).decision.status == "avvist"
            for observation in observations
        )


@pytest.mark.parametrize("kind", ["register", "kontroll"])
@pytest.mark.parametrize("dimension", DIMENSIONS)
@pytest.mark.parametrize("invalid", ["formula", "extra_column"])
def test_formula_and_extra_column_cannot_hide_behind_dimensions(
    tmp_path: Path, kind: str, dimension: str, invalid: str
) -> None:
    database = tmp_path / "cells.sqlite"
    path = tmp_path / "cells.xlsx"
    seed, observations = working_file(database, path, kind)
    book = load_workbook(path)
    page = book.active
    if invalid == "formula":
        headers = [cell.value for cell in page[1]]
        page.cell(page.max_row, headers.index("begrunnelse") + 1, "=1+1")
    else:
        page.cell(page.max_row, page.max_column + 1, "Ekstra opplysning")
    book.save(path)
    book.close()
    misstate_dimensions(path, dimension)
    raw = path.read_bytes()
    imported = (import_register if kind == "register" else import_controls)(database, path)
    assert imported.errors and not imported.new_ids and imported.snapshot_id is None
    if invalid == "formula":
        assert any("Formel/feilcelle" in error.message for error in imported.errors)
    assert read_import(database, imported.id) == imported
    assert read_import_source(database, imported.id) == raw
    assert read_current_register(database).id == seed.id
    assert all(latest_decision(database, observation) is None for observation in observations)


@pytest.mark.parametrize("kind", ["register", "kontroll"])
@pytest.mark.parametrize("invalid", ["row", "cell"])
def test_actual_row_and_cell_limits_apply_with_missing_dimensions(
    tmp_path: Path, kind: str, invalid: str
) -> None:
    database = tmp_path / "limits.sqlite"
    path = tmp_path / "limits.xlsx"
    seed, observations = working_file(database, path, kind)
    book = load_workbook(path)
    page = book.active
    if invalid == "row":
        page.cell(10002, 1, "Sen rad")
    else:
        page.cell(page.max_row, 1, "x" * 30001)
    book.save(path)
    book.close()
    dimension_metadata(path, None)
    raw = path.read_bytes()
    imported = (import_register if kind == "register" else import_controls)(database, path)
    assert imported.errors and not imported.new_ids and imported.snapshot_id is None
    assert read_import_source(database, imported.id) == raw
    assert read_current_register(database).id == seed.id
    assert all(latest_decision(database, observation) is None for observation in observations)


@pytest.mark.parametrize("kind", ["register", "kontroll"])
def test_actual_cell_row_outside_limit_is_preserved_without_publishing(
    tmp_path: Path, kind: str
) -> None:
    database = tmp_path / "coordinate.sqlite"
    path = tmp_path / "coordinate.xlsx"
    seed, observations = working_file(database, path, kind)
    with ZipFile(path) as source:
        members = [(member, source.read(member)) for member in source.infolist()]
    with ZipFile(path, "w") as target:
        for member, content in members:
            if member.filename == "xl/worksheets/sheet1.xml":
                # En liten feilformatert fil: cellekoordinaten avviker fra radens metadata.
                content, count = re.subn(
                    rb'(<c r="[A-Z]+)3("[^>]*>)', rb"\g<1>10002\2", content, count=1
                )
                assert count == 1
            target.writestr(member, content)
    raw = path.read_bytes()
    imported = (import_register if kind == "register" else import_controls)(database, path)
    assert imported.errors and not imported.new_ids and imported.snapshot_id is None
    assert read_import_source(database, imported.id) == raw
    assert read_current_register(database).id == seed.id
    assert all(latest_decision(database, observation) is None for observation in observations)
