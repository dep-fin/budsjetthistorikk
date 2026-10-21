"""B6 gjennom lagrede offentlige grenser; ingen nett eller direkte SQL-endringer."""

import json
from dataclasses import asdict
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_b2 import RAW, synchronize
from test_b4 import ingest

from analyse.session import SessionSelection, build_session_report
from analyse.session_excel import export_session_excel
from innlesning import extract_revision, read_sync
from innlesning.sync import read_revision, stage_sync


def test_global_inventory_is_visible_outside_account_filter_and_without_extraction(
    tmp_path: Path,
) -> None:
    db = tmp_path / "session.sqlite"
    job = synchronize(db)
    revision = read_sync(db, job).entries[0].revision_id
    selection = SessionSelection("prove", "2024-2025", (job,), chapters=("999",))
    report = build_session_report(db, selection)
    assert len(report.inventory) == 1
    assert report.inventory[0].decision_id == "40033944"
    assert report.inventory[0].structure_status == "ikke_uttrukket"
    assert report.coverage.source_complete
    assert not report.coverage.structure_complete and not report.coverage.professional_complete
    extraction = extract_revision(db, revision)
    selection = SessionSelection(
        "prove", "2024-2025", (job,), extraction_ids=(extraction.id,), chapters=("999",)
    )
    report = build_session_report(db, selection)
    assert report.rows == ()
    assert len(report.inventory) == 1 and report.inventory[0].excluded_count == 8
    assert report.sums == ()
    assert not report.coverage.professional_complete
    output = tmp_path / "session.xlsx"
    export_session_excel(db, selection, output)
    with_output = load_workbook(output)
    assert {
        "Sesjonsgrunnlag",
        "Kildeinventar",
        "Dekning",
        "Beløpsobservasjoner",
        "Avstemming",
    } <= set(with_output.sheetnames)
    assert "40033944" in [c.value for row in with_output["Kildeinventar"] for c in row]
    with_output.close()
    spec = tmp_path / "selection.json"
    spec.write_text(json.dumps(asdict(selection)), encoding="utf-8")
    assert RAW


def session_for(db: Path, extractions: tuple, **kwargs) -> SessionSelection:
    jobs = tuple(
        dict.fromkeys(
            read_revision(db, e.source_revision_id).current_observation_id for e in extractions
        )
    )
    main = tuple(j for j in jobs if read_sync(db, j).session == "2024-2025")
    supplement = tuple(j for j in jobs if read_sync(db, j).session != "2024-2025")
    return SessionSelection(
        "prove",
        "2024-2025",
        main,
        extraction_ids=tuple(e.id for e in extractions),
        supplemental_inventory_job_ids=supplement,
        **kwargs,
    )


def test_filters_use_or_within_and_between_dimensions_unknown_department_is_visible(
    tmp_path: Path,
) -> None:
    db = tmp_path / "filters.sqlite"
    change = ingest(db, "40034443")
    selection = session_for(db, (change,), chapters=("900", "1760"), posts=("32",))
    report = build_session_report(db, selection)
    assert [(r.observation.chapter, r.observation.post) for r in report.rows] == [("900", "32")]
    report = build_session_report(db, SessionSelection(**{**asdict(selection), "fdep_nrs": ("8",)}))
    assert len(report.rows) == 1 and report.rows[0].selection_status == "uavklart_treff"
    assert report.rows[0].observation.organization.fdep_nr is None
    assert report.sums[0].complete_total_kr is None
    assert "manglende_start" in report.sums[0].issues


def test_department_filter_keeps_excluded_opening_in_chain_basis(tmp_path: Path) -> None:
    from dataclasses import replace

    from test_b3 import synthetic_html

    from innlesning import read_source
    from innlesning.observations import Organization, save_extraction

    db = tmp_path / "chain-filter.sqlite"
    extractions = []
    for identifier, department, phrase, purpose, amount in (
        ("40032850", "8", "På statsbudsjettet for 2024 bevilges under:", "Drift", 10),
        ("40033805", "9", "I statsbudsjettet for 2024 gjøres følgende endringer:", "økes med", 5),
    ):
        old = synthetic_html(
            db,
            identifier,
            f"<div>{phrase}<table><tr><td>Kap.</td><td>Post</td><td>Formål</td>"
            "<td>Kroner</td></tr><tr><td>41</td><td></td><td>Stortinget</td><td></td></tr>"
            f"<tr><td></td><td>1</td><td>{purpose}</td><td>{amount}</td></tr></table></div>",
        )
        interpretation = replace(old.interpretation, reason="Syntetisk historisk departement")
        alternative = replace(
            old,
            interpretation=interpretation,
            observations=tuple(
                replace(
                    o,
                    interpretation_id=interpretation.id,
                    organization=Organization(fdep_nr=department),
                )
                for o in old.observations
            ),
        )
        save_extraction(db, alternative, read_source(db, old.artifact.sha256))
        extractions.append(alternative)
    report = build_session_report(db, session_for(db, tuple(extractions), fdep_nrs=("9",)))
    assert len(report.rows) == 1 and report.rows[0].observation.meaning == "endring"
    assert len(report.unfiltered_rows) == 2
    assert report.unfiltered_rows[0].selection_status == "utenfor"
    assert "manglende_start" not in report.rows[0].issues
    assert "manglende_start" not in report.issues


def test_session_and_budget_year_are_separate_supplement_has_separate_coverage(
    tmp_path: Path,
) -> None:
    db = tmp_path / "years.sqlite"
    original = ingest(db, "40032850")
    change = ingest(db, "40033805")
    next_year = ingest(db, "40034443")
    report = build_session_report(db, session_for(db, (original, change, next_year)))
    assert {r.observation.budget_year for r in report.rows} == {2024, 2025}
    assert report.supplemental_coverage[0].session == "2023-2024"
    assert any(r.supplemental for r in report.inventory)
    assert any(
        r.observation.api_id == "40033805" and r.observation.budget_year == 2024
        for r in report.rows
    )
    assert not report.coverage.source_complete  # Delvise autentiske utdrag er ingen hel sesjon.
    assert all(s.complete_total_kr is None for s in report.sums)


@pytest.mark.parametrize("kind", ["pending", "failed", "missing", "empty"])
def test_pending_failed_missing_and_empty_inventory_never_claim_full_coverage(
    tmp_path: Path,
    kind: str,
) -> None:
    from lxml import etree
    from test_b2 import metadata

    db = tmp_path / "attempt.sqlite"
    first = synchronize(db)
    extraction = extract_revision(db, read_sync(db, first).entries[0].revision_id)
    if kind == "pending":
        stage_sync(db, RAW, metadata(RAW, tick=1), historical_sample=True)
        chosen = first
    elif kind == "failed":
        chosen = synchronize(db, RAW[:-20], tick=1)
    else:
        root = etree.fromstring(RAW)
        listing = root.find("{http://data.stortinget.no}stortingsvedtak_liste")
        assert listing is not None
        for node in list(listing):
            listing.remove(node)
        chosen = synchronize(db, etree.tostring(root), tick=1)
    selection = SessionSelection(
        "prove",
        "2024-2025",
        (chosen,),
        extraction_ids=(extraction.id,) if kind == "pending" else (),
    )
    report = build_session_report(db, selection)
    assert len(report.inventory) == 1
    assert not report.coverage.source_complete
    assert not report.coverage.professional_complete
    assert report.inventory[0].issues


@pytest.mark.parametrize("pending", [True, False])
def test_excel_attempt_context_has_named_inventory_columns_without_empty_history_headers(
    tmp_path: Path, pending: bool
) -> None:
    from test_b2 import metadata

    db = tmp_path / "attempt-columns.sqlite"
    first = synchronize(db)
    extraction = extract_revision(db, read_sync(db, first).entries[0].revision_id)
    if pending:
        stage_sync(db, RAW, metadata(RAW, tick=1), historical_sample=True)
    else:
        synchronize(db, RAW[:-20], tick=1)
    selection = SessionSelection("prove", "2024-2025", (first,), extraction_ids=(extraction.id,))
    output = tmp_path / "attempt-columns.xlsx"
    report = export_session_excel(db, selection, output)
    book = load_workbook(output)
    try:
        for name, expected_columns in (("Kildeinventar", 27), ("Historikk", 33)):
            values = list(book[name].values)
            assert len(values[0]) == expected_columns
            assert all(value is not None for value in values[0])
        assert list(book["Historikk"].values)[0][-1] == "Avvik"
        inventory = list(book["Kildeinventar"].values)
        mapped = dict(zip(inventory[0], inventory[1], strict=True))
        context = report.inventory[0].context
        assert json.loads(mapped["Siste forsøks feil"]) == list(context.latest_attempt_errors)
        assert mapped["Siste forsøks rå SHA-256"] == context.latest_attempt_source_sha256
        assert mapped["Siste forsøks omfang"] == context.latest_attempt_scope
        assert mapped["Forsøksstatus"] == ("venter" if pending else "feilet")
    finally:
        book.close()


def test_reconciliation_uses_union_keys_missing_is_not_zero_and_rejects_other_concepts(
    tmp_path: Path,
) -> None:
    from analyse.session_reconciliation import AmountScope, Reconciliation, ReferenceValue

    db = tmp_path / "reconcile.sqlite"
    extraction = ingest(db, "40034443")
    target = next(o for o in extraction.observations if o.chapter == "900")
    scope = AmountScope(2025, "2025-01-01", "2025-12-31", "avgrenset_endring")
    reference = ReferenceValue(
        "900",
        "32",
        454_000_000,
        False,
        "https://www.stortinget.no/",
        extraction.artifact.sha256,
        target.pointer.locator,
        "Autentisk prøve",
    )
    missing = ReferenceValue(
        "999", "32", 0, False, "https://example.invalid/", "a" * 64, "rad 1", "Syntetisk nullprøve"
    )
    check = Reconciliation("Avgrenset prøve", scope, scope, (target.id,), (reference, missing))
    selection = session_for(db, (extraction,), reconciliations=(check,))
    report = build_session_report(db, selection)
    assert {r.status for r in report.reconciliations[0].rows} == {
        "ikke_vurderbar",
        "mangler_i_uttrekk",
    }
    wrong = AmountScope(2025, "2025-01-01", "2025-12-31", "rekneskap")
    check = Reconciliation("Feil begrep", scope, wrong, (target.id,), (reference,))
    report = build_session_report(db, session_for(db, (extraction,), reconciliations=(check,)))
    assert report.reconciliations[0].rows[0].status == "uforenlig_omfang"
    assert report.reconciliations[0].rows[0].actual_amount_kr is None


def test_independent_2024_control_recomputes_raw_fields_and_retains_union_residuals() -> None:
    from analyse.preserved_controls import audit_preserved_2024

    result = audit_preserved_2024(Path(__file__).parents[1] / "docs/research/issue-2")
    assert result.dfo_row_count == 4268
    assert result.dfo_post_count == 1638 and result.pdf_post_count == 1655
    assert result.common_count == 1625 and result.amount_difference_count == 0
    assert len(result.residuals) == 43
    assert sum(r.dfo_amount_kr is None for r in result.residuals) == 30
    assert sum(r.pdf_amount_kr is None for r in result.residuals) == 13
    assert result.original_post_count == 1574 and result.original_difference_count == 0
    assert result.session_inventory_count == 2183
    assert result.professional_chain_complete is False


def test_unknown_account_and_year_are_unresolved_filter_hits(tmp_path: Path) -> None:
    from dataclasses import replace

    from test_b4 import synthetic

    from innlesning.observations import save_extraction
    from innlesning.storage import read_source

    db = tmp_path / "unknown.sqlite"
    extraction = synthetic(
        db,
        "<table><tr><td>Kap.</td><td>Post</td><td>Formål</td><td>Kroner</td></tr>"
        "<tr><td>800</td><td></td><td>Forrige kapittel</td><td></td></tr>"
        "<tr><td></td><td colspan='2'>900</td><td>Kapittel</td></tr>"
        "<tr><td></td><td>32</td><td>Beløp</td><td>2</td></tr></table>",
    )
    interpretation = replace(extraction.interpretation, reason="Syntetisk ukjent år")
    observations = tuple(
        replace(o, interpretation_id=interpretation.id, budget_year=None)
        for o in extraction.observations
    )
    alternative = replace(extraction, interpretation=interpretation, observations=observations)
    save_extraction(db, alternative, read_source(db, extraction.artifact.sha256))
    job = read_revision(db, extraction.source_revision_id).current_observation_id
    report = build_session_report(
        db,
        SessionSelection(
            "prove",
            "2023-2024",
            (job,),
            extraction_ids=(alternative.id,),
            chapters=("900",),
            budget_years=(2024,),
        ),
    )
    assert report.rows and all(r.selection_status == "uavklart_treff" for r in report.rows)
    assert all(r.observation.budget_year is None for r in report.rows)
    assert all(s.complete_total_kr is None for s in report.sums)


def test_known_id_outside_selected_inventory_is_an_explicit_global_deviation(tmp_path):
    from copy import deepcopy

    from lxml import etree

    db = tmp_path / "known.sqlite"
    selected = synchronize(db)
    root = etree.fromstring(RAW)
    listing = root.find("{http://data.stortinget.no}stortingsvedtak_liste")
    another = deepcopy(listing[0])
    for node in another.iter():
        if node.text:
            node.text = node.text.replace("40033944", "49999999")
    listing.append(another)
    synchronize(db, etree.tostring(root), tick=1)
    report = build_session_report(
        db, SessionSelection("prove", "2024-2025", (selected,), chapters=("999",))
    )
    assert {r.decision_id for r in report.inventory} == {"40033944", "49999999"}
    assert (
        "kjent_id_utenfor_valgt_inventar"
        in next(r for r in report.inventory if r.decision_id == "49999999").issues
    )
    assert not report.coverage.source_complete


def test_first_empty_full_response_is_not_complete_professional_or_source_coverage(tmp_path):
    from lxml import etree

    db = tmp_path / "empty.sqlite"
    root = etree.fromstring(RAW)
    listing = root.find("{http://data.stortinget.no}stortingsvedtak_liste")
    for node in list(listing):
        listing.remove(node)
    job = synchronize(db, etree.tostring(root))
    report = build_session_report(db, SessionSelection("prove", "2024-2025", (job,)))
    assert report.inventory == ()
    assert not report.coverage.source_complete and not report.coverage.professional_complete
    assert any("tomt_inventar" in issue for issue in report.issues)
