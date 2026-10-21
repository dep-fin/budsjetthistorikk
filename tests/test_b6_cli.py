"""CLI → gjenåpnet SQLite → kontrollert Excel, med filvern."""

import hashlib
import json
import sys
from dataclasses import asdict
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_b2 import RAW, synchronize

from analyse.__main__ import main
from analyse.session import SessionSelection, selection_from_dict
from analyse.session_excel import export_session_excel
from innlesning import extract_revision, read_sync


def test_cli_roundtrip_exports_exact_selected_inventory_and_source_text(
    tmp_path, monkeypatch, capsys
):
    db = tmp_path / "cli.sqlite"
    job = synchronize(db)
    extraction = extract_revision(db, read_sync(db, job).entries[0].revision_id)
    selection = SessionSelection(
        "prove", "2024-2025", (job,), extraction_ids=(extraction.id,), chapters=("42", "3051")
    )
    spec = tmp_path / "selection.json"
    spec.write_text(json.dumps(asdict(selection)), encoding="utf-8")
    output = tmp_path / "cli.xlsx"
    before = hashlib.sha256(db.read_bytes()).hexdigest()
    monkeypatch.setattr(
        sys,
        "argv",
        ["analyse", "--database", str(db), "--sesjonsutvalg", str(spec), "--excel", str(output)],
    )
    assert main() == 0
    message = json.loads(capsys.readouterr().out)
    assert message["avstemming"] == "ikke utført" and message["observasjoner"] == "3"
    book = load_workbook(output)
    rows = list(book["Beløpsobservasjoner"].iter_rows(min_row=2, values_only=True))
    assert {r[0] for r in rows} == {
        o.id for o in extraction.observations if o.chapter in {"42", "3051"}
    }
    assert "Vedtaksfulltekst" in book.sheetnames and "Historikk" in book.sheetnames
    assert book["Avstemming"]["A2"].value == "ikke utført"
    book.close()
    assert hashlib.sha256(db.read_bytes()).hexdigest() == before


@pytest.mark.parametrize("target", ["database", "input", "existing", "source"])
def test_session_export_never_overwrites_database_input_source_or_existing_output(
    tmp_path: Path, target: str
):
    db = tmp_path / "guard.sqlite"
    job = synchronize(db)
    selection = SessionSelection("prove", "2024-2025", (job,))
    spec = tmp_path / "selection.json"
    spec.write_text(json.dumps(asdict(selection)), encoding="utf-8")
    output = {"database": db, "input": spec}.get(target, tmp_path / "existing.xlsx")
    if target in {"existing", "source"}:
        output.write_bytes(RAW if target == "source" else b"Bevar meg")
    before = output.read_bytes()
    with pytest.raises(ValueError):
        export_session_excel(db, selection, output, input_path=spec)
    assert output.read_bytes() == before


@pytest.mark.parametrize(
    "edit",
    [
        {"inventory_job_ids": "a"},
        {"budget_years": [True]},
        {"chapters": ["01"]},
        {"reconciliations": [{}]},
        {"unknown": 1},
        {"session": "2024-2026"},
    ],
)
def test_invalid_json_contract_is_rejected_before_export(edit):
    data = {"dataset": "prove", "session": "2024-2025", "inventory_job_ids": ["job"]}
    data.update(edit)
    with pytest.raises((ValueError, TypeError)):
        selection_from_dict(data)


def test_unextracted_fulltext_keeps_every_character_and_formula_like_title_is_text(tmp_path):
    from lxml import etree

    db = tmp_path / "fulltext.sqlite"
    root = etree.fromstring(RAW)
    node = root.find("{http://data.stortinget.no}stortingsvedtak_liste")[0]
    title = '=HYPERLINK("https://example.invalid","Bare tekst")'
    node.find("{http://data.stortinget.no}stortingsvedtak_tittel").text = title
    fulltext = "<div><p>" + ("æøå økonomisk fullmakt\n" * 4000) + "</p></div>"
    node.find("{http://data.stortinget.no}stortingsvedtak_tekst").text = fulltext
    job = synchronize(db, etree.tostring(root))
    output = tmp_path / "fulltext.xlsx"
    export_session_excel(db, SessionSelection("prove", "2024-2025", (job,)), output)
    book = load_workbook(output)
    text = "".join(
        row[5] for row in book["Inventarfulltekst"].iter_rows(min_row=2, values_only=True)
    )
    assert text == fulltext
    cell = next(c for row in book["Kildeinventar"] for c in row if c.value == title)
    assert cell.data_type == "s"
    assert not any(c.data_type == "f" for sheet in book for row in sheet for c in row)
    book.close()
