"""B5-ferskhet, deldekning og effektive verdier helt ut i B6-Excel."""

import hashlib
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_b5_controls import FIXTURES, prepared
from test_b5_register import edit_rows

from analyse.controlled import read_controlled, reprocess_controlled
from analyse.history import HistorySelection
from analyse.session import SessionSelection, build_session_report
from analyse.session_excel import export_session_excel
from analyse.session_reconciliation import SourceExpectation
from innlesning.controls import (
    activate_basis,
    define_control_basis,
    export_controls,
    import_controls,
    store_evidence,
)
from innlesning.observations import extract_revision
from innlesning.sync import read_revision


def approved_case(db: Path, directory: Path):
    extraction, register, basis = prepared(db)
    target = next(
        o for o in extraction.observations if o.post == "70" and o.meaning == "etter_belop"
    )
    evidence = store_evidence(
        db,
        FIXTURES / "referat-5538.xml",
        version="refs-202526-03-26",
        locator="/table/tgroup/tbody/row[3]/entry[3]/A",
        role="kontrollkilde",
        url="https://data.stortinget.no/eksport/publikasjon?publikasjonid=refs-202526-03-26",
        parent_sha256="4eed3a96b705a6bc80f9f7776efbde1c3dab07f36b898ed4722617ff59ea5b56",
    )
    path = directory / "approve.csv"
    export_controls(
        db,
        (extraction.id,),
        path,
        register_snapshot=register.id,
        control_basis=basis,
        observation_ids=(target.id,),
    )
    edit_rows(
        path,
        lambda rows: rows[0].update(
            status="godkjent",
            godkjente_felter="amount_kr",
            godkjent_bruk="analysebelop",
            korrigert_belop_kr="1615833000",
            bevis_id=evidence.id,
            kontrollor="Prøvekontrollør",
            begrunnelse="Eksakt referatcelle",
        ),
    )
    decision = import_controls(db, path)
    assert not decision.errors
    partial = reprocess_controlled(
        db,
        (extraction.id,),
        register.id,
        basis,
        selection=HistorySelection("prove", (extraction.id,), 2026, "5538", "70"),
    )
    assert len(partial.rows) < len(extraction.observations)
    job = read_revision(db, extraction.source_revision_id).current_observation_id
    expected = SourceExpectation(
        target.id,
        None,
        1_615_833_000,
        extraction.artifact.sha256,
        target.pointer.locator,
        target.pointer.amount_field,
        "https://data.stortinget.no/eksport/publikasjon?publikasjonid=refs-202526-03-26",
        "Fullt tall i referatet",
    )
    selection = SessionSelection(
        "prove",
        "2025-2026",
        (job,),
        extraction_ids=(extraction.id,),
        controlled_result_ids=(partial.id,),
        source_expectations=(expected,),
    )
    return extraction, register, basis, target, partial, selection


def test_corrected_direct_decision_and_partial_b5_coverage_survive_reopen_excel(
    tmp_path: Path,
) -> None:
    db = tmp_path / "direct.sqlite"
    extraction, register, basis, target, partial, selection = approved_case(db, tmp_path)
    before = hashlib.sha256(db.read_bytes()).hexdigest()
    report = build_session_report(db, selection)
    row = next(r for r in report.rows if r.observation.id == target.id)
    assert row.observation.raw_value == "I 615 833" and row.observation.amount_kr is None
    assert row.corrected_amount_kr == row.approved_amount_kr == 1_615_833_000
    assert row.controlled_result_ids == (partial.id,)
    assert report.source_checks[0].status == "samsvar"
    assert any("mangler_valgt_kontrollresultat" in r.issues for r in report.rows)
    assert not report.coverage.professional_complete
    output = tmp_path / "direct.xlsx"
    export_session_excel(db, selection, output)
    book = load_workbook(output)
    page = book["Beløpsobservasjoner"]
    headers = [c.value for c in page[1]]
    cells = next(row for row in page.iter_rows(min_row=2) if row[0].value == target.id)
    assert cells[headers.index("Gjeldende godkjent beløp kr")].value == 1_615_833_000
    assert cells[headers.index("Råbeløp kr")].value is None
    assert cells[headers.index("Avgjørelse")].value == row.controlled.decision_id
    assert not any(c.data_type == "f" for sheet in book for values in sheet for c in values)
    book.close()
    assert read_controlled(db, partial.id).rows == partial.rows
    assert hashlib.sha256(db.read_bytes()).hexdigest() == before


def test_filter_keeps_raw_b5_and_chain_problems_in_global_report_and_excel(tmp_path: Path):
    db = tmp_path / "filtered-problems.sqlite"
    extraction, register, basis, target, partial, selection = approved_case(db, tmp_path)
    selection = replace(selection, posts=("99",), source_expectations=())
    output = tmp_path / "filtered-problems.xlsx"
    report = export_session_excel(db, selection, output)
    assert not report.rows
    assert {
        "ugyldig_belop",
        "manglende_start",
        "mangler_valgt_kontrollresultat",
        "representasjonsavvik_bevart",
    } <= set(report.issues)
    unresolved = next(r for r in report.unfiltered_rows if r.observation.id == target.id)
    assert unresolved.selection_status == "utenfor"
    assert unresolved.observation.raw_value == "I 615 833"
    assert "ugyldig_belop" in unresolved.issues
    book = load_workbook(output)
    try:
        values = list(book["Avvik"].values)
        records = [dict(zip(values[0], row, strict=True)) for row in values[1:]]
        raw_error = next(
            row
            for row in records
            if row["Grunnlag"] == target.id and row["Avvik"] == "ugyldig_belop"
        )
        assert raw_error["Filterstatus"] == "utenfor"
        assert raw_error["Kapittel"] == "5538" and raw_error["Post"] == "70"
        assert raw_error["Råverdi"] == "I 615 833"
    finally:
        book.close()


@pytest.mark.parametrize("status", ["avvist", "til_kontroll"])
def test_new_reject_or_review_blocks_selected_older_result_without_losing_historical_amount(
    tmp_path: Path,
    status: str,
) -> None:
    db = tmp_path / "newer.sqlite"
    extraction, register, basis, target, partial, selection = approved_case(db, tmp_path)
    path = tmp_path / "newer.csv"
    export_controls(
        db,
        (extraction.id,),
        path,
        register_snapshot=register.id,
        control_basis=basis,
        observation_ids=(target.id,),
    )
    edit_rows(
        path,
        lambda rows: rows[0].update(
            status=status, kontrollor="Kontrollør 2", begrunnelse="Ny vurdering"
        ),
    )
    assert not import_controls(db, path).errors
    report = build_session_report(db, selection)
    row = next(r for r in report.rows if r.observation.id == target.id)
    assert row.historical_effective_amount_kr == 1_615_833_000
    assert row.approved_amount_kr is None and row.freshness == "utdatert"
    assert report.source_checks[0].status == "ikke_vurderbar"


@pytest.mark.parametrize("change", ["control", "interpretation", "register"])
def test_b6_does_not_revive_approval_after_activation_a_b_a(tmp_path: Path, change: str) -> None:
    from innlesning.registers import export_register, import_register

    db = tmp_path / "aba.sqlite"
    extraction, register, basis, target, partial, selection = approved_case(db, tmp_path)
    if change == "control":
        other_basis = define_control_basis(db, "Annet kontrollgrunnlag")
        activate_basis(db, (extraction.id,), register.id, other_basis)
    elif change == "interpretation":
        other = extract_revision(
            db,
            extraction.source_revision_id,
            interpretation=replace(extraction.interpretation, reason="Annen uttrykkelig tolkning"),
        )
        activate_basis(db, (other.id,), register.id, basis)
    else:
        path = tmp_path / "register.csv"
        export_register(db, register.id, path)
        edit_rows(path, lambda rows: rows[1].update(begrunnelse="Annet registergrunnlag"))
        other_register = import_register(db, path).snapshot_id
        other = extract_revision(
            db,
            extraction.source_revision_id,
            interpretation=replace(extraction.interpretation, register_snapshot=other_register),
        )
        activate_basis(db, (other.id,), other_register, basis)
    activate_basis(db, (extraction.id,), register.id, basis)
    report = build_session_report(db, selection)
    assert next(r for r in report.rows if r.observation.id == target.id).approved_amount_kr is None
    assert report.controlled_results[0].freshness == "utdatert"


def test_zero_observation_control_context_remains_stale_after_a_b_a_and_fulltext_is_exported(
    tmp_path: Path,
) -> None:
    db = tmp_path / "text.sqlite"
    extraction, register, basis = prepared(db, "40035429")
    assert not extraction.observations
    result = reprocess_controlled(db, (extraction.id,), register.id, basis)
    another = define_control_basis(db, "Et annet grunnlag for tekst")
    activate_basis(db, (extraction.id,), register.id, another)
    activate_basis(db, (extraction.id,), register.id, basis)
    job = read_revision(db, extraction.source_revision_id).current_observation_id
    selection = SessionSelection(
        "prove",
        "2024-2025",
        (job,),
        extraction_ids=(extraction.id,),
        controlled_result_ids=(result.id,),
    )
    report = build_session_report(db, selection)
    assert report.controlled_results[0].freshness == "utdatert"
    assert any("kontrollresultat_utdatert" in i for i in report.issues)
    output = tmp_path / "text.xlsx"
    export_session_excel(db, selection, output)
    book = load_workbook(output)
    assert "fullmaktsramme" in [c.value for row in book["Tekstopplysninger"] for c in row]
    book.close()


@pytest.mark.parametrize("transition", ["pending", "source_a_b_a"])
def test_source_attempt_or_return_never_revives_old_approved_values(
    tmp_path: Path, transition: str
):
    import json

    from innlesning import sync_saved
    from innlesning.sync import stage_sync

    db = tmp_path / "source.sqlite"
    extraction, register, basis, target, partial, selection = approved_case(db, tmp_path)
    directory = Path(__file__).parent / "fixtures/b3"
    raw = (directory / "40036289.xml").read_bytes()
    meta = json.loads((directory / "40036289.json").read_text(encoding="utf-8"))
    meta["retrieved_utc"] = "2026-10-21T12:00:00+00:00"
    if transition == "pending":
        stage_sync(db, raw, meta, historical_sample=True, decision_id="40036289")
    else:
        changed = raw.replace(b"I 615 833", b"I 615 834")
        meta.update(sha256=hashlib.sha256(changed).hexdigest(), bytes=len(changed))
        sync_saved(db, changed, meta, historical_sample=True, decision_id="40036289")
        meta.update(
            retrieved_utc="2026-10-21T12:01:00+00:00",
            sha256=hashlib.sha256(raw).hexdigest(),
            bytes=len(raw),
        )
        returned = sync_saved(db, raw, meta, historical_sample=True, decision_id="40036289")
        selection = replace(selection, inventory_job_ids=(returned,))
    report = build_session_report(db, selection)
    row = next(r for r in report.rows if r.observation.id == target.id)
    assert row.approved_amount_kr is None and row.historical_effective_amount_kr == 1_615_833_000
    assert row.freshness == ("ikke_vurderbar" if transition == "pending" else "utdatert")


@pytest.mark.parametrize("amount,hidden_post", [(0, False), (0, True), (10, True)])
def test_approved_values_and_filtered_subtotals_keep_zero_and_global_control_gaps(
    tmp_path: Path, amount: int, hidden_post: bool
):
    from test_b4 import synthetic

    from analyse.session_reconciliation import AmountScope, Reconciliation, ReferenceValue
    from innlesning import read_source, read_sync, sync_saved
    from innlesning.observations import Interpretation
    from innlesning.registers import PostEntry, export_register, import_register, seed_register

    db = tmp_path / "zero.sqlite"
    original = synthetic(
        db,
        "<table><tr><td>Kap.</td><td>Post</td><td>Formål</td><td>Kroner</td></tr>"
        "<tr><td>41</td><td></td><td>Stortinget</td><td></td></tr>"
        f"<tr><td></td><td>1</td><td>Drift</td><td>{amount}</td></tr>"
        + ("<tr><td></td><td>2</td><td>Uavklart post</td><td>5</td></tr>" if hidden_post else "")
        + "</table>",
        year=2027,
    )
    first_job = read_revision(db, original.source_revision_id).current_observation_id
    metadata = dict(read_sync(db, first_job).source_metadata)
    metadata.update(session_coverage="complete", retrieved_utc="2026-10-20T15:01:00+00:00")
    job = sync_saved(
        db, read_source(db, original.artifact.sha256), metadata, historical_sample=True
    )
    register = seed_register(db)
    catalog = tmp_path / "register.csv"
    export_register(db, register.id, catalog)
    edit_rows(
        catalog,
        lambda rows: rows.append(
            {
                "radtype": "post",
                **{
                    k: str(v)
                    for k, v in PostEntry(
                        "41", "1", 2027, 2027, "Syntetisk kilde", "table[1]/tr[3]", "Kun test"
                    )
                    .to_row()
                    .items()
                },
            }
        ),
    )
    register_id = import_register(db, catalog).snapshot_id
    extraction = extract_revision(
        db,
        original.source_revision_id,
        interpretation=Interpretation(
            adapter_version="storting-belop-v2", register_snapshot=register_id
        ),
    )
    basis = define_control_basis(db, "Syntetisk prøve av uttrykkelig null")
    activate_basis(db, (extraction.id,), register_id, basis)
    proof = store_evidence(
        db,
        FIXTURES / "referat-5538.xml",
        version="syntetisk-nullkontroll",
        locator="/table",
        role="kontrollkilde",
        url="https://example.invalid/synthetic",
    )
    path = tmp_path / "zero.csv"
    export_controls(db, (extraction.id,), path, register_snapshot=register_id, control_basis=basis)
    edit_rows(
        path,
        lambda rows: rows[0].update(
            status="godkjent",
            godkjente_felter="amount_kr,organization,catalogue",
            godkjent_bruk="kjedegrunnlag",
            bevis_id=proof.id,
            kontrollor="Test",
            begrunnelse="Syntetisk null",
        ),
    )
    assert not import_controls(db, path).errors
    result = reprocess_controlled(
        db,
        (extraction.id,),
        register_id,
        basis,
        selection=HistorySelection("prove", (extraction.id,), 2027, "41", "1"),
    )
    assert result.approved_chain_balance == amount
    scope = AmountScope(2027, "2027-01-01", "2027-12-31", "arets_bevilgning")
    reference = ReferenceValue(
        "41",
        "1",
        amount,
        False,
        "https://example.invalid/synthetic",
        extraction.artifact.sha256,
        "table[1]/tr[3]",
        "Syntetisk null",
    )
    check = Reconciliation(
        "Årsbevilgning for syntetisk prøvekonto", scope, scope, (), (reference,), (result.id,)
    )
    selection = SessionSelection(
        "prove",
        "2023-2024",
        (job,),
        extraction_ids=(extraction.id,),
        controlled_result_ids=(result.id,),
        history_calculation_ids=(result.technical_calculation_id,),
        reconciliations=(check,),
        posts=("1",) if hidden_post else (),
    )
    report = build_session_report(db, selection)
    assert report.coverage.source_complete
    assert len(report.rows) == 1
    assert report.rows[0].approved_amount_kr == amount
    assert report.sums[0].approved_subtotal_kr == amount
    assert report.reconciliations[0].rows[0].actual_amount_kr == amount
    assert report.reconciliations[0].status == "samsvar"
    if hidden_post:
        assert report.sums[0].complete_total_kr is None
        assert "filtrert_delsum" in report.sums[0].issues
        assert "mangler_valgt_kontrollresultat" in report.issues
        outside = next(r for r in report.unfiltered_rows if r.observation.post == "2")
        assert outside.selection_status == "utenfor"
        assert "ikke_godkjent_analysebelop" in outside.issues
    else:
        assert report.sums[0].complete_total_kr == amount
    output = tmp_path / "zero.xlsx"
    export_session_excel(db, selection, output)
    book = load_workbook(output)
    assert book["Beløpsobservasjoner"].cell(2, 22).value == amount
    if hidden_post:
        rows = list(book["Summer"].values)
        assert dict(zip(rows[0], rows[1], strict=True))["Sumomfang"].startswith("Delsum")
    book.close()
