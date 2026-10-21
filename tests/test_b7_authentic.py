"""Hele B7-løpet på bevarte autentiske objekter i syntetiske fullistekonvolutter.

Originalfixturer endres ikke. Konvoluttene er testinventarer, ikke komplette ekte sesjoner.
"""

import hashlib
from copy import deepcopy
from datetime import UTC, datetime
from pathlib import Path

from lxml import etree
from openpyxl import load_workbook
from test_b7 import FakeGitHub

from analyse.session_serialization import report_from_json
from drift import Config, run_daily
from drift.store import read_frozen_report
from innlesning.organization_register import REGISTER_SNAPSHOT

NS = "{http://data.stortinget.no}"


class PreservedObjects:
    def __init__(self) -> None:
        self.documents: dict[str, etree._Element] = {}
        self.manifest: dict[str, str] = {}
        for identifier in (
            "40032850",
            "40033025",
            "40033719",
            "40034064",
            "40034226",
            "40034287",
            "40034048",
            "40034229",
            "40036289",
        ):
            folder = "b6" if identifier in ("40034048", "40034229") else "b3"
            raw = (Path(__file__).parent / f"fixtures/{folder}/{identifier}.xml").read_bytes()
            self.manifest[identifier] = hashlib.sha256(raw).hexdigest()
            root = etree.fromstring(raw)
            session = root.findtext(NS + "sesjon_id")
            if session not in self.documents:
                self.documents[session] = deepcopy(root)
                self.documents[session].find(NS + "stortingsvedtak_liste").clear()
            target = self.documents[session].find(NS + "stortingsvedtak_liste")
            for node in root.find(NS + "stortingsvedtak_liste"):
                target.append(deepcopy(node))
        self.calls: list[str] = []

    def sessions(self, timeout: int):
        root = etree.Element(NS + "sesjoner_oversikt", nsmap={None: NS[1:-1]})
        current = etree.SubElement(root, NS + "innevaerende_sesjon")
        listing = etree.SubElement(root, NS + "sesjoner_liste")
        for value in sorted(self.documents):
            year = int(value[:4])
            item = etree.SubElement(listing, NS + "sesjon")
            for key, text in (
                ("id", value),
                ("fra", f"{year}-10-01T00:00:00"),
                ("til", f"{year + 1}-09-30T23:59:59"),
            ):
                etree.SubElement(item, NS + key).text = text
            if value == max(self.documents):
                for node in item:
                    current.append(deepcopy(node))
        raw = etree.tostring(root)
        return raw, {
            "url": "https://data.stortinget.no/eksport/sesjoner",
            "status": 200,
            "bytes": len(raw),
            "sha256": hashlib.sha256(raw).hexdigest(),
            "retrieved_utc": "2026-10-21T05:00:00+00:00",
        }

    def session(self, session: str, timeout: int):
        self.calls.append(session)
        raw = etree.tostring(self.documents[session])
        return raw, {
            "url": f"https://data.stortinget.no/eksport/stortingsvedtak?sesjonid={session}",
            "status": 200,
            "bytes": len(raw),
            "sha256": hashlib.sha256(raw).hexdigest(),
            "retrieved_utc": "2026-10-21T05:00:00+00:00",
            "fixture_kind": "synthetic_test",
            "session_coverage": "complete",
            "preserved_object_fixtures": self.manifest,
        }


def test_all_types_missing_phrase_raw_i_and_three_conflicts_reach_frozen_excel(
    tmp_path: Path,
) -> None:
    source = PreservedObjects()
    config = Config(
        "authentic",
        tmp_path / "authentic.sqlite",
        tmp_path / "reports",
        REGISTER_SNAPSHOT,
        dataset="prove",
        first_session="2023-2024",
    )
    github = FakeGitHub()
    github.fail = True  # Levering er en separat grense; unngå tusenvis av testissues.
    result = run_daily(
        config, source=source, github=github, now=datetime(2026, 10, 21, 6, tzinfo=UTC)
    )
    assert result["outcome"] == "full" and len(result["reports"]) == 3
    assert source.calls == ["2023-2024", "2024-2025", "2025-2026"]
    report = report_from_json(read_frozen_report(config.database, result["reports"][1]["id"])[0])
    inventory = {i.decision_id: i for i in report.inventory}
    assert inventory["40033719"].api_type == inventory["40034287"].api_type == "ANNET"
    phrase_rows = [r for r in report.rows if r.observation.api_id in {"40034048", "40034229"}]
    assert len({r.observation.step_id for r in phrase_rows}) == 76
    raw_i = next(
        r
        for r in report.rows
        if r.observation.api_id == "40036289"
        and r.observation.meaning == "etter_belop"
        and r.observation.post == "70"
    )
    assert raw_i.observation.raw_value == "I 615 833" and raw_i.observation.amount_kr is None
    for chapter, post in (("1860", "1"), ("4860", "1"), ("5999", "90")):
        affected = [
            r for r in report.rows if (r.observation.chapter, r.observation.post) == (chapter, post)
        ]
        assert any("kjedekonflikt" in r.issues or "etter_konflikt" in r.issues for r in affected)
        assert all(r.approved_amount_kr is None for r in affected)
    assert not report.coverage.professional_complete
    assert len(report.supplemental_coverage) == 2
    book = load_workbook(result["reports"][1]["excel"], read_only=True)
    assert book["Sesjonsgrunnlag"].cell(2, 2).value == report.id
    assert "40034048" in {
        str(r[1]) for r in book["Kildeinventar"].iter_rows(min_row=2, values_only=True)
    }
    book.close()
