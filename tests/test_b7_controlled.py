"""B7 bruker manuell kontroll kun når valgt aktivt grunnlag passer eksakt."""

from datetime import UTC, datetime
from pathlib import Path

import pytest
from test_b4 import synthetic
from test_b5_register import edit_rows
from test_b7 import FakeGitHub, FakeSource, session_list

from analyse.session_serialization import report_from_json
from drift import Config, run_daily
from drift.store import read_frozen_report
from innlesning import extract_revision, read_source, read_sync
from innlesning.controls import (
    activate_basis,
    define_control_basis,
    export_controls,
    import_controls,
    read_activation,
    store_evidence,
)
from innlesning.observations import Interpretation
from innlesning.registers import PostEntry, export_register, import_register, seed_register
from innlesning.sync import read_revision


def clean_case(tmp_path: Path):
    db = tmp_path / "clean.sqlite"
    original = synthetic(
        db,
        "<table><tr><td>Kap.</td><td>Post</td><td>Formål</td><td>Kroner</td></tr>"
        "<tr><td>41</td><td></td><td>Stortinget</td><td></td></tr>"
        "<tr><td></td><td>1</td><td>Drift</td><td>0</td></tr></table>",
        year=2027,
    )
    register = seed_register(db)
    path = tmp_path / "catalog.csv"
    export_register(db, register.id, path)
    edit_rows(
        path,
        lambda rows: rows.append(
            {
                "radtype": "post",
                **{
                    k: str(v)
                    for k, v in PostEntry(
                        "41", "1", 2027, 2027, "Syntetisk kilde", "table[1]/tr[3]", "Kun test"
                    )
                    .to_row()
                    .items()
                },
            }
        ),
    )
    snapshot = import_register(db, path).snapshot_id
    extraction = extract_revision(
        db,
        original.source_revision_id,
        interpretation=Interpretation(
            adapter_version="storting-belop-v2", register_snapshot=snapshot
        ),
    )
    basis = define_control_basis(db, "Syntetisk eksplisitt nullkontroll")
    activate_basis(db, (extraction.id,), snapshot, basis)
    proof = store_evidence(
        db,
        Path(__file__).parent / "fixtures/b5/referat-5538.xml",
        version="syntetisk",
        locator="/table",
        role="kontrollkilde",
        url="https://example.invalid/synthetic",
    )
    control = tmp_path / "approve.csv"
    export_controls(db, (extraction.id,), control, register_snapshot=snapshot, control_basis=basis)
    edit_rows(
        control,
        lambda rows: rows[0].update(
            status="godkjent",
            godkjente_felter="amount_kr,organization,catalogue",
            godkjent_bruk="kjedegrunnlag",
            bevis_id=proof.id,
            kontrollor="Test",
            begrunnelse="Syntetisk null",
        ),
    )
    assert not import_controls(db, control).errors
    job = read_revision(db, extraction.source_revision_id).current_observation_id
    meta = dict(read_sync(db, job).source_metadata)
    meta.update(session_coverage="complete", retrieved_utc="2026-10-21T05:00:00+00:00")
    raw = read_source(db, extraction.artifact.sha256)
    session = read_sync(db, job).session

    class Source(FakeSource):
        def sessions(self, timeout: int):
            import hashlib

            result = session_list(session, False)
            return result, {
                "url": "https://data.stortinget.no/eksport/sesjoner",
                "status": 200,
                "bytes": len(result),
                "sha256": hashlib.sha256(result).hexdigest(),
                "retrieved_utc": "2026-10-21T05:00:00+00:00",
            }

        def session(self, selected: str, timeout: int):
            self.calls.append(selected)
            return raw, dict(meta)

    config = Config(
        "clean",
        db,
        tmp_path / "reports",
        snapshot,
        dataset="prove",
        first_session=session,
        control_basis=basis,
    )
    return config, Source(), extraction, control


def test_unchanged_approved_zero_gives_only_reports_and_no_issue(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    config, source, extraction, _ = clean_case(tmp_path)
    github = FakeGitHub()
    first = run_daily(
        config, source=source, github=github, now=datetime(2026, 10, 21, 6, tzinfo=UTC)
    )
    assert first["outcome"] == "full"
    assert first["source_changes"] == {"ny": 0, "rettelse": 0, "uendret": 1, "mangler": 0}
    assert first["control_needs"] == 0 and first["delivery"] == "ingen_avvik"
    assert github.creates == 0
    frozen = report_from_json(read_frozen_report(config.database, first["reports"][0]["id"])[0])
    assert frozen.rows[0].approved_amount_kr == 0 and not frozen.coverage.professional_complete
    assert read_activation(config.database, extraction.id)["generation"] == 1
    book = load_workbook(first["reports"][0]["excel"])
    page = book["Beløpsobservasjoner"]
    headers = [c.value for c in page[1]]
    assert page.cell(2, headers.index("Gjeldende godkjent beløp kr") + 1).value == 0
    assert page.cell(2, 1).value == frozen.rows[0].observation.id
    assert not any(c.data_type == "f" for sheet in book for row in sheet for c in row)
    book.close()


@pytest.mark.parametrize("github_available", [True, False])
def test_run_json_failure_is_durable_notified_and_reuses_issue_across_retries_and_days(
    tmp_path: Path, monkeypatch, github_available: bool
) -> None:
    import json
    import os

    from drift import read_daily
    from drift.notifications import marker
    from drift.store import pending_notifications

    config, source, _, _ = clean_case(tmp_path)
    github = FakeGitHub()
    monkeypatch.delenv(config.token_environment, raising=False)
    rename, link = os.rename, os.link
    attempts = []

    def reject_json(operation):
        def publish(source, destination, *args, **kwargs):
            if Path(destination).suffix == ".json":
                attempts.append(destination)
                raise OSError("synthetic-secret")
            return operation(source, destination, *args, **kwargs)

        return publish

    with monkeypatch.context() as fault:
        fault.setattr(os, "rename", reject_json(rename))
        fault.setattr(os, "link", reject_json(link))
        first = run_daily(
            config,
            source=source,
            github=github if github_available else None,
            now=datetime(2026, 10, 21, 6, tzinfo=UTC),
        )
        assert first["errors"] == ["kjorerapport_fil_feilet"]
        assert first["outcome"] == "partial" and "run_report" not in first
        assert first["control_needs"] == 1
        assert first["delivery"] == ("levert" if github_available else "venter")
        pending = pending_notifications(config.database)
        assert len(pending) == (0 if github_available else 1)
        if pending:
            first_marker = marker(pending[0]["key"])
        else:
            first_marker = str(github.issues[1]["body"]).splitlines()[0]
        assert read_daily(config.database, config.pipeline_id)[0]["result"] == first
        for instant in (
            datetime(2026, 10, 21, 7, tzinfo=UTC),
            datetime(2026, 10, 22, 6, tzinfo=UTC),
        ):
            repeated = run_daily(config, source=source, github=github, now=instant)
            assert repeated["errors"] == ["kjorerapport_fil_feilet"]
            assert repeated["control_needs"] == 1 and repeated["delivery"] == "levert"
            assert github.creates == 1 and not pending_notifications(config.database)
            assert first_marker in str(github.issues[1]["body"])
        assert len(attempts) == 3
        assert len(source.calls) == 2
        assert "synthetic-secret" not in json.dumps(first) + json.dumps(github.issues)
    recovered = run_daily(
        config, source=source, github=github, now=datetime(2026, 10, 22, 7, tzinfo=UTC)
    )
    assert recovered["outcome"] == "full" and recovered["errors"] == []
    assert recovered["control_needs"] == 1 and github.creates == 1
    assert Path(recovered["run_report"]).is_file() and len(source.calls) == 2


@pytest.mark.parametrize("transition", ["new_reject", "basis_a_b_a", "timeout"])
def test_new_rejection_or_activation_return_or_prestage_failure_blocks_old_values(
    tmp_path: Path, transition: str
) -> None:
    from test_b5_register import edit_rows

    config, source, extraction, control = clean_case(tmp_path)
    run_daily(config, source=source, github=FakeGitHub(), now=datetime(2026, 10, 21, 6, tzinfo=UTC))
    if transition == "new_reject":
        control = tmp_path / "reject.csv"
        export_controls(
            config.database,
            (extraction.id,),
            control,
            register_snapshot=config.register_snapshot,
            control_basis=config.control_basis,
        )
        edit_rows(
            control,
            lambda rows: rows[0].update(
                status="avvist",
                godkjente_felter="",
                godkjent_bruk="",
                bevis_id="",
                kontrollor="Test",
                begrunnelse="Ny avvisning",
            ),
        )
        assert not import_controls(config.database, control).errors
    elif transition == "basis_a_b_a":
        other = define_control_basis(config.database, "Annet kontrollgrunnlag")
        activate_basis(config.database, (extraction.id,), config.register_snapshot, other)
        activate_basis(
            config.database, (extraction.id,), config.register_snapshot, config.control_basis
        )
    else:
        # Nettgrensen feiler etter at dagens forsøk er bevart.
        def fail(selected: str, timeout: int):
            raise TimeoutError()

        source.session = fail
    result = run_daily(
        config, source=source, github=FakeGitHub(), now=datetime(2026, 10, 22, 6, tzinfo=UTC)
    )
    frozen = report_from_json(read_frozen_report(config.database, result["reports"][0]["id"])[0])
    assert frozen.rows[0].approved_amount_kr is None
    assert result["control_needs"] > 0


def test_source_a_b_a_returns_same_revision_without_reviving_old_approval(tmp_path: Path) -> None:
    import hashlib

    from lxml import etree

    config, source, extraction, _ = clean_case(tmp_path)
    original_fetch = source.session
    run_daily(config, source=source, github=FakeGitHub(), now=datetime(2026, 10, 21, 6, tzinfo=UTC))

    def changed(selected: str, timeout: int):
        raw, meta = original_fetch(selected, timeout)
        root = etree.fromstring(raw)
        node = root.find(".//{http://data.stortinget.no}stortingsvedtak_tekst")
        node.text = node.text.replace("<td>0</td>", "<td>1</td>")
        raw = etree.tostring(root)
        meta.update(
            sha256=hashlib.sha256(raw).hexdigest(),
            bytes=len(raw),
            retrieved_utc="2026-10-22T05:00:00+00:00",
        )
        return raw, meta

    source.session = changed
    middle = run_daily(
        config, source=source, github=FakeGitHub(), now=datetime(2026, 10, 22, 6, tzinfo=UTC)
    )
    assert middle["source_changes"]["rettelse"] == 1
    assert read_activation(config.database, extraction.id)["extraction_id"] == extraction.id

    def returned(selected: str, timeout: int):
        raw, meta = original_fetch(selected, timeout)
        meta["retrieved_utc"] = "2026-10-23T05:00:00+00:00"
        return raw, meta

    source.session = returned
    final = run_daily(
        config, source=source, github=FakeGitHub(), now=datetime(2026, 10, 23, 6, tzinfo=UTC)
    )
    frozen = report_from_json(read_frozen_report(config.database, final["reports"][0]["id"])[0])
    assert frozen.rows[0].extraction_id == extraction.id
    assert frozen.rows[0].approved_amount_kr is None
    assert frozen.rows[0].freshness == "utdatert"
    assert final["source_changes"]["rettelse"] == 1
