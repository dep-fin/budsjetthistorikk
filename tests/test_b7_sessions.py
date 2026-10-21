"""Sesjonsskifte og sene ID-er/rettelser uten dato-, type- eller frasefilter."""

import hashlib
from copy import deepcopy
from datetime import UTC, datetime
from pathlib import Path

import pytest
from lxml import etree
from test_b2 import RAW, metadata
from test_b7 import FakeGitHub, FakeSource, session_list

from analyse.session_serialization import report_from_json
from drift import Config, run_daily
from drift.sessions import select_sessions
from drift.store import read_frozen_report
from innlesning import read_sync, stage_sync
from innlesning.organization_register import REGISTER_SNAPSHOT

NS = "{http://data.stortinget.no}"


def test_next_session_and_late_old_id_and_correction_are_all_observed(tmp_path: Path):
    class Rollover(FakeSource):
        later = False

        def sessions(self, timeout: int):
            raw = session_list()
            return raw, {
                "url": "https://data.stortinget.no/eksport/sesjoner",
                "status": 200,
                "bytes": len(raw),
                "sha256": hashlib.sha256(raw).hexdigest(),
                "retrieved_utc": "2027-10-21T05:00:00+00:00",
            }

        def session(self, session: str, timeout: int):
            self.calls.append(session)
            root = etree.fromstring(RAW)
            root.find(NS + "sesjon_id").text = session
            listing = root.find(NS + "stortingsvedtak_liste")
            if self.later and session == "2026-2027":
                text = listing[0].find(NS + "stortingsvedtak_tekst")
                text.text = text.text.replace("10 478 000", "10 479 000")
                old = deepcopy(listing[0])
                old.find(NS + "id").text = "40033945"
                link = old.find(NS + "stortingsvedtak_lenke_url")
                link.text = link.text.replace("40033944", "40033945")
                listing.append(old)
            raw = etree.tostring(root)
            meta = metadata(raw)
            meta.update(
                url=f"https://data.stortinget.no/eksport/stortingsvedtak?sesjonid={session}",
                retrieved_utc="2027-10-21T05:00:00+00:00"
                if self.later
                else "2026-10-21T05:00:00+00:00",
            )
            return raw, meta

    config = Config(
        "roll",
        tmp_path / "roll.sqlite",
        tmp_path / "reports",
        REGISTER_SNAPSHOT,
        dataset="prove",
        first_session="2026-2027",
    )
    source, github = Rollover(), FakeGitHub()
    first = run_daily(
        config, source=source, github=github, now=datetime(2026, 10, 21, 6, tzinfo=UTC)
    )
    assert len(first["reports"]) == 1 and source.calls == ["2026-2027"]
    source.later = True
    second = run_daily(
        config, source=source, github=github, now=datetime(2027, 10, 21, 6, tzinfo=UTC)
    )
    assert second["source_changes"] == {"ny": 2, "rettelse": 1, "uendret": 0, "mangler": 0}
    assert source.calls == ["2026-2027", "2026-2027", "2027-2028"]
    report = report_from_json(read_frozen_report(config.database, second["reports"][1]["id"])[0])
    assert {(i.session, i.decision_id) for i in report.inventory} == {
        ("2026-2027", "40033944"),
        ("2026-2027", "40033945"),
        ("2027-2028", "40033944"),
    }
    assert len(report.supplemental_coverage) == 1


def test_pending_b2_is_resumed_before_even_session_list_http(tmp_path: Path):
    db = tmp_path / "pending.sqlite"
    staged = stage_sync(db, RAW, metadata(RAW), historical_sample=True)

    class Inspecting(FakeSource):
        def sessions(self, timeout: int):
            assert read_sync(db, staged).status == "fullfort"
            return super().sessions(timeout)

    config = Config(
        "pending",
        db,
        tmp_path / "reports",
        REGISTER_SNAPSHOT,
        dataset="prove",
        first_session="2024-2025",
    )
    assert (
        run_daily(
            config,
            source=Inspecting(),
            github=FakeGitHub(),
            now=datetime(2026, 10, 21, 6, tzinfo=UTC),
        )["outcome"]
        == "full"
    )


def test_omitted_started_session_stays_in_http_excel_and_stable_catalogue_fault(tmp_path: Path):
    class OmittingCatalogue(FakeSource):
        omitted = False

        def sessions(self, timeout: int):
            raw = session_list("2027-2028", not self.omitted)
            root = etree.fromstring(raw)
            future = deepcopy(root.find(NS + "innevaerende_sesjon"))
            future.tag = NS + "sesjon"
            future.find(NS + "id").text = "2028-2029"
            future.find(NS + "fra").text = "2028-10-01T00:00:00"
            future.find(NS + "til").text = "2029-09-30T23:59:59"
            root.find(NS + "sesjoner_liste").append(future)
            raw = etree.tostring(root)
            return raw, {
                "url": "https://data.stortinget.no/eksport/sesjoner",
                "status": 200,
                "bytes": len(raw),
                "sha256": hashlib.sha256(raw).hexdigest(),
                "retrieved_utc": "2027-10-21T05:00:00+00:00",
            }

        def session(self, session: str, timeout: int):
            self.calls.append(session)
            root = etree.fromstring(RAW)
            root.find(NS + "sesjon_id").text = session
            if self.omitted and session == "2026-2027":
                listing = root.find(NS + "stortingsvedtak_liste")
                text = listing[0].find(NS + "stortingsvedtak_tekst")
                text.text = text.text.replace("10 478 000", "10 479 000")
                late = deepcopy(listing[0])
                late.find(NS + "id").text = "40033945"
                link = late.find(NS + "stortingsvedtak_lenke_url")
                link.text = link.text.replace("40033944", "40033945")
                listing.append(late)
            raw = etree.tostring(root)
            meta = metadata(raw)
            meta.update(
                url=f"https://data.stortinget.no/eksport/stortingsvedtak?sesjonid={session}",
                retrieved_utc="2027-10-22T05:00:00+00:00"
                if self.omitted
                else "2027-10-21T05:00:00+00:00",
            )
            return raw, meta

    config = Config(
        "omitted",
        tmp_path / "omitted.sqlite",
        tmp_path / "reports",
        REGISTER_SNAPSHOT,
        dataset="prove",
        first_session="2026-2027",
    )
    source, github = OmittingCatalogue(), FakeGitHub()
    first = run_daily(
        config, source=source, github=github, now=datetime(2027, 10, 21, 6, tzinfo=UTC)
    )
    assert first["outcome"] == "full" and len(first["reports"]) == 2
    assert set(first["jobs"]) == {"2026-2027", "2027-2028"}
    source.omitted = True
    second = run_daily(
        config, source=source, github=github, now=datetime(2027, 10, 22, 6, tzinfo=UTC)
    )
    assert source.calls == ["2026-2027", "2027-2028", "2026-2027", "2027-2028"]
    assert second["source_changes"] == {"ny": 1, "rettelse": 1, "uendret": 1, "mangler": 0}
    assert second["outcome"] == "partial"
    assert "sesjon_mangler_i_katalog:2026-2027" in second["errors"]
    assert len(second["reports"]) == 2
    report = report_from_json(read_frozen_report(config.database, second["reports"][1]["id"])[0])
    assert {(i.session, i.decision_id) for i in report.inventory} == {
        ("2026-2027", "40033944"),
        ("2026-2027", "40033945"),
        ("2027-2028", "40033944"),
    }
    from openpyxl import load_workbook

    book = load_workbook(second["reports"][1]["excel"])
    assert book["Kildeinventar"].max_row == 4
    book.close()
    fault_issues = [
        n for n, issue in github.issues.items() if "sesjon_mangler_i_katalog" in issue["title"]
    ]
    assert len(fault_issues) == 1
    for instant in (datetime(2027, 10, 22, 7, tzinfo=UTC), datetime(2027, 10, 23, 6, tzinfo=UTC)):
        repeated = run_daily(config, source=source, github=github, now=instant)
        assert repeated["outcome"] == "partial"
        assert "sesjon_mangler_i_katalog:2026-2027" in repeated["errors"]
    assert [
        n for n, issue in github.issues.items() if "sesjon_mangler_i_katalog" in issue["title"]
    ] == fault_issues


@pytest.mark.parametrize("kind", ["nil", "duplicate", "doctype", "wrong_namespace"])
def test_session_list_rejects_ambiguous_or_unsafe_shape(kind: str):
    raw = session_list()
    if kind == "nil":
        raw = raw.replace(
            b"<sesjoner_liste>",
            b"<sesjoner_liste xmlns:i='http://www.w3.org/2001/XMLSchema-instance' i:nil='true'>",
        )
    elif kind == "duplicate":
        raw = raw.replace(
            b"</sesjoner_liste>",
            b"<sesjon><id>2026-2027</id><fra>2026-10-01T00:00:00</fra><til>2027-09-30T23:59:59</til></sesjon></sesjoner_liste>",
        )
    elif kind == "doctype":
        raw = b"<!DOCTYPE sesjoner_oversikt>" + raw
    else:
        raw = raw.replace(b"http://data.stortinget.no", b"https://untrusted.invalid")
    with pytest.raises(ValueError):
        select_sessions(raw, datetime(2026, 10, 21, 6, tzinfo=UTC), "2026-2027")
