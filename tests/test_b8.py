"""B8s offentlige grenser; autentiske originaler og merkede syntetiske feil."""

import hashlib
import json
import re
import sqlite3
import subprocess
import sys
from dataclasses import asdict
from pathlib import Path

import pytest

from innlesning.saksdokumenter import import_case_response, read_case_source

FIXTURES = Path(__file__).parent / "fixtures"


def case_source(case_id="100927"):
    raw = (FIXTURES / "b8" / f"sak-{case_id}.xml").read_bytes()
    metadata = json.loads((FIXTURES / "b8" / f"sak-{case_id}.json").read_text("utf-8"))
    return raw, metadata


def test_authentic_case_keeps_nil_raw_labels_unique_references_and_case_relations(tmp_path):
    raw, metadata = case_source()
    db = tmp_path / "b8.sqlite"
    result = import_case_response(db, "100927", raw, metadata, dataset="prove", session="2024-2025")
    assert result.status == "observert"
    assert read_case_source(db, result.id) == raw
    assert result.content is not None
    assert len(result.content.publications) == 5
    prop, _, innst, referat, voting = result.content.publications
    assert (prop.kind, prop.number, prop.series, prop.session) == (
        "proposisjon",
        "16",
        "S",
        "2024-2025",
    )
    assert prop.fields["eksport_id"].nil == "true"
    assert prop.identity_status == "identifisert" and not prop.issues
    assert prop.fields["lenke_url"].text.startswith("http://")
    assert prop.url.startswith("https://")
    assert (innst.kind, innst.number) == ("innstilling", "81")
    assert referat.fields["eksport_id"] == voting.fields["eksport_id"]
    assert referat.id != voting.id and referat.url != voting.url
    assert result.content.relations[0].fields["relatert_sak_id"].text == "100889"


def seed_decision(db):
    from innlesning import extract_revision, read_sync, sync_saved
    from innlesning.observations import Interpretation

    directory = FIXTURES / "b6"
    raw = (directory / "40034048.xml").read_bytes()
    meta = json.loads((directory / "40034048.json").read_text("utf-8"))
    job = sync_saved(db, raw, meta, historical_sample=True, decision_id="40034048")
    revision = read_sync(db, job).entries[0].revision_id
    extraction = extract_revision(
        db, revision, interpretation=Interpretation(adapter_version="storting-belop-v2")
    )
    return job, revision, extraction


def test_real_source_to_public_clis_sqlite_excel_and_frozen2_without_historical_start(tmp_path):
    from openpyxl import load_workbook

    from analyse.saksdokumenter import report_from_json
    from analyse.session import SessionSelection
    from innlesning.saksdokumenter import read_enrichment

    db = tmp_path / "new.sqlite"
    job, revision, extraction = seed_decision(db)
    command = subprocess.run(
        [
            sys.executable,
            "-m",
            "innlesning",
            "saksdokumenter",
            "--database",
            str(db),
            "--revisjon",
            revision,
            "--response",
            str(FIXTURES / "b8/sak-100927.xml"),
            "--metadata",
            str(FIXTURES / "b8/sak-100927.json"),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert command.returncode == 0, command.stdout + command.stderr
    identifier = json.loads(command.stdout)["saksberikelse"]
    enrichment = read_enrichment(db, identifier)
    assert enrichment.formal_number == "210" and enrichment.api_id == "40034048"
    assert [(k.post, k.words, k.status) for k in enrichment.keywords] == [
        ("1", (), "ingen_stikkord_funnet"),
        ("22", ("kan overføres",), "funnet"),
    ]
    selection = tmp_path / "selection.json"
    selection.write_text(
        json.dumps(
            {
                "base": asdict(
                    SessionSelection("prove", "2024-2025", (job,), extraction_ids=(extraction.id,))
                ),
                "enrichment_ids": [identifier],
            }
        ),
        encoding="utf-8",
    )
    excel, frozen = tmp_path / "result.xlsx", tmp_path / "result.json"
    command = subprocess.run(
        [
            sys.executable,
            "-m",
            "analyse",
            "--database",
            str(db),
            "--saksutvalg",
            str(selection),
            "--excel",
            str(excel),
            "--json",
            str(frozen),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert command.returncode == 0, command.stdout + command.stderr
    report = report_from_json(frozen.read_text("utf-8"))
    assert json.loads(frozen.read_text("utf-8"))["format"] == "b6-frozen-2"
    assert len(report.base.rows) == 6
    assert {
        (r.observation.post, r.observation.meaning): r.observation.amount_kr
        for r in report.base.rows
    } == {
        ("1", "endring"): 5_796_000,
        ("1", "for_belop"): 160_243_000,
        ("1", "etter_belop"): 166_039_000,
        ("22", "endring"): -3_000_000,
        ("22", "for_belop"): 32_185_000,
        ("22", "etter_belop"): 29_185_000,
    }
    book = load_workbook(excel)
    try:
        sheet = book["Beløpsobservasjoner"]
        headers = {c.value: c.column for c in sheet[1]}
        assert headers["Stikkord"] == headers["Post"] + 1
        assert {sheet.cell(n, headers["Formelt vedtak"]).value for n in range(2, 8)} == {
            "210 (2024-2025)"
        }
        assert {sheet.cell(n, headers["Teknisk API-ID"]).value for n in range(2, 8)} == {"40034048"}
        assert sheet.cell(2, headers["Innstilling 1"]).hyperlink.target.startswith(
            "https://www.stortinget.no/"
        )
        assert sheet.cell(2, headers["Proposisjon 1"]).value == "Prop. 16 S (2024-2025)"
        assert {"Saksdokumenter", "Dokumentrelasjoner", "Dokumentavvik"} <= set(book.sheetnames)
        assert book["Historikk"].max_row == 3
    finally:
        book.close()


def altered_case(raw, metadata, *, day, old_text=None, new_text=None):
    """Syntetiske mutasjoner; originalfixturen endres aldri."""
    value = re.sub(rb"2026-10-22T", f"2026-10-{day:02}T".encode(), raw)
    if old_text is not None:
        value = value.replace(old_text.encode(), new_text.encode())
    meta = {
        **metadata,
        "retrieved_utc": f"2026-10-{day:02}T12:00:00+00:00",
        "fixture_kind": "synthetic_mutation",
        "sha256": hashlib.sha256(value).hexdigest(),
        "bytes": len(value),
    }
    return value, meta


def test_case_repeats_corrections_return_to_old_content_and_old_timestamp_are_distinct(tmp_path):
    from innlesning.saksdokumenter import enrich_revision, read_case_context, read_enrichment

    db = tmp_path / "history.sqlite"
    job, revision, extraction = seed_decision(db)
    raw, meta = case_source()
    a = import_case_response(db, "100927", raw, meta, dataset="prove", session="2024-2025")
    old = enrich_revision(db, revision, a.id)
    repeat = import_case_response(
        db, "100927", *altered_case(raw, meta, day=23), dataset="prove", session="2024-2025"
    )
    assert repeat.content.revision_id == a.content.revision_id
    assert (
        read_case_context(db, dataset="prove", session="2024-2025", case_id="100927").generation
        == 1
    )
    b = import_case_response(
        db,
        "100927",
        *altered_case(raw, meta, day=24, old_text="Innst. 81 S", new_text="Innst. 82 S"),
        dataset="prove",
        session="2024-2025",
    )
    assert b.content.revision_id != a.content.revision_id
    returned = import_case_response(
        db, "100927", *altered_case(raw, meta, day=25), dataset="prove", session="2024-2025"
    )
    context = read_case_context(db, dataset="prove", session="2024-2025", case_id="100927")
    assert context.generation == 3 and context.current_attempt_id == returned.id
    assert (a.generation, repeat.generation, b.generation, returned.generation) == (1, 1, 2, 3)
    assert returned.content.revision_id == a.content.revision_id
    assert read_enrichment(db, old.id) == old
    replay = import_case_response(db, "100927", raw, meta, dataset="prove", session="2024-2025")
    assert replay.status == "eldre_kilde" and read_case_source(db, replay.id) == raw
    assert read_case_context(db, dataset="prove", session="2024-2025", case_id="100927") == context
    from analyse.saksdokumenter import SaksSelection, build_saks_report
    from analyse.session import SessionSelection

    report = build_saks_report(
        db,
        SaksSelection(
            SessionSelection("prove", "2024-2025", (job,), extraction_ids=(extraction.id,)),
            (old.id,),
        ),
    )
    assert report.contexts[0].source_freshness == "historisk_observasjon_samme_innhold"


@pytest.mark.parametrize("failure", ["timeout", "parse", "http", "provenance"])
def test_failed_attempt_preserves_received_bytes_and_never_blocks_amount_export(tmp_path, failure):
    from analyse.saksdokumenter import SaksSelection, build_saks_report, export_built_saks_excel
    from analyse.session import SessionSelection
    from innlesning.saksdokumenter import enrich_revision, fetch_case_attempt, read_case_context

    db = tmp_path / "failure.sqlite"
    job, revision, extraction = seed_decision(db)
    raw, meta = case_source()
    good = import_case_response(db, "100927", raw, meta, dataset="prove", session="2024-2025")
    enriched = enrich_revision(db, revision, good.id)
    if failure == "timeout":

        def fetch(case_id, *, timeout):
            assert (
                read_case_context(
                    db, dataset="prove", session="2024-2025", case_id=case_id
                ).latest_status
                == "venter"
            )
            assert timeout == 7
            raise TimeoutError("Syntetisk timeout")

        failed = fetch_case_attempt(
            db, "100927", dataset="prove", session="2024-2025", timeout=7, fetch=fetch
        )
        assert failed.source_sha256 is None
        with pytest.raises(ValueError, match="ingen mottatt"):
            read_case_source(db, failed.id)
    else:
        received, metadata = altered_case(raw, meta, day=26)
        if failure == "parse":
            received = b"<broken>"
            metadata.update(sha256=hashlib.sha256(received).hexdigest(), bytes=len(received))
        elif failure == "http":
            metadata["status"] = 503
        else:
            metadata["sha256"] = "0" * 64
        failed = import_case_response(
            db, "100927", received, metadata, dataset="prove", session="2024-2025"
        )
        assert read_case_source(db, failed.id) == received
    assert failed.status == "feilet"
    selection = SaksSelection(
        SessionSelection("prove", "2024-2025", (job,), extraction_ids=(extraction.id,)),
        (enriched.id,),
    )
    report = build_saks_report(db, selection)
    assert report.contexts[0].source_freshness == "historiske_referanser_etter_feilet"
    assert len(report.base.rows) == 6
    export_built_saks_excel(db, report, tmp_path / "failure.xlsx")


@pytest.mark.parametrize("replacement", ["bytes", "metadata"])
def test_pending_case_attempt_locks_first_response_and_metadata_during_concurrent_save(
    tmp_path, monkeypatch, replacement
):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event, current_thread

    from innlesning.saksdokumenter import (
        begin_case_attempt,
        read_case_attempt,
        read_case_context,
        save_case_response,
    )

    db = tmp_path / "response-binding.sqlite"
    attempt_id = begin_case_attempt(db, "100927", dataset="prove", session="2024-2025")
    raw, metadata = case_source()
    staged, resume = Event(), Event()
    connect = sqlite3.connect

    class Connection(sqlite3.Connection):
        def __exit__(self, *args):
            result = super().__exit__(*args)
            # The SQLite port pauses only after committing received bytes, with no lock held.
            if current_thread().name.startswith("first-response") and not staged.is_set():
                row = self.execute(
                    "SELECT sha256 FROM b8_attempts WHERE id=?", (attempt_id,)
                ).fetchone()
                if row and row[0]:
                    staged.set()
                    assert resume.wait(10), "The controlled test did not release the first save"
            return result

    def connected(*args, **kwargs):
        return connect(*args, **kwargs, factory=Connection)

    monkeypatch.setattr(sqlite3, "connect", connected)
    with ThreadPoolExecutor(max_workers=1, thread_name_prefix="first-response") as executor:
        first = executor.submit(save_case_response, db, attempt_id, raw, metadata)
        assert staged.wait(10), "The first response was not committed"
        try:
            assert read_case_source(db, attempt_id) == raw
            other_raw, other_metadata = altered_case(
                raw, metadata, day=23, old_text="Innst. 81 S", new_text="Innst. 82 S"
            )
            if replacement == "metadata":
                other_raw = raw
                other_metadata = {**metadata, "retrieved_utc": "2026-10-23T12:00:00+00:00"}
            with pytest.raises(ValueError, match="bundet"):
                save_case_response(db, attempt_id, other_raw, other_metadata)
            pending = read_case_attempt(db, attempt_id)
            assert pending.status == "venter"
            assert json.loads(pending.metadata_json) == metadata
        finally:
            resume.set()
        result = first.result(timeout=10)
    assert result.status == "observert"
    assert result.source_sha256 == result.content.source_sha256 == metadata["sha256"]
    assert json.loads(result.metadata_json) == metadata
    assert read_case_source(db, result.id) == raw
    assert (
        read_case_context(
            db, dataset="prove", session="2024-2025", case_id="100927"
        ).current_attempt_id
        == result.id
    )


def test_interrupted_case_save_resumes_only_the_identical_received_response(tmp_path, monkeypatch):
    from innlesning.saksdokumenter import begin_case_attempt, read_case_attempt, save_case_response

    db = tmp_path / "resume.sqlite"
    attempt_id = begin_case_attempt(db, "100927", dataset="prove", session="2024-2025")
    raw, metadata = case_source()
    connect = sqlite3.connect
    interrupted = False

    class Connection(sqlite3.Connection):
        def __exit__(self, *args):
            nonlocal interrupted
            result = super().__exit__(*args)
            if not interrupted:
                row = self.execute(
                    "SELECT sha256 FROM b8_attempts WHERE id=?", (attempt_id,)
                ).fetchone()
                if row and row[0]:
                    interrupted = True
                    raise InterruptedError("Syntetisk avbrudd etter varig rålagring")
            return result

    def connected(*args, **kwargs):
        return connect(*args, **kwargs, factory=Connection)

    monkeypatch.setattr(sqlite3, "connect", connected)
    with pytest.raises(InterruptedError):
        save_case_response(db, attempt_id, raw, metadata)
    assert read_case_attempt(db, attempt_id).status == "venter"
    with pytest.raises(ValueError, match="bundet"):
        save_case_response(db, attempt_id, raw, {**metadata, "headers": {}})
    resumed = save_case_response(db, attempt_id, raw, metadata)
    assert resumed.status == "observert"
    assert read_case_source(db, attempt_id) == raw
    assert json.loads(resumed.metadata_json) == metadata


@pytest.mark.parametrize("body_state", ["oserror", "timeout", "partial", "empty", "complete"])
def test_http_error_body_read_always_closes_and_finishes_a_truthful_enrichable_attempt(
    tmp_path, body_state
):
    from http.client import IncompleteRead
    from urllib.error import HTTPError

    from analyse.saksdokumenter import SaksSelection, build_saks_report
    from analyse.session import SessionSelection
    from innlesning.ingestion import MAX_BYTES
    from innlesning.saksdokumenter import (
        case_url,
        enrich_revision,
        fetch_case_attempt,
        read_case_attempt,
    )

    db = tmp_path / "http-error.sqlite"
    job, revision, extraction = seed_decision(db)
    closed = False
    received = b"<partial" if body_state == "partial" else b"Service unavailable"

    class Body:
        def read(self, size):
            assert size == MAX_BYTES + 1
            if body_state == "oserror":
                raise OSError("Syntetisk lesefeil")
            if body_state == "timeout":
                raise TimeoutError("Syntetisk timeout ved lesing")
            if body_state == "partial":
                raise IncompleteRead(received, 100)
            return b"" if body_state == "empty" else received

        def close(self):
            nonlocal closed
            closed = True

    def fetch(case_id, *, timeout):
        raise HTTPError(case_url(case_id), 503, "Service unavailable", {"Retry-After": "1"}, Body())

    failed = fetch_case_attempt(db, "100927", dataset="prove", session="2024-2025", fetch=fetch)
    assert closed
    assert failed.status == "feilet"
    assert read_case_attempt(db, failed.id) == failed
    metadata = json.loads(failed.metadata_json)
    assert (metadata["status"], metadata["url"], metadata["headers"]) == (
        503,
        case_url("100927"),
        {"Retry-After": "1"},
    )
    if body_state in {"partial", "complete"}:
        assert read_case_source(db, failed.id) == received
        assert metadata["sha256"] == failed.source_sha256 == hashlib.sha256(received).hexdigest()
        assert metadata["bytes"] == len(received)
        assert metadata["body_status"] == body_state
    else:
        assert failed.source_sha256 is None
        assert "sha256" not in metadata
        with pytest.raises(ValueError, match="ingen mottatt"):
            read_case_source(db, failed.id)
    if body_state in {"oserror", "timeout", "partial"}:
        assert "body_read_error" in metadata
        assert "lesing" in failed.reason
    enriched = enrich_revision(db, revision, failed.id)
    assert any("saksforsok_feilet" in issue for issue in enriched.issues)
    report = build_saks_report(
        db,
        SaksSelection(
            SessionSelection("prove", "2024-2025", (job,), extraction_ids=(extraction.id,)),
            (enriched.id,),
        ),
    )
    assert len(report.base.rows) == 6


@pytest.mark.parametrize(
    "mutation,expected",
    [
        ("missing", "mangler_publikasjon_referanse_liste"),
        ("nil", "mangler_publikasjon_referanse_liste"),
        ("unknown", "uavklart_dokumentidentitet"),
        ("conflict", "motstrid_mellom_etikett_og_type"),
        ("many", "flere_innstillinger"),
        ("unsafe", "mangler_eller_utrygg_dokumentlenke"),
    ],
)
def test_document_uncertainty_and_global_anomalies_survive_filters(tmp_path, mutation, expected):
    from lxml import etree

    from analyse.saksdokumenter import SaksSelection, build_saks_report
    from analyse.session import SessionSelection
    from innlesning.ingestion import NS
    from innlesning.saksdokumenter import enrich_revision

    db = tmp_path / "uncertain.sqlite"
    job, revision, extraction = seed_decision(db)
    raw, meta = case_source()
    root = etree.fromstring(raw)
    pubs = root.find(NS + "publikasjon_referanse_liste")
    if mutation == "missing":
        root.remove(pubs)
    elif mutation == "nil":
        pubs.clear()
        pubs.set("{http://www.w3.org/2001/XMLSchema-instance}nil", "true")
    elif mutation == "unknown":
        pubs[0].find(NS + "lenke_tekst").text = "Ukjent regjeringsdokument"
        pubs[0].find(NS + "undertype").text = "ukjent"
    elif mutation == "conflict":
        pubs[0].find(NS + "type").text = "innstilling"
    elif mutation == "many":
        from copy import deepcopy

        pubs.append(deepcopy(pubs[2]))
    else:
        pubs[0].find(NS + "lenke_url").text = "javascript:alert(1)"
    received = etree.tostring(root)
    meta.update(
        sha256=hashlib.sha256(received).hexdigest(),
        bytes=len(received),
        fixture_kind="synthetic_mutation",
    )
    result = import_case_response(
        db, "100927", received, meta, dataset="prove", session="2024-2025"
    )
    enriched = enrich_revision(db, revision, result.id)
    report = build_saks_report(
        db,
        SaksSelection(
            SessionSelection(
                "prove", "2024-2025", (job,), extraction_ids=(extraction.id,), posts=("99",)
            ),
            (enriched.id,),
        ),
    )
    assert not report.base.rows
    assert any(expected in issue for issue in report.issues)
    assert len(report.base.unfiltered_rows) == 6


def test_frozen2_reexports_without_database_reads_and_frozen1_stays_readable(tmp_path):
    from analyse.saksdokumenter import (
        SaksSelection,
        build_saks_report,
        export_built_saks_excel,
        report_from_json,
        report_to_json,
    )
    from analyse.session import SessionSelection
    from analyse.session_serialization import report_from_json as old_read
    from analyse.session_serialization import report_to_json as old_write
    from innlesning.saksdokumenter import enrich_revision

    db = tmp_path / "offline.sqlite"
    job, revision, extraction = seed_decision(db)
    raw, meta = case_source()
    attempt = import_case_response(db, "100927", raw, meta, dataset="prove", session="2024-2025")
    enriched = enrich_revision(db, revision, attempt.id)
    report = build_saks_report(
        db,
        SaksSelection(
            SessionSelection("prove", "2024-2025", (job,), extraction_ids=(extraction.id,)),
            (enriched.id,),
        ),
    )
    assert old_read(old_write(report.base)) == report.base
    saved = report_from_json(report_to_json(report))
    db.write_bytes(b"database intentionally unavailable after freezing")
    export_built_saks_excel(db, saved, tmp_path / "offline.xlsx")
    with pytest.raises(ValueError, match="nytt filnavn"):
        export_built_saks_excel(db, saved, tmp_path / "offline.xlsx")
    for field in ("format", "type", "content"):
        data = json.loads(report_to_json(saved))
        if field == "format":
            data["format"] = "future"
        elif field == "type":
            data["report"]["contexts"][0]["case_context"]["generation"] = True
        else:
            data["report"]["enrichments"][0]["formal_number"] = "211"
        with pytest.raises(ValueError):
            report_from_json(json.dumps(data))


def test_related_case_never_inherits_a_post_link_and_missing_selection_is_explicit(tmp_path):
    from analyse.saksdokumenter import SaksSelection, build_saks_report
    from analyse.session import SessionSelection
    from innlesning.saksdokumenter import enrich_revision

    db = tmp_path / "related.sqlite"
    job, revision, extraction = seed_decision(db)
    other = import_case_response(
        db, "100889", *case_source("100889"), dataset="prove", session="2024-2025"
    )
    assert other.content.publications[2].number == "95"
    with pytest.raises(ValueError, match="hører ikke"):
        enrich_revision(db, revision, other.id)
    report = build_saks_report(
        db,
        SaksSelection(
            SessionSelection("prove", "2024-2025", (job,), extraction_ids=(extraction.id,))
        ),
    )
    assert report.contexts[0].enrichment_id is None
    assert report.contexts[0].source_freshness == "ikke_valgt"
    assert "saksberikelse_ikke_valgt" in report.contexts[0].issues
    assert len(report.base.rows) == 6


@pytest.mark.parametrize(
    "marker,words,status",
    [
        (
            "kan overføres</span>, <em>overslagsbevilgning</em>",
            ("kan overføres", "overslagsbevilgning"),
            "funnet",
        ),
        ("ukjent fullmakt</span>", (), "uavklart"),
    ],
)
def test_separate_stikkord_interpretation_keeps_multiple_words_order_and_unknowns(
    tmp_path, marker, words, status
):
    from innlesning import extract_revision, read_sync, sync_saved
    from innlesning.observations import Interpretation
    from innlesning.saksdokumenter import enrich_revision

    db = tmp_path / "keywords.sqlite"
    raw = (FIXTURES / "b6/40034048.xml").read_bytes()
    meta = json.loads((FIXTURES / "b6/40034048.json").read_text("utf-8"))
    raw = raw.replace(
        b"kan overf\xc3\xb8res&lt;/span&gt;",
        marker.replace("<", "&lt;").replace(">", "&gt;").encode(),
    )
    meta.update(
        sha256=hashlib.sha256(raw).hexdigest(), bytes=len(raw), fixture_kind="synthetic_test"
    )
    job = sync_saved(db, raw, meta, historical_sample=True, decision_id="40034048")
    revision = read_sync(db, job).entries[0].revision_id
    extraction = extract_revision(
        db, revision, interpretation=Interpretation(adapter_version="storting-belop-v2")
    )
    attempt = import_case_response(
        db, "100927", *case_source(), dataset="prove", session="2024-2025"
    )
    result = enrich_revision(db, revision, attempt.id)
    keyword = next(k for k in result.keywords if k.post == "22")
    assert (keyword.words, keyword.status) == (words, status)
    assert keyword.source_sha256 == meta["sha256"] and keyword.locator.endswith("/td[3]")
    assert (
        "fullmakt" in keyword.raw_text
        if status == "uavklart"
        else "overslagsbevilgning" in keyword.raw_text
    )
    assert len(extraction.observations) == 6


@pytest.mark.parametrize(
    "fragment,words",
    [
        ("</span>kan overføres", ()),
        ("kan overføres</span>, overslagsbevilgning", ("kan overføres",)),
        ("kan overføres</span>, ukjent fullmakt", ("kan overføres",)),
        ("kan overføres, overslagsbevilg</span>", ("kan overføres",)),
    ],
)
def test_stikkord_checks_text_outside_markers_and_exports_known_words_with_uncertainty(
    tmp_path, fragment, words
):
    from openpyxl import load_workbook

    from analyse.saksdokumenter import SaksSelection, build_saks_report, export_built_saks_excel
    from analyse.session import SessionSelection
    from innlesning import extract_revision, read_sync, sync_saved
    from innlesning.observations import Interpretation
    from innlesning.saksdokumenter import enrich_revision

    db = tmp_path / "unmarked.sqlite"
    raw = (FIXTURES / "b6/40034048.xml").read_bytes()
    meta = json.loads((FIXTURES / "b6/40034048.json").read_text("utf-8"))
    raw = raw.replace(
        "kan overføres&lt;/span&gt;".encode(),
        fragment.replace("<", "&lt;").replace(">", "&gt;").encode(),
    )
    meta.update(
        sha256=hashlib.sha256(raw).hexdigest(), bytes=len(raw), fixture_kind="synthetic_test"
    )
    job = sync_saved(db, raw, meta, historical_sample=True, decision_id="40034048")
    revision = read_sync(db, job).entries[0].revision_id
    extraction = extract_revision(
        db, revision, interpretation=Interpretation(adapter_version="storting-belop-v2")
    )
    attempt = import_case_response(
        db, "100927", *case_source(), dataset="prove", session="2024-2025"
    )
    enriched = enrich_revision(db, revision, attempt.id)
    keyword = next(k for k in enriched.keywords if k.post == "22")
    assert (keyword.words, keyword.status) == (words, "uavklart")
    assert keyword.source_sha256 == meta["sha256"]
    assert keyword.locator == "table[1]/tr[5]/td[3]"
    assert fragment in keyword.raw_html
    assert "kan overføres" in keyword.raw_text
    report = build_saks_report(
        db,
        SaksSelection(
            SessionSelection("prove", "2024-2025", (job,), extraction_ids=(extraction.id,)),
            (enriched.id,),
        ),
    )
    output = tmp_path / "unmarked.xlsx"
    export_built_saks_excel(db, report, output)
    book = load_workbook(output)
    try:
        sheet = book["Beløpsobservasjoner"]
        headers = {c.value: c.column for c in sheet[1]}
        displayed = [
            sheet.cell(n, headers["Stikkord"]).value
            for n in range(2, 8)
            if sheet.cell(n, headers["Post"]).value == "22"
        ]
        assert len(displayed) == 3
        assert all("Uavklart:" in value for value in displayed)
        assert all(all(word in value for word in words) for value in displayed)
    finally:
        book.close()


@pytest.mark.parametrize("historical", ["absent", "corrupt", "conflicting"])
def test_new_official_source_is_independent_of_historical_workbook_state(tmp_path, historical):
    from test_historical_excel import workbook_bytes

    from analyse.saksdokumenter import SaksSelection, build_saks_report
    from analyse.session import SessionSelection
    from innlesning.historical_excel import import_historical_workbook
    from innlesning.saksdokumenter import enrich_revision

    db = tmp_path / "independent.sqlite"
    if historical != "absent":
        path = tmp_path / "legacy.xlsx"
        path.write_bytes(
            b"invalid OOXML" if historical == "corrupt" else workbook_bytes(amount="-99999")
        )
        result = import_historical_workbook(db, path, budget_year=2024)
        assert result.status == ("feilet" if historical == "corrupt" else "lagret")
    job, revision, extraction = seed_decision(db)
    attempt = import_case_response(
        db, "100927", *case_source(), dataset="prove", session="2024-2025"
    )
    enriched = enrich_revision(db, revision, attempt.id)
    report = build_saks_report(
        db,
        SaksSelection(
            SessionSelection("prove", "2024-2025", (job,), extraction_ids=(extraction.id,)),
            (enriched.id,),
        ),
    )
    assert len(report.base.rows) == 6
    assert (
        next(
            r.observation.amount_kr
            for r in report.base.rows
            if r.observation.post == "22" and r.observation.meaning == "endring"
        )
        == -3_000_000
    )


def test_old_writers_preserve_b8_schema5_and_v1_v2_results_and_reject_future_schema(tmp_path):
    from test_b2 import synchronize

    from innlesning import extract_revision, read_extraction, read_revision
    from innlesning.observations import Interpretation
    from innlesning.saksdokumenter import enrich_revision, read_enrichment

    db = tmp_path / "compat.sqlite"
    _, revision, v2 = seed_decision(db)
    v1 = extract_revision(db, revision)
    old_decision = read_revision(db, revision).report
    attempt = import_case_response(
        db, "100927", *case_source(), dataset="prove", session="2024-2025"
    )
    enriched = enrich_revision(db, revision, attempt.id)
    synchronize(db)
    extract_revision(
        db, revision, interpretation=Interpretation(adapter_version="storting-belop-v2")
    )
    assert read_extraction(db, v1.id) == v1 and read_extraction(db, v2.id) == v2
    assert read_revision(db, revision).report == old_decision
    assert read_enrichment(db, enriched.id) == enriched
    with sqlite3.connect(db) as connection:
        assert connection.execute("PRAGMA user_version").fetchone()[0] == 5
        connection.execute("PRAGMA user_version = 6")
    with pytest.raises(ValueError, match="Ukjent databaseskjema"):
        synchronize(db)


def test_fullmakt211_is_preserved_without_an_extra_financial_event_and_many_docs_do_not_inflate(
    tmp_path,
):
    from copy import deepcopy

    from lxml import etree
    from openpyxl import load_workbook

    from analyse.saksdokumenter import SaksSelection, build_saks_report, export_built_saks_excel
    from analyse.session import SessionSelection
    from innlesning import extract_revision, read_sync, sync_saved
    from innlesning.ingestion import NS
    from innlesning.observations import Interpretation
    from innlesning.saksdokumenter import enrich_revision

    db = tmp_path / "fullmakt.sqlite"
    job, revision, extraction = seed_decision(db)
    raw = (FIXTURES / "b8/40034049.xml").read_bytes()
    meta = json.loads((FIXTURES / "b8/40034049.json").read_text("utf-8"))
    assert hashlib.sha256(raw).hexdigest() == meta["sha256"]
    authority_job = sync_saved(db, raw, meta, historical_sample=True, decision_id="40034049")
    authority_revision = read_sync(db, authority_job).entries[0].revision_id
    authority = extract_revision(
        db, authority_revision, interpretation=Interpretation(adapter_version="storting-belop-v2")
    )
    assert not authority.observations
    case_raw, case_meta = case_source()
    root = etree.fromstring(case_raw)
    pubs = root.find(NS + "publikasjon_referanse_liste")
    pubs.append(deepcopy(pubs[2]))
    case_raw = etree.tostring(root)
    case_meta.update(
        sha256=hashlib.sha256(case_raw).hexdigest(),
        bytes=len(case_raw),
        fixture_kind="synthetic_mutation",
    )
    attempt = import_case_response(
        db, "100927", case_raw, case_meta, dataset="prove", session="2024-2025"
    )
    selected = tuple(enrich_revision(db, i, attempt.id).id for i in (revision, authority_revision))
    report = build_saks_report(
        db,
        SaksSelection(
            SessionSelection(
                "prove",
                "2024-2025",
                (job, authority_job),
                extraction_ids=(extraction.id, authority.id),
            ),
            selected,
        ),
    )
    assert len(report.base.rows) == 6 and len(report.base.inventory) == 2
    assert {e.formal_number for e in report.enrichments} == {"210", "211"}
    assert all(r.observation.api_id == "40034048" for r in report.base.rows)
    assert any("flere_innstillinger" in i for i in report.issues)
    output = tmp_path / "many.xlsx"
    export_built_saks_excel(db, report, output)
    book = load_workbook(output)
    try:
        assert book["Beløpsobservasjoner"].max_row == 7
        history = book["Historikk"]
        assert history.max_row == 3
        history_headers = {c.value: c.column for c in history[1]}
        expected = {
            "1": (160_243_000, 5_796_000, 166_039_000),
            "22": (32_185_000, -3_000_000, 29_185_000),
        }
        for n in (2, 3):
            post = history.cell(n, history_headers["Post"]).value
            assert (
                tuple(
                    history.cell(n, history_headers[m + " / Rå kr"]).value
                    for m in ("for_belop", "endring", "etter_belop")
                )
                == expected[post]
            )
        assert "Innstilling 2" in [c.value for c in book["Historikk"][1]]
        assert any(
            "Stikkordsfullmakt" in str(c.value) for row in book["Vedtaksfulltekst"] for c in row
        )
    finally:
        book.close()


def test_conflicting_document_export_identity_and_case_session_are_unresolved(tmp_path):
    raw, meta = case_source()
    raw, meta = altered_case(
        raw, meta, day=27, old_text="inns-202425-081s", new_text="inns-202425-082s"
    )
    attempt = import_case_response(
        tmp_path / "conflict.sqlite", "100927", raw, meta, dataset="prove", session="2024-2025"
    )
    innst = attempt.content.publications[2]
    assert innst.identity_status == "motstrid"
    assert "motstrid_dokumentidentiteter" in innst.issues


@pytest.mark.parametrize(
    "url,status,issue",
    [
        (
            "http://www.regjeringen.no/id/PRP202420250017000DDDEPIS",
            "motstrid",
            "motstrid_dokumentidentiteter",
        ),
        (
            "http://www.regjeringen.no/id/PRP202320240016000DDDEPIS",
            "motstrid",
            "motstrid_dokumentidentiteter",
        ),
        (
            "http://www.regjeringen.no/id/PRP202420250016000DDDEPIL",
            "motstrid",
            "motstrid_dokumentidentiteter",
        ),
        ("//www.stortinget.no/no/inns-202425-016s/", "motstrid", "motstrid_dokumentidentiteter"),
        ("http://www.regjeringen.no/id/ukjent", "uavklart", "uavklart_lenkeidentitet"),
    ],
)
def test_proposition_label_does_not_hide_conflicting_or_unknown_link_identity(
    tmp_path, url, status, issue
):
    raw, meta = case_source()
    raw, meta = altered_case(
        raw,
        meta,
        day=27,
        old_text="http://www.regjeringen.no/id/PRP202420250016000DDDEPIS",
        new_text=url,
    )
    db = tmp_path / "prop-identity.sqlite"
    attempt = import_case_response(db, "100927", raw, meta, dataset="prove", session="2024-2025")
    prop = attempt.content.publications[0]
    assert prop.identity_status == status
    assert issue in prop.issues
    assert prop.fields["eksport_id"].nil == "true"
    assert prop.fields["lenke_url"].text == url
    assert prop.source_sha256 == meta["sha256"]
    assert prop.locator.endswith("publikasjon_referanse[1]")
    assert read_case_source(db, attempt.id) == raw


def test_case_failures_leave_existing_professional_decision_and_controlled_values_unchanged(
    tmp_path,
):
    from test_b5_register import edit_rows

    from analyse.controlled import read_controlled, reprocess_controlled
    from innlesning import extract_revision, read_extraction
    from innlesning.controls import (
        activate_basis,
        define_control_basis,
        export_controls,
        import_controls,
        read_decision,
        store_evidence,
    )
    from innlesning.observations import Interpretation
    from innlesning.registers import seed_register
    from innlesning.saksdokumenter import enrich_revision, fetch_case_attempt

    db = tmp_path / "approval.sqlite"
    _, revision, old = seed_decision(db)
    register = seed_register(db)
    extraction = extract_revision(
        db,
        revision,
        interpretation=Interpretation(
            adapter_version="storting-belop-v2", register_snapshot=register.id
        ),
    )
    basis = define_control_basis(db, "Eksplisitt prøvekontroll av enkeltbeløp")
    activate_basis(db, (extraction.id,), register.id, basis)
    proof = store_evidence(
        db,
        FIXTURES / "b8/sak-100927.xml",
        version="syntetisk kontroll",
        locator="/detaljert_sak",
        role="kontrollkilde",
        url="https://data.stortinget.no/eksport/sak?sakid=100927",
    )
    target = next(o for o in extraction.observations if o.post == "22" and o.meaning == "endring")
    control = tmp_path / "approve.csv"
    export_controls(
        db,
        (extraction.id,),
        control,
        register_snapshot=register.id,
        control_basis=basis,
        observation_ids=(target.id,),
    )
    edit_rows(
        control,
        lambda rows: rows[0].update(
            status="godkjent",
            godkjente_felter="amount_kr",
            godkjent_bruk="analysebelop",
            kontrollor="Test",
            begrunnelse="Prøvens kjente råbeløp",
            bevis_id=proof.id,
        ),
    )
    imported = import_controls(db, control)
    assert not imported.errors
    approved = read_decision(db, imported.new_ids[0])
    controlled = reprocess_controlled(db, (extraction.id,), register.id, basis)

    def timeout(*args, **kwargs):
        raise TimeoutError("Saksmetadata er utilgjengelig")

    failed = fetch_case_attempt(db, "100927", dataset="prove", session="2024-2025", fetch=timeout)
    enrich_revision(db, revision, failed.id)
    assert read_decision(db, approved.decision.id) == approved
    assert read_controlled(db, controlled.id) == controlled
    assert read_extraction(db, old.id) == old


@pytest.mark.parametrize("series", ["S", "L", "LS"])
def test_document_series_and_raw_label_are_preserved_without_losing_session(tmp_path, series):
    raw, meta = case_source()
    raw, meta = altered_case(
        raw,
        meta,
        day=28,
        old_text="Prop. 16 S (2024-2025)",
        new_text=f"Prop. 16 {series} (2024-2025)",
    )
    raw, meta = altered_case(
        raw,
        meta,
        day=28,
        old_text="PRP202420250016000DDDEPIS",
        new_text=f"PRP202420250016000DDDEPI{series}",
    )
    result = import_case_response(
        tmp_path / "series.sqlite", "100927", raw, meta, dataset="prove", session="2024-2025"
    )
    prop = result.content.publications[0]
    assert (prop.series, prop.session, prop.number) == (series, "2024-2025", "16")
    assert prop.fields["lenke_tekst"].text == f"Prop. 16 {series} (2024-2025)"
    assert prop.identity_status == "identifisert" and not prop.issues


def test_old_decision_revision_is_explicit_and_its_selected_amounts_are_not_replaced(tmp_path):
    from analyse.saksdokumenter import SaksSelection, build_saks_report
    from analyse.session import SessionSelection
    from innlesning import read_sync, sync_saved
    from innlesning.saksdokumenter import enrich_revision, read_enrichment

    db = tmp_path / "decision-correction.sqlite"
    job, revision, extraction = seed_decision(db)
    attempt = import_case_response(
        db, "100927", *case_source(), dataset="prove", session="2024-2025"
    )
    enrichment = enrich_revision(db, revision, attempt.id)
    raw = (FIXTURES / "b6/40034048.xml").read_bytes().replace(b"5 796 000", b"5 796 001")
    raw = raw.replace(b"2026-10-20T", b"2026-10-21T")
    meta = json.loads((FIXTURES / "b6/40034048.json").read_text("utf-8"))
    meta.update(
        retrieved_utc="2026-10-21T12:00:00+00:00",
        fixture_kind="synthetic_test",
        sha256=hashlib.sha256(raw).hexdigest(),
        bytes=len(raw),
    )
    newer = sync_saved(db, raw, meta, historical_sample=True, decision_id="40034048")
    assert read_sync(db, newer).entries[0].revision_id != revision
    report = build_saks_report(
        db,
        SaksSelection(
            SessionSelection("prove", "2024-2025", (job,), extraction_ids=(extraction.id,)),
            (enrichment.id,),
        ),
    )
    assert report.contexts[0].decision_freshness == "historisk_eller_ikke_vurderbar"
    assert read_enrichment(db, enrichment.id) == enrichment
    assert (
        next(
            r.observation.amount_kr
            for r in report.base.rows
            if r.observation.post == "1" and r.observation.meaning == "endring"
        )
        == 5_796_000
    )


@pytest.mark.parametrize(
    "url",
    [
        "javascript:alert(1)",
        "https://www.stortinget.no.evil.test/path",
        "https://user:pass@www.stortinget.no/path",
        "file:///C:/Windows/win.ini",
        "https://www.stortinget.no:445/a",
        "https://www.stortinget.no/\nfoo",
    ],
)
def test_document_urls_never_become_active_unsafe_hyperlinks(url):
    from innlesning.saksdokumenter import safe_document_url

    assert safe_document_url(url) is None
