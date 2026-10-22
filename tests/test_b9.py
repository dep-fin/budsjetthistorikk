"""B9 gjennom offentlig frosset rapport og faktisk gjenlest Excel, uten historikkbase."""

import gzip
import hashlib
import json
from copy import deepcopy
from dataclasses import asdict, replace
from datetime import datetime
from pathlib import Path

import pytest
from lxml import etree
from openpyxl import load_workbook
from test_b8 import case_source, seed_decision

from analyse.saksdokumenter import (
    SaksSelection,
    build_saks_report,
    export_built_saks_excel,
    report_from_json,
    report_to_json,
    verify_saks_excel,
)
from analyse.session import SessionSelection
from analyse.session_excel import export_built_session_excel
from innlesning.ingestion import NS
from innlesning.observations import identity
from innlesning.saksdokumenter import enrich_revision, import_case_response

HEADERS = [
    "Kapittelnummer",
    "Postnummer",
    "Kapittelnavn",
    "Postnavn",
    "Stikkord",
    "Dato",
    "Vedtaksnummer",
    "Innstillingsnummer",
    "Proposisjon",
    "Beløp før vedtak",
    "Vedtatt endring",
    "Beløp etter vedtak",
    "Kontrollstatus",
]


def source_report(tmp_path, **filters):
    db = tmp_path / "b9.sqlite"
    job, revision, extraction = seed_decision(db)
    attempt = import_case_response(
        db, "100927", *case_source(), dataset="prove", session="2024-2025"
    )
    enrichment = enrich_revision(db, revision, attempt.id)
    report = build_saks_report(
        db,
        SaksSelection(
            SessionSelection(
                "prove", "2024-2025", (job,), extraction_ids=(extraction.id,), **filters
            ),
            (enrichment.id,),
        ),
    )
    return db, report


def frozen_variant(report, rows, *, global_issues=()):
    """Merket syntetisk snapshot ved den offentlige frosset-v2-grensen."""
    base = replace(report.base, rows=tuple(rows), unfiltered_rows=tuple(rows), issues=global_issues)
    candidate = replace(report, base=base, id="")
    candidate = replace(candidate, id=identity("b6-frozen-2", asdict(candidate)))
    return report_from_json(report_to_json(candidate))


def test_210_main_view_has_exact_13_columns_source_names_dates_and_amounts(tmp_path):
    db, original = source_report(tmp_path)
    frozen = report_to_json(original)
    report = report_from_json(frozen)
    # Export from the old frozen contract with an unavailable database and no historical files.
    db.write_bytes(b"database unavailable after freezing")
    output = tmp_path / "b9.xlsx"
    export_built_saks_excel(db, report, output)
    verify_saks_excel(output, report)
    assert report_to_json(report) == frozen
    assert report.id == original.id
    book = load_workbook(output)
    try:
        sheet = book["Historikk"]
        assert book.sheetnames[0] == book.active.title == "Historikk"
        assert [c.value for c in sheet[1]] == HEADERS
        assert sheet.max_column == 13 and sheet.max_row == 3
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        assert rows[0][:12] == (
            "1429",
            "1",
            "Riksantikvaren",
            "Driftsutgifter",
            "Ingen stikkord funnet",
            datetime(2024, 12, 16),
            "210 (2024-2025)",
            "Innst. 81 S (2024-2025)",
            "Prop. 16 S (2024-2025)",
            160_243_000,
            5_796_000,
            166_039_000,
        )
        assert rows[1][:12] == (
            "1429",
            "22",
            "Riksantikvaren",
            "Flerårige prosjekter kulturminneforvaltning",
            "kan overføres",
            datetime(2024, 12, 16),
            "210 (2024-2025)",
            "Innst. 81 S (2024-2025)",
            "Prop. 16 S (2024-2025)",
            32_185_000,
            -3_000_000,
            29_185_000,
        )
        assert sheet["F2"].is_date and sheet["F2"].number_format == "dd.mm.yyyy"
        assert all(sheet.cell(2, c).data_type == "n" for c in (10, 11, 12))
        assert all(sheet.cell(2, c).number_format == '#,##0" kr"' for c in (10, 11, 12))
        assert sheet["H2"].hyperlink.target.startswith("https://www.stortinget.no/")
        assert sheet["I2"].hyperlink.target.startswith("https://www.regjeringen.no/")
        assert sheet.freeze_panes == "C2" and sheet.auto_filter.ref == "A1:M3"
        assert all(not d.hidden for d in sheet.column_dimensions.values())
        assert all(s.sheet_state == "visible" for s in book)
        for row in rows:
            assert "Ikke faglig godkjent" in row[12]
            assert "Mangler historisk startgrunnlag" in row[12]
            assert "Mangler organisasjonsoppslag" in row[12]
            assert "Mangler valgt kontrollresultat" in row[12]
            assert "Saksrelasjon" in row[12] and "postvis" in row[12]
            assert "Globalt" in row[12]
        details = book["Historikkdetaljer"]
        assert details.max_column == 39 and details.max_row == 3
        assert "Observasjons-ID-er" in [c.value for c in details[1]]
        assert book["Beløpsobservasjoner"].max_row == 7
        assert book["Saksdokumenter"].max_row == 6
    finally:
        book.close()


def test_all_old_history_fields_and_raw_source_ids_survive_in_visible_reserve(tmp_path):
    db, report = source_report(tmp_path)
    old, new = tmp_path / "b6.xlsx", tmp_path / "b9.xlsx"
    export_built_session_excel(db, report.base, old)
    export_built_saks_excel(db, report, new)
    before, after = load_workbook(old), load_workbook(new)
    try:
        source = before["Historikk"]
        reserve = after["Historikkdetaljer"]
        previous = [c.value for c in source[1]]
        actual = [c.value for c in reserve[1]]
        assert len(previous) == 33 and len(actual) == 39
        extra = {
            "Stikkord",
            "Formelt vedtak",
            "Ferskhet for saksgrunnlag",
            "Relasjonsnivå",
            "Innstilling 1",
            "Proposisjon 1",
        }
        assert set(actual) == {"Teknisk API-ID" if h == "Vedtak" else h for h in previous} | extra
        for row_number in (2, 3):
            for column, header in enumerate(previous, 1):
                new_header = "Teknisk API-ID" if header == "Vedtak" else header
                assert (
                    source.cell(row_number, column).value
                    == reserve.cell(row_number, actual.index(new_header) + 1).value
                )
            assert (
                reserve.cell(row_number, actual.index("Formelt vedtak") + 1).value
                == "210 (2024-2025)"
            )
            assert reserve.cell(row_number, actual.index("Innstilling 1") + 1).hyperlink.target
            assert reserve.cell(row_number, actual.index("Proposisjon 1") + 1).hyperlink.target
        amount_sheet = after["Beløpsobservasjoner"]
        amount_headers = [c.value for c in amount_sheet[1]]
        saved = {
            r[0].value: json.loads(r[amount_headers.index("Alle observasjonsfelt")].value)
            for r in amount_sheet.iter_rows(min_row=2)
        }
        assert saved == {
            r.observation.id: asdict(r.observation)
            | {
                "documents": list(asdict(r.observation)["documents"]),
                "relations": list(asdict(r.observation)["relations"]),
                "anomalies": list(asdict(r.observation)["anomalies"]),
                "null_reasons": [list(v) for v in r.observation.null_reasons],
            }
            for r in report.base.rows
        }
        assert {r[6].value for r in after["Stikkordgrunnlag"].iter_rows(min_row=2)} == {
            k.raw_text for e in report.enrichments for k in e.keywords
        }
    finally:
        before.close()
        after.close()


def test_multiple_documents_link_to_all_references_without_multiplying_amount_rows(tmp_path):
    db, original = source_report(tmp_path)
    raw, metadata = case_source()
    root = etree.fromstring(raw)
    publications = root.find(NS + "publikasjon_referanse_liste")
    for index, old_number, new_number in ((2, "81", "82"), (0, "16", "17")):
        duplicate = deepcopy(publications[index])
        for element in duplicate.iter():
            if element.text:
                element.text = element.text.replace(old_number, new_number)
        publications.append(duplicate)
    raw = etree.tostring(root)
    metadata.update(
        sha256=hashlib.sha256(raw).hexdigest(), bytes=len(raw), fixture_kind="synthetic_mutation"
    )
    attempt = import_case_response(
        db, "100927", raw, metadata, dataset="prove", session="2024-2025"
    )
    enrichment = enrich_revision(db, original.contexts[0].decision_revision_id, attempt.id)
    report = build_saks_report(db, SaksSelection(original.selection.base, (enrichment.id,)))
    output = tmp_path / "many.xlsx"
    export_built_saks_excel(db, report, output)
    book = load_workbook(output)
    try:
        main = book["Historikk"]
        assert main.max_row == 3 and main.max_column == 13
        assert "Innst. 81 S" in main["H2"].value and "Innst. 82 S" in main["H2"].value
        assert "Prop. 16 S" in main["I2"].value and "Prop. 17 S" in main["I2"].value
        for coordinate in ("H2", "I2"):
            assert main[coordinate].hyperlink.location.startswith("'Saksdokumenter'!")
        assert "Flere innstillinger" in main["M2"].value
        assert "Flere proposisjoner" in main["M2"].value
        references = list(book["Saksdokumenter"].iter_rows(min_row=2))
        assert len(references) == 7
        assert {r[9].hyperlink.target for r in references} == {
            p.url for p in enrichment.content.publications
        }
        reserve = book["Historikkdetaljer"]
        assert reserve.max_column == 41
        assert {"Innstilling 2", "Proposisjon 2"} <= {c.value for c in reserve[1]}
        assert book["Beløpsobservasjoner"].max_row == 7
    finally:
        book.close()


def test_freshness_and_new_unknown_global_anomalies_are_visible_even_without_amount_rows(tmp_path):
    db, report = source_report(tmp_path, posts=("99",))
    context = replace(report.contexts[0], source_freshness="historiske_referanser_etter_feilet")
    report = replace(report, contexts=(context,))
    report = frozen_variant(report, (), global_issues=("ny_ukjent_kildekonflikt:utenfor_filter",))
    output = tmp_path / "empty.xlsx"
    export_built_saks_excel(db, report, output)
    book = load_workbook(output)
    try:
        main = book["Historikk"]
        assert main.max_column == 13 and main.max_row == 2
        assert all(main.cell(2, n).value is None for n in range(1, 13))
        assert "Ingen beløpsrader" in main["M2"].value
        assert "ny ukjent kildekonflikt:utenfor filter" in main["M2"].value
        assert "historiske referanser etter feilet" in main["M2"].value
        assert main["M2"].hyperlink.location == "'Avvik'!A1"
        assert book["Historikkdetaljer"].max_row == 1
    finally:
        book.close()


@pytest.mark.parametrize(
    "mutation",
    ["missing_before", "level", "duplicate", "unknown", "component", "scope", "interpretation"],
)
def test_uncertain_amount_roles_and_distinct_source_scopes_are_never_silently_selected(
    tmp_path, mutation
):
    db, original = source_report(tmp_path)
    rows = [r for r in original.base.rows if r.observation.post == "1"]
    delta = next(r for r in rows if r.observation.meaning == "endring")
    if mutation == "missing_before":
        rows = [r for r in rows if r.observation.meaning != "for_belop"]
    elif mutation in {"level", "unknown"}:
        rows = [
            replace(
                delta,
                observation=replace(
                    delta.observation, meaning="nivaa" if mutation == "level" else "uavklart"
                ),
            )
        ]
    else:
        changes = {
            "duplicate": {"amount_kr": 9_000_000},
            "component": {
                "detail_level": "underpost",
                "subpost": "1",
                "component_of": delta.observation.step_id,
            },
            "scope": {"scope": "uavklart"},
            "interpretation": {"interpretation_id": "syntetisk-alternativ-tolkning"},
        }[mutation]
        rows.append(
            replace(
                delta,
                observation=replace(
                    delta.observation,
                    pointer=replace(delta.observation.pointer, locator="syntetisk-ekstra-rad"),
                    **changes,
                ),
            )
        )
    report = frozen_variant(original, rows)
    output = tmp_path / "uncertain.xlsx"
    export_built_saks_excel(db, report, output)
    book = load_workbook(output)
    try:
        main = book["Historikk"]
        if mutation == "missing_before":
            assert main["J2"].value is None and main["K2"].value == 5_796_000
            assert "Mangler entydig kildeobservert førbeløp" in main["M2"].value
        elif mutation in {"level", "unknown"}:
            assert all(main.cell(2, n).value is None for n in (10, 11, 12))
            assert "Annen beløpsbetydning" in main["M2"].value
            assert "5796000" in main["M2"].value
        elif mutation == "duplicate":
            assert main.max_row == 2
            assert main["K2"].value is None
            assert "Flere endringsbeløp; ingen verdi valgt" in main["M2"].value
        else:
            assert main.max_row == 3
            statuses = " ".join(main.cell(n, 13).value for n in (2, 3))
            if mutation == "component":
                assert "ikke posttotal" in statuses
            elif mutation == "scope":
                assert "ikke additivt avklart" in statuses
            else:
                assert "Flere adskilte kildegrunnlag" in statuses
        assert book["Beløpsobservasjoner"].max_row == len(rows) + 1
    finally:
        book.close()


@pytest.mark.parametrize("mutation", ["missing", "conflict", "unresolved", "bad_date", "zero"])
def test_missing_conflicting_names_keywords_and_dates_keep_source_evidence_visible(
    tmp_path, mutation
):
    db, original = source_report(tmp_path)
    rows = [r for r in original.base.rows if r.observation.post == "22"]
    if mutation == "missing":
        rows = [
            replace(r, observation=replace(r.observation, chapter_name=None, purpose=None))
            for r in rows
        ]
    elif mutation == "conflict":
        rows[0] = replace(
            rows[0], observation=replace(rows[0].observation, chapter_name="Alternativt kildenavn")
        )
    elif mutation == "unresolved":
        rows = [
            replace(
                r,
                observation=replace(
                    r.observation, purpose=r.observation.purpose + ", ukjent fullmakt"
                ),
            )
            for r in rows
        ]
    elif mutation == "bad_date":
        rows = [replace(r, observation=replace(r.observation, decision_date=None)) for r in rows]
    else:
        rows = [
            replace(r, observation=replace(r.observation, amount_kr=0))
            if r.observation.meaning == "for_belop"
            else r
            for r in rows
        ]
    report = frozen_variant(original, rows)
    output = tmp_path / "names.xlsx"
    export_built_saks_excel(db, report, output)
    book = load_workbook(output)
    try:
        main = book["Historikk"]
        if mutation == "missing":
            assert main["C2"].value.startswith("Uavklart")
            assert main["D2"].value.startswith("Uavklart")
            assert "Uavklart kapittelnavn" in main["M2"].value
            assert "Uavklart postnavn" in main["M2"].value
        elif mutation == "conflict":
            assert "Riksantikvaren" in main["C2"].value
            assert "Alternativt kildenavn" in main["C2"].value
            assert "motstridende kildenavn" in main["M2"].value
        elif mutation == "unresolved":
            assert "ukjent fullmakt" in main["D2"].value
            assert "Uavklart" in main["E2"].value
            assert "stikkord" in main["M2"].value
        elif mutation == "bad_date":
            assert main["F2"].value is None
            assert "Mangler entydig vedtaksdato" in main["M2"].value
        else:
            assert main["J2"].value == 0 and main["J2"].data_type == "n"
        amount_headers = [c.value for c in book["Beløpsobservasjoner"][1]]
        assert len({r[0].value for r in book["Beløpsobservasjoner"].iter_rows(min_row=2)}) == len(
            rows
        )
        assert all(
            r[amount_headers.index("Alle observasjonsfelt")].value
            for r in book["Beløpsobservasjoner"].iter_rows(min_row=2)
        )
    finally:
        book.close()


@pytest.mark.parametrize(
    "cell,value", [("H2", "bad_link"), ("F2", "2024-12-16"), ("K2", 1), ("reserve", "missing-id")]
)
def test_public_verifier_rejects_corrupted_main_links_types_amounts_and_reserve(
    tmp_path, cell, value
):
    db, report = source_report(tmp_path)
    output = tmp_path / "tampered.xlsx"
    export_built_saks_excel(db, report, output)
    book = load_workbook(output)
    if cell == "H2":
        book["Historikk"][cell].hyperlink = "https://www.stortinget.no/wrong-document"
    elif cell == "reserve":
        sheet = book["Historikkdetaljer"]
        headers = {c.value: c.column for c in sheet[1]}
        sheet.cell(2, headers["Observasjons-ID-er"]).value = value
    else:
        book["Historikk"][cell].value = value
    book.save(output)
    book.close()
    with pytest.raises(ValueError):
        verify_saks_excel(output, report)


def test_long_detail_fields_keep_continuation_addresses_after_column_insertion_and_rename(tmp_path):
    db, original = source_report(tmp_path)
    rows = [r for r in original.base.rows if r.observation.post == "1"]
    rows[0] = replace(
        rows[0],
        observation=replace(
            rows[0].observation,
            pointer=replace(rows[0].observation.pointer, raw_record="syntetisk råtekst " * 2500),
        ),
    )
    report = frozen_variant(original, rows)
    output = tmp_path / "long.xlsx"
    export_built_saks_excel(db, report, output)
    book = load_workbook(output)
    try:
        reserve = book["Historikkdetaljer"]
        headers = {c.value: c.column_letter for c in reserve[1]}
        continuation = book["Historikkdetaljer fortsettelse"]
        assert continuation["A2"].value == f"Historikkdetaljer!{headers['Lokatorer']}2"
        reconstructed = reserve[f"{headers['Lokatorer']}2"].value + "".join(
            r[2].value for r in continuation.iter_rows(min_row=2)
        )
        assert [v["raw_record"] for v in json.loads(reconstructed)] == [
            asdict(r.observation.pointer)["raw_record"] for r in rows
        ]
    finally:
        book.close()


def long_main_report(tmp_path, length=67500):
    db, original = source_report(tmp_path)
    # Each 30,000-character chunk starts with formula-looking source text.
    chapter_name = ("=1+" + "a" * 29997 + "=2+" + "b" * 29997 + "=3+" + "c" * 7497)[:length]
    rows = [
        replace(r, observation=replace(r.observation, chapter_name=chapter_name))
        for r in original.base.rows
        if r.observation.post == "1"
    ]
    return db, frozen_variant(original, rows), chapter_name


@pytest.mark.parametrize("length", [37500, 67500])
def test_long_main_text_reverifies_completely_and_preserves_snapshot_and_formula_safety(
    tmp_path, length
):
    db, report, chapter_name = long_main_report(tmp_path, length)
    frozen = report_to_json(report)
    output = tmp_path / "long-main.xlsx"
    export_built_saks_excel(db, report, output)
    verify_saks_excel(output, report)
    assert report_to_json(report) == frozen
    book = load_workbook(output)
    try:
        main = book["Historikk"]["C2"]
        continuation = book["Historikk fortsettelse"]
        assert main.value + "".join(r[2].value for r in continuation.iter_rows(min_row=2)) == (
            chapter_name
        )
        assert list(continuation.values) == [
            ("Kilde-celle", "Del", "Tekst"),
            ("Historikk!C2", 2, "=2+" + "b" * (length - 30003 if length == 37500 else 29997)),
            *([("Historikk!C2", 3, "=3+" + "c" * 7497)] if length == 67500 else []),
        ]
        assert main.data_type == "s"
        assert all(c.data_type != "f" for sheet in book for row in sheet for c in row)
    finally:
        book.close()


@pytest.mark.parametrize(
    "mutation",
    [
        "missing_sheet",
        "sheet_name",
        "header",
        "origin",
        "part",
        "part_order",
        "text",
        "missing_row",
        "extra_row",
        "extra_column",
        "extra_sheet",
        "formula",
        "main_formula",
        "main_text",
    ],
)
def test_public_verifier_rejects_damaged_long_main_text_continuations(tmp_path, mutation):
    db, report, _ = long_main_report(tmp_path)
    output = tmp_path / "damaged-main.xlsx"
    export_built_saks_excel(db, report, output)
    book = load_workbook(output)
    continuation = book["Historikk fortsettelse"]
    if mutation == "missing_sheet":
        book.remove(continuation)
    elif mutation == "sheet_name":
        continuation.title = "Historikk fortsettelse1"
    elif mutation == "header":
        continuation["A1"] = "Feil overskrift"
    elif mutation == "origin":
        continuation["A2"] = "Historikk!D2"
    elif mutation == "part":
        continuation["B2"] = 3
    elif mutation == "part_order":
        second, third = list(continuation.values)[1:]
        for number, values in ((2, third), (3, second)):
            for column, value in enumerate(values, 1):
                continuation.cell(number, column).value = value
    elif mutation == "text":
        continuation["C3"] = "Skadet avslutning"
    elif mutation == "missing_row":
        continuation.delete_rows(3)
    elif mutation == "extra_row":
        continuation.append(["Historikk!C2", 4, "Ekstra tekst"])
    elif mutation == "extra_column":
        continuation["D2"] = "Ekstra tekst"
    elif mutation == "extra_sheet":
        book.copy_worksheet(continuation)
    elif mutation == "formula":
        continuation["C2"].data_type = "f"
    elif mutation == "main_formula":
        book["Historikk"]["C2"].data_type = "f"
    else:
        book["Historikk"]["C2"] = "Skadet innledning"
    book.save(output)
    book.close()
    with pytest.raises(ValueError):
        verify_saks_excel(output, report)


def test_public_verifier_rejects_unexpected_main_continuation_without_long_fields(tmp_path):
    db, report = source_report(tmp_path)
    output = tmp_path / "unexpected-main.xlsx"
    export_built_saks_excel(db, report, output)
    book = load_workbook(output)
    sheet = book.create_sheet("Historikk fortsettelse")
    sheet.append(["Kilde-celle", "Del", "Tekst"])
    sheet.append(["Historikk!C2", 2, "Uventet tekst"])
    book.save(output)
    book.close()
    with pytest.raises(ValueError):
        verify_saks_excel(output, report)


@pytest.mark.parametrize(
    "mutation,expected",
    [
        ("missing", "mangler publikasjon referanse liste"),
        ("unsafe", "mangler eller utrygg dokumentlenke"),
        ("conflict", "motstrid"),
        ("unselected", "Saksberikelse er ikke valgt"),
    ],
)
def test_missing_or_unresolved_document_fields_remain_visible_on_main(tmp_path, mutation, expected):
    db, original = source_report(tmp_path)
    if mutation == "unselected":
        report = build_saks_report(db, SaksSelection(original.selection.base))
    else:
        raw, metadata = case_source()
        root = etree.fromstring(raw)
        publications = root.find(NS + "publikasjon_referanse_liste")
        if mutation == "missing":
            root.remove(publications)
        elif mutation == "unsafe":
            publications[0].find(NS + "lenke_url").text = "javascript:alert(1)"
        else:
            publications[0].find(NS + "lenke_tekst").text = "Prop. 17 S (2024-2025)"
        raw = etree.tostring(root)
        metadata.update(
            sha256=hashlib.sha256(raw).hexdigest(),
            bytes=len(raw),
            fixture_kind="synthetic_mutation",
        )
        attempt = import_case_response(
            db, "100927", raw, metadata, dataset="prove", session="2024-2025"
        )
        enrichment = enrich_revision(db, original.contexts[0].decision_revision_id, attempt.id)
        report = build_saks_report(db, SaksSelection(original.selection.base, (enrichment.id,)))
    output = tmp_path / "documents.xlsx"
    export_built_saks_excel(db, report, output)
    book = load_workbook(output)
    try:
        main = book["Historikk"]
        assert main.max_column == 13 and main.max_row == 3
        assert expected in main["M2"].value
        assert "Saksrelasjon" in main["M2"].value
        assert main["K2"].value == 5_796_000
        assert book["Dokumentavvik"].max_row > 1
        assert all(c.data_type != "f" for sheet in book for row in sheet for c in row)
        if mutation == "unsafe":
            assert all(
                not c.hyperlink
                or not c.hyperlink.target
                or not c.hyperlink.target.startswith("javascript:")
                for sheet in book
                for row in sheet
                for c in row
            )
    finally:
        book.close()


def test_pre_b9_detailed_excel_layout_is_still_verifiable_without_rewriting_the_input():
    output, report = original_b8_fixture()
    book = load_workbook(output)
    try:
        assert book.active.title == "Sesjonsgrunnlag"
        assert "Historikkdetaljer" not in book
        # Original B8 wrote the pointer before inserting its six document columns.
        assert book["Historikk fortsettelse"]["A2"].value == "Historikk!AB2"
        assert book["Historikk"]["AH1"].value == "Lokatorer"
    finally:
        book.close()
    before = output.read_bytes()
    verify_saks_excel(output, report)
    assert output.read_bytes() == before


def original_b8_fixture():
    """Immutable output from a7fc7d9's actual writer; CI needs neither Git nor network."""
    directory = Path(__file__).parent / "fixtures/b9"
    manifest = json.loads((directory / "legacy-b8.manifest.json").read_text("utf-8"))
    output = directory / "legacy-b8.xlsx"
    compressed = (directory / "legacy-b8.json.gz").read_bytes()
    frozen = gzip.decompress(compressed)
    for raw, label in (
        (output.read_bytes(), "xlsx"),
        (compressed, "compressed_json"),
        (frozen, "uncompressed_json"),
    ):
        assert len(raw) == manifest[label]["bytes"]
        assert hashlib.sha256(raw).hexdigest() == manifest[label]["sha256"]
    return output, report_from_json(frozen.decode("utf-8"))


@pytest.mark.parametrize(
    "mutation",
    [
        "text",
        "part",
        "pointer",
        "final_pointer",
        "missing_sheet",
        "sheet_name",
        "header",
        "part_order",
        "missing_row",
        "extra_row",
        "extra_column",
        "extra_sheet",
        "formula",
        "origin_formula",
        "pointer_formula",
    ],
)
def test_original_b8_long_detail_compatibility_still_rejects_corruption(tmp_path, mutation):
    fixture, report = original_b8_fixture()
    output = tmp_path / "damaged-b8.xlsx"
    output.write_bytes(fixture.read_bytes())
    book = load_workbook(output)
    continued = book["Historikk fortsettelse"]
    if mutation == "text":
        continued["C4"] = "Skadet avslutning"
    elif mutation == "part":
        continued["B2"] = 3
    elif mutation == "pointer":
        continued["A2"] = "Historikk!AC2"
    elif mutation == "final_pointer":
        continued["A2"] = "Historikk!AH2"
    elif mutation == "missing_sheet":
        book.remove(continued)
    elif mutation == "sheet_name":
        continued.title = "Historikk fortsettelse1"
    elif mutation == "header":
        continued["A1"] = "Feil overskrift"
    elif mutation == "part_order":
        second, third = list(continued.values)[1:3]
        for number, values in ((2, third), (3, second)):
            for column, value in enumerate(values, 1):
                continued.cell(number, column).value = value
    elif mutation == "missing_row":
        continued.delete_rows(4)
    elif mutation == "extra_row":
        continued.append(["Historikk!AB2", 6, "Ekstra tekst"])
    elif mutation == "extra_column":
        continued["D2"] = "Ekstra tekst"
    elif mutation == "extra_sheet":
        book.copy_worksheet(continued)
    elif mutation == "formula":
        continued["C2"].data_type = "f"
    elif mutation == "origin_formula":
        book["Historikk"]["AH2"].data_type = "f"
    elif mutation == "pointer_formula":
        continued["A2"].data_type = "f"
    book.save(output)
    book.close()
    with pytest.raises(ValueError):
        verify_saks_excel(output, report)


def test_b9_details_reject_the_old_b8_pointer_convention(tmp_path):
    _, report = original_b8_fixture()
    output = tmp_path / "b9-with-legacy-pointer.xlsx"
    export_built_saks_excel(tmp_path / "absent.sqlite", report, output)
    book = load_workbook(output)
    book["Historikkdetaljer fortsettelse"]["A2"] = "Historikkdetaljer!AB2"
    book.save(output)
    book.close()
    with pytest.raises(ValueError):
        verify_saks_excel(output, report)


def test_control_freshness_is_visible_separately_from_current_document_source(tmp_path):
    db, original = source_report(tmp_path)
    rows = [replace(r, freshness="ny_faglig_avgjorelse") for r in original.base.rows]
    report = frozen_variant(original, rows)
    output = tmp_path / "control-freshness.xlsx"
    export_built_saks_excel(db, report, output)
    book = load_workbook(output)
    try:
        assert "Kontrollgrunnlagets ferskhet: ny faglig avgjorelse" in book["Historikk"]["M2"].value
        assert "saksgrunnlag gjeldende" in book["Historikk"]["M2"].value
        assert "Ikke faglig godkjent" in book["Historikk"]["M2"].value
    finally:
        book.close()
