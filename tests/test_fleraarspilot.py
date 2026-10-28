"""T4/#48: autentiske år, samme grunnlag og avtalte offentlige filgrenser."""

import gzip
import hashlib
import subprocess
import sys
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from analyse.fleraarspilot import (
    build_report,
    contents,
    go_back,
    import_basis,
    open_accounts,
    open_document,
    read_json,
    report_from_json,
)
from analyse.fleraarspilot_excel import export_pair, read_excel
from innlesning.visningspilot_model import PilotSelection


def test_original_201_has_2025_levels_complete_inventory_and_unknown_continuity():
    basis = import_basis(Path("."))
    report = build_report(basis, PilotSelection(years=(2025,), chapters=("1429",)))
    assert len(report.view.observation_ids) == 2
    assert {
        (o.post, o.amount_kr, o.budget_year, o.meaning)
        for o in basis.observations
        if o.source_system == "fleraarspilot"
    } == {
        ("1", 163_173_000, 2025, "nivaa"),
        ("22", 33_408_000, 2025, "nivaa"),
    }
    doc = next(d for d in basis.documents if d.key == "vedtak201")
    assert sum(i.document_id == doc.id for i in basis.inventory) == 153
    assert not report.calculations and not report.view.comparisons
    assert all(
        any(account in issue and "kontinuitet" in issue for issue in basis.issues)
        for account in ("1429/1", "1429/22")
    )


@pytest.fixture(scope="module")
def authentic_basis():
    raw = gzip.decompress(Path("tests/fixtures/t2/authentic-t1.json.gz").read_bytes()).decode(
        "utf8"
    )
    return import_basis(Path("."), t1_json=raw)


@pytest.fixture(scope="module")
def saved_pair(tmp_path_factory, authentic_basis):
    directory = tmp_path_factory.mktemp("t4-saved")
    report = build_report(authentic_basis, PilotSelection(years=(2024, 2025), chapters=("1429",)))
    frozen, excel = directory / "both.json", directory / "both.xlsx"
    export_pair(report, frozen, excel)
    return frozen, excel, report


def test_saved_multiyear_pair_reuses_old_ids_and_keeps_2024_calculations(saved_pair):
    from analyse.forlopspilot import build_report as t3_report
    from analyse.forlopspilot import import_basis as t3_basis
    from ops.komitepilot_proeve import serialized_counts

    frozen, excel, report = saved_pair
    assert read_json(frozen) == read_excel(excel) == report
    assert (
        len(report.view.observation_ids),
        len(report.view.adopted_observation_ids),
        len(report.view.adopted_event_ids),
    ) == (25, 13, 7)
    old = t3_basis(Path("."), t1_json=report.basis.legacy_json)
    old_report = t3_report(old, PilotSelection(years=(2024,), chapters=("1429",)))
    assert report.calculations == old_report.calculations
    assert report.basis.observations[:5] == old.observations
    assert len(contents(report.basis).observations) == 25
    assert {c.budget_year for c in report.calculations} == {2024}
    counts = serialized_counts(frozen.read_text("utf8"))
    assert counts["new"]["serialized_observation_objects"] == 7
    assert counts["new"]["unique_observations"] == 7
    assert set(counts["new"]["observation_occurrences"].values()) == {1}
    assert counts["new"]["serialized_raw_payloads"] == 5
    assert set(counts["new"]["payload_occurrences"].values()) == {1}
    assert counts["legacy"]["serialized_observation_objects"] == 36
    assert counts["legacy"]["unique_observations"] == 18
    assert counts["legacy"]["serialized_raw_payloads"] == 23
    assert counts["legacy"]["unique_raw_hashes"] == 21
    assert len([p for p in report.basis.payloads if p.legacy_source_id]) == 2
    assert all(
        p.content_base64 is None and p.legacy_source_id
        for p in report.basis.payloads
        if p.sha256
        in {
            "1ea276a5730a5196f4def7471e0bf6ae51a7878a4eac86922cc87a9a71200a6b",
            "4ccc151e0ddd84c11b9c5e973c67762a6470a7bb5f2cd59a542d2847e6c5cd28",
        }
    )
    book = load_workbook(excel, read_only=True)
    try:
        groups = [(r[1], r[0]) for r in book["Forløp"].iter_rows(min_row=2, values_only=True)]
        assert groups == sorted(groups, key=lambda x: (x[0], int(x[1].split("/")[1])))
        assert len([r for r in groups if r[0] == 2025]) == 2
    finally:
        book.close()


def test_every_year_account_selection_and_both_navigation_directions_use_actual_pairs(
    tmp_path,
    saved_pair,
):
    from ops.fleraarspilot_proeve import selection_cases

    _, _, both = saved_pair
    for name, selection, counts in selection_cases(both.basis):
        report = build_report(both.basis, selection)
        frozen, excel = tmp_path / (name + ".json"), tmp_path / (name + ".xlsx")
        export_pair(report, frozen, excel)
        saved = read_excel(excel)
        assert saved == read_json(frozen) == report
        assert (
            len(saved.view.observation_ids),
            len(saved.view.adopted_observation_ids),
            len(saved.view.adopted_event_ids),
        ) == counts
        assert saved.basis == both.basis
        if selection.years == (2025,) or name == "document-vedtak201":
            assert not saved.calculations and not saved.view.comparisons
    doc = next(d for d in contents(both.basis).documents if d.key == "vedtak201")
    opened = open_document(both, doc.id)
    export_pair(opened, tmp_path / "opened.json", tmp_path / "opened.xlsx")
    saved = read_excel(tmp_path / "opened.xlsx")
    assert len(saved.view.inventory_ids) == 153 and go_back(saved) == both
    account = open_accounts(saved, PilotSelection(years=(2025,), accounts=(("1429", "22"),)))
    export_pair(account, tmp_path / "account.json", tmp_path / "account.xlsx")
    assert go_back(read_excel(tmp_path / "account.xlsx")) == build_report(
        both.basis, opened.selection
    )
    repeated = tmp_path / "again.json", tmp_path / "again.xlsx"
    export_pair(both, *repeated)
    assert repeated[0].read_bytes() == saved_pair[0].read_bytes()
    assert read_excel(repeated[1]) == both
    before = [hashlib.sha256(p.read_bytes()).hexdigest() for p in repeated]
    with pytest.raises(ValueError):
        export_pair(both, *repeated)
    assert before == [hashlib.sha256(p.read_bytes()).hexdigest() for p in repeated]


@pytest.mark.parametrize(
    "mutation",
    [
        "year",
        "amount",
        "meaning",
        "date",
        "role",
        "raw-name",
        "locator",
        "parent",
        "recipe",
        "reference",
        "hidden-gaps",
        "continuity",
        "selection",
        "calculation",
        "calculation-ref",
        "old-field",
        "inventory-cap",
        "legacy-reference",
        "extra-field",
    ],
)
def test_rehashed_negative_json_and_excel_pairs_are_rejected(tmp_path, saved_pair, mutation):
    from ops.fleraarspilot_proeve import mutate
    from ops.forlopspilot_proeve import _snapshot

    frozen, excel, _ = saved_pair
    raw = mutate(frozen.read_text("utf8"), mutation)
    path = tmp_path / ("synthetic-" + mutation + ".json")
    path.write_text(raw, "utf8")
    with pytest.raises(ValueError):
        read_json(path)
    changed = path.with_suffix(".xlsx")
    _snapshot(excel, changed, raw)
    with pytest.raises(ValueError):
        read_excel(changed)


@pytest.mark.parametrize("condition", ["missing", "failed", "corrupt", "contradictory"])
def test_public_ordinary_368_survives_bad_t4_and_empty_multiyear_selection(tmp_path, condition):
    from ops.fleraarspilot_proeve import bad_sources
    from ops.visningspilot_proeve import _ordinary

    root = bad_sources(tmp_path / "sources", condition)
    basis = import_basis(root)
    assert len(basis.observations) == 5  # Tilgjengelig T3 bevares uavhengig av T4-feilen.
    report = build_report(basis, PilotSelection(years=(2024, 2025), accounts=(("999", "1"),)))
    export_pair(report, tmp_path / "empty.json", tmp_path / "empty.xlsx")
    assert read_json(tmp_path / "empty.json") == read_excel(tmp_path / "empty.xlsx") == report
    assert not report.view.observation_ids
    assert any("201_forelder_eller_transport_mangler" in i for i in report.basis.issues)
    ordinary = _ordinary(tmp_path)
    assert ordinary["observations"] == 222 and ordinary["events"] == 74
    assert ordinary["raw_preserved"]


def test_offline_reading_cli_and_original_recovery_without_network_database_or_root(
    tmp_path,
    saved_pair,
):
    frozen, _, both = saved_pair
    script = """import socket, sqlite3, sys
from pathlib import Path
def unavailable(*args, **kwargs): raise AssertionError('network/database forbidden')
socket.create_connection=unavailable
socket.socket.connect=unavailable
socket.socket.connect_ex=unavailable
sqlite3.connect=unavailable
sys.argv=['analyse.fleraarspilot','--frosset',sys.argv[1],'--rot',sys.argv[2],
          '--json',sys.argv[3],'--excel',sys.argv[4],'--originaler',sys.argv[5]]
from analyse.fleraarspilot import main
main()
"""
    path = tmp_path / "offline.json"
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            script,
            str(frozen),
            str(tmp_path / "no-root"),
            str(path),
            str(path.with_suffix(".xlsx")),
            str(tmp_path / "originals"),
        ],
        capture_output=True,
        text=True,
        encoding="utf8",
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert read_json(path) == read_excel(path.with_suffix(".xlsx")) == both
    original = Path("docs/research/to-visninger-37/originals/aar/40034038.xml").read_bytes()
    assert (tmp_path / "originals/derived-40034038.xml").read_bytes() == original


def test_new_adapter_has_no_analysis_dependency_and_year_guard_is_unchanged(
    tmp_path, authentic_basis
):
    from analyse.forlopspilot import calculate_path
    from innlesning.fleraarspilot import derived_xml

    script = """import importlib.abc, sys
from pathlib import Path
class NoAnalysis(importlib.abc.MetaPathFinder):
 def find_spec(self, fullname, path, target=None):
  if fullname == 'analyse' or fullname.startswith('analyse.'):
   raise RuntimeError('analysis unavailable')
sys.meta_path.insert(0, NoAnalysis())
from innlesning.fleraarspilot import import_basis, validate_basis
basis=import_basis(Path('.')); validate_basis(basis)
assert len(basis.observations)==7
"""
    result = subprocess.run(
        [sys.executable, "-c", script], capture_output=True, text=True, encoding="utf8"
    )
    assert result.returncode == 0, result.stdout + result.stderr
    data = contents(authentic_basis)
    assert calculate_path(data.observations, data.contexts, account="1429/1").status == "sperret"
    first = data.observations[0]
    changed = replace(first, budget_year=2025)
    observations = (changed,) + tuple(o for o in data.observations[1:] if o.budget_year == 2024)
    contexts = tuple(
        replace(c, observation_id=changed.id) if c.observation_id == first.id else c
        for c in data.contexts
    )
    assert calculate_path(observations, contexts, account="1429/1").amount_kr is None
    for raw in (
        '{"format":"visningspilot-4","format":"visningspilot-4","report":{}}',
        '{"format":"visningspilot-4","report":NaN}',
    ):
        with pytest.raises(ValueError):
            report_from_json(raw)
    with pytest.raises(ValueError):
        derived_xml(b"synthetic false parent")


def test_201_raw_name_keywords_source_year_and_t3_objects_survive_saved_files(saved_pair):
    from analyse import forlopspilot as t3

    frozen, excel, both = saved_pair
    saved = read_json(frozen)
    old = t3.import_basis(Path("."), t1_json=both.basis.legacy_json)
    for field in ("documents", "inventory", "statements", "observations", "contexts", "relations"):
        assert getattr(saved.basis, field)[: len(getattr(old, field))] == getattr(old, field)
    new = [o for o in saved.basis.observations if o.source_system == "fleraarspilot"]
    assert [o.pointer.locator for o in new] == ["table[1]/tr[101]/td[4]", "table[1]/tr[103]/td[4]"]
    assert new[1].purpose == "Flerårige prosjekt kulturmiljøforvaltning, kan overføres"
    assert all(
        o.decision_date == "2024-12-16"
        and o.budget_year == 2025
        and o.professional_status == "ikke_faglig_godkjent"
        for o in new
    )
    assert read_excel(excel) == saved


@pytest.mark.parametrize(
    "mutation",
    [
        "amount",
        "hidden-row",
        "hidden-column",
        "hyperlink",
        "filter",
        "doctype",
        "numberformat",
        "hidden-sheet",
    ],
)
def test_false_or_hidden_saved_excel_presentation_is_rejected(tmp_path, saved_pair, mutation):
    from ops.fleraarspilot_proeve import mutate_display

    path = tmp_path / ("synthetic-display-" + mutation + ".xlsx")
    mutate_display(saved_pair[1], path, mutation)
    with pytest.raises(ValueError):
        read_excel(path)
