"""T3/#46: avtalte offentlige original-/filgrenser, autentisk 2024-fasit."""

import gzip
import hashlib
import json
import shutil
import subprocess
import sys
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

import pytest
from lxml import etree
from openpyxl import load_workbook

from analyse.forlopspilot import (
    build_report,
    calculate_path,
    contents,
    export_json,
    go_back,
    import_basis,
    open_accounts,
    open_document,
    read_json,
    report_from_json,
)
from analyse.forlopspilot_excel import export_pair, read_excel
from innlesning.forlopspilot import CATALOG, MAX_INVENTORY, META, PARENT, derived_xml
from innlesning.observations import encode, identity
from innlesning.visningspilot_model import PilotSelection


def test_originals_bind_levels_signed_rnb_and_complete_inventory():
    basis = import_basis(Path("."))
    report = build_report(
        basis, PilotSelection(years=(2024,), accounts=(("1429", "1"), ("1429", "22")))
    )
    assert len(basis.observations) == 5
    assert len(report.view.adopted_observation_ids) == 5
    assert len(report.view.adopted_event_ids) == 3
    assert {(o.round, o.meaning, o.amount_kr) for o in basis.observations} == {
        ("opprinnelig", "nivaa", 161_243_000),
        ("opprinnelig", "nivaa", 32_185_000),
        ("rnb", "for_belop", 161_243_000),
        ("rnb", "endring", -1_000_000),
        ("rnb", "etter_belop", 160_243_000),
    }
    docs = {d.key: d for d in basis.documents}
    for key, count in (("vedtak341", 153), ("vedtak857", 1488)):
        assert sum(i.document_id == docs[key].id for i in basis.inventory) == count
    assert any("857_1429_22" in issue for issue in basis.issues)
    assert all(c.status == "sperret" and c.amount_kr is None for c in report.calculations)


@pytest.fixture(scope="module")
def authentic_basis():
    raw = gzip.decompress(Path("tests/fixtures/t2/authentic-t1.json.gz").read_bytes()).decode(
        "utf8"
    )
    return import_basis(Path("."), t1_json=raw)


def test_inventory_keeps_totals_components_and_absence_out_of_post_membership(authentic_basis):
    docs = {d.key: d for d in authentic_basis.documents}
    assert "1429/22" not in docs["vedtak857"].accounts
    assert "1429/1" in docs["vedtak857"].accounts
    for key, row, role in (
        ("vedtak341", 116, "total_eller_kontekst"),
        ("vedtak341", 153, "total_eller_kontekst"),
        ("vedtak857", 1368, "ustrukturert_spesifikasjon"),
        ("vedtak857", 1373, "total_eller_kontekst"),
    ):
        item = next(
            i
            for i in authentic_basis.inventory
            if i.document_id == docs[key].id and i.locator == f"table[1]/tr[{row}]"
        )
        assert item.post == "" and item.role == role and not item.observation_ids


@pytest.fixture(scope="module")
def saved_pair(tmp_path_factory, authentic_basis):
    directory = tmp_path_factory.mktemp("t3-saved-union")
    report = build_report(authentic_basis, PilotSelection(years=(2024,), chapters=("1429",)))
    frozen, excel = directory / "both.json", directory / "both.xlsx"
    export_pair(report, frozen, excel)
    return frozen, excel, report


def test_saved_union_json_excel_preserves_old_fields_and_bounded_calculations(saved_pair):
    from analyse.visningspilot import report_from_json as read_t1

    frozen, excel, report = saved_pair
    before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (frozen, excel)}
    saved = read_json(frozen)
    assert read_excel(excel) == saved == report
    assert (
        len(saved.view.observation_ids),
        len(saved.view.adopted_observation_ids),
        len(saved.view.adopted_event_ids),
    ) == (23, 11, 5)
    assert {(c.account, c.status, c.amount_kr, c.observed_kr) for c in saved.calculations} == {
        ("1429/1", "avgrenset kontroll", 166_039_000, 166_039_000),
        ("1429/22", "avgrenset kontroll", 29_185_000, 29_185_000),
    }
    old = read_t1(report.basis.legacy_json)
    assert all("T1-kapsel: " + issue in contents(saved.basis).issues for issue in old.basis.issues)
    assert (
        tuple(
            o
            for o in contents(saved.basis).observations
            if o.id in {p.id for p in old.basis.observations}
        )
        == old.basis.observations
    )
    assert len(saved.basis.observations) == 5
    book = load_workbook(excel, read_only=True, data_only=False)
    try:
        date = book["Forløp"]["D2"]
        assert date.value == datetime(2023, 12, 19)
        # Escapet punktum vises som dato også i Artifact Tool; uescapet ble serienummer.
        assert date.number_format == 'dd"."mm"."yyyy'
        literal_equals = [
            c for row in book["Omtale"] for c in row if isinstance(c.value, str) and "=" in c.value
        ]
        assert literal_equals and all(c.data_type == "s" for c in literal_equals)
    finally:
        book.close()
    with pytest.raises(ValueError):
        export_pair(report, frozen, excel)
    assert before == {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in before}


def test_selection_counts_navigation_and_actual_offline_files(tmp_path, saved_pair):
    _, _, both = saved_pair
    data = contents(both.basis)
    docs = {d.key: d.id for d in data.documents}
    cases = [
        ("post1", PilotSelection(years=(2024,), accounts=(("1429", "1"),)), (13, 7, 3)),
        ("post22", PilotSelection(years=(2024,), accounts=(("1429", "22"),)), (10, 4, 2)),
        (
            "pairs-no-crossproduct",
            PilotSelection(years=(2024,), accounts=(("1429", "1"), ("315", "22"))),
            (13, 7, 3),
        ),
        ("empty", PilotSelection(), (0, 0, 0)),
        ("wrong-year", PilotSelection(years=(2023,), chapters=("1429",)), (0, 0, 0)),
        ("fullmakt", PilotSelection(years=(2024,), accounts=(("1429", "60"),)), (0, 0, 0)),
        (
            "both",
            PilotSelection(years=(2024,), accounts=(("1429", "1"), ("1429", "22"))),
            (23, 11, 5),
        ),
        ("341", PilotSelection(mode="document", document_id=docs["vedtak341"]), (2, 2, 2)),
        ("857", PilotSelection(mode="document", document_id=docs["vedtak857"]), (3, 3, 1)),
        ("210", PilotSelection(mode="document", document_id=docs["vedtak210"]), (6, 6, 2)),
    ]
    for name, selection, expected in cases:
        report = build_report(both.basis, selection)
        path = tmp_path / (name + ".json")
        export_json(report, path)
        saved = read_json(path)
        assert saved.basis == both.basis
        assert (
            len(saved.view.observation_ids),
            len(saved.view.adopted_observation_ids),
            len(saved.view.adopted_event_ids),
        ) == expected
    document = open_document(both, docs["vedtak857"])
    export_pair(document, tmp_path / "document.json", tmp_path / "document.xlsx")
    opened = read_excel(tmp_path / "document.xlsx")
    assert go_back(opened) == both
    assert len(opened.view.inventory_ids) == 1488
    account = open_accounts(opened, cases[0][1])
    export_pair(account, tmp_path / "account.json", tmp_path / "account.xlsx")
    assert go_back(read_excel(tmp_path / "account.xlsx")) == build_report(
        both.basis, document.selection
    )
    empty = build_report(both.basis, PilotSelection())
    export_pair(empty, tmp_path / "empty-pair.json", tmp_path / "empty-pair.xlsx")
    assert read_excel(tmp_path / "empty-pair.xlsx").basis == both.basis
    book = load_workbook(tmp_path / "empty-pair.xlsx", read_only=True)
    try:
        assert book.sheetnames[:5] == [
            "Forløp",
            "Dokumenter",
            "Omtale",
            "Sammenligninger",
            "Dekning og avvik",
        ]
        assert book["Forløp"].max_row == 1
        assert book["Dokumenter"].max_row > 1641
        assert book["Omtale"].max_row > 54
        assert any(
            "periodeinventar" in str(row[0])
            for row in book["Dekning og avvik"].iter_rows(values_only=True)
        )
    finally:
        book.close()


@pytest.mark.parametrize(
    "field,value",
    [
        ("budget_year", 2025),
        ("currency", "EUR"),
        ("multiplier", "1000"),
        ("scope", "periodeaggregat"),
        ("meaning", "endring"),
        ("chapter", "1428"),
        ("round", "nysaldering"),
        ("round", "uavklart"),
        ("amount_kr", None),
        ("figure_type", "regjeringsforslag"),
        ("decision_date", "2024-12-17"),
    ],
)
def test_synthetic_incompatible_inputs_block_bounded_sum(authentic_basis, field, value):
    data = contents(authentic_basis)
    first = data.observations[0]
    changed = replace(first, **{field: value})
    contexts = tuple(
        replace(c, observation_id=changed.id) if c.observation_id == first.id else c
        for c in data.contexts
    )
    result = calculate_path((changed,) + data.observations[1:], contexts, account="1429/1")
    assert result.status == "sperret" and result.reason and result.amount_kr is None


def rehashed_mutation(raw, mutation):
    """Merkede syntetiske negativfiler; nye hash er ikke kildefasit."""
    data = json.loads(raw)
    report, basis = data["report"], data["report"]["basis"]
    if mutation == "amount":
        basis["observations"][0]["amount_kr"] += 1
    elif mutation == "sign":
        basis["observations"][3]["amount_kr"] = 1_000_000
    elif mutation in {"year", "round", "meaning", "role"}:
        key, value = {
            "year": ("budget_year", 2025),
            "round": ("round", "nysaldering"),
            "meaning": ("meaning", "endring"),
            "role": ("figure_type", "regjeringsforslag"),
        }[mutation]
        basis["observations"][0][key] = value
    elif mutation == "locator":
        basis["observations"][0]["pointer"]["locator"] = "table[1]/tr[103]/td[4]"
    elif mutation == "parent":
        source = next(s for s in basis["sources"] if s["name"].startswith("derived:"))
        properties = json.loads(source["metadata_json"])
        properties["parent_sha256"] = (
            "554e9960c3e90932594106463cca9c285abb53e9272c1ae391d6fb8093baf0bf"
        )
        source["metadata_json"] = encode(properties)
    elif mutation == "dangling":
        basis["relations"][0]["to_id"] = "synthetic-missing"
    elif mutation == "order":
        basis["documents"][0]["date"] = "2024-12-20"
    elif mutation == "hidden-gaps":
        basis["issues"] = []
    elif mutation == "absence":
        next(s for s in basis["statements"] if "fravaer" in s["role"])["text"] = (
            "ingen RNB i hele året"
        )
    elif mutation == "calculation":
        report["calculations"][0]["amount_kr"] += 1
    elif mutation == "calculation-status":
        report["calculations"][0]["status"] = "godkjent aarsbalanse"
    elif mutation == "calculation-ref":
        report["calculations"][0]["input_ids"] = []
    elif mutation == "selection":
        report["selection"]["years"] = [2025]
    elif mutation == "inventory-cap":
        basis["inventory"] = [basis["inventory"][0]] * (MAX_INVENTORY + 1)
    elif mutation == "old-field":
        old = json.loads(basis["legacy_json"])
        old["report"]["basis"]["observations"][0]["amount_kr"] += 1
        old["report"]["basis"]["id"] = identity(
            "visningsgrunnlag-1", {**old["report"]["basis"], "id": ""}
        )
        old["report"]["id"] = identity("visningspilot-1", {**old["report"], "id": ""})
        basis["legacy_json"] = encode(old)
    for c in report["calculations"]:
        c["id"] = ""
        c["id"] = identity("forlopskontroll-1", c)
    basis["id"] = ""
    basis["id"] = identity("forlopsgrunnlag-1", basis)
    report["id"] = ""
    report["id"] = identity("visningspilot-3", report)
    return encode(data)


def replace_snapshot(source, destination, raw):
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(source) as original, ZipFile(destination, "w") as changed:
        for entry in original.infolist():
            content = original.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet7.xml":
                root = etree.fromstring(content)
                rows = root.find(namespace + "sheetData")
                for row in list(rows)[1:]:
                    rows.remove(row)
                for n, start in enumerate(range(0, len(raw), 30000), 2):
                    row = etree.SubElement(rows, namespace + "row", r=str(n))
                    a = etree.SubElement(row, namespace + "c", r=f"A{n}", t="n")
                    etree.SubElement(a, namespace + "v").text = str(n - 1)
                    b = etree.SubElement(row, namespace + "c", r=f"B{n}", t="inlineStr")
                    inline = etree.SubElement(b, namespace + "is")
                    etree.SubElement(inline, namespace + "t").text = raw[start : start + 30000]
                content = etree.tostring(root)
            changed.writestr(entry, content)


@pytest.mark.parametrize(
    "mutation",
    [
        "amount",
        "sign",
        "year",
        "round",
        "meaning",
        "role",
        "locator",
        "parent",
        "dangling",
        "order",
        "hidden-gaps",
        "absence",
        "calculation",
        "calculation-status",
        "calculation-ref",
        "selection",
        "inventory-cap",
        "old-field",
    ],
)
def test_rehashed_synthetic_actual_json_and_excel_reject_false_binding(
    tmp_path, saved_pair, mutation
):
    frozen, excel, _ = saved_pair
    raw = rehashed_mutation(frozen.read_text("utf8"), mutation)
    json_path = tmp_path / ("synthetic-" + mutation + ".json")
    excel_path = tmp_path / ("synthetic-" + mutation + ".xlsx")
    json_path.write_text(raw, "utf8")
    with pytest.raises(ValueError):
        read_json(json_path)
    replace_snapshot(excel, excel_path, raw)
    with pytest.raises(ValueError):
        read_excel(excel_path)


@pytest.mark.parametrize(
    "mutation",
    [
        "amount",
        "numberformat",
        "hidden-row",
        "hidden-column",
        "hidden-sheet",
        "hyperlink",
        "filter",
        "doctype",
    ],
)
def test_synthetic_excel_false_or_hidden_presentation_is_rejected(tmp_path, saved_pair, mutation):
    _, excel, _ = saved_pair
    path = tmp_path / ("synthetic-display-" + mutation + ".xlsx")
    s = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(excel) as original, ZipFile(path, "w") as changed:
        for entry in original.infolist():
            raw = original.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                root = etree.fromstring(raw)
                if mutation == "amount":
                    root.find(f".//{s}c[@r='G2']/{s}v").text = "999"
                elif mutation == "hidden-row":
                    root.find(f".//{s}row[@r='2']").set("hidden", "1")
                elif mutation == "hidden-column":
                    root.find(f".//{s}col").set("hidden", "1")
                elif mutation == "hyperlink":
                    root.find(f".//{s}hyperlink").set("location", "'Omtale'!A1")
                elif mutation == "filter":
                    root.find(s + "autoFilter").set("ref", "A1:N2")
                raw = etree.tostring(root)
                if mutation == "doctype":
                    raw = b'<!DOCTYPE worksheet [<!ENTITY synthetic "bad">]>' + raw
            if mutation == "numberformat" and entry.filename == "xl/styles.xml":
                raw = raw.replace(b'numFmtId="3"', b'numFmtId="1"')
            if mutation == "hidden-sheet" and entry.filename == "xl/workbook.xml":
                raw = raw.replace(b'state="visible"', b'state="hidden"', 1)
            changed.writestr(entry, raw)
    with pytest.raises(ValueError):
        read_excel(path)


@pytest.mark.parametrize("condition", ["missing", "failed", "corrupt", "contradictory"])
def test_new_authentic_368_public_core_survives_bad_t3_and_empty_selection(tmp_path, condition):
    from ops.visningspilot_proeve import _ordinary

    source_root = tmp_path / "sources"
    for name in CATALOG:
        destination = source_root / name
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(name, destination)
    source = source_root / (PARENT if condition == "corrupt" else META)
    if condition in {"missing", "failed"}:
        source.unlink()
        if condition == "failed":
            source.mkdir()
    elif condition == "corrupt":
        source.write_bytes(b"synthetic invalid gzip")
    else:
        source.write_text(source.read_text("utf8").replace("6816165", "6816164"), "utf8")
    basis = import_basis(source_root)
    assert not basis.observations
    empty = build_report(basis, PilotSelection())
    export_pair(empty, tmp_path / "empty.json", tmp_path / "empty.xlsx")
    saved = read_excel(tmp_path / "empty.xlsx")
    assert not saved.view.observation_ids and any(
        "mangler" in s or "motstrid" in s for s in saved.basis.issues
    )
    ordinary = _ordinary(tmp_path)
    assert ordinary["observations"] == 222 and ordinary["events"] == 74
    assert ordinary["raw_preserved"]


def test_ingestion_adapter_without_analysis_and_strict_resource_guards(tmp_path):
    script = """import importlib.abc, sys
from pathlib import Path
class NoAnalysis(importlib.abc.MetaPathFinder):
 def find_spec(self, fullname, path, target=None):
  if fullname == "analyse" or fullname.startswith("analyse."):
   raise RuntimeError("analysis unavailable")
sys.meta_path.insert(0, NoAnalysis())
from innlesning.forlopspilot import import_basis, validate_basis
basis=import_basis(Path("."));validate_basis(basis)
assert len(basis.observations)==5
"""
    result = subprocess.run(
        [sys.executable, "-c", script], capture_output=True, text=True, encoding="utf8"
    )
    assert result.returncode == 0, result.stdout + result.stderr
    with pytest.raises(ValueError):
        report_from_json('{"format":"visningspilot-3","format":"visningspilot-3","report":{}}')
    with pytest.raises(ValueError):
        report_from_json('{"format":"visningspilot-3","report":NaN}')
    with pytest.raises(ValueError):
        derived_xml(b"synthetic wrong parent", "40033719")
