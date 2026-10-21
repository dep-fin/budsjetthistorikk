"""Offentlig pilotgrense: syntetisk OOXML → SQLite → gjenåpning → faktisk pakke."""

import hashlib
import json
import os
import sqlite3
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook

from analyse.historical_package import export_historical_package, verify_historical_package
from innlesning.historical_excel import (
    import_historical_workbook,
    read_historical_workbook,
    resume_historical_import,
    stage_historical_workbook,
)
from innlesning.historical_xlsx import Limits
from innlesning.storage import read_source


def workbook_bytes(*, amount: str = "14.768000000000001", extra: str = "") -> bytes:
    stream = BytesIO()
    namespace = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", '<Types xmlns="urn:synthetic"/>')
        archive.writestr(
            "xl/workbook.xml",
            f'<workbook {namespace} xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
            '<sheets><sheet name="Rådata" sheetId="1" r:id="rId7"/>'
            '<sheet name="Dokumentiltknyntning2024" sheetId="2" r:id="rId8"/>'
            "</sheets></workbook>",
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId7" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"'
            ' Target="worksheets/other.xml"/>'
            '<Relationship Id="rId8" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"'
            ' Target="worksheets/docs.xml"/>'
            "</Relationships>",
        )
        archive.writestr(
            "xl/worksheets/other.xml",
            f'<worksheet {namespace}><dimension ref="A1:N3"/>'
            '<cols><col min="1" max="16384" style="1" hidden="1"/></cols>'
            '<sheetData><row r="1"><c r="H1" t="inlineStr"><is><t>Saldert (f120)</t></is></c></row>'
            '<row r="2" hidden="1"><c r="A2"><v>15</v></c><c r="B2"><v>1</v></c>'
            '<c r="C2"><v>1</v></c><c r="D2"><v>1</v></c>'
            f'<c r="H2"><v>{amount}</v></c><c r="G2"><v>0</v></c>'
            '<c r="I2" t="n"/>'
            '</row><row r="3"><c r="B3"><v>2</v></c><c r="C3"><v>1</v></c>'
            '<c r="D3"><v>0</v></c><c r="H3"><f>1+1</f><v>2</v></c></row>'
            f"{extra}</sheetData></worksheet>",
        )
        archive.writestr(
            "xl/worksheets/docs.xml",
            f'<worksheet {namespace}><dimension ref="A1:O2"/><sheetData>'
            '<row r="1"><c r="D1" t="inlineStr"><is><t>belop</t></is></c></row>'
            '<row r="2"><c r="A2"><v>1</v></c><c r="B2"><v>1</v></c>'
            '<c r="C2"><v>0</v></c><c r="D2"><v>-7</v></c>'
            '<c r="E2" t="inlineStr"><is><t>p1/23-24</t></is></c>'
            '<c r="F2" t="inlineStr"><is><t>i2/23-24</t></is></c>'
            '<c r="G2" t="inlineStr"><is><t>VedtakRNB</t></is></c>'
            '<c r="H2" t="inlineStr"><is><t>Driftsutgifter</t></is></c>'
            '<c r="I2" t="inlineStr"><is><t>kan overføres</t></is></c>'
            '<c r="J2" t="inlineStr"><is><t>Statsministerens kontor</t></is></c>'
            '<c r="K2"><v>92</v></c><c r="L2"><v>16</v></c>'
            "</row></sheetData></worksheet>",
        )
    return stream.getvalue()


def test_document_names_keep_post_keyword_and_chapter_distinct(tmp_path: Path) -> None:
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    source.write_bytes(workbook_bytes())
    result = import_historical_workbook(database, source, budget_year=2024)
    snapshot = read_historical_workbook(database, result.extraction_id)
    document = next(
        o
        for o in snapshot.extraction.observations
        if o.pointer.locator == "Dokumentiltknyntning2024!D2"
    )
    assert document.chapter_name == "Statsministerens kontor"
    assert document.purpose == "Driftsutgifter"
    assert json.loads(document.pointer.raw_record or "{}")["I"] == "kan overføres"


@pytest.mark.parametrize(
    "version,old_id",
    [
        ("v1", "7875b8bafed25b14dd8fe5072ec72ff744168d20e3fa2dc15232032302fffe11"),
        ("v2", "7f1fe49e087fb4760eb96cdf54b0268d2ef5968ae59a1b6860adb96f53743c82"),
    ],
)
def test_saved_versions_remain_readable_beside_a_new_v3_predecessor(
    tmp_path: Path, version: str, old_id: str
) -> None:
    database = tmp_path / "pilot.sqlite"
    with sqlite3.connect(database) as connection:
        connection.executescript(
            (Path(__file__).parent / f"fixtures/historical-{version}.sql").read_text(
                encoding="utf-8"
            )
        )
    old = read_historical_workbook(database, old_id)
    assert old.extraction.interpretation.adapter_version == f"historisk-excel-2024-{version}"
    source = tmp_path / "same-source.xlsx"
    source.write_bytes(read_source(database, old.extraction.artifact.sha256))
    result = import_historical_workbook(database, source, budget_year=2024, predecessor_id=old_id)
    new = read_historical_workbook(database, result.extraction_id)
    assert new.extraction.interpretation.adapter_version == "historisk-excel-2024-v3"
    assert new.extraction.id != old_id and new.predecessor_id == old_id
    assert read_historical_workbook(database, old_id) == old
    assert (
        import_historical_workbook(
            database, source, budget_year=2024, predecessor_id=old_id
        ).extraction_id
        == new.extraction.id
    )


@pytest.mark.parametrize(
    "formula", ["<f>1+1</f>", '<f t="shared" si="0"/>', '<f t="array" ref="H3:H4"/>', "<f/>"]
)
def test_every_formula_presence_blocks_cached_numeric_candidate(
    tmp_path: Path, formula: str
) -> None:
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    source.write_bytes(
        mutate_member(workbook_bytes(), "xl/worksheets/other.xml", b"<f>1+1</f>", formula.encode())
    )
    result = import_historical_workbook(database, source, budget_year=2024)
    snapshot = read_historical_workbook(database, result.extraction_id)
    cell = next(c for c in snapshot.cells if c["sheet"] == "Rådata" and c["cell"] == "H3")
    observation = next(
        o for o in snapshot.extraction.observations if o.pointer.locator == "Rådata!H3"
    )
    candidate = next(c for c in snapshot.candidates if c["observation_id"] == observation.id)
    assert candidate["candidate_kr"] is None
    assert cell["formula_present"] is True
    assert cell["formula"] == ("1+1" if formula == "<f>1+1</f>" else None)
    assert cell["cached_value"] == "2" and observation.cached_value == "2"
    assert observation.formula is not None  # empty formula still has an explicit marker
    assert observation.amount_kr is None


def formula_range_workbook(formula_type: str = "array", ref: str = "H2:H3") -> bytes:
    raw = mutate_member(
        workbook_bytes(amount="2"),
        "xl/worksheets/other.xml",
        b'<c r="H2"><v>2</v></c>',
        f'<c r="H2"><f t="{formula_type}" ref="{ref}" si="0">1+1</f><v>2</v></c>'.encode(),
    )
    return mutate_member(raw, "xl/worksheets/other.xml", b"<f>1+1</f>", b"")


@pytest.mark.parametrize("formula_type", ["array", "shared", "dataTable"])
def test_formula_region_blocks_covered_cache_and_preserves_own_formula_source(
    tmp_path: Path, formula_type: str
) -> None:
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    raw = formula_range_workbook(formula_type)
    source.write_bytes(raw)
    result = import_historical_workbook(database, source, budget_year=2024)
    assert result.status == "lagret"
    snapshot = read_historical_workbook(database, result.extraction_id)
    cells = {c["cell"]: c for c in snapshot.cells if c["sheet"] == "Rådata"}
    observations = {o.pointer.cell: o for o in snapshot.extraction.observations}
    candidates = {c["observation_id"]: c for c in snapshot.candidates}
    for ref in ("H3", "H2"):
        assert candidates[observations[ref].id]["candidate_kr"] is None
        assert cells[ref]["formula_result"] is True
        assert cells[ref]["formula_region"] == {
            "anchor": "H2",
            "ref": "H2:H3",
            "type": formula_type,
        }
        assert cells[ref]["cached_value"] == observations[ref].cached_value == "2"
        assert observations[ref].formula is not None
        assert "formelresultat" in {a.code for a in observations[ref].anomalies}
    assert cells["H3"]["formula_present"] is False
    assert cells["H3"]["formula"] is None and cells["H3"]["formula_attributes"] is None
    assert cells["H2"]["formula_present"] is True and cells["H2"]["formula"] == "1+1"
    assert candidates[observations["G2"].id]["candidate_kr"] == "0"
    assert read_source(database, result.source_sha256 or "") == raw


def test_formula_region_is_independent_of_xml_order_and_sheet_dimension(tmp_path: Path) -> None:
    raw = formula_range_workbook(ref="H2:I3")
    member = "xl/worksheets/other.xml"
    with ZipFile(BytesIO(raw)) as archive:
        xml = archive.read(member)
    start, middle, end = (
        xml.index(b'<row r="2"'),
        xml.index(b'<row r="3"'),
        xml.index(b"</sheetData>"),
    )
    raw = mutate_member(raw, member, xml[start:end], xml[middle:end] + xml[start:middle])
    raw = mutate_member(raw, member, b"A1:N3", b"A1:I2")
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    source.write_bytes(raw)
    result = import_historical_workbook(database, source, budget_year=2024)
    assert result.status == "lagret"
    snapshot = read_historical_workbook(database, result.extraction_id)
    cells = {c["cell"]: c for c in snapshot.cells if c["sheet"] == "Rådata"}
    assert cells["H3"]["formula_result"] is True  # physical cache outside dimension, seen first
    assert cells["I2"]["formula_result"] is True and cells["I2"]["state"] == "blank"
    assert "I3" not in cells  # formula ranges do not materialize absent positions
    assert cells["G2"]["formula_result"] is False


def test_scalar_array_and_absent_positions_keep_source_and_region(tmp_path: Path) -> None:
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    source.write_bytes(formula_range_workbook(ref="H2"))
    result = import_historical_workbook(database, source, budget_year=2024)
    snapshot = read_historical_workbook(database, result.extraction_id)
    cells = {c["cell"]: c for c in snapshot.cells if c["sheet"] == "Rådata"}
    assert cells["H2"]["formula_attributes"]["ref"] == "H2"
    assert cells["H3"]["formula_result"] is False and cells["H3"]["raw_v"] == "2"
    source.write_bytes(formula_range_workbook(ref="H2:I3"))
    result = import_historical_workbook(database, source, budget_year=2024)
    snapshot = read_historical_workbook(database, result.extraction_id)
    absent = next(c for c in snapshot.cells if c["sheet"] == "Rådata" and c["cell"] == "I3")
    assert absent["state"] == "absent" and absent["formula_result"] is True
    assert absent["formula_present"] is False and absent["cached_value"] is None


@pytest.mark.parametrize("ref", ["", "H:H", "$H$2:$H$3", "H3:H2", "H0:H3", "H2:XFE3", "I2:I3"])
def test_invalid_or_unsupported_formula_regions_fail_with_preserved_original(
    tmp_path: Path, ref: str
) -> None:
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    raw = formula_range_workbook(ref=ref)
    source.write_bytes(raw)
    result = import_historical_workbook(database, source, budget_year=2024)
    assert result.status == "feilet" and result.extraction_id is None
    assert read_source(database, result.source_sha256 or "") == raw
    assert resume_historical_import(database, result.attempt_id) == result


def test_formula_region_area_is_bounded_without_expanding_declared_rectangle(
    tmp_path: Path,
) -> None:
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    raw = formula_range_workbook(ref="H2:H102")  # 101 positions, only two physical caches
    source.write_bytes(raw)
    result = import_historical_workbook(
        database, source, budget_year=2024, limits=replace(Limits(), positions=100)
    )
    assert result.status == "feilet" and result.extraction_id is None
    assert "formelområde" in (result.error or "")
    assert read_source(database, result.source_sha256 or "") == raw


def test_overlapping_formula_regions_are_rejected(tmp_path: Path) -> None:
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    source.write_bytes(
        mutate_member(
            formula_range_workbook(),
            "xl/worksheets/other.xml",
            b'<c r="I2" t="n"/>',
            b'<c r="I2"><f t="array" ref="H2:I3"/><v>3</v></c>',
        )
    )
    result = import_historical_workbook(database, source, budget_year=2024)
    assert result.status == "feilet" and "Overlappende" in (result.error or "")


def test_reused_shared_text_is_bounded_before_row_context_materialization(tmp_path: Path) -> None:
    raw = mutate_member(
        workbook_bytes(),
        "xl/worksheets/other.xml",
        b'<c r="H2"><v>14.768000000000001</v></c>',
        b'<c r="H2" t="s"><v>0</v></c>',
    )
    raw = mutate_member(
        raw, "xl/worksheets/other.xml", b'<c r="G2"><v>0</v></c>', b'<c r="G2" t="s"><v>0</v></c>'
    )
    stream = BytesIO(raw)
    with ZipFile(stream, "a") as archive:
        archive.writestr(
            "xl/sharedStrings.xml",
            '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>'
            + "long " * 25
            + "</t></si></sst>",
        )
    raw = stream.getvalue()
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    source.write_bytes(raw)
    result = import_historical_workbook(
        database, source, budget_year=2024, limits=replace(Limits(), resolved_chars=200)
    )
    assert result.status == "feilet" and result.extraction_id is None
    assert "Samlet" in (result.error or "")
    assert read_source(database, result.source_sha256 or "") == raw
    assert resume_historical_import(database, result.attempt_id) == result


def test_repeated_raw_context_has_a_separate_serialization_budget(tmp_path: Path) -> None:
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    source.write_bytes(
        workbook_bytes(
            extra='<row r="4"><c r="E4" t="inlineStr"><is><t>'
            + "context " * 100
            + "</t></is></c></row>"
        )
    )
    result = import_historical_workbook(
        database, source, budget_year=2024, limits=replace(Limits(), serialized_chars=30000)
    )
    assert result.status == "feilet" and result.extraction_id is None
    assert "serialisert" in (result.error or "")


def test_preserved_source_cells_candidates_and_repeat(tmp_path: Path) -> None:
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    raw = workbook_bytes(extra='<row r="4"><c r="Q4" t="e"><v>#N/A</v></c></row>')
    source.write_bytes(raw)
    result = import_historical_workbook(database, source, budget_year=2024)
    assert result.status == "lagret"
    snapshot = read_historical_workbook(database, result.extraction_id)
    extraction = snapshot.extraction
    assert read_source(database, extraction.artifact.sha256) == raw
    assert extraction.dataset == "historisk"
    assert extraction.source_system == "historisk-excel"
    assert len(extraction.observations) == 17  # 2 × 8 råfelt og én dokument-D
    by_cell = {(o.pointer.sheet, o.pointer.cell): o for o in extraction.observations}
    precise = by_cell[("Rådata", "H2")]
    assert precise.amount_kr is None
    assert precise.raw_value == "14.768000000000001"
    candidate = next(c for c in snapshot.candidates if c["observation_id"] == precise.id)
    assert candidate["candidate_kr"] == "14768000.000000001000000"
    assert candidate["controlled_kr"] is None
    assert len(by_cell[("Dokumentiltknyntning2024", "D2")].documents) == 3
    assert by_cell[("Dokumentiltknyntning2024", "D2")].event_id is None
    assert by_cell[("Dokumentiltknyntning2024", "D2")].organization.fdep_nr == "16"
    cells = {(c["sheet"], c["cell"]): c for c in snapshot.cells}
    assert cells[("Rådata", "G2")]["state"] == "value"
    assert cells[("Rådata", "I2")]["state"] == "blank"
    assert cells[("Rådata", "J2")]["state"] == "absent"
    assert cells[("Rådata", "Q4")]["disposition"] == "uavklart"
    assert len(snapshot.cells) == 73  # 42 + 30 + fysisk Q4, ingen XFD-ekspansjon
    moved = tmp_path / "moved.xlsx"
    moved.write_bytes(raw)
    repeated = import_historical_workbook(database, moved, budget_year=2024)
    assert repeated.attempt_id != result.attempt_id
    assert repeated.extraction_id == result.extraction_id
    assert read_historical_workbook(database, repeated.extraction_id) == snapshot
    assert resume_historical_import(database, result.attempt_id) == result


def test_package_reads_database_and_detects_actual_excel_changes(tmp_path: Path) -> None:
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    source.write_bytes(workbook_bytes())
    result = import_historical_workbook(database, source, budget_year=2024)
    package = tmp_path / "package"
    export_historical_package(database, result.extraction_id, package)
    assert verify_historical_package(database, result.extraction_id, package)["verified"] is True
    assert export_historical_package(database, result.extraction_id, package) == package
    wb = load_workbook(package / "kontroll.xlsx")
    assert {
        "Beløpsobservasjoner",
        "Celler",
        "Kolonnegrunnlag",
        "Kandidater",
        "Sammenligning",
    } <= set(wb.sheetnames)
    wb["Kandidater"]["A2"] = "tampered"
    wb.save(package / "kontroll.xlsx")
    wb.close()
    with pytest.raises(ValueError):
        verify_historical_package(database, result.extraction_id, package)


def test_staged_import_survives_missing_file_and_versions_keep_predecessor(tmp_path: Path) -> None:
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    source.write_bytes(workbook_bytes())
    attempt_id = stage_historical_workbook(database, source, budget_year=2024)
    source.unlink()
    result = resume_historical_import(database, attempt_id)
    original = read_historical_workbook(database, result.extraction_id)
    source.write_bytes(workbook_bytes(amount="12"))
    changed = import_historical_workbook(
        database, source, budget_year=2024, predecessor_id=result.extraction_id
    )
    assert changed.extraction_id != result.extraction_id
    assert (
        read_historical_workbook(database, changed.extraction_id).predecessor_id
        == result.extraction_id
    )
    new_rule = import_historical_workbook(
        database,
        source,
        budget_year=2024,
        predecessor_id=changed.extraction_id,
        interpretation_note="ny uttrykkelig kontrolltolkning",
    )
    assert new_rule.extraction_id != changed.extraction_id
    assert read_historical_workbook(database, result.extraction_id) == original


def mutate_member(raw: bytes, member: str, old: bytes, new: bytes) -> bytes:
    stream = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(stream, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item)
            target.writestr(
                item.filename, content.replace(old, new) if item.filename == member else content
            )
    return stream.getvalue()


@pytest.mark.parametrize(
    "raw",
    [
        b"not a zip",
        mutate_member(
            workbook_bytes(),
            "xl/workbook.xml",
            b"<sheets>",
            b'<!DOCTYPE foo [<!ENTITY x "a">]><sheets>',
        ),
        mutate_member(
            workbook_bytes(),
            "xl/_rels/workbook.xml.rels",
            b"worksheets/other.xml",
            b"../../outside.xml",
        ),
        mutate_member(workbook_bytes(), "xl/worksheets/other.xml", b'<c r="G2">', b'<c r="H2">'),
        mutate_member(workbook_bytes(), "xl/worksheets/other.xml", b"A1:N3", b"A1:XFD1048576"),
        mutate_member(
            workbook_bytes(), "xl/worksheets/other.xml", b"<v>0</v>", b"<v>0</v><v>1</v>"
        ),
    ],
)
def test_unsafe_inputs_preserve_received_bytes_without_observations(
    tmp_path: Path, raw: bytes
) -> None:
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    source.write_bytes(raw)
    result = import_historical_workbook(database, source, budget_year=2024)
    assert result.status == "feilet"
    assert result.extraction_id is None
    assert result.source_sha256 == hashlib.sha256(raw).hexdigest()
    assert read_source(database, result.source_sha256) == raw
    assert resume_historical_import(database, result.attempt_id) == result


def test_failed_file_read_has_no_invented_hash_and_limits_are_explicit(tmp_path: Path) -> None:
    source, database = tmp_path / "missing.xlsx", tmp_path / "pilot.sqlite"
    result = import_historical_workbook(database, source, budget_year=2024)
    assert result.status == "feilet" and result.source_sha256 is None
    source.write_bytes(workbook_bytes())
    oversized = import_historical_workbook(
        database, source, budget_year=2024, limits=replace(Limits(), max_bytes=10)
    )
    assert oversized.status == "feilet" and oversized.source_sha256 is None
    for limits in (
        replace(Limits(), members=2),
        replace(Limits(), member_bytes=50),
        replace(Limits(), inflated_bytes=100),
        replace(Limits(), positions=10),
        replace(Limits(), cells=2),
        replace(Limits(), xml_nodes=2),
        replace(Limits(), text_chars=4),
    ):
        failure = import_historical_workbook(database, source, budget_year=2024, limits=limits)
        assert failure.status == "feilet"


@pytest.mark.parametrize("amount", ["NaN", "Infinity", "1e99999999", "9" * 130, "1,5"])
def test_unknown_numeric_lexemes_are_retained_without_candidates(
    tmp_path: Path, amount: str
) -> None:
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    source.write_bytes(workbook_bytes(amount=amount))
    result = import_historical_workbook(database, source, budget_year=2024)
    snapshot = read_historical_workbook(database, result.extraction_id)
    obs = next(o for o in snapshot.extraction.observations if o.pointer.locator == "Rådata!H2")
    assert obs.raw_value == amount and obs.amount_kr is None
    assert (
        next(c for c in snapshot.candidates if c["observation_id"] == obs.id)["candidate_kr"]
        is None
    )


def test_atomic_publication_long_text_and_lost_ack(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source, database, package = (
        tmp_path / "input.xlsx",
        tmp_path / "pilot.sqlite",
        tmp_path / "package",
    )
    long_text = "=" + "preserved " * 7000
    source.write_bytes(
        workbook_bytes(
            extra=f'<row r="4"><c r="Q4" t="inlineStr"><is><t>{long_text}</t></is></c></row>'
        )
    )
    result = import_historical_workbook(database, source, budget_year=2024)
    rename = os.rename

    def lost_ack(src: str | Path, dst: str | Path) -> None:
        rename(src, dst)
        raise OSError("synthetic lost acknowledgement")

    with monkeypatch.context() as patch:
        patch.setattr(os, "rename", lost_ack)
        with pytest.raises(OSError):
            export_historical_package(database, result.extraction_id, package)
    export_historical_package(database, result.extraction_id, package)
    wb = load_workbook(package / "kontroll.xlsx")
    assert wb["Fortsettelser"].max_row > 1
    assert not any(c.data_type == "f" for s in wb for row in s for c in row)
    wb["Fortsettelser"]["E2"] = "tampered continuation"
    wb.save(package / "kontroll.xlsx")
    wb.close()
    manifest_path = package / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    excel_bytes = (package / "kontroll.xlsx").read_bytes()
    manifest["files"]["kontroll.xlsx"].update(
        sha256=hashlib.sha256(excel_bytes).hexdigest(), bytes=len(excel_bytes)
    )
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="Excel-celle"):
        verify_historical_package(database, result.extraction_id, package)


@pytest.mark.parametrize("tamper", ["jsonl", "manifest_id", "unexpected_file", "removed_file"])
def test_package_checks_database_binding_even_when_file_hash_is_updated(
    tmp_path: Path, tamper: str
) -> None:
    source, database, package = (
        tmp_path / "input.xlsx",
        tmp_path / "pilot.sqlite",
        tmp_path / "package",
    )
    source.write_bytes(workbook_bytes())
    result = import_historical_workbook(database, source, budget_year=2024)
    export_historical_package(database, result.extraction_id, package)
    manifest_path = package / "manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if tamper == "jsonl":
        path = package / "kandidater.jsonl"
        lines = path.read_text(encoding="utf-8").splitlines()
        row = json.loads(lines[0])
        row["candidate_kr"] = "42"
        lines[0] = json.dumps(row, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        raw = path.read_bytes()
        manifest["files"][path.name].update(sha256=hashlib.sha256(raw).hexdigest(), bytes=len(raw))
    elif tamper == "manifest_id":
        manifest["extraction_id"] = "wrong"
    elif tamper == "unexpected_file":
        (package / "extra.json").write_text("{}")
    else:
        (package / "observasjoner.jsonl").unlink()
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, sort_keys=True, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError):
        verify_historical_package(database, result.extraction_id, package)


def test_publication_failure_leaves_no_visible_package_and_does_not_overwrite_source(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source, database, package = (
        tmp_path / "input.xlsx",
        tmp_path / "pilot.sqlite",
        tmp_path / "package",
    )
    raw = workbook_bytes()
    source.write_bytes(raw)
    result = import_historical_workbook(database, source, budget_year=2024)

    def fail(src: str | Path, dst: str | Path) -> None:
        raise OSError("synthetic publication failure")

    with monkeypatch.context() as patch:
        patch.setattr(os, "rename", fail)
        with pytest.raises(OSError):
            export_historical_package(database, result.extraction_id, package)
    assert not package.exists()
    for target in (source, database, tmp_path):
        with pytest.raises(ValueError):
            export_historical_package(database, result.extraction_id, target)
    assert source.read_bytes() == raw
    assert resume_historical_import(database, result.attempt_id) == result
    export_historical_package(database, result.extraction_id, package)


def test_duplicate_zip_members_and_external_sheet_are_rejected(tmp_path: Path) -> None:
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    stream = BytesIO(workbook_bytes())
    with pytest.warns(UserWarning, match="Duplicate name"), ZipFile(stream, "a") as archive:
        archive.writestr("xl/workbook.xml", "duplicate")
    source.write_bytes(stream.getvalue())
    assert import_historical_workbook(database, source, budget_year=2024).status == "feilet"
    source.write_bytes(
        mutate_member(
            workbook_bytes(),
            "xl/_rels/workbook.xml.rels",
            b'Id="rId7"',
            b'Id="rId7" TargetMode="External"',
        )
    )
    assert import_historical_workbook(database, source, budget_year=2024).status == "feilet"


def test_shared_strings_and_exact_integer_limit_stay_raw(tmp_path: Path) -> None:
    raw = mutate_member(
        workbook_bytes(), "xl/worksheets/other.xml", b'<c r="H2">', b'<c r="H2" t="s">'
    )
    raw = mutate_member(raw, "xl/worksheets/other.xml", b"14.768000000000001", b"0")
    stream = BytesIO(raw)
    with ZipFile(stream, "a") as archive:
        archive.writestr(
            "xl/sharedStrings.xml",
            '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            "<si><r><t>=untrusted text</t></r></si></sst>",
        )
    source, database = tmp_path / "input.xlsx", tmp_path / "pilot.sqlite"
    source.write_bytes(stream.getvalue())
    result = import_historical_workbook(database, source, budget_year=2024)
    snapshot = read_historical_workbook(database, result.extraction_id)
    cell = next(c for c in snapshot.cells if c["sheet"] == "Rådata" and c["cell"] == "H2")
    assert cell["raw_v"] == "0" and cell["resolved_text"] == "=untrusted text"
    source.write_bytes(workbook_bytes(amount="9007199254740992"))
    result = import_historical_workbook(database, source, budget_year=2024)
    snapshot = read_historical_workbook(database, result.extraction_id)
    observation = next(
        o for o in snapshot.extraction.observations if o.pointer.locator == "Rådata!H2"
    )
    assert observation.amount_kr is None
    assert "utenfor_eksakt_excel_heltall" in {a.code for a in observation.anomalies}
