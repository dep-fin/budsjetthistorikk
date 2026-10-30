"""T6s avtalte offentlige original-, navigasjons- og filgrenser."""

import gzip
import hashlib
import json
import shutil
import subprocess
import sys
from dataclasses import asdict, replace
from pathlib import Path

import pytest

from innlesning.observations import encode
from innlesning.omtale import identified


@pytest.fixture(scope="module")
def t5_basis():
    from analyse.proposisjon16 import import_basis

    return import_basis(
        gzip.decompress(Path("tests/fixtures/t2/authentic-t1.json.gz").read_bytes()).decode(),
        Path("."),
    )


@pytest.fixture(scope="module")
def t6_basis(t5_basis):
    from analyse import innstilling95 as t6
    from analyse.proposisjon16 import build_report, report_to_json
    from innlesning.visningspilot_model import PilotSelection

    return t6.import_basis(report_to_json(build_report(t5_basis, PilotSelection())), Path("."))


def test_complete_fixed_originals_preserve_one_t5_and_reconstruct_six_decisions(t5_basis):
    from analyse.visningspilot import source_input
    from innlesning.innstilling95 import import_basis, reconstruct

    basis = import_basis(t5_basis, source_input(t5_basis.t1_basis.o2b_json), Path("."))
    assert basis.t5_basis == t5_basis
    assert len(basis.observations) == 216
    assert len(basis.documents) == 6
    assert len(basis.recipes) == 6
    for recipe in basis.recipes:
        assert (
            reconstruct(basis.t5_basis, basis.sources, recipe)
            == Path(
                "docs/research/to-visninger-37/originals/forgrening/"
                + str(40034220 + int(recipe.document_key.removeprefix("vedtak")) - 359)
                + ".xml"
            ).read_bytes()
        )
    assert len(basis.authorizations) == 16
    assert sum(a.roman == "IV" for a in basis.authorizations) == 2
    assert {a.frame_kr for a in basis.authorizations if a.roman == "IV"} == {16_300_000}


def test_both_entries_missing_stage_zero_and_exact_return(t5_basis):
    from analyse import innstilling95 as t6
    from analyse.proposisjon16 import build_report, report_to_json
    from innlesning.visningspilot_model import PilotSelection

    basis = t6.import_basis(report_to_json(build_report(t5_basis, PilotSelection())), Path("."))
    data = t6.contents(basis)
    assert len(data.observations) == 339
    assert len({o.step_id for o in data.observations if o.figure_type == "stortingsvedtak"}) == 38
    selection = PilotSelection(years=(2024,), accounts=(("1420", "32"), ("4400", "2")))
    report = t6.build_report(basis, selection)
    assert len(report.view.observation_ids) == 15
    assert len(report.view.adopted_event_ids) == 2
    gaps = [c for c in report.view.comparisons if c.difference_kr is None]
    assert len(gaps) == 3 and all(c.account == "1420/32" and c.left_id is None for c in gaps)
    zero = [
        o
        for o in data.observations
        if o.chapter == "4400" and o.meaning == "etter_belop" and o.post == "2"
    ]
    assert len(zero) == 3 and all(o.amount_kr == 0 for o in zero)
    ids = {d.key: d.id for d in data.documents}
    opened = t6.open_document(report, ids["inns95"])
    assert len(opened.view.observation_ids) == 108
    assert t6.go_back(t6.report_from_json(t6.report_to_json(opened))) == report
    assert t6.go_back(t6.open_accounts(opened, selection)) == t6.build_report(
        basis, opened.selection
    )
    for account in (("1429", "1"), ("1533", "45")):
        view = t6.build_report(basis, PilotSelection(years=(2024,), accounts=(account,))).view
        assert ids["inns95"] not in view.document_ids
        assert ids["vedtak359"] not in view.document_ids


def test_saved_pair_and_offline_originals_preserve_visible_zero_and_all_fields(tmp_path, t5_basis):
    from openpyxl import load_workbook

    from analyse import innstilling95 as t6
    from analyse.innstilling95_excel import export_originals, export_pair, read_excel
    from analyse.proposisjon16 import build_report, report_to_json
    from innlesning.visningspilot_model import PilotSelection

    basis = t6.import_basis(report_to_json(build_report(t5_basis, PilotSelection())), Path("."))
    report = t6.build_report(basis, PilotSelection(years=(2024,), accounts=(("4400", "2"),)))
    frozen, excel = tmp_path / "zero.json", tmp_path / "zero.xlsx"
    export_pair(report, frozen, excel)
    assert t6.read_json(frozen) == read_excel(excel) == report
    book = load_workbook(excel)
    try:
        assert [r[6].value for r in book["Forløp"].iter_rows(min_row=2)] == [518000, -518000, 0] * 3
        assert book["Forløp"]["G10"].data_type == "n"
        assert book["Utvalg"]["B11"].value == 114
        assert book["Utvalg"]["B12"].value == 38
        assert len(list(book["Fullmakter"].iter_rows(min_row=2))) == 25
    finally:
        book.close()
    directory = tmp_path / "originals"
    export_originals(report, directory)
    for number in range(40034220, 40034226):
        assert (directory / "vedtak" / (str(number) + ".xml")).read_bytes() == Path(
            "docs/research/to-visninger-37/originals/forgrening/" + str(number) + ".xml"
        ).read_bytes()


def test_all_fields_against_independent_original_only_oracle(t6_basis):
    from analyse import innstilling95 as t6
    from innlesning.visningspilot_model import PilotSelection

    oracle = json.loads(
        Path("docs/research/innstilling95-64/independent/original-fields.json").read_text("utf8")
    )
    fields = {
        (o.source_unit_id, o.chapter + "/" + o.post, o.meaning): o for o in t6_basis.observations
    }
    for key, source_key in (("inns95", "innstilling_rows"), ("vedtak359", "decision_rows")):
        for row in oracle[source_key]:
            for meaning, value_key in (
                ("for_belop", "before_kr"),
                ("endring", "change_kr"),
                ("etter_belop", "after_kr"),
            ):
                o = fields[key, row["account"], meaning]
                assert o.amount_kr == row["values"][value_key]
                assert o.budget_year == row["budget_year"] == 2024
                assert o.pointer.artifact_id == row["pointers"][value_key]["source_sha256"]
                assert o.raw_value == row["pointers"][value_key]["raw_number"]
    report = t6.build_report(
        t6_basis,
        PilotSelection(
            years=(2024,),
            chapters=tuple(
                str(i)
                for i in (
                    1400,
                    1410,
                    1411,
                    1412,
                    1420,
                    1423,
                    1425,
                    1428,
                    1429,
                    1471,
                    1473,
                    1474,
                    1481,
                    1482,
                    4400,
                    4420,
                    4471,
                    4481,
                )
            ),
        ),
    )
    assert (
        len(report.view.observation_ids),
        len(report.view.adopted_observation_ids),
        len(report.view.adopted_event_ids),
    ) == (339, 114, 38)
    assert len([c for c in report.view.comparisons if c.difference_kr == 0]) == 225
    assert len([c for c in report.view.comparisons if c.difference_kr is None]) == 3
    assert (
        len({a.document_id for a in t6_basis.authorizations if a.role.startswith("vedtak:")}) == 5
    )
    assert all(o.chapter != "1429" for o in t6_basis.observations)


@pytest.fixture(scope="module")
def saved_cases(tmp_path_factory, t6_basis):
    from analyse import innstilling95 as t6
    from analyse.innstilling95_excel import export_pair
    from innlesning.visningspilot_model import PilotSelection

    directory = tmp_path_factory.mktemp("t6-files")
    ids = {d.key: d.id for d in t6.contents(t6_basis).documents}
    cases = {
        "post32": PilotSelection(years=(2024,), accounts=(("1420", "32"),)),
        "utgift": PilotSelection(years=(2024,), accounts=(("1400", "1"),)),
        "null": PilotSelection(years=(2024,), accounts=(("4400", "2"),)),
        "fullmakt": PilotSelection(years=(2024,), accounts=(("1420", "66"),)),
        "delt-ramme": PilotSelection(
            years=(2024,), accounts=(("1420", "39"), ("1420", "69"), ("1420", "79"))
        ),
        "1429-1": PilotSelection(years=(2024,), accounts=(("1429", "1"),)),
        "1429-22": PilotSelection(years=(2024,), accounts=(("1429", "22"),)),
        "1429-60": PilotSelection(years=(2024,), accounts=(("1429", "60"),)),
        "1533-45": PilotSelection(years=(2024,), accounts=(("1533", "45"),)),
        "flere": PilotSelection(
            years=(2024,), chapters=("1429",), accounts=(("1420", "32"), ("4400", "2"))
        ),
        "kapittel": PilotSelection(years=(2024,), chapters=("4400",)),
        "tomt": PilotSelection(years=(2024,)),
        "mangler": PilotSelection(years=(2024,), accounts=(("999", "99"),)),
        "feil-aar": PilotSelection(years=(2025,), chapters=("1420",)),
        "uten-aar": PilotSelection(chapters=("1420",)),
    }
    cases.update(
        {
            key: PilotSelection(mode="document", document_id=ids[key])
            for key in (
                "prop16",
                "inns95",
                "vedtak359",
                "vedtak360",
                "vedtak361",
                "vedtak362",
                "vedtak363",
                "vedtak364",
            )
        }
    )
    reports = {name: t6.build_report(t6_basis, selection) for name, selection in cases.items()}
    reports["konto-dokument"] = t6.open_document(reports["flere"], ids["inns95"])
    reports["dokument-konto"] = t6.open_accounts(reports["inns95"], cases["flere"])
    result = {}
    for name, report in reports.items():
        frozen, excel = directory / (name + ".json"), directory / (name + ".xlsx")
        export_pair(report, frozen, excel)
        result[name] = report, frozen, excel
    return result


def test_actual_closed_saved_matrix_and_navigation(saved_cases, t5_basis):
    from analyse import innstilling95 as t6
    from analyse.innstilling95_excel import read_excel

    expected = {
        "post32": (6, 3, 1),
        "utgift": (9, 3, 1),
        "null": (9, 3, 1),
        "1429-1": (9, 3, 1),
        "1429-22": (9, 3, 1),
        "flere": (33, 12, 4),
        "kapittel": (18, 6, 2),
        "prop16": (111, 0, 0),
        "inns95": (108, 0, 0),
        "vedtak359": (108, 108, 36),
    }
    for name, (report, frozen, excel) in saved_cases.items():
        actual = t6.read_json(frozen)
        assert read_excel(excel) == actual == report
        assert actual.basis.t5_basis == t5_basis
        if name not in {"delt-ramme", "konto-dokument", "dokument-konto"}:
            assert (
                len(actual.view.observation_ids),
                len(actual.view.adopted_observation_ids),
                len(actual.view.adopted_event_ids),
            ) == expected.get(name, (0, 0, 0))
        assert "uberiktiget95_og_rettelsesmelding_ikke_funnet" in actual.basis.issues
    assert t6.go_back(t6.read_json(saved_cases["konto-dokument"][1])) == saved_cases["flere"][0]
    assert t6.go_back(read_excel(saved_cases["dokument-konto"][2])) == saved_cases["inns95"][0]


def _poison(report, mutation):
    basis = report.basis
    if mutation in {"amount", "role", "year", "account", "zero", "null32", "monetary"}:
        old = next(
            o
            for o in basis.observations
            if (o.chapter, o.post, o.meaning) == ("4400", "2", "etter_belop")
        )
        changes = {
            "amount": {"amount_kr": 1},
            "role": {"proposal_role": "alternativ"},
            "year": {"budget_year": 2025},
            "account": {"chapter": "1429", "post": "22"},
            "zero": {"amount_kr": None},
            "null32": {"chapter": "1420", "post": "32", "source_unit_id": "prop16", "amount_kr": 0},
            "monetary": {"figure_type": "stortingsvedtak", "information_type": "fullmakt"},
        }[mutation]
        new = replace(old, **changes)
        contexts = tuple(
            identified("innstilling95-context", replace(c, observation_id=new.id))
            if c.observation_id == old.id
            else c
            for c in basis.contexts
        )
        basis = replace(
            basis,
            observations=tuple(new if o == old else o for o in basis.observations),
            contexts=contexts,
        )
    elif mutation == "frame":
        auth = next(a for a in basis.authorizations if a.roman == "IV")
        multiplied = tuple(
            identified("innstilling95-authorization", replace(auth, accounts=(account,)))
            for account in auth.accounts
        )
        basis = replace(
            basis, authorizations=tuple(a for a in basis.authorizations if a != auth) + multiplied
        )
    elif mutation in {"roman", "branch"}:
        relation = basis.relations[0]
        change = (
            {"accounts": relation.accounts + ("1533/45", "1429/1")}
            if mutation == "branch"
            else {"role": "95II→vedtak363"}
        )
        basis = replace(
            basis,
            relations=(
                identified("innstilling95-relation", replace(relation, **change)),
                *basis.relations[1:],
            ),
        )
    elif mutation == "old":
        old_t5 = basis.t5_basis
        old_t1 = replace(old_t5.t1_basis, issues=())
        basis = replace(
            basis,
            t5_basis=identified(
                "proposisjonsgrunnlag16-1",
                replace(old_t5, t1_basis=identified("visningsgrunnlag-1", old_t1)),
            ),
        )
    elif mutation == "coverage":
        basis = replace(
            basis,
            coverage=(
                identified(
                    "innstilling95-coverage",
                    replace(basis.coverage[0], status="full_faglig_strukturering"),
                ),
                *basis.coverage[1:],
            ),
        )
    elif mutation == "correction":
        basis = replace(basis, issues=("uberiktiget95_rekonstruert",))
    elif mutation == "hidden":
        basis = replace(basis, issues=())
    else:
        raise AssertionError(mutation)
    return identified(
        "innstilling95-1", replace(report, basis=identified("innstilling95-basis", basis))
    )


@pytest.mark.parametrize(
    "mutation",
    [
        "amount",
        "role",
        "year",
        "account",
        "zero",
        "null32",
        "monetary",
        "frame",
        "roman",
        "branch",
        "old",
        "coverage",
        "correction",
        "hidden",
    ],
)
def test_rehashed_semantic_poison_passes_container_before_both_readers_reject(
    tmp_path, saved_cases, mutation
):
    from analyse import innstilling95 as t6
    from analyse.innstilling95_excel import SHEET_NAMES, _book, read_excel
    from analyse.omtale_excel_guard import inspect_excel

    # Tomt utvalg holder angrepet uavhengig av en ærlig avledet radvisning.
    poisoned = _poison(saved_cases["tomt"][0], mutation)
    raw = encode({"format": "innstilling95-1", "report": asdict(poisoned)})
    path = tmp_path / "synthetic.json"
    path.write_text(raw, "utf8")
    with pytest.raises(ValueError):
        t6.read_json(path)
    book = _book(poisoned, raw)
    excel = tmp_path / "synthetic.xlsx"
    book.save(excel)
    book.close()
    assert (
        inspect_excel(
            excel,
            sheet_names=SHEET_NAMES,
            snapshot_name="Visningsgrunnlag",
            max_expanded_bytes=128 * 1024 * 1024,
        )[0]
        == raw
    )
    with pytest.raises(ValueError):
        read_excel(excel)


@pytest.mark.parametrize("condition", ["missing", "failed", "corrupt", "contradictory"])
def test_optional_t6_errors_preserved_beside_real_ordinary_new_registration(
    tmp_path, t5_basis, condition
):
    from analyse import innstilling95 as t6
    from analyse.innstilling95_excel import export_pair, read_excel
    from analyse.proposisjon16 import build_report, report_to_json
    from innlesning.innstilling95 import ROOT, SOURCE_CATALOG
    from innlesning.observations import read_extraction
    from innlesning.visningspilot_model import PilotSelection

    root = tmp_path / "optional"
    for name in SOURCE_CATALOG:
        path = root / ROOT / name
        path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(ROOT + name, path)
    path = root / ROOT / "40034220.json"
    if condition in {"missing", "failed"}:
        path.unlink()
        if condition == "failed":
            path.mkdir()
    else:
        path.write_bytes(
            b"synthetic invalid JSON"
            if condition == "corrupt"
            else path.read_bytes().replace(b"2172545", b"2172546")
        )
    basis = t6.import_basis(report_to_json(build_report(t5_basis, PilotSelection())), root)
    assert basis.sources[0].status == (
        "mangler" if condition in {"missing", "failed"} else "motstrid"
    )
    report = t6.build_report(basis, PilotSelection(years=(2025,), chapters=("1420",)))
    frozen, excel = tmp_path / "error.json", tmp_path / "error.xlsx"
    export_pair(report, frozen, excel)
    assert t6.read_json(frozen) == read_excel(excel) == report
    database = tmp_path / "ordinary.sqlite"

    def cli(*args):
        result = subprocess.run(
            [sys.executable, "-m", *args], capture_output=True, text=True, encoding="utf8"
        )
        assert result.returncode in (0, 2), result.stdout + result.stderr
        return json.loads(result.stdout)

    summary = cli(
        "innlesning",
        "synkroniser",
        "--database",
        str(database),
        "--sesjon",
        "2024-2025",
        "--historisk-proeve",
        "--vedtak-id",
        "40034229",
        "--response",
        "tests/fixtures/b6/40034229.xml",
        "--metadata",
        "tests/fixtures/b6/40034229.json",
    )
    assert summary["antall"]["ny"] == 1 and not summary["feil"]
    extraction_id = cli(
        "innlesning",
        "tolk",
        "--database",
        str(database),
        "--revisjon",
        summary["vedtak"][0]["revision_id"],
        "--tolkning",
        "storting-belop-v2",
    )["belopsuttrekk"]
    extraction = read_extraction(database, extraction_id)
    assert len(extraction.observations) == 222
    assert len({o.step_id for o in extraction.observations}) == 74
    ordinary_excel = tmp_path / "ordinary.xlsx"
    cli(
        "analyse",
        "--database",
        str(database),
        "--belopsuttrekk",
        extraction_id,
        "--excel",
        str(ordinary_excel),
    )
    assert ordinary_excel.exists()


def test_offline_cli_no_network_database_and_no_clobber(tmp_path, saved_cases, monkeypatch):
    import socket
    import sqlite3

    from analyse import innstilling95 as t6
    from analyse.innstilling95_excel import export_pair, read_excel

    report, frozen, excel = saved_cases["null"]
    before = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (frozen, excel)]
    with pytest.raises(ValueError):
        export_pair(report, frozen, excel)
    assert before == [hashlib.sha256(p.read_bytes()).hexdigest() for p in (frozen, excel)]

    def forbidden(*args, **kwargs):
        raise AssertionError("Offline skal ikke åpne nett/database")

    monkeypatch.setattr(socket, "socket", forbidden)
    monkeypatch.setattr(sqlite3, "connect", forbidden)
    assert t6.read_json(frozen) == read_excel(excel) == report
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "analyse.innstilling95",
            "--frosset",
            str(frozen),
            "--rot",
            str(tmp_path / "does-not-exist"),
            "--json",
            str(tmp_path / "offline.json"),
            "--excel",
            str(tmp_path / "offline.xlsx"),
        ],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert (
        t6.read_json(tmp_path / "offline.json") == read_excel(tmp_path / "offline.xlsx") == report
    )


@pytest.mark.parametrize(
    "mutation", ["numeric", "formula", "link", "hide", "style", "theme", "zip", "doctype", "chunk"]
)
def test_actual_excel_presentation_and_package_defenses(tmp_path, saved_cases, mutation):
    from zipfile import ZipFile

    from openpyxl import load_workbook

    from analyse.innstilling95_excel import read_excel

    path = tmp_path / "synthetic-presentation.xlsx"
    source = saved_cases["null"][2]
    if mutation in {"zip", "doctype", "theme"}:
        with ZipFile(source) as archive:
            files = {n: archive.read(n) for n in archive.namelist()}
        if mutation == "zip":
            files["unexpected.bin"] = b"synthetic"
        elif mutation == "doctype":
            files["xl/workbook.xml"] = (
                b'<!DOCTYPE workbook [<!ENTITY x "unsafe">]>' + files["xl/workbook.xml"]
            )
        else:
            files["xl/theme/theme1.xml"] = files["xl/theme/theme1.xml"].replace(
                b"000000", b"FFFFFF"
            )
        with ZipFile(path, "w") as archive:
            for name, raw in files.items():
                archive.writestr(name, raw)
    else:
        book = load_workbook(source)
        if mutation == "numeric":
            book["Forløp"]["G4"] = "0"
        elif mutation == "formula":
            book["Fulltekst"]["K2"] = "=1+1"
        elif mutation == "link":
            book["Forløp"]["B2"].hyperlink.location = "'Dokumenter'!B3"
        elif mutation == "hide":
            book["T6-dekning"].sheet_state = "hidden"
        elif mutation == "style":
            book["Forløp"]["G2"].number_format = ";;;"
        else:
            book["Visningsgrunnlag"]["A2"] = 2
        book.save(path)
        book.close()
    with pytest.raises(ValueError):
        read_excel(path)


@pytest.mark.parametrize(
    "field,limit",
    [
        ("sources", 64),
        ("documents", 16),
        ("inventory", 4096),
        ("statements", 1024),
        ("observations", 512),
        ("relations", 4096),
    ],
)
def test_own_resource_boundaries_preserve_old_profiles(t6_basis, field, limit):
    from analyse.innstilling95 import validate_basis

    old = t6_basis.t5_basis.t1_basis
    old_count = len(getattr(old, field))
    if field in {"sources", "observations"}:
        old_count += len(getattr(t6_basis.t5_basis, field))
    if field == "sources":
        from analyse.visningspilot import source_input

        old_count += len(source_input(old.o2b_json).basis.sources)
    values = getattr(t6_basis, field)
    at_limit = replace(t6_basis, **{field: (values[0],) * (limit - old_count)})
    with pytest.raises(ValueError, match="avviker|kildekatalog"):
        validate_basis(at_limit)
    over_limit = replace(t6_basis, **{field: (values[0],) * (limit + 1 - old_count)})
    with pytest.raises(ValueError, match="ressursprofil"):
        validate_basis(over_limit)


def test_authentic_old_t5_and_t1_files_are_unchanged_and_read_by_old_readers():
    from analyse import innstilling95 as t6
    from analyse import proposisjon16 as t5
    from analyse import visningspilot as t1

    first = Path("tests/fixtures/t2/authentic-t1.json.gz").read_bytes()
    second = Path("tests/fixtures/t6/authentic-t5.json.gz").read_bytes()
    meta = json.loads(Path("tests/fixtures/t6/authentic-t5.meta.json").read_text("utf8"))
    assert hashlib.sha256(second).hexdigest() == meta["gzip_sha256"]
    raw_t5 = gzip.decompress(second)
    assert hashlib.sha256(raw_t5).hexdigest() == meta["raw_sha256"]
    old_t1 = t1.report_from_json(gzip.decompress(first).decode())
    old_t5 = t5.report_from_json(raw_t5.decode())
    assert old_t5.basis.t1_basis == old_t1.basis
    basis = t6.import_basis(raw_t5.decode())
    assert basis.t5_basis == old_t5.basis
    assert basis.t5_basis.t1_basis == old_t1.basis
    assert all(
        o.round == "uavklart"
        for o in old_t1.basis.observations
        if o.figure_type == "stortingsvedtak"
    )
    assert Path("tests/fixtures/t6/authentic-t5.json.gz").read_bytes() == second
    assert Path("tests/fixtures/t2/authentic-t1.json.gz").read_bytes() == first


def test_strict_json_and_64_mib_boundary():
    from analyse.innstilling95 import MAX_REPORT_BYTES, report_from_json

    for raw in (
        '{"format":"innstilling95-1","format":"innstilling95-1","report":{}}',
        '{"format":"innstilling95-1","report":NaN}',
        '{"format":"innstilling95-1","report":Infinity}',
        '{"format":"innstilling95-1","report":{},"extra":true}',
    ):
        with pytest.raises(ValueError):
            report_from_json(raw)
    with pytest.raises(json.JSONDecodeError):
        report_from_json(" " * MAX_REPORT_BYTES)
    with pytest.raises(ValueError, match="overstiger"):
        report_from_json(" " * (MAX_REPORT_BYTES + 1))


def test_new_raw_source_16_mib_gate_keeps_failure_visible(tmp_path, t5_basis):
    from analyse.visningspilot import source_input
    from innlesning.innstilling95 import ROOT, import_basis
    from innlesning.visningspilot import MAX_SOURCE_BYTES

    path = tmp_path / ROOT / "40034220.json"
    path.parent.mkdir(parents=True)
    path.write_bytes(b" " * MAX_SOURCE_BYTES)
    basis = import_basis(t5_basis, source_input(t5_basis.t1_basis.o2b_json), tmp_path)
    assert basis.sources[0].status == "motstrid"
    with path.open("ab") as handle:
        handle.write(b" ")
    oversized = import_basis(t5_basis, source_input(t5_basis.t1_basis.o2b_json), tmp_path)
    assert oversized.sources[0].status == "mangler"
    assert json.loads(oversized.sources[0].metadata_json)["error_type"] == "ValueError"
    assert any("40034220.json" in issue for issue in oversized.issues)


def test_missing_plenum_keeps_raw_decision_but_never_claims_authorization_coverage(
    tmp_path, t5_basis
):
    from analyse import innstilling95 as t6
    from analyse.innstilling95_excel import export_pair, read_excel
    from analyse.proposisjon16 import build_report, report_to_json
    from innlesning.innstilling95 import ROOT, SOURCE_CATALOG
    from innlesning.visningspilot_model import PilotSelection

    for name in SOURCE_CATALOG:
        if name == "votering-181224.html":
            continue
        path = tmp_path / ROOT / name
        path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(ROOT + name, path)
    basis = t6.import_basis(report_to_json(build_report(t5_basis, PilotSelection())), tmp_path)
    document = next(d for d in basis.documents if d.key == "vedtak359")
    assert document.raw_coverage == "komplett_rekonstruert"
    assert document.structured_coverage == "ikke_strukturert"
    coverage = next(c for c in basis.coverage if c.document_id == document.id)
    assert coverage.status == "tabellstadium_mangler"
    assert "komplett" not in coverage.scope
    assert "vedtakelsesbevis_mangler" in basis.issues
    report = t6.build_report(basis, PilotSelection(mode="document", document_id=document.id))
    assert not report.view.observation_ids
    frozen, excel = tmp_path / "no-plenum.json", tmp_path / "no-plenum.xlsx"
    export_pair(report, frozen, excel)
    assert t6.read_json(frozen) == read_excel(excel) == report
