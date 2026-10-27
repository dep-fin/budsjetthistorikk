"""T2s offentlige original/utvalg/offline-grense, med autentiske og merkede feilprøver."""

import gzip
import hashlib
import json
import shutil
import subprocess
import sys
from collections import Counter
from copy import copy
from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill

from analyse.komitepilot import (
    build_report,
    compare_observations,
    contents,
    go_back,
    import_basis,
    open_accounts,
    open_document,
    read_json,
    report_from_json,
    report_to_json,
)
from analyse.komitepilot_excel import export_originals, export_pair, read_excel
from innlesning.komitepilot import CATALOG, PLENUM, K, derived_xml, payload_bytes
from innlesning.observations import encode, identity
from innlesning.visningspilot_model import PilotSelection


def test_standalone_originals_bind_nine_fields_and_one_adopted_event():
    basis = import_basis(Path("."))
    report = build_report(basis, PilotSelection(years=(2024,), accounts=(("315", "78"),)))
    assert len(report.view.observation_ids) == 9
    assert len(report.view.adopted_observation_ids) == 3
    assert len(report.view.adopted_event_ids) == 1
    assert [o.amount_kr for o in basis.observations] == [
        29_050_000,
        119_000_000,
        148_050_000,
        29_050_000,
        129_000_000,
        158_050_000,
        29_050_000,
        129_000_000,
        158_050_000,
    ]
    assert [(c.meaning, c.difference_kr) for c in report.view.comparisons] == [
        ("for_belop", 0),
        ("for_belop", 0),
        ("endring", 10_000_000),
        ("endring", 0),
        ("etter_belop", 10_000_000),
        ("etter_belop", 0),
    ]


def test_saved_json_excel_and_original_byte_registry(tmp_path):
    import base64
    import json
    from collections import Counter

    from openpyxl import load_workbook

    from analyse.komitepilot import read_json
    from analyse.komitepilot_excel import export_pair, read_excel

    basis = import_basis(Path("."))
    report = build_report(basis, PilotSelection(years=(2024,), accounts=(("315", "78"),)))
    frozen, excel = tmp_path / "standalone.json", tmp_path / "standalone.xlsx"
    export_pair(report, frozen, excel)
    assert read_json(frozen) == read_excel(excel) == report
    data = json.loads(frozen.read_text("utf8"))
    assert not data["report"]["basis"]["legacy_json"]
    payloads = data["report"]["basis"]["payloads"]
    assert len({p["sha256"] for p in payloads}) == len(payloads)
    assert all(p["legacy_source_id"] is None for p in payloads)
    book = load_workbook(excel)
    try:
        assert book.sheetnames[:5] == [
            "Forløp",
            "Dokumenter",
            "Omtale",
            "Sammenligninger",
            "Dekning og avvik",
        ]
        assert "Originaler" not in book.sheetnames
        assert book["Forløp"].max_row == 4
        assert [book["Forløp"].cell(r, 8).value for r in (2, 3, 4)] == [
            119_000_000,
            129_000_000,
            129_000_000,
        ]
        assert book["Forløp"]["C2"].hyperlink.location.startswith("'Dokumenter'!")
        comparison = book["Sammenligninger"]
        assert comparison["C4"].value == "Prop. 25 S (2024–2025)"
        assert comparison["D4"].value == "Innst. 85 S (2024–2025)"
        assert [comparison.cell(4, column).value for column in (5, 6, 7)] == [
            119_000_000,
            129_000_000,
            10_000_000,
        ]
        assert comparison["C4"].hyperlink.location.startswith("'Dokumenter'!")
        snapshot = "".join(
            str(row[1]) for row in book["Offlinegrunnlag"].iter_rows(min_row=2, values_only=True)
        )
        assert snapshot == frozen.read_text("utf8")
        saved_text = {}
        saved_html = {}
        for row in book["Omtale"].iter_rows(min_row=2, values_only=True):
            saved_text.setdefault(row[0], []).append(row[8] or "")
            saved_html.setdefault(row[0], []).append(row[9] or "")
        for statement in basis.statements:
            assert "".join(saved_text[statement.id]) == statement.text
            assert "".join(saved_html[statement.id]) == statement.html
        for p in payloads:
            content = p["content_base64"]
            assert snapshot.count(content) == 1
            assert base64.b64decode(content)
        assert Counter(x.source_unit_id for x in basis.observations) == {
            "prop25": 3,
            "inns85": 3,
            "vedtak208": 3,
        }
    finally:
        book.close()


@pytest.fixture(scope="module")
def authentic_basis():
    return import_basis(Path("."))


@pytest.fixture(scope="module")
def union_basis():
    archive = Path("tests/fixtures/t2/authentic-t1.json.gz")
    raw = gzip.decompress(archive.read_bytes())
    assert (
        hashlib.sha256(raw).hexdigest()
        == "99e871e0ab780cbb548e15699f3093621d0f25229e3fa06c29cfdb5552803f53"
    )
    return import_basis(Path("."), t1_json=raw.decode("utf8"))


@pytest.fixture(scope="module")
def saved_pair(tmp_path_factory, authentic_basis):
    directory = tmp_path_factory.mktemp("saved-t2")
    report = build_report(authentic_basis, PilotSelection(years=(2024,), accounts=(("315", "78"),)))
    export_pair(report, directory / "post.json", directory / "post.xlsx")
    return directory, report


def test_source_roles_inventory_fullmakt_and_post_exception_are_separate(authentic_basis):
    b = authentic_basis
    docs = {d.key: d for d in b.documents}
    assert [len(docs[k].accounts) for k in ("prop25", "inns85", "vedtak208")] == [30, 30, 30]
    assert all(d.structured_coverage.startswith("delvis") for d in b.documents[:3])
    assert Counter(i.structure_status for i in b.inventory) == {
        "strukturert": 3,
        "ikke_strukturert": 87,
    }
    statements = {s.role: s for s in b.statements}
    extra = statements["flertallsbegrunnelse_LNU_tillegg"]
    assert extra.actor == "Ap_Sp_SV" and extra.accounts == ("315/78",)
    assert "10 mill. kroner ut over regjeringens forslag" in " ".join(extra.text.split())
    assert extra.hierarchy == (
        "2. Komiteens merknader",
        "2.2 Frivillighetstiltak",
        "2.2.1 Kap. 315 post 78 Frivillighetstiltak",
    )
    exception = statements["tilraading_med_postunntak"]
    assert "med unntak av kap. 315 post 78" in " ".join(exception.text.split())
    assert statements["generell_merknad"].actor == "H_V"
    assert statements["generell_merknad"].accounts == ()
    prop = next(s for s in b.statements if s.role == "regjeringens_begrunnelse" and s.accounts)
    assert prop.hierarchy[-2:] == ("Kap. 315 Frivillighetsformål", "Post 78 Frivillighetstiltak")
    vote = statements["postvis_votering_63_35_194948"]
    text = " ".join(vote.text.split())
    assert all(
        word in text for word in ("315 post 78", "63 mot 35", "19.49.48", "Fremskrittspartiet")
    )
    fullmakt = statements["vedtatt_fullmakt_II"]
    assert fullmakt.document_id == docs["vedtak209"].id and fullmakt.accounts == ()
    assert "342 mill. kroner" in " ".join(fullmakt.text.split())
    assert not any(o.source_unit_id == "vedtak209" for o in b.observations)
    assert not any(o.proposal_role == "alternativ" for o in b.observations)
    opposition = next(r for r in b.relations if r.role == "varslet_motstand_H_FrP_V")
    assert opposition.accounts == ("315/78",) and opposition.status == "kildeverifisert"
    assert len({o.id for o in b.observations}) == 9


def test_union_selection_navigation_offline_and_legacy_objects(tmp_path, union_basis):
    from analyse.visningspilot import report_from_json as read_t1

    b = union_basis
    selection = PilotSelection(years=(2024,), accounts=(("315", "78"), ("1429", "1")))
    start = build_report(b, selection)
    assert (
        len(start.view.observation_ids),
        len(start.view.adopted_observation_ids),
        len(start.view.adopted_event_ids),
    ) == (18, 6, 2)
    assert {
        (o.chapter, o.post) for o in contents(b).observations if o.id in start.view.observation_ids
    } == {("315", "78"), ("1429", "1")}
    old = read_t1(b.legacy_json)
    assert [o for o in contents(b).observations if o.source_system != "komitepilot"] == list(
        old.basis.observations
    )
    assert all(
        o.round == "uavklart" for o in old.basis.observations if o.figure_type == "stortingsvedtak"
    )
    shared = next(
        p
        for p in b.payloads
        if p.sha256 == "22170a60f769e89642b2e52525899903020de25938aa2b566ea28d11f85fb2e5"
    )
    assert shared.content_base64 is None and shared.legacy_source_id
    variants = {
        "union": start,
        "empty": build_report(b, PilotSelection()),
        "chapter": build_report(b, PilotSelection(years=(2024,), chapters=("315",))),
        "years": build_report(b, PilotSelection(years=(2024, 2025), accounts=selection.accounts)),
    }
    docs = {d.key: d for d in contents(b).documents}
    for key in ("prop25", "inns85", "vedtak208", "vedtak209", "prop16", "inns81"):
        document = open_document(start, docs[key].id)
        variants["doc-" + key] = document
        assert go_back(document) == start
        account = open_accounts(document, selection)
        variants["from-" + key] = account
        assert go_back(account) == build_report(b, document.selection)
    assert len(variants["doc-prop25"].view.inventory_ids) == 30
    assert len(variants["doc-inns85"].view.inventory_ids) == 30
    assert len(variants["doc-vedtak208"].view.observation_ids) == 3
    assert not variants["doc-vedtak209"].view.observation_ids
    for name, report in variants.items():
        frozen, excel = tmp_path / (name + ".json"), tmp_path / (name + ".xlsx")
        export_pair(report, frozen, excel)
        assert read_json(frozen) == read_excel(excel) == report
        assert report.basis == b
    empty = variants["empty"]
    assert not empty.view.observation_ids and not empty.view.adopted_event_ids
    assert empty.basis.issues == b.issues
    raw = report_to_json(start)
    data = json.loads(raw)
    assert (
        len(data["report"]["basis"]["observations"]) == 9
    )  # gamle objekter ligger bare i kapselen
    assert len(data["report"]["basis"]["legacy_json"]) == len(b.legacy_json)
    assert all(
        "observations" not in x for x in data["report"]["view"].values() if isinstance(x, dict)
    )
    originals = tmp_path / "originals"
    export_originals(read_excel(tmp_path / "union.xlsx"), originals)
    assert (originals / (shared.sha256 + ".bin")).read_bytes() == Path(PLENUM).read_bytes()


@pytest.mark.parametrize(
    "field,value",
    [
        ("budget_year", 2025),
        ("chapter", "1429"),
        ("post", "1"),
        ("currency", "USD"),
        ("multiplier", "1000"),
        ("meaning", "niva"),
        ("detail_level", "underpost"),
        ("scope", "aggregat"),
        ("round", "rnb"),
        ("figure_type", "regnskap"),
        ("proposal_role", "alternativ"),
        ("proposal_author", "samlet_komite"),
        ("information_type", "regnskap"),
    ],
)
def test_synthetic_incompatible_comparison_has_named_block(authentic_basis, field, value):
    left = authentic_basis.observations[1]
    right = authentic_basis.observations[4]
    context = {c.observation_id: c for c in authentic_basis.contexts}
    changed = replace(right, **{field: value})
    result = compare_observations(
        left,
        changed,
        context[left.id],
        replace(context[right.id], observation_id=changed.id),
        account="315/78",
        meaning="endring",
        relation_bound=True,
    )
    assert result.status == "sperret" and result.difference_kr is None and result.null_reason


def _rehash(data):
    b = data["report"]["basis"]
    b["id"] = ""
    b["id"] = identity("komitegrunnlag-1", b)
    data["report"]["id"] = ""
    data["report"]["id"] = identity("visningspilot-2", data["report"])
    return encode(data)


@pytest.mark.parametrize(
    "mutation",
    [
        "amount",
        "actor",
        "vote",
        "case",
        "relation",
        "dangling",
        "locator",
        "year",
        "sourcehash",
        "payload",
        "parallel",
        "status",
        "legacy",
    ],
)
def test_synthetic_rehashed_actual_files_cannot_rebind_source_proofs(
    tmp_path, saved_pair, mutation
):
    directory, _ = saved_pair
    data = json.loads((directory / "post.json").read_text("utf8"))
    b = data["report"]["basis"]
    if mutation == "amount":
        b["observations"][4]["amount_kr"] = 119_000_000
    elif mutation == "actor":
        b["contexts"][4]["actor"] = "samlet_komite"
    elif mutation == "vote":
        next(s for s in b["statements"] if s["role"] == "postvis_votering_63_35_194948")["text"] = (
            "synthetic unanimous 315/78"
        )
    elif mutation == "case":
        next(d for d in b["documents"] if d["key"] == "sak100887")["title"] = "Sak 100927"
    elif mutation == "relation":
        b["relations"][1]["accounts"] = ["1429/1"]
    elif mutation == "dangling":
        b["relations"][1]["evidence_ids"] = ["synthetic-unknown-proof"]
    elif mutation == "locator":
        b["observations"][4]["pointer"]["locator"] = "//table/tbody/tr[1]/td[4]"
    elif mutation == "year":
        b["observations"][4]["budget_year"] = 2025
    elif mutation == "sourcehash":
        b["sources"][3]["expected_sha256"] = "0" * 64
    elif mutation == "payload":
        b["payloads"][0]["sha256"] = "0" * 64
    elif mutation == "parallel":
        parallel = json.loads(json.dumps(b["observations"][4]))
        parallel["proposal_role"] = "alternativ"
        parallel["proposal_author"] = "uavklart"
        b["observations"].append(parallel)
    elif mutation == "status":
        b["issues"] = []
    elif mutation == "legacy":
        b["legacy_json"] = '{"format":"synthetic"}'
    raw = _rehash(data)
    path = tmp_path / "synthetic-rehashed.json"
    path.write_text(raw, "utf8")
    with pytest.raises(ValueError):
        read_json(path)
    excel = tmp_path / "synthetic-rehashed.xlsx"
    shutil.copyfile(directory / "post.xlsx", excel)
    book = load_workbook(excel)
    sheet = book["Offlinegrunnlag"]
    sheet.delete_rows(2, sheet.max_row)
    for n, start in enumerate(range(0, len(raw), 30000), 2):
        sheet.cell(n, 1, n - 1)
        sheet.cell(n, 2, raw[start : start + 30000])
    book.save(excel)
    book.close()
    with pytest.raises(ValueError):
        read_excel(excel)


@pytest.mark.parametrize(
    "mutation",
    [
        "numberformat",
        "font",
        "fill",
        "row",
        "column",
        "filter",
        "hidden",
        "conditional",
        "hyperlink",
        "parentref",
        "drawing",
        "ziplimit",
        "doctype",
    ],
)
def test_synthetic_excel_presentation_and_archive_manipulations_are_rejected(
    tmp_path, saved_pair, mutation
):
    directory, _ = saved_pair
    output = tmp_path / "synthetic-view.xlsx"
    shutil.copyfile(directory / "post.xlsx", output)
    if mutation in {"drawing", "ziplimit", "doctype"}:
        with ZipFile(output) as archive:
            members = {n: archive.read(n) for n in archive.namelist()}
        if mutation == "drawing":
            members["xl/drawings/drawing1.xml"] = b"<synthetic/>"
        elif mutation == "ziplimit":
            members.update({f"synthetic-{i}": b"x" for i in range(260)})
        else:
            members["xl/worksheets/sheet1.xml"] = (
                b'<!DOCTYPE worksheet [<!ENTITY e "hidden">]>' + members["xl/worksheets/sheet1.xml"]
            )
        with ZipFile(output, "w") as archive:
            for name, raw in members.items():
                archive.writestr(name, raw)
    else:
        book = load_workbook(output)
        sheet = book["Forløp"]
        if mutation == "numberformat":
            sheet["H2"].number_format = '"0"'
        elif mutation == "font":
            sheet["H2"].font = Font(name="Arial", color="FFFFFF", size=1)
        elif mutation == "fill":
            sheet["H2"].fill = PatternFill("solid", fgColor="000000")
        elif mutation == "row":
            book["Dekning og avvik"].row_dimensions[12].hidden = True
        elif mutation == "column":
            sheet.column_dimensions["H"].width = 0.2
        elif mutation == "filter":
            sheet.auto_filter.ref = "A1:A1"
        elif mutation == "hidden":
            book["Dekning og avvik"].sheet_state = "hidden"
        elif mutation == "conditional":
            sheet.conditional_formatting.add(
                "H2", CellIsRule(operator="equal", formula=["119000000"], font=Font(color="FFFFFF"))
            )
        elif mutation == "hyperlink":
            sheet["C2"].hyperlink = "https://example.invalid/synthetic"
        elif mutation == "parentref":
            sheet["C2"].hyperlink = copy(sheet["C2"].hyperlink)
            sheet["C2"].hyperlink.location = "'Dokumenter'!C999"
        book.save(output)
        book.close()
    with pytest.raises(ValueError):
        read_excel(output)


@pytest.mark.parametrize("condition", ["missing", "failed", "corrupt", "contradictory"])
def test_new_authentic_368_is_independent_of_bad_t2_enrichment(tmp_path, condition):
    from innlesning import read_source
    from innlesning.observations import read_extraction

    root = tmp_path / "sources"
    for name in CATALOG:
        destination = root / name
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(name, destination)
    changed = root / (K + "inns85s.html")
    if condition == "missing":
        changed.unlink()
    elif condition == "failed":
        changed.unlink()
        changed.mkdir()
    elif condition == "corrupt":
        changed.write_bytes(b"synthetic corrupt HTML")
    else:
        changed.write_text(changed.read_text("utf8").replace("129 000 000", "119 000 000"), "utf8")
    basis = import_basis(root)
    assert not any(o.source_unit_id == "inns85" for o in basis.observations)
    report = build_report(basis, PilotSelection(years=(2024,), accounts=(("315", "78"),)))
    assert len(report.view.observation_ids) == 6
    assert all(c.status == "hull" and c.difference_kr is None for c in report.view.comparisons)
    export_pair(report, tmp_path / "gap.json", tmp_path / "gap.xlsx")
    assert read_excel(tmp_path / "gap.xlsx") == report
    empty = build_report(basis, PilotSelection())
    export_pair(empty, tmp_path / "empty.json", tmp_path / "empty.xlsx")
    assert not read_excel(tmp_path / "empty.xlsx").view.observation_ids
    assert any("inns85s.html" in issue for issue in basis.issues)
    database = tmp_path / "fresh.sqlite"

    def cli(*args):
        result = subprocess.run(
            [sys.executable, "-m", *args], capture_output=True, text=True, encoding="utf8"
        )
        assert result.returncode in (0, 2), result.stdout + result.stderr
        return json.loads(result.stdout)

    sync = cli(
        "innlesning",
        "synkroniser",
        "--database",
        str(database),
        "--sesjon",
        "2024-2025",
        "--historisk-proeve",
        "--vedtak-id",
        "40034229",
        "--response",
        "tests/fixtures/b6/40034229.xml",
        "--metadata",
        "tests/fixtures/b6/40034229.json",
    )
    assert sync["antall"]["ny"] == 1 and not sync["feil"]
    revision = sync["vedtak"][0]["revision_id"]
    parsed = cli(
        "innlesning",
        "tolk",
        "--database",
        str(database),
        "--revisjon",
        revision,
        "--tolkning",
        "storting-belop-v2",
    )
    extraction = read_extraction(database, parsed["belopsuttrekk"])
    assert (
        len(extraction.observations) == 222
        and len({o.step_id for o in extraction.observations}) == 74
    )
    assert (
        read_source(database, extraction.artifact.sha256)
        == Path("tests/fixtures/b6/40034229.xml").read_bytes()
    )
    output = tmp_path / "ordinary.xlsx"
    cli(
        "analyse",
        "--database",
        str(database),
        "--belopsuttrekk",
        extraction.id,
        "--excel",
        str(output),
    )
    book = load_workbook(output)
    try:
        rows = list(book["Beløpsobservasjoner"].iter_rows(min_row=2, values_only=True))
        assert len(rows) == 222 and {r[0] for r in rows} == {o.id for o in extraction.observations}
    finally:
        book.close()


def test_shared_adapter_runs_with_analysis_unavailable(tmp_path):
    script = tmp_path / "no_analysis.py"
    script.write_text(
        """import importlib.abc, sys
from pathlib import Path
class NoAnalysis(importlib.abc.MetaPathFinder):
 def find_spec(self, fullname, path, target=None):
  if fullname == "analyse" or fullname.startswith("analyse."):
   raise RuntimeError("analysis unavailable")
sys.meta_path.insert(0, NoAnalysis())
from innlesning.komitepilot import import_basis, validate_basis
basis=import_basis(Path("."));validate_basis(basis)
assert len(basis.observations)==9
""",
        "utf8",
    )
    result = subprocess.run(
        [sys.executable, "-c", script.read_text("utf8")],
        capture_output=True,
        text=True,
        encoding="utf8",
    )
    assert result.returncode == 0, result.stdout + result.stderr


def test_strict_json_and_derivation_parent_guard(saved_pair):
    directory, report = saved_pair
    raw = (directory / "post.json").read_text("utf8")
    for bad in (
        raw.replace('"format":', '"format":"visningspilot-2","format":', 1),
        raw.replace("119000000", "NaN", 1),
        raw.replace('"version":"visningspilot-2"', '"version":true', 1),
    ):
        with pytest.raises(ValueError):
            report_from_json(bad)
    with pytest.raises(ValueError):
        derived_xml(b"synthetic parent", "40034045")
    assert all(payload_bytes(p) for p in report.basis.payloads)
