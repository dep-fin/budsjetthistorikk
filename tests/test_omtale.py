"""O2a ved den avtalte original → faktisk JSON/Excel-grensen. Ingen live nett."""

import hashlib
import json
import re
import shutil
import subprocess
import sys
from dataclasses import asdict, replace
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from lxml import etree, html
from openpyxl import load_workbook
from test_b9 import source_report

from analyse.omtale import build_report, export_pair, read_excel, read_json
from innlesning.omtale import import_omtale, source_bytes

SOURCES = Path("docs/research/omtale-o1")
SUPPORT = Path("tests/fixtures")


@pytest.mark.parametrize("over_limit", [False, True])
def test_supporting_manifest_is_read_once_with_a_byte_limit(tmp_path, monkeypatch, over_limit):
    """Små byte og instrumentering; en uavgrenset read stanser før faktisk lesing."""
    import innlesning.omtale as omtale_module

    root = tmp_path / "synthetic-o1"
    shutil.copytree(SOURCES, root)
    manifest = root / "supporting-manifest.json"
    limit = len(manifest.read_bytes())
    if over_limit:
        manifest.write_bytes(manifest.read_bytes() + b" ")
    monkeypatch.setattr(omtale_module, "MAX_METADATA_BYTES", limit)
    original_open = Path.open
    reads = []
    opens = []

    class BoundedInput:
        def __init__(self, handle):
            self.handle = handle

        def __enter__(self):
            return self

        def __exit__(self, *args):
            self.handle.close()

        def read(self, size=-1):
            reads.append(size)
            assert 0 < size <= limit + 1, "Manifest må aldri leses uten bytegrense"
            return self.handle.read(size)

    def guarded_open(path, *args, **kwargs):
        handle = original_open(path, *args, **kwargs)
        if path == manifest:
            opens.append(path)
            return BoundedInput(handle)
        return handle

    monkeypatch.setattr(Path, "open", guarded_open)
    basis = import_omtale(root, SUPPORT)
    sidecars = [s for s in basis.sources if s.name.endswith(".json")]
    expected = "motstrid" if over_limit else "kildeverifisert"
    assert [s.status for s in sidecars] == [expected, expected]
    assert all(source_bytes(s) == (SUPPORT / s.name).read_bytes() for s in sidecars)
    assert opens == [manifest]
    assert reads == [limit + 1]


@pytest.mark.parametrize("shape", ["object", "scalar", "mixed", "missing_path", "duplicate"])
def test_synthetic_supporting_manifest_shape_stays_visible_in_saved_files(tmp_path, shape):
    root = tmp_path / "synthetic-o1"
    shutil.copytree(SOURCES, root)
    path = root / "supporting-manifest.json"
    valid = json.loads(path.read_text("utf-8"))
    variants = {
        "object": {"repository_path": "tests/fixtures/b6/40034048.json"},
        "scalar": 1,
        "mixed": [*valid, None],
        "missing_path": [{"bytes": 786}],
        "duplicate": [*valid, *valid],
    }
    path.write_text(json.dumps(variants[shape]), "utf-8")
    _, base = source_report(tmp_path, posts=("1",))
    frozen, excel = tmp_path / "synthetic-manifest.json", tmp_path / "synthetic-manifest.xlsx"
    export_pair(build_report(base, import_omtale(root, SUPPORT)), frozen, excel)
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    assert len(saved.basis.sources) == 10
    for name in ("b6/40034048.json", "b8/sak-100927.json"):
        sidecar = next(s for s in saved.basis.sources if s.name == name)
        assert sidecar.status == "motstrid"
        assert source_bytes(sidecar) == (SUPPORT / name).read_bytes()
        assert "motstrid:" + name in saved.issues
    assert not any(r.status == "kildeverifisert" for r in saved.relations)


@pytest.fixture(scope="module")
def range_test_excel(tmp_path_factory):
    directory = tmp_path_factory.mktemp("range-source")
    _, base = source_report(directory)
    frozen, excel = directory / "valid.json", directory / "valid.xlsx"
    export_pair(build_report(base, import_omtale(SOURCES, SUPPORT)), frozen, excel)
    return excel


@pytest.mark.parametrize(
    "part,element,attribute,value",
    [
        ("sheet1", "hyperlink", "ref", "A1:XFD1048576"),
        ("sheet1", "hyperlink", "ref", "A1:A1"),
        ("sheet1", "hyperlink", "ref", "IV20000"),
        ("sheet1", "dimension", "ref", "A1:IV20001"),
        ("sheet1", "dimension", "ref", "B2:A1"),
        ("sheet1", "autoFilter", "ref", "A1:XFD1048576"),
        ("sheet1", "autoFilter", "ref", "A1:A1 B2:B2"),
        ("sheet1", "selection", "sqref", "A1:XFD1048576"),
        ("sheet1", "selection", "sqref", "A1 B2"),
        ("sheet1", "selection", "activeCell", "A20001"),
        ("sheet1", "pane", "topLeftCell", "IW1"),
        ("sheet1", "pane", "xSplit", "16384"),
        ("sheet1", "pane", "ySplit", "1048576"),
        ("sheet1", "col", "min", "0"),
        ("sheet1", "col", "max", "16384"),
        ("sheet1", "row", "r", "1048576"),
        ("sheet1", "row", "spans", "1:16384"),
        ("sheet1", "c", "s", "99999999"),
        ("sheet1", "c", "t", "s"),
        ("sheet1", "sheetView", "workbookViewId", "64"),
        ("workbook", "definedName", "localSheetId", "64"),
        ("workbook", "definedName", "text", "'Historikk'!$A$1:$XFD$1048576"),
        ("workbook", "definedName", "name", "_xlnm.Print_Area"),
        ("workbook", "workbookView", "activeTab", "64"),
        ("workbook", "sheet", "sheetId", "99999999"),
        ("styles", "xf", "fontId", "99999999"),
        ("styles", "xf", "fillId", "99999999"),
        ("styles", "xf", "borderId", "99999999"),
        ("styles", "xf", "xfId", "99999999"),
        ("styles", "xf", "numFmtId", "99999999"),
        ("styles", "numFmts", "count", "99999999"),
        ("styles", "numFmt", "numFmtId", "99999999"),
    ],
)
def test_synthetic_range_and_index_is_rejected_before_input_load(
    tmp_path, range_test_excel, monkeypatch, part, element, attribute, value
):
    """Liten XML-fixtur; sentinel stanser alltid før eventuell allokering."""
    import analyse.omtale as omtale_module
    import analyse.saksdokumenter as saks_module

    damaged = tmp_path / "synthetic-range.xlsx"
    member_path = f"xl/worksheets/{part}.xml" if part == "sheet1" else f"xl/{part}.xml"
    with ZipFile(range_test_excel) as source, ZipFile(damaged, "w", ZIP_DEFLATED) as target:
        for member in source.infolist():
            raw = source.read(member)
            if member.filename == member_path:
                root = etree.fromstring(raw)
                node = next(
                    root.iter(
                        "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}" + element
                    )
                )
                if attribute == "text":
                    node.text = value
                else:
                    node.set(attribute, value)
                raw = etree.tostring(root)
            target.writestr(member, raw)
    loads = []
    original_load = load_workbook

    def guarded_load(filename, *args, **kwargs):
        if Path(filename).resolve() == damaged.resolve():
            loads.append(str(filename))
            raise AssertionError("Sentinel: input må ikke lastes eller materialiseres")
        return original_load(filename, *args, **kwargs)

    monkeypatch.setattr(omtale_module, "load_workbook", guarded_load)
    monkeypatch.setattr(saks_module, "load_workbook", guarded_load)
    with pytest.raises(ValueError, match="Excel"):
        read_excel(damaged)
    assert not loads


@pytest.mark.parametrize(
    "field,value",
    [
        ("chapter_raw", "1428"),
        ("post_raw", "01"),
        ("subpost_raw", "1"),
        ("chapter_name", "Syntetisk kapittelnavn"),
        ("purpose", "Driftsutgifter, reduseres med"),
        ("decision_date", "2024-12-15"),
        ("decision_session", "2023-2024"),
        ("stage_raw", "Syntetisk vedtakstittel"),
        ("year_evidence", "Syntetisk årsbelegg"),
        ("unit_raw", "Tusen kroner"),
        ("currency", "USD"),
        ("multiplier", "1000"),
        ("event_id", "synthetic-other-event"),
        ("documents", ()),
    ],
)
def test_synthetic_consistent_raw_context_stays_unresolved_in_both_saved_files(
    tmp_path, field, value
):
    """Både utvalgsrader og uttrekk lyver likt; XML-konteksten er autoritativ."""
    from innlesning.observations import identity

    _, base = source_report(tmp_path, posts=("1",))

    def forged(observation):
        return replace(observation, **{field: value}) if observation.post == "1" else observation

    base = replace(
        base,
        source_text=replace(
            base.source_text,
            extractions=tuple(
                replace(e, observations=tuple(forged(o) for o in e.observations))
                for e in base.source_text.extractions
            ),
        ),
        base=replace(
            base.base,
            rows=tuple(replace(r, observation=forged(r.observation)) for r in base.base.rows),
            unfiltered_rows=tuple(
                replace(r, observation=forged(r.observation)) for r in base.base.unfiltered_rows
            ),
        ),
        id="",
    )
    base = replace(base, id=identity("b6-frozen-2", asdict(base)))
    report = build_report(base, import_omtale(SOURCES, SUPPORT))
    frozen, excel = tmp_path / "synthetic-raw.json", tmp_path / "synthetic-raw.xlsx"
    export_pair(report, frozen, excel)
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    assert "210_post1_grunnlag_mangler_eller_motstrid" in saved.issues
    assert not saved.relations


@pytest.mark.parametrize(
    "field,value",
    [
        ("id", "40034049"),
        ("number", "211"),
        ("case_id", "100889"),
        ("date", "2024-12-15"),
        ("date_raw", "2024-12-15T00:00:00"),
        ("api_type", "BUDSJETT"),
        ("title", "Syntetisk metadataendring"),
        ("url", "https://www.stortinget.no/synthetic"),
        ("url_raw", "//www.stortinget.no/synthetic"),
        ("budget_year", 2023),
    ],
)
def test_synthetic_decision_metadata_is_bound_to_original_xml_in_saved_files(
    tmp_path, field, value
):
    from innlesning.observations import identity

    _, base = source_report(tmp_path, posts=("1",))
    base = replace(
        base,
        source_text=replace(
            base.source_text,
            decisions=tuple(
                (r, replace(d, **{field: value})) for r, d in base.source_text.decisions
            ),
        ),
        id="",
    )
    base = replace(base, id=identity("b6-frozen-2", asdict(base)))
    frozen, excel = tmp_path / "synthetic-metadata.json", tmp_path / "synthetic-metadata.xlsx"
    export_pair(build_report(base, import_omtale(SOURCES, SUPPORT)), frozen, excel)
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    assert any("vedtaksbinding_motstrid" in issue for issue in saved.issues)
    assert not any(r.status == "kildeverifisert" for r in saved.relations)


@pytest.mark.parametrize(
    "field,value",
    [
        ("url", "https://data.stortinget.no/eksport/stortingsvedtak?sesjonid=2023-2024"),
        ("retrieved_utc", "2026-10-20T11:37:10+00:00"),
        ("headers", {"Content-Type": "application/json"}),
        ("decision_id", "40034049"),
        ("status", 200.0),
        ("source_response_bytes", 7479944),
        ("synthetic_note", "Uttrykkelig syntetisk metadataavvik"),
        ("provenance", None),
    ],
)
def test_synthetic_xml_metadata_disagreement_stays_visible_in_saved_files(tmp_path, field, value):
    """Bare XML-kildens metadata endres; bevart sidevogn og original er uendret."""
    from innlesning.observations import encode
    from innlesning.omtale import identified, structure

    _, base = source_report(tmp_path, posts=("1",))
    basis = import_omtale(SOURCES, SUPPORT)
    original = next(s for s in basis.sources if s.name == "b6/40034048.xml")
    metadata = json.loads(original.metadata_json)
    if value is None:
        del metadata[field]
    else:
        metadata[field] = value
    changed = identified("omtale-kilde-1", replace(original, metadata_json=encode(metadata)))
    basis = structure(
        tuple(changed if s == original else s for s in basis.sources), basis.extractor
    )
    frozen, excel = (
        tmp_path / "synthetic-xml-metadata.json",
        tmp_path / "synthetic-xml-metadata.xlsx",
    )
    export_pair(build_report(base, basis), frozen, excel)
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    assert any("vedtaksbinding_motstrid" in issue for issue in saved.issues)
    assert not any(r.status == "kildeverifisert" for r in saved.relations)
    preserved = next(s for s in saved.basis.sources if s.name == "b6/40034048.xml")
    assert preserved == changed
    assert source_bytes(preserved) == (SUPPORT / original.name).read_bytes()
    sidecar = next(s for s in saved.basis.sources if s.name == "b6/40034048.json")
    assert source_bytes(sidecar) == (SUPPORT / sidecar.name).read_bytes()


@pytest.mark.parametrize("duplicate", [False, True])
def test_xml_metadata_json_representation_is_checked_without_rewriting_raw_data(
    tmp_path, duplicate
):
    from innlesning.omtale import identified, structure

    _, base = source_report(tmp_path, posts=("1",))
    basis = import_omtale(SOURCES, SUPPORT)
    original = next(s for s in basis.sources if s.name == "b6/40034048.xml")
    metadata = json.loads(original.metadata_json)
    raw = json.dumps(dict(reversed(tuple(metadata.items()))), ensure_ascii=False, indent=2)
    if duplicate:
        raw = raw[:-1] + ', "status": 200}'
    changed = identified("omtale-kilde-1", replace(original, metadata_json=raw))
    basis = structure(
        tuple(changed if s == original else s for s in basis.sources), basis.extractor
    )
    frozen, excel = (
        tmp_path / "metadata-representation.json",
        tmp_path / "metadata-representation.xlsx",
    )
    export_pair(build_report(base, basis), frozen, excel)
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    preserved = next(s for s in saved.basis.sources if s.name == "b6/40034048.xml")
    assert preserved.metadata_json == raw
    if duplicate:
        assert any("vedtaksbinding_motstrid" in issue for issue in saved.issues)
        assert not any(r.status == "kildeverifisert" for r in saved.relations)
    else:
        assert len(saved.relations) == 6
        assert all(r.status == "kildeverifisert" for r in saved.relations)


def test_synthetic_session_relabel_cannot_verify_another_event_in_saved_files(tmp_path):
    """Sesjon endres konsistent i snapshot/uttrekk/rader, men XML er uendret."""
    from innlesning.observations import identity

    _, base = source_report(tmp_path, posts=("1",))
    unit = "storting:2023-2024:40034048"

    def forged(observation):
        return replace(
            observation,
            decision_session="2023-2024",
            source_unit_id=unit,
            step_id=identity(unit, observation.source_revision_id, "table[1]/tr[3]"),
        )

    base = replace(
        base,
        source_text=replace(
            base.source_text,
            decisions=tuple(
                (r, replace(d, session="2023-2024")) for r, d in base.source_text.decisions
            ),
            extractions=tuple(
                replace(
                    e, source_unit_id=unit, observations=tuple(forged(o) for o in e.observations)
                )
                for e in base.source_text.extractions
            ),
        ),
        base=replace(
            base.base,
            rows=tuple(replace(r, observation=forged(r.observation)) for r in base.base.rows),
            unfiltered_rows=tuple(
                replace(r, observation=forged(r.observation)) for r in base.base.unfiltered_rows
            ),
        ),
        id="",
    )
    base = replace(base, id=identity("b6-frozen-2", asdict(base)))
    report = build_report(base, import_omtale(SOURCES, SUPPORT))
    frozen, excel = tmp_path / "synthetic-session.json", tmp_path / "synthetic-session.xlsx"
    export_pair(report, frozen, excel)
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    assert any("vedtaksbinding_motstrid" in issue for issue in saved.issues)
    assert not any(r.status == "kildeverifisert" for r in saved.relations)


@pytest.mark.parametrize("scope", ["rows", "rows_and_extraction", "selected_only"])
def test_synthetic_rehashed_event_and_locator_cannot_create_verified_relations(tmp_path, scope):
    """Merket syntetisk B8; original og kildebundet uttrekk er uendret."""
    from innlesning.observations import identity

    _, base = source_report(tmp_path, posts=("1",))

    def forged(row):
        if row.observation.post != "1":
            return row
        observation = replace(
            row.observation,
            step_id="synthetic-forged-event-999",
            pointer=replace(row.observation.pointer, locator="table[999]/tr[999]"),
        )
        return replace(row, observation=observation)

    session = replace(
        base.base,
        rows=tuple(forged(r) for r in base.base.rows),
        unfiltered_rows=(
            base.base.unfiltered_rows
            if scope == "selected_only"
            else tuple(forged(r) for r in base.base.unfiltered_rows)
        ),
        id="",
    )
    session = replace(
        session,
        id=identity(
            "sesjonsrapport-v1",
            asdict(session.selection),
            asdict(session.snapshot),
            [(r.observation.id, r.approved_amount_kr, r.freshness) for r in session.rows],
        ),
    )
    base = replace(base, base=session, id="")
    if scope == "rows_and_extraction":
        base = replace(
            base,
            source_text=replace(
                base.source_text,
                extractions=tuple(
                    replace(
                        e,
                        observations=tuple(
                            forged(replace(session.unfiltered_rows[0], observation=o)).observation
                            for o in e.observations
                        ),
                    )
                    for e in base.source_text.extractions
                ),
            ),
        )
    base = replace(base, id=identity("b6-frozen-2", asdict(base)))
    report = build_report(base, import_omtale(SOURCES, SUPPORT))
    frozen, excel = tmp_path / "synthetic-forged.json", tmp_path / "synthetic-forged.xlsx"
    export_pair(report, frozen, excel)
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    assert "210_post1_grunnlag_mangler_eller_motstrid" in saved.issues
    assert not saved.relations


def test_authentic_post1_from_reopened_originals_to_self_contained_exports(tmp_path):
    db, base = source_report(tmp_path, posts=("1",))
    basis = import_omtale(SOURCES, SUPPORT)
    report = build_report(base, basis)
    output = tmp_path / "omtale.json"
    excel = tmp_path / "omtale.xlsx"
    export_pair(report, output, excel)
    db.write_bytes(b"unavailable")
    saved = read_json(output)
    assert saved == report == read_excel(excel)
    assert json.loads(output.read_text("utf-8"))["format"] == "omtale-pilot-1"
    assert len(base.base.unfiltered_rows) == 6
    assert len({r.observation.step_id for r in base.base.unfiltered_rows}) == 2
    sections = {s.key: s for s in saved.basis.sections}
    p1 = sections["P1"]
    assert p1.accounts == ("1429/1", "1429/22")
    assert "2,8\xa0mill." in p1.text and "3\xa0mill." in p1.text
    assert p1.hierarchy[1].text == "Kap. 1429 Riksantikvaren"
    assert sections["P1_PDF"].locator == "pdf:page=3"
    assert {r.role for r in saved.relations} >= {
        "regjeringsforklaring",
        "gjengivelse_av_regjeringen",
        "samlet_stotte",
        "tilrading",
        "vedtakelse",
    }
    assert all(r.professional_status == "ikke_faglig_godkjent" for r in saved.relations)
    assert all(r.status == "kildeverifisert" for r in saved.relations)
    assert all(r.event_id == base.base.rows[0].observation.step_id for r in saved.relations)
    for source in saved.basis.sources:
        path = (SOURCES if source.name.startswith("originals/") else SUPPORT) / source.name
        assert source_bytes(source) == path.read_bytes()
    book = load_workbook(excel)
    try:
        assert book["Historikk"].max_column == 13
        assert book["Historikk"].max_row == 2
        assert tuple(book["Historikk"].cell(2, c).value for c in (10, 11, 12)) == (
            160_243_000,
            5_796_000,
            166_039_000,
        )
        assert {"Omtale", "Relasjoner", "Omtalekilder", "Omtaleavvik"} <= set(book.sheetnames)
    finally:
        book.close()


def test_later_synthetic_envelope_keeps_original_revision_artifact_and_verified_relations(tmp_path):
    """B2 gjenbruker objektet; ny responstid er ingen ny innholdsrevisjon."""
    from analyse.saksdokumenter import build_saks_report
    from innlesning import read_sync, sync_saved

    database, original = source_report(tmp_path, posts=("1",))
    raw = (
        (SUPPORT / "b6/40034048.xml")
        .read_bytes()
        .replace(b"2026-10-20T13:36:10.2847284+02:00", b"2026-10-20T13:37:10.2847284+02:00")
    )
    metadata = json.loads((SUPPORT / "b6/40034048.json").read_text("utf-8"))
    metadata.update(
        sha256=hashlib.sha256(raw).hexdigest(),
        bytes=len(raw),
        retrieved_utc="2026-10-20T11:37:10+00:00",
        fixture_kind="synthetic_test",
        synthetic_note="Kun senere responskonvolutt; originalt vedtaksobjekt er byteidentisk",
    )
    job = sync_saved(database, raw, metadata, historical_sample=True, decision_id="40034048")
    assert read_sync(database, job).entries[0].revision_id == original.source_text.decisions[0][0]
    base = build_saks_report(
        database,
        replace(
            original.selection,
            base=replace(original.selection.base, inventory_job_ids=(job,)),
        ),
    )
    assert base.source_text == original.source_text
    frozen, excel = tmp_path / "later-envelope.json", tmp_path / "later-envelope.xlsx"
    export_pair(build_report(base, import_omtale(SOURCES, SUPPORT)), frozen, excel)
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    assert len(saved.relations) == 6
    assert all(r.status == "kildeverifisert" for r in saved.relations)


def synthetic_sources(tmp_path, filename, old, new):
    """Merket syntetisk kildeversjon, aldri endring av autentiske originaler."""
    root = tmp_path / "synthetic-o1"
    shutil.copytree(SOURCES, root)
    path = root / "originals" / filename
    raw = path.read_bytes()
    assert old.encode() in raw
    raw = raw.replace(old.encode(), new.encode())
    path.write_bytes(raw)
    manifest = json.loads((root / "source-manifest.json").read_text("utf-8"))
    for m in manifest:
        if m.get("file") == "originals/" + filename and m.get("status") == 200:
            m.update(
                sha256=hashlib.sha256(raw).hexdigest(),
                bytes=len(raw),
                fixture_kind="synthetic_mutation",
            )
    (root / "source-manifest.json").write_text(json.dumps(manifest), "utf-8")
    return root


@pytest.mark.parametrize(
    "filename,selector,key",
    [
        ("inns81s.html", "//*[@id='m3']//table", "KT"),
        ("votering-161224.html", "//*[@id='2024-12-16-25']//table", "V"),
    ],
)
def test_synthetic_wrong_local_table_chapter_is_unresolved_in_both_saved_files(
    tmp_path, filename, selector, key
):
    """1429 finnes fortsatt i innledning/post 60, men ikke i beløpstabellen."""
    original = (SOURCES / "originals" / filename).read_bytes().decode("utf-8")
    doc = html.fromstring(original)
    chapter = doc.xpath(selector + "//tr[td[normalize-space()='1429']]/td[1]")
    assert len(chapter) == 1
    chapter[0].xpath(".//p")[0].text = "9999"
    root = synthetic_sources(tmp_path, filename, original, html.tostring(doc, encoding="unicode"))
    _, base = source_report(tmp_path)
    report = build_report(base, import_omtale(root, SUPPORT))
    frozen, excel = tmp_path / "synthetic-table.json", tmp_path / "synthetic-table.xlsx"
    export_pair(report, frozen, excel)
    saved = read_excel(excel)
    assert saved == read_json(frozen)
    assert "konto_eller_forslagsbelop_motstrid:" + key in saved.issues
    assert all(r.status == "uavklart" for r in saved.relations)


@pytest.mark.parametrize(
    "filename,selector,key",
    [
        ("prop16s-ch4.html", "//tr[td[normalize-space()='1429']]/following-sibling::tr[1]", "PT"),
        ("inns81s.html", "//*[@id='m3']//tr[td[2][normalize-space()='1']]", "KT"),
        ("votering-161224.html", "//*[@id='2024-12-16-25']//tr[td[2][normalize-space()='1']]", "V"),
    ],
)
def test_synthetic_table_evidence_cannot_borrow_account_or_levels_from_other_rows(
    tmp_path, filename, selector, key
):
    original = (SOURCES / "originals" / filename).read_bytes().decode("utf-8")
    doc = html.fromstring(original)
    delta = doc.xpath(selector)[0]
    if key == "PT":
        delta.xpath("./td[1]/p")[0].text = "9999"
    else:
        level = delta.xpath("following-sibling::tr[1]/td[3]/p")[0]
        level.text = "fra kr 0 til kr 0"
        donor = html.Element("p")
        donor.text = "Syntetisk nabokontekst: fra kr 160 243 000 til kr 166 039 000"
        delta.xpath("ancestor::table")[0].addnext(donor)
    root = synthetic_sources(tmp_path, filename, original, html.tostring(doc, encoding="unicode"))
    _, base = source_report(tmp_path)
    report = build_report(base, import_omtale(root, SUPPORT))
    frozen, excel = tmp_path / "synthetic-local.json", tmp_path / "synthetic-local.xlsx"
    export_pair(report, frozen, excel)
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    assert "konto_eller_forslagsbelop_motstrid:" + key in saved.issues
    assert all(r.status == "uavklart" for r in saved.relations)


def test_synthetic_named_representative_does_not_inherit_collective_committee_role(tmp_path):
    original = (SOURCES / "originals/inns81s.html").read_bytes().decode("utf-8")
    doc = html.fromstring(original)
    actor = doc.xpath("//*[@id='m2']//p/span")[0]
    actor.text = "Representanten Grunde Almeland"
    root = synthetic_sources(
        tmp_path, "inns81s.html", original, html.tostring(doc, encoding="unicode")
    )
    _, base = source_report(tmp_path)
    report = build_report(base, import_omtale(root, SUPPORT))
    frozen, excel = tmp_path / "synthetic-actor.json", tmp_path / "synthetic-actor.xlsx"
    export_pair(report, frozen, excel)
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    ks = next(s for s in saved.basis.sections if s.key == "KS")
    assert (ks.actor, ks.role, ks.status) == ("ukjent", "ukjent_rolle", "ukjent_rolle")
    assert "ukjent_rolle:KS" in saved.issues
    assert all(r.status == "uavklart" for r in saved.relations)


@pytest.mark.parametrize("padding", ["extra_sheet", "unbound_cell", "unbound_value", "xml_content"])
def test_synthetic_workbook_padding_is_rejected_at_the_saved_file_boundary(tmp_path, padding):
    _, base = source_report(tmp_path)
    report = build_report(base, import_omtale(SOURCES, SUPPORT))
    frozen, excel = tmp_path / "valid.json", tmp_path / "valid.xlsx"
    export_pair(report, frozen, excel)
    book = load_workbook(excel)
    if padding == "extra_sheet":
        sheet = book.create_sheet("synthetic-unbound")
        for n in range(1, 200_001):
            sheet.cell(n, 1, n)
    elif padding == "unbound_cell":
        book["Saksgrunnlag"].cell(999, 1, "synthetic unbound content")
    elif padding == "unbound_value":
        book["Sesjonsgrunnlag"].cell(1, 1, "synthetic unbound value")
    padded = tmp_path / "synthetic-padded.xlsx"
    book.save(padded)
    book.close()
    if padding == "xml_content":
        rewritten = tmp_path / "synthetic-xml.xlsx"
        with ZipFile(padded) as source, ZipFile(rewritten, "w", ZIP_DEFLATED) as target:
            for member in source.infolist():
                raw = source.read(member)
                if member.filename == "xl/worksheets/sheet1.xml":
                    raw = raw.replace(
                        b"</worksheet>", b"<synthetic>unbound</synthetic></worksheet>"
                    )
                target.writestr(member, raw)
        padded = rewritten
    with pytest.raises(ValueError, match="ark|celle|grense"):
        read_excel(padded)


def test_large_synthetic_original_still_continues_and_reopens_without_source_files(tmp_path):
    """Høy komprimerbarhet er legitimt i en komplett base64-original."""
    original = (SOURCES / "originals/prop16s-ch2.html").read_bytes().decode("utf-8")
    root = synthetic_sources(
        tmp_path, "prop16s-ch2.html", original, original + "<!--" + "a" * (5 * 1024 * 1024) + "-->"
    )
    _, base = source_report(tmp_path)
    report = build_report(base, import_omtale(root, SUPPORT))
    frozen, excel = tmp_path / "synthetic-large.json", tmp_path / "synthetic-large.xlsx"
    export_pair(report, frozen, excel)
    shutil.rmtree(root)
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    source = next(s for s in saved.basis.sources if s.name == "originals/prop16s-ch2.html")
    assert len(source_bytes(source)) > 5 * 1024 * 1024
    assert all(r.status == "kildeverifisert" for r in saved.relations)


@pytest.mark.parametrize("limit", ["cells", "dimension", "zip_members", "xml_metadata"])
def test_synthetic_xlsx_resource_limits_are_checked_before_materializing_workbook(tmp_path, limit):
    _, base = source_report(tmp_path)
    report = build_report(base, import_omtale(SOURCES, SUPPORT))
    frozen, excel = tmp_path / "valid.json", tmp_path / "valid.xlsx"
    export_pair(report, frozen, excel)
    damaged = tmp_path / "synthetic-resource.xlsx"
    if limit == "cells":
        book = load_workbook(excel)
        sheet = book["Saksgrunnlag"]
        for row in range(1, 701):
            for column in range(1, 101):
                sheet.cell(row, column, 1)
        book.save(damaged)
        book.close()
    else:
        with ZipFile(excel) as source, ZipFile(damaged, "w", ZIP_DEFLATED) as target:
            for member in source.infolist():
                raw = source.read(member)
                if limit == "dimension" and member.filename == "xl/worksheets/sheet1.xml":
                    raw = re.sub(rb"<dimension[^>]+/>", b'<dimension ref="A1:XFD1048576"/>', raw)
                if limit == "xml_metadata" and member.filename == "xl/styles.xml":
                    raw += b" " * (3 * 1024 * 1024)
                target.writestr(member, raw)
            if limit == "zip_members":
                for n in range(260):
                    target.writestr(f"synthetic-{n}.xml", b"<synthetic/>")
    with pytest.raises(ValueError, match="grense"):
        read_excel(damaged)


@pytest.mark.parametrize(
    "filename,old,new,issue",
    [
        ("prop16s-ch2.html", "Kap. 1429 Riksantikvaren", "Kap. 1428 Riksantikvaren", "konto"),
        ("inns81s.html", "støtter regjeringens", "har en ukjent holdning til", "ukjent_rolle"),
        ("inns81s.html", "Innst. 81 S", "Innst. 82 S", "dokumentidentitet"),
    ],
)
def test_synthetic_account_unknown_role_and_document_version_stay_unresolved_after_readback(
    tmp_path, filename, old, new, issue
):
    _, base = source_report(tmp_path)
    root = synthetic_sources(tmp_path, filename, old, new)
    report = build_report(base, import_omtale(root, SUPPORT))
    excel, frozen = tmp_path / "synthetic.xlsx", tmp_path / "synthetic.json"
    export_pair(report, frozen, excel)
    report = read_excel(excel)
    assert read_json(frozen) == report
    assert any(issue in i for i in report.issues)
    assert all(r.status == "uavklart" for r in report.relations)


@pytest.mark.parametrize(
    "field,value",
    [
        ("locator", "/html/body/p[999]"),
        ("text", "Syntetisk endret tekst"),
        ("accounts", ["1429/22"]),
        ("actor", "mindretall"),
        ("document_id", "inns-82-s-2024-2025"),
    ],
)
def test_rehashed_synthetic_json_does_not_replace_actual_source_evidence(tmp_path, field, value):
    _, base = source_report(tmp_path)
    report = build_report(base, import_omtale(SOURCES, SUPPORT))
    data = {"format": "omtale-pilot-1", "report": asdict(report)}
    data = json.loads(json.dumps(data))
    section = next(s for s in data["report"]["basis"]["sections"] if s["key"] == "P1")
    section[field] = value
    # Recompute all outer hashes: independent source replay must still reject the mutation.
    from innlesning.observations import identity

    for obj, namespace in (
        (section, "omtale-seksjon-1"),
        (data["report"]["basis"], "omtale-grunnlag-1"),
        (data["report"], "omtale-pilot-1"),
    ):
        obj["id"] = ""
        obj["id"] = identity(namespace, obj)
    frozen = tmp_path / "synthetic-forged.json"
    frozen.write_text(json.dumps(data), "utf-8")
    with pytest.raises(ValueError, match="gjenleste"):
        read_json(frozen)


def test_long_formula_like_synthetic_text_continues_and_excel_tampering_is_detected(tmp_path):
    _, base = source_report(tmp_path)
    text = "=1+1 " + "lang syntetisk tekst " * 4000 + "For å redusere"
    root = synthetic_sources(tmp_path, "prop16s-ch2.html", "For å redusere", text)
    report = build_report(base, import_omtale(root, SUPPORT))
    frozen, excel = tmp_path / "long.json", tmp_path / "long.xlsx"
    export_pair(report, frozen, excel)
    assert read_excel(excel) == read_json(frozen)
    book = load_workbook(excel)
    row = next(
        n for n in range(2, book["Omtale"].max_row + 1) if book["Omtale"].cell(n, 2).value == "P1"
    )
    assert book["Omtale"].cell(row, 11).data_type == "s"
    assert "Omtale fortsettelse" in book
    book["Omtale fortsettelse"].cell(2, 3).value = "Syntetisk forfalsket fortsettelse"
    damaged = tmp_path / "synthetic-damaged.xlsx"
    book.save(damaged)
    book.close()
    with pytest.raises(ValueError, match="videreført"):
        read_excel(damaged)


def test_repeat_and_new_extractor_keep_first_frozen_export_and_stable_amount_ids(tmp_path):
    _, base = source_report(tmp_path)
    first = build_report(base, import_omtale(SOURCES, SUPPORT))
    assert first == build_report(base, import_omtale(SOURCES, SUPPORT))
    frozen, excel = tmp_path / "first.json", tmp_path / "first.xlsx"
    export_pair(first, frozen, excel)
    hashes = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (frozen, excel)]
    second = build_report(base, import_omtale(SOURCES, SUPPORT, extractor="o2a-html-pdf-1-recheck"))
    assert first.basis.id != second.basis.id
    assert first.base_json == second.base_json
    assert {s.id for s in first.basis.sections}.isdisjoint(s.id for s in second.basis.sections)
    assert read_excel(excel) == read_json(frozen) == first
    assert [hashlib.sha256(p.read_bytes()).hexdigest() for p in (frozen, excel)] == hashes
    with pytest.raises(ValueError, match="nye"):
        export_pair(second, frozen, excel)


@pytest.mark.parametrize("condition", ["missing", "corrupt", "contradictory", "failed_manifest"])
@pytest.mark.parametrize("extractor", ["o2a-html-pdf-1", "o2b-html-pdf-1"])
def test_new_authentic_decision_ingestion_and_export_ignore_bad_omtale_without_historical_excel(
    tmp_path, condition, extractor
):
    _, base = source_report(tmp_path, posts=("99",))
    root = tmp_path / "missing"
    if condition == "contradictory":
        root = synthetic_sources(tmp_path, "prop16s-ch2.html", "Kap. 1429", "Kap. 9999")
    elif condition in {"corrupt", "failed_manifest"}:
        root = tmp_path / "synthetic-failed"
        shutil.copytree(SOURCES, root)
        if condition == "corrupt":
            (root / "originals/prop16s-ch2.html").write_bytes(b"synthetic corrupt source")
        else:
            (root / "source-manifest.json").write_text("synthetic invalid JSON", "utf-8")
    report = build_report(base, import_omtale(root, SUPPORT, extractor=extractor))
    frozen, excel = tmp_path / "empty.json", tmp_path / "empty.xlsx"
    export_pair(report, frozen, excel)
    assert read_excel(excel) == read_json(frozen)
    assert not report.visible_section_ids and report.issues
    book = load_workbook(excel)
    assert not base.base.rows
    assert "Ingen beløpsrader" in book["Historikk"].cell(2, 13).value
    assert book["Omtaleavvik"].max_row > 1 and book["Omtalekilder"].max_row >= 11
    book.close()
    # Ordinær eksisterende CLI tar en annen autentisk ny kilde inn i en helt ny database.
    # --historisk-proeve gjelder Stortingets eldre sesjon, ingen historisk Excel brukes.
    db = tmp_path / "new-official.sqlite"
    ordinary = tmp_path / "ordinary.xlsx"
    completed = subprocess.run(
        [
            sys.executable,
            "-m",
            "innlesning",
            "synkroniser",
            "--database",
            str(db),
            "--sesjon",
            "2024-2025",
            "--historisk-proeve",
            "--vedtak-id",
            "40034229",
            "--response",
            str(SUPPORT / "b6/40034229.xml"),
            "--metadata",
            str(SUPPORT / "b6/40034229.json"),
            "--excel",
            str(ordinary),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert completed.returncode in (0, 2), completed.stdout + completed.stderr
    summary = json.loads(completed.stdout)
    assert summary["antall"]["ny"] == 1 and not summary["feil"]
    revision = summary["vedtak"][0]["revision_id"]
    interpreted = subprocess.run(
        [
            sys.executable,
            "-m",
            "innlesning",
            "tolk",
            "--database",
            str(db),
            "--revisjon",
            revision,
            "--tolkning",
            "storting-belop-v2",
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert interpreted.returncode == 0, interpreted.stdout + interpreted.stderr
    extraction = json.loads(interpreted.stdout)["belopsuttrekk"]
    v2_excel = tmp_path / "ordinary-v2.xlsx"
    exported = subprocess.run(
        [
            sys.executable,
            "-m",
            "analyse",
            "--database",
            str(db),
            "--belopsuttrekk",
            extraction,
            "--excel",
            str(v2_excel),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert exported.returncode == 0, exported.stdout + exported.stderr
    book = load_workbook(v2_excel)
    try:
        assert book["Beløpsobservasjoner"].max_row > 1
    finally:
        book.close()


@pytest.mark.parametrize("extractor", ["o2a-html-pdf-1", "o2b-html-pdf-1"])
def test_explicit_cli_and_frozen_reexport_need_no_original_sources_or_database(tmp_path, extractor):
    from analyse.saksdokumenter import report_to_json

    _, base = source_report(tmp_path, posts=("1",))
    selected = tmp_path / "base.json"
    selected.write_text(report_to_json(base), "utf-8")
    frozen, excel = tmp_path / "cli.json", tmp_path / "cli.xlsx"
    first = subprocess.run(
        [
            sys.executable,
            "-m",
            "analyse.omtale",
            "--base",
            str(selected),
            "--kilder",
            str(SOURCES),
            "--stotte",
            str(SUPPORT),
            "--uttrekksversjon",
            extractor,
            "--json",
            str(frozen),
            "--excel",
            str(excel),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert first.returncode == 0, first.stdout + first.stderr
    selected.unlink()
    second_json, second_excel = tmp_path / "offline.json", tmp_path / "offline.xlsx"
    second = subprocess.run(
        [
            sys.executable,
            "-m",
            "analyse.omtale",
            "--frosset",
            str(frozen),
            "--json",
            str(second_json),
            "--excel",
            str(second_excel),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    assert second.returncode == 0, second.stdout + second.stderr
    assert read_json(second_json) == read_excel(second_excel) == read_excel(excel)
