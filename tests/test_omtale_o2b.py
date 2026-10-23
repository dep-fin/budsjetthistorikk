"""#35 ved den godkjente original → skrevet/gjenlest JSON og Excel-grensen."""

import hashlib
import json
import re
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_b9 import source_report
from test_omtale import synthetic_sources

from analyse.omtale import build_report, export_pair, read_excel, read_json
from analyse.omtale_begge import propose_relation
from analyse.saksdokumenter import build_saks_report
from analyse.saksdokumenter import report_from_json as read_base
from innlesning.observations import identity
from innlesning.omtale import import_omtale

SOURCES = Path("docs/research/omtale-o1")
SUPPORT = Path("tests/fixtures")
EXTRACTOR = "o2b-html-pdf-1"


def test_authentic_both_posts_keep_shared_section_and_two_events_after_saved_readback(tmp_path):
    database, base = source_report(tmp_path, posts=("1", "22"))
    report = build_report(base, import_omtale(SOURCES, SUPPORT, extractor=EXTRACTOR))
    frozen, excel = tmp_path / "both.json", tmp_path / "both.xlsx"
    export_pair(report, frozen, excel)
    database.rename(tmp_path / "unavailable.sqlite")
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    sections = {s.key: s for s in saved.basis.sections}
    assert {"P22", "K22", "PT22"} <= sections.keys()
    assert len([s for s in saved.basis.sections if s.key == "P1"]) == 1
    p1_relations = [r for r in saved.relations if r.section_id == sections["P1"].id]
    assert {r.accounts for r in p1_relations} == {("1429/1",), ("1429/22",)}
    assert len({r.event_id for r in p1_relations}) == 2
    assert len(saved.relations) == 13
    assert all(r.status == "kildeverifisert" for r in saved.relations)
    assert "kulturmiljøforvaltning" in sections["P22"].text
    assert "kulturminneforvaltning" in sections["PT22"].text
    assert "lønnsforklaringen gjelder bare post 1" in p1_relations[1].reason
    assert "Vega" in sections["PN"].text
    assert sections["PN"].role == "nabokontekst_fullmakt"
    assert not any(r.section_id == sections["PN"].id for r in saved.relations)
    underlying = read_base(saved.base_json)
    assert len(underlying.base.unfiltered_rows) == 6
    assert len({r.observation.step_id for r in underlying.base.unfiltered_rows}) == 2
    book = load_workbook(excel)
    try:
        assert book["Historikk"].max_column == 13
        assert book["Historikk"].max_row == 3
        assert tuple(book["Historikk"].cell(3, c).value for c in (10, 11, 12)) == (
            32_185_000,
            -3_000_000,
            29_185_000,
        )
        assert json.loads(book["Omtale"].cell(2, 8).value) == []  # PDF representation
    finally:
        book.close()


@pytest.mark.parametrize(
    "posts,visible_count", [(("1",), 3), (("22",), 3), (("1", "22"), 6), (("99",), 0)]
)
def test_saved_filters_keep_global_inventory_shared_ids_and_issues(tmp_path, posts, visible_count):
    db, original = source_report(tmp_path, posts=("1", "22"))
    basis = import_omtale(SOURCES, SUPPORT, extractor=EXTRACTOR)
    both = build_report(original, basis)
    sections = {s.key: s for s in basis.sections}
    target = next(r for r in both.relations if r.accounts == ("1429/1",))
    report = propose_relation(
        both,
        section_id=sections["K22"].id,
        event_id=target.event_id,
        evidence_ids=(sections["K22"].id,),
        reason=(
            "Syntetisk hypotese: post 22-sammendrag kan belyse mottakerposten; "
            "ikke dokumentert som egen forklaring."
        ),
    )
    selected = build_saks_report(
        db, replace(original.selection, base=replace(original.selection.base, posts=posts))
    )
    from analyse.omtale_begge import build_both_report

    report = build_both_report(
        selected, basis, proposals=tuple(r for r in report.relations if r.status == "uavklart")
    )
    frozen, excel = tmp_path / "filter.json", tmp_path / "filter.xlsx"
    export_pair(report, frozen, excel)
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    selected_base = read_base(saved.base_json)
    assert len(selected_base.base.rows) == visible_count
    assert len(selected_base.base.unfiltered_rows) == 6
    assert len(saved.relations) == 14
    assert len([r for r in saved.relations if r.status == "kildeverifisert"]) == 13
    assert any(i.startswith("uavklart_relasjonsforslag:") for i in saved.issues)
    assert {s.key: s.id for s in saved.basis.sections} == {s.key: s.id for s in both.basis.sections}
    assert (sections["P1"].id in saved.visible_section_ids) == bool(visible_count)
    if posts == ("1",):
        assert {r.observation.post for r in selected_base.base.rows} == {"1"}
        assert sections["P22"].id not in saved.visible_section_ids
    if not visible_count:
        assert not saved.visible_section_ids
    book = load_workbook(excel)
    try:
        assert book["Omtaleavvik"].max_row >= 3
        assert book["Omtalekilder"].max_row >= 11
        assert book["Relasjoner"].max_row == 15
    finally:
        book.close()


@pytest.mark.parametrize(
    "field", ["wrong_post", "salary_scope", "section_text", "section_locator", "status_upgrade"]
)
def test_synthetic_rehashed_false_binding_is_rejected_from_written_json_and_excel(tmp_path, field):
    _, base = source_report(tmp_path)
    report = build_report(base, import_omtale(SOURCES, SUPPORT, extractor=EXTRACTOR))
    good_json, good_excel = tmp_path / "original.json", tmp_path / "original.xlsx"
    export_pair(report, good_json, good_excel)
    data = json.loads(good_json.read_text("utf-8"))
    p22 = next(s for s in data["report"]["basis"]["sections"] if s["key"] == "P22")
    relation = next(r for r in data["report"]["relations"] if r["section_id"] == p22["id"])
    if field == "wrong_post":
        post1 = next(r for r in data["report"]["relations"] if r["accounts"] == ["1429/1"])
        for name in ("event_id", "decision_revision_id", "observation_ids", "accounts"):
            relation[name] = post1[name]
    elif field == "salary_scope":
        relation["reason"] = (
            "Syntetisk forfalskning: lønnsoppgjøret begrunner reduksjonen på post 22."
        )
    elif field == "status_upgrade":
        relation["professional_status"] = "godkjent"
    else:
        p22["text" if field == "section_text" else "locator"] = (
            "Syntetisk tekst/lokator uten kildebevis"
        )
        p22["id"] = ""
        p22["id"] = identity("omtale-seksjon-1", p22)
    for obj, namespace in (
        (relation, "omtale-relasjon-1"),
        (data["report"]["basis"], "omtale-grunnlag-1"),
        (data["report"], "omtale-pilot-1"),
    ):
        obj["id"] = ""
        obj["id"] = identity(namespace, obj)
    raw = json.dumps(data, ensure_ascii=False)
    forged_json = tmp_path / "synthetic-forged.json"
    forged_json.write_text(raw, "utf-8")
    with pytest.raises(ValueError, match="gjenleste"):
        read_json(forged_json)
    # Samme forfalskede snapshot er faktisk skrevet inn i den fysiske arbeidsboken.
    book = load_workbook(good_excel)
    snapshot = book["Omtalegrunnlag"]
    snapshot.delete_rows(2, snapshot.max_row)
    for i in range(0, len(raw), 30_000):
        snapshot.append([i // 30_000 + 1, raw[i : i + 30_000]])
    forged_excel = tmp_path / "synthetic-forged.xlsx"
    book.save(forged_excel)
    book.close()
    with pytest.raises(ValueError):
        read_excel(forged_excel)


@pytest.mark.parametrize(
    "old,new",
    [
        ("kap. 1429, post 01", "kap. 1429, post 22"),
        (
            "som foreslås omdisponert til kap. 1429, post 01.",
            "Syntetisk løs omtale: 3 mill. kroner og kap. 1429 post 1 finnes også her.",
        ),
    ],
)
def test_synthetic_post22_conflict_or_account_amount_coincidence_never_verifies(tmp_path, old, new):
    root = synthetic_sources(tmp_path, "prop16s-ch2.html", old, new)
    _, base = source_report(tmp_path, posts=("22",))
    report = build_report(base, import_omtale(root, SUPPORT, extractor=EXTRACTOR))
    frozen, excel = tmp_path / "synthetic.json", tmp_path / "synthetic.xlsx"
    export_pair(report, frozen, excel)
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    assert "forklaring_omfang_motstrid:P22" in saved.issues
    assert not any(r.status == "kildeverifisert" for r in saved.relations)
    assert len(read_base(saved.base_json).base.unfiltered_rows) == 6


def test_unresolved_formula_like_long_proposal_keeps_evidence_without_replacing_chain(tmp_path):
    _, base = source_report(tmp_path)
    report = build_report(base, import_omtale(SOURCES, SUPPORT, extractor=EXTRACTOR))
    sections = {s.key: s for s in report.basis.sections}
    target = next(r for r in report.relations if r.accounts == ("1429/1",))
    reason = "=1+1 Syntetisk uavklart forslag. " + "Belegg må vurderes. " * 4000
    proposed = propose_relation(
        report,
        section_id=sections["K22"].id,
        event_id=target.event_id,
        evidence_ids=(sections["K22"].id,),
        reason=reason,
    )
    frozen, excel = tmp_path / "proposal.json", tmp_path / "proposal.xlsx"
    export_pair(proposed, frozen, excel)
    hashes = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (frozen, excel)]
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    assert saved.relations[:-1] == report.relations
    proposal = saved.relations[-1]
    assert proposal.reason == reason and proposal.status == "uavklart"
    assert proposal.evidence_ids == (sections["K22"].id,)
    assert proposal.professional_status == "ikke_faglig_godkjent"
    assert [hashlib.sha256(p.read_bytes()).hexdigest() for p in (frozen, excel)] == hashes
    book = load_workbook(excel)
    try:
        assert book["Relasjoner"].cell(15, 12).data_type == "s"
        assert "Relasjoner fortsettelse" in book.sheetnames
    finally:
        book.close()
    with pytest.raises(ValueError):
        propose_relation(
            report,
            section_id=sections["PN"].id,
            event_id=target.event_id,
            evidence_ids=(sections["PN"].id,),
            reason="Syntetisk feil fullmaktskobling",
        )


def test_new_explicit_b_revision_keeps_old_saved_result_and_stable_amount_identity(tmp_path):
    _, base = source_report(tmp_path)
    basis = import_omtale(SOURCES, SUPPORT, extractor=EXTRACTOR)
    first = build_report(base, basis)
    assert first == build_report(base, import_omtale(SOURCES, SUPPORT, extractor=EXTRACTOR))
    frozen, excel = tmp_path / "first.json", tmp_path / "first.xlsx"
    export_pair(first, frozen, excel)
    hashes = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (frozen, excel)]
    second = build_report(base, import_omtale(SOURCES, SUPPORT, extractor=EXTRACTOR + "-recheck"))
    assert len(second.relations) == 13 and len(second.basis.sections) == 11
    assert first.basis.id != second.basis.id
    assert first.base_json == second.base_json
    assert {s.id for s in first.basis.sections}.isdisjoint(s.id for s in second.basis.sections)
    assert read_json(frozen) == read_excel(excel) == first
    assert [hashlib.sha256(p.read_bytes()).hexdigest() for p in (frozen, excel)] == hashes
    with pytest.raises(ValueError, match="nye"):
        export_pair(second, frozen, excel)


def test_synthetic_incomplete_local_table_stays_visible_after_both_saved_readbacks(tmp_path):
    original = (SOURCES / "originals/inns81s.html").read_bytes().decode("utf-8")
    match = re.search(r"<table\b.*?</table>", original, re.S)
    assert match is not None
    # Ingen kapittelrad, og en kort rad før den tilsynelatende post 22-gruppen.
    table = (
        "<table><tbody><tr><td>Syntetisk ufullstendig tabell</td></tr>"
        "<tr><td></td><td></td><td></td><td></td></tr>"
        "<tr><td></td><td>22</td><td>Flerårige prosjekter kulturminneforvaltning, "
        "kan overføres, reduseres med</td><td>3 000 000</td></tr>"
        "<tr><td></td><td></td><td>fra kr 32 185 000 til kr 29 185 000</td><td></td></tr>"
        "</tbody></table>"
    )
    root = synthetic_sources(tmp_path, "inns81s.html", match.group(0), table)
    _, base = source_report(tmp_path)
    report = build_report(base, import_omtale(root, SUPPORT, extractor=EXTRACTOR))
    frozen, excel = tmp_path / "incomplete.json", tmp_path / "incomplete.xlsx"
    export_pair(report, frozen, excel)
    saved = read_json(frozen)
    assert saved == read_excel(excel)
    assert "konto_eller_forslagsbelop_motstrid_post22:KT" in saved.issues
    assert not any(r.status == "kildeverifisert" for r in saved.relations)
