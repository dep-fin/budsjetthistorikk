"""P1: kildebundet register gjennom de avtalte offentlige filgrensene."""

import gzip
from pathlib import Path

import pytest

from analyse.perioderegister import build_report, import_basis
from innlesning.visningspilot_model import PilotSelection


@pytest.fixture(scope="module")
def authentic_basis():
    old = gzip.decompress(Path("tests/fixtures/t2/authentic-t1.json.gz").read_bytes()).decode()
    return import_basis(Path("."), t1_json=old)


def test_authentic_authorization_and_global_candidates_do_not_create_amounts(authentic_basis):
    basis = authentic_basis
    report = build_report(basis, PilotSelection(years=(2024,), chapters=("1429",)))
    fullmakt = next(e for e in basis.entries if e.key == "original-346")
    assert (fullmakt.decision_api_id, fullmakt.budget_year, fullmakt.accounts) == (
        "40032933",
        2024,
        ("1429/1",),
    )
    assert fullmakt.role == "adopted_revenue_authorization"
    assert not fullmakt.observation_ids and not fullmakt.event_ids
    assert (
        len(report.view.monetary.adopted_observation_ids),
        len(report.view.monetary.adopted_event_ids),
    ) == (11, 5)
    assert (len(basis.case_candidates), len(basis.decision_candidates), len(basis.gaps)) == (
        67,
        25,
        7,
    )
    assert basis.year_coverage == "partial" and basis.professional_annual_balance_approval is False
    doc = next(d for d in basis.documents if d.key == "prop1")
    assert any(
        "161 243 000" in s.text and "32 185 000" in s.text
        for s in basis.statements
        if s.document_id == doc.id
    )
    control = next(d for d in basis.documents if d.key == "prop36")
    assert any(
        "prop36-ch10.html" in s.text for s in basis.statements if s.document_id == control.id
    )


def test_authorization_reopens_from_actual_json_and_excel_without_amount_cells(
    tmp_path, authentic_basis
):
    from openpyxl import load_workbook

    from analyse.perioderegister import open_document, read_json
    from analyse.perioderegister_excel import export_pair, read_excel

    original = build_report(
        authentic_basis, PilotSelection(years=(2024,), accounts=(("1429", "1"),))
    )
    fullmakt = next(e for e in authentic_basis.entries if e.key == "original-346")
    report = open_document(original, fullmakt.document_id)
    frozen, excel = tmp_path / "fullmakt.json", tmp_path / "fullmakt.xlsx"
    export_pair(report, frozen, excel)
    assert read_json(frozen) == read_excel(excel) == report
    assert not report.view.monetary.adopted_event_ids
    book = load_workbook(excel)
    try:
        rows = list(book["Forløp"].iter_rows(min_row=2, values_only=True))
        assert len(rows) == 1 and rows[0][6:10] == (None, None, None, None)
        assert rows[0][4] == "opprinnelig / adopted_revenue_authorization"
        assert {i.chapter + "/" + i.post for i in authentic_basis.inventory if i.chapter} == {
            "1429/1",
            "4429/2",
            "4429/9",
            "1633/1",
        }
    finally:
        book.close()


@pytest.mark.parametrize(
    "posts,years,counts",
    [
        (("1",), (2024,), (7, 3)),
        (("22",), (2024,), (4, 2)),
        (("1", "22"), (2024,), (11, 5)),
        (("1", "22"), (2025,), (0, 0)),
        (("1", "22"), (2023,), (0, 0)),
        ((), (), (0, 0)),
    ],
)
def test_selection_preserves_monetary_boundary_and_global_register(
    authentic_basis, posts, years, counts
):
    report = build_report(
        authentic_basis, PilotSelection(years=years, accounts=tuple(("1429", p) for p in posts))
    )
    assert (
        len(report.view.monetary.adopted_observation_ids),
        len(report.view.monetary.adopted_event_ids),
    ) == counts
    assert (
        len(report.view.case_candidate_ids),
        len(report.view.decision_candidate_ids),
        len(report.view.gap_ids),
    ) == (67, 25, 7)


def test_both_navigation_directions_keep_unknown_index_year_and_return_selection(authentic_basis):
    from analyse.perioderegister import go_back, open_accounts, open_document

    original = build_report(
        authentic_basis, PilotSelection(years=(2024,), accounts=(("1429", "22"),))
    )
    candidate = next(
        c for c in authentic_basis.case_candidates if c.status == "open_title_candidate"
    )
    assert (
        candidate.budget_year is None
        and candidate.accounts == ()
        and candidate.post_link == "unknown"
    )
    document = open_document(original, candidate.document_id)
    assert not document.view.entry_ids and not document.view.monetary.observation_ids
    assert go_back(document) == original
    reverse = open_accounts(document, original.selection)
    assert go_back(reverse) == build_report(authentic_basis, document.selection)
    assert document.basis is original.basis


def test_complete_index_occurrences_and_journal_reasons_are_preserved(authentic_basis):
    import json
    from collections import Counter

    root = Path("docs/research/k3-arsinventar-50")
    source = json.loads((root / "case-candidates.json").read_text("utf8"))
    assert len({c.case_id for c in authentic_basis.case_candidates}) == 66
    assert Counter(c.status for c in authentic_basis.case_candidates) == {
        "investigated_case_references": 7,
        "open_title_candidate": 60,
    }
    assert [
        (c.index_session, c.ordinal_1based, json.loads(c.metadata_json), c.status, c.reason)
        for c in authentic_basis.case_candidates
    ] == [
        (r["session"], r["ordinal_1based"], r["metadata"], r["status"], r["reason"]) for r in source
    ]
    journal = json.loads((root / "coverage-journal.json").read_text("utf8"))["decisions"]
    assert [
        (c.decision_api_id, c.status, c.reason) for c in authentic_basis.decision_candidates
    ] == [(r["id"], r["status"], r["reason"]) for r in journal]
    assert {
        c.decision_api_id: c.budget_year
        for c in authentic_basis.decision_candidates
        if c.decision_api_id in {"40032936", "40033007", "40034038"}
    } == {"40032936": 2023, "40033007": 2023, "40034038": 2025}


def test_frozen_json_rejects_rehashed_fullmakt_and_global_gap_lies(authentic_basis):
    from analyse.perioderegister import report_from_json, report_to_json
    from ops.perioderegister_proeve import MUTATIONS, mutate

    report = build_report(authentic_basis, PilotSelection(years=(2024,), chapters=("1429",)))
    raw = report_to_json(report)
    for name in MUTATIONS:
        with pytest.raises(ValueError):
            report_from_json(mutate(raw, name))


def test_new_resource_caps_and_payload_duplication_are_enforced(authentic_basis):
    from dataclasses import replace

    from analyse.perioderegister import validate_basis

    for name, cap in (
        ("entries", 256),
        ("case_candidates", 67),
        ("decision_candidates", 25),
        ("gaps", 7),
        ("evidence", 4096),
        ("sources", 512),
    ):
        invalid = replace(
            authentic_basis, **{name: (getattr(authentic_basis, name)[0],) * (cap + 1)}
        )
        with pytest.raises(ValueError, match="ressursgrense"):
            validate_basis(invalid)
    with pytest.raises(ValueError, match="Duplisert P1-payload"):
        validate_basis(
            replace(
                authentic_basis, payloads=authentic_basis.payloads + authentic_basis.payloads[:1]
            )
        )


def test_rehashed_excel_attack_reaches_semantics_before_rejection(tmp_path, authentic_basis):
    import json

    from analyse.omtale_excel_guard import inspect_excel
    from analyse.perioderegister import report_to_json
    from analyse.perioderegister_excel import SHEET_NAMES, _book, read_excel
    from analyse.session_serialization import _decode
    from innlesning.perioderegister_model import Report
    from ops.perioderegister_proeve import mutate

    valid = build_report(authentic_basis, PilotSelection(years=(2024,), chapters=("1429",)))
    raw = mutate(report_to_json(valid), "positive-coverage")
    poisoned = _decode(Report, json.loads(raw)["report"])
    book = _book(poisoned, raw)
    path = tmp_path / "semantic.xlsx"
    book.save(path)
    book.close()
    capsule, _ = inspect_excel(
        path,
        sheet_names=SHEET_NAMES,
        snapshot_name="Offlinegrunnlag",
        max_expanded_bytes=128 * 1024 * 1024,
    )
    assert capsule == raw
    with pytest.raises(ValueError, match="P1-grunnlaget avviker"):
        read_excel(path)


@pytest.fixture(scope="module")
def saved_pair(tmp_path_factory, authentic_basis):
    from analyse.perioderegister_excel import export_pair

    directory = tmp_path_factory.mktemp("saved-p1")
    report = build_report(authentic_basis, PilotSelection(years=(2024,), chapters=("1429",)))
    frozen, excel = directory / "both.json", directory / "both.xlsx"
    export_pair(report, frozen, excel)
    return frozen, excel, report


def test_external_package_main_document_is_rejected(tmp_path, saved_pair):
    from zipfile import ZIP_DEFLATED, ZipFile

    from lxml import etree

    from analyse.perioderegister_excel import read_excel

    path = tmp_path / "external-main-document.xlsx"
    with ZipFile(saved_pair[1]) as source, ZipFile(path, "w", ZIP_DEFLATED) as target:
        for member in source.infolist():
            raw = source.read(member)
            if member.filename == "_rels/.rels":
                root = etree.fromstring(raw)
                root[0].set("Target", "https://example.invalid/synthetic.xlsx")
                root[0].set("TargetMode", "External")
                raw = etree.tostring(root)
            target.writestr(member, raw)
    with pytest.raises(ValueError, match="P1-Excel.*relasjon"):
        read_excel(path)


def test_disabled_raw_alignment_is_rejected(tmp_path, saved_pair):
    from zipfile import ZIP_DEFLATED, ZipFile

    from lxml import etree

    from analyse.perioderegister_excel import read_excel

    path = tmp_path / "disabled-alignment.xlsx"
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(saved_pair[1]) as source, ZipFile(path, "w", ZIP_DEFLATED) as target:
        for member in source.infolist():
            raw = source.read(member)
            if member.filename == "xl/styles.xml":
                root = etree.fromstring(raw)
                for xf in root.find(namespace + "cellXfs"):
                    xf.set("applyAlignment", "0")
                raw = etree.tostring(root)
            target.writestr(member, raw)
    with pytest.raises(ValueError, match="P1-Excel.*rå.*stil"):
        read_excel(path)


@pytest.mark.parametrize(
    "mutation",
    [
        "duplicate-root-id",
        "missing-root-id",
        "duplicate-main-document",
        "wrong-main-target",
        "wrong-main-type",
        "external-style-target",
        "worksheet-relationship",
        "wrong-workbook-content-type",
        "duplicate-content-type",
        "missing-content-type",
    ],
)
def test_package_relationships_and_content_types_are_bound_to_local_parts(
    tmp_path, saved_pair, mutation
):
    from copy import deepcopy
    from zipfile import ZIP_DEFLATED, ZipFile

    from lxml import etree

    from analyse.perioderegister_excel import read_excel

    path = tmp_path / (mutation + ".xlsx")
    with ZipFile(saved_pair[1]) as source, ZipFile(path, "w", ZIP_DEFLATED) as target:
        for member in source.infolist():
            raw = source.read(member)
            if member.filename not in {
                "_rels/.rels",
                "xl/_rels/workbook.xml.rels",
                "[Content_Types].xml",
            }:
                target.writestr(member, raw)
                continue
            root = etree.fromstring(raw)
            if member.filename == "_rels/.rels":
                if mutation == "duplicate-root-id":
                    root[1].set("Id", root[0].get("Id"))
                elif mutation == "missing-root-id":
                    del root[0].attrib["Id"]
                elif mutation == "duplicate-main-document":
                    root.append(deepcopy(root[0]))
                    root[-1].set("Id", "another-main")
                elif mutation == "wrong-main-target":
                    root[0].set("Target", "docProps/core.xml")
                elif mutation == "wrong-main-type":
                    root[0].set("Type", root[2].get("Type"))
            elif (
                member.filename == "xl/_rels/workbook.xml.rels"
                and mutation == "external-style-target"
            ):
                root[-2].set("TargetMode", "External")
                root[-2].set("Target", "https://example.invalid/styles.xml")
            elif member.filename == "[Content_Types].xml":
                if mutation == "wrong-workbook-content-type":
                    root[-1].set("ContentType", "application/xml")
                elif mutation == "duplicate-content-type":
                    root.append(deepcopy(root[-1]))
                elif mutation == "missing-content-type":
                    root.remove(root[-1])
            target.writestr(member, etree.tostring(root))
        if mutation == "worksheet-relationship":
            target.writestr(
                "xl/worksheets/_rels/sheet1.xml.rels",
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/'
                '2006/relationships/hyperlink" Target="https://example.invalid/" '
                'TargetMode="External"/></Relationships>',
            )
    with pytest.raises(ValueError):
        read_excel(path)


@pytest.mark.parametrize(
    "flag,value",
    [
        ("applyAlignment", None),
        ("applyFont", "0"),
        ("applyFont", "1"),
        ("applyFill", "0"),
        ("applyFill", "1"),
        ("applyNumberFormat", "0"),
        ("applyNumberFormat", "1"),
        ("applyProtection", "0"),
    ],
)
def test_raw_style_application_flags_are_preserved(tmp_path, saved_pair, flag, value):
    from zipfile import ZIP_DEFLATED, ZipFile

    from lxml import etree

    from analyse.perioderegister_excel import read_excel

    path = tmp_path / (flag + str(value) + ".xlsx")
    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(saved_pair[1]) as source, ZipFile(path, "w", ZIP_DEFLATED) as target:
        for member in source.infolist():
            raw = source.read(member)
            if member.filename == "xl/styles.xml":
                root = etree.fromstring(raw)
                # Første egentlige cellestil er synlig overskrift med skrift,
                # fyll og justering. Endre bare råflagget, ikke stilelementene.
                styles = root.find(namespace + "cellXfs")
                xf = (
                    next(e for e in styles if e.get("numFmtId") != "0")
                    if flag == "applyNumberFormat"
                    else styles[1]
                )
                if value is None:
                    del xf.attrib[flag]
                else:
                    xf.set(flag, value)
                raw = etree.tostring(root)
            target.writestr(member, raw)
    with pytest.raises(ValueError, match="P1-Excel.*rå.*stil"):
        read_excel(path)


@pytest.mark.parametrize(
    "mutation",
    [
        "amount",
        "hidden-row",
        "hidden-column",
        "hyperlink",
        "filter",
        "doctype",
        "numberformat",
        "hidden-sheet",
    ],
)
def test_false_hidden_or_active_excel_content_is_rejected(tmp_path, saved_pair, mutation):
    from analyse.perioderegister_excel import read_excel
    from ops.perioderegister_proeve import mutate_display

    path = tmp_path / (mutation + ".xlsx")
    mutate_display(saved_pair[1], path, mutation)
    with pytest.raises(ValueError):
        read_excel(path)


def test_old_ids_calculations_and_one_unchanged_t1_capsule_survive_actual_files(saved_pair):
    from analyse.fleraarspilot import build_report as old_report
    from analyse.perioderegister import read_json
    from analyse.perioderegister_excel import read_excel

    frozen, excel, report = saved_pair
    old = old_report(
        report.basis.t4_basis,
        PilotSelection(years=(2024,), accounts=(("1429", "1"), ("1429", "22"))),
    )
    assert read_json(frozen) == read_excel(excel) == report
    assert report.view.monetary == old.view
    assert report.calculations == old.calculations
    assert [c.amount_kr for c in report.calculations] == [166039000, 29185000]
    assert {c.status for c in report.calculations} == {"avgrenset kontroll"}


def test_frozen_input_can_run_public_cli_in_empty_directory_with_network_sqlite_blocked(
    tmp_path, saved_pair
):
    import os
    import subprocess
    import sys

    from analyse.perioderegister import read_json
    from analyse.perioderegister_excel import read_excel

    script = """import sys,socket,sqlite3
def forbidden(*a,**kw): raise AssertionError('offline boundary')
socket.create_connection=forbidden
socket.socket.connect=forbidden
sqlite3.connect=forbidden
sys.argv=['analyse.perioderegister','--frosset',sys.argv[1],'--rot',sys.argv[2],
          '--json',sys.argv[3],'--excel',sys.argv[4]]
from analyse.perioderegister import main
main()
"""
    environment = dict(os.environ)
    environment["PYTHONPATH"] = os.pathsep.join(
        (str(Path.cwd()), str(Path(".local/python").resolve()))
    )
    frozen, excel = tmp_path / "offline.json", tmp_path / "offline.xlsx"
    process = subprocess.run(
        [
            sys.executable,
            "-c",
            script,
            str(saved_pair[0]),
            str(tmp_path / "missing"),
            str(frozen),
            str(excel),
        ],
        cwd=tmp_path,
        env=environment,
        capture_output=True,
        text=True,
    )
    assert process.returncode == 0, process.stdout + process.stderr
    assert read_json(frozen) == read_excel(excel) == saved_pair[2]


def test_reused_plenary_original_is_a_reference_to_the_nested_old_source(authentic_basis):
    # Autentisk plenumskilde ligger allerede i den uendrede O2b/T1-kapselen.
    payload = next(
        p
        for p in authentic_basis.payloads
        if p.sha256 == "22170a60f769e89642b2e52525899903020de25938aa2b566ea28d11f85fb2e5"
    )
    assert payload.content_base64 is None and payload.legacy_source_id is not None
