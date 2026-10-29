"""P2s avtalte kilde-, utvalgs- og faktiske filgrenser."""

import gzip
from pathlib import Path

import pytest

from innlesning.visningspilot_model import PilotSelection


@pytest.fixture(scope="module")
def authentic_basis():
    from analyse.perioderegister_r2 import import_basis

    old = gzip.decompress(Path("tests/fixtures/t2/authentic-t1.json.gz").read_bytes()).decode()
    return import_basis(Path("."), t1_json=old)


def test_three_branches_keep_the_original_p1_and_source_roles(authentic_basis):
    from analyse.perioderegister import build_report as p1_report
    from analyse.perioderegister_r2 import build_report

    basis = authentic_basis
    selected = PilotSelection(years=(2024,), chapters=("1429",))
    report = build_report(basis, selected)
    old = p1_report(basis.p1_basis, selected)
    assert report.view.p1 == old.view and report.calculations == old.calculations
    assert (len(basis.evidence), len(basis.findings), len(basis.cross_references)) == (74, 11, 12)
    assert (
        len(old.view.monetary.adopted_observation_ids),
        len(old.view.monetary.adopted_event_ids),
    ) == (11, 5)
    findings = {f.key: f for f in basis.findings}
    assert findings["D230"].accounts == ("1560/71",)
    assert findings["D231"].amounts_json == "null"
    assert findings["FULLMAKT429"].accounts == ()
    assert findings["FULLMAKT429"].conditional_accounts == ("1429/1",)
    assert findings["FULLMAKT429"].transfer_target_year == 2025
    assert findings["FULLMAKT429"].amounts_json == "null"
    assert findings["FIN-1429-1"].role == "kildeobservert_lønnskompensasjonsanslag"
    assert basis.year_coverage == "partial" and basis.professional_annual_balance_approval is False
    (update,) = basis.status_updates
    candidate = next(c for c in basis.p1_basis.case_candidates if c.case_id == "101000")
    assert update.case_candidate_id == candidate.id
    assert candidate.status == "open_title_candidate"
    assert (update.case_id, update.index_session, update.ordinal_1based) == (
        "101000",
        "2024-2025",
        624,
    )


def test_written_files_keep_the_proposed_value_and_conditional_authorization(
    tmp_path, authentic_basis
):
    from openpyxl import load_workbook

    from analyse.perioderegister_r2 import build_report, read_json
    from analyse.perioderegister_r2_excel import export_pair, read_excel

    report = build_report(authentic_basis, PilotSelection(years=(2024,), chapters=("1429",)))
    frozen, excel = tmp_path / "both.json", tmp_path / "both.xlsx"
    export_pair(report, frozen, excel)
    assert read_json(frozen) == read_excel(excel) == report
    book = load_workbook(excel)
    try:
        rows = list(book["Forløp"].iter_rows(min_row=2, values_only=True))
        estimate = next(r for r in rows if r[18] == "FIN-1429-1")
        assert estimate[6:10] == (None, None, None, None)
        assert estimate[19] == 2796000
        assert estimate[20] == "foreslått kompensasjonsanslag, kroner"
        authorization = next(r for r in rows if r[18] == "FULLMAKT429")
        assert authorization[6:10] == (None, None, None, None)
        assert authorization[19] is None and authorization[22] == 2025
        assert authorization[21] == '["1429/1"]'
        coverage = list(book["Dekning og avvik"].values)
        assert all(
            any(str(r[0]).startswith(g.key + " ") for r in coverage)
            for g in authentic_basis.p1_basis.gaps
        )
        assert any("hovedsak_undersokt_faktisk_postvirkning_uavklart" in str(r) for r in coverage)
    finally:
        book.close()


@pytest.mark.parametrize(
    "posts,years,counts,has_salary",
    [
        (("1",), (2024,), (7, 3), True),
        (("22",), (2024,), (4, 2), False),
        (("1", "22"), (2024,), (11, 5), True),
        (("1", "22"), (2025,), (0, 0), False),
        ((), (), (0, 0), False),
    ],
)
def test_accounts_keep_roles_and_all_global_ids(authentic_basis, posts, years, counts, has_salary):
    from analyse.perioderegister_r2 import build_report

    basis = authentic_basis
    report = build_report(
        basis, PilotSelection(years=years, accounts=tuple(("1429", p) for p in posts))
    )
    view = report.view.p1
    assert (
        len(view.monetary.adopted_observation_ids),
        len(view.monetary.adopted_event_ids),
    ) == counts
    assert view.case_candidate_ids == tuple(c.id for c in basis.p1_basis.case_candidates)
    assert view.decision_candidate_ids == tuple(c.id for c in basis.p1_basis.decision_candidates)
    assert view.gap_ids == tuple(g.id for g in basis.p1_basis.gaps)
    assert report.view.status_update_ids == tuple(s.id for s in basis.status_updates)
    assert report.view.cross_reference_ids == tuple(c.id for c in basis.cross_references)
    keys = {f.key for f in basis.findings if f.id in report.view.finding_ids}
    assert ("FIN-1429-1" in keys) is has_salary
    assert ("FULLMAKT429" in keys) is has_salary
    assert ("D231" in keys) is has_salary
    if years == (2024,) and posts == ("22",):
        assert keys == {"O1-OMDISPONERING"}
    elif not has_salary:
        assert keys == set()


def test_document_inventory_and_both_navigation_directions(authentic_basis):
    from analyse.perioderegister_r2 import (
        build_report,
        contents,
        go_back,
        open_accounts,
        open_document,
    )

    basis = authentic_basis
    original = build_report(basis, PilotSelection(years=(2024,), accounts=(("1429", "22"),)))
    docs = {d.key: d for d in contents(basis).documents}
    for key in (
        "prop34",
        "inns91",
        "vedtak230",
        "vedtak231",
        "vedtak429",
        "prop36",
        "inns90",
        "forslag23945",
        "votering23945",
    ):
        document = open_document(original, docs[key].id)
        assert go_back(document) == original
        account = open_accounts(document, original.selection)
        assert go_back(account) == build_report(basis, document.selection)
        assert document.basis is basis
        if key != "prop36":
            assert not document.view.p1.monetary.adopted_event_ids
    assert {
        i.chapter + "/" + i.post
        for i in basis.inventory
        if i.document_id == docs["prop34"].id and i.chapter
    } == {"1560/71"}
    assert {
        i.chapter + "/" + i.post
        for i in basis.inventory
        if i.document_id == docs["inns91"].id and i.chapter
    } == {"1560/71"}
    assert not docs["vedtak231"].accounts and not docs["vedtak429"].accounts
    p1_prop36 = next(d for d in basis.p1_basis.documents if d.key == "prop36")
    assert docs["prop36"] == p1_prop36


def test_rehashed_lies_are_rejected_at_the_semantic_file_boundary(authentic_basis):
    from analyse.perioderegister_r2 import build_report, report_from_json, report_to_json
    from ops.perioderegister_r2_proeve import MUTATIONS, mutate

    raw = report_to_json(
        build_report(authentic_basis, PilotSelection(years=(2024,), chapters=("1429",)))
    )
    for name in MUTATIONS:
        with pytest.raises(ValueError):
            report_from_json(mutate(raw, name))


def test_strict_json_and_resource_caps_reject_untrusted_extension(authentic_basis):
    from dataclasses import replace

    from analyse.perioderegister_r2 import report_from_json, validate_basis

    for raw in (
        '{"format":"perioderegister-r2-1","format":"perioderegister-r2-1","report":{}}',
        '{"format":"perioderegister-r2-1","report":NaN}',
        '{"format":"perioderegister-r2-1","report":{},"unknown":1}',
    ):
        with pytest.raises(ValueError):
            report_from_json(raw)
    for field, cap in (
        ("sources", 128),
        ("payloads", 128),
        ("evidence", 128),
        ("findings", 32),
        ("status_updates", 1),
        ("cross_references", 12),
    ):
        invalid = replace(
            authentic_basis, **{field: (getattr(authentic_basis, field)[0],) * (cap + 1)}
        )
        with pytest.raises(ValueError, match="ressursgrense"):
            validate_basis(invalid)


@pytest.fixture(scope="module")
def saved_pair(tmp_path_factory, authentic_basis):
    from analyse.perioderegister_r2 import build_report
    from analyse.perioderegister_r2_excel import export_pair

    output = tmp_path_factory.mktemp("actual-p2")
    report = build_report(authentic_basis, PilotSelection(years=(2024,), chapters=("1429",)))
    frozen, excel = output / "saved.json", output / "saved.xlsx"
    export_pair(report, frozen, excel)
    return frozen, excel


@pytest.mark.parametrize(
    "attack",
    [
        "research-amount",
        "hidden-row",
        "hidden-column",
        "filter",
        "numberformat",
        "hidden-sheet",
        "formula",
        "root-relationship",
        "content-type",
        "raw-style",
    ],
)
def test_excel_semantics_and_p1_audit_guards_still_bind_the_extension(saved_pair, tmp_path, attack):
    from zipfile import ZIP_DEFLATED, ZipFile

    from lxml import etree
    from openpyxl import load_workbook

    from analyse.perioderegister_r2_excel import read_excel
    from ops.perioderegister_proeve import mutate_display

    source = saved_pair[1]
    target = tmp_path / (attack + ".xlsx")
    if attack in {"hidden-row", "hidden-column", "filter", "numberformat", "hidden-sheet"}:
        mutate_display(source, target, attack)
    elif attack in {"research-amount", "formula"}:
        book = load_workbook(source)
        try:
            sheet = book["Forløp"]
            row = next(r for r in sheet.iter_rows(min_row=2) if r[18].value == "FIN-1429-1")
            row[19].value = 2796001 if attack == "research-amount" else "=1+1"
            book.save(target)
        finally:
            book.close()
    else:
        with ZipFile(source) as before, ZipFile(target, "w", ZIP_DEFLATED) as after:
            for entry in before.infolist():
                raw = before.read(entry.filename)
                if attack == "root-relationship" and entry.filename == "_rels/.rels":
                    tree = etree.fromstring(raw)
                    tree[0].set("Target", "https://example.invalid/synthetic")
                    tree[0].set("TargetMode", "External")
                    raw = etree.tostring(tree)
                elif attack == "content-type" and entry.filename == "[Content_Types].xml":
                    tree = etree.fromstring(raw)
                    tree[0].set("ContentType", "application/synthetic")
                    raw = etree.tostring(tree)
                elif attack == "raw-style" and entry.filename == "xl/styles.xml":
                    tree = etree.fromstring(raw)
                    tree.find("{*}cellXfs")[0].set("applyAlignment", "0")
                    raw = etree.tostring(tree)
                after.writestr(entry, raw)
    with pytest.raises(ValueError):
        read_excel(target)


@pytest.mark.parametrize(
    "attack",
    ["estimate-adopted", "authorization-used", "authorization-2025", "post22", "hidden-gaps"],
)
def test_rehashed_excel_pairs_pass_container_before_semantic_rejection(
    saved_pair, tmp_path, attack
):
    import json

    from analyse.omtale_excel_guard import inspect_excel
    from analyse.perioderegister_r2 import read_json
    from analyse.perioderegister_r2_excel import SHEET_NAMES, _book, read_excel
    from analyse.session_serialization import _decode
    from innlesning.perioderegister_r2_model import Report
    from ops.perioderegister_r2_proeve import mutate

    poisoned = mutate(saved_pair[0].read_text("utf8"), attack)
    frozen, excel = tmp_path / (attack + ".json"), tmp_path / (attack + ".xlsx")
    frozen.write_text(poisoned, encoding="utf8")
    book = _book(_decode(Report, json.loads(poisoned)["report"]), poisoned)
    try:
        book.save(excel)
    finally:
        book.close()
    capsule, _ = inspect_excel(
        excel,
        sheet_names=SHEET_NAMES,
        snapshot_name="Offlinegrunnlag",
        max_expanded_bytes=128 * 1024 * 1024,
    )
    assert capsule == poisoned
    with pytest.raises(ValueError, match="grunnlaget avviker"):
        read_json(frozen)
    with pytest.raises(ValueError, match="grunnlaget avviker"):
        read_excel(excel)
