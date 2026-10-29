"""T5 prøves ved original/typet basis og offentlige visnings-/filgrenser."""

import gzip
import hashlib
import json
import shutil
import subprocess
import sys
from dataclasses import asdict, replace
from pathlib import Path

import pytest

from innlesning.observations import encode
from innlesning.omtale import identified


@pytest.fixture(scope="module")
def original_t1():
    from analyse.visningspilot import report_from_json

    return report_from_json(
        gzip.decompress(Path("tests/fixtures/t2/authentic-t1.json.gz").read_bytes()).decode()
    )


@pytest.fixture(scope="module")
def t5_basis(original_t1):
    from analyse.proposisjon16 import import_basis
    from analyse.visningspilot import report_to_json

    return import_basis(report_to_json(original_t1), Path("."))


def test_complete_table_preserves_old_capsule_and_matches_independent_source(original_t1):
    from analyse.proposisjon16 import contents, import_basis
    from analyse.visningspilot import report_to_json

    basis = import_basis(report_to_json(original_t1), Path("."))
    assert basis.t1_basis == original_t1.basis
    assert len(basis.observations) == 105
    whole = contents(basis)
    assert len(whole.observations) == 123
    reference = json.loads(
        Path("docs/research/to-visninger-37/originals/forgrening/inventory.json").read_text("utf8")
    )
    assert reference["summary"]["prop_change_posts"] == 37
    proposed = [o for o in whole.observations if o.figure_type == "regjeringsforslag"]
    assert len(proposed) == 111
    zero = next(
        o for o in proposed if (o.chapter, o.post, o.meaning) == ("4400", "2", "etter_belop")
    )
    assert zero.amount_kr == 0 and zero.raw_value == "0"
    assert "amount_kr" not in dict(zero.null_reasons)
    for entry in reference["account_comparison"]:
        source = entry["proposal"]
        if source is None:
            continue
        rows = {
            o.meaning: o
            for o in proposed
            if (o.chapter, o.post) == (source["chapter"], source["post"])
        }
        for meaning, key in (
            ("for_belop", "before"),
            ("endring", "delta"),
            ("etter_belop", "after"),
        ):
            assert rows[meaning].amount_kr == source[key]
        assert rows["endring"].pointer.locator == source["xpath"] + "/td[4]"
    prop = next(d for d in basis.t1_basis.documents if d.key == "prop16")
    items = [i for i in basis.t1_basis.inventory if i.document_id == prop.id]
    assert len({(i.chapter, i.post) for i in items}) == 45
    assert len({o.chapter for o in proposed}) == 15
    assert sum(o.meaning == "endring" and int(o.chapter) < 3000 for o in proposed) == 27
    assert sum(o.meaning == "endring" and int(o.chapter) >= 3000 for o in proposed) == 10
    assert prop.structured_coverage == "delvis_1429"
    assert basis.coverage[0].table_fields == 111
    assert {b.inventory_id for b in basis.bindings} == {
        i.id for i in items if i.role == "endringstabell"
    }


def test_both_entries_and_return_use_same_identities_through_saved_json(tmp_path, original_t1):
    from analyse import proposisjon16 as t5
    from analyse.visningspilot import report_to_json
    from innlesning.visningspilot_model import PilotSelection

    basis = t5.import_basis(report_to_json(original_t1), Path("."))
    selection = PilotSelection(
        years=(2024,), chapters=("1429",), accounts=(("4400", "2"), ("1420", "66"))
    )
    report = t5.build_report(basis, selection)
    assert len(report.view.observation_ids) == 21
    assert len(report.view.adopted_event_ids) == 2
    prop_id = next(d.id for d in basis.t1_basis.documents if d.key == "prop16")
    opened = t5.open_document(report, prop_id)
    assert len(opened.view.observation_ids) == 111
    assert len(opened.view.inventory_ids) > 45
    path = tmp_path / "opened.json"
    t5.export_json(opened, path)
    assert t5.go_back(t5.read_json(path)) == report
    assert t5.go_back(t5.open_accounts(opened, selection)) == t5.build_report(
        basis, opened.selection
    )


def test_authorizations_keep_one_shared_frame_and_exact_scoped_branches(original_t1):
    from analyse.proposisjon16 import import_basis
    from analyse.visningspilot import report_to_json

    basis = import_basis(report_to_json(original_t1), Path("."))
    assert len(basis.authorizations) == 9
    assert {a.roman for a in basis.authorizations} == {"II", "III", "IV", "V", "VI"}
    shared = [a for a in basis.authorizations if a.roman == "IV"]
    assert len(shared) == 1
    assert shared[0].accounts == ("1420/39", "1420/69", "1420/79")
    assert shared[0].frame_kr == 16_300_000
    v3 = next(b for b in basis.branches if b.point == "V.3")
    assert (v3.case_id, v3.committee_point, v3.decision_ids) == ("100927", "II", ("40034049",))
    v12 = next(b for b in basis.branches if b.point == "V.1-2")
    assert (v12.case_id, v12.committee_point, v12.decision_ids) == ("100889", "V", ("40034224",))
    branch95 = next(b for b in basis.branches if b.point == "I" and b.case_id == "100889")
    assert len(branch95.accounts) == 35 and branch95.amount_stage == "ikke_strukturert_T6"
    assert all("1533/45" not in b.accounts for b in basis.branches)


def test_two_beriktiget_representations_keep_distinct_hashes_without_correction_history(
    original_t1,
):
    from analyse.proposisjon16 import import_basis
    from analyse.visningspilot import report_to_json

    basis = import_basis(report_to_json(original_t1), Path("."))
    reused = next(s for s in basis.sources if s.media_type == "text/html")
    assert reused.sha256 == "1f159a3ca1df3e0bc35617644d0b36f7c92fd5684947f6eca5510595af533a03"
    current = next(
        s
        for s in basis.t1_basis.sources
        if s.document_key == "inns95" and s.media_type == "text/html"
    )
    assert current.sha256 == "ae61227fa83adf04fb45956b1bd714a75592bbc1d6c70da300e554afef806aa6"
    assert "inns95_to_beriktigede_html:lik_hovedtekst_ukjent_rettelseshistorikk" in basis.issues


def test_observed_income_zero_is_numeric_and_visible_in_saved_pair(tmp_path, original_t1):
    from openpyxl import load_workbook

    from analyse import proposisjon16 as t5
    from analyse.proposisjon16_excel import export_pair, read_excel
    from analyse.visningspilot import report_to_json
    from innlesning.visningspilot_model import PilotSelection

    basis = t5.import_basis(report_to_json(original_t1), Path("."))
    report = t5.build_report(basis, PilotSelection(years=(2024,), accounts=(("4400", "2"),)))
    frozen, excel = tmp_path / "income.json", tmp_path / "income.xlsx"
    export_pair(report, frozen, excel)
    assert t5.read_json(frozen) == read_excel(excel) == report
    book = load_workbook(excel)
    try:
        assert [r[6].value for r in book["Forløp"].iter_rows(min_row=2)] == [518000, -518000, 0]
        assert book["Forløp"]["G4"].data_type == "n"
        assert all(s.sheet_state == "visible" for s in book)
        assert len(list(book["Fullmakter"].iter_rows(min_row=2))) == 9
        assert len(list(book["Grenbevis"].iter_rows(min_row=2))) == 8
        assert book["Fullmakter"]["L1"].value == "Valgt"
        assert all(r[11].value == "nei" for r in book["Fullmakter"].iter_rows(min_row=2))
    finally:
        book.close()


def test_context_reference_does_not_inherit_sister_branch(original_t1):
    from analyse import proposisjon16 as t5
    from analyse.visningspilot import report_to_json
    from innlesning.visningspilot_model import PilotSelection

    basis = t5.import_basis(report_to_json(original_t1), Path("."))
    report = t5.build_report(basis, PilotSelection(years=(2024,), accounts=(("1533", "45"),)))
    documents = {d.key: d.id for d in basis.t1_basis.documents}
    assert documents["prop16"] in report.view.document_ids
    assert documents["inns95"] not in report.view.document_ids
    assert not report.view.observation_ids


@pytest.fixture(scope="module")
def saved_cases(tmp_path_factory, t5_basis):
    from analyse import proposisjon16 as t5
    from analyse.proposisjon16_excel import export_pair
    from innlesning.visningspilot_model import PilotSelection

    directory = tmp_path_factory.mktemp("t5-files")
    prop_id = next(d.id for d in t5_basis.t1_basis.documents if d.key == "prop16")
    selections = {
        "utgift": PilotSelection(years=(2024,), accounts=(("1400", "1"),)),
        "inntekt-null": PilotSelection(years=(2024,), accounts=(("4400", "2"),)),
        "tekst": PilotSelection(years=(2024,), accounts=(("1420", "66"),)),
        "flere": PilotSelection(
            years=(2024,), chapters=("1429",), accounts=(("4400", "2"), ("1420", "22"))
        ),
        "tomt": PilotSelection(years=(2024,)),
        "hele-prop": PilotSelection(mode="document", document_id=prop_id),
    }
    reports = {name: t5.build_report(t5_basis, selection) for name, selection in selections.items()}
    reports["konto-til-dokument"] = t5.open_document(reports["flere"], prop_id)
    reports["dokument-til-konto"] = t5.open_accounts(reports["hele-prop"], selections["flere"])
    reports["retur-konto"] = t5.go_back(reports["konto-til-dokument"])
    reports["retur-dokument"] = t5.go_back(reports["dokument-til-konto"])
    result = {}
    for name, report in reports.items():
        frozen, excel = directory / (name + ".json"), directory / (name + ".xlsx")
        export_pair(report, frozen, excel)
        result[name] = report, frozen, excel
    return result


def test_complete_saved_pairs_preserve_selection_counts_and_legacy(
    t5_basis, original_t1, saved_cases
):
    from analyse import proposisjon16 as t5
    from analyse.proposisjon16_excel import read_excel
    from analyse.visningspilot import report_from_json, report_to_json

    expected = {
        "utgift": 3,
        "inntekt-null": 3,
        "tekst": 0,
        "flere": 21,
        "tomt": 0,
        "hele-prop": 111,
        "konto-til-dokument": 111,
        "dokument-til-konto": 21,
        "retur-konto": 21,
        "retur-dokument": 111,
    }
    for name, (report, frozen, excel) in saved_cases.items():
        assert t5.read_json(frozen) == read_excel(excel) == report
        assert len(report.view.observation_ids) == expected[name]
        assert report.basis == t5_basis
        assert report.basis.coverage[0].table_fields == 111
    assert saved_cases["retur-konto"][0] == saved_cases["flere"][0]
    assert saved_cases["retur-dokument"][0] == saved_cases["hele-prop"][0]
    assert report_from_json(report_to_json(original_t1)) == original_t1
    assert t5_basis.t1_basis.observations == original_t1.basis.observations
    assert not set(o.id for o in t5_basis.observations) & set(
        o.id for o in original_t1.basis.observations
    )
    from openpyxl import load_workbook

    book = load_workbook(saved_cases["tekst"][2])
    try:
        selected = [r for r in book["Fullmakter"].iter_rows(min_row=2) if r[11].value == "ja"]
        assert [(r[0].value, r[1].value, r[2].value) for r in selected] == [("V", "1", "1420/66")]
        assert book["Kontoer"]["G1"].value == "Historisk T1-strukturstatus"
        assert book["Kontoer"]["H1"].value == "Gjeldende T5-tabellstatus"
    finally:
        book.close()


def _poison(report, mutation):
    """Merkede syntetiske semantikkangrep; alle ytre objekt-ID-er rehashes."""
    basis = report.basis
    if mutation in {"amount", "role", "year", "zero", "account"}:
        index = next(
            i
            for i, o in enumerate(basis.observations)
            if (o.chapter, o.post, o.meaning) == ("4400", "2", "etter_belop")
        )
        row = basis.observations[index]
        changes = {
            "amount": {"amount_kr": 1},
            "role": {"figure_type": "stortingsvedtak"},
            "year": {"budget_year": 2025},
            "zero": {"amount_kr": None},
            "account": {"chapter": "1429", "post": "22"},
        }[mutation]
        observations = list(basis.observations)
        observations[index] = replace(row, **changes)
        basis = replace(basis, observations=tuple(observations))
    elif mutation == "branch":
        first = basis.branches[1]
        branch = identified(
            "proposisjonsgren16-1", replace(first, accounts=first.accounts + ("1533/45",))
        )
        basis = replace(basis, branches=(basis.branches[0], branch, *basis.branches[2:]))
    elif mutation == "frame":
        auth = next(a for a in basis.authorizations if a.roman == "IV")
        split = tuple(
            identified("proposisjonsfullmakt16-1", replace(auth, accounts=(account,)))
            for account in auth.accounts
        )
        basis = replace(
            basis, authorizations=tuple(a for a in basis.authorizations if a != auth) + split
        )
    elif mutation == "old_id":
        old = basis.t1_basis
        old = identified(
            "visningsgrunnlag-1",
            replace(
                old,
                observations=tuple(
                    replace(o, step_id="f" * 64) if o.figure_type == "stortingsvedtak" else o
                    for o in old.observations
                ),
            ),
        )
        basis = replace(basis, t1_basis=old)
    elif mutation == "coverage":
        coverage = identified(
            "proposisjonsdekning16-1",
            replace(
                basis.coverage[0], status="full_semantisk_strukturering", scope="hele_dokumentet"
            ),
        )
        basis = replace(basis, coverage=(coverage,))
    elif mutation == "hide_gap":
        basis = replace(basis, issues=())
    else:
        raise AssertionError(mutation)
    basis = identified("proposisjonsgrunnlag16-1", basis)
    return identified("proposisjon16-1", replace(report, basis=basis))


@pytest.mark.parametrize(
    "mutation",
    [
        "amount",
        "role",
        "year",
        "zero",
        "account",
        "branch",
        "frame",
        "old_id",
        "coverage",
        "hide_gap",
    ],
)
def test_rehashed_semantic_poison_reaches_both_actual_readers(tmp_path, saved_cases, mutation):
    from analyse import proposisjon16 as t5
    from analyse.omtale_excel_guard import inspect_excel
    from analyse.proposisjon16_excel import SHEET_NAMES, _book, read_excel

    poisoned = _poison(saved_cases["inntekt-null"][0], mutation)
    raw = encode({"format": "proposisjon16-1", "report": asdict(poisoned)})
    damaged_json = tmp_path / "synthetic-semantic.json"
    damaged_json.write_text(raw, "utf8")
    with pytest.raises(ValueError, match="avviker"):
        t5.read_json(damaged_json)
    # Konsistent forfalskning av hele snapshot OG presentasjon, ikke bare ytre hash/chunks.
    book = _book(poisoned, raw)
    path = tmp_path / "synthetic-semantic.xlsx"
    book.save(path)
    book.close()
    assert (
        inspect_excel(
            path,
            sheet_names=SHEET_NAMES,
            snapshot_name="Visningsgrunnlag",
            max_expanded_bytes=128 * 1024 * 1024,
        )[0]
        == raw
    )
    with pytest.raises(ValueError, match="avviker"):
        read_excel(path)


@pytest.mark.parametrize("condition", ["missing", "failed", "corrupt", "contradictory"])
def test_bad_optional_t5_never_blocks_real_ordinary_registration(tmp_path, original_t1, condition):
    from analyse import proposisjon16 as t5
    from analyse.visningspilot import report_to_json
    from innlesning.observations import read_extraction
    from innlesning.proposisjon16 import SOURCE_CATALOG

    root = tmp_path / "optional"
    for name in SOURCE_CATALOG:
        path = root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(name, path)
    path = root / next(iter(SOURCE_CATALOG))
    if condition == "missing":
        path.unlink()
    elif condition == "failed":
        path.unlink()
        path.mkdir()
    elif condition == "corrupt":
        path.write_bytes(b"synthetic invalid HTML")
    else:
        path.write_bytes(path.read_bytes().replace(b"1429", b"1420"))
    basis = t5.import_basis(report_to_json(original_t1), root)
    assert basis.sources[0].status == (
        "mangler" if condition in {"missing", "failed"} else "motstrid"
    )
    assert basis.t1_basis == original_t1.basis
    database = tmp_path / "ordinary.sqlite"

    def cli(*args):
        result = subprocess.run(
            [sys.executable, "-m", *args], capture_output=True, text=True, encoding="utf8"
        )
        assert result.returncode in (0, 2), result.stdout + result.stderr
        return json.loads(result.stdout)

    summary = cli(
        "innlesning",
        "synkroniser",
        "--database",
        str(database),
        "--sesjon",
        "2024-2025",
        "--historisk-proeve",
        "--vedtak-id",
        "40034229",
        "--response",
        "tests/fixtures/b6/40034229.xml",
        "--metadata",
        "tests/fixtures/b6/40034229.json",
    )
    assert summary["antall"]["ny"] == 1 and not summary["feil"]
    extraction_id = cli(
        "innlesning",
        "tolk",
        "--database",
        str(database),
        "--revisjon",
        summary["vedtak"][0]["revision_id"],
        "--tolkning",
        "storting-belop-v2",
    )["belopsuttrekk"]
    extraction = read_extraction(database, extraction_id)
    assert len(extraction.observations) == 222
    assert len({o.step_id for o in extraction.observations}) == 74
    assert (
        extraction.artifact.sha256
        == "66208daf51111356ad73c26985819870ab9c2f97ec7b070024d89e150cd772b9"
    )
    excel = tmp_path / "ordinary.xlsx"
    cli(
        "analyse",
        "--database",
        str(database),
        "--belopsuttrekk",
        extraction_id,
        "--excel",
        str(excel),
    )
    assert excel.exists()


def test_offline_cli_originals_no_clobber_and_strict_json(tmp_path, saved_cases, monkeypatch):
    import socket
    import sqlite3

    from analyse import proposisjon16 as t5
    from analyse.proposisjon16_excel import export_originals, export_pair, read_excel

    report, frozen, excel = saved_cases["inntekt-null"]
    before = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (frozen, excel)]
    with pytest.raises(ValueError):
        export_pair(report, frozen, excel)
    assert before == [hashlib.sha256(p.read_bytes()).hexdigest() for p in (frozen, excel)]
    for invalid in (
        '{"format":"proposisjon16-1","format":"proposisjon16-1","report":{}}',
        '{"format":"proposisjon16-1","report":NaN}',
        '{"format":"proposisjon16-1","report":Infinity}',
    ):
        with pytest.raises(ValueError):
            t5.report_from_json(invalid)

    def forbidden(*args, **kwargs):
        raise AssertionError("Offline skal ikke åpne database eller nett")

    monkeypatch.setattr(socket, "socket", forbidden)
    monkeypatch.setattr(sqlite3, "connect", forbidden)
    assert read_excel(excel) == t5.read_json(frozen) == report
    originals = tmp_path / "originals"
    export_originals(report, originals)
    assert (originals / "o2b/originals/prop16s.pdf").read_bytes() == Path(
        "docs/research/omtale-o1/originals/prop16s.pdf"
    ).read_bytes()
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "analyse.proposisjon16",
            "--frosset",
            str(frozen),
            "--rot",
            str(tmp_path / "does-not-exist"),
            "--json",
            str(tmp_path / "offline.json"),
            "--excel",
            str(tmp_path / "offline.xlsx"),
        ],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert t5.read_json(tmp_path / "offline.json") == report


@pytest.mark.parametrize(
    "mutation",
    [
        "numeric_type",
        "formula",
        "link",
        "hide",
        "text_part",
        "style_flag",
        "zip_member",
        "doctype",
        "chunk_order",
    ],
)
def test_saved_excel_defenses_bind_presentation_and_container(tmp_path, saved_cases, mutation):
    from zipfile import ZipFile

    from lxml import etree
    from openpyxl import load_workbook

    from analyse.proposisjon16_excel import read_excel

    source = saved_cases["inntekt-null"][2]
    damaged = tmp_path / "synthetic-presentation.xlsx"
    if mutation in {"style_flag", "zip_member", "doctype"}:
        with ZipFile(source) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        if mutation == "style_flag":
            root = etree.fromstring(files["xl/styles.xml"])
            ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            root.xpath("//s:cellXfs/s:xf", namespaces=ns)[0].set("applyNumberFormat", "0")
            files["xl/styles.xml"] = etree.tostring(root)
        elif mutation == "zip_member":
            files["unexpected.bin"] = b"synthetic"
        else:
            files["xl/workbook.xml"] = (
                b'<!DOCTYPE workbook [<!ENTITY x "unsafe">]>' + files["xl/workbook.xml"]
            )
        with ZipFile(damaged, "w") as archive:
            for name, raw in files.items():
                archive.writestr(name, raw)
    else:
        book = load_workbook(source)
        if mutation == "numeric_type":
            book["Forløp"]["G4"] = "0"
        elif mutation == "formula":
            book["Fulltekst"]["K2"] = "=1+1"
        elif mutation == "link":
            book["Forløp"]["B2"].hyperlink.location = "'Dokumenter'!B3"
        elif mutation == "hide":
            book["T5-dekning"].sheet_state = "hidden"
        elif mutation == "text_part":
            book["Fulltekst"]["J2"] = 2
        else:
            book["Visningsgrunnlag"]["A2"] = 2
        book.save(damaged)
        book.close()
    with pytest.raises(ValueError):
        read_excel(damaged)


@pytest.mark.parametrize("mutation", ["srgb-white", "system-white", "minor-font"])
def test_saved_excel_rejects_theme_that_hides_current_status(tmp_path, saved_cases, mutation):
    from zipfile import ZIP_DEFLATED, ZipFile

    from lxml import etree
    from openpyxl import load_workbook

    from analyse.omtale_excel_guard import inspect_excel
    from analyse.proposisjon16_excel import SHEET_NAMES, read_excel

    report, frozen, source = saved_cases["inntekt-null"]
    assert read_excel(source) == report
    with ZipFile(source) as archive:
        original = {name: archive.read(name) for name in archive.namelist()}
    files = dict(original)
    drawing = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    theme = etree.fromstring(files["xl/theme/theme1.xml"])
    if mutation == "minor-font":
        latin = theme.find(".//" + drawing + "minorFont/" + drawing + "latin")
        assert latin is not None
        latin.set("typeface", "Wingdings")
    else:
        dark = theme.find(".//" + drawing + "clrScheme/" + drawing + "dk1")
        assert dark is not None
        for color in list(dark):
            dark.remove(color)
        if mutation == "srgb-white":
            etree.SubElement(dark, drawing + "srgbClr", val="FFFFFF")
        else:
            etree.SubElement(dark, drawing + "sysClr", val="window", lastClr="FFFFFF")
    files["xl/theme/theme1.xml"] = etree.tostring(theme)
    assert [name for name in files if files[name] != original[name]] == ["xl/theme/theme1.xml"]
    damaged = tmp_path / "synthetic-theme-white.xlsx"
    with ZipFile(damaged, "w", ZIP_DEFLATED) as archive:
        for name, raw in files.items():
            archive.writestr(name, raw)
    # Uendret, gyldig kapsel passerer fortsatt første beholderkontroll.
    assert inspect_excel(
        damaged,
        sheet_names=SHEET_NAMES,
        snapshot_name="Visningsgrunnlag",
        max_expanded_bytes=128 * 1024 * 1024,
    )[0] == frozen.read_text("utf8")
    book = load_workbook(damaged)
    try:
        for title, coordinate in (("Inventar", "M2"), ("Inventar", "N2"), ("Kontoer", "H2")):
            cell = book[title][coordinate]
            assert cell.font.color.type == "theme" and cell.font.color.theme == 1
            assert cell.font.scheme == "minor"
            assert cell.fill.patternType is None
    finally:
        book.close()
    with pytest.raises(ValueError, match="T5-Excel.*tema"):
        read_excel(damaged)


def test_adapter_accepts_typed_input_without_analysis_import(tmp_path, t5_basis):
    import pickle

    from analyse.visningspilot import source_input

    prepared = tmp_path / "typed.pickle"
    prepared.write_bytes(
        pickle.dumps(
            (t5_basis.t1_basis, source_input(t5_basis.t1_basis.o2b_json), t5_basis.sources)
        )
    )
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            """
import importlib.abc, pickle, sys
from pathlib import Path
class NoAnalysis(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname == 'analyse' or fullname.startswith('analyse.'):
            raise ModuleNotFoundError(fullname)
sys.meta_path.insert(0, NoAnalysis())
from innlesning.proposisjon16 import structure_basis, validate_basis
with Path(sys.argv[1]).open('rb') as handle:
    old, prepared, sources = pickle.load(handle)
basis = structure_basis(old, prepared, sources)
validate_basis(basis, prepared)
assert len(basis.observations) == 105
assert not any(n == 'analyse' or n.startswith('analyse.') for n in sys.modules)
print(basis.id)
""",
            str(prepared),
        ],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert result.stdout.strip() == t5_basis.id


def test_exact_year_intersection_of_chapters_and_pairs_has_no_cross_product(t5_basis):
    from analyse import proposisjon16 as t5
    from innlesning.visningspilot_model import PilotSelection

    selected = t5.build_report(
        t5_basis, PilotSelection(years=(2024,), accounts=(("1429", "1"), ("1420", "22")))
    )
    assert len(selected.view.observation_ids) == 9
    observations = [
        o for o in t5.contents(t5_basis).observations if o.id in selected.view.observation_ids
    ]
    assert {(o.chapter, o.post) for o in observations} == {("1429", "1")}
    assert not t5.build_report(
        t5_basis, PilotSelection(years=(2025,), chapters=("1429", "4400"))
    ).view.observation_ids
    assert not t5.build_report(t5_basis, PilotSelection(chapters=("1429",))).view.observation_ids
    for account, fields, adopted, events in (("1", 9, 3, 1), ("22", 9, 3, 1), ("60", 0, 0, 0)):
        view = t5.build_report(
            t5_basis, PilotSelection(years=(2024,), accounts=(("1429", account),))
        ).view
        assert (
            len(view.observation_ids),
            len(view.adopted_observation_ids),
            len(view.adopted_event_ids),
        ) == (fields, adopted, events)


def test_long_formula_like_legacy_text_stays_literal_and_lossless_in_t5(tmp_path, original_t1):
    from openpyxl import load_workbook

    from analyse import proposisjon16 as t5
    from analyse import visningspilot as t1
    from analyse.omtale import report_from_json as read_o2
    from analyse.omtale import report_to_json as freeze_o2
    from analyse.omtale_begge import propose_relation
    from analyse.proposisjon16_excel import export_pair, read_excel
    from innlesning.visningspilot_model import PilotSelection

    old = read_o2(original_t1.basis.o2b_json)
    target = old.relations[0]
    reason = "=SYNTHETIC('" + "Lang syntetisk kontrolltekst øæå " * 2200 + "')"
    proposed = propose_relation(
        old,
        section_id=old.basis.sections[1].id,
        event_id=target.event_id,
        evidence_ids=(old.basis.sections[1].id,),
        reason=reason,
    )
    legacy = t1.build_report(t1.import_basis(freeze_o2(proposed), Path(".")), PilotSelection())
    report = t5.build_report(
        t5.import_basis(t1.report_to_json(legacy), Path(".")), PilotSelection()
    )
    frozen, excel = tmp_path / "synthetic-long.json", tmp_path / "synthetic-long.xlsx"
    export_pair(report, frozen, excel)
    assert t5.read_json(frozen) == read_excel(excel) == report
    book = load_workbook(excel)
    try:
        assert "Relasjoner fortsettelse" in book.sheetnames
        assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
        assert reason in report.basis.t1_basis.o2b_json
    finally:
        book.close()


def test_initial_cli_uses_explicit_t1_without_changing_legacy_file(tmp_path, original_t1):
    from analyse import proposisjon16 as t5
    from analyse.visningspilot import report_to_json

    legacy = tmp_path / "legacy.json"
    legacy.write_text(report_to_json(original_t1), "utf8")
    before = hashlib.sha256(legacy.read_bytes()).hexdigest()
    frozen, excel = tmp_path / "initial.json", tmp_path / "initial.xlsx"
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "analyse.proposisjon16",
            "--t1",
            str(legacy),
            "--rot",
            ".",
            "--aar",
            "2024",
            "--konto",
            "4400/2",
            "--json",
            str(frozen),
            "--excel",
            str(excel),
        ],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert len(t5.read_json(frozen).view.observation_ids) == 3
    assert hashlib.sha256(legacy.read_bytes()).hexdigest() == before
