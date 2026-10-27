"""T1: avtalte vertikale grenser fra autentisk kilde til faktiske filer."""

import base64
import hashlib
import json
import shutil
import subprocess
import sys
from copy import copy
from dataclasses import asdict, replace
from pathlib import Path

import pytest
from test_b9 import source_report

from analyse.omtale import build_report, report_to_json
from analyse.visningspilot import import_basis
from innlesning.omtale import import_omtale
from innlesning.visningspilot_model import PilotSelection


@pytest.fixture(scope="module")
def pilot_basis(tmp_path_factory):
    root = tmp_path_factory.mktemp("pilot")
    _, base = source_report(root, posts=("1", "22"))
    old = build_report(
        base,
        import_omtale(
            Path("docs/research/omtale-o1"),
            Path("tests/fixtures"),
            extractor="o2b-html-pdf-1",
        ),
    )
    return import_basis(report_to_json(old), Path("."))


def test_authentic_basis_has_18_fields_and_complete_inventory(pilot_basis):
    assert len(pilot_basis.observations) == 18
    assert len([o for o in pilot_basis.observations if o.figure_type == "stortingsvedtak"]) == 6
    docs = {d.key: d for d in pilot_basis.documents}
    assert docs["inns81"].accounts == ("1429/1", "1429/22", "1429/60")
    prop = [r for r in pilot_basis.inventory if r.document_id == docs["prop16"].id]
    assert len([r for r in prop if r.role == "endringstabell"]) == 37
    assert len({r.chapter + "/" + r.post for r in prop}) == 45
    assert docs["prop16"].structured_coverage == "delvis_1429"
    assert (
        len({o.step_id for o in pilot_basis.observations if o.figure_type == "stortingsvedtak"})
        == 2
    )


def test_source_adapter_runs_with_analysis_package_unavailable(tmp_path, pilot_basis):
    import pickle

    from analyse.visningspilot import source_input

    prepared = tmp_path / "validated-input.pickle"
    prepared.write_bytes(pickle.dumps(source_input(pilot_basis.o2b_json)))
    result = subprocess.run(
        [
            sys.executable,
            "-c",
            """
import importlib.abc, pickle, sys
from pathlib import Path
class NoAnalysis(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname == 'analyse' or fullname.startswith('analyse.'):
            raise ModuleNotFoundError(fullname)
sys.meta_path.insert(0, NoAnalysis())
from innlesning.visningspilot import import_basis, structure_basis
with Path(sys.argv[1]).open('rb') as handle:
    prepared = pickle.load(handle)
basis = import_basis(prepared, Path('.'))
assert len(basis.observations) == 18
assert basis == structure_basis(prepared, basis.sources)
assert not any(k == 'analyse' or k.startswith('analyse.') for k in sys.modules)
print(basis.id)
""",
            str(prepared),
        ],
        capture_output=True,
        text=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert result.stdout.strip() == pilot_basis.id


def test_new_statement_hierarchies_follow_original_sections_in_both_saved_views(
    tmp_path, pilot_basis
):
    from lxml import html
    from openpyxl import load_workbook

    from analyse.visningspilot import build_report as build
    from analyse.visningspilot import read_json
    from analyse.visningspilot_excel import export_pair, read_excel

    original = html.fromstring(
        Path("docs/research/omtale-o1/originals/inns81s.html").read_text("utf8")
    )
    # Uavhengig kildefasit: overskrifter med originalens egne ankre, ikke adapterens helper.
    headings = {
        key: " ".join(original.get_element_by_id(key + "-title").text_content().split())
        for key in ("m1", "m1_1", "m1_1_1", "m1_1_2", "m1_1_3", "m2", "m3")
    }
    assert headings["m1"] == "Sammendrag"
    assert headings["m1_1"] == "Kap. 1429 Riksantikvaren"
    assert headings["m1_1_3"] == "Post 60 Kulturminnekompetanse i kommunene, kan overføres"
    docs = {d.key: d for d in pilot_basis.documents}
    for name, selection in (
        ("document", PilotSelection(mode="document", document_id=docs["inns81"].id)),
        ("accounts", PilotSelection(years=(2024,), chapters=("1429",))),
    ):
        report = build(pilot_basis, selection)
        frozen, excel = tmp_path / (name + ".json"), tmp_path / (name + ".xlsx")
        export_pair(report, frozen, excel)
        assert read_json(frozen) == read_excel(excel) == report
        for key in headings:
            locator = original.getroottree().getpath(original.get_element_by_id(key))
            statement = next(
                s
                for s in report.basis.statements
                if s.locator == locator
                and s.role
                in {
                    "sammendrag_innledning",
                    "kapitteloverskrift",
                    "gjengivelse_av_regjeringen",
                    "samlet_stotte",
                    "tilraading",
                }
            )
            prefixes = ["_".join(key.split("_")[:n]) for n in range(1, len(key.split("_")) + 1)]
            assert statement.hierarchy == tuple(headings[prefix] for prefix in prefixes)
        ii = next(s for s in report.basis.statements if s.role == "stikkordsfullmakt_tilraading")
        assert ii.hierarchy == (headings["m3"], "II")
        prop = next(s for s in report.basis.statements if s.role == "stikkordsfullmakt_forslag")
        assert prop.hierarchy == (
            "Forslag til vedtak om endringer i statsbudsjettet 2024 "
            "under Klima- og miljødepartementet",
            "Andre fullmakter",
            "V",
            "Stikkordsfullmakt",
        )
        auth = [s for s in report.basis.statements if s.role == "fullmaktsforslag"]
        assert [s.hierarchy[-2:] for s in auth] == [
            ("II", "Bestillingsfullmakt"),
            ("III", "Tilsagnsfullmakter"),
            ("IV", "Fullmakt til å inngå forpliktelser"),
            ("V", "Stikkordsfullmakt"),
            ("VI", "Fullmakt til forskuttering i oppdragsprosjekter"),
        ]
        book = load_workbook(excel)
        post60 = next(
            s
            for s in report.basis.statements
            if s.hierarchy
            == (
                headings["m1"],
                headings["m1_1"],
                headings["m1_1_3"],
            )
        )
        rows = [
            r for r in book["Fulltekst"].iter_rows(min_row=2, values_only=True) if r[0] == post60.id
        ]
        assert rows and json.loads(rows[0][8]) == list(post60.hierarchy)
        book.close()


@pytest.mark.parametrize(
    "post,meaning,field,value",
    [
        ("1", "endring", "amount_kr", 999_999_999),
        ("1", "for_belop", "amount_kr", 0),
        ("1", "etter_belop", "amount_kr", 0),
        ("22", "for_belop", "amount_kr", 0),
        ("22", "endring", "amount_kr", 3_000_000),
        ("22", "etter_belop", "amount_kr", 0),
        ("1", "endring", "budget_year", 2025),
        ("22", "endring", "meaning", "niva"),
        ("1", "endring", "chapter", "1420"),
        ("22", "endring", "step_id", "syntetisk-hendelse"),
        ("1", "endring", "interpretation_id", "syntetisk-tolkning"),
        ("22", "endring", "source_revision_id", "syntetisk-revisjon"),
        ("1", "endring", "pointer", "table[1]/tr[999]/td[4]"),
        ("22", "endring", "currency", "EUR"),
        ("1", "endring", "round", "rnb"),
    ],
)
@pytest.mark.parametrize("coordinated", [False, True])
def test_transferred_adopted_fields_are_replayed_before_t1_acceptance(
    tmp_path,
    pilot_basis,
    post,
    meaning,
    field,
    value,
    coordinated,
):
    from analyse.omtale import report_from_json as read_old
    from analyse.omtale_begge import build_both_report
    from analyse.saksdokumenter import report_from_json as read_base
    from analyse.visningspilot import source_input
    from innlesning.omtale import identified
    from innlesning.visningspilot import structure_basis

    old = read_old(pilot_basis.o2b_json)
    base = read_base(old.base_json)
    target = next(
        r.observation
        for r in base.base.unfiltered_rows
        if (r.observation.post, r.observation.meaning) == (post, meaning)
    )
    if field == "pointer":
        value = replace(target.pointer, locator=value)
    changed = replace(target, **{field: value})
    rows = tuple(
        replace(r, observation=changed) if r.observation.id == target.id else r
        for r in base.base.unfiltered_rows
    )
    source_text = base.source_text
    if coordinated:
        source_text = replace(
            source_text,
            extractions=tuple(
                replace(
                    e,
                    observations=tuple(changed if o.id == target.id else o for o in e.observations),
                )
                for e in source_text.extractions
            ),
        )
    bad_base = identified(
        "b6-frozen-2",
        replace(base, base=replace(base.base, unfiltered_rows=rows), source_text=source_text),
    )
    bad_old = build_both_report(bad_base, old.basis)
    frozen_old = report_to_json(bad_old)
    assert read_old(frozen_old) == bad_old  # Den gamle, videre kontrakten består.
    # Gamle konvolutt-ID-er er beregnet på nytt. Avvisning må komme i T1s nye port.
    with pytest.raises(ValueError, match="vedtaks"):
        prepared = source_input(frozen_old)
        structure_basis(prepared, pilot_basis.sources)


@pytest.fixture(scope="module")
def saved_pilot_pair(tmp_path_factory, pilot_basis):
    from analyse.visningspilot import build_report as build
    from analyse.visningspilot_excel import export_pair

    root = tmp_path_factory.mktemp("saved-presentation")
    report = build(pilot_basis, PilotSelection(years=(2024,), chapters=("1429",)))
    export_pair(report, root / "genuine.json", root / "genuine.xlsx")
    return root


@pytest.mark.parametrize(
    "mutation",
    [
        "constant_format",
        "empty_format",
        "hidden_sheet",
        "very_hidden_sheet",
        "hidden_row",
        "hidden_column",
        "tiny_row",
        "tiny_column",
        "white_font",
        "tiny_font",
        "matching_fill",
        "rotated_text",
        "row_default_height",
        "sheet_gridlines",
        "filter_range",
        "merged_cells",
        "conditional_format",
    ],
)
def test_saved_excel_presentation_cannot_hide_or_relabel_source_values(
    tmp_path, saved_pilot_pair, mutation
):
    from openpyxl import load_workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import PatternFill

    from analyse.visningspilot_excel import read_excel

    frozen = saved_pilot_pair / "genuine.json"
    before = hashlib.sha256(frozen.read_bytes()).hexdigest()
    book = load_workbook(saved_pilot_pair / "genuine.xlsx")
    sheet = book["Forløp"]
    cell = sheet["G3"]
    if mutation == "constant_format":
        cell.number_format = '"999 999 999"'
    elif mutation == "empty_format":
        cell.number_format = ";;;"
    elif mutation in {"hidden_sheet", "very_hidden_sheet"}:
        book["Dekning"].sheet_state = "hidden" if mutation == "hidden_sheet" else "veryHidden"
    elif mutation == "hidden_row":
        sheet.row_dimensions[3].hidden = True
    elif mutation == "hidden_column":
        sheet.column_dimensions["G"].hidden = True
    elif mutation == "tiny_row":
        sheet.row_dimensions[3].height = 0.1
    elif mutation == "tiny_column":
        sheet.column_dimensions["G"].width = 0.1
    elif mutation == "white_font":
        replace_font = copy(cell.font)
        replace_font.color = "FFFFFF"
        cell.font = replace_font
    elif mutation == "tiny_font":
        font = copy(cell.font)
        font.sz = 1
        cell.font = font
    elif mutation == "matching_fill":
        cell.fill = PatternFill("solid", fgColor="000000")
    elif mutation == "rotated_text":
        alignment = copy(cell.alignment)
        alignment.textRotation = 180
        cell.alignment = alignment
    elif mutation == "row_default_height":
        sheet.sheet_format.defaultRowHeight = 0.1
    elif mutation == "sheet_gridlines":
        sheet.sheet_view.showGridLines = True
    elif mutation == "filter_range":
        sheet.auto_filter.ref = "A1:B2"
    elif mutation == "merged_cells":
        sheet.merge_cells("G3:H3")
    else:
        sheet.conditional_formatting.add(
            "G3", FormulaRule(formula=["TRUE"], fill=PatternFill("solid", fgColor="000000"))
        )
    damaged = tmp_path / (mutation + ".xlsx")
    book.save(damaged)
    book.close()
    with pytest.raises(ValueError):
        read_excel(damaged)
    assert hashlib.sha256(frozen.read_bytes()).hexdigest() == before


def test_rehashed_adopted_amount_in_saved_json_and_excel_is_rejected(tmp_path, saved_pilot_pair):
    from openpyxl import load_workbook

    from analyse.omtale import report_from_json as read_old
    from analyse.omtale_begge import build_both_report
    from analyse.saksdokumenter import report_from_json as read_base
    from analyse.visningspilot import read_json
    from analyse.visningspilot_excel import read_excel
    from innlesning.observations import encode, identity
    from innlesning.omtale import identified

    data = json.loads((saved_pilot_pair / "genuine.json").read_text("utf8"))
    basis = data["report"]["basis"]
    old = read_old(basis["o2b_json"])
    base = read_base(old.base_json)
    rows = tuple(
        replace(r, observation=replace(r.observation, amount_kr=999_999_999))
        if (r.observation.post, r.observation.meaning) == ("1", "endring")
        else r
        for r in base.base.unfiltered_rows
    )
    bad_base = identified(
        "b6-frozen-2", replace(base, base=replace(base.base, unfiltered_rows=rows))
    )
    basis["o2b_json"] = report_to_json(build_both_report(bad_base, old.basis))
    for observation in basis["observations"]:
        if (observation["figure_type"], observation["post"], observation["meaning"]) == (
            "stortingsvedtak",
            "1",
            "endring",
        ):
            observation["amount_kr"] = 999_999_999
    basis["id"] = identity("visningsgrunnlag-1", {**basis, "id": ""})
    data["report"]["id"] = identity("visningspilot-1", {**data["report"], "id": ""})
    raw = encode(data)
    bad_json = tmp_path / "bad-adopted.json"
    bad_json.write_text(raw, "utf8")
    with pytest.raises(ValueError, match="vedtaksfelt"):
        read_json(bad_json)
    book = load_workbook(saved_pilot_pair / "genuine.xlsx")
    sheet = book["Visningsgrunnlag"]
    sheet.delete_rows(2, sheet.max_row)
    for start in range(0, len(raw), 30_000):
        sheet.append([start // 30_000 + 1, raw[start : start + 30_000]])
    damaged = tmp_path / "bad-adopted.xlsx"
    book.save(damaged)
    book.close()
    with pytest.raises(ValueError, match="vedtaksfelt"):
        read_excel(damaged)


def test_source_adapter_rejects_rebound_old_interpretation_and_extraction(pilot_basis):
    from analyse.visningspilot import source_input
    from innlesning.visningspilot import structure_basis

    prepared = source_input(pilot_basis.o2b_json)
    for field, value in (("normalization_rule", "syntetisk-enhet"), ("round", "rnb")):
        interpretation = replace(
            prepared.extraction.interpretation,
            **{field: value},
            round_evidence="syntetisk" if field == "round" else None,
        )
        extraction = replace(prepared.extraction, interpretation=interpretation)
        with pytest.raises(ValueError, match="vedtaksuttrekk/tolkning"):
            structure_basis(replace(prepared, extraction=extraction), pilot_basis.sources)


def test_account_selection_and_both_navigation_directions_survive_json(tmp_path, pilot_basis):
    from analyse.visningspilot import (
        build_report as build,
    )
    from analyse.visningspilot import (
        export_json,
        go_back,
        open_accounts,
        open_document,
        read_json,
    )

    docs = {d.key: d for d in pilot_basis.documents}
    for pairs, counts in (
        ((("1429", "1"),), (9, 3, 1)),
        ((("1429", "22"),), (9, 3, 1)),
        ((("1429", "1"), ("1429", "22")), (18, 6, 2)),
        ((), (0, 0, 0)),
        ((("1429", "60"),), (0, 0, 0)),
    ):
        selection = PilotSelection(years=(2024,), accounts=pairs)
        report = build(pilot_basis, selection)
        assert (
            len(report.view.observation_ids),
            len(report.view.adopted_observation_ids),
            len(report.view.adopted_event_ids),
        ) == counts
        document = open_document(report, docs["inns81"].id)
        assert len(document.view.inventory_ids) == 3
        assert len(document.view.observation_ids) == 6
        assert go_back(document) == report
        path = tmp_path / (
            "selection-" + str(len(pairs)) + "-" + (pairs[0][1] if pairs else "empty") + ".json"
        )
        export_json(document, path)
        saved = read_json(path)
        assert go_back(saved) == report
        assert saved.basis.id == report.basis.id
    document = build(pilot_basis, PilotSelection(mode="document", document_id=docs["vedtak211"].id))
    account = open_accounts(document, PilotSelection(years=(2024,), accounts=(("1429", "60"),)))
    assert docs["vedtak211"].id in account.view.document_ids
    assert go_back(account) == document
    assert all(
        c.difference_kr == 0
        for c in build(
            pilot_basis, PilotSelection(years=(2024,), chapters=("1429",))
        ).view.comparisons
    )


def test_actual_saved_excel_has_cells_links_originals_and_fulltext(tmp_path, pilot_basis):
    from openpyxl import load_workbook

    from analyse.visningspilot import build_report as build
    from analyse.visningspilot_excel import export_pair, read_excel

    report = build(pilot_basis, PilotSelection(years=(2024,), chapters=("1429",)))
    frozen, excel = tmp_path / "both.json", tmp_path / "both.xlsx"
    export_pair(report, frozen, excel)
    assert read_excel(excel) == report
    book = load_workbook(excel)
    try:
        sheet = book["Forløp"]
        assert sheet.max_row == 19
        assert sheet["G2"].value == 160_243_000
        assert sheet["G3"].value == 5_796_000
        assert sheet["B2"].hyperlink.location.startswith("'Dokumenter'!")
        assert book["Fulltekst"].max_row > 25
        assert book["Originaler"].max_row > 19
        assert all(c.data_type != "f" for s in book for row in s for c in row)
    finally:
        book.close()


def test_complete_original_text_fullmakts_chain_branches_and_unchanged_old_ids(pilot_basis):
    from analyse.omtale import report_from_json as old_reader
    from analyse.saksdokumenter import report_from_json as base_reader
    from analyse.visningspilot import build_report as build
    from innlesning.omtale import compact

    docs = {d.key: d for d in pilot_basis.documents}
    texts = [s for s in pilot_basis.statements if s.document_id == docs["inns81"].id]
    full = next(s for s in texts if s.locator == "pdf:page=2")
    assert "Oslo" in full.text and "Grunde Almeland" in full.text
    assert "9. desember 2024" in compact(full.text)
    p60 = next(
        s for s in texts if s.accounts == ("1429/60",) and s.role == "gjengivelse_av_regjeringen"
    )
    assert "totalt 10 mill." in compact(p60.text)
    assert "2025" in p60.period and "historisk RNB" in p60.period
    assert not any(o.post == "60" for o in pilot_basis.observations)
    fullmakter = [r for r in pilot_basis.relations if r.role == "dokumentert_stikkordsfullmakt"]
    assert len(fullmakter) == 2
    assert fullmakter[0].to_id == fullmakter[1].from_id
    branch = next(r for r in pilot_basis.relations if r.scope == "unntatt_1429")
    assert branch.to_id == docs["inns95"].id
    assert all(not a.startswith("1429/") for a in branch.accounts)
    selected = build(pilot_basis, PilotSelection(years=(2024,), accounts=(("1429", "1"),)))
    assert docs["inns95"].id not in selected.view.document_ids
    other = build(pilot_basis, PilotSelection(years=(2024,), accounts=(("1420", "1"),)))
    assert docs["inns95"].id in other.view.document_ids
    assert docs["inns81"].id not in other.view.document_ids
    refs81 = [s for s in texts if s.role == "saksreferanse_kapittel_1429"]
    assert len(refs81) == 5
    assert refs81[3].text != refs81[4].text
    old = old_reader(pilot_basis.o2b_json)
    base = base_reader(old.base_json)
    assert tuple(
        o for o in pilot_basis.observations if o.figure_type == "stortingsvedtak"
    ) == tuple(r.observation for r in base.base.unfiltered_rows)
    assert {o.round for o in pilot_basis.observations if o.figure_type == "stortingsvedtak"} == {
        "uavklart"
    }
    assert len([s for s in old.basis.sections if s.key == "P1"]) == 1
    assert "210_byteeksakt_rekonstruert_fra_9cfeaef5" in pilot_basis.issues
    assert "211_byteeksakt_rekonstruert_fra_c5f3b900" in pilot_basis.issues
    assert all(o.amount_kr != 2_796_000 for o in pilot_basis.observations)


@pytest.mark.parametrize(
    "field,value",
    [
        ("budget_year", 2025),
        ("chapter", "1420"),
        ("post", "2"),
        ("currency", "EUR"),
        ("multiplier", "1000000"),
        ("meaning", "nivaa"),
        ("detail_level", "underpost"),
        ("scope", "aggregat"),
        ("figure_type", "regnskap"),
        ("proposal_role", "alternativ"),
        ("round", "rnb"),
    ],
)
def test_incompatible_source_fields_block_stage_difference(pilot_basis, field, value):
    from analyse.visningspilot import compare_observations

    left = next(
        o
        for o in pilot_basis.observations
        if o.figure_type == "regjeringsforslag" and o.post == "1" and o.meaning == "endring"
    )
    right = next(
        o
        for o in pilot_basis.observations
        if o.figure_type == "komiteforslag" and o.post == "1" and o.meaning == "endring"
    )
    contexts = {c.observation_id: c for c in pilot_basis.contexts}
    mutated = replace(right, **{field: value})
    rc = replace(contexts[right.id], observation_id=mutated.id)
    result = compare_observations(
        left,
        mutated,
        contexts[left.id],
        rc,
        account="1429/1",
        meaning="endring",
        relation_bound=True,
    )
    assert result.status == "sperret" and result.difference_kr is None
    assert result.null_reason


def test_missing_stage_observed_zero_and_context_guard_are_distinct(tmp_path, pilot_basis):
    from analyse.visningspilot import compare_observations
    from innlesning.observations import encode

    left = next(
        o
        for o in pilot_basis.observations
        if o.figure_type == "regjeringsforslag" and o.post == "1" and o.meaning == "endring"
    )
    right = next(
        o
        for o in pilot_basis.observations
        if o.figure_type == "komiteforslag" and o.post == "1" and o.meaning == "endring"
    )
    contexts = {c.observation_id: c for c in pilot_basis.contexts}
    missing = compare_observations(
        left,
        None,
        contexts[left.id],
        None,
        account="1429/1",
        meaning="endring",
        relation_bound=False,
    )
    assert missing.status == "hull" and missing.difference_kr is None
    assert missing.right_id is None and missing.null_reason == "manglende stadium"
    equal = compare_observations(
        left,
        right,
        contexts[left.id],
        contexts[right.id],
        account="1429/1",
        meaning="endring",
        relation_bound=True,
    )
    assert equal.difference_kr == 0 and equal.null_reason is None
    for lc, rc, bound in (
        (None, contexts[right.id], True),
        (contexts[left.id], replace(contexts[right.id], actor="mindretall"), True),
        (contexts[left.id], contexts[right.id], False),
    ):
        blocked = compare_observations(
            left, right, lc, rc, account="1429/1", meaning="endring", relation_bound=bound
        )
        assert blocked.status == "sperret" and blocked.difference_kr is None
    # Reelle skrevne/gjenleste kontrollverdier: JSON-null kan ikke bli et observert nullbeløp.
    path = tmp_path / "syntetiske-sammenligningskontroller.json"
    path.write_text(encode([asdict(missing), asdict(equal)]), "utf-8")
    read = json.loads(path.read_text("utf-8"))
    assert [r["difference_kr"] for r in read] == [None, 0]


def test_missing_new_numeric_adapter_stage_is_a_gap_in_saved_files(tmp_path, pilot_basis):
    from analyse.visningspilot import build_report as build
    from analyse.visningspilot import read_json
    from analyse.visningspilot_excel import export_pair, read_excel
    from innlesning.visningspilot import SOURCE_CATALOG

    root = tmp_path / "synthetic-missing-stage"
    for name in SOURCE_CATALOG:
        target = root / name
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(name, target)
    missing = root / "docs/research/omtale-o1/originals/inns81s.html"
    missing.unlink(missing_ok=True)
    report = build(
        import_basis(pilot_basis.o2b_json, root),
        PilotSelection(years=(2024,), accounts=(("1429", "1"),)),
    )
    frozen, excel = tmp_path / "missing-stage.json", tmp_path / "missing-stage.xlsx"
    export_pair(report, frozen, excel)
    actual = read_json(frozen)
    assert read_excel(excel) == actual
    assert len(actual.view.observation_ids) == 6
    assert len(actual.view.adopted_event_ids) == 1
    assert len(actual.view.comparisons) == 6
    assert all(c.status == "hull" and c.difference_kr is None for c in actual.view.comparisons)
    assert any("inns81s.html" in issue for issue in actual.basis.issues)
    # Det tidligere O2b-grunnlaget og gamle vedtaksfelt består selv om den nye adapteren mangler.
    assert actual.basis.o2b_json == pilot_basis.o2b_json


@pytest.mark.parametrize(
    "mutation", ["role", "text", "locator", "relation", "observation", "context", "pin", "source"]
)
def test_rehashed_json_cannot_change_source_binding(tmp_path, pilot_basis, mutation):
    from analyse.visningspilot import build_report as build
    from analyse.visningspilot import report_from_json, report_to_json
    from innlesning.observations import encode

    report = build(pilot_basis, PilotSelection(years=(2024,), chapters=("1429",)))
    data = json.loads(report_to_json(report))
    basis = data["report"]["basis"]
    if mutation in {"role", "text", "locator"}:
        statement = basis["statements"][0]
        statement[{"role": "role", "text": "text", "locator": "locator"}[mutation]] = (
            "syntetisk endret"
        )
    elif mutation == "relation":
        basis["relations"][0]["scope"] = "unntatt_1429"
    elif mutation == "observation":
        basis["observations"][0]["amount_kr"] = 0
    elif mutation == "context":
        basis["contexts"][0]["round"] = "rnb"
    elif mutation == "pin":
        basis["sources"][0]["expected_sha256"] = "0" * 64
    else:
        source = basis["sources"][0]
        raw = b"syntetisk manipulasjon med alle egne hasher beregnet pa nytt"
        source.update(
            content_base64=base64.b64encode(raw).decode(),
            sha256=hashlib.sha256(raw).hexdigest(),
            byte_count=len(raw),
            expected_sha256=hashlib.sha256(raw).hexdigest(),
            status="kildeverifisert",
        )
    # Ikke bare topphash: alle nye identiteter med eksplisitt id-felt beregnes på nytt.
    from innlesning.observations import identity

    for name, namespace in (
        ("sources", "visningskilde-1"),
        ("documents", "visningsdokument-1"),
        ("statements", "visningsutsagn-1"),
        ("relations", "visningsrelasjon-1"),
        ("contexts", "visningskontekst-1"),
    ):
        for obj in basis[name]:
            obj["id"] = identity(namespace, {**obj, "id": ""})
    basis["id"] = identity("visningsgrunnlag-1", {**basis, "id": ""})
    data["report"]["id"] = identity("visningspilot-1", {**data["report"], "id": ""})
    path = tmp_path / "syntetisk-manipulert.json"
    path.write_text(encode(data), "utf-8")
    with pytest.raises(ValueError):
        report_from_json(path.read_text("utf-8"))


@pytest.mark.parametrize(
    "sheet,cell,value",
    [
        ("Forløp", "G2", 0),
        ("Fulltekst", "K2", "syntetisk tekst"),
        ("Originaler", "K2", "syntetiske byte"),
        ("Relasjoner", "E2", "unntatt_1429"),
        ("Beløpsgrunnlag", "H2", "stortingsvedtak"),
        ("Visningsgrunnlag", "A2", 99),
    ],
)
def test_actual_excel_changes_are_rejected(tmp_path, pilot_basis, sheet, cell, value):
    from openpyxl import load_workbook

    from analyse.visningspilot import build_report as build
    from analyse.visningspilot_excel import export_excel, read_excel

    path = tmp_path / "valid.xlsx"
    export_excel(build(pilot_basis, PilotSelection(years=(2024,), chapters=("1429",))), path)
    book = load_workbook(path)
    book[sheet][cell] = value
    damaged = tmp_path / "syntetisk-manipulert.xlsx"
    book.save(damaged)
    book.close()
    with pytest.raises(ValueError):
        read_excel(damaged)


def test_long_formula_like_source_text_is_safe_and_lossless_in_real_excel(tmp_path, pilot_basis):
    from openpyxl import load_workbook

    from analyse.omtale import report_from_json as read_old
    from analyse.omtale import report_to_json as write_old
    from analyse.omtale_begge import propose_relation
    from analyse.visningspilot import build_report as build
    from analyse.visningspilot_excel import export_pair, read_excel

    old = read_old(pilot_basis.o2b_json)
    target = old.relations[0]
    proposal = propose_relation(
        old,
        section_id=old.basis.sections[1].id,
        event_id=target.event_id,
        evidence_ids=(old.basis.sections[1].id,),
        reason="=SYNTHETIC('" + "Lang syntetisk kontrolltekst øæå " * 2200 + "')",
    )
    report = build(
        import_basis(write_old(proposal), Path(".")), PilotSelection(years=(2024,), accounts=())
    )
    frozen, excel = tmp_path / "synthetic-long.json", tmp_path / "synthetic-long.xlsx"
    export_pair(report, frozen, excel)
    assert read_excel(excel) == report
    book = load_workbook(excel)
    try:
        assert "Relasjoner fortsettelse" in book.sheetnames
        assert all(c.data_type != "f" for s in book for row in s for c in row)
    finally:
        book.close()


@pytest.mark.parametrize("condition", ["missing", "failed", "corrupt", "contradictory"])
def test_new_authentic_decision_is_independent_of_four_bad_pilot_sources(
    tmp_path, pilot_basis, condition
):
    from openpyxl import load_workbook

    from analyse.visningspilot import build_report as build
    from analyse.visningspilot_excel import export_pair, read_excel
    from innlesning import read_source
    from innlesning.observations import read_extraction
    from innlesning.visningspilot import SOURCE_CATALOG

    root = tmp_path / "synthetic-sources"
    for name in SOURCE_CATALOG:
        destination = root / name
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(name, destination)
    source = root / "tests/fixtures/b8/40034049.xml"
    if condition == "missing":
        source.unlink()
    elif condition == "failed":
        source.unlink()
        source.mkdir()  # faktisk feilet filåpning, ikke bare oppdiktet feiltekst
    elif condition == "corrupt":
        source.write_bytes(b"synthetic invalid XML")
    else:
        source.write_text(
            Path("tests/fixtures/b8/40034049.xml").read_text("utf8").replace("1429", "1420"), "utf8"
        )
    basis = import_basis(pilot_basis.o2b_json, root)
    report = build(basis, PilotSelection(years=(2024,), accounts=()))
    frozen, excel = tmp_path / "empty.json", tmp_path / "empty.xlsx"
    export_pair(report, frozen, excel)
    assert read_excel(excel) == report and not report.view.observation_ids
    assert any("40034049.xml" in issue for issue in basis.issues)
    assert any("stikkordsfullmaktkjede_innlegg_mangler" in issue for issue in basis.issues)

    database = tmp_path / "ordinary.sqlite"

    def cli(*args):
        result = subprocess.run(
            [sys.executable, "-m", *args], capture_output=True, text=True, encoding="utf8"
        )
        assert result.returncode in (0, 2), result.stdout + result.stderr
        return json.loads(result.stdout)

    summary = cli(
        "innlesning",
        "synkroniser",
        "--database",
        str(database),
        "--sesjon",
        "2024-2025",
        "--historisk-proeve",
        "--vedtak-id",
        "40034229",
        "--response",
        "tests/fixtures/b6/40034229.xml",
        "--metadata",
        "tests/fixtures/b6/40034229.json",
    )
    assert summary["antall"]["ny"] == 1 and not summary["feil"]
    revision = summary["vedtak"][0]["revision_id"]
    interpreted = cli(
        "innlesning",
        "tolk",
        "--database",
        str(database),
        "--revisjon",
        revision,
        "--tolkning",
        "storting-belop-v2",
    )
    extraction = read_extraction(database, interpreted["belopsuttrekk"])
    assert len({o.step_id for o in extraction.observations}) == 74
    assert (
        extraction.artifact.sha256
        == "66208daf51111356ad73c26985819870ab9c2f97ec7b070024d89e150cd772b9"
    )
    ordinary = tmp_path / "ordinary.xlsx"
    cli(
        "analyse",
        "--database",
        str(database),
        "--belopsuttrekk",
        extraction.id,
        "--excel",
        str(ordinary),
    )
    book = load_workbook(ordinary)
    try:
        rows = list(book["Beløpsobservasjoner"].iter_rows(min_row=2, values_only=True))
        assert {r[0] for r in rows} == {o.id for o in extraction.observations}
        assert {
            (o.meaning, o.amount_kr)
            for o in extraction.observations
            if o.chapter == "700" and o.post == "1"
        } == {("for_belop", 304_351_000), ("endring", 6_324_000), ("etter_belop", 310_675_000)}
        headers = {c.value: n for n, c in enumerate(book["Beløpsobservasjoner"][1])}
        selected = [
            r for r in rows if r[headers["Kapittel"]] == "700" and r[headers["Post"]] == "1"
        ]
        assert {r[headers["Kroner"]] for r in selected} == {304_351_000, 6_324_000, 310_675_000}
        assert (
            read_source(database, extraction.artifact.sha256)
            == Path("tests/fixtures/b6/40034229.xml").read_bytes()
        )
    finally:
        book.close()


def test_union_has_no_cross_product_years_stay_separate_and_offline_repeat_is_stable(
    tmp_path, pilot_basis, monkeypatch
):
    import socket
    import sqlite3

    from analyse.visningspilot import build_report as build
    from analyse.visningspilot import read_json
    from analyse.visningspilot_excel import export_originals, export_pair, read_excel

    def forbidden(*args, **kwargs):
        raise AssertionError("Offline-gjenlesing skal ikke åpne database eller nett")

    monkeypatch.setattr(sqlite3, "connect", forbidden)
    monkeypatch.setattr(socket, "socket", forbidden)
    selection = PilotSelection(years=(2024,), accounts=(("1429", "1"), ("1420", "22")))
    report = build(pilot_basis, selection)
    assert len(report.view.observation_ids) == 9
    selected_items = [i for i in pilot_basis.inventory if i.id in report.view.inventory_ids]
    assert {(i.chapter, i.post) for i in selected_items} <= {("1429", "1"), ("1420", "22")}
    future = build(pilot_basis, PilotSelection(years=(2025,), accounts=(("1429", "60"),)))
    assert not future.view.observation_ids
    assert future.view.statement_ids
    assert not any(i.id in future.view.inventory_ids for i in pilot_basis.inventory)
    frozen, excel = tmp_path / "offline.json", tmp_path / "offline.xlsx"
    export_pair(report, frozen, excel)
    hashes = [hashlib.sha256(p.read_bytes()).hexdigest() for p in (frozen, excel)]
    assert read_json(frozen) == read_excel(excel) == report
    export_pair(read_excel(excel), tmp_path / "repeat.json", tmp_path / "repeat.xlsx")
    assert (tmp_path / "repeat.json").read_bytes() == frozen.read_bytes()
    with pytest.raises(ValueError):
        export_pair(report, frozen, excel)
    assert hashes == [hashlib.sha256(p.read_bytes()).hexdigest() for p in (frozen, excel)]
    originals = tmp_path / "originals"
    export_originals(read_json(frozen), originals)
    assert (originals / "o2b/originals/inns81s.pdf").read_bytes() == Path(
        "docs/research/omtale-o1/originals/inns81s.pdf"
    ).read_bytes()


@pytest.mark.parametrize("mutation", ["extra_sheet", "link", "formula", "fulltext_part"])
def test_excel_structure_and_navigation_binding_are_validated_before_acceptance(
    tmp_path, pilot_basis, mutation
):
    from openpyxl import load_workbook

    from analyse.visningspilot import build_report as build
    from analyse.visningspilot_excel import export_excel, read_excel

    path = tmp_path / "valid.xlsx"
    export_excel(build(pilot_basis, PilotSelection(years=(2024,), chapters=("1429",))), path)
    book = load_workbook(path)
    if mutation == "extra_sheet":
        book.create_sheet("Skjult falskt grunnlag")
    elif mutation == "link":
        book["Forløp"]["B2"].hyperlink.location = "'Dokumenter'!B3"
    elif mutation == "formula":
        book["Fulltekst"]["K2"] = "=1+1"
    else:
        book["Fulltekst"]["J2"] = 2
    damaged = tmp_path / "synthetic-structure.xlsx"
    book.save(damaged)
    book.close()
    with pytest.raises(ValueError):
        read_excel(damaged)
